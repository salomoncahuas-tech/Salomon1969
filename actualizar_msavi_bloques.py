# -*- coding: utf-8 -*-
"""
AUTORIDAD NACIONAL DE INFRAESTRUCTURA - ANIN
DIRECCION DE INTERVENCIONES MULTISECTORIALES Y DE EMERGENCIA - DIME
SUBDIRECCION DE ESTUDIOS DE INVERSION - SESDI

PROYECTO IN PIURA | CUI 2669244
Recuperacion del servicio de regulacion de riesgos naturales y recuperacion de
ecosistemas degradados en la Cuenca Alta del Rio Piura.

TRASPASO AUTOMATIZADO DE AREAS MSAVI POR CLASE (DN) A LAS PLANTILLAS DE BLOQUE
-----------------------------------------------------------------------------
Traslada la distribucion areal por clase espectral (campo DN del shapefile
MSAVI V6) desde el consolidado AREAS_MSAVI_BLOQUES_V6_REV_HSCM.xlsx hacia las
117 plantillas Excel de bloque, y propaga la actualizacion a las hojas que
dependen de ese dato:

  1. Hoja "Cobertura MSAVI-NDVI", seccion A  -> distribucion areal por clase,
     total clasificado y superficie bajo el umbral 0.4976 (indicador de brecha).
  2. Hoja "Resumen", seccion 3               -> superficie clasificada MSAVI,
     desviacion frente al catalogo y superficie bajo umbral.
  3. Hoja "Control de consistencia"          -> registro de las verificaciones
     D-nn correspondientes y recalculo del cuadro RESUMEN.

Correspondencia DN -> clase de umbral del proyecto (verificada empiricamente
reconstruyendo la media del MSAVI a partir de los puntos medios de clase
ponderados por superficie: error absoluto medio 0.0131 frente a 0.2814 con el
orden inverso, sobre los 117 bloques):

     DN 5 = > 0.6139           (Vigor alto)
     DN 4 = 0.4976 - 0.6139    (Vigor moderado)
     DN 3 = 0.3813 - 0.4976    (Vigor bajo)
     DN 2 = 0.2650 - 0.3813    (Vigor muy bajo)
     DN 1 = <= 0.2650          (Suelo desnudo / no vegetal)

Sistema de coordenadas: UTM WGS 84 Zona 17S (EPSG:32717).

Uso:
    python actualizar_msavi_bloques.py --fuente AREAS_MSAVI_BLOQUES_V6_REV_HSCM.xlsx \
                                       --plantillas "Plantillas Excel 117 bloques" \
                                       --salida "Plantillas Excel 117 bloques - MSAVI V6" \
                                       --reporte Reporte_Traspaso_MSAVI_V6.xlsx

Elaborado por: Ing. Hector Salomon Cahuas Miller - SESDI/DIME/ANIN
"""

from __future__ import annotations

import argparse
import copy
import os
import re
import shutil
import sys
from collections import Counter, defaultdict

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# Identidad institucional ANIN
# --------------------------------------------------------------------------- #
VERDE_ANIN = "1B4D2E"
VERDE_CLARO = "E8F0EA"
GRIS_ETIQUETA = "F2F2F2"
FUENTE = "Arial"

_LADO = Side(style="thin", color="B7C9BC")
BORDE_FINO = Border(left=_LADO, right=_LADO, top=_LADO, bottom=_LADO)

ENCABEZADOS_ANIN = (
    "AUTORIDAD NACIONAL DE INFRAESTRUCTURA - ANIN",
    "DIRECCIÓN DE INTERVENCIONES MULTISECTORIALES Y DE EMERGENCIA - DIME",
    "SUBDIRECCIÓN DE ESTUDIOS DE INVERSIÓN - SESDI",
)

# --------------------------------------------------------------------------- #
# Parametros del modelo MSAVI
# --------------------------------------------------------------------------- #
UMBRAL_BRECHA = 0.4976
TOLERANCIA_AREA_PCT = 2.0  # tolerancia planimetrica aplicada en las plantillas

# Fila de la seccion A de "Cobertura MSAVI-NDVI" -> clase DN de origen.
# La seccion lista las clases de mayor a menor vigor, el campo DN al reves.
FILAS_CLASE = [
    (10, 5, "> 0.6139"),
    (11, 4, "0.4976 - 0.6139"),
    (12, 3, "0.3813 - 0.4976"),
    (13, 2, "0.2650 - 0.3813"),
    (14, 1, "<= 0.2650"),
]
FILA_TOTAL_MSAVI = 15        # fila nueva: TOTAL CLASIFICADO
FILA_BRECHA_MSAVI = 16       # fila nueva: SUPERFICIE BAJO UMBRAL
FILAS_INSERTADAS_COBERTURA = 2
FILA_MEDIA_ORIGINAL = 15     # fila de la media del bloque antes del desplazamiento

# Clases DN que quedan por debajo del umbral de brecha 0.4976
DN_BAJO_UMBRAL = (1, 2, 3)

HOJA_COBERTURA = "Cobertura MSAVI-NDVI"
HOJA_RESUMEN = "Resumen"
HOJA_CONTROL = "Control de consistencia"

RE_PLANTILLA = re.compile(r"^Plantilla_Excel_Bloque_(.+)_IN_Piura\.xlsx$")


# --------------------------------------------------------------------------- #
# Utilidades de hoja de calculo
# --------------------------------------------------------------------------- #
def insertar_filas(ws, fila, n, ancho):
    """Inserta n filas vacias en 'fila', desplazando hacia abajo valores,
    estilos, formulas (traducidas), rangos combinados y alturas de fila.

    openpyxl.insert_rows() no reubica rangos combinados, alturas ni traduce
    formulas, por lo que el desplazamiento se hace aqui de forma explicita.
    """
    ultima = ws.max_row

    # 1) desarmar los rangos combinados que seran desplazados
    a_recombinar = []
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= fila:
            a_recombinar.append((m.min_row, m.min_col, m.max_row, m.max_col))
            ws.unmerge_cells(str(m))

    # 2) desplazar celdas de abajo hacia arriba
    for r in range(ultima, fila - 1, -1):
        for c in range(1, ancho + 1):
            org = ws.cell(row=r, column=c)
            dst = ws.cell(row=r + n, column=c)
            valor = org.value
            if isinstance(valor, str) and valor.startswith("="):
                valor = Translator(valor, origin=org.coordinate).translate_formula(
                    "%s%d" % (get_column_letter(c), r + n)
                )
            dst.value = valor
            if org.has_style:
                dst._style = copy.copy(org._style)
            org.value = None

    # 3) dejar limpias las filas liberadas
    for r in range(fila, fila + n):
        for c in range(1, ancho + 1):
            celda = ws.cell(row=r, column=c)
            celda.value = None
            celda.font = Font(name=FUENTE, size=9)
            celda.fill = PatternFill(fill_type=None)
            celda.border = Border()
            celda.alignment = Alignment()
            celda.number_format = "General"

    # 4) rehacer los rangos combinados desplazados
    for r1, c1, r2, c2 in a_recombinar:
        ws.merge_cells(start_row=r1 + n, start_column=c1,
                       end_row=r2 + n, end_column=c2)

    # 5) desplazar alturas de fila
    alturas = {r: d.height for r, d in ws.row_dimensions.items()
               if d.height and r >= fila}
    for r in list(ws.row_dimensions.keys()):
        if r >= fila:
            ws.row_dimensions[r].height = None
    for r, h in alturas.items():
        ws.row_dimensions[r + n].height = h


def copiar_estilo_fila(ws, fila_origen, fila_destino, ancho):
    """Replica el formato de una fila modelo en otra (sin copiar valores)."""
    for c in range(1, ancho + 1):
        org = ws.cell(row=fila_origen, column=c)
        dst = ws.cell(row=fila_destino, column=c)
        if org.has_style:
            dst._style = copy.copy(org._style)


def escribir(ws, fila, columna, valor, formato=None):
    celda = ws.cell(row=fila, column=columna)
    celda.value = valor
    if formato:
        celda.number_format = formato
    return celda


# --------------------------------------------------------------------------- #
# Lectura de la fuente MSAVI
# --------------------------------------------------------------------------- #
def leer_fuente_msavi(ruta):
    """Devuelve {codigo_bloque: dict con areas por DN, total y n de poligonos}."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)

    if "Resumen_por_Bloque" not in wb.sheetnames:
        raise SystemExit("La fuente no contiene la hoja 'Resumen_por_Bloque'.")

    # --- areas agregadas por bloque y clase DN -----------------------------
    datos = {}
    ws = wb["Resumen_por_Bloque"]
    encabezado_visto = False
    for fila in ws.iter_rows(values_only=True):
        if not encabezado_visto:
            encabezado_visto = fila and fila[0] == "BLOQUE"
            continue
        if not fila or fila[0] is None:
            continue
        codigo = str(fila[0]).strip()
        if codigo.upper() == "TOTAL":
            continue
        m2 = {dn: float(fila[dn] or 0) for dn in range(1, 6)}
        total_m2 = float(fila[6] or 0)
        datos[codigo] = {
            "m2": m2,
            "total_m2": total_m2,
            "total_ha": total_m2 / 10000.0,
            "n_poligonos": 0,
        }

    # --- numero de poligonos por bloque (hoja de detalle) ------------------
    if "Datos" in wb.sheetnames:
        conteo = Counter()
        encabezado_visto = False
        for fila in wb["Datos"].iter_rows(values_only=True):
            if not encabezado_visto:
                encabezado_visto = fila and fila[0] == "FID"
                continue
            if not fila or fila[1] is None:
                continue
            conteo[str(fila[1]).strip()] += 1
        for codigo, n in conteo.items():
            if codigo in datos:
                datos[codigo]["n_poligonos"] = n

    wb.close()
    return datos


def metricas_bloque(reg, catalogo_ha):
    """Calcula superficies en hectareas, porcentajes y desviacion planimetrica."""
    total_ha = reg["total_ha"]
    ha = {dn: reg["m2"][dn] / 10000.0 for dn in range(1, 6)}
    ha_bajo = sum(ha[dn] for dn in DN_BAJO_UMBRAL)
    pct_bajo = (ha_bajo / total_ha * 100.0) if total_ha else 0.0
    if catalogo_ha:
        desviacion = (total_ha - catalogo_ha) / catalogo_ha * 100.0
    else:
        desviacion = None
    return {
        "ha": ha,
        "total_ha": total_ha,
        "ha_bajo": ha_bajo,
        "pct_bajo": pct_bajo,
        "catalogo_ha": catalogo_ha,
        "desviacion": desviacion,
        "n_poligonos": reg["n_poligonos"],
    }


# --------------------------------------------------------------------------- #
# 1. Hoja "Cobertura MSAVI-NDVI"
# --------------------------------------------------------------------------- #
def actualizar_cobertura(ws, met, media_msavi):
    """Escribe la distribucion areal por clase MSAVI en la seccion A."""
    # Filas nuevas para el total clasificado y la superficie bajo umbral.
    insertar_filas(ws, FILA_TOTAL_MSAVI, FILAS_INSERTADAS_COBERTURA, 5)

    fila_total = FILA_TOTAL_MSAVI
    ref_total = "$B$%d" % fila_total

    # --- superficies y porcentajes por clase -------------------------------
    for fila, dn, _clase in FILAS_CLASE:
        escribir(ws, fila, 2, round(met["ha"][dn], 4), "0.0000")
        escribir(ws, fila, 3,
                 "=IF(%s=0,0,ROUND(B%d/%s*100,2))" % (ref_total, fila, ref_total),
                 "0.00")

    # --- fila TOTAL CLASIFICADO (modelo de estilo: fila de clase par) ------
    copiar_estilo_fila(ws, 11, fila_total, 5)
    for c in range(1, 6):
        celda = ws.cell(row=fila_total, column=c)
        celda.font = Font(name=FUENTE, size=9, bold=True)
        celda.fill = PatternFill("solid", fgColor=VERDE_CLARO)
    escribir(ws, fila_total, 1, "TOTAL CLASIFICADO MSAVI 2024")
    escribir(ws, fila_total, 2, "=ROUND(SUM(B10:B14),4)", "0.0000")
    escribir(ws, fila_total, 3, "=ROUND(SUM(C10:C14),2)", "0.00")
    escribir(ws, fila_total, 4,
             "Superficie del bloque clasificada por MSAVI 2024 (campo DN)")
    if met["catalogo_ha"] and met["desviacion"] is not None:
        escribir(ws, fila_total, 5,
                 "Catálogo V5/V6: %s ha · desviación %+0.2f %%"
                 % (_num(met["catalogo_ha"]), met["desviacion"]))
    else:
        escribir(ws, fila_total, 5, "Superficie de catálogo no disponible")

    # --- fila SUPERFICIE BAJO UMBRAL (indicador de brecha) -----------------
    fila_brecha = FILA_BRECHA_MSAVI
    copiar_estilo_fila(ws, fila_total, fila_brecha, 5)
    for c in range(1, 6):
        ws.cell(row=fila_brecha, column=c).fill = PatternFill(
            "solid", fgColor=GRIS_ETIQUETA)
    escribir(ws, fila_brecha, 1, "SUPERFICIE BAJO UMBRAL MSAVI %0.4f" % UMBRAL_BRECHA)
    escribir(ws, fila_brecha, 2, "=ROUND(SUM(B12:B14),4)", "0.0000")
    escribir(ws, fila_brecha, 3, "=ROUND(SUM(C12:C14),2)", "0.00")
    escribir(ws, fila_brecha, 4,
             "Ecosistema degradado que requiere recuperación (clases DN 1 a DN 3)")
    escribir(ws, fila_brecha, 5, "Indicador de brecha · R.M. N.° 00213-2024-MINAM")

    # --- nota metodologica de la seccion A ---------------------------------
    fila_nota = 17 + FILAS_INSERTADAS_COBERTURA  # nota original en la fila 17
    escribir(ws, fila_nota, 1, _nota_cobertura(met, media_msavi))
    ws.row_dimensions[fila_nota].height = 108.0

    return fila_total, fila_brecha


def _num(x, dec=4):
    """Formatea un numero suprimiendo ceros finales innecesarios."""
    if x is None:
        return "no disponible"
    if isinstance(x, str):
        return x
    txt = ("%0." + str(dec) + "f") % x
    return txt.rstrip("0").rstrip(".") if "." in txt else txt


def _nota_cobertura(met, media_msavi):
    if met["desviacion"] is None:
        juicio = ("No se dispone de superficie de catálogo para contrastar la "
                  "superficie clasificada.")
    elif abs(met["desviacion"]) <= TOLERANCIA_AREA_PCT:
        juicio = ("La superficie clasificada (%s ha) concuerda con la de catálogo "
                  "V5/V6 (%s ha) dentro del ±%0.0f %% (%+0.2f %%), lo que valida "
                  "la geometría del bloque."
                  % (_num(met["total_ha"]), _num(met["catalogo_ha"]),
                     TOLERANCIA_AREA_PCT, met["desviacion"]))
    else:
        juicio = ("ATENCIÓN: la superficie clasificada (%s ha) difiere de la de "
                  "catálogo V5/V6 (%s ha) en %+0.2f %%, fuera de la tolerancia de "
                  "±%0.0f %%. La distribución por clase se reporta sobre la "
                  "superficie efectivamente clasificada; el contraste queda "
                  "registrado en la hoja «Control de consistencia»."
                  % (_num(met["total_ha"]), _num(met["catalogo_ha"]),
                     met["desviacion"], TOLERANCIA_AREA_PCT))

    return (
        "NOTA METODOLÓGICA. La media del MSAVI 2024 del bloque (%s) procede del "
        "catálogo maestro de bloques V5/V6 y es dato oficial. La DISTRIBUCIÓN "
        "AREAL por clase de MSAVI 2024 procede de la estadística zonal del "
        "ráster clasificado MSAVI V6 (shapefile «MSAVI V6.shp», campo DN, UTM "
        "WGS 84 Zona 17S), consolidada en «AREAS_MSAVI_BLOQUES_V6_REV_HSCM.xlsx», "
        "hoja «Resumen_por_Bloque»: %d polígonos del bloque agregados por clase. "
        "Correspondencia DN → clase de umbral: DN 5 = > 0.6139; DN 4 = 0.4976 - "
        "0.6139; DN 3 = 0.3813 - 0.4976; DN 2 = 0.2650 - 0.3813; DN 1 = <= 0.2650. "
        "La fila resaltada indica la clase en la que cae la media del bloque. La "
        "superficie situada bajo el umbral 0.4976 (clases DN 1 a DN 3) constituye "
        "la superficie de ecosistema degradado que requiere recuperación y es la "
        "base del indicador de brecha del proyecto (R.M. N.° 00213-2024-MINAM). %s"
        % (_num(media_msavi, 6), met["n_poligonos"], juicio)
    )


# --------------------------------------------------------------------------- #
# 2. Hoja "Resumen"
# --------------------------------------------------------------------------- #
def actualizar_resumen(ws, met):
    """Incorpora la sintesis MSAVI en la seccion 3 (indices de vegetacion)."""
    fila = 31  # inmediatamente despues de "Superficie clasificada NDVI 2025 (ha)"
    insertar_filas(ws, fila, 2, 6)

    for offset in (0, 1):
        copiar_estilo_fila(ws, 30, fila + offset, 6)

    escribir(ws, fila, 1, "Superficie clasificada MSAVI 2024 (ha)")
    escribir(ws, fila, 2, round(met["total_ha"], 4), "0.0000")
    escribir(ws, fila, 3, "Desviación MSAVI frente al catálogo (%)")
    if met["desviacion"] is None:
        escribir(ws, fila, 4, "No disponible")
    else:
        escribir(ws, fila, 4, round(met["desviacion"], 2), "0.00")

    escribir(ws, fila + 1, 1,
             "Superficie bajo umbral MSAVI %0.4f (ha)" % UMBRAL_BRECHA)
    escribir(ws, fila + 1, 2, round(met["ha_bajo"], 4), "0.0000")
    escribir(ws, fila + 1, 3, "Proporción del bloque bajo umbral (%)")
    escribir(ws, fila + 1, 4, round(met["pct_bajo"], 2), "0.00")

    return fila, fila + 1


# --------------------------------------------------------------------------- #
# 3. Hoja "Control de consistencia"
# --------------------------------------------------------------------------- #
def actualizar_control(ws, met):
    """Registra las verificaciones del traspaso y recalcula el cuadro RESUMEN."""
    fila_resumen = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "RESUMEN":
            fila_resumen = r
            break
    if fila_resumen is None:
        return None, None, None

    primera_d = 10
    ultima_d = fila_resumen - 2  # entre la ultima D-nn y RESUMEN hay una fila vacia

    # correlativo siguiente a partir de los codigos D-nn existentes
    numeros = []
    for r in range(primera_d, ultima_d + 1):
        cod = str(ws.cell(row=r, column=1).value or "")
        m = re.match(r"^D-(\d+)$", cod.strip())
        if m:
            numeros.append(int(m.group(1)))
    siguiente = (max(numeros) + 1) if numeros else 1

    fila_nueva = ultima_d + 1
    insertar_filas(ws, fila_nueva, 2, 5)
    for offset in (0, 1):
        copiar_estilo_fila(ws, ultima_d, fila_nueva + offset, 5)

    # --- D-nn: incorporacion de la distribucion areal ----------------------
    escribir(ws, fila_nueva, 1, "D-%02d" % siguiente)
    escribir(ws, fila_nueva, 2, "Distribución areal MSAVI 2024")
    escribir(ws, fila_nueva, 3,
             "La distribución de superficie por clase de MSAVI 2024 se consignaba "
             "como «Por determinar»: la estadística zonal por clase no estaba "
             "disponible en los insumos del entregable")
    escribir(ws, fila_nueva, 4, "CORREGIDO")
    escribir(ws, fila_nueva, 5,
             "Se incorpora la distribución areal por clase espectral (campo DN, 5 "
             "clases) del shapefile MSAVI V6, consolidada en "
             "«AREAS_MSAVI_BLOQUES_V6_REV_HSCM.xlsx», hoja «Resumen_por_Bloque» "
             "(%d polígonos del bloque). Trasladada a la hoja «Cobertura "
             "MSAVI-NDVI», sección A, y sintetizada en la hoja «Resumen», sección 3."
             % met["n_poligonos"])

    # --- D-nn+1: contraste planimetrico ------------------------------------
    escribir(ws, fila_nueva + 1, 1, "D-%02d" % (siguiente + 1))
    escribir(ws, fila_nueva + 1, 2, "Superficie clasificada MSAVI vs. catálogo")
    if met["desviacion"] is None:
        escribir(ws, fila_nueva + 1, 3,
                 "La plantilla no consigna superficie de catálogo V5/V6 "
                 "contrastable con la superficie clasificada por MSAVI 2024 "
                 "(%s ha)" % _num(met["total_ha"]))
        escribir(ws, fila_nueva + 1, 4, "SUSTANTIVA")
        escribir(ws, fila_nueva + 1, 5,
                 "Debe reponerse la superficie de catálogo del bloque antes de "
                 "validar la distribución areal.")
    elif abs(met["desviacion"]) <= TOLERANCIA_AREA_PCT:
        escribir(ws, fila_nueva + 1, 3,
                 "La superficie clasificada por MSAVI 2024 (%s ha) coincide con la "
                 "de catálogo V5/V6 (%s ha) dentro del ±%0.0f %% (%+0.2f %%)"
                 % (_num(met["total_ha"]), _num(met["catalogo_ha"]),
                    TOLERANCIA_AREA_PCT, met["desviacion"]))
        escribir(ws, fila_nueva + 1, 4, "CONFORME")
        escribir(ws, fila_nueva + 1, 5,
                 "Sin acción: la clasificación espectral valida la geometría del "
                 "bloque.")
    else:
        escribir(ws, fila_nueva + 1, 3,
                 "La superficie clasificada por MSAVI 2024 (%s ha) difiere de la de "
                 "catálogo V5/V6 (%s ha) en %+0.2f %%, fuera de la tolerancia de "
                 "±%0.0f %%"
                 % (_num(met["total_ha"]), _num(met["catalogo_ha"]),
                    met["desviacion"], TOLERANCIA_AREA_PCT))
        escribir(ws, fila_nueva + 1, 4, "SUSTANTIVA")
        escribir(ws, fila_nueva + 1, 5,
                 "Se conserva la superficie de catálogo como referencia "
                 "planimétrica oficial y se declara la discrepancia. Debe "
                 "verificarse el recorte del ráster MSAVI V6 al polígono y la "
                 "vigencia del límite en el catálogo maestro antes del expediente "
                 "técnico. La distribución por clase se reporta sobre la superficie "
                 "efectivamente clasificada; no debe usarse para metas físicas "
                 "hasta conciliar la geometría.")

    # --- recalculo del cuadro RESUMEN --------------------------------------
    fila_resumen += 2
    ultima_d += 2
    califs = Counter()
    for r in range(primera_d, ultima_d + 1):
        val = ws.cell(row=r, column=4).value
        if val:
            califs[str(val).strip().upper()] += 1
    total = sum(califs.values())
    escribir(ws, fila_resumen, 2, "%d verificaciones" % total)
    escribir(ws, fila_resumen, 3,
             " · ".join("%s: %d" % (k, califs[k]) for k in sorted(califs)))

    return "D-%02d" % siguiente, "D-%02d" % (siguiente + 1), total


# --------------------------------------------------------------------------- #
# Proceso por plantilla
# --------------------------------------------------------------------------- #
def procesar_plantilla(ruta_entrada, ruta_salida, registro):
    wb = openpyxl.load_workbook(ruta_entrada)

    faltantes = [h for h in (HOJA_COBERTURA, HOJA_RESUMEN, HOJA_CONTROL)
                 if h not in wb.sheetnames]
    if faltantes:
        raise ValueError("hojas ausentes: %s" % ", ".join(faltantes))

    ws_res = wb[HOJA_RESUMEN]
    ws_cob = wb[HOJA_COBERTURA]

    # Guardas: la plantilla debe conservar el layout esperado.
    if ws_cob["A9"].value != "Clase MSAVI":
        raise ValueError("la sección A de «%s» no tiene el layout esperado"
                         % HOJA_COBERTURA)
    for fila, _dn, clase in FILAS_CLASE:
        if str(ws_cob.cell(row=fila, column=1).value).strip() != clase:
            raise ValueError("clase inesperada en la fila %d: se esperaba «%s»"
                             % (fila, clase))
    if ws_res["A28"].value != "MSAVI 2024 — media del bloque":
        raise ValueError("la sección 3 de «%s» no tiene el layout esperado"
                         % HOJA_RESUMEN)

    catalogo_ha = ws_res["B14"].value
    catalogo_ha = float(catalogo_ha) if isinstance(catalogo_ha, (int, float)) else None
    media_msavi = ws_cob["B%d" % FILA_MEDIA_ORIGINAL].value
    if not isinstance(media_msavi, (int, float)):
        media_msavi = ws_res["B28"].value

    met = metricas_bloque(registro, catalogo_ha)

    actualizar_cobertura(ws_cob, met, media_msavi)
    actualizar_resumen(ws_res, met)
    cod1, cod2, n_verif = actualizar_control(wb[HOJA_CONTROL], met)

    wb.save(ruta_salida)
    wb.close()

    met["codigos_control"] = (cod1, cod2)
    met["n_verificaciones"] = n_verif
    met["media_msavi"] = media_msavi
    return met


# --------------------------------------------------------------------------- #
# Reporte consolidado de traspaso (formato ANIN)
# --------------------------------------------------------------------------- #
def generar_reporte(resultados, ruta, nombre_fuente):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Traspaso MSAVI"

    columnas = [
        ("BLOQUE", 14), ("N.° POLÍGONOS", 13),
        ("DN 5 · > 0.6139 (ha)", 19), ("DN 4 · 0.4976-0.6139 (ha)", 22),
        ("DN 3 · 0.3813-0.4976 (ha)", 22), ("DN 2 · 0.2650-0.3813 (ha)", 22),
        ("DN 1 · <= 0.2650 (ha)", 19), ("TOTAL MSAVI (ha)", 16),
        ("CATÁLOGO V5/V6 (ha)", 18), ("DESVIACIÓN (%)", 14),
        ("BAJO UMBRAL 0.4976 (ha)", 21), ("BAJO UMBRAL (%)", 15),
        ("MSAVI MEDIA", 12), ("CONTROL DE CONSISTENCIA", 24),
        ("CALIFICACIÓN PLANIMÉTRICA", 24),
    ]

    for i, texto in enumerate(ENCABEZADOS_ANIN, start=1):
        c = ws.cell(row=i, column=1, value=texto)
        c.font = Font(name=FUENTE, size=10, bold=True, color=VERDE_ANIN)
        ws.merge_cells(start_row=i, start_column=1,
                       end_row=i, end_column=len(columnas))
        c.alignment = Alignment(horizontal="center")

    c = ws.cell(row=4, column=1,
                value="TRASPASO DE ÁREAS MSAVI POR CLASE (DN) A LAS PLANTILLAS "
                      "DE BLOQUE — PROYECTO IN PIURA (CUI 2669244)")
    c.font = Font(name=FUENTE, size=11, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=VERDE_ANIN)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=len(columnas))
    ws.row_dimensions[4].height = 26

    c = ws.cell(row=5, column=1,
                value="Fuente: %s, hoja «Resumen_por_Bloque» (shapefile MSAVI V6, "
                      "campo DN) · Sistema de referencia UTM WGS 84 Zona 17S "
                      "(EPSG:32717) · Umbral de brecha del proyecto MSAVI = %0.4f"
                      % (nombre_fuente, UMBRAL_BRECHA))
    c.font = Font(name=FUENTE, size=8, italic=True)
    c.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=len(columnas))

    fila_enc = 7
    for j, (titulo, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=fila_enc, column=j, value=titulo)
        c.font = Font(name=FUENTE, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERDE_ANIN)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE_FINO
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.row_dimensions[fila_enc].height = 32

    fila = fila_enc + 1
    for r in sorted(resultados, key=_orden_bloque):
        met = resultados[r]
        calif = ("SIN CATÁLOGO" if met["desviacion"] is None
                 else "CONFORME" if abs(met["desviacion"]) <= TOLERANCIA_AREA_PCT
                 else "SUSTANTIVA")
        valores = [
            r, met["n_poligonos"],
            round(met["ha"][5], 4), round(met["ha"][4], 4), round(met["ha"][3], 4),
            round(met["ha"][2], 4), round(met["ha"][1], 4),
            round(met["total_ha"], 4),
            round(met["catalogo_ha"], 4) if met["catalogo_ha"] else "—",
            round(met["desviacion"], 2) if met["desviacion"] is not None else "—",
            round(met["ha_bajo"], 4), round(met["pct_bajo"], 2),
            round(met["media_msavi"], 4) if isinstance(met["media_msavi"], (int, float)) else "—",
            " / ".join(x for x in met["codigos_control"] if x),
            calif,
        ]
        for j, v in enumerate(valores, start=1):
            c = ws.cell(row=fila, column=j, value=v)
            c.font = Font(name=FUENTE, size=9,
                          bold=(calif == "SUSTANTIVA" and j in (1, 10, 15)))
            c.border = BORDE_FINO
            c.alignment = Alignment(
                horizontal="left" if j in (1, 14, 15) else "center")
            if 3 <= j <= 9 or j == 11:
                c.number_format = "0.0000"
            elif j in (10, 12, 13):
                c.number_format = "0.00"
            if fila % 2 == 0:
                c.fill = PatternFill("solid", fgColor=VERDE_CLARO)
            if calif == "SUSTANTIVA" and j in (10, 15):
                c.fill = PatternFill("solid", fgColor="FCE4D6")
        fila += 1

    ultima = fila - 1
    c = ws.cell(row=fila, column=1, value="TOTAL / PROMEDIO")
    for j in range(1, len(columnas) + 1):
        c = ws.cell(row=fila, column=j)
        c.font = Font(name=FUENTE, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERDE_ANIN)
        c.border = BORDE_FINO
        c.alignment = Alignment(horizontal="center")
    ws.cell(row=fila, column=1, value="TOTAL / PROMEDIO").alignment = \
        Alignment(horizontal="left")
    for j in list(range(2, 12)):
        letra = get_column_letter(j)
        ws.cell(row=fila, column=j,
                value="=ROUND(SUM(%s%d:%s%d),4)" % (letra, fila_enc + 1, letra, ultima))
        ws.cell(row=fila, column=j).number_format = "0.0000"
    ws.cell(row=fila, column=2).value = "=SUM(B%d:B%d)" % (fila_enc + 1, ultima)
    ws.cell(row=fila, column=2).number_format = "0"
    for j, dec in ((10, 2), (12, 2), (13, 4)):
        letra = get_column_letter(j)
        ws.cell(row=fila, column=j,
                value="=ROUND(AVERAGE(%s%d:%s%d),%d)"
                      % (letra, fila_enc + 1, letra, ultima, dec))
        ws.cell(row=fila, column=j).number_format = "0." + "0" * dec
    ws.cell(row=fila, column=14, value="%d bloques procesados" % len(resultados))
    n_sust = sum(1 for m in resultados.values()
                 if m["desviacion"] is None or abs(m["desviacion"]) > TOLERANCIA_AREA_PCT)
    ws.cell(row=fila, column=15,
            value="CONFORME: %d · SUSTANTIVA: %d" % (len(resultados) - n_sust, n_sust))

    ws.freeze_panes = "B%d" % (fila_enc + 1)
    ws.auto_filter.ref = "A%d:%s%d" % (fila_enc, get_column_letter(len(columnas)), ultima)
    ws.sheet_view.showGridLines = False

    wb.save(ruta)
    wb.close()


def _orden_bloque(codigo):
    """Ordena los bloques numericos primero y luego los codificados M#B#."""
    if codigo.isdigit():
        return (0, int(codigo), 0, "")
    m = re.match(r"^M(\d+)B(\d+)(.*)$", codigo)
    if m:
        return (1, int(m.group(1)), int(m.group(2)), m.group(3))
    return (2, 0, 0, codigo)


# --------------------------------------------------------------------------- #
# Programa principal
# --------------------------------------------------------------------------- #
def main(argv=None):
    ap = argparse.ArgumentParser(
        description="Traspasa las areas MSAVI por clase (DN) a las plantillas "
                    "Excel de bloque del Proyecto IN Piura.")
    ap.add_argument("--fuente", required=True,
                    help="Excel consolidado de areas MSAVI por bloque.")
    ap.add_argument("--plantillas", required=True,
                    help="Carpeta con las plantillas Plantilla_Excel_Bloque_*.xlsx")
    ap.add_argument("--salida",
                    help="Carpeta de salida. Si se omite, se actualiza en el sitio.")
    ap.add_argument("--reporte",
                    help="Ruta del Excel de reporte consolidado del traspaso.")
    args = ap.parse_args(argv)

    fuente = leer_fuente_msavi(args.fuente)
    print("Fuente: %s — %d bloques con áreas MSAVI." % (args.fuente, len(fuente)))

    archivos = sorted(f for f in os.listdir(args.plantillas)
                      if RE_PLANTILLA.match(f))
    if not archivos:
        raise SystemExit("No se encontraron plantillas en %s" % args.plantillas)

    destino = args.salida or args.plantillas
    if args.salida:
        os.makedirs(destino, exist_ok=True)

    resultados, sin_fuente, errores = {}, [], []
    for nombre in archivos:
        codigo = RE_PLANTILLA.match(nombre).group(1)
        entrada = os.path.join(args.plantillas, nombre)
        salida = os.path.join(destino, nombre)
        if codigo not in fuente:
            sin_fuente.append(codigo)
            if args.salida:
                shutil.copy2(entrada, salida)
            continue
        try:
            resultados[codigo] = procesar_plantilla(entrada, salida, fuente[codigo])
        except Exception as exc:                      # noqa: BLE001
            errores.append((codigo, str(exc)))
            if args.salida and not os.path.exists(salida):
                shutil.copy2(entrada, salida)

    print("Plantillas actualizadas: %d de %d." % (len(resultados), len(archivos)))
    if sin_fuente:
        print("Sin datos MSAVI en la fuente (%d): %s"
              % (len(sin_fuente), ", ".join(sin_fuente)))
    if errores:
        print("Con error (%d):" % len(errores))
        for codigo, msg in errores:
            print("   - %s: %s" % (codigo, msg))

    fuera = [(c, m["desviacion"]) for c, m in resultados.items()
             if m["desviacion"] is None or abs(m["desviacion"]) > TOLERANCIA_AREA_PCT]
    print("Bloques con desviación planimétrica > ±%0.0f %%: %d"
          % (TOLERANCIA_AREA_PCT, len(fuera)))

    if args.reporte and resultados:
        generar_reporte(resultados, args.reporte, os.path.basename(args.fuente))
        print("Reporte consolidado: %s" % args.reporte)

    return 1 if errores else 0


if __name__ == "__main__":
    sys.exit(main())
