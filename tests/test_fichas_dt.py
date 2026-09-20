"""Pruebas de la analitica de las fichas F-DT actualizadas.

Las pruebas de lectura corren sobre los 117 libros de campo que viajan con
el aplicativo, igual que las de los resumenes por bloque: es la unica forma
de comprobar que el lector sigue entendiendo los archivos reales, que no
tienen una convencion unica de nombre ni de disposicion de celdas. Las
pruebas de agregacion usan fichas sinteticas, donde el valor esperado se
puede calcular a mano.
"""

import io
import os
import sys
import tempfile
import unittest

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import fichas_dt as fdt  # noqa: E402
import resumenes_bloques as rb  # noqa: E402


def ficha_sintetica(codigo, distrito="Chalaco", provincia="Morropon",
                    **campos):
    """Ficha ya normalizada, como la deja `parsear_ficha`."""
    ficha = {
        "codigo_bloque": codigo,
        "nombre_archivo": f"DT_{codigo}.xlsx",
        "distrito": distrito,
        "provincia": provincia,
        "microcuenca": "C1096-Q9564",
        "area_ha": None,
    }
    ficha.update(campos)
    for clave, _e, _u, _a in fdt.INDICADORES_NUM:
        ficha.setdefault(clave + "_num", None)
    for clave, _e in fdt.INDICADORES_SINO:
        ficha.setdefault(clave + "_si", None)
    for clave, _e in fdt.INDICADORES_CONTEO:
        ficha.setdefault(clave, 0)
    for tipo, _e, _c, _a, _co in fdt.INVENTARIOS:
        ficha.setdefault(f"inv_{tipo}", [])
    return ficha


# ══════════════════════════════════════════════════════════════════════════
# Normalizacion
# ══════════════════════════════════════════════════════════════════════════

class TestNormalizacion(unittest.TestCase):

    def test_las_declaraciones_de_ausencia_no_son_datos(self):
        for ausente in ("", "   ", "Por determinar", "por verificar", "N/A",
                        "Sin registro", "—", "No aplica"):
            self.assertFalse(fdt.declarado(ausente), ausente)
        for presente in ("Moderada (surcos)", "0", "Sí", "12.5"):
            self.assertTrue(fdt.declarado(presente), presente)

    def test_el_cero_es_un_dato_y_no_una_ausencia(self):
        ficha = {"dt02_num_carcavas": 0}
        self.assertEqual(fdt.valor_num(ficha, "dt02_num_carcavas"), 0.0)

    def test_si_no_distingue_la_respuesta_de_la_falta_de_respuesta(self):
        self.assertIs(fdt.es_si("Sí"), True)
        self.assertIs(fdt.es_si("si"), True)
        self.assertIs(fdt.es_si("No"), False)
        self.assertIsNone(fdt.es_si("Por verificar"))
        self.assertIsNone(fdt.es_si(""))

    def test_la_categoria_ausente_se_rotula_en_vez_de_quedar_vacia(self):
        self.assertEqual(fdt.categoria({"x": "Por determinar"}, "x"),
                         fdt.SIN_DATO)
        self.assertEqual(fdt.categoria({"x": "Ladera media"}, "x"),
                         "Ladera media")

    def test_las_filas_en_blanco_del_inventario_no_son_registros(self):
        filas = [{"tipo": "Manantial / Puquio", "regimen": "Permanente"},
                 {"tipo": "", "regimen": ""},
                 {"tipo": "Por determinar", "regimen": ""}]
        self.assertEqual(len(fdt.filas_con_contenido(filas, ("tipo", "regimen"))), 1)

    def test_el_codigo_sale_del_nombre_cuando_la_ficha_no_lo_declara(self):
        casos = {
            "Plantilla_DT_Campo_Check_Validada_V5_Bloque23.xlsx": "23",
            "DT_Bloque67_V5.xlsx": "67",
            "F-DT_M6B10_Diagnostico_Territorial_rev1_dron.xlsx": "M6B10",
            "DT_B13_IN_Piura.xlsx": "13",
            "Plantilla_Excel_M6B2-1_IN_Piura.xlsx": "M6B2-1",
        }
        for nombre, esperado in casos.items():
            self.assertEqual(fdt.codigo_desde_nombre(nombre), esperado, nombre)


# ══════════════════════════════════════════════════════════════════════════
# Agregacion
# ══════════════════════════════════════════════════════════════════════════

class TestAgregacion(unittest.TestCase):

    def setUp(self):
        self.fichas = [
            ficha_sintetica("M1B1", distrito="Frias", provincia="Ayabaca",
                            dt03_cobertura_total_num=80.0,
                            dt02_num_carcavas_num=2.0,
                            n_taxones=10, n_fuentes_agua=2,
                            dt02_nivel_erosion_sintesis="Moderada (surcos)",
                            dt01_escarpes_activos_si=True),
            ficha_sintetica("M1B2", distrito="Frías", provincia="Ayabaca",
                            dt03_cobertura_total_num=60.0,
                            dt02_num_carcavas_num=4.0,
                            n_taxones=6, n_fuentes_agua=1,
                            dt02_nivel_erosion_sintesis="Moderada (surcos)",
                            dt01_escarpes_activos_si=False),
            ficha_sintetica("M2B1", distrito="Chalaco", provincia="Morropon",
                            n_taxones=4,
                            dt02_nivel_erosion_sintesis="Por determinar"),
        ]

    def test_distrito_agrupa_pese_a_la_tilde(self):
        grupos = fdt.agrupar(self.fichas, "distrito")
        self.assertEqual([e for e, _g in grupos], ["Chalaco", "Frias"])
        self.assertEqual(len(dict(grupos)["Frias"]), 2)

    def test_los_cuatro_niveles_responden(self):
        self.assertEqual(len(fdt.agrupar(self.fichas, "bloque")), 3)
        self.assertEqual(len(fdt.agrupar(self.fichas, "distrito")), 2)
        self.assertEqual(len(fdt.agrupar(self.fichas, "provincia")), 2)
        self.assertEqual(len(fdt.agrupar(self.fichas, "total")), 1)

    def test_un_nivel_desconocido_no_pasa_en_silencio(self):
        with self.assertRaises(ValueError):
            fdt.agrupar(self.fichas, "microcuenca")

    def test_el_promedio_solo_toma_los_bloques_que_declaran(self):
        total = fdt.tabla_resumen(self.fichas, "total")[0]
        # 80 y 60 declarados; el tercer bloque no declara y no promedia.
        self.assertEqual(total["dt03_cobertura_total"], 70.0)
        self.assertEqual(total["dt03_cobertura_total_n"], 2)
        self.assertEqual(total["n_bloques"], 3)

    def test_los_conteos_se_suman_sobre_todo_el_grupo(self):
        total = fdt.tabla_resumen(self.fichas, "total")[0]
        self.assertEqual(total["n_taxones"], 20)
        self.assertEqual(total["n_fuentes_agua"], 3)
        self.assertEqual(total["dt02_num_carcavas"], 6.0)

    def test_el_si_no_se_reporta_sobre_los_que_responden(self):
        total = fdt.tabla_resumen(self.fichas, "total")[0]
        self.assertEqual(total["dt01_escarpes_activos_si"], 1)
        self.assertEqual(total["dt01_escarpes_activos_n"], 2)
        self.assertEqual(total["dt01_escarpes_activos_pct"], 50.0)

    def test_un_indicador_sin_ningun_dato_queda_en_nulo(self):
        total = fdt.tabla_resumen(self.fichas, "total")[0]
        self.assertIsNone(total["dt03_dap_promedio"])
        self.assertEqual(total["dt03_dap_promedio_n"], 0)

    def test_la_distribucion_cierra_con_lo_no_declarado(self):
        filas = fdt.distribucion(self.fichas, "dt02_nivel_erosion_sintesis",
                                 "total")
        self.assertEqual(filas[0]["Categoria"], "Moderada (surcos)")
        self.assertEqual(filas[0]["Bloques"], 2)
        self.assertAlmostEqual(filas[0]["% del grupo"], 66.7, places=1)
        self.assertEqual(filas[-1]["Categoria"], fdt.SIN_DATO)

    def test_el_ranking_cuenta_registros_y_bloques(self):
        fichas = [
            ficha_sintetica("A", inv_floristica=[
                {"nombre_comun": "Higueron"}, {"nombre_comun": "Higueron"},
                {"nombre_comun": "Faique"}]),
            ficha_sintetica("B", inv_floristica=[{"nombre_comun": "Higueron"}]),
        ]
        rank = fdt.ranking(fichas, "floristica", "nombre_comun")
        self.assertEqual(rank[0], {"Valor": "Higueron", "Registros": 3,
                                   "Bloques": 2})
        self.assertEqual(rank[1]["Valor"], "Faique")

    def test_las_causas_presentes_excluyen_las_no_declaradas(self):
        fichas = [ficha_sintetica("A", inv_causas=[
            {"causa": "Sobrepastoreo", "presencia": "Sí"},
            {"causa": "Tala para lena", "presencia": "No"},
            {"causa": "Quema", "presencia": "Por determinar"}])]
        self.assertEqual(fdt.causas_presentes(fichas),
                         [{"Valor": "Sobrepastoreo", "Registros": 1,
                           "Bloques": 1}])

    def test_el_catalogo_del_aplicativo_manda_sobre_la_ficha(self):
        ficha = {"codigo_bloque": "M7B3", "provincia": "90",
                 "distrito": "Salitral"}
        fdt.resolver_localidad(ficha, {"M7B3": {
            "provincia": "Morropon", "distrito": "Salitral",
            "microcuenca": "C1077-Q9566", "area_ha": 120.5}})
        self.assertEqual(ficha["provincia"], "Morropon")
        self.assertEqual(ficha["provincia_ficha"], "90")
        self.assertEqual(ficha["area_ha"], 120.5)

    def test_sin_catalogo_se_conserva_lo_declarado_en_la_ficha(self):
        ficha = {"codigo_bloque": "M7B3", "provincia": "Morropon",
                 "distrito": ""}
        fdt.resolver_localidad(ficha, {})
        self.assertEqual(ficha["provincia"], "Morropon")
        self.assertEqual(ficha["distrito"], "")


# ══════════════════════════════════════════════════════════════════════════
# Lectura de los 117 libros de campo del repositorio
# ══════════════════════════════════════════════════════════════════════════

class TestFichasDelRepositorio(unittest.TestCase):
    """Los libros F-DT actualizados que el aplicativo trae consigo."""

    @classmethod
    def setUpClass(cls):
        cls.libros = fdt.fichas_del_repositorio()
        cls.fichas, cls.errores = fdt.cargar_fichas(cls.libros)

    def test_el_repositorio_trae_los_117_libros(self):
        self.assertEqual(len(self.libros), 117)

    def test_carpeta_ausente_no_rompe_la_carga(self):
        self.assertEqual(fdt.fichas_del_repositorio("/no/existe"), [])
        self.assertEqual(fdt.fichas_del_zip("/no/existe.zip"), [])

    def test_la_carpeta_extraida_manda_sobre_el_zip(self):
        with tempfile.TemporaryDirectory() as tmp:
            nombre, contenido = self.libros[0]
            with open(os.path.join(tmp, nombre), "wb") as fh:
                fh.write(contenido)
            with open(os.path.join(tmp, "~$temporal.xlsx"), "wb") as fh:
                fh.write(b"basura")
            self.assertEqual([n for n, _c in fdt.fichas_del_repositorio(tmp)],
                             [nombre])

    def test_las_117_fichas_se_leen_sin_errores(self):
        self.assertEqual(self.errores, [])
        self.assertEqual(len(self.fichas), 117)

    def test_cada_libro_trae_las_cinco_fichas_oficiales(self):
        for ficha in self.fichas:
            self.assertEqual(sorted(ficha["fichas_leidas"]),
                             ["F-DT-01", "F-DT-02", "F-DT-03", "F-DT-04",
                              "F-DT-05"], ficha["nombre_archivo"])

    def test_los_codigos_son_los_mismos_117_bloques_de_los_resumenes(self):
        codigos = {f["codigo_bloque"] for f in self.fichas}
        self.assertEqual(len(codigos), 117)
        self.assertEqual(codigos, set(rb.codigos_esperados()))

    def test_los_inventarios_llegan_sin_las_filas_en_blanco(self):
        for tipo, _e, _c, _a, _co in fdt.INVENTARIOS:
            filas = fdt.inventario(self.fichas, tipo)
            self.assertTrue(filas, tipo)
            for fila in filas[:200]:
                self.assertTrue(fila["Bloque"])

    def test_la_analitica_responde_en_los_cuatro_niveles(self):
        for nivel in fdt.NIVELES:
            filas = fdt.tabla_resumen(self.fichas, nivel)
            self.assertTrue(filas, nivel)
            self.assertEqual(sum(f["n_bloques"] for f in filas), 117, nivel)

    def test_ningun_promedio_se_sostiene_sobre_cero_bloques(self):
        total = fdt.tabla_resumen(self.fichas, "total")[0]
        for clave, etiqueta, _u, _a in fdt.INDICADORES_NUM:
            if total[clave] is not None:
                self.assertGreater(total[clave + "_n"], 0, etiqueta)


# ══════════════════════════════════════════════════════════════════════════
# Exportaciones
# ══════════════════════════════════════════════════════════════════════════

class TestExportaciones(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        libros = fdt.fichas_del_repositorio()[:8]
        cls.fichas, _errores = fdt.cargar_fichas(libros)

    def test_el_excel_trae_resumen_distribuciones_e_inventarios(self):
        salida = fdt.generar_excel_fdt(self.fichas, "distrito")
        wb = load_workbook(io.BytesIO(salida))
        self.assertIn("Resumen Distrito", wb.sheetnames)
        self.assertIn("Distribuciones", wb.sheetnames)
        self.assertIn("Rankings", wb.sheetnames)
        self.assertTrue([h for h in wb.sheetnames if h.startswith("Inv. ")])
        self.assertTrue(wb["Resumen Distrito"]._charts)

    def test_el_excel_responde_en_los_cuatro_niveles(self):
        for nivel in fdt.NIVELES:
            salida = fdt.generar_excel_fdt(self.fichas, nivel,
                                           incluir_inventarios=False)
            wb = load_workbook(io.BytesIO(salida))
            self.assertIn(f"Resumen {fdt.ETIQUETA_NIVEL[nivel]}"[:31],
                          wb.sheetnames, nivel)

    def test_el_pdf_responde_en_los_cuatro_niveles(self):
        for nivel in fdt.NIVELES:
            salida = fdt.generar_pdf_fdt(self.fichas, nivel)
            self.assertTrue(salida.startswith(b"%PDF"), nivel)

    def test_un_nivel_desconocido_no_genera_reporte(self):
        with self.assertRaises(ValueError):
            fdt.generar_excel_fdt(self.fichas, "microcuenca")
        with self.assertRaises(ValueError):
            fdt.generar_pdf_fdt(self.fichas, "microcuenca")


if __name__ == "__main__":
    unittest.main(verbosity=2)
