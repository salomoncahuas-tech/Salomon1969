"""F-DS-02: separacion "Nombre del actor" / "Cargo" sin perder datos antiguos."""
import io
import json
import os
import sys
import types

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# database.py exige st.secrets["DATABASE_URL"] al importarse; estas pruebas no tocan la base.
try:
    import database  # noqa: F401
except Exception:
    sys.modules["database"] = types.ModuleType("database")

import analitica_social as ans  # noqa: E402
import export_diagnosticos as exp  # noqa: E402
import fds_actores as FA  # noqa: E402
import reports  # noqa: E402

FILA_ANTIGUA = {
    "Nombre del actor / Organizacion": "Municipalidad Distrital de Huancabamba",
    "Tipo": "Gobierno Local (Municipalidad)",
    "Rol / Funcion frente al proyecto": "Articulación",
    "Influencia": "Medio", "Interes": "Alto",
    "Posicion": "Neutral / Sin posición definida",
    "Observaciones / Historial": "Apoyo en convenio 2024",
}
FILA_NUEVA = {
    "Nombre del actor": "Juan Pérez Huamán", "Cargo": "Presidente de la CC",
    "Tipo": "Comunidad Campesina",
    "Rol / Funcion frente al proyecto": "Titular del predio",
    "Influencia": "Alto", "Interes": "Alto", "Posicion": "A favor del proyecto",
}


def _registro(filas):
    form = {"f2_actores": filas}
    return {"id": 1, "ficha": "F-DS-02", "bloque_codigo": "M5-B1",
            "centro_poblado": "Chungayo", "fecha_evaluacion": "2026-03-15",
            "ds02_data_v3": json.dumps(form, ensure_ascii=False),
            "ds02_registro_actores": json.dumps(filas, ensure_ascii=False)}


def test_migrar_fila_antigua_conserva_texto_y_campos():
    fila = FA.migrar_fila(FILA_ANTIGUA)
    assert fila["Nombre del actor"] == "Municipalidad Distrital de Huancabamba"
    assert fila["Cargo"] == ""
    assert "Nombre del actor / Organizacion" not in fila
    assert list(fila)[:2] == ["Nombre del actor", "Cargo"]
    assert fila["Observaciones / Historial"] == "Apoyo en convenio 2024"


def test_migrar_fila_nueva_sin_cambios():
    assert FA.migrar_fila(FILA_NUEVA) == FILA_NUEVA


def test_migrar_tabla_solo_afecta_actores():
    otra = [{"Nombre del actor / Organizacion": "x"}]
    assert FA.migrar_tabla("f4_part", otra) is otra
    assert FA.migrar_tabla("f2_actores", otra)[0]["Nombre del actor"] == "x"


def test_pdf_lee_formato_nuevo_y_antiguo():
    filas = reports._actores_ds02(_registro([FILA_ANTIGUA, FILA_NUEVA]))
    assert filas[0]["nombre"] == "Municipalidad Distrital de Huancabamba"
    assert filas[0]["cargo"] == ""
    assert filas[1]["nombre"] == "Juan Pérez Huamán"
    assert filas[1]["cargo"] == "Presidente de la CC"
    assert filas[1]["posicion"] == "A favor del proyecto"


def test_pdf_legacy_claves_cortas():
    ds = {"ds02_registro_actores": json.dumps(
        [{"nombre": "Ronda de Chungayo", "tipo": "Ronda Campesina",
          "rol": "Acceso", "relacion": "A favor", "influencia": "Alto"}])}
    fila = reports._actores_ds02(ds)[0]
    assert fila["nombre"] == "Ronda de Chungayo"
    assert fila["posicion"] == "A favor"


def test_excel_consolidado_separa_columnas():
    datos = exp.exportar_fds_consolidado([_registro([FILA_ANTIGUA, FILA_NUEVA])])
    ws = load_workbook(io.BytesIO(datos))["Actores clave"]
    cab = [c.value for c in ws[1]]
    assert "Nombre del actor" in cab and "Cargo" in cab
    assert "Nombre del actor / Organizacion" not in cab
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    i_nom, i_car = cab.index("Nombre del actor"), cab.index("Cargo")
    assert filas[0][i_nom] == "Municipalidad Distrital de Huancabamba"
    assert filas[1][i_nom] == "Juan Pérez Huamán"
    assert filas[1][i_car] == "Presidente de la CC"


def test_analitica_detalle_con_nombre_y_cargo():
    sec = ans._seccion_actores([_registro([FILA_ANTIGUA, FILA_NUEVA])])
    detalle = dict(sec["tablas"])["Actores clave"]
    assert detalle[0]["Nombre del actor"] == "Municipalidad Distrital de Huancabamba"
    assert detalle[1]["Cargo"] == "Presidente de la CC"
