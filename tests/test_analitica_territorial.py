"""Pruebas de la analitica del Diagnostico Territorial.

Cubren lo que sostiene el informe: que las secciones se construyan solo con
lo declarado, que el color comunique severidad y no identidad, que la
agregacion consolidada cuadre con los bloques que la componen y que los tres
renderizadores -Altair, Excel y PDF- consuman la misma declaracion sin
contradecirse.
"""

import io
import os
import sys
import unittest

from openpyxl import load_workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import analitica_series as ase          # noqa: E402
import analitica_territorial as ate     # noqa: E402
import dt_campo as dtc                  # noqa: E402
import resumenes_bloques as rbq         # noqa: E402
from test_dt_campo import ficha_campo                          # noqa: E402
from test_resumenes_bloques import (construir_libro,           # noqa: E402
                                    graficos_encimados)


def datos_bloque(codigo="M9B1", **cambios):
    """Bloque con distribucion areal del MSAVI, que es el caso vigente."""
    datos = rbq.completar_sintesis_msavi(rbq.parsear_resumen_bloque(
        construir_libro(codigo=codigo, con_msavi_areal=True),
        f"Plantilla_Excel_Bloque_{codigo}_IN_Piura.xlsx"))
    datos.update(cambios)
    return datos


# ══════════════════════════════════════════════════════════════════════════
# Secciones de un bloque
# ══════════════════════════════════════════════════════════════════════════

class TestSeccionesDeBloque(unittest.TestCase):

    def setUp(self):
        self.datos = datos_bloque()
        self.integrado = dtc.integrar_bloque(self.datos, ficha_campo())

    def test_el_informe_trae_las_secciones_con_contenido(self):
        informe = ate.indicadores_bloque(self.datos, self.integrado)
        titulos = [s["titulo"] for s in informe["secciones"]]
        self.assertIn("Índices de vegetación", titulos)
        self.assertIn("Control de consistencia", titulos)
        self.assertIn("Procedencia de los datos", titulos)
        for seccion in informe["secciones"]:
            self.assertTrue(seccion["series"] or seccion["tablas"],
                            seccion["titulo"])

    def test_toda_serie_declara_lo_que_el_motor_necesita(self):
        informe = ate.indicadores_bloque(self.datos, self.integrado)
        for seccion in informe["secciones"]:
            for serie in seccion["series"]:
                self.assertTrue(serie["id"])
                self.assertTrue(serie["titulo"])
                self.assertTrue(serie["filas"])
                self.assertIn(serie["forma"],
                              ("barras_h", "apiladas", "agrupadas", "mapa_calor"))
                for fila in serie["filas"]:
                    self.assertIn(serie["cat"], fila)
                    self.assertIn(serie["val"], fila)

    def test_las_clases_msavi_se_colorean_por_el_umbral(self):
        """Bajo el umbral, tono de severidad; sobre él, tono favorable."""
        serie = ate.series_de_bloque(self.datos)["dt_msavi"]
        tabla = {r["clase"]: r for r in self.datos["msavi_tabla"]}
        for clase, color in serie["colores"].items():
            bajo = "bajo" in tabla[clase]["condicion"].lower()
            rampa = ate.RAMPA_CRITICA if bajo else ate.RAMPA_FAVORABLE
            self.assertIn(color, rampa, clase)

    def test_no_se_grafica_lo_que_el_libro_no_declara(self):
        datos = datos_bloque()
        datos["msavi_tabla"] = []
        datos["msavi_bajo_umbral_ha"] = None
        datos["msavi_sobre_umbral_ha"] = None
        series = ate.series_de_bloque(datos)
        self.assertNotIn("dt_msavi", series)
        self.assertNotIn("dt_brecha", series)
        informe = ate.indicadores_bloque(datos)
        self.assertTrue(any("distribución areal" in a for a in informe["avisos"]))

    def test_el_bloque_de_la_ficha_se_destaca_en_su_microcuenca(self):
        serie = ate.series_de_bloque(self.datos).get("dt_micro_area")
        if serie is None:
            self.skipTest("El fixture no trae contexto intramicrocuenca")
        destacados = [c for c in serie["colores"] if c.startswith("►")]
        self.assertEqual(len(destacados), 1)

    def test_sin_ficha_de_campo_el_informe_lo_declara(self):
        informe = ate.indicadores_bloque(self.datos)
        self.assertTrue(any("Sin ficha DT de campo" in a
                            for a in informe["avisos"]))
        self.assertNotIn("Verificación de campo",
                         [s["titulo"] for s in informe["secciones"]])

    def test_con_ficha_de_campo_suma_la_seccion_de_verificacion(self):
        informe = ate.indicadores_bloque(self.datos, self.integrado)
        self.assertIn("Verificación de campo",
                      [s["titulo"] for s in informe["secciones"]])

    def test_las_cifras_de_cabecera_no_inventan_valores(self):
        informe = ate.indicadores_bloque(datos_bloque(area_ha_num=None),
                                         self.integrado)
        superficie = next(m for m in informe["metricas"]
                          if m["etiqueta"] == "Superficie de catálogo")
        self.assertEqual(superficie["valor"], "s/d")

    def test_bloque_vacio_no_rompe(self):
        informe = ate.indicadores_bloque({})
        self.assertEqual(informe["alcance"], "bloque")
        self.assertTrue(informe["avisos"])


# ══════════════════════════════════════════════════════════════════════════
# Consolidado
# ══════════════════════════════════════════════════════════════════════════

class TestConsolidado(unittest.TestCase):

    def setUp(self):
        self.lista = [
            datos_bloque("M9B1"),
            datos_bloque("M9B2", distrito="Frías", area_ha_num=200.0),
            datos_bloque("M9B3", provincia="Ayabaca", distrito="Frías"),
        ]

    def test_agrega_por_los_cuatro_niveles(self):
        for nivel in ate.AGRUPACIONES:
            informe = ate.indicadores_consolidado(self.lista, agrupacion=nivel)
            self.assertTrue(informe["secciones"])
            bloques = next(m for m in informe["metricas"]
                           if m["etiqueta"] == "Bloques en el ámbito")
            self.assertEqual(bloques["valor"], "3")

    def test_agrupacion_desconocida_falla(self):
        with self.assertRaises(ValueError):
            ate.indicadores_consolidado(self.lista, agrupacion="departamento")

    def test_la_superficie_agregada_cuadra_con_los_bloques(self):
        informe = ate.indicadores_consolidado(self.lista, agrupacion="distrito")
        serie = next(s for sec in informe["secciones"] for s in sec["series"]
                     if s["id"] == "dtc_area")
        total_serie = sum(f["valor"] for f in serie["filas"])
        total_bloques = sum(d["area_ha_num"] for d in self.lista)
        self.assertAlmostEqual(total_serie, total_bloques, places=2)

    def test_las_apiladas_declaran_su_dominio_ordenado(self):
        informe = ate.indicadores_consolidado(self.lista, agrupacion="provincia")
        apiladas = [s for sec in informe["secciones"] for s in sec["series"]
                    if s["forma"] == "apiladas"]
        self.assertTrue(apiladas)
        for serie in apiladas:
            self.assertTrue(serie["orden_sub"])
            self.assertEqual(set(serie["orden_sub"]),
                             {f["sub"] for f in serie["filas"]})

    def test_el_detalle_por_bloque_lista_todos_los_bloques(self):
        informe = ate.indicadores_consolidado(self.lista, agrupacion="provincia")
        detalle = next(s for s in informe["secciones"]
                       if s["titulo"] == "Detalle por bloque")
        self.assertEqual(len(detalle["tablas"][0][1]), 3)

    def test_ambito_vacio_no_rompe(self):
        informe = ate.indicadores_consolidado([])
        self.assertTrue(informe["avisos"])


# ══════════════════════════════════════════════════════════════════════════
# Renderizadores
# ══════════════════════════════════════════════════════════════════════════

class TestSalidas(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        datos = datos_bloque()
        cls.informe = ate.indicadores_bloque(
            datos, dtc.integrar_bloque(datos, ficha_campo()))
        cls.consolidado = ate.indicadores_consolidado(
            [datos_bloque("M9B1"), datos_bloque("M9B2", distrito="Frías")],
            etiqueta="2 bloques", agrupacion="distrito")

    def test_el_libro_trae_portada_y_una_hoja_por_seccion(self):
        wb = load_workbook(io.BytesIO(ate.generar_excel(self.informe)))
        self.assertEqual(wb.sheetnames[0], "Resumen DT")
        for seccion in self.informe["secciones"]:
            if seccion["series"]:
                self.assertTrue(
                    any(h.startswith(seccion["titulo"][:20])
                        for h in wb.sheetnames), seccion["titulo"])

    def test_el_libro_no_encima_sus_graficos(self):
        wb = load_workbook(io.BytesIO(ate.generar_excel(self.informe)))
        for ws in wb.worksheets:
            self.assertEqual(graficos_encimados(ws), [], ws.title)

    def test_las_hojas_de_respaldo_quedan_filtrables(self):
        wb = load_workbook(io.BytesIO(ate.generar_excel(self.informe)))
        respaldo = [ws for ws in wb.worksheets if ws.title.startswith("T ")]
        self.assertTrue(respaldo)
        for ws in respaldo:
            self.assertIsNotNone(ws.auto_filter.ref, ws.title)
            self.assertIsNotNone(ws.freeze_panes, ws.title)

    def test_el_consolidado_tampoco_encima_sus_graficos(self):
        wb = load_workbook(io.BytesIO(ate.generar_excel(self.consolidado)))
        for ws in wb.worksheets:
            self.assertEqual(graficos_encimados(ws), [], ws.title)

    def test_pdf_de_bloque_y_consolidado(self):
        self.assertTrue(ate.generar_pdf(self.informe).startswith(b"%PDF"))
        self.assertTrue(ate.generar_pdf(self.consolidado).startswith(b"%PDF"))

    def test_pdf_de_un_informe_sin_series(self):
        self.assertTrue(ate.generar_pdf(ate.indicadores_bloque({}))
                        .startswith(b"%PDF"))

    def test_toda_serie_se_dibuja_en_los_dos_temas(self):
        for informe in (self.informe, self.consolidado):
            for seccion in informe["secciones"]:
                for serie in seccion["series"]:
                    for tema in ("claro", "oscuro"):
                        self.assertIsNotNone(
                            ase.grafico_altair(serie, tema=tema),
                            f"{serie['id']} / {tema}")
                    self.assertFalse(ase.tabla_serie(serie).empty, serie["id"])

    def test_los_nombres_de_archivo_distinguen_alcance_y_nivel(self):
        self.assertIn("Bloque_M9B1", ate.nombre_excel(self.informe))
        self.assertIn("por_Distrito", ate.nombre_excel(self.consolidado))
        self.assertTrue(ate.nombre_pdf(self.informe).endswith(".pdf"))


if __name__ == "__main__":
    unittest.main()
