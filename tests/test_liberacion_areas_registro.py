# -*- coding: utf-8 -*-
"""
Pruebas de las tres vías de registro de Liberación de Áreas (Kobo, digitación en el aplicativo y plantilla Excel)
y de la edición / eliminación / restauración con bitácora. Datos SINTÉTICOS.
- Lógica pura: `python -m pytest tests/test_liberacion_areas_registro.py -q`
- Integración PostgreSQL: LA_TEST_DB="postgresql://…" (base VACÍA de pruebas: se borra el esquema public).
"""
import io
import os
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "tests"))
from liberacion_areas import la_campos as lc  # noqa: E402
from liberacion_areas import la_core as core  # noqa: E402
from liberacion_areas import la_plantillas as lp  # noqa: E402
import test_liberacion_areas as base  # noqa: E402

E0, N0 = base.E0, base.N0


def _cat():
    return base.catalogo()


# ------------------------------------------------------------------ digitación en el aplicativo
def test_registro_app_equivale_a_kobo():
    """El mismo predio digitado en UTM produce el mismo resultado de validación que el envío Kobo."""
    cat = _cat()
    kobo = base.casos()["f_la_03_inspeccion"][0]
    rk = core.validar_envio("f_la_03_inspeccion", core.aplanar(kobo), cat, base.GEOMS, set())
    reg, avisos = lc.construir_registro(
        "f_la_03_inspeccion",
        {"hoy": "2026-10-05", "asistente": "AP-1", "unidad": "27", "n_predio": 1, "fecha_insp": "2026-10-05",
         "punto_interior_este": E0 + 400, "punto_interior_norte": N0 + 400, "punto_interior_prec": 4,
         "interferencias": ["ninguna"], "conclusion_campo": "Sin restricción"},
        cat, {"r_puntos": [{"pt_gps_este": E0 + 450, "pt_gps_norte": N0 + 300, "pt_tipo": "Vértice"}]},
        kobo_uuid=lc.nuevo_uuid())
    assert not avisos, avisos
    ra = core.validar_envio("f_la_03_inspeccion", reg, cat, base.GEOMS, set())
    assert ra.estado_import == rk.estado_import == "NUEVO"
    assert ra.cod_predio == rk.cod_predio == "27-P01"
    assert ra.datos["validacion_espacial"] == rk.datos["validacion_espacial"] == "DENTRO"
    assert abs(ra.datos["este"] - rk.datos["este"]) < 0.05 and abs(ra.datos["norte"] - rk.datos["norte"]) < 0.05
    assert reg["conclusion_campo"] == "sin_restriccion" and reg["r_puntos"][0]["pt_tipo"] == "vertice"
    assert ra.kobo_uuid.startswith("app-")


def test_registro_app_sus_y_avisos():
    cat = _cat()
    lado = 123
    verts = [(E0 + 5050, N0 + 50), (E0 + 5050 + lado, N0 + 50), (E0 + 5050 + lado, N0 + 50 + lado), (E0 + 5050, N0 + 50 + lado)]
    reg, avisos = lc.construir_registro(
        "f_la_03_inspeccion",
        {"hoy": "2026-10-06", "asistente": "AP-6", "unidad": "SUS-059", "n_predio": 1, "fecha_insp": "2026-10-06",
         "punto_interior_este": E0 + 5100, "punto_interior_norte": N0 + 100, "conclusion_campo": "sin_restriccion"},
        cat, {"poligono_sus": [{"orden": i, "pt_gps_este": e, "pt_gps_norte": n} for i, (e, n) in enumerate(verts, 1)]},
        kobo_uuid=lc.nuevo_uuid())
    assert reg["tipo_unidad"] == "lote_sus" and reg["bloque_ref"] == "M20B1"
    r = core.validar_envio("f_la_03_inspeccion", reg, cat, base.GEOMS, set(), core.sus_areas_desde_catalogo(cat))
    assert abs(r.datos["area_sus_ha"] - 1.513) < 0.01 and r.datos["posicion_sus_calc"] == "DENTRO"
    # UTM fuera del ámbito y obligatorios faltantes
    _, av = lc.construir_registro("f_la_03_inspeccion", {"unidad": "27", "punto_interior_este": 100, "punto_interior_norte": 5},
                                  cat)
    assert any("fuera del rango UTM" in a for a in av) and any("Falta Fecha de inspección" in a for a in av)


def test_edicion_ida_y_vuelta():
    cat = _cat()
    for fid, regs in base.casos().items():
        for k in regs:
            plano = core.aplanar(k)
            val, reps = lc.registro_a_valores(fid, plano)
            nuevo, _ = lc.construir_registro(fid, val, cat, reps, kobo_uuid="x", base=plano)
            a = core.validar_envio(fid, plano, cat, base.GEOMS, set(), core.sus_areas_desde_catalogo(cat))
            b = core.validar_envio(fid, nuevo, cat, base.GEOMS, set(), core.sus_areas_desde_catalogo(cat))
            assert a.estado_import == b.estado_import, (fid, a.motivos, b.motivos)
            assert a.cod_predio == b.cod_predio


# ------------------------------------------------------------------ plantilla Excel
def _llenar(ws, filas: list[dict]):
    claves = [ws.cell(lp.FILA_CLAVE, j).value for j in range(1, ws.max_column + 1)]
    for i, f in enumerate(filas, lp.FILA_DATOS):
        for k, v in f.items():
            ws.cell(i, claves.index(k) + 1, v)


def _plantilla_llenada() -> bytes:
    cat = _cat()
    wb = load_workbook(io.BytesIO(lp.generar_plantilla_excel(cat)))
    assert {"Instrucciones", "F-LA-01", "F-LA-02", "F-LA-03", "F-LA-03_puntos", "F-LA-03_vertices_SUS", "F-LA-04",
            "F-LA-04_firmantes", "F-LA-05", "F-LA-06", "Catalogo_unidades", "_Listas"} <= set(wb.sheetnames)
    _llenar(wb["F-LA-02"], [{"id_registro": 1, "hoy": "05/10/2026", "asistente": "AP-1", "unidad": "27", "n_predio": 3,
                             "consentimiento": "Sí", "tipo_titularidad": "Privado (con título)", "tit_nombres": "LUIS (FICTICIO)",
                             "tit_apellidos": "PRUEBA TRES", "tit_dni": 123, "docs_exhibidos__dni": "SI",
                             "docs_exhibidos__titulo": "si", "predio_area_unidad_ha": 12.5, "aceptacion": "Favorable"}])
    _llenar(wb["F-LA-03"], [{"id_registro": 1, "hoy": "06/10/2026", "asistente": "AP-1", "unidad": 27, "n_predio": 3,
                             "fecha_insp": "06/10/2026", "punto_interior_este": E0 + 200, "punto_interior_norte": N0 + 200,
                             "punto_interior_prec": 3, "interferencias__ninguna": "SI", "conclusion_campo": "Sin restricción"}])
    _llenar(wb["F-LA-03_puntos"], [{"id_registro": 1, "pt_gps_este": E0 + 210, "pt_gps_norte": N0 + 210, "pt_tipo": "Hito"},
                                   {"id_registro": 1, "pt_gps_este": E0 + 220, "pt_gps_norte": N0 + 220, "pt_tipo": "Vértice"}])
    _llenar(wb["F-LA-04"], [{"id_registro": 1, "hoy": "07/10/2026", "asistente": "AP-1", "unidad": "27", "n_predio": 3,
                             "tipo_acta": "A-04 · Acta del titular / posesionario", "fecha_acta": "07/10/2026",
                             "area_comprometida_ha": 12.5, **{f"checklist__{c}": "SI" for c, _ in lc.OPCIONES["checklist"]},
                             "fedatario_tipo": "Juez de paz", "estado_acta": "Completa"}])
    _llenar(wb["F-LA-04_firmantes"], [{"id_registro": 1, "fir_nombre": "LUIS (FICTICIO) PRUEBA TRES", "fir_dni": "00000123",
                                       "fir_rol": "Titular", "fir_huella": "Sí"}])
    _llenar(wb["F-LA-05"], [{"id_registro": 1, "unidad": "27", "n_predio": 3, "tipo": "Constancia de búsqueda",
                             "entidad": "SUNARP", "fecha": "02/10/2026", "resultado": "Negativo", "registrado_por": "pytest"}])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_plantilla_excel_lectura():
    cat = _cat()
    b = _plantilla_llenada()
    r = lp.leer_plantilla_excel(io.BytesIO(b), cat)
    assert not r["errores"], r["errores"]
    f2 = r["envios"]["f_la_02_titular"][0]["registro"]
    assert f2["cod_predio"] == "27-P03" and f2["tit_dni"] == "00000123" and f2["consentimiento"] == "si"
    assert f2["docs_exhibidos"] == "titulo dni" and f2["tipo_titularidad"] == "privado" and f2["hoy"] == "2026-10-05"
    f3 = r["envios"]["f_la_03_inspeccion"][0]["registro"]
    assert len(f3["r_puntos"]) == 2 and f3["r_puntos"][0]["pt_tipo"] == "hito"
    v3 = core.validar_envio("f_la_03_inspeccion", f3, cat, base.GEOMS, set())
    assert v3.estado_import == "NUEVO" and v3.datos["validacion_espacial"] == "DENTRO", v3.motivos
    f4 = r["envios"]["f_la_04_actas"][0]["registro"]
    v4 = core.validar_envio("f_la_04_actas", f4, cat, base.GEOMS, set())
    assert v4.estado_import == "NUEVO", v4.motivos
    assert f4["cod_doc"] == "A04-27-P03" and f4["r_firmantes"][0]["fir_rol"] == "titular"
    doc = r["documentos"][0]["datos"]
    assert doc["cod_predio"] == "27-P03" and doc["tipo"] == "constancia_busqueda" and doc["resultado"] == "negativo"
    # determinista: leer dos veces el mismo archivo da los mismos identificadores
    r2 = lp.leer_plantilla_excel(io.BytesIO(b), cat)
    assert [x["registro"]["meta_instanceID"] for x in r2["envios"]["f_la_03_inspeccion"]] == \
           [x["registro"]["meta_instanceID"] for x in r["envios"]["f_la_03_inspeccion"]]


def test_plantilla_por_asistente_y_errores():
    cat = _cat()
    wb = load_workbook(io.BytesIO(lp.generar_plantilla_excel(cat, asistente="AP-6")))
    unidades = [c.value for c in wb["_Listas"]["A"]]
    col = [c for c in wb["_Listas"][1] if c.value == "unidad"][0].column
    lista = [wb["_Listas"].cell(i, col).value for i in range(2, wb["_Listas"].max_row + 1)]
    lista = [x for x in lista if x]
    assert lista and all(cat[x]["asistente"] == "AP-6" for x in lista), unidades
    _llenar(wb["F-LA-03_puntos"], [{"id_registro": 9, "pt_gps_este": E0, "pt_gps_norte": N0}])
    buf = io.BytesIO()
    wb.save(buf)
    r = lp.leer_plantilla_excel(io.BytesIO(buf.getvalue()), cat)
    assert any("sin fila en F-LA-03" in e for e in r["errores"])


def test_xlsform_kobo():
    cat = _cat()
    for fid, f in lc.FORMULARIOS.items():
        if not f.kobo:
            continue
        h = lp.xlsform_hojas(fid)
        nombres = set(h["survey"]["name"])
        # los nombres que lee el importador existen en el XLSForm
        assert {"hoy", "asistente", "cod_unidad"} <= nombres, fid
        if f.con_predio:
            assert {"unidad", "n_predio", "cod_predio", "tipo_unidad"} <= nombres
    assert "poligono_sus" in set(lp.xlsform_hojas("f_la_03_inspeccion")["survey"]["name"])
    assert len(pd.read_csv(io.BytesIO(lp.unidades_csv(cat)))) == len(cat)
    try:
        from pyxform.xls2xform import convert
    except ImportError:
        return
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        for fid, f in lc.FORMULARIOS.items():
            if f.kobo:
                p = Path(d) / f"{fid}.xlsx"
                p.write_bytes(lp.generar_xlsform(fid))
                assert "<h:html" in convert(xlsform=str(p), validate=False).xform


# ------------------------------------------------------------------ integración: registrar, editar, eliminar, restaurar
def test_integracion_edicion_eliminacion():
    url = os.environ.get("LA_TEST_DB")
    if not url:
        return
    base.test_integracion_postgres()              # esquema limpio + catálogo + envíos Kobo de prueba
    from liberacion_areas import la_db as db
    from liberacion_areas import la_reportes as rep
    conn = db.conectar(url)
    with conn.cursor() as cur:
        for cod, g in base.GEOMS.items():
            cur.execute("UPDATE la_unidades SET geom_wkt=%s WHERE codigo=%s", (g.wkt, cod))
    conn.commit()
    cat, geoms = db.catalogo(conn), db.geometrias(conn)

    # vía 3: plantilla Excel
    r = lp.leer_plantilla_excel(io.BytesIO(_plantilla_llenada()), cat)
    for fid in ("f_la_02_titular", "f_la_03_inspeccion", "f_la_04_actas"):
        res = [core.validar_envio(fid, x["registro"], cat, geoms, db.uuids_existentes(conn), db.sus_areas(conn, cat))
               for x in r["envios"][fid]]
        db.importar(conn, res, "PLANTILLA", fid, "pytest")
    for d in r["documentos"]:
        assert db.documento_existe(conn, d["datos"]) is None
        db.registrar_documento(conn, d["datos"])
        assert db.documento_existe(conn, d["datos"])
    m = db.df(conn, "SELECT estado_la FROM la_predios WHERE cod_predio='27-P03'")
    assert m["estado_la"][0] == "LA-4", m          # titular + reunión (27) + inspección dentro + acta conforme
    assert set(db.listar_envios(conn)["origen"]) == {"KOBO", "PLANTILLA"}
    # re-importar el mismo archivo → DUPLICADO
    res = [core.validar_envio("f_la_03_inspeccion", x["registro"], cat, geoms, db.uuids_existentes(conn))
           for x in r["envios"]["f_la_03_inspeccion"]]
    assert all(x.estado_import == "DUPLICADO" for x in res)

    # vía 2: digitación en el aplicativo
    reg, _ = lc.construir_registro("f_la_02_titular", {"hoy": "2026-10-08", "asistente": "AP-1", "unidad": "27", "n_predio": 4,
                                                       "consentimiento": "si", "tipo_titularidad": "posesionario",
                                                       "tit_nombres": "ROSA (FICTICIA)", "tit_apellidos": "PRUEBA",
                                                       "tit_dni": "00000444", "aceptacion": "favorable"}, cat,
                                   kobo_uuid=lc.nuevo_uuid())
    rr = core.validar_envio("f_la_02_titular", reg, cat, geoms, db.uuids_existentes(conn))
    db.importar(conn, [rr], "APP", "f_la_02_titular", "pytest")
    uid = rr.kobo_uuid
    assert db.obtener_envio(conn, uid)["origen"] == "APP"

    # editar: cambia el nombre y la aceptación (el titular se corrige, no se duplica)
    env = db.obtener_envio(conn, uid)
    val, reps = lc.registro_a_valores("f_la_02_titular", env["payload"])
    val.update({"tit_nombres": "ROSA MARÍA (FICTICIA)", "aceptacion": "rechazo"})
    nuevo, _ = lc.construir_registro("f_la_02_titular", val, cat, reps, kobo_uuid=uid, base=env["payload"])
    rr2 = core.validar_envio("f_la_02_titular", nuevo, cat, geoms, db.uuids_existentes(conn) - {uid})
    db.actualizar_envio(conn, rr2, "f_la_02_titular", "editor")
    t = db.df(conn, "SELECT nombre FROM la_titulares WHERE dni_ruc='00000444'")
    assert len(t) == 1 and t["nombre"][0].startswith("ROSA MARÍA")
    assert db.df(conn, "SELECT estado_la FROM la_predios WHERE cod_predio='27-P04'")["estado_la"][0] == "NEG"
    assert db.obtener_envio(conn, uid)["editado_por"] == "editor"

    # eliminar una inspección con foto → estado baja; restaurar → vuelve y la foto se reenlaza
    insp = db.df(conn, "SELECT kobo_uuid FROM la_inspecciones WHERE cod_predio='27-P01'")["kobo_uuid"][0]
    assert len(db.adjuntos_envio(conn, insp)) == 1
    db.eliminar_envio(conn, insp, "editor")
    assert db.obtener_envio(conn, insp) is None
    assert db.df(conn, "SELECT estado_la FROM la_predios WHERE cod_predio='27-P01'")["estado_la"][0] == "LA-2"
    assert db.df(conn, "SELECT count(*) n FROM la_adjuntos")["n"][0] == 5          # la foto no se borra
    b = db.bitacora(conn, solo_eliminados=True)
    db.restaurar(conn, int(b[b.clave == insp]["id"].iloc[0]), "editor")
    assert db.obtener_envio(conn, insp) is not None and len(db.adjuntos_envio(conn, insp)) == 1
    assert db.df(conn, "SELECT estado_la FROM la_predios WHERE cod_predio='27-P01'")["estado_la"][0] == "LA-4"

    # documentos: editar, eliminar y restaurar sin reutilizar el código
    doc = db.df(conn, "SELECT * FROM la_documentos").to_dict("records")[0]
    db.actualizar_documento(conn, doc["id"], {**doc, "resultado": "positivo", "n_partida": "P-001"}, "editor")
    assert db.obtener_documento(conn, doc["id"])["resultado"] == "positivo"
    db.eliminar_documento(conn, doc["id"], "editor")
    otro = db.registrar_documento(conn, {**doc, "cod_doc": None})
    assert otro != doc["cod_doc"]
    b = db.bitacora(conn, solo_eliminados=True)
    db.restaurar(conn, int(b[b.clave == doc["cod_doc"]]["id"].iloc[0]), "editor")
    assert doc["cod_doc"] in set(db.df(conn, "SELECT cod_doc FROM la_documentos")["cod_doc"])

    # predios: editar núcleo / área y eliminar solo si no tiene registros
    assert db.actualizar_predios(conn, [{"cod_predio": "27-P02", "nucleo": True, "area_unidad_ha": 30}], "editor") == 1
    assert db.actualizar_predios(conn, [{"cod_predio": "27-P02", "nucleo": True, "area_unidad_ha": 30.0}], "editor") == 0
    try:
        db.eliminar_predio(conn, "27-P02", "editor")
        raise AssertionError("debió impedir la eliminación")
    except ValueError:
        pass

    # reportes
    xls = rep.excel_consolidado(conn)
    wb = load_workbook(io.BytesIO(xls))
    assert {"Resumen", "Matriz_predial", "F-LA-01", "F-LA-05", "Titulares", "Observados"} <= set(wb.sheetnames)
    hdr = [c.value for c in wb["Titulares"][7]]
    assert "dni_ruc" not in hdr and "celular" not in hdr
    assert rep.pdf_expediente_unidad(conn, "27")[:4] == b"%PDF"
    assert load_workbook(io.BytesIO(rep.excel_consolidado(conn, asistente="AP-1")))
    # la re-inicialización (arranque del aplicativo) no borra nada
    antes = db.df(conn, "SELECT count(*) n FROM la_envios_raw")["n"][0]
    db.inicializar_la(conn, base.CSV)
    assert db.df(conn, "SELECT count(*) n FROM la_envios_raw")["n"][0] == antes
    conn.close()
