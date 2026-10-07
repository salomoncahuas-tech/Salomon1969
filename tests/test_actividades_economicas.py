"""Pruebas de las actividades economicas de la F-DS-01 en % de las familias.

Reproducen el reporte observado (octubre 2026): la tabla 6 de la ficha
mezclaba el N.° de familias ("240") con porcentajes ("100 %") y el grafico
los sumaba como familias (935 familias "mayormente autoconsumo" en caserios
de unas decenas de viviendas). La regla vigente:

  - todo se expresa en % (nunca mas de 100 %) de las familias / viviendas
    del CP: la moda entre sus fichas (numeral 2) o, sin ella, la mediana; a
    falta del dato, las viviendas INEI del catalogo, una estimacion con la
    poblacion declarada y, como ultimo recurso, el mayor N.° de familias
    declarado entre sus actividades;
  - el ambito pondera cada CP por esas familias / viviendas;
  - la tabla de actividades y el Excel de la base de datos llevan el distrito.
"""

import io
import json
import os
import sys
import unittest

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import analitica_social as an  # noqa: E402
import export_diagnosticos as ex  # noqa: E402


_ID = [0]
AUTO = "Mayormente autoconsumo (>70%)"
MIXTO = "Mixto autoconsumo/mercado (30-70%)"
MERCADO = "100% Mercado"


def _reg(form, cp, bloque="M7B2", distrito="Salitral", **extra):
    _ID[0] += 1
    registro = {
        "id": extra.pop("id", _ID[0]), "ficha": "F-DS-01",
        "bloque_codigo": bloque, "centro_poblado": cp,
        "comunidad_campesina": "Ninguna", "distrito": distrito,
        "provincia": "Morropon",
        "fecha_evaluacion": extra.pop("fecha_evaluacion", "2026-08-04"),
        "evaluador": extra.pop("evaluador", "Stefany Campos"),
        "ds01_data_v3": json.dumps(form, ensure_ascii=False),
    }
    registro.update(extra)
    return registro


def _activ(*filas):
    return [{"Actividad / Rubro": a, "N fam.": n, "Destino": d,
             "Productos principales": "", "Ingreso (S/./mes)": ""}
            for a, n, d in filas]


def _unidad(registros, etiqueta=None):
    unidades = an.unidades_por_cp(registros, "F-DS-01")
    if etiqueta is None:
        return unidades[0]
    return next(u for u in unidades if u["etiqueta"] == etiqueta)


def _serie(registros):
    seccion = an._seccion_socioeconomica(registros)
    return next(s for s in seccion["series"] if s["id"] == "f1_actividades"), seccion


def _matriz(serie):
    categorias, subclases, matriz = an._pivote(serie)
    return {c: dict(zip(subclases, fila)) for c, fila in zip(categorias, matriz)}


def _detalle(seccion):
    return dict(seccion["tablas"])["Actividades económicas"]


def _textos_control(seccion):
    return [o["Detalle"] for o in seccion["control"]]


# ══════════════════════════════════════════════════════════════════════════
# Lectura de la casilla "N fam."
# ══════════════════════════════════════════════════════════════════════════

class ValorDeLaCasilla(unittest.TestCase):

    def test_el_signo_porcentaje_manda(self):
        self.assertEqual(an._valor_actividad("40 %"), ("pct", 40.0))
        self.assertEqual(an._valor_actividad("40%"), ("pct", 40.0))
        self.assertEqual(an._valor_actividad("40"), ("n", 40.0))

    def test_numero_de_familias_en_texto_libre(self):
        self.assertEqual(an._valor_actividad("1,250"), ("n", 1250.0))
        self.assertEqual(an._valor_actividad("aprox. 30 familias"), ("n", 30.0))
        self.assertEqual(an._valor_actividad(25), ("n", 25.0))

    def test_fraccion_de_la_plantilla_excel_es_porcentaje(self):
        """La plantilla guarda "40%" como 0.4: no son 0.4 familias."""
        self.assertEqual(an._valor_actividad("0.4"), ("pct", 40.0))
        self.assertEqual(an._valor_actividad(0.25), ("pct", 25.0))
        self.assertEqual(an._valor_actividad("1"), ("n", 1.0))
        self.assertEqual(an._valor_actividad("0"), ("n", 0.0))

    def test_sin_cifra_o_negativo_no_es_dato(self):
        for texto in ("", None, "muchas", "-5"):
            self.assertIsNone(an._valor_actividad(texto), texto)


# ══════════════════════════════════════════════════════════════════════════
# Familias / viviendas de referencia del CP
# ══════════════════════════════════════════════════════════════════════════

class ReferenciaDelCP(unittest.TestCase):

    def _fichas(self, *nfams, **form):
        return [_reg(dict(form, f1_nfam=n, f1_pob_t="300"), "La Peña",
                     fecha_evaluacion=f"2026-08-{10 + i:02d}")
                for i, n in enumerate(nfams)]

    def test_moda_entre_las_fichas(self):
        u = _unidad(self._fichas("50", "80", "50"))
        valor, criterio, origen = an._referencia_actividades(u)
        self.assertEqual((valor, origen), (50.0, "ficha"))
        self.assertIn("moda (2 de 3 fichas)", criterio)

    def test_fichas_que_coinciden(self):
        u = _unidad(self._fichas("50", "50"))
        self.assertIn("las 2 fichas coinciden", an._referencia_actividades(u)[1])

    def test_sin_moda_se_usa_la_mediana(self):
        """La ficha atipica (120, la comunidad entera) no arrastra la base."""
        u = _unidad(self._fichas("40", "120", "50"))
        valor, criterio, _origen = an._referencia_actividades(u)
        self.assertEqual(valor, 50.0)
        self.assertIn("mediana de 3 fichas", criterio)

    def test_mediana_de_dos_fichas_es_una_cifra_declarada(self):
        """40 (antigua) y 60 (reciente) no se repiten: de los dos valores
        centrales manda la ficha mas reciente. Nunca un promedio (50) que no
        declara nadie, y la misma cifra en actividades, demografia y control."""
        regs = self._fichas("40", "60")
        u = _unidad(regs)
        self.assertEqual(an._referencia_actividades(u)[0], 60.0)
        self.assertEqual(u["form"]["f1_nfam"], "60")
        seccion = an._seccion_socioeconomica(regs)
        demografia = dict(seccion["tablas"])["Demografía por centro poblado"][0]
        self.assertEqual(demografia["Familias / viviendas"], 60.0)
        self.assertIn("mediana", demografia["Criterio (familias / viviendas)"])
        texto = next(t for t in _textos_control(seccion)
                     if t.startswith("Familias / viviendas"))
        self.assertIn("se usa la mediana, 60 (de los dos valores centrales, "
                      "el de la ficha más reciente)", texto)
        self.assertIn("declaran 40 (1)", texto)

    def test_mediana_no_inventa_un_conflicto_con_juntos(self):
        """Antes: 20 y 25 daban 22.5 familias y 'JUNTOS: 23 supera 22.5'."""
        regs = self._fichas("20", "25", f1_juntos="23")
        seccion = an._seccion_socioeconomica(regs)
        self.assertFalse(any(t.startswith("JUNTOS") for t in _textos_control(seccion)))

    def test_mediana_conserva_la_deteccion_de_fichas_duplicadas(self):
        """La ficha reciente de La Peña registrada como El Pino se detecta."""
        regs = self._fichas("20", "25") + [
            _reg({"f1_nfam": "25", "f1_pob_t": "300"}, "El Pino")]
        avisos = an._control_duplicados_probables(an.unidades_por_cp(regs))
        self.assertEqual(len(avisos), 2)

    def test_control_cuenta_las_fichas_como_el_criterio(self):
        """'20 familias' y '20' son la misma cifra: 2 de 3, no 1 de 3."""
        seccion = an._seccion_socioeconomica(self._fichas("35", "20", "20 familias"))
        texto = next(t for t in _textos_control(seccion)
                     if t.startswith("Familias / viviendas"))
        self.assertIn("se usa 20 (2 de 3 fichas con dato)", texto)
        self.assertIn("declaran 35 (1)", texto)
        coinciden = an._seccion_socioeconomica(self._fichas("20", "20 viviendas"))
        self.assertFalse(any(t.startswith("Familias / viviendas")
                             for t in _textos_control(coinciden)))

    def test_un_cero_no_es_un_dato(self):
        u = _unidad(self._fichas("0", "35"))
        valor, criterio, _origen = an._referencia_actividades(u)
        self.assertEqual(valor, 35.0)
        self.assertIn("única ficha con dato", criterio)
        seccion = an._seccion_socioeconomica(self._fichas("0", "0", "35"))
        self.assertFalse(any(t.startswith("Familias / viviendas")
                             for t in _textos_control(seccion)))

    def test_poblacion_negativa_no_es_referencia(self):
        """'+/- 240' en una ficha antigua se lee -240: no sirve de base."""
        regs = [_reg({"f1_pob_t": "+/- 240", "f1_activ": _activ(
            ("Agricultura de secano", "20", AUTO))}, "Caserio Nuevo", bloque="1")]
        resultado = an._actividades_cp(_unidad(regs))
        self.assertEqual(resultado["origen"], "actividades")
        self.assertEqual(resultado["por_clave"][("Agricultura de secano", AUTO)], 100.0)

    def test_la_moda_conserva_el_texto_de_la_ficha(self):
        u = _unidad(self._fichas("1.500", "1500", "900"))
        self.assertEqual(an._entero(u["form"]["f1_nfam"]), 1500.0)

    def test_sin_dato_en_la_ficha_usa_las_viviendas_inei(self):
        """Mamayaco (bloque 1): 75 viviendas en el catalogo INEI."""
        regs = [_reg({"f1_pob_t": "204", "f1_activ": _activ(
            ("Agricultura de secano", "30", AUTO))}, "Mamayaco", bloque="1",
            distrito="Ayabaca")]
        u = _unidad(regs)
        valor, criterio, origen = an._referencia_actividades(u)
        self.assertEqual((valor, origen), (75.0, "inei"))
        self.assertIn("catálogo INEI", criterio)
        resultado = an._actividades_cp(u)
        self.assertAlmostEqual(resultado["por_clave"][("Agricultura de secano", AUTO)], 40.0)
        self.assertTrue(any("no consigna el N.° de familias" in t
                            for t in resultado["observaciones"]))

    def test_solo_porcentajes_no_dice_que_se_divide_entre_el_inei(self):
        regs = [_reg({"f1_activ": _activ(("Agricultura de secano", "40%", AUTO))},
                     "Mamayaco", bloque="1", distrito="Ayabaca")]
        resultado = an._actividades_cp(_unidad(regs))
        self.assertEqual(resultado["origen"], "inei")
        self.assertEqual(resultado["observaciones"], [])

    def test_sin_dato_ni_catalogo_estima_con_la_poblacion(self):
        """Bloque 1: 238 hab. en 95 viviendas INEI (2.51 hab. por vivienda)."""
        regs = [_reg({"f1_pob_t": "100", "f1_activ": _activ(
            ("Agricultura de secano", "20", AUTO))}, "Caserio Nuevo", bloque="1")]
        valor, criterio, origen = an._referencia_actividades(_unidad(regs))
        self.assertEqual(origen, "estimada")
        self.assertAlmostEqual(valor, 100 / (238 / 95))
        self.assertIn("2.51 hab. por vivienda", criterio)

    def test_ultimo_recurso_el_mayor_numero_de_familias(self):
        regs = [_reg({"f1_activ": _activ(("Agricultura de secano", "40", AUTO),
                                         ("Ganadería vacuna", "10", MIXTO))},
                     "Sin Nombre INEI", bloque="ZZ9")]
        resultado = an._actividades_cp(_unidad(regs))
        self.assertEqual((resultado["referencia"], resultado["origen"]),
                         (40.0, "actividades"))
        self.assertEqual(resultado["por_clave"][("Agricultura de secano", AUTO)], 100.0)
        self.assertEqual(resultado["por_clave"][("Ganadería vacuna", MIXTO)], 25.0)

    def test_sin_ninguna_referencia_se_usan_los_porcentajes(self):
        regs = [_reg({"f1_activ": _activ(("Apicultura", "15 %", MERCADO))},
                     "Sin Nombre INEI", bloque="ZZ9")]
        resultado = an._actividades_cp(_unidad(regs))
        self.assertIsNone(resultado["referencia"])
        self.assertEqual(resultado["por_clave"][("Apicultura", MERCADO)], 15.0)
        self.assertIn("porcentajes declarados", resultado["criterio"])


# ══════════════════════════════════════════════════════════════════════════
# Conversion a porcentaje dentro del CP
# ══════════════════════════════════════════════════════════════════════════

class ConversionAPorcentaje(unittest.TestCase):

    def _cp(self, *filas, nfam="50"):
        return an._actividades_cp(_unidad([_reg(
            {"f1_nfam": nfam, "f1_activ": _activ(*filas)}, "La Peña")]))

    def test_numeros_y_porcentajes_quedan_en_la_misma_escala(self):
        r = self._cp(("Agricultura de secano", "40%", AUTO),
                     ("Ganadería vacuna", "25", MIXTO))
        self.assertEqual(r["por_clave"][("Agricultura de secano", AUTO)], 40.0)
        self.assertEqual(r["por_clave"][("Ganadería vacuna", MIXTO)], 50.0)
        self.assertEqual(r["observaciones"], [])

    def test_familias_por_encima_de_las_del_cp_se_recortan(self):
        """El caso del reporte: 240 familias en un caserio de 20 viviendas."""
        r = self._cp(("Agricultura bajo riego", "240", AUTO), nfam="20")
        self.assertEqual(r["por_clave"][("Agricultura bajo riego", AUTO)], 100.0)
        self.assertTrue(any("240 familias superan las 20" in t
                            for t in r["observaciones"]))

    def test_porcentaje_mayor_a_cien_se_recorta(self):
        r = self._cp(("Apicultura", "150 %", MERCADO))
        self.assertEqual(r["por_clave"][("Apicultura", MERCADO)], 100.0)
        self.assertTrue(any("supera el 100 %" in t for t in r["observaciones"]))

    def test_los_destinos_de_una_actividad_no_pasan_de_cien(self):
        r = self._cp(("Agricultura de secano", "80%", AUTO),
                     ("Agricultura de secano", "60%", MERCADO))
        auto = r["por_clave"][("Agricultura de secano", AUTO)]
        mercado = r["por_clave"][("Agricultura de secano", MERCADO)]
        self.assertAlmostEqual(auto + mercado, 100.0)
        self.assertAlmostEqual(auto, 80 * 100 / 140)
        self.assertTrue(any("se reparten en proporción" in t
                            for t in r["observaciones"]))

    def test_fila_repetida_toma_el_mayor_no_la_suma(self):
        r = self._cp(("Ganadería porcina", "10", MIXTO),
                     ("Ganadería porcina", "20", MIXTO))
        self.assertEqual(r["por_clave"][("Ganadería porcina", MIXTO)], 40.0)

    def test_fraccion_se_observa(self):
        r = self._cp(("Ganadería vacuna", "0.4", MIXTO))
        self.assertEqual(r["por_clave"][("Ganadería vacuna", MIXTO)], 40.0)
        self.assertTrue(any("«0.4» se lee como 40 %" in t for t in r["observaciones"]))

    def test_texto_sin_cifra_se_observa(self):
        r = self._cp(("Comercio local", "varias", MERCADO))
        self.assertEqual(r["por_clave"], {})
        self.assertTrue(any("«varias» no es un N.° de familias" in t
                            for t in r["observaciones"]))

    def test_varias_fichas_del_cp_no_se_suman(self):
        """La actividad sale de una sola ficha (la de referencia) y la base
        es la moda de las viviendas, no la suma de las fichas."""
        regs = [_reg({"f1_nfam": "60", "f1_pob_t": "200",
                      "f1_activ": _activ(("Agricultura de secano", "30", AUTO))},
                     "La Peña", fecha_evaluacion="2026-08-10"),
                _reg({"f1_nfam": "60", "f1_pob_t": "200",
                      "f1_activ": _activ(("Agricultura de secano", "30", AUTO))},
                     "La Peña", fecha_evaluacion="2026-08-11",
                     evaluador="Stefany Campos A")]
        r = an._actividades_cp(_unidad(regs))
        self.assertEqual(r["referencia"], 60.0)
        self.assertEqual(r["por_clave"][("Agricultura de secano", AUTO)], 50.0)


# ══════════════════════════════════════════════════════════════════════════
# Grafico del ambito
# ══════════════════════════════════════════════════════════════════════════

def _reporte_con_errores():
    """Tres caserios con la tabla 6 llenada de las dos formas (como el
    reporte observado): numeros que exceden las viviendas, % y fichas sin
    actividades."""
    return [
        _reg({"f1_nfam": "20", "f1_activ": _activ(
            ("Agricultura bajo riego", "240", AUTO),
            ("Agricultura bajo riego", "100", MIXTO),
            ("Ganadería menor (aves, cuyes)", "100%", AUTO))},
            "Platanal", distrito="Salitral"),
        _reg({"f1_nfam": "80", "f1_activ": _activ(
            ("Agricultura bajo riego", "30%", AUTO),
            ("Ganadería menor (aves, cuyes)", "60", AUTO),
            ("Jornalero agrícola", "140", "Mayormente mercado (>70%)"))},
            "Serrán", distrito="Salitral"),
        _reg({"f1_nfam": "45", "f1_pob_t": "150"}, "Pampa Larga",
             distrito="San Juan de Bigote"),
    ]


class Ambito(unittest.TestCase):

    def test_promedio_ponderado_por_las_familias_de_cada_cp(self):
        regs = [_reg({"f1_nfam": "100", "f1_activ": _activ(
                    ("Agricultura de secano", "50", AUTO))}, "Grande"),
                _reg({"f1_nfam": "300", "f1_activ": _activ(
                    ("Agricultura de secano", "30", AUTO))}, "Chico")]
        serie, _seccion = _serie(regs)
        # (50 % x 100 + 10 % x 300) / 400 = 20 %
        self.assertAlmostEqual(_matriz(serie)["Agricultura de secano"][AUTO], 20.0)
        self.assertIn("400 familias", serie["descripcion"])

    def test_cp_sin_actividades_no_diluye_el_porcentaje(self):
        regs = [_reg({"f1_nfam": "100", "f1_activ": _activ(
                    ("Agricultura de secano", "50", AUTO))}, "Grande"),
                _reg({"f1_nfam": "900"}, "Sin Tabla")]
        serie, _seccion = _serie(regs)
        self.assertAlmostEqual(_matriz(serie)["Agricultura de secano"][AUTO], 50.0)
        self.assertIn("de 1 centro(s) poblado(s)", serie["descripcion"])

    def test_ningun_valor_ni_total_supera_el_cien(self):
        serie, _seccion = _serie(_reporte_con_errores())
        self.assertEqual(serie["unidad"], "%")
        for actividad, fila in _matriz(serie).items():
            for valor in fila.values():
                self.assertLessEqual(valor, 100.0, actividad)
            self.assertLessEqual(sum(fila.values()), 100.0 + 1e-9, actividad)
        tabla = an.tabla_serie(serie)
        self.assertLessEqual(tabla["Familias en la actividad (%)"].max(), 100.0)

    def test_reporte_observado_queda_en_porcentaje(self):
        """Platanal (20 viv.): 240 y 100 familias de riego eran el 100 %
        del CP, repartido entre sus dos destinos; Serran (80 viv.), 30 %."""
        serie, seccion = _serie(_reporte_con_errores())
        riego = _matriz(serie)["Agricultura bajo riego"]
        # Platanal: 100 % + 100 % -> 50 % y 50 %; Serran: 30 % autoconsumo.
        self.assertAlmostEqual(riego[AUTO], (50 * 20 + 30 * 80) / 100)
        self.assertAlmostEqual(riego[MIXTO], 50 * 20 / 100)
        jornal = _matriz(serie)["Jornalero agrícola"]["Mayormente mercado (>70%)"]
        self.assertAlmostEqual(jornal, 100 * 80 / 100)
        control = _textos_control(seccion)
        self.assertTrue(any("240 familias superan las 20" in t for t in control))
        self.assertTrue(any("140 familias superan las 80" in t for t in control))

    def test_cp_solo_con_porcentajes_pesa_como_el_promedio(self):
        regs = [_reg({"f1_nfam": "100", "f1_activ": _activ(
                    ("Apicultura", "10", MERCADO))}, "Con Base"),
                _reg({"f1_activ": _activ(("Apicultura", "50%", MERCADO))},
                     "Sin Nombre INEI", bloque="ZZ9")]
        serie, seccion = _serie(regs)
        self.assertAlmostEqual(_matriz(serie)["Apicultura"][MERCADO], 30.0)
        self.assertTrue(any("pesa como el CP promedio" in t
                            for t in _textos_control(seccion)))

    def test_serie_en_porcentaje_sin_total_por_columnas(self):
        serie, _seccion = _serie(_reporte_con_errores())
        self.assertEqual(an._modo_totales(serie), "filas")
        self.assertEqual(serie["maximo"], 100)
        self.assertIsNone(an._base_porcentaje(serie))

    def test_grafico_llega_al_cien_por_ciento(self):
        serie, _seccion = _serie(_reporte_con_errores())
        spec = an.grafico_altair(serie).to_dict()
        self.assertEqual(spec["encoding"]["x"]["scale"]["domain"], [0, 100.0])
        self.assertEqual(spec["encoding"]["x"]["axis"]["format"], ",.0f")
        self.assertEqual(spec["encoding"]["x"]["axis"]["values"],
                         list(range(0, 101, 10)))


# ══════════════════════════════════════════════════════════════════════════
# Tablas y libros Excel
# ══════════════════════════════════════════════════════════════════════════

class Salidas(unittest.TestCase):

    def test_tabla_de_actividades_lleva_distrito_y_porcentaje(self):
        _serie_act, seccion = _serie(_reporte_con_errores())
        filas = _detalle(seccion)
        self.assertEqual(list(filas[0])[:3],
                         ["Centro poblado / ámbito", "Distrito", "Bloque(s)"])
        platanal = [f for f in filas if f["Centro poblado / ámbito"] == "Platanal"]
        self.assertTrue(all(f["Distrito"] == "Salitral" for f in platanal))
        self.assertTrue(all(f["Familias del CP en la actividad (%)"] <= 100
                            for f in filas))
        riego = next(f for f in platanal if f["Familias declaradas (N.°)"] == 240)
        self.assertEqual(riego["Familias / viviendas de referencia del CP"], 20.0)
        self.assertEqual(riego["Familias del CP en la actividad (%)"], 50.0)
        aves = next(f for f in platanal if f["Familias declaradas (%)"] == 100)
        self.assertIsNone(aves["Familias declaradas (N.°)"])

    def test_excel_del_grafico_sin_fila_total(self):
        informe = an.indicadores_bloque(
            {"codigo": "M7B2"}, _reporte_con_errores(), {})
        wb = load_workbook(io.BytesIO(an.generar_excel_social(informe)))
        ws = next(ws for ws in wb.worksheets if ws.title.startswith("F-DS-01"))
        titulo = next(r for r in range(1, ws.max_row + 1)
                      if str(ws.cell(r, 1).value or "").endswith(
                          "destino de la producción (%)"))
        cab = next(r for r in range(titulo, titulo + 4)
                   if ws.cell(r, 1).value == "Actividad / rubro")
        cabeceras = [ws.cell(cab, c).value for c in range(1, 10)]
        self.assertIn("Familias en la actividad (%)", cabeceras)
        fila = cab + 1
        while ws.cell(fila, 1).value and not str(
                ws.cell(fila, 1).value).startswith("Fuente"):
            self.assertNotEqual(ws.cell(fila, 1).value, "TOTAL")
            fila += 1
        grafico = next(ch for ch in ws._charts
                       if "destino de la producción (%)" in str(
                           ch.title.tx.rich.p[0].r[0].t))
        self.assertEqual((grafico.y_axis.scaling.min, grafico.y_axis.scaling.max),
                         (0, 100.0))
        hoja_tabla = next(ws for ws in wb.worksheets
                          if ws.title.startswith("T Actividades"))
        fila_cab = next(fila for fila in hoja_tabla.iter_rows()
                        if fila[0].value == "Centro poblado / ámbito")
        cab_tabla = [c.value for c in fila_cab]
        self.assertIn("Distrito", cab_tabla)
        col = cab_tabla.index("Familias del CP en la actividad (%)") + 1
        formatos = {hoja_tabla.cell(r, col).number_format
                    for r in range(fila_cab[0].row + 1, hoja_tabla.max_row + 1)
                    if hoja_tabla.cell(r, col).value is not None}
        self.assertEqual(formatos, {"0.0"})

    def test_excel_de_la_base_de_datos_lleva_distrito(self):
        regs = _reporte_con_errores()
        wb = load_workbook(io.BytesIO(ex.exportar_fds_consolidado(regs)))
        ws = wb["Actividades economicas"]
        cabeceras = [c.value for c in ws[1]]
        self.assertIn("Distrito", cabeceras)
        col = cabeceras.index("Distrito")
        distritos = {fila[col].value for fila in ws.iter_rows(min_row=2)}
        self.assertEqual(distritos, {"Salitral"})
        self.assertIn("Actividades por CP (%)", wb.sheetnames)
        ws_pct = wb["Actividades por CP (%)"]
        cab_pct = [c.value for c in ws_pct[1]]
        self.assertIn("Distrito", cab_pct)
        col_pct = cab_pct.index("Familias del CP en la actividad (%)")
        self.assertTrue(all(fila[col_pct].value <= 100
                            for fila in ws_pct.iter_rows(min_row=2)
                            if fila[col_pct].value is not None))

    def test_distrito_del_bloque_si_la_ficha_no_lo_trae(self):
        regs = [_reg({"f1_nfam": "30", "f1_activ": _activ(
            ("Apicultura", "3", MERCADO))}, "Cerezal", distrito="",
            bloque_distrito="Chalaco")]
        wb = load_workbook(io.BytesIO(ex.exportar_fds_consolidado(regs)))
        ws = wb["Actividades economicas"]
        cabeceras = [c.value for c in ws[1]]
        self.assertEqual(ws.cell(2, cabeceras.index("Distrito") + 1).value, "Chalaco")

    def test_pdf_con_actividades_en_porcentaje(self):
        informe = an.indicadores_bloque(
            {"codigo": "M7B2"}, _reporte_con_errores(), {})
        self.assertTrue(an.generar_pdf_social(informe).startswith(b"%PDF"))
        self.assertEqual(an._rotulo_pdf(37.25, {"unidad": "%", "decimales": 1}),
                         "37.2 %")


if __name__ == "__main__":
    unittest.main()
