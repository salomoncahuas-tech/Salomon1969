# -*- coding: utf-8 -*-
"""
Pruebas de la Ficha Técnica Predial imprimible (liberacion_areas/la_ficha_predial.py).
`python -m pytest tests/test_ficha_predial.py -q`
"""
import csv
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
from liberacion_areas import la_campos as lc  # noqa: E402
from liberacion_areas import la_ficha_predial as fp  # noqa: E402


def _libro(tmp_path):
    return load_workbook(fp.generar_ficha_predial(tmp_path / "ficha.xlsx"))


def _lista(wb, nombre):
    hoja, rango = wb.defined_names[nombre].attr_text.split("!")
    return [c.value for fila in wb[hoja][rango.replace("$", "")] for c in fila if c.value is not None]


def _celda_kobo(wb, campo):
    for fila in wb["INSTRUCCIONES"].iter_rows(values_only=True):
        if fila[3] == campo:
            return fila[2]
    raise KeyError(campo)


def test_hojas_y_catalogo(tmp_path):
    wb = _libro(tmp_path)
    assert wb.sheetnames == ["FICHA PREDIAL", "INSTRUCCIONES", "CATALOGO_V6", "LISTAS", "_MAPA"]
    assert wb["_MAPA"].sheet_state == "hidden"
    cat = fp.cargar_catalogo()
    tipos = Counter(u["tipo"] for u in cat)
    assert tipos["Bloque"] == 117 and tipos["Lote SUS"] == 60
    assert round(sum(u["area_ha"] for u in cat if u["tipo"] == "Bloque"), 3) == 12270.235
    assert len({u["distrito"] for u in cat}) == 15


def test_unidades_por_asistente(tmp_path):
    """Cada asistente ve exactamente sus unidades del catálogo V6; ESP ve todas."""
    wb = _libro(tmp_path)
    with open(fp.CATALOGO_CSV, encoding="utf-8-sig", newline="") as f:
        filas = list(csv.DictReader(f))
    for codigo, _ in lc.OPCIONES["asistente"]:
        esperadas = {r["name"] for r in filas if codigo == "ESP" or r["asistente"] == codigo}
        assert set(_lista(wb, fp._nombre_rango("U_", codigo))) == esperadas, codigo
    assert _lista(wb, "L_asistente") == [c for c, _ in lc.OPCIONES["asistente"]]


def test_listas_iguales_a_kobo(tmp_path):
    wb = _libro(tmp_path)
    for k in fp.LISTAS_KOBO:
        assert _lista(wb, "L_" + k) == [e for _, e in lc.OPCIONES[k]], k
    assert "Otro" in _lista(wb, "L_uso_actual")
    assert _lista(wb, "L_personeria") == ["Natural", "Jurídica"]
    assert _lista(wb, "D_Morropon") and "Frías" in _lista(wb, "D_Ayabaca")


def test_validaciones_dependientes(tmp_path):
    wb = _libro(tmp_path)
    ws = wb["FICHA PREDIAL"]
    asis, unidad = _celda_kobo(wb, "asistente"), _celda_kobo(wb, "unidad")
    dvs = {str(c): dv for dv in ws.data_validations.dataValidation for c in dv.sqref.ranges}
    assert dvs[asis].formula1 == "L_asistente"
    assert dvs[unidad].formula1 == f'INDIRECT("U_"&SUBSTITUTE({asis},"-","_"))'
    assert any(dv.formula1.startswith('INDIRECT("D_"') for dv in dvs.values())
    assert any(dv.formula1 == "L_tipo_titularidad" for dv in dvs.values())
    assert any(dv.type == "decimal" and dv.formula1 == str(fp.ESTE_MIN) and dv.formula2 == str(fp.ESTE_MAX)
               for dv in dvs.values())
    assert all("#REF" not in (dv.formula1 or "") for dv in dvs.values())
    assert ws[_celda_kobo(wb, "cod_predio (calculado)")].value.startswith(f'=IF(OR({unidad}=""')
