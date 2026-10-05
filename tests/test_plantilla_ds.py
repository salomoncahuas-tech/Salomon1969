"""Plantilla oficial del Diagnostico Social (V4): listas y campos F-DS-01.

Verifica que la lista de niveles de IE de la hoja oculta `_Listas` coincide
con `fds_listas` (incluida la opcion Inicial + Primaria + Secundaria) y que
la mano de obra disponible consignada en la plantilla se importa al formulario.
"""

import io
import os
import sys
import unittest
import warnings
import zipfile

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import excel_diagnostico_social as eds  # noqa: E402
import fds_listas as FL  # noqa: E402


def _libro():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(eds.PLANTILLA_DS_PATH)


class PlantillaDS(unittest.TestCase):

    def test_lista_de_niveles_de_ie_coincide_con_la_plantilla(self):
        self.assertIn("Inicial + Primaria + Secundaria", FL.L_NIV_EDU_IE)
        ws = _libro()["_Listas"]
        col = next(c for c in range(1, ws.max_column + 1)
                   if ws.cell(1, c).value == "L_NIV_EDU_IE")
        valores = [ws.cell(r, col).value for r in range(2, ws.max_row + 1)]
        self.assertEqual([v for v in valores if v], FL.L_NIV_EDU_IE)

    def test_la_ficha_impresa_ofrece_la_nueva_opcion(self):
        ws = _libro()["F-DS-01"]
        self.assertEqual(ws["G88"].value, "Inicial + Primaria + Secundaria")
        self.assertEqual(ws["H24"].value, "Mano de obra disponible")

    def test_se_conservan_los_desplegables_nativos(self):
        with zipfile.ZipFile(eds.PLANTILLA_DS_PATH) as zf:
            n = sum(zf.read(i).count(b"<x14:dataValidation ")
                    for i in zf.namelist() if i.startswith("xl/worksheets/"))
        self.assertGreaterEqual(n, 10)

    def test_importa_la_mano_de_obra(self):
        wb = _libro()
        ws = wb["F-DS-01"]
        ws["L24"] = 42
        buf = io.BytesIO()
        wb.save(buf)
        res = eds.parsear_excel_ds(buf.getvalue(), ficha="F-DS-01")
        self.assertEqual(res[0]["datos"]["form"].get("f1_mano_obra"), "42")


if __name__ == "__main__":
    unittest.main()
