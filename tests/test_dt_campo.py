"""Pruebas de la integracion de las fichas DT de campo con los resumenes.

Cubren las tres piezas que sostienen el cruce: la lectura de la plantilla de
campo (incluido el rescate del codigo de bloque desde el nombre del
archivo), la regla que evita declarar un mismo hecho dos veces, y la
actualizacion del libro de resumen con su control de consistencia
regenerado. Las plantillas reales del repositorio se usan cuando estan
presentes; las pruebas de logica no dependen de ningun archivo.
"""

import io
import os
import sys
import unittest

from openpyxl import load_workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import dt_campo as dtc            # noqa: E402
import resumenes_bloques as rbq   # noqa: E402
from test_resumenes_bloques import construir_libro  # noqa: E402


# ══════════════════════════════════════════════════════════════════════════
# Fixture: ficha de campo equivalente a la que produce la plantilla V5
# ══════════════════════════════════════════════════════════════════════════

def ficha_campo(**cambios):
    """Registro de campo con la forma que devuelve `parsear_ficha_campo`."""
    base = {
        "nombre_archivo": "DT_Campo_Bloque_M9B1.xlsx",
        "fichas_leidas": list(dtc.FICHAS_CAMPO),
        "fichas_ausentes": [],
        "codigo_bloque": "M9B1",
        "microcuenca": "C1096-Q9545",
        "provincia": "Morropon",
        "distrito": "Chulucanas",
        "evaluador": "Ing. Juan Carlos Dominguez Varas",
        "fecha_evaluacion": "03/07/2026",
        "utm_este_dt": "595600",
        "utm_norte_dt": "9451300",
        "altitud_gps": "210",
        "centro_poblado_cercano": "La Peña",
        "forma_terreno": "Colinoso",
        "pendiente": "15-25% (Fuert. inclinado)",
        "posicion_fisiografica": "Ladera baja",
        "exposicion_orientacion": "Noreste",
        "rango_altitudinal": "<1000 m (Yunga)",
        "dt01_afloramientos_rocosos": "Sí",
        "dt01_escarpes_activos": "No",
        "dt01_remociones_masa_activas": "No",
        "dt01_reptacion_suelo": "Sí",
        "dt02_nivel_erosion_general": "Moderada (surcos)",
        "dt02_num_carcavas": "3",
        "dt02_longitud_total_carcavas": "120",
        "dt02_urgencia_control": "Alta",
        "dt03_tipo_ecosistema": "Bosque Estacionalmente Seco Colina/Montaña (Bes-cm)",
        "dt03_estado_conservacion_eco": "Medianamente alterado",
        "dt03_uso_dominante": "Pastoreo extensivo",
        "dt03_cobertura_total": "90",
        "dt03_suelo_desnudo": "10",
        "dt03_cobertura_dosel": "40",
        "dt03_regeneracion_natural": "Regular (10-50 pl./100m²)",
        "dt03_estado_sanitario": "Sano",
        "dt04_causa_subyacente": "Economia de subsistencia sin ordenamiento",
        "dt04_velocidad_degradacion": "Moderada",
        "dt04_reversibilidad": "Parcialmente reversible",
        "dt04_urgencia_intervencion": "Media",
        "dt05_zona_recarga": "Sí",
        "dt05_modalidad_acceso": "Vehicular + caminata <30 min",
        "dt05_senal_celular": "Parcial",
        "carcavas": [], "floristica": [], "especies_clave": [],
        "causas": [], "causas_activas": [], "indicadores": [],
        "fuentes_agua": [],
        "n_carcavas_inventariadas": 0, "n_taxones": 9,
        "n_especies_clave": 0, "n_fuentes_agua": 0, "n_causas_activas": 0,
        "n_indicadores": 0, "completitud_pct": 96.4,
    }
    base.update(cambios)
    return base


def resumen_de(codigo="M9B1", **cambios):
    """Diccionario de la ficha de resumen, leido del libro de produccion."""
    datos = rbq.completar_sintesis_msavi(rbq.parsear_resumen_bloque(
        construir_libro(codigo=codigo),
        f"Plantilla_Excel_Bloque_{codigo}_IN_Piura.xlsx"))
    datos.update(cambios)
    return datos


# ══════════════════════════════════════════════════════════════════════════
# Codigos de bloque
# ══════════════════════════════════════════════════════════════════════════

class TestCodigos(unittest.TestCase):

    def test_codigo_desde_nombres_heterogeneos(self):
        casos = {
            "DT_B13_IN_Piura.xlsx": "13",
            "DT_Bloque67_V5.xlsx": "67",
            "Plantilla_DT_Campo_Check_Validada_V5_Bloque23.xlsx": "23",
            "FDT_V5_Bloque87_IN_Piura.xlsx": "87",
            "DT_M1B1_llenado.xlsx": "M1B1",
            "Diagnostico_Territorial_M18B1_IN_Piura.xlsx": "M18B1",
            "FDT_M6B2-1_Diagnostico_Territorial_V5.xlsx": "M6B2-1",
            "Plantilla_DT_Campo_V5_BloqueM3B6.xlsx": "M3B6",
            "Plantilla_DT_Campo_Check_M17B10_IN_Piura_V5.xlsx": "M17B10",
            "M17B5.xlsx": "M17B5",
        }
        for nombre, esperado in casos.items():
            self.assertEqual(dtc.codigo_desde_nombre_campo(nombre), esperado,
                             nombre)

    def test_nombre_sin_bloque_no_inventa_codigo(self):
        self.assertEqual(dtc.codigo_desde_nombre_campo("Plantilla_V5.xlsx"), "")

    def test_normalizacion_de_codigos(self):
        self.assertEqual(dtc.normalizar_codigo("013"), "13")
        self.assertEqual(dtc.normalizar_codigo(" m6b2-1 "), "M6B2-1")
        self.assertEqual(dtc.normalizar_codigo(None), "")

    def test_orden_natural(self):
        codigos = ["M10B4", "10", "2", "M1B1"]
        self.assertEqual(sorted(codigos, key=dtc._orden_codigo),
                         ["2", "10", "M1B1", "M10B4"])


# ══════════════════════════════════════════════════════════════════════════
# Cruce de fuentes
# ══════════════════════════════════════════════════════════════════════════

class TestIntegracion(unittest.TestCase):

    def setUp(self):
        self.integrado = dtc.integrar_bloque(resumen_de(), ficha_campo())

    def test_cada_hecho_se_declara_una_sola_vez(self):
        claves = [c["clave"] for c in self.integrado["campos"]]
        self.assertEqual(len(claves), len(set(claves)))
        self.assertEqual(len(claves), len(dtc.CAMPOS_INTEGRADOS))

    def test_todo_hecho_declarado_tiene_una_sola_procedencia(self):
        for registro in self.integrado["campos"]:
            if registro["estado"] == dtc.PENDIENTE:
                self.assertEqual(registro["valor"], "")
            else:
                self.assertIn(registro["fuente_valor"], dtc.FUENTES)
                self.assertTrue(registro["valor"])

    def test_campo_actualiza_el_valor_de_gabinete(self):
        erosion = self.integrado["por_clave"]["nivel_erosion"]
        self.assertEqual(erosion["estado"], dtc.ACTUALIZADO)
        self.assertEqual(erosion["valor"], "Moderada (surcos)")
        self.assertEqual(erosion["fuente_valor"], dtc.CAMPO)
        # El valor anterior no se pierde: queda como trazabilidad.
        self.assertEqual(erosion["valor_alterno"],
                         "FICHA F-DT-02 SIN CONTENIDO")

    def test_el_catalogo_manda_sobre_la_ficha_de_campo(self):
        integrado = dtc.integrar_bloque(
            resumen_de(), ficha_campo(distrito="Morropon"))
        distrito = integrado["por_clave"]["distrito"]
        self.assertEqual(distrito["estado"], dtc.DISCREPANTE)
        self.assertEqual(distrito["valor"], "Chulucanas")
        self.assertEqual(distrito["fuente_valor"], dtc.OFICIAL)
        self.assertEqual(distrito["valor_alterno"], "Morropon")

    def test_hecho_solo_de_campo_se_adopta_sin_duplicar(self):
        reptacion = self.integrado["por_clave"]["reptacion_suelo"]
        self.assertEqual(reptacion["valor"], "Sí")
        self.assertEqual(reptacion["fuente_valor"], dtc.CAMPO)
        self.assertEqual(reptacion["valor_resumen"], "")

    def test_misma_fecha_en_distinto_formato_no_es_discrepancia(self):
        fecha = self.integrado["por_clave"]["fecha_evaluacion"]
        self.assertEqual(fecha["estado"], dtc.CONFORME)

    def test_misma_clase_con_otra_redaccion_no_es_discrepancia(self):
        integrado = dtc.integrar_bloque(
            resumen_de(tipo_ecosistema="Bosque seco de colina y montaña (Bes-cm)"),
            ficha_campo())
        self.assertEqual(
            integrado["por_clave"]["tipo_ecosistema"]["estado"], dtc.CONFORME)

    def test_marcador_de_ausencia_no_cuenta_como_dato(self):
        peligro = self.integrado["por_clave"]["peligro_integrado"]
        self.assertEqual(peligro["estado"], dtc.PENDIENTE)

    def test_no_aplica_si_es_una_declaracion_con_contenido(self):
        integrado = dtc.integrar_bloque(
            resumen_de(comunidad_campesina="No aplica – tierras del Estado"),
            ficha_campo())
        self.assertNotEqual(
            integrado["por_clave"]["comunidad_campesina"]["estado"],
            dtc.PENDIENTE)

    def test_distancia_al_centroide_y_validacion_utm(self):
        self.assertAlmostEqual(self.integrado["distancia_centroide_m"], 91.6,
                               places=1)
        self.assertEqual(self.integrado["validacion_utm"], "Conforme")

    def test_bloque_sin_ficha_de_campo_se_integra_igual(self):
        integrado = dtc.integrar_bloque(resumen_de(), None)
        self.assertFalse(integrado["tiene_campo"])
        self.assertEqual(integrado["por_clave"]["distrito"]["valor"],
                         "Chulucanas")
        self.assertEqual(integrado["fichas_ausentes"], list(dtc.FICHAS_CAMPO))

    def test_bloque_sin_ficha_de_resumen_se_integra_igual(self):
        integrado = dtc.integrar_bloque({}, ficha_campo(), "M9B1")
        self.assertFalse(integrado["tiene_resumen"])
        self.assertEqual(integrado["por_clave"]["forma_terreno"]["valor"],
                         "Colinoso")

    def test_conteos_cuadran_con_el_numero_de_hechos(self):
        self.assertEqual(sum(self.integrado["conteo_estado"].values()),
                         self.integrado["n_campos"])
        self.assertEqual(sum(self.integrado["conteo_fuente"].values()),
                         self.integrado["n_declarados"])


# ══════════════════════════════════════════════════════════════════════════
# Control de consistencia
# ══════════════════════════════════════════════════════════════════════════

class TestConsistencia(unittest.TestCase):

    def test_califica_con_el_vocabulario_de_la_ficha(self):
        integrado = dtc.integrar_bloque(resumen_de(), ficha_campo())
        for registro in integrado["consistencia"]:
            self.assertIn(registro["calificacion"], rbq.CALIFICACIONES)
            self.assertTrue(registro["codigo"])
            self.assertTrue(registro["discrepancia"])
            self.assertTrue(registro["tratamiento"])

    def test_peligro_integrado_sin_dato_es_sustantivo(self):
        integrado = dtc.integrar_bloque(resumen_de(), ficha_campo())
        peligro = [r for r in integrado["consistencia"]
                   if "Peligro integrado" in r["campo"]]
        self.assertEqual(len(peligro), 1)
        self.assertEqual(peligro[0]["calificacion"], "SUSTANTIVA")

    def test_una_verificacion_por_hecho_actualizado(self):
        integrado = dtc.integrar_bloque(resumen_de(), ficha_campo())
        actualizados = sum(1 for c in integrado["campos"]
                           if c["estado"] == dtc.ACTUALIZADO)
        filas = [r for r in integrado["consistencia"]
                 if r["codigo"].startswith("I-")]
        self.assertEqual(len(filas), actualizados)

    def test_coordenada_fuera_de_rango_se_declara(self):
        integrado = dtc.integrar_bloque(
            resumen_de(), ficha_campo(utm_este_dt="123456"))
        self.assertNotEqual(integrado["validacion_utm"], "Conforme")
        fila = [r for r in integrado["consistencia"]
                if "Sistema de referencia" in r["campo"]][0]
        self.assertEqual(fila["calificacion"], "SUSTANTIVA")

    def test_conserva_las_verificaciones_redactadas_en_la_ficha_anterior(self):
        heredada = {"codigo": "C-01", "campo": "Duplicidad de fichas",
                    "discrepancia": "El archivo trae dos juegos de fichas.",
                    "calificacion": "SUSTANTIVA",
                    "tratamiento": "Se adopta el juego 2."}
        integrado = dtc.integrar_bloque(
            resumen_de(consistencia=[heredada]), ficha_campo())
        filas = [r for r in integrado["consistencia"]
                 if r["campo"] == "Duplicidad de fichas"]
        self.assertEqual(len(filas), 1)
        self.assertTrue(filas[0]["codigo"].startswith("H-"))

    def test_no_duplica_una_verificacion_ya_regenerada(self):
        heredada = {"codigo": "D-01", "campo": "Código de microcuenca",
                    "discrepancia": "Coinciden.", "calificacion": "CONFORME",
                    "tratamiento": "Sin acción."}
        integrado = dtc.integrar_bloque(
            resumen_de(consistencia=[heredada]), ficha_campo())
        filas = [r for r in integrado["consistencia"]
                 if "microcuenca" in r["campo"].lower()]
        self.assertEqual(len(filas), 1)


# ══════════════════════════════════════════════════════════════════════════
# Actualizacion del libro de resumen
# ══════════════════════════════════════════════════════════════════════════

class TestActualizacionLibro(unittest.TestCase):

    def setUp(self):
        self.original = construir_libro(codigo="M9B1")
        self.integrado = dtc.integrar_bloque(resumen_de(), ficha_campo())
        self.nuevo = dtc.actualizar_libro(self.original, self.integrado)

    def test_conserva_las_hojas_y_agrega_las_de_integracion(self):
        wb = load_workbook(io.BytesIO(self.nuevo))
        for hoja in ("Resumen", "Cobertura MSAVI-NDVI", "Microcuenca",
                     "Control de consistencia", "Integración de campo",
                     "Registro de campo F-DT", "Gráficos de integración"):
            self.assertIn(hoja, wb.sheetnames)

    def test_el_libro_actualizado_sigue_siendo_legible(self):
        datos = rbq.parsear_resumen_bloque(
            self.nuevo, "Plantilla_Excel_Bloque_M9B1_IN_Piura.xlsx")
        self.assertEqual(datos["codigo_bloque"], "M9B1")
        self.assertEqual(datos["consistencia_resumen"]["total"],
                         len(self.integrado["consistencia"]))

    def test_escribe_en_la_hoja_resumen_el_valor_de_campo(self):
        datos = rbq.parsear_resumen_bloque(self.nuevo, "x.xlsx")
        self.assertEqual(datos["nivel_erosion"], "Moderada (surcos)")
        self.assertIn("nivel_erosion", self.integrado["campos_reescritos"])

    def test_no_toca_los_valores_de_gabinete(self):
        datos = rbq.parsear_resumen_bloque(self.nuevo, "x.xlsx")
        self.assertEqual(datos["pendiente_pct_num"], 10.17)
        self.assertEqual(datos["msavi_2024_num"], 0.3143)

    def test_reintegrar_no_acumula_verificaciones(self):
        segunda = dtc.integrar_bloque(
            rbq.completar_sintesis_msavi(
                rbq.parsear_resumen_bloque(self.nuevo, "x.xlsx")),
            ficha_campo())
        tercera = dtc.integrar_bloque(
            rbq.completar_sintesis_msavi(
                rbq.parsear_resumen_bloque(
                    dtc.actualizar_libro(self.nuevo, segunda), "x.xlsx")),
            ficha_campo())
        self.assertEqual(len(segunda["consistencia"]),
                         len(tercera["consistencia"]))


# ══════════════════════════════════════════════════════════════════════════
# Agregacion y reportes
# ══════════════════════════════════════════════════════════════════════════

class TestReportes(unittest.TestCase):

    def setUp(self):
        self.integrados = [
            dtc.integrar_bloque(resumen_de(), ficha_campo()),
            dtc.integrar_bloque(resumen_de(codigo="M9B2", distrito="Frías"),
                                ficha_campo(codigo_bloque="M9B2",
                                            distrito="Frías")),
        ]

    def test_agrupa_por_los_cuatro_niveles(self):
        for nivel in dtc.AGRUPACIONES:
            grupos = dtc.agrupar(self.integrados, nivel)
            self.assertTrue(grupos)
            self.assertEqual(sum(g["n_bloques"] for g in grupos),
                             len(self.integrados))

    def test_agrupacion_desconocida_falla(self):
        with self.assertRaises(ValueError):
            dtc.agrupar(self.integrados, "departamento")

    def test_totales_suman_los_bloques(self):
        totales = dtc.totales(self.integrados)
        self.assertEqual(totales["n_bloques"], 2)
        self.assertEqual(totales["bloques_con_campo"], 2)

    def test_excel_consolidado_trae_las_hojas_del_formato(self):
        contenido = dtc.generar_excel_consolidado(self.integrados, "provincia")
        wb = load_workbook(io.BytesIO(contenido))
        for hoja in ("Síntesis", "Por provincia", "Bloques",
                     "Hechos actualizados", "Control de consistencia",
                     "Metodología"):
            self.assertIn(hoja, wb.sheetnames)

    def test_pdf_por_bloque_y_consolidado(self):
        self.assertTrue(dtc.generar_pdf_bloque(self.integrados[0])
                        .startswith(b"%PDF"))
        self.assertTrue(dtc.generar_pdf_consolidado(self.integrados, "distrito")
                        .startswith(b"%PDF"))

    def test_tabla_para_el_aplicativo(self):
        filas = dtc.tabla_integrados(self.integrados)
        self.assertEqual(len(filas), 2)
        self.assertIn("Hechos de campo", filas[0])


# ══════════════════════════════════════════════════════════════════════════
# Plantillas reales del repositorio
# ══════════════════════════════════════════════════════════════════════════

class TestPlantillasDelRepositorio(unittest.TestCase):
    """Se ejecutan solo si las carpetas de trabajo estan en el despliegue."""

    @classmethod
    def setUpClass(cls):
        cls.fichas = dtc.fichas_del_repositorio()
        if not cls.fichas:
            raise unittest.SkipTest("Sin carpeta %s" % dtc.CARPETA_CAMPO)

    def test_todas_las_fichas_resuelven_su_bloque(self):
        sin_codigo = [n for n, _ in self.fichas
                      if not dtc.codigo_desde_nombre_campo(n)]
        self.assertEqual(sin_codigo, [])

    def test_el_manifiesto_cubre_las_fichas_de_la_carpeta(self):
        esperados = dtc.codigos_campo_esperados()
        if not esperados:
            self.skipTest("Sin manifiesto de fichas de campo")
        self.assertEqual(len(esperados), len(set(esperados)))
        self.assertEqual(len(esperados), len(self.fichas))

    def test_lectura_de_una_plantilla_real(self):
        nombre, contenido = self.fichas[0]
        registro = dtc.parsear_ficha_campo(contenido, nombre)
        self.assertTrue(registro["codigo_bloque"])
        self.assertEqual(registro["fichas_ausentes"], [])
        self.assertGreater(registro["completitud_pct"], 0)


class TestLibrosIntegradosDelRepositorio(unittest.TestCase):
    """Los 117 libros vigentes (V7) que `integrar_dt_campo.py` deja al dia."""

    @classmethod
    def setUpClass(cls):
        cls.libros = rbq.libros_del_repositorio()
        if not cls.libros:
            raise unittest.SkipTest("Sin carpeta %s" % rbq.CARPETA_LIBROS)

    def test_trae_un_libro_por_bloque_del_manifiesto(self):
        nombres = {n for n, _ in self.libros}
        esperados = {b["archivo"] for b in rbq.cargar_manifiesto()["bloques"]}
        self.assertEqual(nombres, esperados)

    def test_cada_libro_suma_las_hojas_de_la_integracion(self):
        for nombre, contenido in self.libros:
            wb = load_workbook(io.BytesIO(contenido), read_only=True)
            try:
                for hoja in ("Resumen", "Control de consistencia",
                             "Integración de campo", "Registro de campo F-DT",
                             "Gráficos de integración"):
                    self.assertIn(hoja, wb.sheetnames, nombre)
            finally:
                wb.close()

    def test_el_control_de_consistencia_esta_regenerado(self):
        for nombre, contenido in self.libros:
            datos = rbq.parsear_resumen_bloque(contenido, nombre)
            registros = datos["consistencia"]
            self.assertTrue(registros, nombre)
            self.assertEqual(datos["consistencia_resumen"]["total"],
                             len(registros), nombre)
            codigos = [r["codigo"] for r in registros]
            # La serie G (controles de geometria y coherencia espectral) se
            # regenera completa y correlativa en cada libro.
            serie_g = [c for c in codigos if c.startswith("G-")]
            self.assertEqual(serie_g,
                             ["G-%02d" % i for i in range(1, len(serie_g) + 1)],
                             nombre)
            for registro in registros:
                self.assertIn(registro["calificacion"], rbq.CALIFICACIONES,
                              nombre)

    def test_no_pierde_los_parametros_de_gabinete(self):
        for nombre, contenido in self.libros:
            datos = rbq.parsear_resumen_bloque(contenido, nombre)
            self.assertTrue(datos.get("codigo_bloque"), nombre)
            self.assertIsNotNone(datos.get("msavi_2024_num"), nombre)
            self.assertIsNotNone(datos.get("area_ha_num"), nombre)
            self.assertEqual(len(datos.get("msavi_tabla") or []), 5, nombre)


if __name__ == "__main__":
    unittest.main()
