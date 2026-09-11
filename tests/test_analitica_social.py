"""Pruebas de la analitica del Diagnostico Social (F-DS-01..F-DS-07).

El fixture reproduce registros tal como los devuelve
`database.obtener_diagnosticos_sociales_por_bloque()`: cabecera con bloque,
centro poblado y responsable, mas el formulario V4 serializado en la columna
`dsNN_data_v3`. Se cubren el caso completo, el bloque con fichas parciales y
los campos numericos escritos como texto libre, que es como llegan de campo.
"""

import io
import json
import os
import sys
import unittest

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import analitica_social as an  # noqa: E402


# ══════════════════════════════════════════════════════════════════════════
# Fixture
# ══════════════════════════════════════════════════════════════════════════

def _registro(ficha, form, centro_poblado="Chungayo", **extra):
    num = ficha.split("-")[-1]
    registro = {
        "id": extra.pop("id", 1), "ficha": ficha,
        "bloque_codigo": extra.pop("bloque_codigo", "M5-B1"),
        "centro_poblado": centro_poblado,
        "comunidad_campesina": extra.pop("comunidad_campesina", "CC Segunda y Cajas"),
        "distrito": extra.pop("distrito", "Huancabamba"),
        "provincia": extra.pop("provincia", "Huancabamba"),
        "fecha_evaluacion": extra.pop("fecha_evaluacion", "2026-03-15"),
        "evaluador": extra.pop("evaluador", "H. Cahuas"),
        f"ds{num}_data_v3": json.dumps(form, ensure_ascii=False),
    }
    registro.update(extra)
    return registro


def _fds01(centro_poblado="Chungayo", **campos):
    form = {
        "f1_nombre_oficial": centro_poblado,
        # Los numeros llegan como texto libre: miles con coma, unidades y
        # porcentajes conviven en el mismo formulario.
        "f1_nfam": "120", "f1_pob_t": "1,450", "f1_pob_h": "740",
        "f1_pob_m": "710", "f1_pob_men18": "520", "f1_pob_may65": "130",
        "f1_idioma": "Español", "f1_nivel_edu": "Primaria completa",
        "f1_migracion": "Alto", "f1_presencia_estatal": "Bajo",
        "f1_junta_vig": "Sí", "f1_ronda": "Sí", "f1_reglamento": "No",
        "f1_comite_rrnn": "No", "f1_pdc": "Sí", "f1_ongs": "No",
        "f1_agrorural": "Sí", "f1_prodern": "No",
        "f1_tenencia": "Comunal (tierras comunales)",
        "f1_pct_tituladas": "35 %",
        "f1_agua": ["JASS / Sistema local", "Manantial / Quebrada directa"],
        "f1_agua_cob": "68", "f1_sanea": ["Letrina seca"],
        "f1_energia": ["Red pública", "Panel solar domiciliario"],
        "f1_energia_cob": "82",
        "f1_activ": [
            {"Actividad / Rubro": "Agricultura de secano", "N fam.": "90",
             "Productos principales": "Maíz, arveja",
             "Destino": "Mayormente autoconsumo (>70%)", "Ingreso (S/./mes)": "350"},
            {"Actividad / Rubro": "Ganadería vacuna", "N fam.": "45",
             "Productos principales": "Leche, queso",
             "Destino": "Mixto autoconsumo/mercado (30-70%)",
             "Ingreso (S/./mes)": "S/ 620"},
        ],
        "f1_juntos": "48", "f1_pension65": "26", "f1_beca18": "3",
    }
    form.update(campos)
    return form


def _fds02():
    return {
        "f2_actores": [
            {"Nombre del actor / Organizacion": "Comunidad Campesina Segunda y Cajas",
             "Tipo": "Comunidad Campesina", "Rol / Funcion frente al proyecto": "Titular del predio",
             "Influencia": "Alto", "Interes": "Alto", "Posicion": "A favor del proyecto",
             "Nivel territorial": "Comunal"},
            {"Nombre del actor / Organizacion": "Ronda Campesina de Chungayo",
             "Tipo": "Ronda Campesina", "Rol / Funcion frente al proyecto": "Control de acceso",
             "Influencia": "Alto", "Interes": "Medio",
             "Posicion": "Reticente (requiere persuasión)", "Nivel territorial": "Comunal"},
            {"Nombre del actor / Organizacion": "Municipalidad Distrital de Huancabamba",
             "Tipo": "Gobierno Local (Municipalidad)", "Rol / Funcion frente al proyecto": "Articulación",
             "Influencia": "Medio", "Interes": "Alto", "Posicion": "Neutral / Sin posición definida",
             "Nivel territorial": "Distrital"},
        ],
        "f2_favor": "Comunidad Campesina Segunda y Cajas",
    }


def _fds04():
    return {
        "f4_lugar": "Local comunal de Chungayo", "f4_conv": "ANIN - DIME",
        "f4_fecha": "2026-03-15", "f4_conv_n": "80", "f4_h": "34",
        "f4_m": "28", "f4_jov": "12", "f4_am": "9", "f4_tot": "62",
        "f4_metod": ["Mapa parlante", "Mesas de trabajo / Grupos focales"],
        "f4_idioma": "Español",
        "f4_part": [{"Nombres y Apellidos": f"Participante {i}", "Sexo": "M"}
                    for i in range(5)],
        "f4_acuerdos": [
            {"Acuerdo / Compromiso": "Autorizar el ingreso de la brigada",
             "Responsable": "Junta Directiva", "Plazo": "abril 2026"},
            {"Acuerdo / Compromiso": "Convocar asamblea de titulares",
             "Responsable": "Presidente", "Plazo": "mayo 2026"},
        ],
    }


def _fds05():
    return {
        "f5_conflictos": [
            {"Tipo": "Socioambiental minero", "Actores involucrados": "Comuneros, empresa",
             "Estado": "Latente (no manifiesto)", "Antiguedad": "Histórico (>20 años)",
             "Descripcion / Causa raiz": "Antecedente Río Blanco",
             "Impacto potencial en el proyecto": "Desconfianza inicial"},
            {"Tipo": "Tenencia / Linderos de tierras", "Actores involucrados": "Familias colindantes",
             "Estado": "En negociación / Mesa de diálogo", "Antiguedad": "3-5 años",
             "Descripcion / Causa raiz": "Linderos sin saneamiento"},
            {"Tipo": "Socioambiental hídrico", "Actores involucrados": "JASS y regantes",
             "Estado": "Resuelto / Inactivo", "Antiguedad": "1-3 años"},
        ],
        "f5_oportunidades": [
            {"Oportunidad identificada": "Mesa de gestión de la microcuenca",
             "Actores relacionados": "Municipalidad, JASS", "Potencial": "Alto"},
            {"Oportunidad identificada": "Convenio con la ronda campesina",
             "Actores relacionados": "Ronda", "Potencial": "Medio"},
        ],
        "f5_rb1": "Sí", "f5_rb2": "No", "f5_rb3": "Sí", "f5_rb4": "No",
        "f5_rb5": "Sí", "f5_polar": "Medio", "f5_confglob": "Medio",
        "f5_viab": "Media — requiere acercamiento reforzado",
        "f5_plazo": "Medio (3-6 meses)", "f5_mesa": "Sí",
        "f5_estrategia": "Acercamiento a través de la ronda campesina.",
    }


def _fds06():
    return {
        "f6_fuente": "12 personas consultadas",
        "f6_peligros": [
            {"Peligro observado": "Movimientos en masa (deslizamientos)",
             "¿Ocurre?": "Sí", "Frecuencia": "F (Frecuente)", "Magnitud": "A (Alta)",
             "Tendencia": "Sube (Aumentando)", "Ultimo evento (año)": "2025",
             "Principales daños observados": "Pérdida de chacras y camino"},
            {"Peligro observado": "Erosión hídrica (cárcavas)", "¿Ocurre?": "Sí",
             "Frecuencia": "MF (Muy frecuente)", "Magnitud": "M (Media)",
             "Tendencia": "Sube (Aumentando)"},
            {"Peligro observado": "Heladas / Friaje", "¿Ocurre?": "No",
             "Frecuencia": "SR (Sin registro)"},
        ],
        "f6_cambios": [
            {"Cambio observado": "Lluvias más intensas", "¿Se percibe?": "Sí",
             "Intensidad": "Alta", "Año aprox. de inicio": "2015"},
            {"Cambio observado": "Pérdida de manantiales", "¿Se percibe?": "Sí",
             "Intensidad": "Media"},
            {"Cambio observado": "Heladas más frecuentes", "¿Se percibe?": "No"},
        ],
        "f6_medidas": "Sí", "f6_alerta": "No", "f6_saberes": "Sí",
        "f6_apoyo": "Sí",
        "f6_p1": "Movimientos en masa (deslizamientos)",
        "f6_p2": "Erosión hídrica (cárcavas)",
        "f6_p3": "Sequía / Déficit hídrico",
    }


def _fds07(nombre="Juan Pérez", disposicion="Totalmente dispuesto/a"):
    form = {
        "f7_tipo_prop": "Titular individual del predio", "f7_nombre": nombre,
        "f7_dni": "44556677", "f7_edad": "52", "f7_superficie": "12.5",
        "f7_docs": ["Constancia de posesion (municipal)"],
        "f7_linderos": "No", "f7_residente": "Sí", "f7_disp": disposicion,
        "f7_aut_ing": "Sí", "f7_aut_foto": "Sí", "f7_aut_nom": "No",
        "f7_consultar": "No",
    }
    form.update({clave: "Sí" for clave, _ in an._PUNTOS_CPI})
    form["f7_info_material"] = "No"
    return form


def registros_completos():
    """Un bloque con las siete fichas y dos centros poblados."""
    return [
        _registro("F-DS-01", _fds01("Chungayo"), "Chungayo", id=1),
        _registro("F-DS-01", _fds01("Sapalache", f1_pob_t="820", f1_pob_h="410",
                                    f1_pob_m="410", f1_ronda="No",
                                    f1_agua_cob="41", f1_energia_cob="55",
                                    f1_pct_tituladas="70"),
                  "Sapalache", id=2, evaluador="M. Ruiz"),
        _registro("F-DS-02", _fds02(), "Chungayo", id=3),
        _registro("F-DS-03", {"f3_nombre": "Rosa Calle", "f3_dni": "12345678",
                              "f3_edad": "47", "f3_genero": "F",
                              "f3_cargo": "Presidenta de la Junta",
                              "f3_inst": "Comunidad Campesina", "f3_dur": "55",
                              "f3_c_nom": "Sí", "f3_c_foto": "No"},
                  "Chungayo", id=4),
        _registro("F-DS-04", _fds04(), "Chungayo", id=5),
        _registro("F-DS-05", _fds05(), "Chungayo", id=6),
        _registro("F-DS-06", _fds06(), "Chungayo", id=7),
        _registro("F-DS-07", _fds07("Juan Pérez"), "Chungayo", id=8),
        _registro("F-DS-07", _fds07("Ana Quispe", "No dispuesto/a por el momento"),
                  "Sapalache", id=9, evaluador="M. Ruiz"),
    ]


BLOQUE = {"id": 1, "codigo": "M5-B1", "microcuenca": "Microcuenca 5",
          "provincia": "Huancabamba", "distrito": "Huancabamba"}

DATOS_CP = {"centros_poblados": ["Chungayo", "Sapalache", "El Tambo"],
            "poblacion_total": 2500}


# ══════════════════════════════════════════════════════════════════════════
# Lectura de campos
# ══════════════════════════════════════════════════════════════════════════

class LecturaCampos(unittest.TestCase):

    def test_numeros_en_texto_libre(self):
        """Los campos numericos son cajas de texto: se recupera la cifra."""
        self.assertEqual(an._num("1,450"), 1450.0)
        self.assertEqual(an._num("35 %"), 35.0)
        self.assertEqual(an._num("S/ 620"), 620.0)
        self.assertEqual(an._num("12.5"), 12.5)
        self.assertEqual(an._num("aprox. 40 familias"), 40.0)
        self.assertEqual(an._num("12,5"), 12.5)

    def test_sin_dato_no_es_cero(self):
        """Un campo vacio no puede contarse como cero en los promedios."""
        for vacio in ("", None, "s/d", "  ", "sin dato"):
            self.assertIsNone(an._num(vacio), vacio)

    def test_normalizacion_si_no(self):
        self.assertEqual(an._sino("si"), "Sí")
        self.assertEqual(an._sino("SÍ"), "Sí")
        self.assertEqual(an._sino("No"), "No")
        self.assertEqual(an._sino("No aplica"), "No aplica")
        self.assertEqual(an._sino(""), "")

    def test_columnas_con_alias_y_tildes(self):
        """El encabezado de la tabla puede venir con o sin tilde."""
        fila = {"Posición": "A favor del proyecto", "Antiguedad": "1-3 años"}
        self.assertEqual(an._col(fila, "Posicion", "Posición"),
                         "A favor del proyecto")
        self.assertEqual(an._col(fila, "Antigüedad"), "1-3 años")
        self.assertEqual(an._col(fila, "Inexistente"), "")

    def test_deduplicacion_conserva_el_mas_reciente(self):
        """Una reedicion no debe contar dos veces en los graficos."""
        base = _registro("F-DS-01", _fds01(), "Chungayo", id=2)
        vieja = _registro("F-DS-01", _fds01(), "Chungayo", id=1)
        salida = an.deduplicar([base, vieja])
        self.assertEqual([r["id"] for r in salida], [2])

    def test_deduplicacion_respeta_centros_poblados_distintos(self):
        salida = an.deduplicar([
            _registro("F-DS-01", _fds01(), "Chungayo", id=1),
            _registro("F-DS-01", _fds01(), "Sapalache", id=2)])
        self.assertEqual(len(salida), 2)


# ══════════════════════════════════════════════════════════════════════════
# Construccion del informe
# ══════════════════════════════════════════════════════════════════════════

class Informe(unittest.TestCase):

    def setUp(self):
        self.informe = an.indicadores_bloque(BLOQUE, registros_completos(),
                                             DATOS_CP)

    def _serie(self, id_):
        for seccion in self.informe["secciones"]:
            for serie in seccion["series"]:
                if serie["id"] == id_:
                    return serie
        return None

    def test_secciones_de_las_siete_fichas(self):
        ids = {s["id"] for s in self.informe["secciones"]}
        self.assertEqual(ids, {"COBERTURA"} | set(an.FICHAS_DS))

    def test_todas_las_series_tienen_filas_y_titulo(self):
        for seccion in self.informe["secciones"]:
            self.assertTrue(seccion["series"], seccion["id"])
            for serie in seccion["series"]:
                self.assertTrue(serie["filas"], serie["id"])
                self.assertTrue(serie["titulo"], serie["id"])
                self.assertIn(serie["forma"],
                              ("barras_h", "apiladas", "agrupadas", "mapa_calor"))

    def test_poblacion_suma_los_dos_centros_poblados(self):
        serie = self._serie("f1_poblacion")
        total = sum(f["valor"] for f in serie["filas"])
        self.assertEqual(total, 1450 + 820)

    def test_estructura_etaria_no_supera_la_poblacion(self):
        """El tramo intermedio sale por diferencia: nunca debe ser negativo."""
        serie = self._serie("f1_etaria")
        por_ambito = {}
        for fila in serie["filas"]:
            por_ambito[fila["cat"]] = por_ambito.get(fila["cat"], 0) + fila["valor"]
            self.assertGreaterEqual(fila["valor"], 0)
        self.assertEqual(por_ambito["Chungayo"], 1450)

    def test_matriz_de_actores(self):
        serie = self._serie("f2_matriz")
        celdas = {(f["cat"], f["sub"]): f["valor"] for f in serie["filas"]}
        self.assertEqual(celdas[("Alto", "Alto")], 1)
        self.assertEqual(celdas[("Alto", "Medio")], 1)
        self.assertEqual(celdas[("Medio", "Alto")], 1)

    def test_conflicto_resuelto_no_se_pinta_como_riesgo_vivo(self):
        serie = self._serie("f5_tipos")
        resuelto = [c for c in serie["colores"]
                    if c.lower().startswith("resuelto")]
        self.assertTrue(resuelto)
        self.assertIn(serie["colores"][resuelto[0]], an.RAMPA_FAVORABLE)

    def test_solo_se_grafican_peligros_ocurrentes(self):
        """El peligro marcado 'No ocurre' no entra en la frecuencia."""
        serie = self._serie("f6_frecuencia")
        categorias = {f["cat"] for f in serie["filas"]}
        self.assertIn("Movimientos en masa (deslizamientos)", categorias)
        self.assertNotIn("Heladas / Friaje", categorias)

    def test_indice_de_priorizacion_ponderado(self):
        serie = self._serie("f6_prioridad")
        valores = {f["clase"]: f["valor"] for f in serie["filas"]}
        self.assertEqual(valores["Movimientos en masa (deslizamientos)"], 3)
        self.assertEqual(valores["Erosión hídrica (cárcavas)"], 2)
        self.assertEqual(valores["Sequía / Déficit hídrico"], 1)

    def test_cobertura_marca_el_centro_poblado_sin_ficha(self):
        serie = self._serie("cob_catalogo")
        valores = {f["clase"]: f["valor"] for f in serie["filas"]}
        self.assertEqual(valores["Con al menos una ficha social"], 2)
        self.assertEqual(valores["Sin ficha social registrada"], 1)   # El Tambo

    def test_metricas_de_cabecera(self):
        metricas = {m["etiqueta"]: m["valor"] for m in self.informe["metricas"]}
        self.assertEqual(metricas["Actores mapeados"], "3")
        self.assertEqual(metricas["Conflictos identificados"], "3")
        self.assertEqual(metricas["Titulares dispuestos a participar"], "1 / 2")
        self.assertTrue(metricas["Superficie predial comprometida"].startswith("25"))

    def test_el_nivel_mas_alto_recibe_el_tono_mas_oscuro(self):
        """Las listas de la ficha van de mayor a menor: la rampa se invierte."""
        for id_, rampa in (("f1_percepciones", an.RAMPA_NEUTRA),
                           ("f6_frecuencia", an.RAMPA_CRITICA),
                           ("f6_magnitud", an.RAMPA_CRITICA)):
            colores = self._serie(id_)["colores"]
            orden = self._serie(id_)["orden_sub"]
            posiciones = [rampa.index(colores[s]) for s in orden]
            self.assertEqual(posiciones, sorted(posiciones, reverse=True), id_)

    def test_color_estable_entre_bloques(self):
        """Una clase conserva su color aunque cambien las clases presentes."""
        completo = self._serie("f6_magnitud")["colores"]
        parcial = an.indicadores_bloque(
            BLOQUE, [_registro("F-DS-06", {"f6_peligros": [
                {"Peligro observado": "Caída de rocas", "¿Ocurre?": "Sí",
                 "Magnitud": "A (Alta)", "Frecuencia": "F (Frecuente)"}]})],
            DATOS_CP)
        for seccion in parcial["secciones"]:
            for serie in seccion["series"]:
                if serie["id"] == "f6_magnitud":
                    self.assertEqual(serie["colores"]["A (Alta)"],
                                     completo["A (Alta)"])

    def test_filas_repetidas_se_acumulan_una_sola_vez(self):
        """Dos centros poblados con la misma actividad dan un tramo, no dos."""
        serie = self._serie("f1_actividades")
        df, categorias, subclases = an._datos_grafico(serie)
        self.assertEqual(len(df), len(df.drop_duplicates(["cat", "sub"])))
        self.assertEqual(
            df.loc[df["cat"] == "Agricultura de secano", "valor"].iloc[0], 180)

    def test_eje_de_conteos_usa_marcas_enteras(self):
        """Un eje que llega a 2 no puede rotularse '0 1 1 2'."""
        spec = an.grafico_altair(self._serie("f7_cpi")).to_dict()
        self.assertEqual(spec["encoding"]["x"]["axis"]["values"], [0, 1, 2])

    def test_la_leyenda_no_declara_clases_ausentes(self):
        """Una entrada sin tramo hace dudar de si el dato falta o vale cero."""
        serie = self._serie("f1_gobernanza")
        _df, _categorias, subclases = an._datos_grafico(serie)
        self.assertEqual(subclases, ["Sí", "No"])
        dominio = an.grafico_altair(serie).to_dict()[
            "encoding"]["color"]["scale"]["domain"]
        self.assertEqual(dominio, ["Sí", "No"])

    def test_escala_ordenada_conserva_su_orden_en_el_eje(self):
        """La disposicion a participar se lee de mejor a peor, no por conteo."""
        serie = self._serie("f7_disposicion")
        self.assertEqual(serie["orden_cat"][0], "Totalmente dispuesto/a")

    def test_bateria_sino_se_ordena_por_respuestas_afirmativas(self):
        """La brecha (todo 'No') debe quedar al final de la lista."""
        serie = self._serie("f7_cpi")
        _df, categorias, _sub = an._datos_grafico(serie)
        self.assertEqual(categorias[-1], "Se entregó material informativo")

    def test_colores_declarados_existen_en_la_paleta_validada(self):
        """Ningun grafico puede inventar un color fuera de las paletas."""
        permitidos = set(an.CATEGORICA_CLARA) | set(an.CATEGORICA_OSCURA)
        permitidos |= set(an.RAMPA_NEUTRA) | set(an.RAMPA_CRITICA)
        permitidos |= set(an.RAMPA_FAVORABLE) | set(an.DIVERGENTE)
        permitidos |= {an.SIN_DATO, "#C4A03C"}
        for seccion in self.informe["secciones"]:
            for serie in seccion["series"]:
                for color in (serie.get("colores") or {}).values():
                    self.assertIn(color, permitidos,
                                  f"{serie['id']} usa un color no validado")


class InformeParcial(unittest.TestCase):

    def test_bloque_sin_fichas(self):
        informe = an.indicadores_bloque(BLOQUE, [], DATOS_CP)
        self.assertEqual(informe["n_registros"], 0)
        self.assertEqual(informe["secciones"], [])
        self.assertTrue(informe["avisos"])

    def test_bloque_con_una_sola_ficha(self):
        informe = an.indicadores_bloque(
            BLOQUE, [_registro("F-DS-06", _fds06())], DATOS_CP)
        ids = {s["id"] for s in informe["secciones"]}
        self.assertEqual(ids, {"COBERTURA", "F-DS-06"})
        self.assertTrue(any("F-DS-01" in a for a in informe["avisos"]))

    def test_ficha_vacia_no_rompe_el_informe(self):
        informe = an.indicadores_bloque(
            BLOQUE, [_registro("F-DS-01", {}), _registro("F-DS-05", {})],
            DATOS_CP)
        self.assertIsInstance(informe["secciones"], list)

    def test_consolidado_agrega_el_reparto_por_bloque(self):
        registros = registros_completos()
        for r in registros[:3]:
            r["bloque_codigo"] = "M7-B4"
        informe = an.indicadores_consolidado(registros, etiqueta="2 bloques")
        self.assertEqual(informe["alcance"], "consolidado")
        self.assertEqual(sorted(informe["bloques"]), ["M5-B1", "M7-B4"])
        ids = {s["id"] for sec in informe["secciones"] for s in sec["series"]}
        self.assertIn("cons_bloque", ids)
        self.assertIn("cons_distrito", ids)


# ══════════════════════════════════════════════════════════════════════════
# Salidas: Excel, PDF y graficos interactivos
# ══════════════════════════════════════════════════════════════════════════

class Salidas(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.informe = an.indicadores_bloque(BLOQUE, registros_completos(),
                                            DATOS_CP)

    def test_excel_con_una_hoja_y_un_grafico_por_seccion(self):
        contenido = an.generar_excel_social(self.informe)
        wb = load_workbook(io.BytesIO(contenido))
        self.assertEqual(wb.sheetnames[0], "Resumen DS")
        con_graficos = [ws for ws in wb.worksheets if ws._charts]
        self.assertGreaterEqual(len(con_graficos), 7)
        total = sum(len(ws._charts) for ws in wb.worksheets)
        esperado = sum(len(s["series"]) for s in self.informe["secciones"])
        self.assertEqual(total, esperado)

    def test_excel_lleva_encabezado_institucional(self):
        wb = load_workbook(io.BytesIO(an.generar_excel_social(self.informe)))
        ws = wb["Resumen DS"]
        self.assertEqual(ws.cell(1, 1).value,
                         "AUTORIDAD NACIONAL DE INFRAESTRUCTURA - ANIN")
        self.assertIn("CUI 2669244", ws.cell(4, 1).value)

    def test_excel_totaliza_con_formulas(self):
        """Los totales van como formula para que el usuario pueda recalcular."""
        wb = load_workbook(io.BytesIO(an.generar_excel_social(self.informe)))
        formulas = [c.value for ws in wb.worksheets for fila in ws.iter_rows()
                    for c in fila if isinstance(c.value, str)
                    and c.value.startswith("=SUM(")]
        self.assertTrue(formulas)

    def test_excel_incluye_las_tablas_de_respaldo(self):
        wb = load_workbook(io.BytesIO(an.generar_excel_social(self.informe)))
        hojas = " | ".join(wb.sheetnames)
        self.assertIn("Actores", hojas)
        self.assertIn("Conflictos", hojas)

    def test_excel_del_bloque_vacio_es_valido(self):
        informe = an.indicadores_bloque(BLOQUE, [], DATOS_CP)
        wb = load_workbook(io.BytesIO(an.generar_excel_social(informe)))
        self.assertIn("Resumen DS", wb.sheetnames)

    def test_pdf_se_genera_y_es_un_pdf(self):
        contenido = an.generar_pdf_social(self.informe)
        self.assertTrue(contenido.startswith(b"%PDF"))
        self.assertGreater(len(contenido), 5000)

    def test_pdf_del_bloque_vacio_no_falla(self):
        informe = an.indicadores_bloque(BLOQUE, [], DATOS_CP)
        self.assertTrue(an.generar_pdf_social(informe).startswith(b"%PDF"))

    def test_nombres_de_archivo(self):
        self.assertTrue(an.nombre_excel(self.informe).startswith(
            "Graficos_DS_Bloque_M5-B1"))
        self.assertTrue(an.nombre_pdf(self.informe).endswith(".pdf"))

    def test_grafico_altair_de_cada_serie(self):
        """Toda serie debe producir una especificacion Vega-Lite valida."""
        for seccion in self.informe["secciones"]:
            for serie in seccion["series"]:
                grafico = an.grafico_altair(serie)
                self.assertIsNotNone(grafico, serie["id"])
                spec = grafico.to_dict()
                self.assertIn("$schema", spec)

    def test_grafico_altair_en_tema_oscuro(self):
        serie = self.informe["secciones"][0]["series"][0]
        self.assertIn("$schema", an.grafico_altair(serie, tema="oscuro").to_dict())

    def test_tabla_de_respaldo_de_cada_serie(self):
        """Cada grafico debe poder leerse tambien como tabla."""
        for seccion in self.informe["secciones"]:
            for serie in seccion["series"]:
                df = an.tabla_serie(serie)
                self.assertFalse(df.empty, serie["id"])


if __name__ == "__main__":
    unittest.main()
