# -*- coding: utf-8 -*-
"""
Ficha Técnica Predial completa (FP) ↔ aplicativo: formulario Kobo, ficha Excel llenada y ficha generada.
Cliente Kobo SIMULADO (sin red) y datos SINTÉTICOS (nombres, DNI, RUC y geometrías ficticios).
- Lógica pura: `python -m pytest tests/test_la_fp.py -q`
- Integración PostgreSQL: LA_TEST_DB="postgresql://…" (base VACÍA de pruebas: se borra el esquema public).
"""
import copy
import io
import os
import sys
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "tests"))
from liberacion_areas import la_core as core  # noqa: E402
from liberacion_areas import la_ficha_predial as fpx  # noqa: E402
from liberacion_areas import la_fp as fp  # noqa: E402
from liberacion_areas import la_kobo as kb  # noqa: E402
from liberacion_areas import la_xlsform_ficha as xf  # noqa: E402
import test_liberacion_areas as base  # noqa: E402
from test_liberacion_areas_kobo_sync import KoboFalso  # noqa: E402

E1, N1 = base.E0 + 300, base.N0 + 300          # dentro del bloque 27 sintético


def _jpeg(color=(30, 120, 60)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (640, 480), color).save(buf, format="JPEG")
    return buf.getvalue()


def _envio_fp(uuid: str, n_predio: int = 5, **extra) -> dict:
    """Envío del formulario FP tal como lo entrega la API v2 de Kobo (grupos, repeats y adjuntos)."""
    r = {"meta/instanceID": f"uuid:{uuid}", "_uuid": uuid, "_submission_time": "2026-10-20T10:00:00",
         "g_generales/hoy": "2026-10-20", "g_generales/asistente": "AP-1", "g_generales/unidad": "27",
         "g_generales/cod_unidad": "27", "g_generales/tipo_unidad": "bloque", "g_generales/provincia": "ayabaca",
         "g_generales/distrito": "frias", "g_generales/n_predio": str(n_predio),
         "g_generales/cod_predio": f"27-P{n_predio:02d}", "g_generales/fecha_insp": "2026-10-20",
         "g_generales/consentimiento": "si",
         "g_registral/sit_registral": "inscrito_en_sunarp", "g_registral/doc_partida": "P-0001",
         "g_registral/docs_exhibidos": "titulo dni",
         "g_prediante/tipo_titularidad": "posesionario", "g_prediante/personeria": "natural",
         "g_prediante/g_natural/tit_nombres": "PEDRO (FICTICIO)", "g_prediante/g_natural/tit_apellidos": "PRUEBA FP",
         "g_prediante/g_natural/tit_dni": "00000555", "g_prediante/g_natural/tit_celular": "987654321",
         "g_prediante/g_natural/tit_estado_civil": "soltero",
         "g_prediante/g_domicilio/dom_direccion": "CASERÍO FICTICIO S/N", "g_prediante/g_domicilio/dom_departamento": "piura",
         "g_prediante/g_domicilio/dom_provincia": "ayabaca", "g_prediante/g_domicilio/dom_distrito": "frias",
         "g_predio/predio_nombre": "EL ALGARROBO", "g_predio/metodo_coord": "gps",
         "g_predio/punto_interior": base.gp(E1, N1), "g_predio/predio_departamento": "piura",
         "g_predio/predio_provincia": "ayabaca", "g_predio/predio_distrito": "frias",
         "g_matriz/uso_actual": "matorral", "g_matriz/ocupacion": "pastoreo", "g_matriz/pendiente_clase": "empinada_2550",
         "g_matriz/erosion": "en_surcos", "g_matriz/tipo_predio": "rustico", "g_matriz/predio_area_decl_ha": "12.5",
         "g_colindancias/col_norte": "QUEBRADA SECA", "g_colindancias/col_norte_tipo": "quebrada_rio",
         "conflicto_linderos_det": "Sin conflictos",
         "g_area/predio_area_unidad_ha": "4.2", "g_area/interferencias": "ninguna", "g_area/tiene_afectaciones": "si",
         "r_afectaciones": [{"r_afectaciones/af_descripcion": "CERCO DE PALOS", "r_afectaciones/af_metrado": "120",
                             "r_afectaciones/af_und": "m", "r_afectaciones/af_material": "madera",
                             "r_afectaciones/af_conservacion": "regular"}],
         "g_liberacion/aceptacion": "favorable", "g_liberacion/conclusion_campo": "sin_restriccion",
         "g_liberacion/estado_la_propuesto": "LA-2",
         "r_puntos": [{"r_puntos/pt_gps": base.gp(E1 + 50, N1 + 50), "r_puntos/pt_tipo": "vertice"}],
         "_attachments": [
             {"question_xpath": "g_fotos/foto_1", "mimetype": "image/jpeg",
              "download_medium_url": f"https://kf.kobotoolbox.org/{uuid}_f1.jpg"},
             {"question_xpath": "g_firmas/firma_prediante", "mimetype": "image/png",
              "download_medium_url": f"https://kf.kobotoolbox.org/{uuid}_firma.png"}]}
    r.update(extra)
    return r


# ------------------------------------------------------------------ reconocimiento
def test_reconoce_formulario_fp():
    assert fp.es_formulario_fp(xf.TITULO)
    assert fp.es_formulario_fp("cualquier nombre", "aCyBhgqeCTZ3UKzmBW8tx2")
    assert kb.form_id_kobo(xf.TITULO) == fp.FORM_ID
    assert kb.form_id_kobo("F-LA-02 – Ficha predial y socio-territorial") == "f_la_02_titular"
    assert kb.form_id_kobo("F-LA-03 Inspección") == "f_la_03_inspeccion"
    assert kb.form_id_por_campos([_envio_fp("x")]) == fp.FORM_ID
    assert kb.codigo_formulario(fp.FORM_ID) == "FP" and kb.codigo_formulario("f_la_02_titular") == "F-LA-02"


# ------------------------------------------------------------------ ficha FP → F-LA-02 + F-LA-03
def test_partes_y_validacion():
    reg = core.aplanar(_envio_fp("fp-001"))
    p = fp.partes(reg)
    assert set(p) == {"f_la_02_titular", "f_la_03_inspeccion"}
    assert core.kobo_uuid(p["f_la_02_titular"]) == "fp-001-F02"
    assert core.kobo_uuid(p["f_la_03_inspeccion"]) == "fp-001-F03"
    assert [fp._campo_adjunto(a) for a in p["f_la_02_titular"]["_attachments"]] == ["firma_prediante"]
    assert [fp._campo_adjunto(a) for a in p["f_la_03_inspeccion"]["_attachments"]] == ["foto_1"]
    assert p["f_la_03_inspeccion"]["uso_observado"] == "matorral"
    assert p["f_la_02_titular"]["col_norte"] == "QUEBRADA SECA"          # los datos propios de la ficha se conservan
    res = fp.validar(reg, base.catalogo(), base.GEOMS, set())
    assert [r.form_id for r in res] == ["f_la_02_titular", "f_la_03_inspeccion"]
    assert all(r.estado_import == "NUEVO" for r in res), [r.motivos for r in res]
    assert res[1].datos["validacion_espacial"] == "DENTRO"
    # existentes → DUPLICADO; eliminados → se omiten
    res = fp.validar(reg, base.catalogo(), base.GEOMS, {"fp-001-F02"}, omitir={"fp-001-F03"})
    assert [(r.form_id, r.estado_import) for r in res] == [("f_la_02_titular", "DUPLICADO")]


def test_utm_digitado_y_estatal():
    reg = core.aplanar(_envio_fp("fp-002", **{"g_predio/metodo_coord": "utm", "g_predio/punto_interior": None,
                                               "g_predio/utm_este": str(E1), "g_predio/utm_norte": str(N1),
                                               "g_prediante/tipo_titularidad": "estatal",
                                               "g_prediante/personeria": "juridica",
                                               "g_prediante/g_juridica/cc_nombre": "ENTIDAD FICTICIA",
                                               "g_prediante/g_juridica/cc_presidente": "JEFE FICTICIO"}))
    p = fp.partes(reg)
    pt = core.parse_geopoint(p["f_la_03_inspeccion"]["punto_interior"])
    assert abs(pt["este"] - E1) < 0.5 and abs(pt["norte"] - N1) < 0.5
    assert p["f_la_02_titular"]["est_entidad"] == "ENTIDAD FICTICIA"
    r = fp.validar(reg, base.catalogo(), base.GEOMS, set())[1]
    assert r.datos["validacion_espacial"] == "DENTRO"


# ------------------------------------------------------------------ ficha Excel: generar llena ↔ leer
def _ficha_llena(uuids=None) -> bytes:
    p = fp.partes(core.aplanar(_envio_fp("fp-003")))
    cat = fpx.cargar_catalogo()
    return fp.ficha_excel_llena(cat, [p["f_la_02_titular"], p["f_la_03_inspeccion"]],
                                {"foto_1": _jpeg(), "croquis": _jpeg((200, 200, 200))}, uuids)


def test_ficha_excel_ida_y_vuelta():
    contenido = _ficha_llena()
    wb = load_workbook(io.BytesIO(contenido))
    ws = wb["FICHA PREDIAL"]
    mapa = fpx.mapa_celdas()
    assert ws[mapa["asistente"]].value == "AP-1" and ws[mapa["unidad"]].value == "27"
    assert ws[mapa["tipo_titularidad"]].value == "Posesionario" and ws[mapa["personeria"]].value == "Natural"
    assert ws[mapa["predio_provincia"]].value == "Ayabaca" and ws[mapa["predio_distrito"]].value == "Frías"
    assert ws[mapa["pendiente_clase"]].value == "Empinada (25–50 %)"
    assert abs(ws[mapa["utm_este"]].value - E1) < 0.5
    assert ws[mapa["af_material_1"]].value == "Madera" and ws[mapa["pt_tipo_1"]].value == "Vértice"
    assert len(ws._images) == 2                                       # foto 1 y croquis
    leido = fp.leer_ficha_excel(io.BytesIO(contenido), "ficha.xlsx")
    reg = leido["registro"]
    for k, v in {"asistente": "AP-1", "unidad": "27", "n_predio": 5, "tipo_titularidad": "posesionario",
                 "personeria": "natural", "tit_dni": "00000555", "tit_celular": "987654321", "predio_provincia": "ayabaca",
                 "predio_distrito": "frias", "uso_actual": "matorral", "pendiente_clase": "empinada_2550",
                 "col_norte_tipo": "quebrada_rio", "docs_exhibidos": "titulo dni", "dom_distrito": "frias",
                 "aceptacion": "favorable", "fecha_insp": "2026-10-20", "metodo_coord": "utm"}.items():
        assert reg.get(k) == v, (k, reg.get(k))
    assert reg["r_afectaciones"][0]["af_material"] == "madera" and reg["r_afectaciones"][0]["af_metrado"] == 120
    assert reg["r_puntos"][0]["pt_tipo"] == "vertice"
    assert not [a for a in leido["avisos"] if a.startswith("Falta")], leido["avisos"]
    res = fp.validar(reg, base.catalogo(), base.GEOMS, set())
    assert [r.estado_import for r in res] == ["NUEVO", "NUEVO"], [r.motivos for r in res]
    assert res[1].datos["validacion_espacial"] == "DENTRO"
    # el mismo archivo da el mismo identificador (re-importar → DUPLICADO)
    again = fp.leer_ficha_excel(io.BytesIO(contenido))["registro"]
    assert core.kobo_uuid(again) == core.kobo_uuid(reg)


def test_ficha_generada_por_el_aplicativo_no_duplica():
    contenido = _ficha_llena({"f_la_02_titular": "uuid-a02", "f_la_03_inspeccion": "uuid-a03"})
    leido = fp.leer_ficha_excel(io.BytesIO(contenido))
    res = fp.validar(leido["registro"], base.catalogo(), base.GEOMS, {"uuid-a02", "uuid-a03"})
    assert [r.estado_import for r in res] == ["DUPLICADO", "DUPLICADO"]


def test_ficha_en_blanco_y_sin_hoja_mapa(tmp_path):
    vacia = fpx.generar_ficha_predial(tmp_path / "f.xlsx")
    leido = fp.leer_ficha_excel(str(vacia))
    assert "Falta Asistente predial" in leido["avisos"] and "Falta Código de la unidad" in leido["avisos"]
    # ficha llenada con una copia sin la hoja _MAPA (p. ej. guardada con otro programa): se usa el diseño vigente
    wb = load_workbook(io.BytesIO(_ficha_llena()))
    del wb[fpx.HOJA_MAPA]
    buf = io.BytesIO()
    wb.save(buf)
    reg = fp.leer_ficha_excel(io.BytesIO(buf.getvalue()))["registro"]
    assert reg["unidad"] == "27" and reg["tit_nombres"] == "PEDRO (FICTICIO)"


def test_datos_adicionales_sin_reservados():
    p = fp.partes(core.aplanar(_envio_fp("fp-004")))["f_la_02_titular"]
    extra = fp.datos_adicionales(p)
    valores = dict(zip(extra.campo, extra.valor))
    assert valores["Colindante NORTE"] == "QUEBRADA SECA"
    assert valores["Tipo de colindante NORTE"] == "Quebrada / río"
    assert "987654321" not in " ".join(extra.valor) and "00000555" not in " ".join(extra.valor)
    assert fp.tabla_afectaciones(p).iloc[0]["material"] == "Madera"


# ------------------------------------------------------------------ integración: sincronización Kobo → base → ficha
def test_integracion_sincronizacion_fp():
    url = os.environ.get("LA_TEST_DB")
    if not url:
        return
    base.test_integracion_postgres()              # esquema limpio + catálogo + envíos F-LA de prueba
    from liberacion_areas import la_db as db
    conn = db.conectar(url)
    with conn.cursor() as cur:
        for cod, g in base.GEOMS.items():
            cur.execute("UPDATE la_unidades SET geom_wkt=%s WHERE codigo=%s", (g.wkt, cod))
    conn.commit()
    juridica = _envio_fp("fp-102", n_predio=6, **{
        "g_prediante/tipo_titularidad": "privado", "g_prediante/personeria": "juridica",
        "g_prediante/g_natural/tit_nombres": None, "g_prediante/g_natural/tit_apellidos": None,
        "g_prediante/g_natural/tit_dni": None, "g_prediante/g_juridica/cc_nombre": "EMPRESA FICTICIA SAC",
        "g_prediante/g_juridica/cc_ruc": "20000000001", "g_prediante/g_juridica/cc_presidente": "GERENTE FICTICIO"})
    forms = [{"uid": "aCyBhgqeCTZ3UKzmBW8tx2", "nombre": xf.TITULO, "envios": 2, "desplegado": True},
             {"uid": "aOtro", "nombre": "Encuesta ajena", "envios": 3, "desplegado": True}]
    kobo = KoboFalso(forms, {"aCyBhgqeCTZ3UKzmBW8tx2": [_envio_fp("fp-101"), juridica]})

    rev = kb.revisar_envios(kobo, conn)
    assert [(f["form_id"], len(f["pendientes"])) for f in rev] == [(fp.FORM_ID, 2)]
    res = kb.sincronizar_envios(kobo, conn, "pytest", revision=rev)
    assert res[0]["Formulario"] == "FP" and res[0]["Importados"] == 4 and res[0]["Adjuntos"] == 4, res
    assert set(db.listar_envios(conn)["kobo_uuid"]) >= {"fp-101-F02", "fp-101-F03", "fp-102-F02", "fp-102-F03"}
    firma = db.adjuntos_envio(conn, "fp-101-F02")
    assert list(firma["campo"]) == ["firma_prediante"]
    assert list(db.adjuntos_envio(conn, "fp-101-F03")["campo"]) == ["foto_1"]
    tit = db.df(conn, "SELECT nombre, tipo_titularidad FROM la_titulares WHERE dni_ruc = '20000000001'")
    assert tit["nombre"][0] == "EMPRESA FICTICIA SAC"
    insp = db.df(conn, "SELECT cod_predio, validacion_espacial FROM la_inspecciones WHERE cod_predio IN ('27-P05','27-P06')")
    assert set(insp.validacion_espacial) == {"DENTRO"} and len(insp) == 2
    pred = db.df(conn, "SELECT nombre_predio, area_decl_ha, uso_actual FROM la_predios WHERE cod_predio='27-P05'")
    assert pred["nombre_predio"][0] == "EL ALGARROBO" and float(pred["area_decl_ha"][0]) == 12.5
    # al día: nada pendiente; una parte eliminada en el aplicativo no se vuelve a traer
    assert all(not f["pendientes"] for f in kb.revisar_envios(kobo, conn))
    db.eliminar_envio(conn, "fp-101-F03", "pytest")
    assert all(not f["pendientes"] for f in kb.revisar_envios(kobo, conn))

    # ficha Excel llena desde la base (con la foto de Kobo) y re-importación sin duplicar
    with conn.cursor() as cur:
        cur.execute("UPDATE la_adjuntos SET contenido=%s WHERE kobo_uuid='fp-102-F03'", (_jpeg(),))
    conn.commit()
    xls = fp.ficha_predio_bytes(conn, "27-P06")
    ws = load_workbook(io.BytesIO(xls))["FICHA PREDIAL"]
    mapa = fpx.mapa_celdas()
    assert ws[mapa["cc_nombre"]].value == "EMPRESA FICTICIA SAC" and ws[mapa["personeria"]].value == "Jurídica"
    assert ws[mapa["predio_nombre"]].value == "EL ALGARROBO" and len(ws._images) == 1
    leido = fp.leer_ficha_excel(io.BytesIO(xls))
    cat, geoms = db.catalogo(conn), db.geometrias(conn)
    res = fp.validar(leido["registro"], cat, geoms, db.uuids_existentes(conn))
    assert [r.estado_import for r in res] == ["DUPLICADO", "DUPLICADO"]
    # una ficha Excel nueva (otro predio) entra como F-LA-02 + F-LA-03 con origen PLANTILLA
    reg = copy.deepcopy(leido["registro"])
    reg.pop("_uuid_partes")
    reg.update({"n_predio": 7, "cod_predio": "27-P07", "meta_instanceID": "xls-fp-prueba"})
    n = fp.importar(conn, fp.validar(reg, cat, geoms, db.uuids_existentes(conn)), "PLANTILLA", "pytest")
    assert n == 2
    origenes = db.df(conn, "SELECT origen FROM la_envios_raw WHERE cod_predio='27-P07'")
    assert set(origenes.origen) == {"PLANTILLA"}
    conn.close()
