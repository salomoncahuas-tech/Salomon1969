"""Importacion de la plantilla de Diagnostico Social V4 llenada en campo.

Verifica que:
  * la especificacion celda a celda del importador coincide con la plantilla
    oficial (casillas vacias con su rotulo a la derecha, celdas de valor
    vacias), para que un cambio de plantilla no rompa la importacion en
    silencio;
  * una plantilla en blanco no produce fichas «con datos» (antes se tomaban
    los rotulos vecinos como valores);
  * una plantilla llenada (texto, casillas «X», desplegables y codigos de la
    hoja _Codigos) llega completa a las claves del formulario del aplicativo.
"""

import io
import os
import sys
import unittest
import warnings
from datetime import date, datetime

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import excel_diagnostico_social as eds  # noqa: E402
import fds_listas as FL  # noqa: E402


def _libro():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(eds.PLANTILLA_DS_PATH)


def _bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _por_ficha(res):
    return {r["ficha"]: r for r in res}


def _celda_izq_derecha(coord):
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
    r, c = coordinate_to_tuple(coord)
    return f"{get_column_letter(c + 1)}{r}"


class EspecificacionCoincideConPlantilla(unittest.TestCase):

    def test_casillas_vacias_y_con_rotulo(self):
        wb = _libro()
        for spec in (eds.UNICA, eds.MULTIPLE):
            for hoja, campos in spec.items():
                ws = wb[hoja]
                for clave, opciones in campos.items():
                    for celda, _valor in opciones:
                        with self.subTest(hoja=hoja, clave=clave, celda=celda):
                            self.assertIsNone(ws[celda].value)
                            rotulo = ws[_celda_izq_derecha(celda)].value
                            self.assertTrue(rotulo and str(rotulo).strip(),
                                            f"sin rotulo a la derecha de {celda}")

    def test_valores_de_opciones_existen_en_el_aplicativo(self):
        listas = [getattr(FL, n) for n in dir(FL) if n.isupper()]
        todas = {v for lst in listas if isinstance(lst, list) for v in lst}
        for spec in (eds.UNICA, eds.MULTIPLE):
            for campos in spec.values():
                for clave, opciones in campos.items():
                    for _celda, valor in opciones:
                        with self.subTest(clave=clave, valor=valor):
                            self.assertIn(valor, todas)

    def test_celdas_de_texto_vacias(self):
        wb = _libro()
        grupos = [eds.TEXTOS, eds.TEXTOS_LARGOS, eds.GENERALES]
        for g in grupos:
            for hoja, campos in g.items():
                for clave, celdas in campos.items():
                    for celda in ([celdas] if isinstance(celdas, str) else celdas):
                        with self.subTest(hoja=hoja, clave=clave, celda=celda):
                            self.assertIsNone(wb[hoja][celda].value)
        for hoja, cab in eds.CABECERA.items():
            for celda in cab.values():
                with self.subTest(hoja=hoja, celda=celda):
                    self.assertIsNone(wb[hoja][celda].value)


class DesplegablesFDS05(unittest.TestCase):

    def test_estado_y_antiguedad_en_columnas_visibles(self):
        import zipfile
        with zipfile.ZipFile(eds.PLANTILLA_DS_PATH) as zf:
            xml = zf.read("xl/worksheets/sheet6.xml").decode("utf-8")
        self.assertIn("<xm:sqref>F14 F15 F16 F17 F18 F19 F20 F21</xm:sqref>", xml)
        self.assertIn("<xm:sqref>G14 G15 G16 G17 G18 G19 G20 G21</xm:sqref>", xml)


class PlantillaEnBlanco(unittest.TestCase):

    def test_no_detecta_fichas_ni_toma_rotulos(self):
        res = eds.parsear_excel_ds(_bytes(_libro()))
        self.assertEqual(res, [])


class PlantillaLlenada(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        wb = _libro()
        f1 = wb["F-DS-01"]
        f1["D7"] = datetime(2026, 9, 15)
        f1["G7"] = "007"
        f1["K7"] = "Juan Pérez"
        f1["D8"] = "M9B1"
        f1["I9"] = "Pampa Grande"
        f1["E11"] = 612345.4
        f1["I11"] = 9435678
        f1["M11"] = 1450
        f1["D12"] = "Rosa Chávez"
        f1["F19"] = "Caserío Pampa Grande"
        f1["E21"] = 85
        f1["L21"] = 340
        f1["L24"] = 42
        f1["H16"] = "X"            # Centro poblado menor
        f1["A27"] = "x"            # Bilingüe
        f1["D40"] = "X"            # Junta vigente: Sí
        f1["F42"] = "X"            # Ronda: No
        f1["H49"] = "X"            # Reglamento: N/A
        f1["A67"] = "X"            # Agua: red publica
        f1["H68"] = "X"            # Agua: manantial
        f1["D84"] = "Movistar X"   # operador (sin casilla propia)
        f1["L84"] = "Otros: Tuenti"
        f1["F88"] = "X"            # IE: Inicial + Primaria + Secundaria
        f1["H92"] = "X"            # EESS: III
        f1["J127"] = "X"           # ONGs: No
        f1["B100"] = "Ganadería vacuna"
        f1["E100"] = 12
        f1["K100"] = "X"           # destino mixto
        f1["C101"] = "A07"         # actividad por codigo del catalogo
        f1["A133"] = "Sin novedades"

        f2 = wb["F-DS-02"]
        f2["C7"] = "15-09-2026"
        f2["E7"] = "Ana Ruiz"
        f2["J7"] = 22              # bloque numerico
        f2["B16"] = "Municipalidad Distrital"
        f2["C16"] = "GL"
        f2["E16"] = "A"
        f2["F16"] = "M"
        f2["G16"] = "C"
        f2["H16"] = "D"
        f2["F44"] = "Alcalde"

        f3 = wb["F-DS-03"]
        f3["D7"] = datetime(2026, 9, 16)
        f3["J7"] = "Ana Ruiz"
        f3["D8"] = "M9B1"
        f3["D10"] = "Luis Torres"
        f3["M11"] = "Masculino"
        f3["H16"] = "X"
        f3["A20"] = "Hualtaco y palo santo"
        f3["A21"] = "han disminuido"

        f4 = wb["F-DS-04"]
        f4["C7"] = datetime(2026, 9, 17)
        f4["H7"] = "Ana Ruiz"
        f4["C8"] = "M9B1"
        f4["D13"] = datetime(2026, 9, 17)
        f4["A20"] = "X"
        f4["H22"] = "X"
        f4["A25"] = "X"            # Papelógrafos (rotulo vacio en V4)
        f4["H31"] = "X"
        f4["B35"] = "Pedro López"
        f4["K35"] = "Mujer"
        f4["C58"] = "Presentación"
        f4["B66"] = "Siguiente reunión"

        f5 = wb["F-DS-05"]
        f5["D7"] = datetime(2026, 9, 18)
        f5["J7"] = "Ana Ruiz"
        f5["D8"] = "M9B1"
        f5["B14"] = "SH"
        f5["C14"] = "Comuneros"
        f5["F14"] = "LT"
        f5["G14"] = 2
        f5["M26"] = "X"            # rb3: N/A
        f5["K32"] = "X"            # polarizacion: Medio
        f5["H52"] = "X"            # viabilidad: Media
        f5["B36"] = "Mesa técnica"
        f5["J36"] = "A"

        f6 = wb["F-DS-06"]
        f6["D7"] = datetime(2026, 9, 19)
        f6["J7"] = "Ana Ruiz"
        f6["D8"] = "M9B1"
        f6["C15"] = "Sí"
        f6["E15"] = "MF (Muy frecuente)"
        f6["G15"] = "A (Alta)"
        f6["I15"] = "Sube (Aumentando)"
        f6["K15"] = 2023
        f6["D16"] = "X"            # huaicos: No (casilla)
        f6["F17"] = "X"            # erosion: Rara (casilla)
        f6["E35"] = "Sí"
        f6["H35"] = "X"            # intensidad media (casilla)
        f6["H45"] = "X"
        f6["O46"] = "No"           # alternativa por desplegable
        f6["F52"] = "Huaicos"

        f7 = wb["F-DS-07"]
        f7["D7"] = datetime(2026, 9, 20)
        f7["J7"] = "Ana Ruiz"
        f7["D8"] = "M9B1"
        f7["H11"] = "X"            # Posesionario
        f7["D14"] = "María Quispe"
        f7["F15"] = "Mujer X"
        f7["A20"] = "X"            # Título COFOPRI
        f7["M27"] = "Sí"
        f7["N28"] = "X"            # punto 3.2: No
        f7["A40"] = "X"            # dispuesto con condiciones
        f7["A53"] = "¿Cuándo empiezan?"

        cls.res = _por_ficha(eds.parsear_excel_ds(_bytes(wb)))

    def test_detecta_las_siete_fichas(self):
        self.assertEqual(sorted(self.res), eds.FICHAS_HOJAS)

    def test_cabecera_exacta(self):
        d = self.res["F-DS-01"]["datos"]
        self.assertEqual(d["fecha"], "2026-09-15")
        self.assertEqual(d["evaluador"], "Juan Pérez")
        self.assertEqual(d["codigo_bloque"], "M9B1")
        d2 = self.res["F-DS-02"]["datos"]
        self.assertEqual((d2["fecha"], d2["evaluador"], d2["codigo_bloque"]),
                         ("2026-09-15", "Ana Ruiz", "22"))

    def test_fds01_campos(self):
        d = self.res["F-DS-01"]["datos"]
        f = d["form"]
        self.assertEqual(d["generales"]["ds_fnum"], "007")
        self.assertEqual(d["generales"]["ds_cpob"], "Pampa Grande")
        self.assertEqual(d["generales"]["ds_obs"], "Sin novedades")
        self.assertEqual(f["f1_nfam"], "85")
        self.assertEqual(f["f1_mano_obra"], "42")
        self.assertEqual(f["f1_org_terr"], "Centro poblado menor")
        self.assertEqual(f["f1_idioma"], "Bilingüe español-quechua")
        self.assertEqual(f["f1_junta_vig"], "Sí")
        self.assertEqual(f["f1_ronda"], "No")
        self.assertEqual(f["f1_reglamento"], "No aplica")
        self.assertEqual(f["f1_agua"], ["Red pública domiciliaria", "Manantial / Quebrada directa"])
        self.assertEqual(f["f1_telecom_op"], "Movistar, Tuenti")
        self.assertEqual(f["f1_ie_niveles"], "Inicial + Primaria + Secundaria")
        self.assertEqual(f["f1_eess"], "III")
        self.assertEqual(f["f1_ongs"], "No")
        act = f["f1_activ"]
        self.assertEqual(act[0]["Actividad / Rubro"], "Ganadería vacuna")
        self.assertEqual(act[0]["Destino"], "Mixto autoconsumo/mercado (30-70%)")
        self.assertEqual(act[1]["Actividad / Rubro"], "Apicultura")

    def test_fds02_codigos(self):
        f = self.res["F-DS-02"]["datos"]["form"]
        a = f["f2_actores"][0]
        self.assertEqual(a["Tipo"], "Gobierno Local (Municipalidad)")
        self.assertEqual((a["Influencia"], a["Interes"]), ("Alto", "Medio"))
        self.assertEqual(a["Posicion"], "En contra del proyecto")
        self.assertEqual(a["Nivel territorial"], "Distrital")
        self.assertEqual(f["f2_favor"], "Alcalde")

    def test_fds03(self):
        f = self.res["F-DS-03"]["datos"]["form"]
        self.assertEqual(f["f3_nombre"], "Luis Torres")
        self.assertEqual(f["f3_genero"], "M")
        self.assertEqual(f["f3_c_nom"], "Sí")
        self.assertEqual(f["f3_r1"], "Hualtaco y palo santo\nhan disminuido")

    def test_fds04(self):
        f = self.res["F-DS-04"]["datos"]["form"]
        self.assertEqual(f["f4_fecha"], "17/09/2026")
        self.assertEqual(f["f4_metod"], ["Exposicion magistral", "Lluvia de ideas"])
        self.assertEqual(f["f4_mater"], ["Papelografos"])
        self.assertEqual(f["f4_idioma"], "Bilingue español-quechua")
        self.assertEqual(f["f4_part"][0]["Sexo"], "F")
        self.assertEqual(f["f4_agenda"][0]["Agenda"], "Presentación")
        self.assertEqual(f["f4_acuerdos"][0]["Acuerdo / Compromiso"], "Siguiente reunión")

    def test_fds05_codigos(self):
        f = self.res["F-DS-05"]["datos"]["form"]
        c = f["f5_conflictos"][0]
        self.assertEqual(c["Tipo"], "Socioambiental hídrico")
        self.assertEqual(c["Estado"], "Latente (no manifiesto)")
        self.assertEqual(c["Antiguedad"], "1-3 años")
        self.assertEqual(f["f5_rb3"], "No aplica")
        self.assertEqual(f["f5_polar"], "Medio")
        self.assertEqual(f["f5_viab"], FL.FDS05_VIABILIDAD[1])
        self.assertEqual(f["f5_oportunidades"][0]["Potencial"], "Alto")

    def test_fds06(self):
        f = self.res["F-DS-06"]["datos"]["form"]
        p = {r["Peligro observado"]: r for r in f["f6_peligros"]}
        self.assertEqual(len(f["f6_peligros"]), 11)          # sin la fila «Otro» vacia
        mm = p["Movimientos en masa (deslizamientos)"]
        self.assertEqual((mm["¿Ocurre?"], mm["Frecuencia"], mm["Magnitud"], mm["Tendencia"]),
                         ("Sí", "MF (Muy frecuente)", "A (Alta)", "Sube (Aumentando)"))
        self.assertEqual(mm["Ultimo evento (año)"], "2023")
        self.assertEqual(p["Huaycos / Flujos de detritos"]["¿Ocurre?"], "No")
        self.assertEqual(p["Erosión hídrica (cárcavas)"]["Frecuencia"], "R (Rara)")
        cambio = f["f6_cambios"][0]
        self.assertEqual((cambio["¿Se percibe?"], cambio["Intensidad"]), ("Sí", "Media"))
        self.assertEqual(f["f6_medidas"], "Sí")
        self.assertEqual(f["f6_alerta"], "No")
        self.assertEqual(f["f6_p1"], "Huaicos")

    def test_fds07(self):
        f = self.res["F-DS-07"]["datos"]["form"]
        self.assertEqual(f["f7_tipo_prop"], "Posesionario (sin título)")
        self.assertEqual(f["f7_genero"], "Mujer")
        self.assertEqual(f["f7_docs"], ["Titulo COFOPRI"])
        self.assertEqual(f["f7_info_anin"], "Sí")
        self.assertEqual(f["f7_info_objetivo"], "No")
        self.assertEqual(f["f7_disp"], "Dispuesto/a con condiciones")
        self.assertEqual(f["f7_preg"], "¿Cuándo empiezan?")

    def test_mapeo_a_formulario(self):
        bm = {"M9B1": 1, "22": 2}
        res = self.res["F-DS-01"]
        pend = eds.mapear_a_session_state(res, bm, date(2024, 1, 1), date(2026, 10, 8))
        self.assertEqual(pend["ds_ficha_sel"], "F-DS-01")
        self.assertEqual(pend["ds_bl"], "M9B1")
        self.assertEqual(pend["ds_fecha"], date(2026, 9, 15))
        self.assertEqual(pend["ds_eval"], "Juan Pérez")
        self.assertEqual(pend["ds_fnum"], "007")
        self.assertEqual(pend["ds_este"], 612345.4)
        self.assertEqual(pend["ds_norte"], 9435678.0)
        self.assertEqual(pend["f1_org_terr"], "Centro poblado menor")
        self.assertIn("_dsinit_f1_activ", pend)
        self.assertEqual(res["avisos"], [])
        pend2 = eds.mapear_a_session_state(self.res["F-DS-02"], bm)
        self.assertEqual(pend2["ds_bl"], "22")

    def test_mapeo_avisa_bloque_inexistente_y_fecha_fuera_de_rango(self):
        res = {"ficha": "F-DS-03", "datos": {"fecha": "2019-01-01", "evaluador": "X",
                                             "codigo_bloque": "M99B9", "form": {}},
               "avisos": []}
        pend = eds.mapear_a_session_state(res, {"M9B1": 1})
        self.assertNotIn("ds_bl", pend)
        self.assertNotIn("ds_fecha", pend)
        self.assertEqual(len(res["avisos"]), 2)


if __name__ == "__main__":
    unittest.main()
