"""
IN Piura - Resumenes Excel de Diagnostico Territorial por bloque.

Lee los libros "Plantilla_Excel_Bloque_<codigo>_IN_Piura.xlsx" (117 bloques
con ficha DT) que se generan a partir de las fichas F-DT-01 a F-DT-05, el
catalogo maestro Bloques V5/V6, la estadistica zonal sobre el MDE y los
compuestos Sentinel-2 (MSAVI 2024 / NDVI mediana 2025).

Cada libro trae cinco hojas:
  1. Resumen                 identificacion, parametros fisicos, ecosistema
  2. Cobertura MSAVI-NDVI    clases espectrales y metrado en hectareas
  3. Estaciones fotograficas puntos georreferenciados de la ficha DT
  4. Microcuenca             contexto intramicrocuenca del bloque
  5. Control de consistencia discrepancias campo / gabinete / catalogo

El parseo es dirigido por etiquetas y no por coordenadas de celda: los 117
libros varian en numero de filas (discrepancias, estaciones y bloques por
microcuenca son de largo variable), de modo que fijar posiciones haria el
lector fragil. Se localizan las cabeceras de tabla y las etiquetas conocidas,
y se lee lo contiguo.

Criterio de integridad declarativa del proyecto: no se estima ni se infiere
ningun valor ausente. Los textos "Por determinar", "Por verificar" y "Sin
registro en ficha" se conservan tal cual.
"""

import io
import json
import os
import re
import unicodedata
import zipfile
from datetime import datetime

from openpyxl import load_workbook, Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Identidad institucional ANIN ──────────────────────────────────────────
ANIN_VERDE = "1B4D2E"
ANIN_AZUL = "1B4F72"
ANIN_VERDE_CLARO = "E8F0EA"
ANIN_GRIS = "F2F2F2"

ENCABEZADOS_ANIN = [
    "AUTORIDAD NACIONAL DE INFRAESTRUCTURA - ANIN",
    "DIRECCION DE INTERVENCIONES MULTISECTORIALES Y DE EMERGENCIA - DIME",
    "SUBDIRECCION DE ESTUDIOS DE INVERSION - SESDI",
]
SUBTITULO_PROYECTO = (
    "PROYECTO IN PIURA | CUI 2669244 | Recuperacion del servicio de regulacion "
    "de riesgos naturales y de ecosistemas degradados - Cuenca Alta del Rio Piura"
)

# Umbral de brecha de degradacion (R.M. N.o 00213-2024-MINAM).
UMBRAL_MSAVI = 0.4976

# Validacion UTM WGS 84 Zona 17S (EPSG:32717) para el ambito del proyecto.
UTM_ESTE_MIN, UTM_ESTE_MAX = 450000, 750000
UTM_NORTE_MIN, UTM_NORTE_MAX = 9300000, 9600000

CALIFICACIONES = ["SUSTANTIVA", "NO SUSTANTIVA", "CORREGIDO", "CONFORME"]

# Limite de barrido por hoja. Evita que un rango usado inflado por Excel
# obligue a recorrer decenas de miles de filas vacias por archivo.
MAX_FILAS = 400
MAX_COLUMNAS = 20


# ══════════════════════════════════════════════════════════════════════════
# Utilidades de normalizacion
# ══════════════════════════════════════════════════════════════════════════

def _norm(valor):
    """Normaliza texto para comparar etiquetas: sin acentos, sin signos
    de puntuacion decorativos, minusculas y espacios colapsados."""
    if valor is None:
        return ""
    txt = str(valor)
    txt = unicodedata.normalize("NFKD", txt)
    txt = "".join(c for c in txt if not unicodedata.combining(c))
    txt = txt.replace("—", " ").replace("–", " ").replace("-", " ")
    txt = txt.replace("º", "").replace("°", "")
    txt = re.sub(r"[^0-9a-zA-Z%/().,<>= ]+", " ", txt)
    return re.sub(r"\s+", " ", txt).strip().lower()


def _txt(valor):
    """Texto limpio conservando acentos y mayusculas del original."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return re.sub(r"\s+", " ", str(valor)).strip()


def _num(valor):
    """Convierte a float. Devuelve None si no hay numero interpretable.

    No inventa valores: 'Por determinar', '—' y celdas vacias dan None.
    """
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    txt = str(valor).strip()
    if not txt:
        return None
    # Descarta separadores de miles y toma el primer numero con signo.
    txt = txt.replace(" ", " ").replace(" ", "")
    m = re.search(r"[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?|[+-]?\d*\.?\d+", txt)
    if not m:
        return None
    try:
        return float(m.group(0).replace(",", ""))
    except ValueError:
        return None


# ══════════════════════════════════════════════════════════════════════════
# Rejilla acotada de hoja
# ══════════════════════════════════════════════════════════════════════════

class _Rejilla:
    """Copia acotada de una hoja a una matriz de valores.

    Se materializa una sola vez por hoja para no re-recorrer el iterador de
    openpyxl en cada busqueda de etiqueta.
    """

    def __init__(self, ws, max_filas=MAX_FILAS, max_columnas=MAX_COLUMNAS):
        self.titulo = ws.title
        self.filas = []
        for i, fila in enumerate(ws.iter_rows(max_row=max_filas,
                                              max_col=max_columnas,
                                              values_only=True)):
            self.filas.append(list(fila))
            if i + 1 >= max_filas:
                break
        self.n_filas = len(self.filas)
        self.n_columnas = max((len(f) for f in self.filas), default=0)

    def valor(self, fila, col):
        """Valor en indices base 0. Fuera de rango devuelve None."""
        if 0 <= fila < self.n_filas:
            f = self.filas[fila]
            if 0 <= col < len(f):
                return f[col]
        return None

    def fila_vacia(self, fila):
        return all(_txt(v) == "" for v in (self.filas[fila]
                                           if 0 <= fila < self.n_filas else []))

    def buscar_etiqueta(self, *variantes, columna_max=6):
        """Ubica la primera celda cuyo texto normalizado empieza por alguna
        de las variantes dadas. Devuelve (fila, col) o None."""
        objetivos = [_norm(v) for v in variantes if _norm(v)]
        for r in range(self.n_filas):
            for c in range(min(self.n_columnas, columna_max)):
                celda = _norm(self.valor(r, c))
                if not celda:
                    continue
                for obj in objetivos:
                    if celda == obj or celda.startswith(obj):
                        return (r, c)
        return None

    def buscar_fila_cabecera(self, *columnas_clave):
        """Ubica la fila que actua como cabecera de una tabla: aquella que
        contiene todas las columnas clave indicadas.

        Devuelve (fila, mapa) donde `mapa` indexa TODAS las columnas de esa
        cabecera por su texto normalizado, no solo las de busqueda, para que
        el lector pueda pedir cualquier columna de la tabla.
        """
        objetivos = [_norm(c) for c in columnas_clave]
        for r in range(self.n_filas):
            mapa = {}
            for c in range(self.n_columnas):
                celda = _norm(self.valor(r, c))
                if celda and celda not in mapa:
                    mapa[celda] = c
            if all(any(t == obj or t.startswith(obj) for t in mapa)
                   for obj in objetivos):
                return (r, mapa)
        return None

    @staticmethod
    def columna(mapa, nombre):
        """Indice de una columna de cabecera por prefijo normalizado."""
        objetivo = _norm(nombre)
        if objetivo in mapa:
            return mapa[objetivo]
        for texto, indice in mapa.items():
            if texto.startswith(objetivo):
                return indice
        return None


# ══════════════════════════════════════════════════════════════════════════
# Lectura de la hoja 1: Resumen
# ══════════════════════════════════════════════════════════════════════════

# Etiqueta en la hoja -> clave del diccionario de salida. Las etiquetas se
# comparan normalizadas, de modo que las variantes de acento o de guion
# (— vs -) no rompen la lectura.
_CAMPOS_RESUMEN = {
    # 1. Identificacion y localizacion
    "codigo del bloque": "codigo_bloque",
    "microcuenca (catalogo)": "microcuenca",
    "microcuenca declarada en ficha dt": "microcuenca_ficha",
    "zona de planificacion": "zona_planificacion",
    "departamento": "departamento",
    "provincia": "provincia",
    "distrito": "distrito",
    "capital distrital": "capital_distrital",
    "centro poblado asociado": "centro_poblado",
    "comunidad campesina": "comunidad_campesina",
    "superficie de catalogo (v5/v6), ha": "area_ha",
    "fuente de superficie": "fuente_superficie",
    "centroide utm este (m)": "utm_este",
    "centroide utm norte (m)": "utm_norte",
    "sistema de coordenadas": "sistema_coordenadas",
    "tipo de intervencion": "tipo_intervencion",
    # 2. Parametros fisicos
    "altitud minima (msnm)": "altitud_min",
    "altitud maxima (msnm)": "altitud_max",
    "amplitud altitudinal (m)": "amplitud_altitudinal",
    "piso altitudinal dominante": "piso_altitudinal",
    "pendiente promedio (%)": "pendiente_pct",
    "pendiente promedio (grados)": "pendiente_grados",
    "clase de pendiente equivalente": "clase_pendiente",
    "rango de pendiente declarado en campo": "pendiente_campo",
    "forma predominante del terreno": "forma_terreno",
    "posicion fisiografica": "posicion_fisiografica",
    "exposicion / orientacion": "exposicion",
    "afloramientos rocosos": "afloramientos_rocosos",
    "escarpes activos": "escarpes_activos",
    "remociones en masa activas": "remociones_masa",
    # 3. Indices de vegetacion
    "msavi 2024 media del bloque": "msavi_2024",
    "msavi 2024 clase de la media": "msavi_clase",
    "condicion frente al umbral": "condicion_umbral",
    "ndvi mediana 2025 clase modal": "ndvi_clase_modal",
    "superficie clasificada ndvi 2025 (ha)": "superficie_ndvi_ha",
    "desviacion frente al catalogo (%)": "desviacion_catalogo_pct",
    # Distribucion areal del MSAVI 2024 por clase DN (plantillas V6 en
    # adelante). En los libros anteriores estas filas no existen y las
    # claves quedan ausentes.
    "superficie clasificada msavi 2024 (ha)": "superficie_msavi_ha",
    "desviacion msavi frente al catalogo (%)": "desviacion_msavi_pct",
    "superficie bajo umbral msavi 0.4976 (ha)": "superficie_bajo_umbral_ha",
    "% del bloque bajo umbral (brecha espectral)": "bajo_umbral_pct",
    "clase dn dominante (msavi 2024)": "msavi_clase_dominante",
    "n. de poligonos msavi del bloque": "msavi_poligonos",
    # 4. Ecosistema y estado de conservacion
    "tipo de ecosistema (up)": "tipo_ecosistema",
    "superficie de ecosistema (ha)": "superficie_ecosistema",
    "estado de conservacion": "estado_conservacion",
    "uso actual dominante del suelo": "uso_dominante",
    "tipo de cobertura dominante": "tipo_cobertura",
    "cobertura vegetal total campo (%)": "cobertura_total_pct",
    "suelo desnudo campo (%)": "suelo_desnudo_pct",
    "regeneracion natural": "regeneracion",
    "nivel general de erosion": "nivel_erosion",
    "n. de carcavas registradas": "n_carcavas",
    "elenco floristico (n. de taxones)": "n_taxones",
    "estado sanitario": "estado_sanitario",
    # 5. Responsable y modalidad
    "responsable de la evaluacion": "evaluador",
    "fecha de evaluacion": "fecha_evaluacion",
    "hora de registro": "hora_registro",
    "correlativo de ficha": "correlativo_ficha",
    "entidad": "entidad",
    "fase del estudio": "fase_estudio",
    "instrumento aplicado": "instrumento",
    "parcela de muestreo": "parcela_muestreo",
    "estaciones fotograficas georreferenciadas": "n_estaciones",
    "modalidad de acceso": "modalidad_acceso",
    # 6. Sintesis GdR-CCC
    "causa subyacente principal": "causa_subyacente",
    "velocidad de degradacion": "velocidad_degradacion",
    "reversibilidad tecnica": "reversibilidad",
    "urgencia de intervencion": "urgencia_intervencion",
    "urgencia de control de erosion": "urgencia_erosion",
    "zona de recarga hidrica": "zona_recarga",
    "peligro integrado preliminar (mca ahp)": "peligro_integrado",
    "prioridad de intervencion": "prioridad",
    "estado de verificacion de campo": "estado_verificacion",
    "marco del indicador de brecha": "marco_indicador",
}

# Campos que se exponen tambien como numero, para filtros y graficos.
_CAMPOS_NUMERICOS = [
    "area_ha", "utm_este", "utm_norte", "altitud_min", "altitud_max",
    "amplitud_altitudinal", "pendiente_pct", "pendiente_grados",
    "msavi_2024", "superficie_ndvi_ha", "desviacion_catalogo_pct",
    "cobertura_total_pct", "suelo_desnudo_pct", "n_taxones", "n_estaciones",
    "superficie_msavi_ha", "desviacion_msavi_pct", "superficie_bajo_umbral_ha",
    "bajo_umbral_pct", "msavi_poligonos",
]


def _leer_resumen(rej, datos):
    """Lee los pares etiqueta/valor de la hoja Resumen.

    La hoja dispone dos pares por fila: (A,B) y (C,D). Se recorre celda por
    celda y, cuando el texto coincide con una etiqueta conocida, se toma la
    celda inmediatamente a la derecha como valor.
    """
    for r in range(rej.n_filas):
        for c in range(rej.n_columnas - 1):
            clave = _CAMPOS_RESUMEN.get(_norm(rej.valor(r, c)))
            if clave and clave not in datos:
                datos[clave] = _txt(rej.valor(r, c + 1))

    # Codigo del bloque de respaldo: el titulo de la hoja lo repite.
    if not datos.get("codigo_bloque"):
        pos = rej.buscar_etiqueta("ficha resumen bloque preliminar de intervencion",
                                  columna_max=rej.n_columnas)
        if pos:
            titulo = _txt(rej.valor(*pos))
            m = re.search(r"INTERVENCI[OÓ]N\s+(\S+)\s*$", titulo, re.IGNORECASE)
            if m:
                datos["codigo_bloque"] = m.group(1)

    for clave in _CAMPOS_NUMERICOS:
        datos[clave + "_num"] = _num(datos.get(clave))


# ══════════════════════════════════════════════════════════════════════════
# Lectura de tablas (hojas 2 a 5)
# ══════════════════════════════════════════════════════════════════════════

def _leer_tabla(rej, columnas_clave, campos, fila_inicio=0, filas_corte=()):
    """Lee una tabla delimitada por su fila de cabecera.

    columnas_clave: textos que identifican la cabecera (deben estar todos).
    campos: lista de (clave_normalizada_de_columna, nombre_salida, tipo)
            donde tipo es "texto" o "numero".
    filas_corte: textos que, al aparecer en la primera columna, terminan la
            tabla (TOTAL, notas metodologicas, etc.).
    """
    hallazgo = rej.buscar_fila_cabecera(*columnas_clave)
    if not hallazgo:
        return [], None
    fila_cab, mapa = hallazgo
    if fila_cab < fila_inicio:
        return [], None

    cortes = [_norm(t) for t in filas_corte]
    registros = []
    vacias_seguidas = 0
    for r in range(fila_cab + 1, rej.n_filas):
        if rej.fila_vacia(r):
            vacias_seguidas += 1
            # Una fila en blanco suelta no cierra la tabla; dos si.
            if vacias_seguidas >= 2:
                break
            continue
        vacias_seguidas = 0

        col_primera = rej.columna(mapa, columnas_clave[0]) or 0
        primera = _norm(rej.valor(r, col_primera))
        if any(primera.startswith(c) for c in cortes):
            break
        # Una nota metodologica larga en la primera columna cierra la tabla.
        if len(primera) > 90:
            break

        registro = {}
        for clave_col, salida, tipo in campos:
            idx = rej.columna(mapa, clave_col)
            bruto = rej.valor(r, idx) if idx is not None else None
            if tipo == "numero":
                registro[salida] = _num(bruto)
                registro[salida + "_txt"] = _txt(bruto)
            else:
                registro[salida] = _txt(bruto)
        if any(_txt(v) for v in registro.values()):
            registros.append(registro)
    return registros, fila_cab


def _leer_msavi_ndvi(rej, datos):
    """Hoja 2: clases MSAVI 2024 y distribucion areal NDVI 2025."""
    datos["msavi_tabla"], _ = _leer_tabla(
        rej,
        ["Clase MSAVI", "Superficie (ha)", "% del area clasificada"],
        [("Clase MSAVI", "clase", "texto"),
         ("Superficie (ha)", "superficie_ha", "numero"),
         ("% del area clasificada", "pct", "numero"),
         ("Interpretacion", "interpretacion", "texto"),
         ("Condicion frente al umbral", "condicion", "texto")],
        filas_corte=("MSAVI 2024 - MEDIA DEL BLOQUE", "MSAVI 2024 MEDIA DEL BLOQUE",
                     "TOTAL CLASIFICADO", "Superficie SOBRE umbral",
                     "Superficie BAJO umbral", "Superficie de catalogo",
                     "NOTA METODOLOGICA", "B. NDVI"))

    # Sintesis de la distribucion MSAVI (plantillas V6). Ausente en los
    # libros anteriores, donde la seccion solo tenia las cinco clases.
    for etiqueta, clave in (
            ("total clasificado msavi", "msavi_total_ha"),
            ("superficie sobre umbral", "msavi_sobre_umbral_ha"),
            ("superficie bajo umbral", "msavi_bajo_umbral_ha")):
        pos = rej.buscar_etiqueta(etiqueta, columna_max=rej.n_columnas)
        if pos:
            datos[clave] = _num(rej.valor(pos[0], pos[1] + 1))
            datos[clave.replace("_ha", "_pct")] = _num(
                rej.valor(pos[0], pos[1] + 2))

    datos["ndvi_tabla"], _ = _leer_tabla(
        rej,
        ["Clase NDVI", "Superficie (ha)", "% del area clasificada"],
        [("Clase NDVI", "clase", "texto"),
         ("Superficie (ha)", "superficie_ha", "numero"),
         ("% del area clasificada", "pct", "numero"),
         ("Interpretacion", "interpretacion", "texto"),
         ("Observacion", "observacion", "texto")],
        filas_corte=("TOTAL CLASIFICADO", "Superficie de catalogo",
                     "LECTURA CRITICA", "NOTA METODOLOGICA"))

    pos = rej.buscar_etiqueta("total clasificado", columna_max=rej.n_columnas)
    if pos:
        datos["ndvi_total_ha"] = _num(rej.valor(pos[0], pos[1] + 1))


def completar_sintesis_msavi(datos):
    """Deriva el reparto porcentual del MSAVI cuando el libro no lo trae.

    En la seccion A el porcentaje y los totales son formulas. Un libro
    guardado sin recalcular no lleva el resultado, y quien lee valores (este
    aplicativo, pandas) obtendria celdas vacias. Las superficies por clase si
    son literales, de modo que el reparto se reconstruye sobre ellas y queda
    marcado como calculado. Es idempotente: lo que el libro ya declara no se
    toca.
    """
    tabla = datos.get("msavi_tabla") or []
    areas = [f.get("superficie_ha") for f in tabla]
    if not areas or any(a is None for a in areas):
        return datos

    total = datos.get("msavi_total_ha")
    if total is None:
        total = round(sum(areas), 4)
        datos["msavi_total_ha"] = total
        datos["msavi_sintesis_calculada"] = True
    if not total:
        return datos

    for fila in tabla:
        if fila.get("pct") is None:
            fila["pct"] = round(fila["superficie_ha"] / total * 100, 2)
            fila["pct_txt"] = f"{fila['pct']:.2f}"
            fila["pct_calculado"] = True

    # El umbral 0.4976 separa las clases: la propia tabla declara de que lado
    # cae cada una.
    sobre = sum(f["superficie_ha"] for f in tabla
                if _norm(f.get("condicion")).startswith("sobre"))
    for clave, valor in (("msavi_sobre_umbral_ha", round(sobre, 4)),
                         ("msavi_bajo_umbral_ha", round(total - sobre, 4))):
        if datos.get(clave) is None:
            datos[clave] = valor
            datos[clave.replace("_ha", "_pct")] = round(valor / total * 100, 2)
            datos["msavi_sintesis_calculada"] = True
    return datos


def _leer_estaciones(rej, datos):
    """Hoja 3: inventario de puntos georreferenciados de la ficha DT."""
    datos["estaciones"], _ = _leer_tabla(
        rej,
        ["Codigo", "UTM ESTE (m)", "UTM NORTE (m)"],
        [("Codigo", "codigo", "texto"),
         ("Naturaleza del punto", "naturaleza", "texto"),
         ("UTM ESTE (m)", "utm_este", "numero"),
         ("UTM NORTE (m)", "utm_norte", "numero"),
         ("Dist. al centroide (m)", "dist_centroide", "numero"),
         ("Estado / regimen", "estado", "texto"),
         ("Contenido registrado", "contenido", "texto")],
        filas_corte=("TOTAL", "Centroide del bloque", "CONTROL GEOMETRICO"))

    for etiqueta, clave in (
            ("n. de estaciones fotograficas declaradas", "n_estaciones_declaradas"),
            ("estaciones dentro del poligono", "estaciones_dentro")):
        pos = rej.buscar_etiqueta(etiqueta, columna_max=rej.n_columnas)
        if pos:
            datos[clave] = _num(rej.valor(pos[0], pos[1] + 1))


def _leer_microcuenca(rej, datos):
    """Hoja 4: contexto intramicrocuenca (bloques de la misma microcuenca)."""
    registros, _ = _leer_tabla(
        rej,
        ["Bloque", "Area (ha)", "% microcuenca"],
        [("Bloque", "bloque", "texto"),
         ("Area (ha)", "area_ha", "numero"),
         ("% microcuenca", "pct_microcuenca", "numero"),
         ("Rango altitudinal (msnm)", "rango_altitudinal", "texto"),
         ("Amplitud (m)", "amplitud", "numero"),
         ("Pendiente prom. (%)", "pendiente_pct", "numero"),
         ("MSAVI 2024", "msavi", "numero"),
         ("NDVI 2025", "ndvi_veg_alta", "numero")],
        filas_corte=("TOTAL / PROMEDIO", "TOTAL/PROMEDIO", "LECTURA INTRAMICROCUENCA",
                     "NOTA METODOLOGICA"))

    # El bloque de la ficha viene marcado con un puntero en la hoja.
    for reg in registros:
        etiqueta = reg.get("bloque", "")
        reg["es_actual"] = etiqueta.startswith("►") or etiqueta.startswith(">")
        reg["bloque"] = etiqueta.lstrip("►> ").strip()
    datos["microcuenca_tabla"] = registros


def _leer_consistencia(rej, datos):
    """Hoja 5: control de consistencia campo / gabinete / catalogo."""
    registros, _ = _leer_tabla(
        rej,
        ["Cod.", "Campo afectado", "Calificacion"],
        [("Cod.", "codigo", "texto"),
         ("Campo afectado", "campo", "texto"),
         ("Discrepancia observada", "discrepancia", "texto"),
         ("Calificacion", "calificacion", "texto"),
         ("Tratamiento adoptado", "tratamiento", "texto")],
        filas_corte=("RESUMEN", "NOTA METODOLOGICA"))
    datos["consistencia"] = registros

    conteo = {c: 0 for c in CALIFICACIONES}
    for reg in registros:
        cal = _norm(reg.get("calificacion"))
        for c in CALIFICACIONES:
            if cal == _norm(c):
                conteo[c] += 1
                break
    conteo["total"] = len(registros)
    datos["consistencia_resumen"] = conteo


# ══════════════════════════════════════════════════════════════════════════
# Parseo completo de un libro
# ══════════════════════════════════════════════════════════════════════════

_HOJAS = [
    (("resumen",), _leer_resumen),
    (("cobertura", "msavi"), _leer_msavi_ndvi),
    (("estaciones",), _leer_estaciones),
    (("microcuenca",), _leer_microcuenca),
    (("control de consistencia", "consistencia"), _leer_consistencia),
]


def _elegir_hoja(wb, claves):
    """Primera hoja cuyo nombre contiene alguna de las claves."""
    for nombre in wb.sheetnames:
        n = _norm(nombre)
        for clave in claves:
            if _norm(clave) in n:
                return nombre
    return None


def codigo_desde_nombre(nombre_archivo):
    """Extrae el codigo del bloque del nombre de archivo del generador."""
    base = (nombre_archivo or "").rsplit("/", 1)[-1]
    m = re.match(r"^Plantilla_Excel_Bloque_(.+)_IN_Piura\.xlsx$", base,
                 re.IGNORECASE)
    return m.group(1) if m else ""


def parsear_resumen_bloque(archivo, nombre_archivo=""):
    """Parsea un libro de resumen y devuelve un diccionario de datos.

    `archivo` puede ser bytes, un path o un objeto tipo file. Nunca lanza por
    hojas ausentes: lo que no se pueda leer queda vacio y se registra en
    `hojas_leidas` / `advertencias`, de modo que un libro parcial siga siendo
    cargable y el original quede intacto en la base de datos.
    """
    if isinstance(archivo, (bytes, bytearray)):
        archivo = io.BytesIO(archivo)

    datos = {
        "nombre_archivo": (nombre_archivo or "").rsplit("/", 1)[-1],
        "hojas_leidas": [],
        "advertencias": [],
        "msavi_tabla": [], "ndvi_tabla": [], "estaciones": [],
        "microcuenca_tabla": [], "consistencia": [],
        "consistencia_resumen": {c: 0 for c in CALIFICACIONES},
    }
    datos["consistencia_resumen"]["total"] = 0

    wb = load_workbook(archivo, data_only=True, read_only=True)
    try:
        datos["hojas"] = list(wb.sheetnames)
        for claves, lector in _HOJAS:
            nombre = _elegir_hoja(wb, claves)
            if not nombre:
                datos["advertencias"].append(
                    f"No se encontro la hoja '{claves[0]}'.")
                continue
            try:
                lector(_Rejilla(wb[nombre]), datos)
                datos["hojas_leidas"].append(nombre)
            except Exception as exc:                      # hoja malformada
                datos["advertencias"].append(
                    f"Hoja '{nombre}' no pudo leerse: {exc}")
    finally:
        wb.close()

    if not datos.get("codigo_bloque"):
        datos["codigo_bloque"] = codigo_desde_nombre(datos["nombre_archivo"])

    completar_sintesis_msavi(datos)
    datos["validacion_utm"] = validar_utm(datos.get("utm_este_num"),
                                          datos.get("utm_norte_num"))
    return datos


def validar_utm(este, norte):
    """Verifica que el centroide caiga en el ambito UTM 17S del proyecto."""
    if este is None or norte is None:
        return "Sin coordenadas"
    if not (UTM_ESTE_MIN <= este <= UTM_ESTE_MAX):
        return f"ESTE fuera de rango ({este:,.0f} m)"
    if not (UTM_NORTE_MIN <= norte <= UTM_NORTE_MAX):
        return f"NORTE fuera de rango ({norte:,.0f} m)"
    return "Conforme"


def parsear_lote(archivos):
    """Parsea una lista de (nombre, bytes). Devuelve (resultados, errores).

    Un archivo ilegible no aborta el lote: se acumula en `errores` con su
    motivo para que la carga masiva de los 117 bloques informe que fallo.
    """
    resultados, errores = [], []
    for nombre, contenido in archivos:
        try:
            datos = parsear_resumen_bloque(contenido, nombre)
            resultados.append((nombre, contenido, datos))
        except Exception as exc:
            errores.append((nombre, f"{type(exc).__name__}: {exc}"))
    return resultados, errores


def expandir_zip(contenido, limite_archivos=500):
    """Extrae los .xlsx de un ZIP. Devuelve lista de (nombre, bytes).

    Ignora entradas de macOS (__MACOSX, ._*) y los temporales de Excel (~$).
    """
    salida = []
    with zipfile.ZipFile(io.BytesIO(contenido)) as z:
        for info in z.infolist():
            if info.is_dir() or len(salida) >= limite_archivos:
                continue
            base = info.filename.rsplit("/", 1)[-1]
            if not base.lower().endswith(".xlsx"):
                continue
            if base.startswith("._") or base.startswith("~$"):
                continue
            if "__MACOSX" in info.filename:
                continue
            salida.append((base, z.read(info)))
    return salida


# ══════════════════════════════════════════════════════════════════════════
# Manifiesto de los 117 bloques con ficha DT
# ══════════════════════════════════════════════════════════════════════════

_MANIFIESTO_CACHE = None


def cargar_manifiesto():
    """Catalogo de los 117 libros esperados (codigo, archivo, enlace Drive).

    Permite contrastar lo cargado en la base contra lo que debe existir y
    enlazar cada bloque pendiente a su archivo de origen.
    """
    global _MANIFIESTO_CACHE
    if _MANIFIESTO_CACHE is None:
        import os
        ruta = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                            "datos", "manifiesto_resumenes_117.json")
        try:
            with open(ruta, encoding="utf-8") as fh:
                _MANIFIESTO_CACHE = json.load(fh)
        except (OSError, ValueError):
            _MANIFIESTO_CACHE = {"total_bloques": 0, "bloques": []}
    return _MANIFIESTO_CACHE


# Carpeta del repositorio con los 117 libros vigentes (V6: distribucion areal
# del MSAVI 2024 por clase DN). Viaja con el aplicativo, de modo que la
# recarga masiva no depende de que alguien vuelva a subir los archivos.
CARPETA_LIBROS = "plantillas_117_msavi_v6"


def libros_del_repositorio(carpeta=None):
    """[(nombre, contenido)] de los libros de resumen incluidos en el repo.

    Devuelve lista vacia si la carpeta no esta presente en el despliegue.
    """
    ruta = carpeta or os.path.join(
        os.path.dirname(os.path.abspath(__file__)), CARPETA_LIBROS)
    if not os.path.isdir(ruta):
        return []
    libros = []
    for nombre in sorted(os.listdir(ruta)):
        if not nombre.lower().endswith(".xlsx") or nombre.startswith("~$"):
            continue
        with open(os.path.join(ruta, nombre), "rb") as fh:
            libros.append((nombre, fh.read()))
    return libros


def codigos_esperados():
    return [b["codigo"] for b in cargar_manifiesto().get("bloques", [])]


# ══════════════════════════════════════════════════════════════════════════
# Graficos en Excel
# ══════════════════════════════════════════════════════════════════════════

_BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def _titulo_hoja(ws, titulo, ancho=8):
    """Escribe el encabezado institucional ANIN en una hoja nueva."""
    fila = 1
    for texto in ENCABEZADOS_ANIN:
        ws.cell(fila, 1, texto).font = Font(name="Arial", size=9, bold=True,
                                            color=ANIN_VERDE)
        ws.merge_cells(start_row=fila, start_column=1,
                       end_row=fila, end_column=ancho)
        fila += 1
    ws.cell(fila, 1, SUBTITULO_PROYECTO).font = Font(name="Arial", size=8,
                                                     italic=True)
    ws.merge_cells(start_row=fila, start_column=1,
                   end_row=fila, end_column=ancho)
    fila += 1
    celda = ws.cell(fila, 1, titulo)
    celda.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor=ANIN_VERDE)
    celda.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=fila, start_column=1,
                   end_row=fila, end_column=ancho)
    ws.row_dimensions[fila].height = 22
    return fila + 2


def _escribir_bloque_datos(ws, fila, titulo, cabeceras, filas):
    """Escribe una tabla con estilo ANIN. Devuelve (fila_cabecera, fila_fin)."""
    celda = ws.cell(fila, 1, titulo)
    celda.font = Font(name="Arial", size=10, bold=True, color=ANIN_AZUL)
    fila += 1
    fila_cab = fila
    for i, texto in enumerate(cabeceras, start=1):
        c = ws.cell(fila, i, texto)
        c.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ANIN_VERDE)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = _BORDE
    fila += 1
    for j, registro in enumerate(filas):
        for i, valor in enumerate(registro, start=1):
            c = ws.cell(fila, i, valor)
            c.font = Font(name="Arial", size=9)
            c.border = _BORDE
            if j % 2:
                c.fill = PatternFill("solid", fgColor=ANIN_GRIS)
            if isinstance(valor, float):
                c.number_format = "#,##0.0000" if abs(valor) < 10 else "#,##0.00"
        fila += 1
    return fila_cab, fila


def _agregar_grafico_barras(ws, titulo, fila_cab, n_filas, col_cat, col_val,
                            ancla, eje_y="", ancho=16, alto=8):
    if n_filas <= 0:
        return
    graf = BarChart()
    graf.type = "col"
    graf.style = 10
    graf.title = titulo
    graf.y_axis.title = eje_y
    graf.height, graf.width = alto, ancho
    datos = Reference(ws, min_col=col_val, min_row=fila_cab,
                      max_row=fila_cab + n_filas)
    cats = Reference(ws, min_col=col_cat, min_row=fila_cab + 1,
                     max_row=fila_cab + n_filas)
    graf.add_data(datos, titles_from_data=True)
    graf.set_categories(cats)
    graf.legend = None
    ws.add_chart(graf, ancla)


def _agregar_grafico_torta(ws, titulo, fila_cab, n_filas, col_cat, col_val,
                           ancla, ancho=12, alto=8):
    if n_filas <= 0:
        return
    graf = PieChart()
    graf.title = titulo
    graf.height, graf.width = alto, ancho
    datos = Reference(ws, min_col=col_val, min_row=fila_cab,
                      max_row=fila_cab + n_filas)
    cats = Reference(ws, min_col=col_cat, min_row=fila_cab + 1,
                     max_row=fila_cab + n_filas)
    graf.add_data(datos, titles_from_data=True)
    graf.set_categories(cats)
    ws.add_chart(graf, ancla)


def generar_excel_con_graficos(contenido_original, datos):
    """Devuelve el libro original con una hoja adicional de graficos.

    No modifica las cinco hojas de origen: agrega "Graficos" al final con las
    series y sus graficos nativos de Excel (editables por el usuario).
    """
    wb = load_workbook(io.BytesIO(contenido_original))
    if "Graficos" in wb.sheetnames:
        del wb["Graficos"]
    ws = wb.create_sheet("Graficos")
    _construir_hoja_graficos(ws, datos)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _construir_hoja_graficos(ws, datos):
    """Series y graficos del bloque: NDVI, MSAVI, microcuenca y consistencia."""
    codigo = datos.get("codigo_bloque", "")
    fila = _titulo_hoja(ws, f"GRAFICOS DEL BLOQUE {codigo}", ancho=8)
    for col, ancho in zip("ABCDEFGH", (34, 16, 16, 16, 16, 16, 16, 16)):
        ws.column_dimensions[col].width = ancho

    # ── NDVI 2025: distribucion areal ──
    ndvi = [r for r in datos.get("ndvi_tabla", [])
            if r.get("superficie_ha") is not None]
    if ndvi:
        cab, fin = _escribir_bloque_datos(
            ws, fila, "A. NDVI mediana 2025 - distribucion areal",
            ["Clase NDVI", "Superficie (ha)", "% del area"],
            [[r["clase"], r["superficie_ha"], r.get("pct")] for r in ndvi])
        _agregar_grafico_torta(ws, f"NDVI 2025 - {codigo} (ha)", cab, len(ndvi),
                               1, 2, f"E{cab}")
        _agregar_grafico_barras(ws, f"NDVI 2025 - {codigo} (% del area)", cab,
                                len(ndvi), 1, 3, f"E{cab + 17}", eje_y="%")
        fila = fin + 18

    # ── MSAVI 2024: clases del proyecto ──
    msavi = [r for r in datos.get("msavi_tabla", [])
             if r.get("superficie_ha") is not None]
    if msavi:
        cab, fin = _escribir_bloque_datos(
            ws, fila, f"B. MSAVI 2024 - clases (umbral {UMBRAL_MSAVI})",
            ["Clase MSAVI", "Superficie (ha)", "% del area"],
            [[r["clase"], r["superficie_ha"], r.get("pct")] for r in msavi])
        _agregar_grafico_barras(ws, f"MSAVI 2024 - {codigo} (ha)", cab,
                                len(msavi), 1, 2, f"E{cab}", eje_y="ha")
        fila = fin + 18
    else:
        media = datos.get("msavi_2024_num")
        if media is not None:
            celda = ws.cell(fila, 1,
                            "B. MSAVI 2024 - la distribucion areal por clase no "
                            "figura en los insumos; solo se dispone de la media "
                            f"del bloque ({media:.4f}). No se estima.")
            celda.font = Font(name="Arial", size=9, italic=True)
            fila += 2

    # ── Contexto intramicrocuenca ──
    micro = [r for r in datos.get("microcuenca_tabla", [])
             if r.get("area_ha") is not None]
    if micro:
        cab, fin = _escribir_bloque_datos(
            ws, fila,
            f"C. Contexto intramicrocuenca {datos.get('microcuenca', '')}",
            ["Bloque", "Area (ha)", "Pendiente prom. (%)", "MSAVI 2024"],
            [[("> " + r["bloque"]) if r.get("es_actual") else r["bloque"],
              r["area_ha"], r.get("pendiente_pct"), r.get("msavi")]
             for r in micro])
        _agregar_grafico_barras(ws, "Area por bloque (ha)", cab, len(micro),
                                1, 2, f"F{cab}", eje_y="ha")
        _agregar_grafico_barras(ws, "MSAVI 2024 por bloque", cab, len(micro),
                                1, 4, f"F{cab + 17}", eje_y="MSAVI")
        fila = fin + 18

    # ── Control de consistencia ──
    resumen = datos.get("consistencia_resumen", {})
    filas_cal = [[c, resumen.get(c, 0)] for c in CALIFICACIONES
                 if resumen.get(c, 0)]
    if filas_cal:
        cab, fin = _escribir_bloque_datos(
            ws, fila, "D. Control de consistencia por calificacion",
            ["Calificacion", "N. de verificaciones"], filas_cal)
        _agregar_grafico_torta(ws, "Verificaciones por calificacion", cab,
                               len(filas_cal), 1, 2, f"E{cab}")
        fila = fin + 18

    celda = ws.cell(fila, 1,
                    "Graficos generados por el aplicativo IN Piura sobre los "
                    "valores declarados en el libro de origen. Las clases sin "
                    "estadistica zonal disponible se omiten y no se estiman.")
    celda.font = Font(name="Arial", size=8, italic=True)
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=8)
    ws.sheet_view.showGridLines = False


def generar_excel_consolidado(lista_datos):
    """Libro consolidado de los bloques cargados, con graficos comparativos."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Consolidado"
    fila = _titulo_hoja(
        ws, f"CONSOLIDADO DE RESUMENES DT - {len(lista_datos)} BLOQUES", ancho=13)
    for col, ancho in zip("ABCDEFGHIJKLM",
                          (12, 16, 16, 18, 12, 12, 12, 12, 12, 12, 14, 14, 22)):
        ws.column_dimensions[col].width = ancho

    cabeceras = ["Bloque", "Microcuenca", "Provincia", "Distrito", "Area (ha)",
                 "UTM ESTE", "UTM NORTE", "Alt. min", "Alt. max",
                 "Pend. (%)", "MSAVI 2024", "Sustantivas", "Estado verificacion"]
    filas = []
    for d in lista_datos:
        filas.append([
            d.get("codigo_bloque", ""), d.get("microcuenca", ""),
            d.get("provincia", ""), d.get("distrito", ""),
            d.get("area_ha_num"), d.get("utm_este_num"), d.get("utm_norte_num"),
            d.get("altitud_min_num"), d.get("altitud_max_num"),
            d.get("pendiente_pct_num"), d.get("msavi_2024_num"),
            (d.get("consistencia_resumen") or {}).get("SUSTANTIVA", 0),
            d.get("estado_verificacion", ""),
        ])
    fila_cab, fila_fin = _escribir_bloque_datos(
        ws, fila, "Parametros por bloque", cabeceras, filas)
    ws.freeze_panes = ws.cell(fila_cab + 1, 1)

    # Totales con formulas, no con valores precalculados.
    ws.cell(fila_fin, 1, "TOTAL / PROMEDIO").font = Font(
        name="Arial", size=9, bold=True)
    ini, fin = fila_cab + 1, fila_fin - 1
    if fin >= ini:
        for col, func in ((5, "SUM"), (10, "AVERAGE"), (11, "AVERAGE"),
                          (12, "SUM")):
            letra = get_column_letter(col)
            c = ws.cell(fila_fin, col, f"={func}({letra}{ini}:{letra}{fin})")
            c.font = Font(name="Arial", size=9, bold=True)
            c.number_format = "#,##0.0000" if col == 11 else "#,##0.00"

    _hoja_graficos_consolidado(wb, lista_datos)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _hoja_graficos_consolidado(wb, lista_datos):
    """Agregados por distrito, por clase MSAVI y por calificacion."""
    ws = wb.create_sheet("Graficos")
    fila = _titulo_hoja(ws, "GRAFICOS CONSOLIDADOS", ancho=8)
    for col, ancho in zip("ABCDEFGH", (34, 16, 16, 16, 16, 16, 16, 16)):
        ws.column_dimensions[col].width = ancho

    # Area y numero de bloques por distrito.
    por_distrito = {}
    for d in lista_datos:
        clave = d.get("distrito") or "Sin distrito"
        acu = por_distrito.setdefault(clave, [0, 0.0])
        acu[0] += 1
        acu[1] += d.get("area_ha_num") or 0.0
    filas = sorted(([k, v[0], round(v[1], 2)] for k, v in por_distrito.items()),
                   key=lambda r: -r[2])
    if filas:
        cab, fin = _escribir_bloque_datos(
            ws, fila, "A. Bloques y superficie por distrito",
            ["Distrito", "N. de bloques", "Area (ha)"], filas)
        _agregar_grafico_barras(ws, "Superficie por distrito (ha)", cab,
                                len(filas), 1, 3, f"E{cab}", eje_y="ha")
        _agregar_grafico_barras(ws, "Bloques por distrito", cab, len(filas),
                                1, 2, f"E{cab + 17}", eje_y="bloques")
        fila = fin + 18

    # Distribucion de bloques frente al umbral de brecha MSAVI.
    bajo = sum(1 for d in lista_datos
               if (d.get("msavi_2024_num") or 0) and
               d["msavi_2024_num"] < UMBRAL_MSAVI)
    sobre = sum(1 for d in lista_datos
                if (d.get("msavi_2024_num") or 0) >= UMBRAL_MSAVI)
    sin_dato = len(lista_datos) - bajo - sobre
    filas = [[f"BAJO umbral {UMBRAL_MSAVI}", bajo],
             [f"Sobre umbral {UMBRAL_MSAVI}", sobre]]
    if sin_dato:
        filas.append(["Sin MSAVI declarado", sin_dato])
    cab, fin = _escribir_bloque_datos(
        ws, fila, "B. Bloques frente al umbral de brecha (R.M. 00213-2024-MINAM)",
        ["Condicion", "N. de bloques"], filas)
    _agregar_grafico_torta(ws, "Condicion frente al umbral MSAVI", cab,
                           len(filas), 1, 2, f"E{cab}")
    fila = fin + 18

    # Verificaciones de consistencia acumuladas.
    acumulado = {c: 0 for c in CALIFICACIONES}
    for d in lista_datos:
        for c in CALIFICACIONES:
            acumulado[c] += (d.get("consistencia_resumen") or {}).get(c, 0)
    filas = [[c, acumulado[c]] for c in CALIFICACIONES if acumulado[c]]
    if filas:
        cab, _ = _escribir_bloque_datos(
            ws, fila, "C. Verificaciones de consistencia acumuladas",
            ["Calificacion", "N. de verificaciones"], filas)
        _agregar_grafico_barras(ws, "Verificaciones por calificacion", cab,
                                len(filas), 1, 2, f"E{cab}",
                                eje_y="verificaciones")
    ws.sheet_view.showGridLines = False


# ══════════════════════════════════════════════════════════════════════════
# Exportacion a PDF
# ══════════════════════════════════════════════════════════════════════════

from fpdf import FPDF                                       # noqa: E402

# Helvetica (fuente estandar de FPDF) solo cubre Latin-1. Se reemplazan los
# simbolos Unicode que traen los libros generados (guiones largos, comillas
# tipograficas, el puntero del bloque actual) por equivalentes ASCII.
_REEMPLAZOS_PDF = {
    "—": "-", "–": "-", "‘": "'", "’": "'", "“": '"', "”": '"',
    "…": "...", "•": "-", "►": ">", "≈": "~", "±": "+/-", "≤": "<=",
    "≥": ">=", "→": "->", "«": '"', "»": '"', " ": " ", " ": " ",
    "​": "", "−": "-", "·": "-", "²": "2", "³": "3",
}

# Paleta de series para los graficos vectoriales (verde institucional y
# tonos derivados, legibles en impresion a color y en escala de grises).
_PALETA = [(27, 77, 46), (58, 124, 79), (122, 168, 116), (196, 160, 60),
           (176, 96, 54), (108, 122, 137), (140, 60, 90), (60, 110, 140)]

# En el control de consistencia el color comunica severidad, no solo separa
# series: la discrepancia sustantiva condiciona el cierre del entregable.
_COLOR_CALIFICACION = {
    "SUSTANTIVA": (176, 58, 52),
    "NO SUSTANTIVA": (196, 160, 60),
    "CORREGIDO": (58, 124, 79),
    "CONFORME": (27, 77, 46),
}


def _s(valor):
    """Texto seguro para Latin-1."""
    if valor is None:
        return ""
    txt = str(valor)
    for origen, destino in _REEMPLAZOS_PDF.items():
        if origen in txt:
            txt = txt.replace(origen, destino)
    return txt.encode("latin-1", "replace").decode("latin-1")


class _PDFResumen(FPDF):
    """PDF institucional ANIN para los resumenes de Diagnostico Territorial."""

    def __init__(self, subtitulo="", orientacion="P"):
        super().__init__(orientation=orientacion, unit="mm", format="A4")
        self.subtitulo = subtitulo
        self.set_auto_page_break(auto=True, margin=18)
        self.set_margins(12, 12, 12)

    # ── Cabecera y pie institucionales ──
    def header(self):
        self.set_fill_color(27, 77, 46)
        self.rect(0, 0, self.w, 21, "F")
        self.set_text_color(255, 255, 255)
        self.set_xy(12, 4)
        self.set_font("Helvetica", "B", 9)
        self.cell(0, 4, _s("AUTORIDAD NACIONAL DE INFRAESTRUCTURA - ANIN"), 0, 1)
        self.set_x(12)
        self.set_font("Helvetica", "", 7.5)
        self.cell(0, 3.6, _s("DIRECCION DE INTERVENCIONES MULTISECTORIALES Y "
                             "DE EMERGENCIA - DIME"), 0, 1)
        self.set_x(12)
        self.cell(0, 3.6, _s("SUBDIRECCION DE ESTUDIOS DE INVERSION - SESDI"),
                  0, 1)
        self.set_x(12)
        self.set_font("Helvetica", "I", 7)
        self.cell(0, 3.6, _s("Proyecto IN Piura | CUI 2669244 | UTM WGS 84 "
                             "Zona 17S (EPSG:32717)"), 0, 1)
        self.set_text_color(0, 0, 0)
        self.set_y(25)
        if self.subtitulo:
            self.set_font("Helvetica", "B", 11)
            self.set_fill_color(232, 240, 234)
            self.cell(0, 7, _s(self.subtitulo), 0, 1, "C", True)
            self.ln(2)

    def footer(self):
        self.set_y(-14)
        self.set_draw_color(27, 77, 46)
        self.line(12, self.get_y(), self.w - 12, self.get_y())
        self.set_y(-11)
        self.set_font("Helvetica", "I", 7)
        self.set_text_color(90, 90, 90)
        self.cell(0, 4, _s("ANIN - DIME - SESDI | Preinversion (Perfil) - "
                           "Invierte.pe"), 0, 0, "L")
        self.cell(0, 4, _s(f"Pagina {self.page_no()} de {{nb}}"), 0, 0, "R")
        self.set_text_color(0, 0, 0)

    # ── Bloques de contenido ──
    def seccion(self, titulo):
        self._salto_si_falta(14)
        self.ln(1)
        self.set_font("Helvetica", "B", 9.5)
        self.set_fill_color(27, 77, 46)
        self.set_text_color(255, 255, 255)
        self.cell(0, 6, _s(" " + titulo), 0, 1, "L", True)
        self.set_text_color(0, 0, 0)
        self.ln(1.5)

    def campos(self, pares, columnas=2):
        """Rejilla de etiqueta/valor. Omite los pares sin valor.

        El alto de la fila se calcula sobre el texto mas largo de la fila,
        etiqueta incluida: las etiquetas del formato son largas ("Superficie
        clasificada NDVI (ha)") y, si solo se midiera el valor, se
        desbordarian sobre la fila siguiente.
        """
        pares = [(e, _s(v)) for e, v in pares if _s(v)]
        if not pares:
            return
        ancho_col = (self.w - 24) / columnas
        ancho_etq = ancho_col * 0.46
        ancho_val = ancho_col - ancho_etq
        alto_linea = 3.6

        for i in range(0, len(pares), columnas):
            grupo = pares[i:i + columnas]
            lineas = 1
            for etiqueta, valor in grupo:
                self.set_font("Helvetica", "B", 7.5)
                lineas = max(lineas, len(self.multi_cell(
                    ancho_etq - 2, alto_linea, _s(etiqueta),
                    dry_run=True, output="LINES")))
                self.set_font("Helvetica", "", 7.5)
                lineas = max(lineas, len(self.multi_cell(
                    ancho_val - 2, alto_linea, _s(valor),
                    dry_run=True, output="LINES")))
            alto = lineas * alto_linea + 1.8
            self._salto_si_falta(alto + 2)

            y0 = self.get_y()
            x0 = 12.0
            self.set_draw_color(190, 190, 190)
            for etiqueta, valor in grupo:
                self.set_fill_color(242, 242, 242)
                self.rect(x0, y0, ancho_etq, alto, "DF")
                self.rect(x0 + ancho_etq, y0, ancho_val, alto, "D")
                self.set_xy(x0 + 1, y0 + 0.9)
                self.set_font("Helvetica", "B", 7.5)
                self.multi_cell(ancho_etq - 2, alto_linea, _s(etiqueta), 0, "L")
                self.set_xy(x0 + ancho_etq + 1, y0 + 0.9)
                self.set_font("Helvetica", "", 7.5)
                self.multi_cell(ancho_val - 2, alto_linea, _s(valor), 0, "L")
                x0 += ancho_col
            self.set_y(y0 + alto)

    def tabla(self, cabeceras, filas, anchos_rel=None, alineaciones=None,
              tam=7):
        """Tabla con cabecera verde y filas alternas."""
        if not filas:
            return
        ancho_total = self.w - 24
        anchos_rel = anchos_rel or [1] * len(cabeceras)
        suma = float(sum(anchos_rel))
        anchos = [ancho_total * a / suma for a in anchos_rel]
        alineaciones = alineaciones or ["L"] * len(cabeceras)

        def dibujar_cabecera():
            self.set_font("Helvetica", "B", tam)
            self.set_fill_color(27, 77, 46)
            self.set_text_color(255, 255, 255)
            for ancho, texto in zip(anchos, cabeceras):
                self.cell(ancho, 6, _s(texto), 1, 0, "C", True)
            self.ln()
            self.set_text_color(0, 0, 0)

        self._salto_si_falta(18)
        dibujar_cabecera()
        for i, fila in enumerate(filas):
            celdas = [_s(v) for v in fila]
            self.set_font("Helvetica", "", tam)
            alto_linea = tam * 0.55
            alto = 5
            for ancho, texto in zip(anchos, celdas):
                lineas = max(1, len(self.multi_cell(
                    ancho, alto_linea, texto, dry_run=True, output="LINES")))
                alto = max(alto, lineas * alto_linea + 1.2)
            if self.get_y() + alto > self.page_break_trigger:
                self.add_page(self.cur_orientation)
                dibujar_cabecera()
                self.set_font("Helvetica", "", tam)
            y0 = self.get_y()
            x0 = 12
            relleno = i % 2 == 1
            self.set_fill_color(244, 247, 245)
            for ancho, texto, alin in zip(anchos, celdas, alineaciones):
                self.set_xy(x0, y0)
                self.multi_cell(ancho, alto, texto, 1, alin, relleno)
                x0 += ancho
            self.set_y(y0 + alto)
        self.ln(2)

    def nota(self, texto):
        self.set_font("Helvetica", "I", 6.8)
        self.set_text_color(70, 70, 70)
        self.multi_cell(0, 3.4, _s(texto))
        self.set_text_color(0, 0, 0)
        self.ln(1.5)

    # ── Graficos vectoriales ──
    def grafico_barras(self, titulo, etiquetas, valores, unidad="",
                       alto=46, horizontal=False, colores=None):
        """Grafico de barras dibujado con primitivas: sin dependencias de
        imagen y con texto seleccionable en el PDF."""
        datos = [(e, v) for e, v in zip(etiquetas, valores)
                 if isinstance(v, (int, float))]
        if not datos:
            return
        tinta = _tintas(colores, [e for e, _ in datos])
        self._salto_si_falta(alto + 16)
        self.set_font("Helvetica", "B", 8)
        self.cell(0, 5, _s(titulo), 0, 1, "L")
        x0, y0 = 12.0, self.get_y() + 1
        ancho = self.w - 24
        maximo = max(v for _, v in datos) or 1.0

        if horizontal:
            ancho_etq = ancho * 0.34
            ancho_barra = ancho - ancho_etq - 22
            alto_fila = min(7.0, alto / max(len(datos), 1))
            for i, (etiqueta, valor) in enumerate(datos):
                y = y0 + i * alto_fila
                self.set_xy(x0, y)
                self.set_font("Helvetica", "", 6.5)
                self.cell(ancho_etq, alto_fila, _s(_recortar(etiqueta, 42)), 0,
                          0, "L")
                largo = max(0.4, ancho_barra * valor / maximo)
                self.set_fill_color(*tinta(i, etiqueta))
                self.rect(x0 + ancho_etq, y + alto_fila * 0.18, largo,
                          alto_fila * 0.64, "F")
                self.set_xy(x0 + ancho_etq + largo + 1.5, y)
                self.cell(20, alto_fila, _s(_fmt(valor) + unidad), 0, 0, "L")
            self.set_y(y0 + len(datos) * alto_fila + 3)
            return

        alto_area = alto - 9
        paso = ancho / len(datos)
        ancho_barra = min(paso * 0.6, 26)
        self.set_draw_color(200, 200, 200)
        self.line(x0, y0 + alto_area, x0 + ancho, y0 + alto_area)
        for i, (etiqueta, valor) in enumerate(datos):
            altura = max(0.4, (alto_area - 4) * valor / maximo)
            x = x0 + i * paso + (paso - ancho_barra) / 2
            self.set_fill_color(*tinta(i, etiqueta))
            self.rect(x, y0 + alto_area - altura, ancho_barra, altura, "F")
            self.set_xy(x - paso * 0.2, y0 + alto_area - altura - 4)
            self.set_font("Helvetica", "B", 6)
            self.cell(ancho_barra + paso * 0.4, 3.5, _s(_fmt(valor)), 0, 0, "C")
            self.set_xy(x - paso * 0.2, y0 + alto_area + 1)
            self.set_font("Helvetica", "", 6)
            self.cell(ancho_barra + paso * 0.4, 3.5,
                      _s(_recortar(etiqueta, int(paso / 1.6))), 0, 0, "C")
        self.set_y(y0 + alto + 2)

    def grafico_torta(self, titulo, etiquetas, valores, alto=48,
                      colores=None):
        """Sectores circulares con leyenda a la derecha."""
        datos = [(e, float(v)) for e, v in zip(etiquetas, valores)
                 if isinstance(v, (int, float)) and v > 0]
        total = sum(v for _, v in datos)
        if not datos or total <= 0:
            return
        tinta = _tintas(colores, [e for e, _ in datos])
        self._salto_si_falta(alto + 14)
        self.set_font("Helvetica", "B", 8)
        self.cell(0, 5, _s(titulo), 0, 1, "L")
        y0 = self.get_y() + 1
        radio = min(alto / 2 - 2, 20)
        cx, cy = 12 + radio + 2, y0 + radio + 1

        angulo = 0.0
        for i, (etiqueta, valor) in enumerate(datos):
            barrido = 360.0 * valor / total
            self.set_fill_color(*tinta(i, etiqueta))
            _sector(self, cx, cy, radio, angulo, angulo + barrido)
            angulo += barrido

        x_leyenda = cx + radio + 6
        y = y0
        for i, (etiqueta, valor) in enumerate(datos):
            self.set_fill_color(*tinta(i, etiqueta))
            self.rect(x_leyenda, y + 0.8, 3, 3, "F")
            self.set_xy(x_leyenda + 4.5, y)
            self.set_font("Helvetica", "", 6.8)
            pct = 100.0 * valor / total
            self.cell(0, 4.5, _s(f"{_recortar(etiqueta, 46)}  "
                                 f"{_fmt(valor)} ({pct:.1f}%)"), 0, 0, "L")
            y += 4.8
        self.set_y(max(y, y0 + 2 * radio) + 3)

    def _salto_si_falta(self, alto):
        if self.get_y() + alto > self.page_break_trigger:
            self.add_page(self.cur_orientation)


def _tintas(colores, etiquetas):
    """Devuelve una funcion (indice, etiqueta) -> color RGB.

    Con `colores` se puede fijar el color por etiqueta (por ejemplo para que
    SUSTANTIVA se lea siempre como severidad alta); si no, se recorre la
    paleta institucional.
    """
    def color(indice, etiqueta):
        if colores:
            elegido = colores.get(str(etiqueta).strip().upper())
            if elegido:
                return elegido
        return _PALETA[indice % len(_PALETA)]
    return color


def _sector(pdf, cx, cy, radio, ang_ini, ang_fin, pasos=None):
    """Dibuja un sector circular como poligono relleno.

    fpdf2 no expone arcos rellenos de forma estable entre versiones, de modo
    que se aproxima el arco con segmentos: a 2 grados por segmento el borde
    resulta indistinguible de una curva en impresion.
    """
    import math
    barrido = ang_fin - ang_ini
    pasos = pasos or max(3, int(abs(barrido) / 2) + 1)
    puntos = [(cx, cy)]
    for i in range(pasos + 1):
        ang = math.radians(ang_ini + barrido * i / pasos - 90)
        puntos.append((cx + radio * math.cos(ang), cy + radio * math.sin(ang)))
    if barrido >= 359.9:
        pdf.ellipse(cx - radio, cy - radio, radio * 2, radio * 2, "F")
        return
    pdf.polygon(puntos, style="F")


def _fmt(valor):
    """Formato numerico compacto para etiquetas de grafico."""
    if valor is None:
        return ""
    if isinstance(valor, float):
        if abs(valor) < 1:
            return f"{valor:.4f}".rstrip("0").rstrip(".")
        if abs(valor) >= 1000:
            return f"{valor:,.0f}"
        return f"{valor:,.2f}".rstrip("0").rstrip(".")
    return str(valor)


def _recortar(texto, largo):
    texto = str(texto or "")
    return texto if len(texto) <= largo else texto[:max(1, largo - 1)] + "."


def _pdf_bytes(pdf):
    """Serializa el PDF en memoria."""
    return bytes(pdf.output())


def generar_pdf_bloque(datos, incluir_graficos=True):
    """Ficha PDF del resumen de un bloque, con graficos y control de
    consistencia. Refleja el libro Excel sin agregar ni estimar valores."""
    codigo = datos.get("codigo_bloque", "")
    pdf = _PDFResumen(subtitulo=f"FICHA RESUMEN - BLOQUE {codigo}")
    pdf.alias_nb_pages()
    pdf.add_page()

    pdf.seccion("1. IDENTIFICACION Y LOCALIZACION")
    pdf.campos([
        ("Codigo del bloque", datos.get("codigo_bloque")),
        ("Microcuenca (catalogo)", datos.get("microcuenca")),
        ("Departamento", datos.get("departamento")),
        ("Provincia", datos.get("provincia")),
        ("Distrito", datos.get("distrito")),
        ("Capital distrital", datos.get("capital_distrital")),
        ("Centro poblado asociado", datos.get("centro_poblado")),
        ("Comunidad campesina", datos.get("comunidad_campesina")),
        ("Superficie de catalogo (ha)", datos.get("area_ha")),
        ("Tipo de intervencion", datos.get("tipo_intervencion")),
        ("Centroide UTM ESTE (m)", datos.get("utm_este")),
        ("Centroide UTM NORTE (m)", datos.get("utm_norte")),
        ("Sistema de coordenadas", datos.get("sistema_coordenadas")),
        ("Validacion UTM 17S", datos.get("validacion_utm")),
    ])

    pdf.seccion("2. PARAMETROS FISICOS (ESTADISTICA ZONAL SOBRE MDE)")
    pdf.campos([
        ("Altitud minima (msnm)", datos.get("altitud_min")),
        ("Altitud maxima (msnm)", datos.get("altitud_max")),
        ("Amplitud altitudinal (m)", datos.get("amplitud_altitudinal")),
        ("Piso altitudinal dominante", datos.get("piso_altitudinal")),
        ("Pendiente promedio (%)", datos.get("pendiente_pct")),
        ("Pendiente promedio (grados)", datos.get("pendiente_grados")),
        ("Clase de pendiente equivalente", datos.get("clase_pendiente")),
        ("Pendiente declarada en campo", datos.get("pendiente_campo")),
        ("Forma del terreno", datos.get("forma_terreno")),
        ("Posicion fisiografica", datos.get("posicion_fisiografica")),
        ("Exposicion / orientacion", datos.get("exposicion")),
        ("Afloramientos rocosos", datos.get("afloramientos_rocosos")),
        ("Escarpes activos", datos.get("escarpes_activos")),
        ("Remociones en masa activas", datos.get("remociones_masa")),
    ])

    pdf.seccion("3. INDICES DE VEGETACION (SENTINEL-2)")
    pdf.campos([
        ("MSAVI 2024 - media del bloque", datos.get("msavi_2024")),
        ("MSAVI 2024 - clase de la media", datos.get("msavi_clase")),
        (f"Condicion frente al umbral {UMBRAL_MSAVI}",
         datos.get("condicion_umbral")),
        ("NDVI 2025 - clase modal", datos.get("ndvi_clase_modal")),
        ("Superficie clasificada NDVI (ha)", datos.get("superficie_ndvi_ha")),
        ("Desviacion frente al catalogo (%)",
         datos.get("desviacion_catalogo_pct")),
    ])

    ndvi = [r for r in datos.get("ndvi_tabla", [])
            if r.get("superficie_ha") is not None]
    if ndvi:
        pdf.tabla(["Clase NDVI", "Superficie (ha)", "% del area",
                   "Interpretacion"],
                  [[r["clase"], _fmt(r["superficie_ha"]), _fmt(r.get("pct")),
                    r.get("interpretacion", "")] for r in ndvi],
                  anchos_rel=[1.5, 1.1, 1.0, 3.4],
                  alineaciones=["L", "R", "R", "L"])
        if incluir_graficos:
            pdf.grafico_torta("Distribucion areal NDVI mediana 2025 (ha)",
                              [r["clase"] for r in ndvi],
                              [r["superficie_ha"] for r in ndvi])

    msavi = [r for r in datos.get("msavi_tabla", [])
             if r.get("superficie_ha") is not None]
    if msavi:
        pdf.tabla(["Clase MSAVI", "Superficie (ha)", "% del area", "Condicion"],
                  [[r["clase"], _fmt(r["superficie_ha"]), _fmt(r.get("pct")),
                    r.get("condicion", "")] for r in msavi],
                  anchos_rel=[1.4, 1.1, 1.0, 2.2],
                  alineaciones=["L", "R", "R", "L"])
        if incluir_graficos:
            pdf.grafico_barras("Clases MSAVI 2024 (ha)",
                               [r["clase"] for r in msavi],
                               [r["superficie_ha"] for r in msavi],
                               unidad=" ha", horizontal=True)
    elif datos.get("msavi_2024_num") is not None:
        pdf.nota("La distribucion areal por clase de MSAVI 2024 no esta "
                 "disponible como estadistica zonal en los insumos. Se "
                 "consigna la media del bloque y no se estima la distribucion, "
                 "conforme a la declaracion de integridad de datos.")

    pdf.seccion("4. ECOSISTEMA Y ESTADO DE CONSERVACION")
    pdf.campos([
        ("Tipo de ecosistema (UP)", datos.get("tipo_ecosistema")),
        ("Superficie de ecosistema", datos.get("superficie_ecosistema")),
        ("Estado de conservacion", datos.get("estado_conservacion")),
        ("Uso actual dominante", datos.get("uso_dominante")),
        ("Tipo de cobertura dominante", datos.get("tipo_cobertura")),
        ("Cobertura vegetal total (%)", datos.get("cobertura_total_pct")),
        ("Suelo desnudo (%)", datos.get("suelo_desnudo_pct")),
        ("Regeneracion natural", datos.get("regeneracion")),
        ("Nivel general de erosion", datos.get("nivel_erosion")),
        ("N. de carcavas registradas", datos.get("n_carcavas")),
        ("Elenco floristico (taxones)", datos.get("n_taxones")),
        ("Estado sanitario", datos.get("estado_sanitario")),
    ])

    estaciones = datos.get("estaciones", [])
    if estaciones:
        pdf.seccion("5. PUNTOS GEORREFERENCIADOS DE LA FICHA DT")
        pdf.tabla(["Codigo", "Naturaleza del punto", "UTM ESTE", "UTM NORTE",
                   "Dist. centroide (m)"],
                  [[r.get("codigo", ""), r.get("naturaleza", ""),
                    _fmt(r.get("utm_este")), _fmt(r.get("utm_norte")),
                    _fmt(r.get("dist_centroide"))] for r in estaciones],
                  anchos_rel=[0.8, 3.0, 1.2, 1.2, 1.2],
                  alineaciones=["L", "L", "R", "R", "R"])
        pdf.nota("La distancia al centroide es referencial: no acredita "
                 "inclusion ni exclusion respecto del poligono, que exige "
                 "prueba de inclusion sobre la geometria del bloque. Las "
                 "estaciones a menos de ~30 m del limite deben considerarse "
                 "SOBRE EL LIMITE (error de georreferenciacion +/- 10-15 m).")

    micro = [r for r in datos.get("microcuenca_tabla", [])
             if r.get("area_ha") is not None]
    if micro:
        pdf.seccion(f"6. CONTEXTO INTRAMICROCUENCA {datos.get('microcuenca', '')}")
        pdf.tabla(["Bloque", "Area (ha)", "% microcuenca", "Rango altitudinal",
                   "Pendiente (%)", "MSAVI 2024"],
                  [[(">" + r["bloque"]) if r.get("es_actual") else r["bloque"],
                    _fmt(r.get("area_ha")), _fmt(r.get("pct_microcuenca")),
                    r.get("rango_altitudinal", ""),
                    _fmt(r.get("pendiente_pct")), _fmt(r.get("msavi"))]
                   for r in micro],
                  anchos_rel=[1.4, 1.1, 1.1, 1.6, 1.1, 1.1],
                  alineaciones=["L", "R", "R", "C", "R", "R"])
        if incluir_graficos and len(micro) > 1:
            pdf.grafico_barras("Superficie por bloque en la microcuenca (ha)",
                               [r["bloque"] for r in micro],
                               [r["area_ha"] for r in micro], unidad=" ha")

    consistencia = datos.get("consistencia", [])
    if consistencia:
        pdf.seccion("7. CONTROL DE CONSISTENCIA DE LA INFORMACION")
        resumen = datos.get("consistencia_resumen", {})
        pdf.campos([("Total de verificaciones", resumen.get("total"))] +
                   [(c.capitalize(), resumen.get(c)) for c in CALIFICACIONES
                    if resumen.get(c)], columnas=3)
        if incluir_graficos:
            etiquetas = [c for c in CALIFICACIONES if resumen.get(c)]
            pdf.grafico_torta("Verificaciones por calificacion", etiquetas,
                              [resumen[c] for c in etiquetas],
                              colores=_COLOR_CALIFICACION)
        pdf.tabla(["Cod.", "Campo afectado", "Discrepancia observada",
                   "Calificacion", "Tratamiento adoptado"],
                  [[r.get("codigo", ""), r.get("campo", ""),
                    r.get("discrepancia", ""), r.get("calificacion", ""),
                    r.get("tratamiento", "")] for r in consistencia],
                  anchos_rel=[0.5, 1.4, 3.2, 1.0, 3.0], tam=6.2)
        pdf.nota("Calificacion: SUSTANTIVA / NO SUSTANTIVA / CORREGIDO / "
                 "CONFORME. Una calificacion CONFORME acredita que la "
                 "verificacion se ejecuto y no arrojo discrepancia; no "
                 "equivale a validacion de campo del dato.")

    pdf.seccion("8. SINTESIS PARA LA GESTION DEL RIESGO EN CONTEXTO DE CAMBIO CLIMATICO")
    pdf.campos([
        ("Causa subyacente principal", datos.get("causa_subyacente")),
        ("Velocidad de degradacion", datos.get("velocidad_degradacion")),
        ("Reversibilidad tecnica", datos.get("reversibilidad")),
        ("Urgencia de intervencion", datos.get("urgencia_intervencion")),
        ("Urgencia de control de erosion", datos.get("urgencia_erosion")),
        ("Zona de recarga hidrica", datos.get("zona_recarga")),
        ("Peligro integrado (MCA-AHP)", datos.get("peligro_integrado")),
        ("Prioridad de intervencion", datos.get("prioridad")),
        ("Estado de verificacion de campo", datos.get("estado_verificacion")),
        ("Marco del indicador de brecha", datos.get("marco_indicador")),
        ("Responsable de la evaluacion", datos.get("evaluador")),
        ("Fecha de evaluacion", datos.get("fecha_evaluacion")),
    ])

    pdf.nota("DECLARACION DE INTEGRIDAD DE DATOS. Ningun valor ausente ha sido "
             "estimado o inferido sin declararlo. Los campos que no pudieron "
             "sustentarse en observacion de campo, estadistica zonal o "
             "catalogo oficial se consignan como 'Por determinar' o 'Por "
             "verificar'. Fuente: " + _s(datos.get("nombre_archivo", "")) + ".")
    return _pdf_bytes(pdf)


def generar_pdf_consolidado(lista_datos, incluir_graficos=True):
    """Reporte PDF consolidado de los bloques cargados (orientacion apaisada)."""
    pdf = _PDFResumen(
        subtitulo=f"CONSOLIDADO DE RESUMENES DT - {len(lista_datos)} BLOQUES",
        orientacion="L")
    pdf.alias_nb_pages()
    pdf.add_page()

    area_total = sum(d.get("area_ha_num") or 0 for d in lista_datos)
    con_msavi = [d["msavi_2024_num"] for d in lista_datos
                 if d.get("msavi_2024_num") is not None]
    bajo_umbral = sum(1 for v in con_msavi if v < UMBRAL_MSAVI)
    sustantivas = sum((d.get("consistencia_resumen") or {}).get("SUSTANTIVA", 0)
                      for d in lista_datos)

    pdf.seccion("1. SINTESIS DEL CONJUNTO")
    pdf.campos([
        ("Bloques con resumen cargado", len(lista_datos)),
        ("Bloques esperados (con ficha DT)", len(codigos_esperados()) or ""),
        ("Superficie total de catalogo (ha)", f"{area_total:,.2f}"),
        ("Bloques con MSAVI 2024 declarado", len(con_msavi)),
        (f"Bloques BAJO umbral {UMBRAL_MSAVI}", bajo_umbral),
        ("MSAVI 2024 promedio",
         f"{sum(con_msavi) / len(con_msavi):.4f}" if con_msavi else ""),
        ("Discrepancias sustantivas acumuladas", sustantivas),
        ("Sistema de referencia", "UTM WGS 84 Zona 17S (EPSG:32717)"),
    ], columnas=3)

    if incluir_graficos:
        por_distrito = {}
        for d in lista_datos:
            clave = d.get("distrito") or "Sin distrito"
            acu = por_distrito.setdefault(clave, [0, 0.0])
            acu[0] += 1
            acu[1] += d.get("area_ha_num") or 0.0
        orden = sorted(por_distrito.items(), key=lambda kv: -kv[1][1])
        pdf.grafico_barras("Superficie por distrito (ha)",
                           [k for k, _ in orden], [v[1] for _, v in orden],
                           unidad=" ha", horizontal=True,
                           alto=max(30, 6.5 * len(orden)))
        pdf.grafico_torta(
            f"Bloques frente al umbral de brecha MSAVI {UMBRAL_MSAVI}",
            [f"BAJO umbral {UMBRAL_MSAVI}", f"Sobre umbral {UMBRAL_MSAVI}",
             "Sin MSAVI declarado"],
            [bajo_umbral, len(con_msavi) - bajo_umbral,
             len(lista_datos) - len(con_msavi)])

    pdf.seccion("2. PARAMETROS POR BLOQUE")
    pdf.tabla(["Bloque", "Microcuenca", "Provincia", "Distrito", "Area (ha)",
               "UTM ESTE", "UTM NORTE", "Pend. (%)", "MSAVI 2024",
               "Sustant.", "Verificacion"],
              [[d.get("codigo_bloque", ""), d.get("microcuenca", ""),
                d.get("provincia", ""), d.get("distrito", ""),
                _fmt(d.get("area_ha_num")), _fmt(d.get("utm_este_num")),
                _fmt(d.get("utm_norte_num")), _fmt(d.get("pendiente_pct_num")),
                _fmt(d.get("msavi_2024_num")),
                (d.get("consistencia_resumen") or {}).get("SUSTANTIVA", 0),
                d.get("estado_verificacion", "")] for d in lista_datos],
              anchos_rel=[1.0, 1.3, 1.1, 1.3, 0.9, 0.9, 1.0, 0.8, 0.9, 0.7, 1.2],
              alineaciones=["L", "L", "L", "L", "R", "R", "R", "R", "R", "R", "L"],
              tam=6.4)
    pdf.nota("Fuente: libros de resumen de Diagnostico Territorial por bloque "
             "(catalogo maestro Bloques V5/V6, estadistica zonal sobre MDE y "
             "compuestos Sentinel-2). Umbral de brecha MSAVI 0.4976 conforme a "
             "la R.M. N. 00213-2024-MINAM.")
    return _pdf_bytes(pdf)
