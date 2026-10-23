"""Paquete de respuesta a la solicitud de información sobre el avance del Perfil IN Piura.

Atiende los seis puntos de la solicitud del Programa de Inversión del Río Piura (ámbito, criterios de
selección, población afectada, Unidad Productora, involucrados e información georreferenciada) con los
datos vigentes del repositorio:

  - datos/gis/IN_Piura_area_influencia_bloques_v6.gpkg     bloques V6, área de influencia, distritos
  - mapas/salidas/IN_Piura_ecosistemas_bloques_v6.gpkg      ecosistemas (MINAM 2018) en el ámbito y bloques
  - datos/unidades_liberacion_areas.csv                     catálogo V6 (117 bloques, 12,270.235 ha)
  - centros_poblados.py                                     centros poblados asociados a cada bloque
  - CentrosPoblados-INEI-Bloques_v5(CON SIN) - 24.09.26.xlsx  IDCCPP, distrito y avance de campo social
  - datos/traspaso_msavi_117.json                           MSAVI 2024 por clase DN y bloque
  - Caracterización de cárcavas.xlsx                        55 cárcavas caracterizadas

Salidas en ./salidas:
  - IN_Piura_Anexos_Avance_Perfil.xlsx      anexos A1-A6 con formato ANIN y totales con fórmulas
  - IN_Piura_Paquete_SIG_Avance_Perfil.zip  SHP, KMZ, GPKG, mapas PDF y LEEME
  - IN_Piura_Mapa_Centros_Poblados_V6.pdf/.png
  - cifras_informe.json                     cifras que usa el informe Word (informe_respuesta.js)

Uso:  python entregables/respuesta_solicitud_rio_piura/generar_paquete_respuesta.py
Requiere geopandas, pyogrio, shapely, pyproj, matplotlib y openpyxl.
"""
import json
import re
import shutil
import sys
import unicodedata
import zipfile
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from xml.sax.saxutils import escape

import geopandas as gpd
import matplotlib
import openpyxl
import pandas as pd
from shapely.geometry import Point

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.backends.backend_pdf import PdfPages  # noqa: E402
from matplotlib.lines import Line2D  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[1]
SALIDA = AQUI / "salidas"
sys.path[:0] = [str(RAIZ), str(RAIZ / "mapas"), str(RAIZ / ".claude" / "skills" / "anin-in-piura" / "scripts")]

from anin_catalogo import AREA_TOTAL_HA, cargar_bloques  # noqa: E402
from anin_excel import BORDE, FILA_ALTERNA, VERDE, encabezado_institucional, escribir_tabla  # noqa: E402
from anin_utm import validar_utm  # noqa: E402
from centros_poblados import CENTROS_POBLADOS_BLOQUE  # noqa: E402
import mapa_area_influencia_bloques_v6 as MAI  # noqa: E402

GPKG_AI = RAIZ / "datos" / "gis" / "IN_Piura_area_influencia_bloques_v6.gpkg"
GPKG_ECO = RAIZ / "mapas" / "salidas" / "IN_Piura_ecosistemas_bloques_v6.gpkg"
XLS_CP = RAIZ / "CentrosPoblados-INEI-Bloques_v5(CON SIN) - 24.09.26.xlsx"
XLS_CARCAVAS = RAIZ / "Caracterización de cárcavas.xlsx"
JSON_MSAVI = RAIZ / "datos" / "traspaso_msavi_117.json"
MAPAS_EXISTENTES = [RAIZ / "mapas" / "salidas" / "IN_Piura_Mapa_Area_Influencia_Bloques_V6.pdf",
                    RAIZ / "mapas" / "salidas" / "IN_Piura_Mapa_Ecosistemas_Bloques_V6.pdf"]
CRS = "EPSG:32717"
HOY = date.today()
PROVINCIAS = ["Ayabaca", "Huancabamba", "Morropón"]
UMBRAL_MSAVI = 0.4976  # límite superior de la clase DN 3 (traspaso MSAVI V6)

FUENTE_BLOQUES = "Bloques V6 (Bloques V6.kml) y catálogo V6 (datos/unidades_liberacion_areas.csv), ANIN-DIME-SESDI 2026"
FUENTE_AI = "Área de influencia aprobada (AI_aprobado_2.shp), ANIN-DIME-SESDI 2026; sin los AI de los bloques retirados 74 y 75"
FUENTE_CP = ("INEI, capa de centros poblados con población censal; asociación centro poblado–bloque del equipo social "
             "(CentrosPoblados-INEI-Bloques_v5, 24.09.2026, y bloques adicionales 83-87)")
FUENTE_ECO = "Mapa Nacional de Ecosistemas del Perú, MINAM 2018 (R.M. N.° 440-2018-MINAM), recorte al ámbito"


# ══════════════════════════════════════════════════════════════════════════
# UTILIDADES
# ══════════════════════════════════════════════════════════════════════════
def clave(texto):
    """Mayúsculas sin tildes ni espacios repetidos, para cruzar nombres entre fuentes."""
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().upper()


def cod_bloque(texto):
    """'BLOQUE - M6B2-2' / 'Bloque - 56' / 'Bloque M17B5' -> 'M6B2-2' / '56' / 'M17B5'."""
    return re.sub(r"(?i)^bloque\s*-?\s*", "", str(texto).strip()).strip()


def num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def nombre_oficial(slug_o_mayus, tabla):
    return tabla.get(clave(slug_o_mayus), str(slug_o_mayus).title())


# ══════════════════════════════════════════════════════════════════════════
# 1. DATOS
# ══════════════════════════════════════════════════════════════════════════
def cargar_catalogo():
    bl = cargar_bloques()
    assert len(bl) == 117, f"Se esperaban 117 bloques V6 y hay {len(bl)}"
    assert abs(sum(b["area_ha"] for b in bl) - AREA_TOTAL_HA) < 0.001
    return {b["codigo"]: b for b in bl}


def cargar_capas(cat):
    b = gpd.read_file(GPKG_AI, layer="bloques_v6")
    ai = gpd.read_file(GPKG_AI, layer="area_influencia_v6")
    amb = gpd.read_file(GPKG_AI, layer="ambito_distritos")
    ctx = gpd.read_file(GPKG_AI, layer="contexto_distritos")
    eco_amb = gpd.read_file(GPKG_ECO, layer="ecosistemas_ambito")
    eco_bl = gpd.read_file(GPKG_ECO, layer="ecosistemas_bloques_v6")
    for g in (b, ai, amb, ctx, eco_amb, eco_bl):
        assert g.crs.to_epsg() == 32717, f"CRS inesperado: {g.crs}"
    assert set(b.BLOQUE) == set(cat), "La capa de bloques no coincide con el catálogo V6"
    assert set(ai.BLOQUE) == set(cat), "Hay bloques V6 sin AI (o AI sin bloque)"
    assert abs(b.AREA_HA.sum() - AREA_TOTAL_HA) < 0.01
    return b, ai, amb, ctx, eco_amb, eco_bl


def tablas_nombres(cat):
    """Mayúsculas sin tildes -> nombre oficial, para distritos y provincias."""
    dist = {clave(b["distrito"]): b["distrito"] for b in cat.values()}
    prov = {clave(b["provincia"]): b["provincia"] for b in cat.values()}
    return dist, prov


def tabla_bloques(cat, b, ai, eco_bl, msavi, cp_por_bloque):
    ai_ha = ai.groupby("BLOQUE").AREA_HA.sum()
    ai_n = ai.groupby("BLOQUE").size()
    dom = (eco_bl.sort_values("HA_OFICIAL", ascending=False).drop_duplicates("BLOQUE")
           .set_index("BLOQUE")[["ECOSISTEMA", "HA_OFICIAL"]])
    filas = []
    for r in b.itertuples():
        c = cat[r.BLOQUE]
        assert clave(c["distrito"]) == clave(r.NOMBDIST), f"Distrito distinto en {r.BLOQUE}"
        m = msavi[r.BLOQUE]
        p = r.geometry.representative_point()
        filas.append({
            "Bloque": r.BLOQUE, "Microcuenca": r.MICROC, "Provincia": c["provincia"], "Distrito": c["distrito"],
            "Área bloque (ha)": c["area_ha"], "Polígonos AI": int(ai_n[r.BLOQUE]),
            "Área AI (ha)": round(float(ai_ha[r.BLOQUE]), 3),
            "MSAVI 2024 media": round(m["media_msavi"], 4),
            "Sup. MSAVI ≤ 0.4976 (ha)": round(m["bajo_umbral_ha"], 3),
            "Ecosistema dominante": dom.loc[r.BLOQUE, "ECOSISTEMA"],
            "% ecosistema dominante": round(dom.loc[r.BLOQUE, "HA_OFICIAL"] / c["area_ha"], 4),
            "CCPP asociados": len(cp_por_bloque.get(r.BLOQUE, [])),
            "Punto interior ESTE (m)": round(p.x, 2), "Punto interior NORTE (m)": round(p.y, 2),
        })
    df = pd.DataFrame(filas)
    df["_p"] = df.Provincia.map({p: i for i, p in enumerate(PROVINCIAS)})
    return df.sort_values(["_p", "Distrito", "Bloque"]).drop(columns="_p").reset_index(drop=True)


def cargar_msavi(cat):
    t = json.load(open(JSON_MSAVI, encoding="utf-8"))
    m = {x["bloque"]: x for x in t["bloques"]}
    assert set(m) == set(cat), "El traspaso MSAVI no cubre exactamente los 117 bloques V6"
    return m, t["verificacion"]


def leer_hoja(ruta, hoja):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    filas = list(wb[hoja].iter_rows(values_only=True))
    wb.close()
    return filas


def centros_poblados(cat, dist_of, prov_of):
    """Centros poblados únicos asociados a los 117 bloques V6.

    Base: centros_poblados.py (catálogo del aplicativo, incluye 83-87 y M17B5). Se enriquece con IDCCPP,
    distrito y provincia del Excel INEI-Bloques (24.09.2026) cruzando bloque + nombre. Un mismo centro poblado
    puede atender a varios bloques: se cuenta una sola vez (IDCCPP o, si falta, nombre + distrito)."""
    filas = leer_hoja(XLS_CP, "centrospobladosineibloques_v5")
    cab = filas[0]
    inei = {}
    for f in filas[1:]:
        if f[0] is None:
            continue
        d = dict(zip(cab, f))
        inei[(cod_bloque(d["BLOQUE"]), clave(d["NOMB_CCPP"]))] = d
    # mismo centro poblado registrado en otro bloque (p. ej. datos «inferidos»): nombre + distrito -> fila INEI
    inei_nd = {(clave(d["NOMB_CCPP"]), clave(d["NOMB_DISTR"])): d for d in inei.values()}
    avance = avance_campo(cat)
    dist_av = {(a["Bloque"], clave(a["Centro poblado"])): a["Distrito"] for a in avance}

    unicos, por_bloque, sin_cp = {}, defaultdict(list), []
    for cod in cat:
        datos = CENTROS_POBLADOS_BLOQUE.get(cod)
        if not datos or not datos["demografia"]:
            sin_cp.append(cod)
            continue
        for e in datos["demografia"]:
            d = inei.get((cod, clave(e["centro_poblado"])))
            if not d:
                dist = dist_av.get((cod, clave(e["centro_poblado"]))) or cat[cod]["distrito"]
                d = inei_nd.get((clave(e["centro_poblado"]), clave(dist)))
            if d:
                idccpp = str(d["IDCCPP"])
                dist = nombre_oficial(d["NOMB_DISTR"], dist_of)
                prov = nombre_oficial(d["NOMB_PROVI"], prov_of)
            else:
                idccpp = ""
                dist = dist_av.get((cod, clave(e["centro_poblado"]))) or cat[cod]["distrito"]
                dist = nombre_oficial(dist, dist_of)
                prov = cat[cod]["provincia"]
            k = idccpp or (clave(e["centro_poblado"]), clave(dist))
            sin_desagregar = e["poblacion_total"] and not (e["hombres"] or e["mujeres"])
            este, norte = num(e.get("utm_este")) or None, num(e.get("utm_norte")) or None
            obs = []
            if sin_desagregar:
                obs.append("Solo población total (sin desagregar por sexo, población vulnerable ni viviendas)")
            if not este:
                obs.append("Sin coordenadas en la fuente")
            if e.get("inferido"):
                obs.append("Datos tomados del mismo centro poblado en otro bloque")
            if d and not este:  # coordenadas del mismo centro poblado en la capa INEI
                este, norte = num(d["ESTE_UTM"]), num(d["NORTE_UTM"])
                obs = [o for o in obs if o != "Sin coordenadas en la fuente"] + ["Coordenadas de la capa INEI"]
            # un centro poblado repetido conserva el registro más completo (con desagregación y coordenadas)
            completo = (not sin_desagregar) + bool(este)
            bloques = unicos[k]["_bloques"] if k in unicos else []
            if k not in unicos or completo > unicos[k]["_completo"]:
                unicos[k] = {
                    "IDCCPP": idccpp, "Centro poblado": e["centro_poblado"], "Provincia": prov, "Distrito": dist,
                    "Tipo": e.get("tipo") or "", "Población total": e["poblacion_total"],
                    "Hombres": None if sin_desagregar else e["hombres"],
                    "Mujeres": None if sin_desagregar else e["mujeres"],
                    "Población vulnerable": None if sin_desagregar else e["poblacion_vulnerable"],
                    "Viviendas particulares": None if sin_desagregar else e["viviendas"],
                    "ESTE (m)": este, "NORTE (m)": norte, "Observación": "; ".join(obs),
                    "_bloques": bloques, "_completo": completo,
                }
            unicos[k]["_bloques"].append(cod)
            por_bloque[cod].append(k)
    df = pd.DataFrame(unicos.values())
    df["N.° bloques"] = df._bloques.map(len)
    df["Bloques asociados"] = df._bloques.map(lambda x: ", ".join(sorted(set(x), key=orden_codigo)))
    df["_p"] = df.Provincia.map({p: i for i, p in enumerate(PROVINCIAS)})
    df = (df.sort_values(["_p", "Distrito", "Centro poblado"]).drop(columns=["_bloques", "_p", "_completo"])
          .reset_index(drop=True))
    # control de coordenadas
    ok = []
    for este, norte in zip(df["ESTE (m)"], df["NORTE (m)"]):
        if pd.notna(este) and pd.notna(norte):
            ok.append("OK" if validar_utm(este, norte)[0] else "FUERA DE RANGO")
        else:
            ok.append("SIN COORDENADAS")
    df["Control UTM"] = ok
    cols = ["IDCCPP", "Centro poblado", "Provincia", "Distrito", "Bloques asociados", "N.° bloques", "Tipo",
            "Población total", "Hombres", "Mujeres", "Población vulnerable", "Viviendas particulares",
            "ESTE (m)", "NORTE (m)", "Control UTM", "Observación"]
    return df[cols], por_bloque, sorted(sin_cp, key=orden_codigo), avance


def orden_codigo(c):
    """Numéricos primero en orden numérico; luego M#B# por microcuenca y bloque."""
    if c.isdigit():
        return (0, int(c), 0, "")
    m = re.match(r"M(\d+)B(\d+)(.*)", c)
    return (1, int(m.group(1)), int(m.group(2)), m.group(3)) if m else (2, 0, 0, c)


def avance_campo(cat):
    """Avance del diagnóstico social por centro poblado y bloque (hoja 'Muestra pob inei+sigrid')."""
    dist_of = {clave(v["distrito"]): v["distrito"] for v in cat.values()}
    filas = leer_hoja(XLS_CP, "Muestra pob inei+sigrid")
    salida = []
    for f in filas:
        if f[0] != "Piura" or not f[4]:
            continue
        b = cod_bloque(f[5])
        if b not in cat:
            continue  # bloques retirados del catálogo V6
        fecha = f[21].date() if isinstance(f[21], datetime) else (f[21] or None)
        salida.append({
            "Provincia": nombre_oficial(f[1], {clave(p): p for p in PROVINCIAS}),
            "Distrito": nombre_oficial(f[2], dist_of), "Centro poblado": str(f[4]).strip(), "Bloque": b,
            "Viviendas (N_h)": num(f[6]), "Muestra asignada + ajuste": num(f[10]),
            "F-DS-01 reportadas": num(f[12]), "Estado F-DS-01": f[14] or "SIN REGISTRO",
            "F-DS-03 reportadas": num(f[15]), "Estado F-DS-03": f[16] or "SIN REGISTRO",
            "Saberes ancestrales reportados": num(f[17]), "Estado saberes": f[18] or "SIN REGISTRO",
            "Fecha de evaluación": fecha, "Observaciones de campo": f[22] or "",
        })
    return salida


def cargar_carcavas(cat, dist_of, prov_of):
    filas = leer_hoja(XLS_CARCAVAS, "Caracterízación de cárcavas")
    cab = filas[0]
    df = pd.DataFrame([dict(zip(cab, f)) for f in filas[1:] if f[0]])
    df["Bloque"] = df.Bloque.astype(str)
    fuera = sorted(set(df.Bloque) - set(cat))
    assert not fuera, f"Cárcavas en bloques que no están en V6: {fuera}"
    for r in df.itertuples():
        assert validar_utm(r.Este, r.Norte)[0], f"Cárcava {r.Codigo} fuera de rango UTM"
    df["Dist"] = df.Dist.map(lambda x: nombre_oficial(x, dist_of))
    df["Prov"] = df.Prov.map(lambda x: nombre_oficial(x, prov_of))
    return df


# ══════════════════════════════════════════════════════════════════════════
# 2. EXCEL DE ANEXOS
# ══════════════════════════════════════════════════════════════════════════
F_TIT = Font(name="Arial", size=10, bold=True, color=VERDE)
F_NOTA = Font(name="Arial", size=9, italic=True, color="595959")
F_NORMAL = Font(name="Arial", size=10)


def nota(ws, fila, texto, ncols=8, font=F_NOTA):
    c = ws.cell(fila, 1, texto)
    c.font = font
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ncols)
    ws.row_dimensions[fila].height = max(15, 13 * (len(texto) // (ncols * 14) + 1))
    return fila + 1


def subtitulo(ws, fila, texto, ncols=8):
    c = ws.cell(fila, 1, texto)
    c.font = F_TIT
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ncols)
    return fila + 1


def tabla_df(ws, fila, df, formatos=None, totales=None, anchos=None):
    filas = [[None if (isinstance(v, float) and pd.isna(v)) else v for v in r] for r in df.itertuples(index=False)]
    return escribir_tabla(ws, fila, list(df.columns), filas, anchos=anchos, totales=totales,
                          formato_num=formatos or {})


def rango(hoja, col_letra, f0, f1):
    return f"'{hoja}'!${col_letra}${f0}:${col_letra}${f1}"


def celda_formula(ws, fila, col, formula, fmt=None, bold=False, fill=None):
    c = ws.cell(fila, col, formula)
    c.font = Font(name="Arial", size=10, bold=bold)
    c.border = BORDE
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    return c


def resumen_con_formulas(ws, fila, claves, cabecera, columnas, total_label="TOTAL"):
    """Tabla de resumen cuyas celdas son fórmulas.

    claves: lista de listas con los valores de las primeras columnas (texto).
    columnas: lista de (nombre, fn(fila_excel) -> fórmula, formato, total) con total 'SUM' o una fórmula."""
    ncl = len(claves[0])
    escribir_tabla(ws, fila, cabecera + [c[0] for c in columnas], [], anchos=None)
    r0 = fila + 1
    for i, k in enumerate(claves):
        r = r0 + i
        fill = FILA_ALTERNA if i % 2 else None
        for j, v in enumerate(k, start=1):
            celda_formula(ws, r, j, v, fill=fill)
        for j, (_, fn, fmt, _) in enumerate(columnas, start=ncl + 1):
            celda_formula(ws, r, j, fn(r), fmt, fill=fill)
    rt = r0 + len(claves)
    celda_formula(ws, rt, 1, total_label, bold=True, fill="D9E6DD")
    for j in range(2, ncl + 1):
        celda_formula(ws, rt, j, None, fill="D9E6DD")
    for j, (_, fn, fmt, tot) in enumerate(columnas, start=ncl + 1):
        letra = get_column_letter(j)
        f = f"=SUM({letra}{r0}:{letra}{rt - 1})" if tot == "SUM" else (tot(rt, r0) if callable(tot) else tot)
        celda_formula(ws, rt, j, f, fmt, bold=True, fill="D9E6DD")
    ws.auto_filter.ref = None
    return rt + 2, r0, rt


def anchos(ws, d):
    for col, w in d.items():
        ws.column_dimensions[col].width = w


def excel_anexos(ruta, cat, amb, tb, cp, avance, eco_bl, eco_amb, msavi, verif_msavi, carc, sin_cp, inventario,
                 involucrados, criterios, notas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Índice"
    f = encabezado_institucional(ws, "ANEXOS DEL INFORME DE AVANCE DEL PERFIL IN PIURA — RESPUESTA A SOLICITUD DE "
                                 "INFORMACIÓN", 4, fuente=f"repositorio del Proyecto IN Piura, corte {HOY:%d.%m.%Y}")
    indice = [
        ("A1_Ambito_Distrito", "1. Delimitación del ámbito", "Bloques V6 y área de influencia por distrito y provincia"),
        ("A1_Microcuencas", "1. Delimitación del ámbito", "Bloques V6 por microcuenca (26 microcuencas)"),
        ("A1_Bloques", "1, 2 y 4", "Detalle de los 117 bloques V6: ubicación, AI, MSAVI, ecosistema, CCPP"),
        ("A2_Criterios", "2. Criterios de selección", "Matriz de criterios de selección y priorización"),
        ("A3_Poblacion_CCPP", "3. Población afectada", "Centros poblados asociados a los bloques V6 (únicos)"),
        ("A3_Poblacion_Distrito", "3. Población afectada", "Población de los CCPP asociados por distrito y provincia"),
        ("A3_Avance_Social", "3 y 5", "Diseño muestral y avance de las fichas F-DS por centro poblado"),
        ("A4_UP_Ecosistemas", "4. Unidad Productora", "Ecosistemas en los bloques V6 (MINAM 2018)"),
        ("A4_UP_Eco_Bloque", "4. Unidad Productora", "Ecosistemas por bloque (detalle)"),
        ("A4_UP_MSAVI", "4. Unidad Productora", "Cobertura vegetal MSAVI 2024 por clase y bloque"),
        ("A4_UP_Carcavas", "4. Unidad Productora", "Cárcavas caracterizadas (55)"),
        ("A5_Involucrados", "5. Involucrados", "Matriz preliminar de involucrados"),
        ("A6_Inventario_SIG", "6. Información georreferenciada", "Inventario de capas entregadas y pendientes"),
        ("Notas", "—", "Fuentes, controles y datos pendientes [POR DEFINIR]"),
    ]
    f = escribir_tabla(ws, f, ["Hoja", "Punto de la solicitud", "Contenido"], [list(x) for x in indice],
                       anchos={"Hoja": 26, "Punto de la solicitud": 30, "Contenido": 70})
    for i, (hoja, *_) in enumerate(indice):
        c = ws.cell(f - len(indice) - 1 + i, 1)
        c.hyperlink = f"#'{hoja}'!A1"
        c.font = Font(name="Arial", size=10, color="1B4F72", underline="single")

    # ── A1_Bloques (detalle; las demás hojas de A1 se calculan con fórmulas sobre esta) ──
    wsb = wb.create_sheet("A1_Bloques")
    f0 = encabezado_institucional(wsb, "A1. DETALLE DE LOS 117 BLOQUES DE INTERVENCIÓN V6", len(tb.columns),
                                  fuente=f"{FUENTE_BLOQUES}. {FUENTE_AI}. MSAVI 2024 (traspaso V6). {FUENTE_ECO}")
    fmt = {"Área bloque (ha)": "#,##0.000", "Área AI (ha)": "#,##0.000", "MSAVI 2024 media": "0.0000",
           "Sup. MSAVI ≤ 0.4976 (ha)": "#,##0.000", "% ecosistema dominante": "0.0%",
           "Punto interior ESTE (m)": "#,##0.00", "Punto interior NORTE (m)": "#,##0.00"}
    fin = tabla_df(wsb, f0, tb, fmt, totales={"Área bloque (ha)": "SUM", "Polígonos AI": "SUM",
                                              "Área AI (ha)": "SUM", "Sup. MSAVI ≤ 0.4976 (ha)": "SUM"},
                   anchos={"Ecosistema dominante": 44, "Microcuenca": 14, "Distrito": 24})
    B0, B1 = f0 + 1, f0 + len(tb)
    col = {c: get_column_letter(i + 1) for i, c in enumerate(tb.columns)}
    nota(wsb, fin, "La columna «CCPP asociados» cuenta los centros poblados vinculados al bloque; un mismo centro "
         "poblado puede atender a varios bloques, por lo que no se suma entre bloques. El punto interior es un punto "
         "garantizado dentro del polígono (UTM WGS84 17S), útil para ubicar el bloque.", len(tb.columns))

    def rb(nombre):
        return rango("A1_Bloques", col[nombre], B0, B1)

    # ── A1_Ambito_Distrito ──
    wsd = wb.create_sheet("A1_Ambito_Distrito", 1)
    f = encabezado_institucional(wsd, "A1. ÁMBITO DE ESTUDIO Y ÁREA DE INFLUENCIA — BLOQUES V6 Y AI POR DISTRITO", 9,
                                 fuente=f"{FUENTE_BLOQUES}. {FUENTE_AI}. Límites distritales INEI (referencial)")
    amb_ha = {clave(r.NOMBDIST): r.AREA_HA for r in amb.itertuples()}
    claves = [[b["provincia"], b["distrito"]] for b in sorted(
        {(v["provincia"], v["distrito"]): v for v in cat.values()}.values(),
        key=lambda x: (PROVINCIAS.index(x["provincia"]), x["distrito"]))]
    for k in claves:
        k.append(amb_ha[clave(k[1])])
    columnas = [
        ("N.° bloques", lambda r: f"=COUNTIFS({rb('Distrito')},B{r})", "0", "SUM"),
        ("Área bloques (ha)", lambda r: f"=SUMIFS({rb('Área bloque (ha)')},{rb('Distrito')},B{r})", "#,##0.000", "SUM"),
        ("Polígonos AI", lambda r: f"=SUMIFS({rb('Polígonos AI')},{rb('Distrito')},B{r})", "0", "SUM"),
        ("Área AI (ha)", lambda r: f"=SUMIFS({rb('Área AI (ha)')},{rb('Distrito')},B{r})", "#,##0.000", "SUM"),
        ("Bloques + AI (ha)", lambda r: f"=E{r}+G{r}", "#,##0.000", "SUM"),
        ("% del área de bloques", lambda r: f"=E{r}/$E${{T}}", "0.0%", lambda rt, r0: f"=SUM(I{r0}:I{rt - 1})"),
    ]
    ft, r0, rt = resumen_con_formulas(wsd, f, claves, ["Provincia", "Distrito", "Área del distrito (ha, ref.)"],
                                      columnas)
    for r in range(r0, rt):
        wsd.cell(r, 9).value = wsd.cell(r, 9).value.replace("{T}", str(rt))
        wsd.cell(r, 3).number_format = "#,##0.0"
    wsd.cell(rt, 3).value = f"=SUM(C{r0}:C{rt - 1})"
    wsd.cell(rt, 3).number_format = "#,##0.0"
    wsd.cell(rt, 2).value = "15 distritos"
    ft = subtitulo(wsd, ft, "Resumen por provincia", 9)
    claves_p = [[p] for p in PROVINCIAS]
    cols_p = [
        ("N.° distritos", lambda r: f"=COUNTIFS($A${r0}:$A${rt - 1},A{r})", "0", "SUM"),
        ("Área de distritos (ha, ref.)", lambda r: f"=SUMIFS($C${r0}:$C${rt - 1},$A${r0}:$A${rt - 1},A{r})", "#,##0.0", "SUM"),
        ("N.° bloques", lambda r: f"=SUMIFS($D${r0}:$D${rt - 1},$A${r0}:$A${rt - 1},A{r})", "0", "SUM"),
        ("Área bloques (ha)", lambda r: f"=SUMIFS($E${r0}:$E${rt - 1},$A${r0}:$A${rt - 1},A{r})", "#,##0.000", "SUM"),
        ("Área AI (ha)", lambda r: f"=SUMIFS($G${r0}:$G${rt - 1},$A${r0}:$A${rt - 1},A{r})", "#,##0.000", "SUM"),
        ("Bloques + AI (ha)", lambda r: f"=E{r}+F{r}", "#,##0.000", "SUM"),
    ]
    ft, p0, pt = resumen_con_formulas(wsd, ft, claves_p, ["Provincia"], cols_p)
    ft = subtitulo(wsd, ft, "Controles", 9)
    controles = [
        ("Total de bloques = 117", f'=IF(D{rt}=117,"OK","REVISAR")'),
        ("Área de bloques = 12,270.235 ha", f'=IF(ABS(E{rt}-12270.235)<0.001,"OK","REVISAR")'),
        ("Área AI = suma de A1_Bloques", f'=IF(ABS(G{rt}-SUM({rb("Área AI (ha)")}))<0.001,"OK","REVISAR")'),
        ("Resumen provincial = distrital", f'=IF(ABS(E{pt}-E{rt})<0.001,"OK","REVISAR")'),
    ]
    for txt, fo in controles:
        celda_formula(wsd, ft, 1, txt)
        wsd.merge_cells(start_row=ft, start_column=1, end_row=ft, end_column=3)
        celda_formula(wsd, ft, 4, fo, bold=True)
        ft += 1
    nota(wsd, ft + 1, "El área de influencia (AI) de cada bloque se asigna al distrito de su bloque. El área de los "
         "distritos proviene de la capa INEI simplificada y es referencial. Cifras en hectáreas, calculadas en UTM "
         "WGS84 17S.", 9)
    anchos(wsd, {"A": 16, "B": 26, "C": 16, "D": 12, "E": 16, "F": 16, "G": 16, "H": 16, "I": 14})
    wsd.freeze_panes = wsd.cell(f + 1, 1)

    # ── A1_Microcuencas ──
    wsm = wb.create_sheet("A1_Microcuencas", 2)
    f = encabezado_institucional(wsm, "A1. BLOQUES V6 POR MICROCUENCA PRIORIZADA", 7, fuente=FUENTE_BLOQUES)
    mic = (tb.groupby("Microcuenca").agg(prov=("Provincia", lambda s: ", ".join(sorted(set(s), key=PROVINCIAS.index))),
                                          dist=("Distrito", lambda s: ", ".join(sorted(set(s)))))
           .reset_index().sort_values("Microcuenca"))
    claves_m = [[r.Microcuenca, r.prov, r.dist] for r in mic.itertuples()]
    cols_m = [
        ("N.° bloques", lambda r: f"=COUNTIFS({rb('Microcuenca')},A{r})", "0", "SUM"),
        ("Área bloques (ha)", lambda r: f"=SUMIFS({rb('Área bloque (ha)')},{rb('Microcuenca')},A{r})", "#,##0.000", "SUM"),
        ("Área AI (ha)", lambda r: f"=SUMIFS({rb('Área AI (ha)')},{rb('Microcuenca')},A{r})", "#,##0.000", "SUM"),
        ("Sup. MSAVI ≤ 0.4976 (ha)", lambda r: f"=SUMIFS({rb('Sup. MSAVI ≤ 0.4976 (ha)')},{rb('Microcuenca')},A{r})",
         "#,##0.000", "SUM"),
    ]
    ft, m0, mt = resumen_con_formulas(wsm, f, claves_m, ["Microcuenca (código)", "Provincia(s)", "Distrito(s)"], cols_m)
    wsm.cell(mt, 2).value = f"{len(claves_m)} microcuencas"
    nota(wsm, ft, "Código de microcuenca según la capa de bloques V6 (campo MICROC). La capa de límites de las "
         "microcuencas priorizadas no forma parte de este paquete: [POR DEFINIR] su remisión por el especialista SIG.", 7)
    anchos(wsm, {"A": 18, "B": 26, "C": 50, "D": 12, "E": 16, "F": 16, "G": 18})
    wsm.freeze_panes = wsm.cell(f + 1, 1)

    # ── A2_Criterios ──
    wsc = wb.create_sheet("A2_Criterios")
    f = encabezado_institucional(wsc, "A2. MATRIZ DE CRITERIOS DE SELECCIÓN Y PRIORIZACIÓN DE LOS ÁMBITOS DE "
                                 "INTERVENCIÓN", 6, fuente="Nota Técnica de Meso y Microlocalización IN (ANIN-DIME, "
                                 "2023) y procedimiento AdR-CCC territorial del Proyecto IN Piura")
    escribir_tabla(wsc, f, ["Paso", "Etapa", "Criterio", "Variable / parámetro y regla", "Fuente de información",
                            "Producto / resultado"], criterios,
                   anchos={"Paso": 6, "Etapa": 20, "Criterio": 28, "Variable / parámetro y regla": 60,
                           "Fuente de información": 36, "Producto / resultado": 40})
    for r in range(f + 1, f + 1 + len(criterios)):
        for c in range(1, 7):
            wsc.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")

    # ── A3_Poblacion_CCPP ──
    wsp = wb.create_sheet("A3_Poblacion_CCPP")
    f0p = encabezado_institucional(wsp, "A3. CENTROS POBLADOS ASOCIADOS A LOS BLOQUES V6 — POBLACIÓN Y VIVIENDAS",
                                   len(cp.columns), fuente=FUENTE_CP)
    fmtp = {"Población total": "#,##0", "Hombres": "#,##0", "Mujeres": "#,##0", "Población vulnerable": "#,##0",
            "Viviendas particulares": "#,##0", "ESTE (m)": "#,##0.00", "NORTE (m)": "#,##0.00"}
    fin = tabla_df(wsp, f0p, cp, fmtp, totales={"Población total": "SUM", "Hombres": "SUM", "Mujeres": "SUM",
                                                 "Población vulnerable": "SUM", "Viviendas particulares": "SUM"},
                   anchos={"Centro poblado": 30, "Distrito": 24, "Bloques asociados": 22, "Observación": 60,
                           "Control UTM": 18})
    P0, P1 = f0p + 1, f0p + len(cp)
    colp = {c: get_column_letter(i + 1) for i, c in enumerate(cp.columns)}
    nota(wsp, fin, "Cada centro poblado figura una sola vez aunque esté asociado a varios bloques. La población es la "
         "censal del centro poblado (no la residente dentro del polígono del bloque). Los bloques "
         f"{', '.join(sin_cp)} no tienen centros poblados asociados en la fuente: [POR DEFINIR]. "
         "Año censal de referencia de la capa INEI: [POR DEFINIR] (confirmar con el especialista social).",
         len(cp.columns))

    def rp(nombre):
        return rango("A3_Poblacion_CCPP", colp[nombre], P0, P1)

    # ── A3_Poblacion_Distrito ──
    wspd = wb.create_sheet("A3_Poblacion_Distrito")
    f = encabezado_institucional(wspd, "A3. POBLACIÓN DE LOS CENTROS POBLADOS ASOCIADOS A LOS BLOQUES V6, POR "
                                 "DISTRITO", 9, fuente=FUENTE_CP)
    dist_cp = sorted({(r.Provincia, r.Distrito) for r in cp.itertuples()},
                     key=lambda x: (PROVINCIAS.index(x[0]), x[1]))
    cols_pd = [
        ("N.° CCPP", lambda r: f"=COUNTIFS({rp('Distrito')},B{r})", "0", "SUM"),
        ("Población total", lambda r: f"=SUMIFS({rp('Población total')},{rp('Distrito')},B{r})", "#,##0", "SUM"),
        ("Hombres", lambda r: f"=SUMIFS({rp('Hombres')},{rp('Distrito')},B{r})", "#,##0", "SUM"),
        ("Mujeres", lambda r: f"=SUMIFS({rp('Mujeres')},{rp('Distrito')},B{r})", "#,##0", "SUM"),
        ("Población vulnerable", lambda r: f"=SUMIFS({rp('Población vulnerable')},{rp('Distrito')},B{r})", "#,##0", "SUM"),
        ("Viviendas particulares", lambda r: f"=SUMIFS({rp('Viviendas particulares')},{rp('Distrito')},B{r})", "#,##0", "SUM"),
        ("% de la población", lambda r: f"=D{r}/$D${{T}}", "0.0%", lambda rt, r0: f"=SUM(I{r0}:I{rt - 1})"),
    ]
    ft, d0, dt = resumen_con_formulas(wspd, f, [list(x) for x in dist_cp], ["Provincia", "Distrito"], cols_pd)
    for r in range(d0, dt):
        wspd.cell(r, 9).value = wspd.cell(r, 9).value.replace("{T}", str(dt))
    wspd.cell(dt, 2).value = f"{len(dist_cp)} distritos"
    ft = subtitulo(wspd, ft, "Resumen por provincia", 9)
    cols_pp = [
        ("N.° CCPP", lambda r: f"=SUMIFS($C${d0}:$C${dt - 1},$A${d0}:$A${dt - 1},A{r})", "0", "SUM"),
        ("Población total", lambda r: f"=SUMIFS($D${d0}:$D${dt - 1},$A${d0}:$A${dt - 1},A{r})", "#,##0", "SUM"),
        ("Hombres", lambda r: f"=SUMIFS($E${d0}:$E${dt - 1},$A${d0}:$A${dt - 1},A{r})", "#,##0", "SUM"),
        ("Mujeres", lambda r: f"=SUMIFS($F${d0}:$F${dt - 1},$A${d0}:$A${dt - 1},A{r})", "#,##0", "SUM"),
        ("Población vulnerable", lambda r: f"=SUMIFS($G${d0}:$G${dt - 1},$A${d0}:$A${dt - 1},A{r})", "#,##0", "SUM"),
        ("Viviendas particulares", lambda r: f"=SUMIFS($H${d0}:$H${dt - 1},$A${d0}:$A${dt - 1},A{r})", "#,##0", "SUM"),
    ]
    ft, _, ppt = resumen_con_formulas(wspd, ft, claves_p, ["Provincia"], cols_pp)
    ft = subtitulo(wspd, ft, "Controles", 9)
    for txt, fo in [
        ("Población por distrito = A3_Poblacion_CCPP", f'=IF(D{dt}=SUM({rp("Población total")}),"OK","REVISAR")'),
        ("Resumen provincial = distrital", f'=IF(C{ppt}=D{dt},"OK","REVISAR")'),
    ]:
        celda_formula(wspd, ft, 1, txt)
        wspd.merge_cells(start_row=ft, start_column=1, end_row=ft, end_column=2)
        celda_formula(wspd, ft, 3, fo, bold=True)
        ft += 1
    nota(wspd, ft + 1, "El distrito es el del centro poblado (INEI), que puede diferir del distrito del bloque al que "
         "está asociado. Hombres, mujeres, población vulnerable y viviendas no incluyen los centros poblados de los "
         "bloques 83 a 87, cuya fuente solo trae la población total.", 9)
    anchos(wspd, {"A": 16, "B": 26, "C": 12, "D": 16, "E": 12, "F": 12, "G": 16, "H": 16, "I": 14})
    wspd.freeze_panes = wspd.cell(f + 1, 1)

    # ── A3_Avance_Social ──
    wsa = wb.create_sheet("A3_Avance_Social")
    av = pd.DataFrame(avance)
    f = encabezado_institucional(wsa, "A3/A5. DISEÑO MUESTRAL Y AVANCE DEL DIAGNÓSTICO SOCIAL (F-DS) POR CENTRO "
                                 "POBLADO Y BLOQUE", len(av.columns),
                                 fuente="Hoja de control «Muestra pob inei+sigrid» del equipo social "
                                        "(CentrosPoblados-INEI-Bloques_v5, 24.09.2026); solo bloques V6")
    f = subtitulo(wsa, f, "Diseño muestral (población finita; muestreo estratificado por centro poblado con "
                  "asignación proporcional a las viviendas)", 8)
    params = [("N — viviendas particulares del marco muestral", 8865.9, "#,##0.0"),
              ("Z — nivel de confianza 95 %", 1.96, "0.00"), ("e — margen de error", 0.05, "0%"),
              ("p — proporción esperada", 0.5, "0.00"), ("q — 1 − p", "=1-C{p}", "0.00")]
    pr = {}
    for i, (txt, v, fm) in enumerate(params):
        r = f + i
        celda_formula(wsa, r, 1, txt)
        wsa.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        pr[txt[0]] = r
        celda_formula(wsa, r, 3, v.replace("{p}", str(pr["p"])) if isinstance(v, str) else v, fm)
    r = f + len(params)
    celda_formula(wsa, r, 1, "n — tamaño de muestra = N·Z²·p·q / (e²·(N−1) + Z²·p·q)", bold=True)
    wsa.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    N, Z, e, p, q = (f"C{pr[k]}" for k in "NZepq")
    celda_formula(wsa, r, 3, f"=ROUND({N}*{Z}^2*{p}*{q}/({e}^2*({N}-1)+{Z}^2*{p}*{q}),0)", "0", bold=True)
    fila_n = r
    r += 2
    r = subtitulo(wsa, r, "Avance de campo (bloques V6)", 8)
    ini_res = r
    r += 7  # espacio para el resumen
    fa0 = r
    fmta = {"Viviendas (N_h)": "#,##0.0", "Muestra asignada + ajuste": "0", "F-DS-01 reportadas": "0",
            "F-DS-03 reportadas": "0", "Saberes ancestrales reportados": "0", "Fecha de evaluación": "dd/mm/yyyy"}
    fin = tabla_df(wsa, fa0, av, fmta, totales={"Muestra asignada + ajuste": "SUM", "F-DS-01 reportadas": "SUM",
                                                "F-DS-03 reportadas": "SUM", "Saberes ancestrales reportados": "SUM"},
                   anchos={"Centro poblado": 26, "Distrito": 20, "Observaciones de campo": 80, "Estado F-DS-01": 14,
                           "Estado F-DS-03": 14, "Estado saberes": 14})
    A0, A1 = fa0 + 1, fa0 + len(av)
    cola = {c: get_column_letter(i + 1) for i, c in enumerate(av.columns)}

    def ra(nombre):
        return f"${cola[nombre]}${A0}:${cola[nombre]}${A1}"

    resumen = [
        ("Registros centro poblado–bloque (V6)", f"=COUNTA({ra('Centro poblado')})", "0"),
        ("Tamaño de muestra calculado (n)", f"=C{fila_n}", "0"),
        ("Fichas F-DS-01 (socioeconómica) reportadas", f"=SUM({ra('F-DS-01 reportadas')})", "0"),
        ("Registros con F-DS-01 COMPLETO", f'=COUNTIF({ra("Estado F-DS-01")},"COMPLETO")', "0"),
        ("Registros con F-DS-03 (entrevista a autoridades) COMPLETO", f'=COUNTIF({ra("Estado F-DS-03")},"COMPLETO")', "0"),
        ("Registros con ficha de saberes ancestrales COMPLETO", f'=COUNTIF({ra("Estado saberes")},"COMPLETO")', "0"),
    ]
    for i, (txt, fo, fm) in enumerate(resumen):
        celda_formula(wsa, ini_res + i, 1, txt)
        wsa.merge_cells(start_row=ini_res + i, start_column=1, end_row=ini_res + i, end_column=2)
        celda_formula(wsa, ini_res + i, 3, fo, fm, bold=True)
    nota(wsa, fin, "Estado «-» o «SIN REGISTRO»: centro poblado descartado o reemplazado por el equipo técnico (ver "
         "observaciones). El marco muestral (N) es el de la hoja de control del equipo social. Las fichas F-DS-02 "
         "(actores), F-DS-04 (talleres), F-DS-05 (conflictos), F-DS-06 (percepción de peligros) y F-DS-07 "
         "(consentimiento) se registran en el aplicativo IN Piura: [POR DEFINIR] adjuntar su consolidado exportado.",
         len(av.columns))

    # ── A4 Ecosistemas ──
    wse = wb.create_sheet("A4_UP_Eco_Bloque")
    eb = eco_bl.drop(columns="geometry").copy()
    eb["Provincia"] = eb.BLOQUE.map(lambda c: cat[c]["provincia"])
    eb["Distrito"] = eb.BLOQUE.map(lambda c: cat[c]["distrito"])
    eb = eb.rename(columns={"BLOQUE": "Bloque", "ECOSISTEMA": "Ecosistema", "SIMBOLO": "Símbolo",
                            "MICROC": "Microcuenca", "AREA_BLOQUE_HA": "Área del bloque (ha)",
                            "HA_OFICIAL": "Área del ecosistema en el bloque (ha)"})
    eb["Símbolo"] = eb["Símbolo"].fillna("—")
    eb["_p"] = eb.Provincia.map(PROVINCIAS.index)
    eb = eb.sort_values(["_p", "Distrito", "Bloque", "Área del ecosistema en el bloque (ha)"],
                        ascending=[True, True, True, False])
    eb = eb[["Bloque", "Microcuenca", "Provincia", "Distrito", "Ecosistema", "Símbolo", "Área del bloque (ha)",
             "Área del ecosistema en el bloque (ha)"]].reset_index(drop=True)
    eb["% del bloque"] = None
    f0e = encabezado_institucional(wse, "A4. ECOSISTEMAS POR BLOQUE V6 (DETALLE)", len(eb.columns),
                                   fuente=f"{FUENTE_ECO}; áreas prorrateadas al área oficial del catálogo V6")
    fin = tabla_df(wse, f0e, eb, {"Área del bloque (ha)": "#,##0.000", "Área del ecosistema en el bloque (ha)": "#,##0.000",
                                  "% del bloque": "0.0%"},
                   totales={"Área del ecosistema en el bloque (ha)": "SUM"},
                   anchos={"Ecosistema": 46, "Distrito": 24})
    for r in range(f0e + 1, f0e + 1 + len(eb)):
        wse.cell(r, 9).value = f"=H{r}/G{r}"
        wse.cell(r, 9).number_format = "0.0%"
    E0, E1 = f0e + 1, f0e + len(eb)

    wsee = wb.create_sheet("A4_UP_Ecosistemas", wb.sheetnames.index("A4_UP_Eco_Bloque"))
    f = encabezado_institucional(wsee, "A4. UNIDAD PRODUCTORA — ECOSISTEMAS EN LOS 117 BLOQUES V6", 8,
                                 fuente=f"{FUENTE_ECO}. {FUENTE_BLOQUES}")
    eco_tot = eco_amb.groupby(["ECOSISTEMA", "SIMBOLO"], dropna=False).AREA_HA.sum()
    orden = (eb.groupby("Ecosistema")["Área del ecosistema en el bloque (ha)"].sum().sort_values(ascending=False))
    claves_e = []
    for e_ in orden.index:
        sim = eb.loc[eb.Ecosistema == e_, "Símbolo"].iloc[0]
        amb_ha_e = float(eco_tot[e_].sum()) if e_ in eco_tot.index.get_level_values(0) else None
        claves_e.append([e_, sim, amb_ha_e])
    re_ = f"'A4_UP_Eco_Bloque'!$E${E0}:$E${E1}"
    rh_ = f"'A4_UP_Eco_Bloque'!$H${E0}:$H${E1}"
    cols_e = [
        ("N.° de bloques", lambda r: f"=COUNTIFS({re_},A{r})", "0", None),
        ("Área en bloques (ha)", lambda r: f"=SUMIFS({rh_},{re_},A{r})", "#,##0.000", "SUM"),
        ("% de los bloques", lambda r: f"=E{r}/$E${{T}}", "0.0%", lambda rt, r0: f"=SUM(F{r0}:F{rt - 1})"),
        ("% del ecosistema en el ámbito intervenido", lambda r: f'=IF(C{r}>0,E{r}/C{r},"—")', "0.00%", None),
    ]
    ft, e0, et = resumen_con_formulas(wsee, f, claves_e, ["Ecosistema", "Símbolo", "Área en el ámbito (ha)"], cols_e)
    for r in range(e0, et):
        wsee.cell(r, 6).value = wsee.cell(r, 6).value.replace("{T}", str(et))
        wsee.cell(r, 3).number_format = "#,##0.000"
    wsee.cell(et, 3).value = f"=SUM(C{e0}:C{et - 1})"
    wsee.cell(et, 3).number_format = "#,##0.000"
    celda_formula(wsee, ft, 1, "Control: área en bloques = 12,270.235 ha")
    wsee.merge_cells(start_row=ft, start_column=1, end_row=ft, end_column=3)
    celda_formula(wsee, ft, 4, f'=IF(ABS(E{et}-12270.235)<0.01,"OK","REVISAR")', bold=True)
    otros = sorted(set(eco_amb.ECOSISTEMA) - set(eb.Ecosistema))
    nota(wsee, ft + 2, f"«Área en el ámbito» suma solo los ecosistemas presentes en los bloques; la capa recortada al "
         f"ámbito cubre {eco_amb.AREA_HA.sum():,.3f} ha e incluye además: {', '.join(otros)}.", 8)
    nota(wsee, ft + 4, "La Unidad Productora es el ecosistema en su totalidad; los bloques de intervención se ubican "
         "dentro de él. El Mapa Nacional de Ecosistemas es de escala regional: el ecosistema de cada bloque se "
         "confirma en campo con la ficha F-DT-03. «Zona agrícola» identifica superficie con uso agropecuario dentro "
         "de los bloques, a revisar en el tamizaje predial y en la verificación de campo.", 8)
    anchos(wsee, {"A": 46, "B": 10, "C": 18, "D": 12, "E": 18, "F": 14, "G": 20})
    wsee.freeze_panes = wsee.cell(f + 1, 1)

    # ── A4_UP_MSAVI ──
    wsv = wb.create_sheet("A4_UP_MSAVI")
    mv = pd.DataFrame([{
        "Bloque": c, "Provincia": cat[c]["provincia"], "Distrito": cat[c]["distrito"],
        "DN 1 ≤ 0.2650 (ha)": m["dn1_ha"], "DN 2 0.2650-0.3813 (ha)": m["dn2_ha"],
        "DN 3 0.3813-0.4976 (ha)": m["dn3_ha"], "DN 4 0.4976-0.6139 (ha)": m["dn4_ha"],
        "DN 5 > 0.6139 (ha)": m["dn5_ha"], "Total MSAVI (ha)": None, "MSAVI media": m["media_msavi"],
        "DN dominante": m["dn_dominante"], "Sup. MSAVI ≤ 0.4976 (ha)": None, "% bajo umbral": None,
    } for c, m in msavi.items()])
    mv["_p"] = mv.Provincia.map(PROVINCIAS.index)
    mv = mv.sort_values(["_p", "Distrito", "Bloque"]).drop(columns="_p").reset_index(drop=True)
    f0v = encabezado_institucional(wsv, "A4. COBERTURA VEGETAL — MSAVI 2024 POR CLASE (DN) Y BLOQUE V6",
                                   len(mv.columns), fuente="AREAS_MSAVI_BLOQUES_V6_REV_HSCM.xlsx (exportación SIG) y "
                                   "traspaso MSAVI V6 (datos/traspaso_msavi_117.json)")
    fmv = {c: "#,##0.000" for c in mv.columns if "(ha)" in c}
    fmv.update({"MSAVI media": "0.0000", "% bajo umbral": "0.0%"})
    fin = tabla_df(wsv, f0v, mv, fmv, totales={c: "SUM" for c in mv.columns if "(ha)" in c},
                   anchos={"Distrito": 24})
    for r in range(f0v + 1, f0v + 1 + len(mv)):
        wsv.cell(r, 9).value = f"=SUM(D{r}:H{r})"
        wsv.cell(r, 12).value = f"=SUM(D{r}:F{r})"
        wsv.cell(r, 13).value = f"=L{r}/I{r}"
    rtv = f0v + 1 + len(mv)
    wsv.cell(rtv, 13).value = f"=L{rtv}/I{rtv}"
    wsv.cell(rtv, 13).number_format = "0.0%"
    nota(wsv, fin, "MSAVI: índice de vegetación ajustado al suelo modificado (Sentinel-2, 2024). Clases DN 1 (menor "
         "cobertura) a DN 5 (mayor cobertura) con los cortes indicados; verificación de la equivalencia DN→clase sobre "
         f"117 bloques: r = {verif_msavi['r_directo']}, error absoluto medio = {verif_msavi['eam_directo']}. "
         "«Sup. MSAVI ≤ 0.4976» (clases DN 1 a 3) es un indicador preliminar de baja cobertura vegetal; la condición "
         "de degradación se confirma con la ficha F-DT-04 y el procedimiento MINAM de áreas degradadas. El total MSAVI "
         "puede diferir ligeramente del área del catálogo por el borde del ráster.", len(mv.columns))

    # ── A4_UP_Carcavas ──
    wsk = wb.create_sheet("A4_UP_Carcavas")
    ck = carc.rename(columns={"Codigo": "Código", "Este": "ESTE (m)", "Norte": "NORTE (m)", "Microc": "Microcuenca",
                              "Dist": "Distrito", "Prov": "Provincia", "Long": "Longitud (m)",
                              "Alt_min": "Altitud mín. (msnm)", "Alt_max": "Altitud máx. (msnm)",
                              "Alt_prom": "Altitud media (msnm)", "Pend_prom": "Pendiente media (%)",
                              "Iirreg": "Índice de irregularidad", "NDVI": "NDVI medio",
                              "Cl_Morf": "Clase morfológica", "Cl_NDVI": "Clase de vegetación (NDVI)"})
    ck = ck[["Código", "Bloque", "Microcuenca", "Provincia", "Distrito", "ESTE (m)", "NORTE (m)", "Longitud (m)",
             "Altitud mín. (msnm)", "Altitud máx. (msnm)", "Altitud media (msnm)", "Pendiente media (%)",
             "Índice de irregularidad", "NDVI medio", "Clase morfológica", "Clase de vegetación (NDVI)"]]
    f0k = encabezado_institucional(wsk, "A4. CÁRCAVAS CARACTERIZADAS EN LOS BLOQUES V6", len(ck.columns),
                                   fuente="Caracterización de cárcavas (gabinete SIG: DEM y NDVI), ANIN-DIME-SESDI 2026")
    fin = tabla_df(wsk, f0k, ck, {"ESTE (m)": "#,##0.00", "NORTE (m)": "#,##0.00", "Longitud (m)": "#,##0.00",
                                  "Altitud media (msnm)": "#,##0.0", "Pendiente media (%)": "0.00",
                                  "Índice de irregularidad": "0.00", "NDVI medio": "0.00"},
                   totales={"Longitud (m)": "SUM"}, anchos={"Distrito": 24, "Clase de vegetación (NDVI)": 24,
                                                           "Clase morfológica": 20})
    nota(wsk, fin, "Unidad de la pendiente media según la fuente: [POR DEFINIR] confirmar (% o grados) con el "
         "especialista SIG. El inventario de cárcavas en campo se registra en la ficha F-DT-02.", len(ck.columns))

    # ── A5_Involucrados ──
    wsi = wb.create_sheet("A5_Involucrados")
    f = encabezado_institucional(wsi, "A5. MATRIZ PRELIMINAR DE INVOLUCRADOS", 7,
                                 fuente="Gabinete del equipo formulador y avance de campo F-DS (versión preliminar, a "
                                        "validar con los registros F-DS-02 del aplicativo IN Piura)")
    escribir_tabla(wsi, f, ["Grupo de involucrados", "Actores identificados", "Problemas percibidos", "Intereses",
                            "Estrategias", "Acuerdos y compromisos", "Posición preliminar"], involucrados,
                   anchos={"Grupo de involucrados": 22, "Actores identificados": 40, "Problemas percibidos": 40,
                           "Intereses": 40, "Estrategias": 40, "Acuerdos y compromisos": 36, "Posición preliminar": 16})
    for r in range(f + 1, f + 1 + len(involucrados)):
        for c in range(1, 8):
            wsi.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")

    # ── A6_Inventario_SIG ──
    wsg = wb.create_sheet("A6_Inventario_SIG")
    f = encabezado_institucional(wsg, "A6. INVENTARIO DE INFORMACIÓN GEORREFERENCIADA", 9,
                                 fuente="Paquete IN_Piura_Paquete_SIG_Avance_Perfil.zip")
    escribir_tabla(wsg, f, ["N.°", "Capa / producto", "Contenido", "Geometría", "Registros", "Superficie (ha)",
                            "Formatos entregados", "Fuente", "Estado"], inventario,
                   anchos={"N.°": 6, "Capa / producto": 34, "Contenido": 50, "Geometría": 12, "Registros": 11,
                           "Superficie (ha)": 15, "Formatos entregados": 26, "Fuente": 46, "Estado": 18},
                   formato_num={"Superficie (ha)": "#,##0.000", "Registros": "#,##0"})
    for r in range(f + 1, f + 1 + len(inventario)):
        for c in range(1, 10):
            wsg.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")

    # ── Notas ──
    wsn = wb.create_sheet("Notas")
    f = encabezado_institucional(wsn, "NOTAS, CONTROLES Y DATOS PENDIENTES", 2)
    for txt in notas:
        c = wsn.cell(f, 1, txt)
        c.font = F_NORMAL if not txt.startswith("[") else Font(name="Arial", size=10, bold=True, color="9C2A00")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        wsn.merge_cells(start_row=f, start_column=1, end_row=f, end_column=2)
        wsn.row_dimensions[f].height = 15 * (len(txt) // 120 + 1)
        f += 1
    wsn.column_dimensions["A"].width = 80
    wsn.column_dimensions["B"].width = 60

    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(ruta)


# ══════════════════════════════════════════════════════════════════════════
# 3. PAQUETE SIG (SHP, KMZ, GPKG)
# ══════════════════════════════════════════════════════════════════════════
def kml_coords(coords):
    return " ".join(f"{x:.7f},{y:.7f},0" for x, y, *_ in coords)


def kml_geom(g):
    if g.geom_type == "Point":
        return f"<Point><coordinates>{kml_coords([g.coords[0]])}</coordinates></Point>"
    if g.geom_type == "Polygon":
        s = (f"<Polygon><tessellate>1</tessellate><outerBoundaryIs><LinearRing><coordinates>"
             f"{kml_coords(g.exterior.coords)}</coordinates></LinearRing></outerBoundaryIs>")
        for i in g.interiors:
            s += f"<innerBoundaryIs><LinearRing><coordinates>{kml_coords(i.coords)}</coordinates></LinearRing></innerBoundaryIs>"
        return s + "</Polygon>"
    if g.geom_type in ("MultiPolygon", "GeometryCollection", "MultiPoint"):
        return "<MultiGeometry>" + "".join(kml_geom(x) for x in g.geoms if not x.is_empty) + "</MultiGeometry>"
    raise ValueError(g.geom_type)


def kml_color(hex_rgb, alpha):
    """'#RRGGBB' + alpha 0-1 -> 'aabbggrr' de KML."""
    h = hex_rgb.lstrip("#")
    return f"{int(alpha * 255):02x}{h[4:6]}{h[2:4]}{h[0:2]}"


def kml_carpeta(nombre, gdf, campo_nombre, estilo_de, estilos, etiqueta_punto=False, descripcion=""):
    g = gdf.to_crs("EPSG:4326")
    campos = [c for c in g.columns if c != "geometry"]
    pm = []
    for (_, r), (_, r32) in zip(g.iterrows(), gdf.iterrows()):
        filas = "".join(f"<tr><td><b>{escape(str(c))}</b></td><td>{escape('' if pd.isna(r[c]) else str(r[c]))}</td></tr>"
                        for c in campos)
        geom = kml_geom(r.geometry)
        if etiqueta_punto:
            p = gpd.GeoSeries([r32.geometry.representative_point()], crs=gdf.crs).to_crs("EPSG:4326").iloc[0]
            geom = f"<MultiGeometry>{kml_geom(p)}{geom}</MultiGeometry>"
        pm.append(f"<Placemark><name>{escape(str(r[campo_nombre]))}</name><styleUrl>#{estilo_de(r)}</styleUrl>"
                  f"<description><![CDATA[<table border='1' cellpadding='2' style='border-collapse:collapse;"
                  f"font-family:Arial;font-size:11px'>{filas}</table>]]></description>{geom}</Placemark>")
    return (f"<Folder><name>{escape(nombre)}</name><description>{escape(descripcion)}</description>"
            + "".join(pm) + "</Folder>"), estilos


def estilo_poly(id_, linea, relleno, alpha, ancho=1.5, etiqueta=0.0):
    return (f'<Style id="{id_}"><LineStyle><color>{kml_color(linea, 1)}</color><width>{ancho}</width></LineStyle>'
            f'<PolyStyle><color>{kml_color(relleno, alpha)}</color><fill>{1 if alpha > 0 else 0}</fill></PolyStyle>'
            f'<IconStyle><scale>0</scale></IconStyle><LabelStyle><scale>{etiqueta}</scale></LabelStyle></Style>')


def estilo_punto(id_, color, escala=0.7, etiqueta=0.7):
    return (f'<Style id="{id_}"><IconStyle><color>{kml_color(color, 1)}</color><scale>{escala}</scale><Icon><href>'
            f'http://maps.google.com/mapfiles/kml/shapes/placemark_circle.png</href></Icon></IconStyle>'
            f'<LabelStyle><scale>{etiqueta}</scale></LabelStyle></Style>')


def escribir_kmz(ruta, titulo, carpetas, estilos):
    kml = ('<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document>'
           f"<name>{escape(titulo)}</name>" + "".join(estilos) + "".join(carpetas) + "</Document></kml>")
    with zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("doc.kml", kml)


ECO_COLOR = {"Bosque estacionalmente seco de colina y montaña": "#C9A55A", "Zona agrícola": "#F3E58A",
             "Matorral andino": "#D8B48F", "Bosque relicto montano de vertiente occidental": "#2E7D4F",
             "Páramo": "#9C8FC4"}


def color_eco(nombre):
    return ECO_COLOR.get(nombre) or next((c for k, c in [("urban", "#A6A6A6"), ("río", "#6FA8DC"), ("rio", "#6FA8DC"),
                                                         ("lag", "#6FA8DC"), ("bosque", "#4F9A5B")]
                                          if k in nombre.lower()), "#BFBFBF")


def capas_sig(cat, b, ai, amb, ctx, eco_amb, eco_bl, tb, cp, carc):
    """Capas del paquete (EPSG:32717) con nombres de campo ≤ 10 caracteres (shapefile)."""
    t = tb.set_index("Bloque")
    bl = b[["BLOQUE", "MICROC", "geometry"]].copy()
    bl["DISTRITO"] = bl.BLOQUE.map(lambda c: cat[c]["distrito"])
    bl["PROVINCIA"] = bl.BLOQUE.map(lambda c: cat[c]["provincia"])
    bl["AREA_HA"] = bl.BLOQUE.map(lambda c: cat[c]["area_ha"])
    bl["AI_HA"] = bl.BLOQUE.map(t["Área AI (ha)"])
    bl["MSAVI_MED"] = bl.BLOQUE.map(t["MSAVI 2024 media"])
    bl["MSAVI_BAJ"] = bl.BLOQUE.map(t["Sup. MSAVI ≤ 0.4976 (ha)"])
    bl["ECO_DOM"] = bl.BLOQUE.map(t["Ecosistema dominante"])
    bl["N_CCPP"] = bl.BLOQUE.map(t["CCPP asociados"])
    bl = bl[["BLOQUE", "MICROC", "PROVINCIA", "DISTRITO", "AREA_HA", "AI_HA", "MSAVI_MED", "MSAVI_BAJ", "ECO_DOM",
             "N_CCPP", "geometry"]]

    aic = ai[["CODIGO", "BLOQUE", "MICROC", "AREA_HA", "geometry"]].copy()
    aic["DISTRITO"] = aic.BLOQUE.map(lambda c: cat[c]["distrito"])
    aic["PROVINCIA"] = aic.BLOQUE.map(lambda c: cat[c]["provincia"])
    aic = aic[["CODIGO", "BLOQUE", "MICROC", "PROVINCIA", "DISTRITO", "AREA_HA", "geometry"]]

    dist_of = {clave(v["distrito"]): v["distrito"] for v in cat.values()}
    am = amb.rename(columns={"NOMBDIST": "DISTRITO", "NOMBPROV": "PROVINCIA", "NOMBDEP": "DEPARTAM",
                             "NOM_CAP": "CAPITAL"}).copy()
    am["DISTRITO"] = am.DISTRITO.map(lambda x: nombre_oficial(x, dist_of))
    am["PROVINCIA"] = am.PROVINCIA.map(lambda x: {"AYABACA": "Ayabaca", "HUANCABAMBA": "Huancabamba",
                                                  "MORROPON": "Morropón"}[x])
    am = am[["IDDIST", "DISTRITO", "PROVINCIA", "DEPARTAM", "CAPITAL", "AREA_HA", "geometry"]]

    cpg = cp[cp["ESTE (m)"].notna()].copy()
    cpg = gpd.GeoDataFrame(cpg, geometry=[Point(x, y) for x, y in zip(cpg["ESTE (m)"], cpg["NORTE (m)"])], crs=CRS)
    cpg = cpg.rename(columns={"Centro poblado": "NOMB_CCPP", "Provincia": "PROVINCIA", "Distrito": "DISTRITO",
                              "Bloques asociados": "BLOQUES", "N.° bloques": "N_BLOQUES", "Tipo": "TIPO",
                              "Población total": "POB_TOTAL", "Hombres": "HOMBRES", "Mujeres": "MUJERES",
                              "Población vulnerable": "POB_VULN", "Viviendas particulares": "VIVIENDAS",
                              "ESTE (m)": "ESTE_UTM", "NORTE (m)": "NORTE_UTM", "Observación": "OBSERV"})
    cpg = cpg[["IDCCPP", "NOMB_CCPP", "PROVINCIA", "DISTRITO", "BLOQUES", "N_BLOQUES", "TIPO", "POB_TOTAL", "HOMBRES",
               "MUJERES", "POB_VULN", "VIVIENDAS", "ESTE_UTM", "NORTE_UTM", "OBSERV", "geometry"]]

    ea = eco_amb[["ECOSISTEMA", "SIMBOLO", "AREA_HA", "geometry"]].copy()
    ebl = eco_bl[["BLOQUE", "ECOSISTEMA", "SIMBOLO", "MICROC", "HA_OFICIAL", "PCT", "geometry"]].copy()
    ebl = ebl.rename(columns={"HA_OFICIAL": "AREA_HA", "PCT": "PCT_BLOQUE"})

    kc = carc.copy()
    kc = gpd.GeoDataFrame(kc, geometry=[Point(x, y) for x, y in zip(kc.Este, kc.Norte)], crs=CRS)
    kc = kc.rename(columns={"Codigo": "CODIGO", "Bloque": "BLOQUE", "Microc": "MICROC", "Dist": "DISTRITO",
                            "Prov": "PROVINCIA", "Este": "ESTE_UTM", "Norte": "NORTE_UTM", "Long": "LONG_M",
                            "Alt_min": "ALT_MIN", "Alt_max": "ALT_MAX", "Alt_prom": "ALT_PROM", "Pend_prom": "PEND_PROM",
                            "Iirreg": "I_IRREG", "NDVI": "NDVI", "Cl_Morf": "CL_MORF", "Cl_NDVI": "CL_NDVI"})
    kc = kc[["CODIGO", "BLOQUE", "MICROC", "PROVINCIA", "DISTRITO", "ESTE_UTM", "NORTE_UTM", "LONG_M", "ALT_MIN",
             "ALT_MAX", "ALT_PROM", "PEND_PROM", "I_IRREG", "NDVI", "CL_MORF", "CL_NDVI", "geometry"]]

    cx = ctx.rename(columns={"NOMBDIST": "DISTRITO", "NOMBPROV": "PROVINCIA", "NOMBDEP": "DEPARTAM",
                             "NOM_CAP": "CAPITAL"})
    return [
        ("01_ambito_estudio_distritos", "Ámbito de estudio: 15 distritos (INEI, referencial)", am),
        ("02_bloques_intervencion_v6", "Bloques de intervención V6 (117)", bl),
        ("03_area_influencia_v6", "Área de influencia aprobada (130 polígonos)", aic),
        ("04_centros_poblados_asociados", "Centros poblados asociados a los bloques V6", cpg),
        ("05_ecosistemas_ambito", "Ecosistemas en el ámbito (MINAM 2018)", ea),
        ("06_ecosistemas_bloques_v6", "Ecosistemas dentro de los bloques V6", ebl),
        ("07_carcavas_caracterizadas", "Cárcavas caracterizadas (55)", kc),
        ("08_limites_distritales_contexto", "Límites distritales de contexto (INEI, simplificado)", cx),
    ]


def paquete_sig(capas, cp_sin_coord, ruta_zip, mapas_pdf):
    tmp = SALIDA / "_sig"
    if tmp.exists():
        shutil.rmtree(tmp)
    (tmp / "SHP").mkdir(parents=True)
    (tmp / "KMZ").mkdir()
    (tmp / "GPKG").mkdir()
    (tmp / "MAPAS").mkdir()
    gpkg = tmp / "GPKG" / "IN_Piura_Avance_Perfil_SIG.gpkg"
    inventario_capas = []
    for nombre, titulo, g in capas:
        g = g.reset_index(drop=True)
        assert g.crs.to_epsg() == 32717
        g.to_file(gpkg, layer=nombre, driver="GPKG", engine="pyogrio")
        d = tmp / "SHP" / nombre
        d.mkdir()
        g.to_file(d / f"{nombre}.shp", driver="ESRI Shapefile", engine="pyogrio", encoding="UTF-8")
        assert (d / f"{nombre}.prj").exists() and (d / f"{nombre}.cpg").exists()
        inventario_capas.append((nombre, titulo, g))

    # KMZ por capa (salvo el contexto) y un KMZ integrado
    estilos = [
        estilo_poly("ambito", "#1B4D2E", "#E8F0EA", 0.15, 2.5),
        estilo_poly("bloque", "#7B1E14", "#C0392B", 0.55, 1.5, etiqueta=0.8),
        estilo_poly("ai", "#8A6410", "#F0AE1E", 0.45, 1.0),
        estilo_punto("ccpp", "#1B4F72", 0.6, 0.6),
        estilo_punto("carcava", "#6E2C00", 0.7, 0.6),
    ] + [estilo_poly(f"eco{i}", "#666666", color_eco(e), 0.65, 0.6)
         for i, e in enumerate(sorted(set(capas[4][2].ECOSISTEMA) | set(capas[5][2].ECOSISTEMA)))]
    eco_id = {e: f"eco{i}" for i, e in enumerate(sorted(set(capas[4][2].ECOSISTEMA) | set(capas[5][2].ECOSISTEMA)))}
    defs = {
        "01_ambito_estudio_distritos": ("DISTRITO", lambda r: "ambito", False),
        "02_bloques_intervencion_v6": ("BLOQUE", lambda r: "bloque", True),
        "03_area_influencia_v6": ("CODIGO", lambda r: "ai", False),
        "04_centros_poblados_asociados": ("NOMB_CCPP", lambda r: "ccpp", False),
        "05_ecosistemas_ambito": ("ECOSISTEMA", lambda r: eco_id[r.ECOSISTEMA], False),
        "06_ecosistemas_bloques_v6": ("BLOQUE", lambda r: eco_id[r.ECOSISTEMA], False),
        "07_carcavas_caracterizadas": ("CODIGO", lambda r: "carcava", False),
    }
    todas = []
    for nombre, titulo, g in inventario_capas:
        if nombre not in defs:
            continue
        campo, est, etq = defs[nombre]
        carpeta, _ = kml_carpeta(titulo, g, campo, est, estilos, etiqueta_punto=etq,
                                 descripcion="Proyecto IN Piura · CUI 2669244 · ANIN-DIME-SESDI")
        escribir_kmz(tmp / "KMZ" / f"{nombre}.kmz", f"IN Piura — {titulo}", [carpeta], estilos)
        todas.append(carpeta)
    escribir_kmz(tmp / "KMZ" / "00_IN_Piura_Avance_Perfil_integrado.kmz",
                 "Proyecto IN Piura — Avance del Perfil (CUI 2669244)", todas, estilos)

    for p in mapas_pdf:
        shutil.copy(p, tmp / "MAPAS" / p.name)
    (tmp / "LEEME.txt").write_text(leeme(inventario_capas, cp_sin_coord), encoding="utf-8")

    if ruta_zip.exists():
        ruta_zip.unlink()
    with zipfile.ZipFile(ruta_zip, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(tmp.rglob("*")):
            if p.is_file():
                z.write(p, Path("IN_Piura_Paquete_SIG_Avance_Perfil") / p.relative_to(tmp))
    shutil.rmtree(tmp)
    return inventario_capas


def leeme(capas, cp_sin_coord):
    lineas = [
        "PROYECTO IN PIURA — CUI 2669244",
        "AUTORIDAD NACIONAL DE INFRAESTRUCTURA - ANIN · DIME · SESDI",
        "Paquete de información georreferenciada del avance del Perfil de Inversión",
        f"Fecha de generación: {HOY:%d/%m/%Y}",
        "",
        "SISTEMA DE COORDENADAS",
        "  SHP y GPKG: UTM WGS 84 Zona 17S (EPSG:32717). Cada shapefile incluye .prj y .cpg (UTF-8).",
        "  KMZ: geográficas WGS 84 (EPSG:4326), formato requerido por KML/Google Earth.",
        "  Áreas en hectáreas calculadas en UTM 17S.",
        "",
        "ESTRUCTURA",
        "  SHP/<capa>/      shapefile de cada capa (.shp .shx .dbf .prj .cpg)",
        "  KMZ/             una capa por archivo y 00_..._integrado.kmz con todas las capas",
        "  GPKG/            IN_Piura_Avance_Perfil_SIG.gpkg (todas las capas en un archivo)",
        "  MAPAS/           láminas A3 en PDF (área de influencia, ecosistemas, centros poblados)",
        "",
        "CAPAS",
    ]
    for nombre, titulo, g in capas:
        area = f" · {g.geometry.area.sum() / 1e4:,.3f} ha (geométrica)" if g.geom_type.iloc[0] != "Point" else ""
        lineas.append(f"  {nombre}: {titulo} · {len(g)} registros · {g.geom_type.iloc[0]}{area}")
        lineas.append(f"      Campos: {', '.join(c for c in g.columns if c != 'geometry')}")
    lineas += [
        "",
        "NOTAS",
        "  - AREA_HA de bloques es el área oficial del catálogo V6 (suma 12,270.235 ha); AREA_HA del AI es el atributo",
        "    de AI_aprobado_2.shp (suma 17,139.655 ha). Se excluyen los bloques retirados en V6.",
        "  - Los centros poblados se asocian a los bloques por cercanía y acceso (equipo social); la población es la",
        "    censal del centro poblado (INEI), no la residente dentro del bloque. Un centro poblado puede atender a",
        "    varios bloques (campo BLOQUES).",
        f"  - Centros poblados sin coordenadas en la fuente (no incluidos en la capa de puntos): {cp_sin_coord}.",
        "  - El Mapa Nacional de Ecosistemas (MINAM 2018) es de escala regional; el ecosistema de cada bloque se",
        "    confirma en campo (F-DT-03).",
        "  - MSAVI_MED: MSAVI 2024 medio del bloque; MSAVI_BAJ: ha con MSAVI ≤ 0.4976 (clases DN 1-3).",
        "  - Pendiente de remisión [POR DEFINIR]: capas de peligro (PMM, EPH, PGI y peligro integrado), límites de",
        "    las 26 microcuencas priorizadas, capas de filtros de exclusión e idoneidad (arqueología, concesiones",
        "    mineras, infraestructura eléctrica, ANP, accesibilidad, déficit hídrico), tamizaje predial y",
        "    ortofotos/vuelos de dron. Se remitirán en sus formatos nativos al ser entregadas por el especialista SIG.",
    ]
    return "\n".join(lineas) + "\n"


# ══════════════════════════════════════════════════════════════════════════
# 4. MAPA DE CENTROS POBLADOS
# ══════════════════════════════════════════════════════════════════════════
def mapa_ccpp(b, ai, amb, ctx, cpg, cp, ruta_pdf, ruta_png):
    fig = plt.figure(figsize=(16.54, 11.69))
    n_coord = len(cpg)
    MAI.marco(fig, "CENTROS POBLADOS ASOCIADOS A LOS BLOQUES DE INTERVENCIÓN V6",
              f"{len(cp)} centros poblados ({int(cp['Población total'].sum()):,} habitantes) asociados a los 117 "
              f"bloques V6 — {n_coord} con coordenadas en la fuente", 1)
    ax = fig.add_axes([0.04, 0.075, 0.62, 0.755])
    xlim, ylim = MAI.encuadre(amb, 0.03, aspecto=0.62 * 16.54 / (0.755 * 11.69))
    MAI.dibujar_base(ax, b, ai, amb, ctx, xlim, ylim)
    MAI.etiquetas_distritos(ax, amb, xlim, ylim, tam=7.0)
    tam = (cpg.POB_TOTAL.clip(lower=10) ** 0.5) * 2.2
    ax.scatter(cpg.geometry.x, cpg.geometry.y, s=tam, facecolor="#1B4F72", edgecolor="white", linewidth=0.6,
               alpha=0.9, zorder=9)
    MAI.ejes_utm(ax, paso=20000)
    MAI.norte_y_escala(ax, 20)

    panel = fig.add_axes([0.69, 0.065, 0.29, 0.77])
    panel.axis("off")
    ax_ub = fig.add_axes([0.70, 0.60, 0.27, 0.23])
    MAI.mapa_ubicacion(ax_ub, ctx, amb)
    # cuadro por provincia
    res = cp.groupby("Provincia").agg(n=("Centro poblado", "size"), pob=("Población total", "sum"))
    res = res.reindex(PROVINCIAS)
    filas = [[p, f"{int(r.n)}", f"{int(r.pob):,}"] for p, r in res.iterrows()]
    filas.append(["TOTAL", f"{int(res.n.sum())}", f"{int(res.pob.sum()):,}"])
    alto = 0.0195 * (len(filas) + 1)
    y_top = 0.635
    t = panel.table(cellText=filas, colLabels=["Provincia", "CCPP", "Población"], colWidths=[0.5, 0.2, 0.3],
                    cellLoc="left", bbox=[0.0, y_top - alto, 1.0, alto])
    t.auto_set_font_size(False)
    t.set_fontsize(7.5)
    for (i, j), c in t.get_celld().items():
        c.set_edgecolor("#B0B7B3")
        c.set_linewidth(0.4)
        if j >= 1:
            c.set_text_props(ha="right")
        if i == 0:
            c.set_facecolor(MAI.VERDE)
            c.set_text_props(color="white", fontweight="bold", ha="center")
        elif i == len(filas):
            c.set_facecolor("#D5E8D4")
            c.set_text_props(fontweight="bold")
    panel.text(0.0, y_top + 0.012, "CENTROS POBLADOS ASOCIADOS POR PROVINCIA", fontsize=9, fontweight="bold",
               color=MAI.VERDE, transform=panel.transAxes)
    extra = [Line2D([], [], marker="o", linestyle="", markerfacecolor="#1B4F72", markeredgecolor="white",
                    markersize=s, label=f"Centro poblado · {p:,} hab.")
             for p, s in [(50, 4), (500, 7), (2000, 10)]]
    MAI.leyenda(panel, y_top - alto - 0.01, extra=extra)
    panel.text(0.0, 0.012,
               "Tamaño del símbolo proporcional a la población censal del centro\n"
               "poblado (INEI). Un centro poblado puede atender a varios bloques;\n"
               "se representa una sola vez. Centros poblados sin coordenadas en la\n"
               f"fuente: {len(cp) - n_coord} (bloques 83 a 87).",
               fontsize=7.3, color="#333333", transform=panel.transAxes, va="bottom",
               bbox=dict(boxstyle="round,pad=0.4", facecolor="#F7F3E3", edgecolor=MAI.DORADO, linewidth=0.8))
    fig.text(0.02, 0.034, "Sistema de coordenadas: UTM WGS84 Zona 17S (EPSG:32717)", fontsize=7.5,
             color="#333333", fontweight="bold")
    fig.text(0.02, 0.019, f"Fuente: {FUENTE_CP}. {FUENTE_BLOQUES}. Elaboración: ANIN-DIME-SESDI.", fontsize=6.3,
             color="#555555")
    fig.savefig(ruta_png, dpi=150)
    with PdfPages(ruta_pdf) as pp:
        pp.savefig(fig)
    plt.close(fig)


# ══════════════════════════════════════════════════════════════════════════
# 5. CONTENIDO CUALITATIVO (criterios e involucrados)
# ══════════════════════════════════════════════════════════════════════════
CRITERIOS = [
    [1, "Mesolocalización", "Peligro integrado (PI) en la cuenca alta",
     "Modelamiento GIS por análisis multicriterio (MCA) y proceso analítico jerárquico (AHP) de: peligro por "
     "movimientos en masa (PMM: geología, pendiente, TPI, cobertura vegetal; desencadenante: precipitación acumulada "
     "multianual; parámetro: tendencia a la degradación TAU/NDVI), erosión potencial hídrica (EPH) y peligro "
     "generador de inundaciones (PGI: pendiente, escorrentía total, NDVI). PI en 5 niveles (muy bajo a muy alto).",
     "Sentinel-2, Landsat, DEM, WorldClim, SENAMHI, INGEMMET",
     "Mapa de peligro integrado y microcuencas priorizadas (26 microcuencas con bloques V6)"],
    [2, "Digitalización", "Áreas con PI medio, alto y muy alto",
     "Delimitación semiautomática o manual de bloques sobre las microcuencas priorizadas, en zonas de origen del "
     "peligro (cuenca alta), no en la zona de impacto.", "Mapa de PI; imágenes satelitales",
     "Bloques preliminares de intervención"],
    [3, "Filtros de exclusión", "Restos arqueológicos", "Exclusión de superficies con sitios arqueológicos declarados.",
     "SIGDA — Ministerio de Cultura", "Bloques depurados"],
    [3, "Filtros de exclusión", "Concesiones mineras activas", "Exclusión de superficies con concesiones mineras "
     "vigentes.", "INGEMMET (GEOCATMIN)", "Bloques depurados"],
    [3, "Filtros de exclusión", "Infraestructura eléctrica", "Exclusión de líneas eléctricas con franja (buffer) de 10 m.",
     "OSINERGMIN", "Bloques depurados"],
    [3, "Filtros de exclusión", "Áreas Naturales Protegidas", "Exclusión de superficies dentro de ANP.",
     "SERNANP", "Bloques depurados"],
    [4, "Idoneidad", "Accesibilidad", "Acceso desde la red vial y caminos identificables.",
     "Red vial MTC; imágenes satelitales", "Bloques accesibles"],
    [4, "Idoneidad", "Pendiente máxima", "Pendiente ≤ 75 %.", "DEM", "Bloques técnicamente intervenibles"],
    [4, "Idoneidad", "Cobertura territorial", "Cobertura de los distritos priorizados (15 distritos de Ayabaca, "
     "Huancabamba y Morropón).", "Límites INEI", "Distribución territorial de bloques"],
    [4, "Idoneidad", "Déficit hídrico futuro", "Déficit hídrico proyectado con el modelo MPI-ESM1-2HR, escenario SSP5-8.5, "
     "periodo 2021-2040.", "Escenarios CMIP6", "Priorización por exposición climática"],
    [5, "Tamizaje predial", "Predios interceptados", "Máximo 19 predios rurales interceptados por bloque.",
     "Catastro rural (predios)", "Bloques con viabilidad de saneamiento"],
    [5, "Tamizaje predial", "Uso agrícola", "Exclusión de áreas agrícolas.", "Imágenes satelitales; catastro",
     "Bloques sin conflicto de uso"],
    [6, "Verificación de campo", "Accesibilidad real y uso actual",
     "Verificación de accesos, uso actual del suelo y evidencia de movimientos en masa y procesos erosivos "
     "(fichas F-DT-01 a F-DT-05; elementos expuestos F-EE-01 a F-EE-07).",
     "Brigadas de campo; aplicativo IN Piura / KoboToolbox; vuelos de dron",
     "Bloques verificados (catálogo V6: 117 bloques, 12,270.235 ha)"],
    [6, "Verificación de campo", "Riesgo social y aceptación",
     "Riesgo social, aceptación de titulares y posesionarios (fichas F-DS y F-LA de liberación de áreas: "
     "identificación de prediantes, actas).", "Equipo social y legal", "Bloques con aceptación; codificación M#-B#"],
]

INVOLUCRADOS = [
    ["Población de los centros poblados asociados a los bloques",
     "Habitantes de los centros poblados rurales vinculados a los 117 bloques (ver A3); autoridades de caserío "
     "(tenientes gobernadores, agentes municipales) entrevistadas con la ficha F-DS-03",
     "Pérdida de cobertura vegetal y de suelos; erosión y cárcavas; movimientos en masa que afectan caminos, "
     "viviendas y terrenos; menor disponibilidad de agua en estiaje",
     "Reducir la exposición a peligros; mantener o mejorar la disponibilidad de agua y de pastos; empleo local",
     "Diagnóstico socioeconómico (F-DS-01), talleres participativos (F-DS-04), consulta de disposición a participar "
     "y consentimiento (F-DS-07)",
     "Socialización con autoridades de caserío en curso. Acuerdos formales: [POR DEFINIR]", "A favor / por confirmar"],
    ["Titulares y posesionarios de predios en los bloques",
     "Propietarios, posesionarios, sucesiones, ocupantes y comunidades con tierras dentro de los bloques "
     "(registro F-LA de liberación de áreas)",
     "Incertidumbre sobre el uso futuro del predio; restricción temporal del pastoreo durante la clausura",
     "Conservar derechos sobre el predio; recibir beneficios de la recuperación del ecosistema",
     "Liberación de áreas con identificación de prediantes y actas de libre disponibilidad (F-LA-01 a F-LA-06)",
     "Actas de disponibilidad en proceso: [POR DEFINIR] número suscrito", "Por confirmar"],
    ["Organizaciones comunales y de productores",
     "Comunidades campesinas, rondas campesinas, asociaciones de ganaderos y productores, JASS",
     "Degradación de pastos y fuentes de agua; conflictos por uso de tierras y agua",
     "Manejo de pastizales, protección de fuentes de agua, fortalecimiento organizativo",
     "Planes de manejo comunitario; capacitación en GdR; mecanismos MERESE",
     "[POR DEFINIR] (se consolida con F-DS-02 y F-DS-05)", "Por confirmar"],
    ["Gobiernos locales",
     "Municipalidades provinciales de Ayabaca, Huancabamba y Morropón y municipalidades de los 15 distritos",
     "Daños recurrentes en vías y servicios por huaicos y erosión; limitada capacidad de gestión del riesgo",
     "Reducir daños y gastos de rehabilitación; articular el proyecto con planes locales",
     "Coordinación institucional; participación en operación y mantenimiento y en alerta temprana comunitaria",
     "[POR DEFINIR] convenios o actas de compromiso de O&M", "Por confirmar"],
    ["Gobierno Regional de Piura",
     "Órganos competentes en agricultura, recursos naturales, gestión ambiental y gestión del riesgo",
     "Degradación de ecosistemas de la cuenca alta y su efecto en la cuenca media y baja",
     "Recuperación de ecosistemas y articulación con la gestión integral de la cuenca",
     "Coordinación técnica; opiniones y autorizaciones en el ámbito de su competencia",
     "[POR DEFINIR]", "Por confirmar"],
    ["Entidades del Gobierno Nacional",
     "MINAM (rector de la infraestructura natural; indicador de brecha R.M. N.° 00213-2024-MINAM), SERFOR, ANA y sus "
     "órganos desconcentrados, SERNANP, Ministerio de Cultura, INGEMMET, SENAMHI, CENEPRED, INDECI",
     "Brecha de ecosistemas degradados que requieren recuperación; riesgo de desastres en la cuenca",
     "Cumplimiento de lineamientos sectoriales; información para la gestión del riesgo y del agua",
     "Uso de información oficial; coordinación de opiniones técnicas cuando corresponda",
     "No requiere acuerdos en esta fase", "Neutral / normativo"],
    ["Usuarios de agua y población de la cuenca media y baja",
     "Juntas y comisiones de usuarios de agua de riego; población del valle del río Piura expuesta a inundaciones",
     "Sedimentos e incremento de escorrentía desde la cuenca alta; variabilidad de la disponibilidad hídrica",
     "Regulación hídrica y menor aporte de sedimentos y escorrentía",
     "Difusión de resultados; exploración de mecanismos MERESE", "[POR DEFINIR]", "A favor (beneficiario indirecto)"],
    ["Programa de Inversión del Río Piura",
     "Dependencia solicitante de la información", "Necesidad de integrar el alcance de la infraestructura natural de "
     "la cuenca alta en el programa de inversión",
     "Articular el Perfil IN Piura con el programa; evitar duplicidad de intervenciones",
     "Intercambio de información técnica y geoespacial (este informe); coordinación de alcances",
     "Remisión de la información de avance (este informe)", "A favor"],
    ["Entidad formuladora",
     "ANIN — DIME — SESDI (Unidad Formuladora del Perfil, CUI 2669244)",
     "Brecha de información de campo para el diagnóstico; plazos de formulación",
     "Formular un perfil viable y articulado; asegurar la sostenibilidad",
     "Trabajo de campo (F-DT, F-EE, F-DS, F-LA); coordinación interinstitucional",
     "Ejecución por Administración Directa", "A favor"],
]

NOTAS = [
    "Fuentes: catálogo V6 (datos/unidades_liberacion_areas.csv), capas datos/gis/IN_Piura_area_influencia_bloques_v6.gpkg "
    "y mapas/salidas/IN_Piura_ecosistemas_bloques_v6.gpkg, centros_poblados.py, CentrosPoblados-INEI-Bloques_v5 "
    "(24.09.2026), traspaso MSAVI V6 y caracterización de cárcavas.",
    "Bloques vigentes: 117 bloques V6 con 12,270.235 ha. Se excluyen los 12 códigos retirados del catálogo "
    "(1, 7, 25, 29, 32, 33, 46, 48, 68, 74, 75, M18B5) y sus registros asociados.",
    "El área de intervención directa del proyecto es un rango de 10,000 a 15,000 ha; 12,270.235 ha es la suma de los "
    "117 bloques preliminares.",
    "Sistema de coordenadas: UTM WGS 84 Zona 17S (EPSG:32717); controles de rango ESTE 450,000–750,000 m y NORTE "
    "9,300,000–9,600,000 m.",
    "Los totales y resúmenes de este libro son fórmulas (SUM, COUNTIFS, SUMIFS) sobre las hojas de detalle.",
    "[POR DEFINIR] Criterio técnico de delimitación del área de influencia aprobada (AI_aprobado_2.shp).",
    "[POR DEFINIR] Resultados del modelamiento de peligro (PMM, EPH, PGI y PI) por microcuenca y bloque, y sus capas.",
    "[POR DEFINIR] Capa de límites de las 26 microcuencas priorizadas.",
    "[POR DEFINIR] Centros poblados de los bloques {sin_cp}; coordenadas y desagregación de los "
    "centros poblados de los bloques 83 a 87; año censal de la capa INEI.",
    "[POR DEFINIR] Consolidado de las fichas F-DT, F-EE, F-DS-02 a F-DS-07 y F-LA registradas en el aplicativo IN Piura "
    "(base de datos del aplicativo, no incluida en este corte).",
    "[POR DEFINIR] Matriz de involucrados definitiva y acuerdos suscritos (la de la hoja A5 es preliminar).",
]


# ══════════════════════════════════════════════════════════════════════════
# 6. PRINCIPAL
# ══════════════════════════════════════════════════════════════════════════
def main():
    SALIDA.mkdir(parents=True, exist_ok=True)
    cat = cargar_catalogo()
    dist_of, prov_of = tablas_nombres(cat)
    b, ai, amb, ctx, eco_amb, eco_bl = cargar_capas(cat)
    msavi, verif = cargar_msavi(cat)
    cp, cp_por_bloque, sin_cp, avance = centros_poblados(cat, dist_of, prov_of)
    tb = tabla_bloques(cat, b, ai, eco_bl, msavi, cp_por_bloque)
    carc = cargar_carcavas(cat, dist_of, prov_of)

    capas = capas_sig(cat, b, ai, amb, ctx, eco_amb, eco_bl, tb, cp, carc)
    cpg = capas[3][2]
    mapa_pdf = SALIDA / "IN_Piura_Mapa_Centros_Poblados_V6.pdf"
    mapa_png = SALIDA / "IN_Piura_Mapa_Centros_Poblados_V6.png"
    mapa_ccpp(capas[1][2], capas[2][2], amb, ctx, cpg, cp, mapa_pdf, mapa_png)
    sin_coord = ", ".join(f"{r['Centro poblado']} ({r['Bloques asociados']})"
                          for _, r in cp[cp["ESTE (m)"].isna()].iterrows())
    inv = paquete_sig(capas, sin_coord, SALIDA / "IN_Piura_Paquete_SIG_Avance_Perfil.zip",
                      MAPAS_EXISTENTES + [mapa_pdf])

    fuentes = {
        "01_ambito_estudio_distritos": ("Límites distritales INEI (capa simplificada); área referencial", "Entregado"),
        "02_bloques_intervencion_v6": (FUENTE_BLOQUES, "Entregado"),
        "03_area_influencia_v6": (FUENTE_AI, "Entregado"),
        "04_centros_poblados_asociados": (FUENTE_CP, f"Entregado ({len(cp) - len(cpg)} CCPP sin coordenadas, solo en A3)"),
        "05_ecosistemas_ambito": (FUENTE_ECO, "Entregado"),
        "06_ecosistemas_bloques_v6": (FUENTE_ECO + "; intersección con bloques V6", "Entregado"),
        "07_carcavas_caracterizadas": ("Caracterización de cárcavas (gabinete SIG), ANIN-DIME-SESDI 2026", "Entregado"),
        "08_limites_distritales_contexto": ("INEI (capa simplificada), base de contexto", "Entregado"),
    }
    inventario = []
    for i, (nombre, titulo, g) in enumerate(inv, start=1):
        geom = {"Point": "Punto", "Polygon": "Polígono", "MultiPolygon": "Polígono"}[g.geom_type.iloc[0]]
        area = None if geom == "Punto" else round(float(g.AREA_HA.sum()) if "AREA_HA" in g else g.area.sum() / 1e4, 3)
        if nombre == "08_limites_distritales_contexto":
            area = None
        fmts = "SHP, GPKG" + ("" if nombre.startswith("08") else ", KMZ")
        inventario.append([i, nombre, titulo, geom, len(g), area, fmts, *fuentes[nombre]])
    n = len(inventario)
    inventario += [
        [n + 1, "Mapas temáticos (PDF A3)", "Área de influencia y bloques V6 (4 láminas); ecosistemas y bloques V6 "
         "(4 láminas); centros poblados asociados (1 lámina)", "—", 9, None, "PDF", "ANIN-DIME-SESDI 2026", "Entregado"],
        [n + 2, "Peligro integrado y componentes", "PMM, EPH, PGI y PI (5 niveles)", "Ráster / polígono", None, None,
         "[POR DEFINIR]", "Modelamiento MCA-AHP (mesolocalización)", "Pendiente de remisión"],
        [n + 3, "Microcuencas priorizadas", "Límites de las 26 microcuencas con bloques V6", "Polígono", None, None,
         "[POR DEFINIR]", "Delimitación hidrográfica del proyecto", "Pendiente de remisión"],
        [n + 4, "Capas de exclusión e idoneidad", "Arqueología, concesiones mineras, infraestructura eléctrica, ANP, "
         "accesibilidad, pendiente, déficit hídrico", "Varios", None, None, "[POR DEFINIR]",
         "SIGDA, INGEMMET, OSINERGMIN, SERNANP, MTC, CMIP6", "Pendiente de remisión"],
        [n + 5, "Ráster MSAVI 2024 clasificado", "Clases DN 1-5 (resumen por bloque en A4_UP_MSAVI)", "Ráster", None,
         None, "[POR DEFINIR]", "Sentinel-2 (2024)", "Pendiente de remisión"],
        [n + 6, "Registros de campo georreferenciados", "Fichas F-DT, F-EE, F-DS y F-LA con coordenadas y fotos",
         "Punto", None, None, "[POR DEFINIR] (exportación del aplicativo)", "Aplicativo IN Piura / KoboToolbox",
         "Pendiente de remisión"],
        [n + 7, "Vuelos de dron", "Ortomosaicos de verificación aérea por bloque", "Ráster", None, None,
         "[POR DEFINIR]", "Verificación aérea con dron (MSAVI V6 rev. 1)", "Pendiente de remisión"],
    ]

    ruta_xlsx = SALIDA / "IN_Piura_Anexos_Avance_Perfil.xlsx"
    notas = [n.replace("{sin_cp}", ", ".join(sin_cp)) for n in NOTAS]
    excel_anexos(ruta_xlsx, cat, amb, tb, cp, avance, eco_bl, eco_amb, msavi, verif, carc, sin_cp, inventario,
                 INVOLUCRADOS, CRITERIOS, notas)

    # ── cifras para el informe (las mismas que las fórmulas del Excel) ──
    rd = (tb.groupby(["Provincia", "Distrito"])
          .agg(n=("Bloque", "size"), ha=("Área bloque (ha)", "sum"), ai=("Área AI (ha)", "sum"),
               ai_n=("Polígonos AI", "sum")).reset_index())
    rd["_p"] = rd.Provincia.map(PROVINCIAS.index)
    rd = rd.sort_values(["_p", "Distrito"])
    amb_ha = {clave(r.NOMBDIST): float(r.AREA_HA) for r in amb.itertuples()}
    rp = cp.groupby("Provincia").agg(n=("Centro poblado", "size"), pob=("Población total", "sum"),
                                    h=("Hombres", "sum"), m=("Mujeres", "sum"), vul=("Población vulnerable", "sum"),
                                    viv=("Viviendas particulares", "sum")).reindex(PROVINCIAS)
    rdp = cp.groupby(["Provincia", "Distrito"]).agg(n=("Centro poblado", "size"), pob=("Población total", "sum"),
                                                   viv=("Viviendas particulares", "sum")).reset_index()
    rdp["_p"] = rdp.Provincia.map(PROVINCIAS.index)
    rdp = rdp.sort_values(["_p", "Distrito"])
    av = pd.DataFrame(avance)
    eco = (eco_bl.groupby("ECOSISTEMA").agg(ha=("HA_OFICIAL", "sum"), n=("BLOQUE", "nunique"))
           .sort_values("ha", ascending=False))
    eco_amb_ha = eco_amb.groupby("ECOSISTEMA").AREA_HA.sum()
    mic = tb.groupby("Microcuenca").agg(n=("Bloque", "size"), ha=("Área bloque (ha)", "sum")).sort_values("ha", ascending=False)
    dn = {k: sum(m[f"{k}_ha"] for m in msavi.values()) for k in ["dn1", "dn2", "dn3", "dn4", "dn5"]}
    cifras = {
        "fecha": HOY.isoformat(),
        "bloques": {"n": len(tb), "ha": round(float(tb["Área bloque (ha)"].sum()), 3),
                    "ai_ha": round(float(tb["Área AI (ha)"].sum()), 3), "ai_n": int(tb["Polígonos AI"].sum()),
                    "n_microcuencas": int(tb.Microcuenca.nunique()), "n_distritos": int(tb.Distrito.nunique())},
        "ambito_ha": round(sum(amb_ha.values()), 1),
        "por_distrito": [{"prov": r.Provincia, "dist": r.Distrito, "dist_ha": amb_ha[clave(r.Distrito)], "n": int(r.n),
                          "ha": round(float(r.ha), 3), "ai": round(float(r.ai), 3)} for r in rd.itertuples()],
        "por_provincia": [{"prov": p, "n": int(s.n.sum()), "ha": round(float(s.ha.sum()), 3),
                           "ai": round(float(s.ai.sum()), 3), "nd": len(s)} for p, s in rd.groupby("Provincia", sort=False)],
        "microcuencas": [{"cod": k, "n": int(r.n), "ha": round(float(r.ha), 3)} for k, r in mic.iterrows()],
        "ccpp": {"n": len(cp), "pob": int(cp["Población total"].sum()), "h": int(cp.Hombres.sum()),
                 "m": int(cp.Mujeres.sum()), "vul": int(cp["Población vulnerable"].sum()),
                 "viv": int(cp["Viviendas particulares"].sum()), "con_coord": int(cp["ESTE (m)"].notna().sum()),
                 "multi_bloque": int((cp["N.° bloques"] > 1).sum()), "rurales": int((cp.Tipo == "Rural").sum()),
                 "bloques_sin_cp": sin_cp, "sin_coord": sin_coord,
                 "pob_sin_desagregar": int(cp.loc[cp.Hombres.isna(), "Población total"].sum()),
                 "por_provincia": [{"prov": p, "n": int(r.n), "pob": int(r.pob), "h": int(r.h), "m": int(r.m),
                                    "vul": int(r.vul), "viv": int(r.viv)} for p, r in rp.iterrows()],
                 "por_distrito": [{"prov": r.Provincia, "dist": r.Distrito, "n": int(r.n), "pob": int(r.pob),
                                   "viv": int(r.viv)} for r in rdp.itertuples()]},
        "social": {"registros": len(av), "ds01": int(av["F-DS-01 reportadas"].fillna(0).sum()),
                   "ds01_completo": int((av["Estado F-DS-01"] == "COMPLETO").sum()),
                   "ds03": int(av["F-DS-03 reportadas"].fillna(0).sum()),
                   "ds03_completo": int((av["Estado F-DS-03"] == "COMPLETO").sum()),
                   "saberes_completo": int((av["Estado saberes"] == "COMPLETO").sum()),
                   "descartados": int(av["Estado F-DS-01"].isin(["-", "SIN REGISTRO"]).sum()),
                   "fecha_ult": str(max(d for d in av["Fecha de evaluación"] if isinstance(d, date))),
                   "N": 8865.9, "n": 368},
        "ecosistemas": [{"eco": e, "ha": round(float(r.ha), 3), "n": int(r.n),
                         "pct": round(float(r.ha) / AREA_TOTAL_HA, 4),
                         "amb_ha": round(float(eco_amb_ha.get(e, 0)), 3)} for e, r in eco.iterrows()],
        "eco_ambito": {"total_ha": round(float(eco_amb.AREA_HA.sum()), 3),
                       "otros": sorted(set(eco_amb.ECOSISTEMA) - set(eco_bl.ECOSISTEMA))},
        "msavi": {"media_min": round(min(m["media_msavi"] for m in msavi.values()), 4),
                  "media_max": round(max(m["media_msavi"] for m in msavi.values()), 4),
                  "bajo_umbral_ha": round(sum(m["bajo_umbral_ha"] for m in msavi.values()), 3),
                  "total_ha": round(sum(dn.values()), 3), "dn": {k: round(v, 3) for k, v in dn.items()},
                  "bloques_mayoria_bajo": int(sum(1 for m in msavi.values() if m["bajo_umbral_pct"] > 50)),
                  "r": verif["r_directo"], "eam": verif["eam_directo"]},
        "carcavas": {"n": len(carc), "long_m": round(float(carc.Long.sum()), 2), "bloques": int(carc.Bloque.nunique()),
                     "por_ndvi": carc.Cl_NDVI.value_counts().to_dict()},
        "inventario": [[x if not (isinstance(x, float) and pd.isna(x)) else None for x in r] for r in inventario],
        "criterios": CRITERIOS, "involucrados": INVOLUCRADOS,
    }
    cifras["msavi"]["pct_bajo"] = round(cifras["msavi"]["bajo_umbral_ha"] / cifras["msavi"]["total_ha"], 4)
    (SALIDA / "cifras_informe.json").write_text(json.dumps(cifras, ensure_ascii=False, indent=1, default=str),
                                                 encoding="utf-8")
    print(json.dumps({k: cifras[k] for k in ["bloques", "ambito_ha"]}, ensure_ascii=False))
    print("CCPP:", {k: v for k, v in cifras["ccpp"].items() if not isinstance(v, list)})
    print("Social:", cifras["social"])
    print("MSAVI:", cifras["msavi"])
    print("Cárcavas:", cifras["carcavas"])
    print("Salidas en", SALIDA)


if __name__ == "__main__":
    main()
