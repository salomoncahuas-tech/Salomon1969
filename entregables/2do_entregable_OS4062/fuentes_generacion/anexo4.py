# Anexo 4 — Matrices de seguimiento del Segundo Entregable (información nativa, formato ANIN)
import sys, json, pandas as pd, openpyxl
sys.path.insert(0, '/home/user/Salomon1969/.claude/skills/anin-in-piura/scripts')
from anin_excel import encabezado_institucional, escribir_tabla
from openpyxl.styles import Font, Alignment
S='/tmp/claude-0/-home-user-Salomon1969/03b82ed8-c0e9-54ec-93c9-bff5f93b575b/scratchpad'
D=json.load(open(S+'/build/datos.json')); O=json.load(open(S+'/build/obs.json'))
pred=pd.read_csv(S+'/territorial/predial_v6.csv'); pred=pred[pred.prov.isin(['Huancabamba','Morropón'])]
wb=openpyxl.Workbook()
def hoja(nombre, titulo, cols, filas, fuente, **kw):
    ws=wb.create_sheet(nombre)
    f=encabezado_institucional(ws, titulo, len(cols), fuente=fuente)
    fin=escribir_tabla(ws, f, cols, filas, **kw)
    for row in ws.iter_rows(min_row=f+1, max_row=fin):
        for c in row: c.alignment=Alignment(vertical='top', wrap_text=True)
    return ws
# Índice
ws=wb.active; ws.title='Indice'
f=encabezado_institucional(ws,'ANEXO N° 4 — MATRICES DE SEGUIMIENTO DEL SEGUNDO ENTREGABLE (OS N° 4062-2026 · Informe N° 008-2026-HSCM)',3,
  fuente='Elaboración: Ing. Héctor Salomón Cahuas Miller – Especialista en Infraestructura Verde. Lima, octubre de 2026')
hojas=[('Bloques_107','Actividad b)','Consolidado por bloque de Huancabamba y Morropón (Versión 5)'),
 ('Distritos','Actividad b)','Consolidado por distrito, con totales por fórmula'),
 ('Situacion_predial','Actividad b)','Situación predial de gabinete por bloque (mapas SP v2)'),
 ('Obs_Suelos_E2','Actividad c)','Matriz de observaciones al Entregable N° 2 del Estudio de Suelos'),
 ('Obs_E6_Ecosistemas','Actividad d)','Observaciones al Diagnóstico del Estado de los Ecosistemas'),
 ('Obs_E6_Material_vegetal','Actividad d)','Observaciones al Estudio de Oferta de Material Vegetal'),
 ('Medios_verificadores','b), c), d)','Registro de medios verificadores'),('Notas','—','Reglas de dato y advertencias')]
escribir_tabla(ws,f,['Hoja','Actividad del TdR','Contenido'],hojas,anchos={'Hoja':26,'Actividad del TdR':18,'Contenido':80})
# Bloques
B=sorted(D['bloques'], key=lambda r:(r['prov'],r['d'],r['b']))
hoja('Bloques_107','CONSOLIDADO DE LOS 107 BLOQUES — HUANCABAMBA Y MORROPÓN (VERSIÓN 5)',
 ['Provincia','Distrito','Bloque','Microcuenca','Superficie (ha)','Este (m)','Norte (m)','Régimen','MSAVI 2024','Clase PI','Índice PI','IC (E6)','Clase de estado (E6)','Tipología MSAVI–NDVI','% bajo umbral MSAVI','Situación predial','Predios COFOPRI','Calicata E5','Ficha F-DS','Riesgo (gabinete)','Lámina SP','N° cárcavas'],
 [[r['prov'],r['d'],r['b'],r['mc'],r['ha'],r['e'],r['n'],r['reg'],r['ms'],r['pi'],r['ipp'],r['ic'],r['est'],r['tip'],r['bu'],r['sit'],r['npred'],r['cal'],r['ds'],r['rk'],r['sp'],r['ncv']] for r in B],
 'Volúmenes I y II, Versión 5; Tablero E6; índice de mapas de situación predial v2; Informe de Avance E5; consolidado DS 17/10/2026. UTM WGS 84 Zona 17S.',
 totales={'Superficie (ha)':'SUM','Predios COFOPRI':'SUM','N° cárcavas':'SUM','Bloque':'COUNTA'},
 formato_num={'Superficie (ha)':'#,##0.00','Este (m)':'#,##0','Norte (m)':'#,##0','MSAVI 2024':'0.000','Índice PI':'0.000','IC (E6)':'0.000','% bajo umbral MSAVI':'0.0'},
 anchos={'Distrito':22,'Situación predial':30,'Clase de estado (E6)':16})
# Distritos
hoja('Distritos','CONSOLIDADO POR DISTRITO — HUANCABAMBA Y MORROPÓN',
 ['Provincia','Distrito','Bloques','Superficie (ha)','Bloques bajo umbral (media)','Superficie bajo umbral E6 (ha)','PI Alto + Muy alto (ha)','IC ponderado','Bloques Moderado/Severo','Calicatas de bloque','C. C. (ha)','COFOPRI (ha)','Sin registro (ha)','Predios COFOPRI'],
 [[r['prov'],r['d'],r['n'],r['ha'],r['bajo_n'],r['e6_bajo'],r['pi45'],r['ic'],r['modsev'],r['calB'],r['cc'],r['cof'],r['est'],r['npred']] for r in D['dist']],
 'Volúmenes I y II, Versión 5; Tablero E6; índice de mapas de situación predial v2.',
 totales={k:'SUM' for k in ['Bloques','Superficie (ha)','Bloques bajo umbral (media)','Superficie bajo umbral E6 (ha)','PI Alto + Muy alto (ha)','Bloques Moderado/Severo','Calicatas de bloque','C. C. (ha)','COFOPRI (ha)','Sin registro (ha)','Predios COFOPRI']},
 formato_num={k:'#,##0.00' for k in ['Superficie (ha)','Superficie bajo umbral E6 (ha)','PI Alto + Muy alto (ha)','C. C. (ha)','COFOPRI (ha)','Sin registro (ha)']}|{'IC ponderado':'0.000'},
 anchos={'Distrito':24})
# Predial
hoja('Situacion_predial','SITUACIÓN PREDIAL DE GABINETE POR BLOQUE — HUANCABAMBA Y MORROPÓN (MAPAS SP v2)',
 ['Lámina','Bloque','Provincia','Distrito','Asistente','Superficie (ha)','C. C. (ha)','Predios COFOPRI fuera de C. C. (ha)','Estado y/o posesionarios (ha)','Situación predial','N° predios COFOPRI','N° CCPP asociados','Población CCPP (hab.)','Excede 19 predios (Paso 5)'],
 [[r.lam,str(r.bloque),r.prov,r.dist,r.asist,r.sup,r.cc,r.priv,r.est,r.sit,int(r.npred),int(r.nccpp),(None if pd.isna(r.pob) else int(r.pob)),'Sí' if r.npred>19 else 'No'] for r in pred.itertuples()],
 'Índice de mapas de situación predial de los 117 bloques V6 v2 (Drive ID 1jkNZ0sikCKQXZyUu4jlaiekNU82eQXE2), 11/10/2026. Tolerancia cartográfica 5 %.',
 totales={k:'SUM' for k in ['Superficie (ha)','C. C. (ha)','Predios COFOPRI fuera de C. C. (ha)','Estado y/o posesionarios (ha)','N° predios COFOPRI','N° CCPP asociados']},
 formato_num={k:'#,##0.000' for k in ['Superficie (ha)','C. C. (ha)','Predios COFOPRI fuera de C. C. (ha)','Estado y/o posesionarios (ha)']},
 anchos={'Distrito':24,'Situación predial':32})
# Observaciones
hoja('Obs_Suelos_E2','MATRIZ DE OBSERVACIONES — ENTREGABLE N° 2 DEL ESTUDIO DE SUELOS (GeoSIG Ingenieros)',
 ['Código','Tipo','Sección','Observación','Levantamiento requerido','Se levanta en','Estado'],
 [o+['Formulada – pendiente de levantamiento'] for o in O['OS']],
 'Revisión técnica del Especialista en Infraestructura Verde, en coordinación con el Especialista en Infraestructura Marrón. Tipo: C consistencia; A alcance; F forma.',
 anchos={'Sección':20,'Observación':60,'Levantamiento requerido':55,'Estado':24})
for nom,tit,key in [('Obs_E6_Ecosistemas','MATRIZ DE OBSERVACIONES — DIAGNÓSTICO DEL ESTADO DE LOS ECOSISTEMAS (E6)','OE'),('Obs_E6_Material_vegetal','MATRIZ DE OBSERVACIONES — ESTUDIO DE OFERTA DE MATERIAL VEGETAL (E6)','OM')]:
    hoja(nom,tit,['Código','Documento / sección','Observación','Levantamiento requerido','Estado al corte'],O[key],
     'Revisión técnica del Especialista en Infraestructura Verde, en coordinación con el Especialista en Intervenciones Vegetativas de Uso Sostenible.',
     anchos={'Documento / sección':26,'Observación':60,'Levantamiento requerido':50,'Estado al corte':34})
MV=[
 ['b)','Diagnóstico Huancabamba Versión 5','Artefacto','https://claude.ai/artifact/9828ogNDewfGQLgQDCq7kn','17/10/2026'],
 ['b)','Diagnóstico Morropón Versión 5','Artefacto','https://claude.ai/artifact/Y5QRXCQRsXFvwFgYL1oZ64','17/10/2026'],
 ['b)','Diagnóstico Ayabaca Versión 5 (contraste)','Artefacto','https://claude.ai/artifact/4kF1tRzj2FkuDgbcuTcBye','17/10/2026'],
 ['b)','Índice de mapas de situación predial V6 v2','Google Drive','1jkNZ0sikCKQXZyUu4jlaiekNU82eQXE2','11/10/2026'],
 ['b)','Lámina SP-042 (M4B4)','Google Drive','1g2ax3ZxRLfZvhH-OEcPyh97se6gHGIj-','11/10/2026'],
 ['b)','Lámina SP-091 (bloque 87)','Google Drive','1eVz0dnA5gcYFZDWue7DQS6tWtLlxhyvJ','11/10/2026'],
 ['b)','Informe E2 Vol. I Morropón (B1)','Google Drive','1AnXV3tTvQphvFLGGW4AVWwhRGRp0aJbv','10/10/2026'],
 ['b)','Informe E2 Vol. II Huancabamba (B1)','Google Drive','1DY7r3RBZy5lYEX6DkTakAaO3Pxu9oKt2','10/10/2026'],
 ['b)','Matrices E2 (B1)','Google Drive','11-tKOKJYgPF-uGGF7nRkWuxR5kjfJ6Xg','10/10/2026'],
 ['b)','Mapas y Excel de bloques V6 (ecosistemas, área de influencia)','Repositorio IN Piura','mapas/salidas; commits 0163a3a, 632ee02, 9a9633c, 5af14b9, 2bd1bae, 77a0266','04–05/10/2026'],
 ['b)','Módulo de liberación de áreas y catálogo V6','Repositorio IN Piura','commits 042679b, 9a4d846, b3d9dcf, db4d49e, 33ae71e, ae1f754','27/09–06/10/2026'],
 ['b)','Diagnóstico social: F-DS-01 y tenencia por bloque','Repositorio IN Piura','commits bfae1c9, d71f8ec, 39a9baf, 656d083, e50264a, ed3d260, 2b76973','05–08/10/2026'],
 ['c)','Informe de Avance – Fase de Campo (PDF)','Google Drive','1iWNiXKXVTf5XYVXS0RHFaJBZYSO5iqnE','27/09/2026'],
 ['c)','Informe de Avance – Fase de Campo (Word)','Google Drive','1B5Z7G-scRsjtE7hpchibTXC9_qWg2x9s','27/09/2026'],
 ['c)','Carpeta «Levantamiento de suelos_ANIN» (consultora)','Google Drive','1COyqfGePNgQUa9ioCOFDW_sM5nNu7mRf','10/09/2026'],
 ['c)','TdR del Estudio de Suelos','Google Drive','1VFZfzJdX0_Ow_gfsMoLvKfwIvkoUe8Hr','27/09/2026'],
 ['c), d)','Tablero de avance del Perfil al 03/10/2026','Google Drive','1wRdIvYH8ni8kbj3XsDKa0cgrgSIdVWn6','03/10/2026'],
 ['c), d)','Plan de Trabajo Detallado, Actualización N° 01, Versión 2 (documento rector del E6 por firma espectral)','Google Drive','1QkhMG4cWkPFAyhxOH8qtaWG2YVCCJsGy','16–17/10/2026'],
 ['c), d)','Matrices del Plan de Trabajo (presupuesto ajustado; hoja «E6_Firma_Espectral»)','Google Drive','1EF3u3btVKiK2DcUKrHYcJZv-4co_arCc','16/10/2026'],
 ['d)','E6 v2.0 «In house» (original; antecedente)','Google Drive','1B3QsYMkodQVnK5ZwQU-7BI0t9Yn6xmoR','08/10/2026'],
 ['d)','E6 v2.0 «In house» (revisión hscm, Word)','Google Drive','1rxdf63E-TOJpUEUWWazdkth6WjdwU4HY','08/10/2026'],
 ['d)','E6 v2.0 «In house» (revisión hscm, PDF)','Google Drive','1lbr4-I5-OH8Nqd4hBccIEizEy3H-rV7V','08/10/2026'],
 ['d)','Nota Técnica N° 002-2026 (original)','Google Drive','1EZk6-sdLsbFfkiuI4jma4IxlTVsTSXyT','13/10/2026'],
 ['d)','Nota Técnica N° 002-2026 (revisada, Word / PDF)','Google Drive','137Syk2NZPG9bTZ8rwrxtpTmjbzOtLhTy / 17Vt9Z1NBZPTGF933r4jhqFNu0N7MP-nR','14/10/2026'],
 ['d)','Clasificación no supervisada de tomas S2DR4 (resultados y LEEME)','Google Drive','1oZ-PLyuw4QyIBe-fupK3j7lQY_DE48J8; 1NRGviLNkpnvfyrLfv3dVNZms4w3SbfuS','13–14/10/2026'],
 ['d)','Informe E6 – versión para revisión de SESDI','Google Drive','1NfYjSV3QXI3uA4eHsMNQNTw9vdY5RdDA','15/10/2026'],
 ['d)','Matrices del diagnóstico del ecosistema (E6)','Google Drive','1Uie1il6FeHI4SD8QhKXsfDof7cm4NF8J','15/10/2026'],
 ['d)','Fuentes semilleras de los Bosques Secos (referencia)','Google Drive','1hS8QlbTu_Ks05061BV4FAbAqbfUpikzy','—'],
]
hoja('Medios_verificadores','REGISTRO DE MEDIOS VERIFICADORES DEL SEGUNDO ENTREGABLE',['Actividad','Medio verificador','Repositorio','Identificador / enlace','Fecha'],MV,
 'Elaboración propia.',anchos={'Medio verificador':48,'Identificador / enlace':60,'Repositorio':20})
ws=wb.create_sheet('Notas'); f=encabezado_institucional(ws,'NOTAS Y ADVERTENCIAS',1)
for k,t in enumerate([
 'Regla de dato: ningún valor ha sido estimado, interpolado ni transferido entre bloques; los vacíos se consignan como celda vacía, «s/d» o «Pendiente».',
 'Superficies en hectáreas calculadas en UTM WGS 84 Zona 17S (EPSG:32717); centroides validados (Este 450,000–750,000 m; Norte 9,300,000–9,600,000 m).',
 'Brecha: coexisten el criterio de media del bloque (MSAVI < 0.4976), el criterio areal del E6 y la media del MSAVI por ecosistema que propone el Plan de Trabajo institucional v2; no son sumables entre sí.',
 'E6: se desarrolla con la metodología de firma espectral por el equipo base, conforme al Plan de Trabajo Detallado – Actualización N° 01, Versión 2 (16/10/2026); las observaciones OE-03 y OE-14 a OE-16 alinean el Informe E6 y las matrices con dicho plan.',
 'Situación predial de gabinete: superficie sin registro en GEORURAL y COFOPRI se asume Tierras del Estado y/o posesionarios; requiere verificación de campo y actas.',
 'Los totales de cada hoja se calculan con fórmulas (=SUM, =COUNTA) sobre las filas de la tabla.',
 'Las observaciones al Estudio de Suelos y al E6 se levantan en los entregables indicados; su estado corresponde al corte del Segundo Entregable.'],start=f):
    c=ws.cell(k,1,t); c.font=Font(name='Arial',size=10); c.alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['A'].width=120
out=sys.argv[1]; wb.save(out); print('OK',out)
