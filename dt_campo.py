"""Integracion de las 117 fichas de campo del Diagnostico Territorial
(F-DT-01 a F-DT-05) con las 117 Fichas de Resumen de Bloques V6.

Proyecto IN Piura | CUI 2669244 | ANIN - DIME - SESDI.

El aplicativo maneja dos cuerpos documentales por bloque preliminar de
intervencion:

* la **plantilla DT de campo** (`Plantilla_DT_Campo_Check_Validada_V5`), que
  el equipo llena en la verificacion in situ y que trae las cinco fichas
  F-DT-01 a F-DT-05;
* la **ficha de resumen V6** (`Plantilla_Excel_Bloque_<codigo>_IN_Piura`),
  armada en gabinete a partir del catalogo maestro de bloques, la
  estadistica zonal sobre el MDE y los compuestos Sentinel-2 (NDVI 2025 y
  distribucion areal del MSAVI 2024 por clase DN).

Ambos cuerpos describen en parte los mismos hechos. Este modulo los cruza
**sin duplicar informacion**: cada hecho territorial se declara una sola
vez, con la fuente que manda sobre el (campo, gabinete o fuente oficial),
y el valor de la otra fuente solo se conserva cuando difiere, como
trazabilidad de la discrepancia. De ese cruce salen:

1. el registro integrado del bloque, con procedencia campo por campo;
2. el control de consistencia regenerado (hoja 5 de la ficha de resumen);
3. la ficha de resumen actualizada (V7 = V6 + campo integrado);
4. los reportes Excel y PDF con graficos, por bloque y consolidados por
   provincia, distrito, microcuenca o bloque.

Sistema de referencia: UTM WGS 84 Zona 17S (EPSG:32717).
"""

import io
import json
import math
import os
import re
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import resumenes_bloques as rbq
from excel_diagnostico_territorial import parsear_excel_dt


# ══════════════════════════════════════════════════════════════════════════
# Fuentes de informacion
# ══════════════════════════════════════════════════════════════════════════

# Las tres procedencias que el estudio distingue. Todo campo integrado
# declara cual de ellas manda sobre el, de modo que un mismo hecho nunca se
# escribe dos veces con dos origenes distintos.
CAMPO = "CAMPO"
GABINETE = "GABINETE"
OFICIAL = "OFICIAL"

FUENTES = (CAMPO, GABINETE, OFICIAL)

ETIQUETA_FUENTE = {
    CAMPO: "Dato de campo",
    GABINETE: "Gabinete (GIS / teledetección)",
    OFICIAL: "Fuente oficial",
}

DESCRIPCION_FUENTE = {
    CAMPO: ("Verificación en campo. Fichas F-DT-01 a F-DT-05 de la Plantilla "
            "DT Campo Check Validada V5, suscritas por el evaluador."),
    GABINETE: ("Gabinete. Estadística zonal sobre el MDE y compuestos "
               "Sentinel-2 (NDVI mediana 2025, MSAVI 2024 por clase DN)."),
    OFICIAL: ("Fuente oficial. Catálogo maestro Bloques V5/V6, división "
              "política INEI y codificación de microcuencas de la ANA."),
}

# Estado del cruce de un campo entre las dos fuentes.
CONFORME = "CONFORME"          # ambas fuentes coinciden
COMPLEMENTADO = "COMPLEMENTADO"  # solo una fuente lo trae; se adopta
ACTUALIZADO = "ACTUALIZADO"    # campo corrige lo que traia el libro V6
DISCREPANTE = "DISCREPANTE"    # difieren y manda gabinete / fuente oficial
PENDIENTE = "PENDIENTE"        # ninguna fuente lo declara

ESTADOS = (CONFORME, COMPLEMENTADO, ACTUALIZADO, DISCREPANTE, PENDIENTE)

# Calificacion que cada estado recibe en el control de consistencia, con el
# vocabulario de la hoja 5 de las fichas de resumen.
CALIFICACION_POR_ESTADO = {
    CONFORME: "CONFORME",
    COMPLEMENTADO: "CORREGIDO",
    ACTUALIZADO: "CORREGIDO",
    DISCREPANTE: "NO SUSTANTIVA",
    PENDIENTE: "SUSTANTIVA",
}

# Carpeta del repositorio con las 117 plantillas DT de campo actualizadas.
# Viaja con el aplicativo, igual que la carpeta de fichas de resumen: asi la
# integracion masiva no depende de que alguien vuelva a subir los archivos.
CARPETA_CAMPO = "plantillas_dt_campo"

# Carpeta con las fichas de resumen ya integradas (V7). La produce
# `integrar_dt_campo.py` a partir de la carpeta V6 y de las plantillas de
# campo, y es la que el aplicativo carga como version vigente.
CARPETA_INTEGRADA = "plantillas_117_integradas_v7"

VERSION_INTEGRACION = "V7 - resumen V6 + campo F-DT integrado"

FICHAS_CAMPO = ("F-DT-01", "F-DT-02", "F-DT-03", "F-DT-04", "F-DT-05")


# ══════════════════════════════════════════════════════════════════════════
# Utilidades
# ══════════════════════════════════════════════════════════════════════════

def _txt(valor):
    """Texto limpio; '' para lo ausente."""
    if valor is None:
        return ""
    return re.sub(r"\s+", " ", str(valor)).strip()


def _num(valor):
    """Numero interpretable del texto.

    Delega en el lector de los libros de resumen: si las dos fuentes se
    leyeran con criterios numericos distintos, un mismo valor podria
    aparecer como discrepancia consigo mismo.
    """
    return rbq.numero(valor)


def _comparable(valor):
    """Forma canonica de un texto para decidir si dos fuentes coinciden."""
    texto = _txt(valor).lower()
    texto = (texto.replace("á", "a").replace("é", "e").replace("í", "i")
             .replace("ó", "o").replace("ú", "u").replace("ü", "u")
             .replace("ñ", "n"))
    return re.sub(r"[^a-z0-9]+", " ", texto).strip()


# Formulas con que los dos formatos declaran que un dato todavia no existe.
# Se comparan como prefijo porque suelen ir seguidas de la explicacion ("Por
# determinar en la microlocalizacion", "No disponible en los insumos; debe
# tomarse del modelamiento..."). «No aplica» NO figura aqui: es una
# declaracion con contenido, no una ausencia.
_MARCAS_SIN_DATO = (
    "por determinar", "por definir", "por confirmar", "por verificar",
    "no disponible", "no se dispone", "sin dato", "sin informacion",
    "pendiente de", "s d",
)


def _sin_dato(valor):
    """True si el valor es uno de los marcadores de ausencia del formato."""
    canon = _comparable(valor)
    if not canon:
        return True
    return canon.startswith(_MARCAS_SIN_DATO)


# ══════════════════════════════════════════════════════════════════════════
# Lectura de la plantilla DT de campo
# ══════════════════════════════════════════════════════════════════════════

# Los 117 archivos llegan de campo con nombres heterogeneos
# (`DT_B13_IN_Piura`, `FDT_M6B2-1_Diagnostico_Territorial_V5`,
# `Plantilla_DT_Campo_V5_BloqueM3B6`...). El codigo del bloque se recupera
# del nombre cuando la hoja F-DT-01 no lo declara.
_RE_MICRO_BLOQUE = re.compile(
    r"(?<![A-Za-z0-9])M\s*(\d{1,2})\s*[_\- ]?\s*B\s*(\d{1,2})"
    r"(?:\s*-\s*(\d))?(?![0-9])", re.IGNORECASE)
_RE_BLOQUE = re.compile(r"bloque\s*[_\- ]?\s*(\d{1,3})(?![0-9])", re.IGNORECASE)
_RE_B = re.compile(r"(?<![A-Za-z0-9])B\s*(\d{1,3})(?![0-9])", re.IGNORECASE)


def codigo_desde_nombre_campo(nombre_archivo):
    """Codigo del bloque deducido del nombre del archivo de campo."""
    base = re.sub(r"\.xlsx?$", "", os.path.basename(nombre_archivo or ""),
                  flags=re.IGNORECASE)
    # `BloqueM3B6` -> `Bloque_M3B6`, para que el codigo de microcuenca no
    # quede pegado a la palabra y pueda reconocerse.
    base = re.sub(r"(bloque)(?=[A-Za-z0-9])", r"\1_", base, flags=re.IGNORECASE)
    m = _RE_MICRO_BLOQUE.search(base)
    if m:
        sufijo = "-%s" % m.group(3) if m.group(3) else ""
        return "M%dB%d%s" % (int(m.group(1)), int(m.group(2)), sufijo)
    for regex in (_RE_BLOQUE, _RE_B):
        m = regex.search(base)
        if m:
            return str(int(m.group(1)))
    return ""


def normalizar_codigo(codigo):
    """Codigo canonico del bloque: '013' -> '13', 'm6b2-1' -> 'M6B2-1'."""
    texto = _txt(codigo).upper().replace(" ", "")
    if not texto:
        return ""
    if texto.isdigit():
        return str(int(texto))
    m = _RE_MICRO_BLOQUE.match(texto)
    if m:
        sufijo = "-%s" % m.group(3) if m.group(3) else ""
        return "M%dB%d%s" % (int(m.group(1)), int(m.group(2)), sufijo)
    return texto


# Tablas de las fichas y la clave con que se guardan en el registro.
_TABLAS_CAMPO = (
    ("dt02_carcavas", "carcavas"),
    ("dt03_floristica", "floristica"),
    ("dt03_especies_clave", "especies_clave"),
    ("dt04_causas", "causas"),
    ("dt04_indicadores", "indicadores"),
    ("dt05_fuentes_agua", "fuentes_agua"),
)


def _fila_con_datos(fila):
    """True si la fila de una tabla del formulario trae algo mas que el N.°."""
    for clave, valor in fila.items():
        if clave == "n":
            continue
        if _txt(valor):
            return True
    return False


def parsear_ficha_campo(archivo, nombre_archivo=""):
    """Lee una plantilla DT de campo y devuelve un registro unico del bloque.

    Funde las cinco fichas F-DT-01 a F-DT-05 en un solo diccionario: los
    datos de cabecera (evaluador, fecha, bloque, microcuenca) se repiten en
    las cinco hojas y aqui se consolidan una sola vez, quedandose con el
    primer valor no vacio. Nunca lanza por una hoja ausente: lo que falte se
    registra en `fichas_ausentes`.
    """
    if isinstance(archivo, (bytes, bytearray)):
        archivo = io.BytesIO(archivo)
    resultados = parsear_excel_dt(archivo)
    datos, fichas = {}, []
    for bloque in resultados:
        fichas.append(bloque["ficha"])
        for clave, valor in (bloque.get("datos") or {}).items():
            actual = datos.get(clave)
            vacio_actual = actual in (None, "", []) or (
                isinstance(actual, str) and not actual.strip())
            if vacio_actual:
                datos[clave] = valor

    registro = {
        "nombre_archivo": os.path.basename(nombre_archivo or ""),
        "fichas_leidas": fichas,
        "fichas_ausentes": [f for f in FICHAS_CAMPO if f not in fichas],
    }

    # Tablas: se descartan las filas en blanco que la plantilla deja
    # preformateadas, para no inflar los conteos con filas vacias.
    for clave_origen, clave_destino in _TABLAS_CAMPO:
        filas = [f for f in (datos.get(clave_origen) or []) if _fila_con_datos(f)]
        registro[clave_destino] = filas

    # Campos escalares: todo lo que no sea tabla entra tal cual.
    claves_tabla = {c for c, _ in _TABLAS_CAMPO}
    for clave, valor in datos.items():
        if clave in claves_tabla:
            continue
        registro[clave] = _txt(valor)

    codigo = normalizar_codigo(registro.get("codigo_bloque"))
    if not codigo:
        codigo = normalizar_codigo(codigo_desde_nombre_campo(nombre_archivo))
        if codigo:
            registro["codigo_desde_nombre"] = True
    registro["codigo_bloque"] = codigo

    # Conteos derivados: se calculan aqui una vez y no se vuelven a pedir.
    # Un conteo de cero sobre una tabla vacia NO es un cero declarado: puede
    # ser que el formulario traiga la tabla con otra maqueta y no se haya
    # podido leer. Se deja sin valor para que la integracion no sobrescriba
    # con un cero el dato que si trae la ficha de resumen.
    registro["n_carcavas_inventariadas"] = len(registro["carcavas"]) or ""
    registro["n_taxones"] = len(registro["floristica"]) or ""
    registro["n_especies_clave"] = len(registro["especies_clave"]) or ""
    registro["n_fuentes_agua"] = len(registro["fuentes_agua"]) or ""
    registro["causas_activas"] = _causas_activas(registro["causas"])
    # Aqui el cero si es un hecho: la matriz de F-DT-04 viene con sus 16
    # causas y ninguna quedo marcada como activa.
    registro["n_causas_activas"] = (len(registro["causas_activas"])
                                    if registro["causas"] else "")
    registro["n_indicadores"] = len(registro["indicadores"]) or ""
    # Tablas que el formulario declara siempre y que llegaron sin filas: su
    # ausencia se informa, no se interpreta como que no existan.
    registro["tablas_vacias"] = [
        etiqueta for clave, etiqueta in (
            ("floristica", "F-DT-03 · elenco florístico"),
            ("causas", "F-DT-04 · matriz de causas de degradación"))
        if not registro[clave]]
    registro["completitud_pct"] = _completitud(registro)
    return registro


# Intensidades de la matriz de causas de F-DT-04, ordenadas de menor a
# mayor. Una causa cuenta como activa cuando se declara presente o se le
# asigna una intensidad distinta de "Nula".
_ORDEN_INTENSIDAD = ["nula", "ligera", "moderada", "fuerte", "muy fuerte"]


def _causas_activas(causas):
    """Causas de degradacion declaradas presentes, de mayor a menor peso."""
    activas = []
    for fila in causas:
        presencia = _comparable(fila.get("presencia"))
        intensidad = _comparable(fila.get("intensidad"))
        if presencia in {"no", "n"} :
            continue
        if not presencia and intensidad in {"", "nula"}:
            continue
        if intensidad == "nula":
            continue
        try:
            peso = _ORDEN_INTENSIDAD.index(intensidad)
        except ValueError:
            peso = 1
        activas.append({
            "causa": _txt(fila.get("causa")),
            "intensidad": _txt(fila.get("intensidad")),
            "extension": _txt(fila.get("extension")),
            "antiguedad": _txt(fila.get("antiguedad")),
            "evidencia": _txt(fila.get("evidencia")),
            "peso": peso,
        })
    activas.sort(key=lambda c: (-c["peso"], c["causa"]))
    return activas


# Campos que se esperan llenos en una ficha de campo completa. Sirven para
# medir cuanto trae cada archivo y advertir sobre los que llegaron a medias.
_CLAVES_COMPLETITUD = (
    "fecha_evaluacion", "evaluador", "codigo_bloque", "microcuenca",
    "distrito", "utm_este_dt", "utm_norte_dt", "altitud_gps",
    "forma_terreno", "pendiente", "posicion_fisiografica",
    "exposicion_orientacion", "rango_altitudinal",
    "dt02_nivel_erosion_general", "dt02_urgencia_control",
    "dt03_tipo_ecosistema", "dt03_estado_conservacion_eco",
    "dt03_uso_dominante", "dt03_cobertura_total", "dt03_suelo_desnudo",
    "dt03_regeneracion_natural", "dt04_causa_subyacente",
    "dt04_velocidad_degradacion", "dt04_reversibilidad",
    "dt04_urgencia_intervencion", "dt05_zona_recarga",
    "dt05_modalidad_acceso", "dt05_senal_celular",
)


def _completitud(registro):
    """Porcentaje de campos esperados que la ficha de campo trae llenos."""
    llenos = sum(1 for c in _CLAVES_COMPLETITUD if not _sin_dato(registro.get(c)))
    return round(100.0 * llenos / len(_CLAVES_COMPLETITUD), 1)


def fichas_del_repositorio(carpeta=None):
    """[(nombre, contenido)] de las plantillas de campo incluidas en el repo."""
    ruta = carpeta or os.path.join(
        os.path.dirname(os.path.abspath(__file__)), CARPETA_CAMPO)
    if not os.path.isdir(ruta):
        return []
    archivos = []
    for nombre in sorted(os.listdir(ruta)):
        if not nombre.lower().endswith(".xlsx") or nombre.startswith("~$"):
            continue
        with open(os.path.join(ruta, nombre), "rb") as fh:
            archivos.append((nombre, fh.read()))
    return archivos


_MANIFIESTO_CAMPO = None


def cargar_manifiesto_campo():
    """Catalogo de las 117 plantillas DT de campo (codigo, archivo, Drive)."""
    global _MANIFIESTO_CAMPO
    if _MANIFIESTO_CAMPO is None:
        ruta = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                            "datos", "manifiesto_dt_campo_117.json")
        try:
            with open(ruta, encoding="utf-8") as fh:
                _MANIFIESTO_CAMPO = json.load(fh)
        except (OSError, ValueError):
            _MANIFIESTO_CAMPO = {"total_bloques": 0, "bloques": []}
    return _MANIFIESTO_CAMPO


def codigos_campo_esperados():
    return [b["codigo"] for b in cargar_manifiesto_campo().get("bloques", [])]


# ══════════════════════════════════════════════════════════════════════════
# Registro de campos integrados
# ══════════════════════════════════════════════════════════════════════════
#
# Cada entrada describe UN hecho territorial y de donde sale. `fuente` es la
# procedencia que manda sobre el hecho; `resumen` y `campo` son las claves
# con que lo declara cada cuerpo documental. Un hecho que solo una fuente
# levanta lleva la otra clave en None. Esta tabla es la que evita la
# duplicidad: la integracion recorre estas entradas, no los diccionarios de
# origen, de modo que un mismo hecho no puede escribirse dos veces.

S_IDENT = "1. Identificación y localización"
S_FISICO = "2. Parámetros físicos del terreno"
S_VEGETACION = "3. Índices de vegetación (teledetección)"
S_EROSION = "4. Suelos y erosión (F-DT-02)"
S_ECOSISTEMA = "5. Ecosistema, cobertura y flora (F-DT-03)"
S_DEGRADACION = "6. Degradación, causas y urgencia (F-DT-04)"
S_HIDRO = "7. Hidrología, acceso y logística (F-DT-05)"
S_RESPONSABLE = "8. Responsable y modalidad de la evaluación"

SECCIONES = (S_IDENT, S_FISICO, S_VEGETACION, S_EROSION, S_ECOSISTEMA,
             S_DEGRADACION, S_HIDRO, S_RESPONSABLE)


def _c(clave, etiqueta, seccion, fuente, resumen=None, campo=None,
       tipo="texto", tolerancia=0.0, unidad=""):
    return {"clave": clave, "etiqueta": etiqueta, "seccion": seccion,
            "fuente": fuente, "resumen": resumen, "campo": campo,
            "tipo": tipo, "tolerancia": tolerancia, "unidad": unidad}


CAMPOS_INTEGRADOS = [
    # ── 1. Identificacion y localizacion ──
    _c("codigo_bloque", "Código del bloque", S_IDENT, OFICIAL,
       "codigo_bloque", "codigo_bloque"),
    _c("microcuenca", "Microcuenca (catálogo ANA)", S_IDENT, OFICIAL,
       "microcuenca", "microcuenca"),
    _c("zona_planificacion", "Zona de planificación", S_IDENT, OFICIAL,
       "zona_planificacion"),
    _c("departamento", "Departamento", S_IDENT, OFICIAL, "departamento"),
    _c("provincia", "Provincia", S_IDENT, OFICIAL, "provincia", "provincia"),
    _c("distrito", "Distrito", S_IDENT, OFICIAL, "distrito", "distrito"),
    _c("capital_distrital", "Capital distrital", S_IDENT, OFICIAL,
       "capital_distrital"),
    _c("centro_poblado", "Centro poblado asociado", S_IDENT, CAMPO,
       "centro_poblado", "centro_poblado_cercano"),
    _c("comunidad_campesina", "Comunidad campesina", S_IDENT, CAMPO,
       "comunidad_campesina", "comunidad_campesina_dt"),
    _c("area_ha", "Superficie de catálogo V5/V6 (ha)", S_IDENT, OFICIAL,
       "area_ha", tipo="numero", unidad="ha"),
    _c("utm_este", "Centroide UTM ESTE (m)", S_IDENT, GABINETE,
       "utm_este", tipo="numero", unidad="m"),
    _c("utm_norte", "Centroide UTM NORTE (m)", S_IDENT, GABINETE,
       "utm_norte", tipo="numero", unidad="m"),
    _c("utm_este_campo", "UTM ESTE del punto de muestreo (m)", S_IDENT, CAMPO,
       None, "utm_este_dt", tipo="numero", unidad="m"),
    _c("utm_norte_campo", "UTM NORTE del punto de muestreo (m)", S_IDENT, CAMPO,
       None, "utm_norte_dt", tipo="numero", unidad="m"),
    _c("tipo_intervencion", "Tipo de intervención", S_IDENT, OFICIAL,
       "tipo_intervencion"),

    # ── 2. Parametros fisicos ──
    _c("altitud_min", "Altitud mínima (msnm)", S_FISICO, GABINETE,
       "altitud_min", tipo="numero", unidad="msnm"),
    _c("altitud_max", "Altitud máxima (msnm)", S_FISICO, GABINETE,
       "altitud_max", tipo="numero", unidad="msnm"),
    _c("amplitud_altitudinal", "Amplitud altitudinal (m)", S_FISICO, GABINETE,
       "amplitud_altitudinal", tipo="numero", unidad="m"),
    _c("altitud_gps", "Altitud GPS del punto de muestreo (msnm)", S_FISICO,
       CAMPO, None, "altitud_gps", tipo="numero", unidad="msnm"),
    # El piso dominante del bloque (estadistica zonal sobre el MDE) y el
    # piso del punto de muestreo (F-DT-01) son hechos distintos: un bloque
    # de 600 m de amplitud atraviesa varios pisos. No se cruzan entre si;
    # su coherencia se verifica en el control de consistencia.
    _c("piso_altitudinal", "Piso altitudinal dominante del bloque", S_FISICO,
       GABINETE, "piso_altitudinal"),
    _c("rango_altitudinal_campo", "Piso altitudinal del punto de muestreo",
       S_FISICO, CAMPO, None, "rango_altitudinal"),
    _c("pendiente_pct", "Pendiente promedio (%)", S_FISICO, GABINETE,
       "pendiente_pct", tipo="numero", unidad="%"),
    _c("pendiente_grados", "Pendiente promedio (grados)", S_FISICO, GABINETE,
       "pendiente_grados", tipo="numero", unidad="grados"),
    _c("clase_pendiente", "Clase de pendiente equivalente", S_FISICO, GABINETE,
       "clase_pendiente"),
    _c("pendiente_campo", "Rango de pendiente declarado en campo", S_FISICO,
       CAMPO, "pendiente_campo", "pendiente"),
    _c("forma_terreno", "Forma predominante del terreno", S_FISICO, CAMPO,
       "forma_terreno", "forma_terreno"),
    _c("posicion_fisiografica", "Posición fisiográfica", S_FISICO, CAMPO,
       "posicion_fisiografica", "posicion_fisiografica"),
    _c("exposicion", "Exposición / orientación", S_FISICO, CAMPO,
       "exposicion", "exposicion_orientacion"),
    _c("paisaje_dominante", "Paisaje dominante descrito en campo", S_FISICO,
       CAMPO, None, "paisaje_dominante"),
    _c("afloramientos_rocosos", "Afloramientos rocosos", S_FISICO, CAMPO,
       "afloramientos_rocosos", "dt01_afloramientos_rocosos"),
    _c("escarpes_activos", "Escarpes activos", S_FISICO, CAMPO,
       "escarpes_activos", "dt01_escarpes_activos"),
    _c("reptacion_suelo", "Reptación de suelo", S_FISICO, CAMPO,
       None, "dt01_reptacion_suelo"),
    _c("deslizamientos_antiguos", "Deslizamientos antiguos", S_FISICO, CAMPO,
       None, "dt01_deslizamientos_antiguos"),
    _c("remociones_masa", "Remociones en masa activas", S_FISICO, CAMPO,
       "remociones_masa", "dt01_remociones_masa_activas"),

    # ── 3. Indices de vegetacion (solo gabinete) ──
    _c("msavi_2024", "MSAVI 2024 - media del bloque", S_VEGETACION, GABINETE,
       "msavi_2024", tipo="numero"),
    _c("msavi_clase", "MSAVI 2024 - clase de la media", S_VEGETACION,
       GABINETE, "msavi_clase"),
    _c("condicion_umbral", "Condición frente al umbral %s" % rbq.UMBRAL_MSAVI,
       S_VEGETACION, GABINETE, "condicion_umbral"),
    _c("ndvi_clase_modal", "NDVI mediana 2025 - clase modal", S_VEGETACION,
       GABINETE, "ndvi_clase_modal"),
    _c("superficie_ndvi_ha", "Superficie clasificada NDVI 2025 (ha)",
       S_VEGETACION, GABINETE, "superficie_ndvi_ha", tipo="numero", unidad="ha"),
    _c("superficie_msavi_ha", "Superficie clasificada MSAVI 2024 (ha)",
       S_VEGETACION, GABINETE, "superficie_msavi_ha", tipo="numero", unidad="ha"),
    _c("superficie_bajo_umbral_ha",
       "Superficie bajo umbral MSAVI %s (ha)" % rbq.UMBRAL_MSAVI,
       S_VEGETACION, GABINETE, "superficie_bajo_umbral_ha", tipo="numero",
       unidad="ha"),
    _c("bajo_umbral_pct", "% del bloque bajo umbral (brecha espectral)",
       S_VEGETACION, GABINETE, "bajo_umbral_pct", tipo="numero", unidad="%"),
    _c("msavi_clase_dominante", "Clase DN dominante (MSAVI 2024)",
       S_VEGETACION, GABINETE, "msavi_clase_dominante"),
    _c("msavi_poligonos", "N.° de polígonos MSAVI del bloque", S_VEGETACION,
       GABINETE, "msavi_poligonos", tipo="numero"),

    # ── 4. Suelos y erosion ──
    _c("nivel_erosion", "Nivel general de erosión", S_EROSION, CAMPO,
       "nivel_erosion", "dt02_nivel_erosion_general"),
    _c("nivel_erosion_sintesis", "Nivel de erosión - síntesis de campo",
       S_EROSION, CAMPO, None, "dt02_nivel_erosion_sintesis"),
    _c("sellamiento_costra", "Sellamiento / costra superficial", S_EROSION,
       CAMPO, None, "dt02_sellamiento_costra"),
    _c("compactacion_pisoteo", "Compactación por pisoteo", S_EROSION, CAMPO,
       None, "dt02_compactacion_pisoteo"),
    _c("raices_expuestas", "Raíces expuestas", S_EROSION, CAMPO, None,
       "dt02_raices_expuestas"),
    _c("n_carcavas", "N.° de cárcavas registradas", S_EROSION, CAMPO,
       "n_carcavas", "dt02_num_carcavas", tipo="numero"),
    _c("longitud_carcavas", "Longitud total de cárcavas (m)", S_EROSION,
       CAMPO, None, "dt02_longitud_total_carcavas", tipo="numero", unidad="m"),
    _c("pct_bloque_carcavas", "% del bloque afectado por cárcavas", S_EROSION,
       CAMPO, None, "dt02_pct_bloque_carcavas", tipo="numero", unidad="%"),
    _c("erosion_laminar_pct", "Erosión laminar (% del bloque)", S_EROSION,
       CAMPO, None, "dt02_erosion_laminar_pct", tipo="numero", unidad="%"),
    _c("patron_carcavas", "Patrón de cárcavas", S_EROSION, CAMPO, None,
       "dt02_patron_carcavas"),
    _c("socavamiento_cauce", "Socavamiento de cauce", S_EROSION, CAMPO, None,
       "dt02_socavamiento_cauce"),
    _c("urgencia_erosion", "Urgencia de control de erosión", S_EROSION, CAMPO,
       "urgencia_erosion", "dt02_urgencia_control"),

    # ── 5. Ecosistema, cobertura y flora ──
    _c("tipo_ecosistema", "Tipo de ecosistema (UP)", S_ECOSISTEMA, CAMPO,
       "tipo_ecosistema", "dt03_tipo_ecosistema", tipo="clase"),
    _c("superficie_ecosistema", "Superficie de ecosistema (ha)", S_ECOSISTEMA,
       CAMPO, "superficie_ecosistema", "dt03_superficie_ecosistema"),
    _c("estado_conservacion", "Estado de conservación", S_ECOSISTEMA, CAMPO,
       "estado_conservacion", "dt03_estado_conservacion_eco"),
    _c("uso_dominante", "Uso actual dominante del suelo", S_ECOSISTEMA, CAMPO,
       "uso_dominante", "dt03_uso_dominante"),
    _c("tipo_cobertura", "Tipo de cobertura dominante", S_ECOSISTEMA, CAMPO,
       "tipo_cobertura", "dt03_tipo_cobertura_dom"),
    _c("cobertura_total_pct", "Cobertura vegetal total - campo (%)",
       S_ECOSISTEMA, CAMPO, "cobertura_total_pct", "dt03_cobertura_total",
       tipo="numero", unidad="%"),
    _c("cobertura_dosel", "Cobertura del dosel (%)", S_ECOSISTEMA, CAMPO,
       None, "dt03_cobertura_dosel", tipo="numero", unidad="%"),
    _c("cobertura_arbustiva", "Cobertura arbustiva (%)", S_ECOSISTEMA, CAMPO,
       None, "dt03_cobertura_arbustiva", tipo="numero", unidad="%"),
    _c("cobertura_herbacea", "Cobertura herbácea (%)", S_ECOSISTEMA, CAMPO,
       None, "dt03_cobertura_herbacea", tipo="numero", unidad="%"),
    _c("cobertura_hojarasca", "Cobertura de hojarasca (%)", S_ECOSISTEMA,
       CAMPO, None, "dt03_cobertura_hojarasca", tipo="numero", unidad="%"),
    _c("suelo_desnudo_pct", "Suelo desnudo - campo (%)", S_ECOSISTEMA, CAMPO,
       "suelo_desnudo_pct", "dt03_suelo_desnudo", tipo="numero", unidad="%"),
    _c("altura_estrato_dom", "Altura del estrato dominante (m)", S_ECOSISTEMA,
       CAMPO, None, "dt03_altura_estrato_dom", tipo="numero", unidad="m"),
    _c("altura_max", "Altura máxima registrada (m)", S_ECOSISTEMA, CAMPO,
       None, "dt03_altura_max", tipo="numero", unidad="m"),
    _c("dap_promedio", "DAP promedio (cm)", S_ECOSISTEMA, CAMPO, None,
       "dt03_dap_promedio", tipo="numero", unidad="cm"),
    _c("regeneracion", "Regeneración natural", S_ECOSISTEMA, CAMPO,
       "regeneracion", "dt03_regeneracion_natural"),
    _c("estado_sanitario", "Estado sanitario", S_ECOSISTEMA, CAMPO,
       "estado_sanitario", "dt03_estado_sanitario"),
    _c("presencia_epifitas", "Presencia de epífitas", S_ECOSISTEMA, CAMPO,
       None, "dt03_presencia_epifitas"),
    _c("fenologia_dominante", "Fenología dominante", S_ECOSISTEMA, CAMPO,
       None, "dt03_fenologia_dominante"),
    _c("parcela_muestreo", "Parcela de muestreo", S_ECOSISTEMA, CAMPO,
       "parcela_muestreo", "dt03_parcela_muestreo"),
    _c("dim_parcela", "Dimensiones de la parcela (m)", S_ECOSISTEMA, CAMPO,
       None, "dt03_dim_parcela"),
    _c("pendiente_parcela", "Pendiente de la parcela (%)", S_ECOSISTEMA,
       CAMPO, None, "dt03_pendiente_parcela", tipo="numero", unidad="%"),
    _c("n_taxones", "Elenco florístico (n.o de taxones)", S_ECOSISTEMA, CAMPO,
       "n_taxones", "n_taxones", tipo="numero"),
    _c("n_especies_clave", "Especies clave georreferenciadas", S_ECOSISTEMA,
       CAMPO, None, "n_especies_clave", tipo="numero"),

    # ── 6. Degradacion, causas y urgencia ──
    _c("causas_directas", "Causas directas de la degradación", S_DEGRADACION,
       CAMPO, None, "dt04_causas_directas_texto"),
    _c("causa_subyacente", "Causa subyacente principal", S_DEGRADACION, CAMPO,
       "causa_subyacente", "dt04_causa_subyacente"),
    _c("n_causas_activas", "N.° de causas activas en la matriz F-DT-04",
       S_DEGRADACION, CAMPO, None, "n_causas_activas", tipo="numero"),
    _c("velocidad_degradacion", "Velocidad de degradación", S_DEGRADACION,
       CAMPO, "velocidad_degradacion", "dt04_velocidad_degradacion"),
    _c("reversibilidad", "Reversibilidad técnica", S_DEGRADACION, CAMPO,
       "reversibilidad", "dt04_reversibilidad"),
    _c("urgencia_intervencion", "Urgencia de intervención", S_DEGRADACION,
       CAMPO, "urgencia_intervencion", "dt04_urgencia_intervencion"),
    _c("peligro_integrado", "Peligro integrado preliminar (MCA-AHP)",
       S_DEGRADACION, GABINETE, "peligro_integrado"),
    _c("prioridad", "Prioridad de intervención", S_DEGRADACION, GABINETE,
       "prioridad"),

    # ── 7. Hidrologia, acceso y logistica ──
    _c("zona_recarga", "Zona de recarga hídrica", S_HIDRO, CAMPO,
       "zona_recarga", "dt05_zona_recarga"),
    _c("escorrentia_concentrada", "Escorrentía concentrada", S_HIDRO, CAMPO,
       None, "dt05_escorrentia_concentrada"),
    _c("n_fuentes_agua", "Fuentes de agua inventariadas", S_HIDRO, CAMPO,
       None, "n_fuentes_agua", tipo="numero"),
    _c("dist_captacion", "Distancia a captación (m)", S_HIDRO, CAMPO, None,
       "dt05_dist_captacion", tipo="numero", unidad="m"),
    _c("jass_captacion", "JASS / usuario de la captación", S_HIDRO, CAMPO,
       None, "dt05_jass_captacion"),
    _c("interferencia_riego", "Interferencia con riego", S_HIDRO, CAMPO,
       None, "dt05_interferencia_riego"),
    _c("modalidad_acceso", "Modalidad de acceso", S_HIDRO, CAMPO,
       "modalidad_acceso", "dt05_modalidad_acceso"),
    _c("via_principal", "Vía principal de acceso", S_HIDRO, CAMPO, None,
       "dt05_via_principal"),
    _c("tipo_via_final", "Tipo de vía del tramo final", S_HIDRO, CAMPO, None,
       "dt05_tipo_via_final"),
    _c("transitabilidad_seca", "Transitabilidad en época seca", S_HIDRO,
       CAMPO, None, "dt05_transitabilidad_seca"),
    _c("transitabilidad_lluviosa", "Transitabilidad en época lluviosa",
       S_HIDRO, CAMPO, None, "dt05_transitabilidad_lluviosa"),
    _c("tiempo_dist_capital", "Tiempo a la capital distrital (min)", S_HIDRO,
       CAMPO, None, "dt05_tiempo_dist_capital", tipo="numero", unidad="min"),
    _c("senal_celular", "Señal celular", S_HIDRO, CAMPO, None,
       "dt05_senal_celular"),
    _c("operador_celular", "Operador con cobertura", S_HIDRO, CAMPO, None,
       "dt05_operador_celular"),
    _c("alojamiento", "Alojamiento disponible en la zona", S_HIDRO, CAMPO,
       None, "dt05_alojamiento"),
    _c("requiere_ronda", "Requiere coordinación con ronda / comunidad",
       S_HIDRO, CAMPO, None, "dt05_requiere_ronda"),
    _c("contacto_ronda", "Contacto local de coordinación", S_HIDRO, CAMPO,
       None, "dt05_contacto_ronda"),

    # ── 8. Responsable y modalidad ──
    _c("evaluador", "Responsable de la evaluación", S_RESPONSABLE, CAMPO,
       "evaluador", "evaluador"),
    _c("fecha_evaluacion", "Fecha de evaluación", S_RESPONSABLE, CAMPO,
       "fecha_evaluacion", "fecha_evaluacion", tipo="fecha"),
    _c("hora_registro", "Hora de registro", S_RESPONSABLE, CAMPO,
       "hora_registro", "hora_registro"),
    _c("correlativo_ficha", "Correlativo de ficha", S_RESPONSABLE, CAMPO,
       "correlativo_ficha", "ficha_correlativo"),
    _c("entidad", "Entidad", S_RESPONSABLE, OFICIAL, "entidad"),
    _c("fase_estudio", "Fase del estudio", S_RESPONSABLE, OFICIAL,
       "fase_estudio"),
    _c("instrumento", "Instrumento aplicado", S_RESPONSABLE, OFICIAL,
       "instrumento"),
    _c("n_estaciones", "Estaciones fotográficas georreferenciadas",
       S_RESPONSABLE, GABINETE, "n_estaciones", tipo="numero"),
    _c("estado_verificacion", "Estado de verificación de campo",
       S_RESPONSABLE, OFICIAL, "estado_verificacion"),
    _c("marco_indicador", "Marco del indicador de brecha", S_RESPONSABLE,
       OFICIAL, "marco_indicador"),
]

# Indice por clave, para consultar un campo sin recorrer la lista.
CAMPOS_POR_CLAVE = {c["clave"]: c for c in CAMPOS_INTEGRADOS}


# ══════════════════════════════════════════════════════════════════════════
# Integracion de un bloque
# ══════════════════════════════════════════════════════════════════════════

_FORMATOS_FECHA = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d",
                   "%d/%m/%y", "%Y-%m-%d %H:%M:%S")


def _fecha_iso(valor):
    """Fecha en ISO a partir de las variantes que usan las dos fuentes."""
    texto = _txt(valor)
    if not texto:
        return ""
    texto = texto.split(" ")[0] if re.match(r"^\d", texto) else texto
    for formato in _FORMATOS_FECHA:
        try:
            return datetime.strptime(texto, formato).date().isoformat()
        except ValueError:
            continue
    return texto


_RE_ACRONIMO = re.compile(r"\(([A-Za-z][A-Za-z\-]{1,7})\)")


def _clase_nomenclada(valor):
    """Acronimo de la clase entre parentesis: «... (Bes-cm)» -> «bes-cm».

    Las listas oficiales del proyecto nombran cada clase con un acronimo
    estable y una glosa que cambia de redaccion entre documentos. El
    acronimo es lo que identifica la clase.
    """
    m = _RE_ACRONIMO.search(_txt(valor))
    return m.group(1).lower() if m else ""


def _relacion(valor_a, valor_b, tipo, tolerancia):
    """Relacion entre los valores de las dos fuentes para un mismo hecho.

    Devuelve "igual" cuando declaran lo mismo, "amplia" cuando el segundo
    dice lo mismo con mas detalle (y viceversa con "reduce"), o "distinto"
    cuando se contradicen. La distincion importa: una redaccion mas completa
    del mismo hecho no es una discrepancia y no debe ensuciar el control de
    consistencia.
    """
    if tipo == "fecha":
        return "igual" if _fecha_iso(valor_a) == _fecha_iso(valor_b) else "distinto"
    if tipo == "clase":
        clase_a, clase_b = _clase_nomenclada(valor_a), _clase_nomenclada(valor_b)
        if clase_a and clase_b:
            return "igual" if clase_a == clase_b else "distinto"
    if tipo == "numero":
        n_a, n_b = _num(valor_a), _num(valor_b)
        if n_a is not None and n_b is not None:
            tol = tolerancia or max(abs(n_a), abs(n_b)) * 0.005
            return "igual" if abs(n_a - n_b) <= max(tol, 1e-9) else "distinto"
    canon_a, canon_b = _comparable(valor_a), _comparable(valor_b)
    if canon_a == canon_b:
        return "igual"
    if canon_a and canon_b:
        if canon_b.startswith(canon_a) or canon_a in canon_b:
            return "amplia"
        if canon_a.startswith(canon_b) or canon_b in canon_a:
            return "reduce"
    return "distinto"


def _cruzar(entrada, resumen, campo):
    """Cruza un hecho entre las dos fuentes y decide su valor unico."""
    clave_res, clave_cam = entrada["resumen"], entrada["campo"]
    val_res = _txt(resumen.get(clave_res)) if clave_res else ""
    val_cam = _txt(campo.get(clave_cam)) if clave_cam else ""
    hay_res = bool(clave_res) and not _sin_dato(val_res)
    hay_cam = bool(clave_cam) and not _sin_dato(val_cam)

    fuente_mandante = entrada["fuente"]
    # La ficha de resumen nunca es "campo" por si misma: los valores de
    # campo que ya traia provienen de una version anterior de la misma
    # plantilla DT, de modo que al cruzarlos la version actualizada manda.
    fuente_resumen = GABINETE if fuente_mandante == CAMPO else fuente_mandante

    registro = dict(entrada)
    registro.update({
        "valor_resumen": val_res, "valor_campo": val_cam,
        "valor_alterno": "", "fuente_alterna": "", "nota": "",
    })

    if not hay_res and not hay_cam:
        registro.update({"valor": "", "fuente_valor": fuente_mandante,
                         "estado": PENDIENTE})
        return registro

    if hay_res and not hay_cam:
        registro.update({
            "valor": val_res,
            "fuente_valor": fuente_resumen,
            "estado": CONFORME if fuente_mandante != CAMPO else COMPLEMENTADO,
            "nota": ("" if fuente_mandante != CAMPO else
                     "La ficha de campo vigente no declara este dato; se "
                     "conserva el valor del libro de resumen."),
        })
        return registro

    if hay_cam and not hay_res:
        registro.update({
            "valor": val_cam, "fuente_valor": CAMPO,
            "estado": CONFORME if fuente_mandante == CAMPO else COMPLEMENTADO,
            "nota": ("" if fuente_mandante == CAMPO else
                     "Dato levantado en campo que el libro de resumen no "
                     "consignaba."),
        })
        return registro

    relacion = _relacion(val_res, val_cam, entrada["tipo"],
                         entrada["tolerancia"])
    if relacion == "igual":
        registro.update({
            "valor": val_cam if fuente_mandante == CAMPO else val_res,
            "fuente_valor": fuente_mandante, "estado": CONFORME})
        if _comparable(val_res) != _comparable(val_cam):
            # Mismo hecho escrito de dos maneras (acronimo de clase, formato
            # de fecha, unidades). No es discrepancia, pero la redaccion
            # descartada se conserva para poder rastrearla.
            descartado = val_res if fuente_mandante == CAMPO else val_cam
            registro.update({
                "valor_alterno": descartado,
                "fuente_alterna": (fuente_resumen if fuente_mandante == CAMPO
                                   else CAMPO),
                "nota": "Ambas fuentes declaran lo mismo con distinta "
                        "redacción; se adopta la nomenclatura de la fuente "
                        "que manda sobre el dato."})
        return registro
    if relacion in ("amplia", "reduce"):
        # Misma declaracion con distinto grado de detalle: se conserva la
        # redaccion mas completa y se deja constancia de la otra.
        mas_completo = val_cam if relacion == "amplia" else val_res
        fuente_completo = CAMPO if relacion == "amplia" else fuente_resumen
        registro.update({
            "valor": mas_completo, "fuente_valor": fuente_completo,
            "estado": CONFORME,
            "nota": "Ambas fuentes coinciden; se conserva la redacción más "
                    "completa."})
        return registro

    if fuente_mandante == CAMPO:
        registro.update({
            "valor": val_cam, "fuente_valor": CAMPO, "estado": ACTUALIZADO,
            "valor_alterno": val_res, "fuente_alterna": fuente_resumen,
            "nota": "La verificación de campo vigente actualiza el valor que "
                    "traía el libro de resumen."})
    else:
        registro.update({
            "valor": val_res, "fuente_valor": fuente_mandante,
            "estado": DISCREPANTE, "valor_alterno": val_cam,
            "fuente_alterna": CAMPO,
            "nota": "Se conserva el valor de la fuente que manda sobre el "
                    "dato y se declara la observación de campo."})
    return registro


def integrar_bloque(resumen, campo, codigo=""):
    """Registro integrado de un bloque: un hecho, un valor, una procedencia.

    `resumen` es la salida de `resumenes_bloques.parsear_resumen_bloque` y
    `campo` la de `parsear_ficha_campo`. Cualquiera de los dos puede venir
    vacio: el bloque se integra igual y el faltante queda declarado.
    """
    resumen = resumen or {}
    campo = campo or {}
    codigo = normalizar_codigo(
        codigo or resumen.get("codigo_bloque") or campo.get("codigo_bloque"))

    campos = [_cruzar(e, resumen, campo) for e in CAMPOS_INTEGRADOS]
    por_clave = {c["clave"]: c for c in campos}

    conteo_estado = {e: 0 for e in ESTADOS}
    conteo_fuente = {f: 0 for f in FUENTES}
    for registro in campos:
        conteo_estado[registro["estado"]] += 1
        if registro["estado"] != PENDIENTE:
            conteo_fuente[registro["fuente_valor"]] += 1

    declarados = sum(1 for c in campos if c["estado"] != PENDIENTE)
    integrado = {
        "codigo_bloque": codigo,
        "version": VERSION_INTEGRACION,
        "fecha_integracion": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "campos": campos,
        "por_clave": por_clave,
        "conteo_estado": conteo_estado,
        "conteo_fuente": conteo_fuente,
        "n_campos": len(campos),
        "n_declarados": declarados,
        "cobertura_pct": round(100.0 * declarados / len(campos), 1) if campos else 0.0,
        "tiene_resumen": bool(resumen),
        "tiene_campo": bool(campo),
        "fichas_leidas": campo.get("fichas_leidas", []),
        "fichas_ausentes": campo.get("fichas_ausentes", list(FICHAS_CAMPO)),
        "completitud_campo_pct": campo.get("completitud_pct", 0.0),
        "tablas_vacias": campo.get("tablas_vacias", []),
        "nombre_archivo_campo": campo.get("nombre_archivo", ""),
        "nombre_archivo_resumen": resumen.get("nombre_archivo", ""),
        # Tablas: viajan intactas desde su fuente, sin cruzarse, porque
        # ninguna de las dos las duplica.
        "carcavas": campo.get("carcavas", []),
        "floristica": campo.get("floristica", []),
        "especies_clave": campo.get("especies_clave", []),
        "causas": campo.get("causas", []),
        "causas_activas": campo.get("causas_activas", []),
        "indicadores": campo.get("indicadores", []),
        "fuentes_agua": campo.get("fuentes_agua", []),
        "observaciones_campo": {
            "F-DT-01": _txt(campo.get("dt01_observaciones")),
            "F-DT-02": _txt(campo.get("dt02_observaciones")),
            "F-DT-03": _txt(campo.get("dt03_observaciones")),
            "F-DT-04": _txt(campo.get("dt04_observaciones")),
            "F-DT-05": _txt(campo.get("dt05_observaciones")),
        },
        # Contenido del libro de resumen que la integracion no toca.
        "msavi_tabla": resumen.get("msavi_tabla", []),
        "ndvi_tabla": resumen.get("ndvi_tabla", []),
        "estaciones": resumen.get("estaciones", []),
        "microcuenca_tabla": resumen.get("microcuenca_tabla", []),
        "consistencia_v6": resumen.get("consistencia", []),
        "msavi_total_ha": resumen.get("msavi_total_ha"),
        "msavi_bajo_umbral_ha": resumen.get("msavi_bajo_umbral_ha"),
        "msavi_bajo_umbral_pct": resumen.get("msavi_bajo_umbral_pct"),
        "msavi_sobre_umbral_ha": resumen.get("msavi_sobre_umbral_ha"),
        "msavi_sobre_umbral_pct": resumen.get("msavi_sobre_umbral_pct"),
    }

    # Atajos de uso frecuente en tablas y filtros del aplicativo.
    for clave in ("provincia", "distrito", "microcuenca", "centro_poblado",
                  "tipo_ecosistema", "estado_conservacion", "nivel_erosion",
                  "urgencia_intervencion", "velocidad_degradacion",
                  "modalidad_acceso", "evaluador", "fecha_evaluacion",
                  "estado_verificacion", "peligro_integrado"):
        integrado[clave] = por_clave[clave]["valor"]
    for clave in ("area_ha", "utm_este", "utm_norte", "utm_este_campo",
                  "utm_norte_campo", "msavi_2024", "pendiente_pct",
                  "cobertura_total_pct", "suelo_desnudo_pct",
                  "bajo_umbral_pct", "n_taxones", "n_carcavas",
                  "altitud_gps", "altitud_min", "altitud_max"):
        integrado[clave + "_num"] = _num(por_clave[clave]["valor"])

    integrado["distancia_centroide_m"] = _distancia(
        integrado.get("utm_este_num"), integrado.get("utm_norte_num"),
        integrado.get("utm_este_campo_num"), integrado.get("utm_norte_campo_num"))
    integrado["validacion_utm"] = rbq.validar_utm(
        integrado.get("utm_este_campo_num") or integrado.get("utm_este_num"),
        integrado.get("utm_norte_campo_num") or integrado.get("utm_norte_num"))
    integrado["consistencia"] = control_consistencia(integrado)
    integrado["consistencia_resumen"] = _resumen_consistencia(
        integrado["consistencia"])
    return integrado


def _distancia(este_a, norte_a, este_b, norte_b):
    """Distancia plana en metros entre dos pares UTM. None si falta alguno."""
    if None in (este_a, norte_a, este_b, norte_b):
        return None
    return round(math.hypot(este_a - este_b, norte_a - norte_b), 1)


def _resumen_consistencia(registros):
    conteo = {c: 0 for c in rbq.CALIFICACIONES}
    for registro in registros:
        calificacion = registro.get("calificacion", "")
        if calificacion in conteo:
            conteo[calificacion] += 1
    conteo["total"] = len(registros)
    return conteo


# ══════════════════════════════════════════════════════════════════════════
# Control de consistencia regenerado
# ══════════════════════════════════════════════════════════════════════════

# Hechos cuya ausencia o contradiccion compromete el sustento tecnico del
# bloque ante la DGPMI: se califican como SUSTANTIVA.
CAMPOS_CRITICOS = {
    "microcuenca", "area_ha", "utm_este", "utm_norte", "peligro_integrado",
    "msavi_2024", "tipo_ecosistema", "estado_conservacion", "nivel_erosion",
    "urgencia_intervencion", "provincia", "distrito",
}

# Tolerancia geometrica: un punto de muestreo puede caer lejos del centroide
# sin que ello sea un error, pero mas alla de este radio conviene revisar si
# la estacion quedo dentro del poligono del bloque.
RADIO_CONTROL_M = 1500.0


def _fila_consistencia(codigo, campo, discrepancia, calificacion, tratamiento):
    return {"codigo": codigo, "campo": campo, "discrepancia": discrepancia,
            "calificacion": calificacion, "tratamiento": tratamiento}


def _verificaciones_geometricas(integrado):
    """Controles de geometria y de coherencia espectral del bloque."""
    filas = []
    por_clave = integrado["por_clave"]
    n = 0

    def cod():
        return "G-%02d" % n

    # Microcuenca declarada en campo frente al catalogo de la ANA.
    n += 1
    micro = por_clave["microcuenca"]
    if micro["estado"] == DISCREPANTE:
        filas.append(_fila_consistencia(
            cod(), "Código de microcuenca",
            "La ficha DT declara la microcuenca «%s» y el catálogo maestro "
            "registra «%s»." % (micro["valor_alterno"], micro["valor"]),
            "SUSTANTIVA",
            "Prevalece el código del catálogo de la ANA. Debe corregirse la "
            "ficha de campo en la próxima salida."))
    elif micro["estado"] == PENDIENTE:
        filas.append(_fila_consistencia(
            cod(), "Código de microcuenca",
            "Ninguna de las dos fuentes declara el código de microcuenca.",
            "SUSTANTIVA", "Debe completarse desde el catálogo de la ANA."))
    else:
        filas.append(_fila_consistencia(
            cod(), "Código de microcuenca",
            "Ficha DT y catálogo maestro coinciden en «%s»." % micro["valor"],
            "CONFORME", "Sin acción."))

    # Punto de muestreo frente al centroide del bloque.
    n += 1
    distancia = integrado.get("distancia_centroide_m")
    if distancia is None:
        filas.append(_fila_consistencia(
            cod(), "UTM del punto de muestreo",
            "No es posible contrastar el punto de muestreo con el centroide: "
            "falta una de las dos coordenadas.",
            "SUSTANTIVA",
            "Debe consignarse la coordenada del punto de muestreo en F-DT-01."))
    elif distancia > RADIO_CONTROL_M:
        filas.append(_fila_consistencia(
            cod(), "UTM del punto de muestreo",
            "El punto de muestreo declarado dista %.1f m del centroide del "
            "bloque, por encima del radio de control de %.0f m."
            % (distancia, RADIO_CONTROL_M),
            "NO SUSTANTIVA",
            "Se verifica en gabinete si la estación cae dentro del polígono; "
            "el dato de campo se conserva como referencia de acceso."))
    else:
        filas.append(_fila_consistencia(
            cod(), "UTM del punto de muestreo",
            "El punto de muestreo declarado dista %.1f m del centroide, "
            "dentro del radio de control." % distancia,
            "CONFORME", "Sin acción."))

    # Validacion del sistema de referencia.
    n += 1
    validacion = integrado.get("validacion_utm", "")
    filas.append(_fila_consistencia(
        cod(), "Sistema de referencia UTM 17S",
        "Verificación de rango para EPSG:32717: %s." % (validacion or "sin dato"),
        "CONFORME" if validacion == "Conforme" else "SUSTANTIVA",
        "Sin acción." if validacion == "Conforme" else
        "Debe reproyectarse o corregirse la coordenada antes de la "
        "microlocalización."))

    # Altitud del punto de muestreo frente al rango del MDE.
    n += 1
    gps = integrado.get("altitud_gps_num")
    alt_min = integrado.get("altitud_min_num")
    alt_max = integrado.get("altitud_max_num")
    if gps is None or alt_min is None or alt_max is None:
        filas.append(_fila_consistencia(
            cod(), "Altitud del punto de muestreo",
            "No es posible contrastar la altitud GPS con el rango altitudinal "
            "del bloque: falta uno de los dos datos.",
            "NO SUSTANTIVA",
            "Se completa con la lectura GPS de F-DT-01 o con la estadística "
            "zonal del MDE."))
    elif alt_min - 50 <= gps <= alt_max + 50:
        filas.append(_fila_consistencia(
            cod(), "Altitud del punto de muestreo",
            "La altitud GPS (%.0f msnm) cae dentro del rango altitudinal del "
            "bloque (%.0f - %.0f msnm)." % (gps, alt_min, alt_max),
            "CONFORME", "Sin acción."))
    else:
        filas.append(_fila_consistencia(
            cod(), "Altitud del punto de muestreo",
            "La altitud GPS (%.0f msnm) queda fuera del rango altitudinal del "
            "bloque (%.0f - %.0f msnm)." % (gps, alt_min, alt_max),
            "NO SUSTANTIVA",
            "Se revisa si la estación cayó fuera del polígono; prevalece el "
            "rango del MDE para los parámetros físicos del bloque."))

    # Superficie segmentada frente al catalogo.
    for clave, etiqueta, insumo in (
            ("superficie_ndvi_ha", "Superficie NDVI vs. catálogo",
             "NDVI mediana 2025"),
            ("superficie_msavi_ha", "Superficie MSAVI vs. catálogo",
             "MSAVI 2024")):
        n += 1
        area = integrado.get("area_ha_num")
        segmentada = _num(por_clave[clave]["valor"])
        if area and segmentada:
            desviacion = 100.0 * (segmentada - area) / area
            conforme = abs(desviacion) <= 2.0
            filas.append(_fila_consistencia(
                cod(), etiqueta,
                "La superficie clasificada por el %s (%.3f ha) se desvía "
                "%.2f %% de la superficie de catálogo (%.2f ha)."
                % (insumo, segmentada, desviacion, area),
                "CONFORME" if conforme else "NO SUSTANTIVA",
                "Sin acción: la segmentación valida la geometría del bloque."
                if conforme else
                "Se declara la desviación; prevalece la superficie del "
                "catálogo maestro para metas físicas y costos."))
        else:
            filas.append(_fila_consistencia(
                cod(), etiqueta,
                "No se dispone de la superficie clasificada por el %s para "
                "contrastarla con el catálogo." % insumo,
                "NO SUSTANTIVA",
                "Se completa con la estadística zonal del entregable de "
                "teledetección."))

    # Brecha espectral: superficie bajo el umbral del indicador.
    n += 1
    bajo_pct = integrado.get("bajo_umbral_pct_num")
    if bajo_pct is None:
        filas.append(_fila_consistencia(
            cod(), "Brecha espectral MSAVI",
            "El libro no declara el porcentaje del bloque bajo el umbral %s."
            % rbq.UMBRAL_MSAVI,
            "SUSTANTIVA",
            "Debe completarse: es la base del indicador de brecha de la "
            "R.M. N.° 00213-2024-MINAM."))
    else:
        filas.append(_fila_consistencia(
            cod(), "Brecha espectral MSAVI",
            "El %.2f %% del bloque queda bajo el umbral %s; superficie de "
            "brecha %s ha."
            % (bajo_pct, rbq.UMBRAL_MSAVI,
               _fmt_num(integrado.get("msavi_bajo_umbral_ha"))),
            "CONFORME",
            "Sin acción: alimenta el indicador de brecha del proyecto."))

    # Distribucion areal del MSAVI por clase DN (aporte de la version V6).
    n += 1
    clases = [c for c in (integrado.get("msavi_tabla") or [])
              if c.get("superficie_ha") is not None]
    if clases:
        filas.append(_fila_consistencia(
            cod(), "Distribución areal MSAVI 2024",
            "La hoja de cobertura declara %d clases DN con superficie, %s ha "
            "clasificadas en total."
            % (len(clases), _fmt_num(integrado.get("msavi_total_ha"))),
            "CONFORME",
            "Sin acción: procede de la estadística zonal del raster MSAVI "
            "2024 clasificado."))
    else:
        filas.append(_fila_consistencia(
            cod(), "Distribución areal MSAVI 2024",
            "La distribución por clase de MSAVI no figura en los insumos de "
            "este bloque.",
            "SUSTANTIVA",
            "Debe completarse con la estadística zonal del raster MSAVI 2024 "
            "clasificado; no se estima."))

    # Coherencia entre la media del bloque y su clase DN dominante.
    n += 1
    media = integrado.get("msavi_2024_num")
    dominante = por_clave["msavi_clase_dominante"]["valor"]
    clase_media = por_clave["msavi_clase"]["valor"]
    if media is None or _sin_dato(dominante) or _sin_dato(clase_media):
        filas.append(_fila_consistencia(
            cod(), "Media MSAVI vs. clase DN dominante",
            "No es posible contrastar la media del bloque con su clase DN "
            "dominante: falta uno de los dos datos.",
            "NO SUSTANTIVA",
            "Se completa con la estadística zonal del entregable de "
            "teledetección."))
    else:
        coincide = _comparable(clase_media) in _comparable(dominante)
        filas.append(_fila_consistencia(
            cod(), "Media MSAVI vs. clase DN dominante",
            "La media del bloque (%.4f) cae en la clase «%s» y la clase DN "
            "dominante por superficie es «%s»." % (media, clase_media, dominante),
            "CONFORME" if coincide else "NO SUSTANTIVA",
            "Sin acción: media y clase dominante son coherentes." if coincide
            else "Se declaran ambas lecturas: la media resume el bloque y la "
                 "clase dominante su reparto areal."))

    # Peligro integrado del modelamiento de mesolocalizacion.
    n += 1
    peligro = por_clave["peligro_integrado"]
    if peligro["estado"] == PENDIENTE or _sin_dato(peligro["valor"]):
        filas.append(_fila_consistencia(
            cod(), "Peligro integrado (MCA-AHP)",
            "El nivel de peligro integrado no está disponible en los insumos "
            "de este entregable.",
            "SUSTANTIVA",
            "Debe tomarse del modelamiento de mesolocalización (MCA-AHP de "
            "PMM, EPH y PGI) antes de cerrar el perfil."))
    else:
        filas.append(_fila_consistencia(
            cod(), "Peligro integrado (MCA-AHP)",
            "Peligro integrado preliminar: %s." % peligro["valor"],
            "CONFORME", "Sin acción."))
    return filas


def _verificaciones_integracion(integrado):
    """Una verificacion por cada hecho que las dos fuentes no declaran igual."""
    filas = []
    n = 0
    for registro in integrado["campos"]:
        if registro["estado"] not in (ACTUALIZADO, DISCREPANTE):
            continue
        n += 1
        critico = registro["clave"] in CAMPOS_CRITICOS
        if registro["estado"] == ACTUALIZADO:
            calificacion = "CORREGIDO"
            discrepancia = (
                "El libro de resumen V6 consignaba «%s»; la ficha de campo "
                "vigente declara «%s»."
                % (_recorte(registro["valor_alterno"]), _recorte(registro["valor"])))
            tratamiento = (
                "Se adopta el valor de campo por ser la fuente que manda "
                "sobre el dato y la más reciente. El valor anterior queda "
                "declarado para trazabilidad.")
        else:
            calificacion = "SUSTANTIVA" if critico else "NO SUSTANTIVA"
            discrepancia = (
                "La ficha de campo declara «%s» y la fuente %s registra "
                "«%s»." % (_recorte(registro["valor_alterno"]),
                           ETIQUETA_FUENTE[registro["fuente_valor"]].lower(),
                           _recorte(registro["valor"])))
            tratamiento = (
                "Se conservan ambos valores: prevalece el de %s y la "
                "observación de campo queda declarada%s."
                % (ETIQUETA_FUENTE[registro["fuente_valor"]].lower(),
                   ", con revisión obligatoria antes del cierre del perfil"
                   if critico else ""))
        filas.append(_fila_consistencia(
            "I-%02d" % n, registro["etiqueta"], discrepancia, calificacion,
            tratamiento))
    return filas


def _verificaciones_completitud(integrado):
    """Cobertura documental: fichas ausentes y hechos que nadie declara."""
    filas, n = [], 0

    def cod():
        return "F-%02d" % n

    n += 1
    ausentes = integrado.get("fichas_ausentes") or []
    if not integrado.get("tiene_campo"):
        filas.append(_fila_consistencia(
            cod(), "Plantilla DT de campo",
            "El bloque no tiene plantilla DT de campo integrada.",
            "SUSTANTIVA",
            "Debe programarse la verificación de campo o localizarse la "
            "ficha llenada."))
    elif ausentes:
        filas.append(_fila_consistencia(
            cod(), "Fichas F-DT levantadas",
            "La plantilla de campo no trae %d de las 5 fichas: %s."
            % (len(ausentes), ", ".join(ausentes)),
            "SUSTANTIVA",
            "Debe completarse el levantamiento de las fichas faltantes."))
    else:
        filas.append(_fila_consistencia(
            cod(), "Fichas F-DT levantadas",
            "La plantilla de campo trae las cinco fichas F-DT-01 a F-DT-05 "
            "(completitud %.1f %% de los campos esperados)."
            % integrado.get("completitud_campo_pct", 0.0),
            "CONFORME", "Sin acción."))

    criticos_pendientes = [
        r["etiqueta"] for r in integrado["campos"]
        if r["estado"] == PENDIENTE and r["clave"] in CAMPOS_CRITICOS
        and r["clave"] != "peligro_integrado"]
    if criticos_pendientes:
        n += 1
        filas.append(_fila_consistencia(
            cod(), "Hechos críticos sin declarar",
            "No hay dato en ninguna de las dos fuentes para: %s."
            % "; ".join(criticos_pendientes),
            "SUSTANTIVA",
            "Debe completarse antes de elevar el bloque a microlocalización."))

    vacias = integrado.get("tablas_vacias") or []
    if vacias:
        n += 1
        filas.append(_fila_consistencia(
            cod(), "Tablas del formulario sin filas leídas",
            "La plantilla de campo no entrego filas en: %s." % "; ".join(vacias),
            "NO SUSTANTIVA",
            "Se conserva lo que declara la ficha de resumen y no se sustituye "
            "por un conteo en cero. Debe revisarse la maqueta de esas tablas "
            "en la plantilla de campo."))

    n += 1
    pendientes = integrado["conteo_estado"].get(PENDIENTE, 0)
    filas.append(_fila_consistencia(
        cod(), "Cobertura de la integración",
        "Se declararon %d de %d hechos integrados (%.1f %%); %d quedan sin "
        "dato en ninguna fuente."
        % (integrado["n_declarados"], integrado["n_campos"],
           integrado["cobertura_pct"], pendientes),
        "CONFORME" if integrado["cobertura_pct"] >= 70 else "NO SUSTANTIVA",
        "Sin acción." if integrado["cobertura_pct"] >= 70 else
        "Se prioriza el completado de los campos faltantes en la próxima "
        "salida de campo."))
    return filas


# Códigos que este generador recalcula en cada pasada. Las filas H-, en
# cambio, son análisis heredado de la ficha anterior y se vuelven a
# arrastrar: de otro modo, reintegrar un libro ya integrado borraría el
# trabajo redactado a mano.
_RE_CODIGO_GENERADO = re.compile(r"^[GIF]-\d+$", re.IGNORECASE)

# Rótulos de la ficha V6 que nombran una verificación que el control
# regenerado ya cubre con otro rótulo. Sin este puente, la misma
# verificación aparecería dos veces en la hoja actualizada. Se comparan
# contra texto ya normalizado, de modo que van sin acentos.
_EQUIVALENCIAS_V6 = {
    "superficie del poligono": "superficie ndvi vs catalogo",
    "superficie ndvi vs superficie de catalogo": "superficie ndvi vs catalogo",
    "superficie msavi vs catalogo": "superficie msavi vs catalogo",
    "altitud gps del punto de muestreo": "altitud del punto de muestreo",
    "utm del punto de muestreo declarado": "utm del punto de muestreo",
    "codigo de microcuenca declarado": "codigo de microcuenca",
}


def _verificaciones_heredadas(integrado, regeneradas):
    """Verificaciones de la ficha V6 que el cruce automático no reproduce.

    El control regenerado cubre geometría, cruce de fuentes y completitud,
    pero la ficha V6 trae además observaciones redactadas a mano (duplicidad
    de juegos de fichas en el archivo de campo, adscripciones por
    continuidad con bloques vecinos). Se conservan solo las que no repiten
    un campo ya verificado, para no duplicar la misma observación.
    """
    cubiertos = {_comparable(r["campo"]) for r in regeneradas}
    filas, n = [], 0
    for registro in integrado.get("consistencia_v6") or []:
        # Las filas G-, I- y F- las recalcula este mismo generador: si el
        # libro ya fue integrado antes, volver a integrarlo no debe
        # acumularlas como si fueran análisis heredado.
        if _RE_CODIGO_GENERADO.match(_txt(registro.get("codigo"))):
            continue
        campo = _comparable(registro.get("campo"))
        campo = _EQUIVALENCIAS_V6.get(campo, campo)
        if not campo or any(campo in c or c in campo for c in cubiertos):
            continue
        n += 1
        cubiertos.add(campo)
        tratamiento = _txt(registro.get("tratamiento"))
        filas.append(_fila_consistencia(
            "H-%02d" % n, _txt(registro.get("campo")),
            _txt(registro.get("discrepancia")),
            _txt(registro.get("calificacion")) or "NO SUSTANTIVA",
            (tratamiento + " " if tratamiento else "")
            + "(Verificación heredada de la ficha de resumen V6.)"))
    return filas


def control_consistencia(integrado):
    """Hoja 5 regenerada: geometria, integracion de fuentes y completitud.

    A lo regenerado se suman las verificaciones redactadas a mano en la
    ficha V6 que el cruce automatico no cubre, de modo que actualizar el
    libro no pierda el analisis ya hecho.
    """
    regeneradas = (_verificaciones_geometricas(integrado)
                   + _verificaciones_integracion(integrado)
                   + _verificaciones_completitud(integrado))
    return regeneradas + _verificaciones_heredadas(integrado, regeneradas)


def _recorte(texto, largo=140):
    texto = _txt(texto)
    return texto if len(texto) <= largo else texto[:largo - 1] + "…"


def _fmt_num(valor):
    """Numero con la precision que el dato pide: los indices espectrales se
    leen con cuatro decimales y las superficies con dos."""
    if valor is None:
        return "s/d"
    if not isinstance(valor, float):
        return str(valor)
    return ("%.4f" % valor) if abs(valor) < 10 else ("%.2f" % valor)


# ══════════════════════════════════════════════════════════════════════════
# Actualizacion de la ficha de resumen (V6 -> V7)
# ══════════════════════════════════════════════════════════════════════════

# Indice inverso: clave del resumen -> etiquetas normalizadas con que los
# libros la rotulan. Permite localizar la celda a actualizar sin depender de
# como este escrita la etiqueta en cada version del formato.
_ETIQUETAS_POR_CLAVE = {}
for _etiqueta_norm, _clave in rbq.CAMPOS_RESUMEN.items():
    _ETIQUETAS_POR_CLAVE.setdefault(_clave, set()).add(_etiqueta_norm)

_FUENTE_ARIAL = Font(name="Arial", size=9)
_FUENTE_ARIAL_NEGRITA = Font(name="Arial", size=9, bold=True)


def _hoja_por_clave(wb, claves):
    for nombre in wb.sheetnames:
        norm = rbq.normalizar(nombre)
        if any(clave in norm for clave in claves):
            return wb[nombre]
    return None


def _actualizar_pares_resumen(ws, integrado):
    """Escribe en la hoja Resumen los valores que la integracion cambia.

    Solo toca las celdas de valor cuyo hecho cambio: la maqueta, los
    encabezados institucionales y las filas que no cambiaron quedan
    exactamente como estaban.
    """
    pendientes = {}
    for registro in integrado["campos"]:
        clave = registro["resumen"]
        if not clave or registro["estado"] not in (ACTUALIZADO, COMPLEMENTADO):
            continue
        if not _txt(registro["valor"]):
            continue
        for etiqueta in _ETIQUETAS_POR_CLAVE.get(clave, ()):
            pendientes[etiqueta] = registro

    escritos = []
    filas = min(ws.max_row or 0, 200)
    columnas = min(ws.max_column or 0, 12)
    for fila in range(1, filas + 1):
        for col in range(1, columnas):
            etiqueta = rbq.normalizar(ws.cell(fila, col).value)
            registro = pendientes.get(etiqueta) if etiqueta else None
            if not registro:
                continue
            celda = ws.cell(fila, col + 1)
            celda.value = registro["valor"]
            celda.font = _FUENTE_ARIAL
            celda.alignment = Alignment(vertical="top", wrap_text=True)
            escritos.append(registro["clave"])
    return escritos


def _reescribir_consistencia(wb, integrado):
    """Regenera la hoja de control de consistencia con el cruce vigente."""
    ws = _hoja_por_clave(wb, ("control de consistencia", "consistencia"))
    posicion = wb.sheetnames.index(ws.title) if ws is not None else len(wb.sheetnames)
    titulo = ws.title if ws is not None else "Control de consistencia"
    if ws is not None:
        del wb[ws.title]
    ws = wb.create_sheet(titulo, posicion)

    fila = rbq.titulo_hoja(
        ws, "CONTROL DE CONSISTENCIA — BLOQUE %s" % integrado["codigo_bloque"],
        ancho=5)
    celda = ws.cell(fila - 1, 1,
                    "Cruce de la ficha de campo F-DT vigente con la ficha de "
                    "resumen de gabinete y el catálogo maestro. "
                    "Integración %s." % integrado["fecha_integracion"])
    celda.font = Font(name="Arial", size=8, italic=True)
    ws.merge_cells(start_row=fila - 1, start_column=1, end_row=fila - 1,
                   end_column=5)

    registros = integrado["consistencia"]
    fila_cab, fila = rbq.escribir_tabla(
        ws, fila, "DISCREPANCIAS Y VERIFICACIONES",
        ["Cod.", "Campo afectado", "Discrepancia observada", "Calificación",
         "Tratamiento adoptado"],
        [[r["codigo"], r["campo"], r["discrepancia"], r["calificacion"],
          r["tratamiento"]] for r in registros])

    resumen = integrado["consistencia_resumen"]
    ws.cell(fila + 1, 1, "RESUMEN").font = _FUENTE_ARIAL_NEGRITA
    ws.cell(fila + 1, 2, "%d verificaciones" % resumen.get("total", 0)).font =\
        _FUENTE_ARIAL
    ws.cell(fila + 1, 3, " · ".join(
        "%s: %d" % (c, resumen.get(c, 0)) for c in rbq.CALIFICACIONES
        if resumen.get(c, 0))).font = _FUENTE_ARIAL

    nota = ws.cell(fila + 3, 1, NOTA_CONSISTENCIA)
    nota.font = Font(name="Arial", size=7.5, italic=True)
    nota.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=fila + 3, start_column=1, end_row=fila + 3,
                   end_column=5)

    for columna, ancho in zip("ABCDE", (8, 30, 62, 16, 62)):
        ws.column_dimensions[columna].width = ancho
    ws.freeze_panes = ws.cell(fila_cab + 1, 1)
    return ws


NOTA_CONSISTENCIA = (
    "NOTA METODOLÓGICA. Este control se genera del cruce automático entre la "
    "plantilla DT de campo vigente (fichas F-DT-01 a F-DT-05) y la ficha de "
    "resumen de gabinete. Cada hecho territorial se declara una sola vez, "
    "con la fuente que manda sobre el: campo para lo observado en terreno, "
    "gabinete para lo derivado del MDE y de los compuestos Sentinel-2, y "
    "fuente oficial para el catálogo maestro de bloques, la división "
    "política INEI y la codificación de microcuencas de la ANA. Ningún valor "
    "ausente se estima: lo que ninguna fuente declara se consigna como "
    "pendiente. Sistema de referencia UTM WGS 84 Zona 17S (EPSG:32717)."
)


def _hoja_integracion(wb, integrado):
    """Hoja nueva con lo que la integracion cambia y lo que aporta campo."""
    if "Integración de campo" in wb.sheetnames:
        del wb["Integración de campo"]
    ws = wb.create_sheet("Integración de campo")
    fila = rbq.titulo_hoja(
        ws, "INTEGRACIÓN DE LA FICHA DE CAMPO — BLOQUE %s"
        % integrado["codigo_bloque"], ancho=6)

    actualizados = [r for r in integrado["campos"] if r["estado"] == ACTUALIZADO]
    _, fila = rbq.escribir_tabla(
        ws, fila, "A. HECHOS ACTUALIZADOS POR LA VERIFICACIÓN DE CAMPO",
        ["Sección", "Hecho", "Valor en la ficha V6", "Valor de campo vigente",
         "Fuente que manda"],
        [[r["seccion"], r["etiqueta"], r["valor_alterno"], r["valor"],
          ETIQUETA_FUENTE[r["fuente_valor"]]] for r in actualizados]
        or [["—", "Sin cambios: la ficha de campo confirma el libro V6",
             "", "", ""]])
    fila += 1

    solo_campo = [r for r in integrado["campos"]
                  if not r["resumen"] and r["estado"] != PENDIENTE]
    _, fila = rbq.escribir_tabla(
        ws, fila, "B. HECHOS APORTADOS ÚNICAMENTE POR LA FICHA DE CAMPO",
        ["Sección", "Hecho", "Valor", "Unidad"],
        [[r["seccion"], r["etiqueta"], r["valor"], r["unidad"]]
         for r in solo_campo])
    fila += 1

    _, fila = rbq.escribir_tabla(
        ws, fila, "C. PROCEDENCIA DE LOS HECHOS DECLARADOS",
        ["Fuente", "Hechos", "Alcance de la fuente"],
        [[ETIQUETA_FUENTE[f], integrado["conteo_fuente"].get(f, 0),
          DESCRIPCION_FUENTE[f]] for f in FUENTES])
    fila += 1

    _, fila = rbq.escribir_tabla(
        ws, fila, "D. ESTADO DEL CRUCE ENTRE FUENTES",
        ["Estado", "Hechos", "Significado"],
        [[estado, integrado["conteo_estado"].get(estado, 0),
          _GLOSA_ESTADO[estado]] for estado in ESTADOS])

    for columna, ancho in zip("ABCDEF", (34, 38, 44, 44, 24, 18)):
        ws.column_dimensions[columna].width = ancho
    return ws


_GLOSA_ESTADO = {
    CONFORME: "Las dos fuentes declaran lo mismo, o solo la fuente que manda "
              "sobre el hecho lo declara.",
    COMPLEMENTADO: "Solo una de las fuentes lo declara y se adopta ese valor.",
    ACTUALIZADO: "La verificación de campo vigente corrige lo que traía el "
                 "libro de resumen.",
    DISCREPANTE: "Las fuentes se contradicen; prevalece la que manda sobre el "
                 "hecho y la otra queda declarada.",
    PENDIENTE: "Ninguna fuente lo declara; no se estima valor alguno.",
}


def _hoja_registro_campo(wb, integrado):
    """Hoja nueva con las tablas del formulario de campo, sin resumir."""
    if "Registro de campo F-DT" in wb.sheetnames:
        del wb["Registro de campo F-DT"]
    ws = wb.create_sheet("Registro de campo F-DT")
    fila = rbq.titulo_hoja(
        ws, "REGISTRO DE CAMPO F-DT-01 A F-DT-05 — BLOQUE %s"
        % integrado["codigo_bloque"], ancho=9)

    _, fila = rbq.escribir_tabla(
        ws, fila, "F-DT-02 · INVENTARIO DE CÁRCAVAS",
        ["Código", "Tipo", "UTM E ini.", "UTM N ini.", "Longitud (m)",
         "Prof. (m)", "Ancho (m)", "Estado", "Causa"],
        [[c.get("codigo", ""), c.get("tipo", ""), _num(c.get("utm_e_ini")),
          _num(c.get("utm_n_ini")), _num(c.get("longitud_m")),
          _num(c.get("prof_m")), _num(c.get("ancho_m")), c.get("estado", ""),
          c.get("causa", "")] for c in integrado["carcavas"]]
        or [["Sin cárcavas inventariadas en el tramo verificado", "", "", "",
             "", "", "", "", ""]])
    fila += 1

    _, fila = rbq.escribir_tabla(
        ws, fila, "F-DT-03 · ELENCO FLORÍSTICO",
        ["N.o", "Nombre común", "Nombre científico", "Familia", "Estrato",
         "Origen", "Abundancia", "DAP (cm)", "Altura (m)"],
        [[f.get("n", ""), f.get("nombre_comun", ""),
          f.get("nombre_cientifico", ""), f.get("familia", ""),
          f.get("estrato", ""), f.get("origen", ""), f.get("abundancia", ""),
          _num(f.get("dap_cm")), _num(f.get("altura_m"))]
         for f in integrado["floristica"]])
    fila += 1

    _, fila = rbq.escribir_tabla(
        ws, fila, "F-DT-03 · ESPECIES CLAVE GEORREFERENCIADAS",
        ["N.o", "Especie", "Categoría", "Estado UICN", "UTM ESTE",
         "UTM NORTE", "N.° individuos", "Observación"],
        [[e.get("n", ""), e.get("nombre", ""), e.get("categoria", ""),
          e.get("estado_uicn", ""), _num(e.get("utm_e")), _num(e.get("utm_n")),
          _num(e.get("n_indiv")), e.get("observacion", "")]
         for e in integrado["especies_clave"]])
    fila += 1

    _, fila = rbq.escribir_tabla(
        ws, fila, "F-DT-04 · MATRIZ DE CAUSAS DE DEGRADACIÓN",
        ["N.o", "Causa", "Presencia", "Intensidad", "Extensión", "Antigüedad",
         "Evidencia"],
        [[c.get("n", ""), c.get("causa", ""), c.get("presencia", ""),
          c.get("intensidad", ""), c.get("extension", ""),
          c.get("antiguedad", ""), c.get("evidencia", "")]
         for c in integrado["causas"]])
    fila += 1

    _, fila = rbq.escribir_tabla(
        ws, fila, "F-DT-04 · INDICADORES CUANTITATIVOS",
        ["N.o", "Indicador", "Unidad", "Valor", "Fuente", "Umbral", "Nivel"],
        [[i.get("n", ""), i.get("indicador", ""), i.get("unidad", ""),
          i.get("valor", ""), i.get("fuente", ""), i.get("umbral", ""),
          i.get("nivel", "")] for i in integrado["indicadores"]])
    fila += 1

    _, fila = rbq.escribir_tabla(
        ws, fila, "F-DT-05 · FUENTES DE AGUA INVENTARIADAS",
        ["N.o", "Tipo de fuente", "UTM ESTE", "UTM NORTE", "Régimen",
         "Calidad", "Distancia (m)", "Uso / observación"],
        [[f.get("n", ""), f.get("tipo", ""), _num(f.get("utm_e")),
          _num(f.get("utm_n")), f.get("regimen", ""), f.get("calidad", ""),
          _num(f.get("distancia_m")), f.get("uso_obs", "")]
         for f in integrado["fuentes_agua"]])
    fila += 2

    observaciones = [(k, v) for k, v in integrado["observaciones_campo"].items()
                     if v]
    if observaciones:
        _, fila = rbq.escribir_tabla(
            ws, fila, "OBSERVACIONES DEL EVALUADOR POR FICHA",
            ["Ficha", "Observación registrada en campo"],
            [[k, v] for k, v in observaciones])

    for columna, ancho in zip("ABCDEFGHI", (8, 30, 30, 22, 20, 16, 16, 20, 34)):
        ws.column_dimensions[columna].width = ancho
    return ws


def _hoja_graficos_integracion(wb, integrado):
    """Series y graficos nativos de Excel sobre la integracion del bloque."""
    if "Gráficos de integración" in wb.sheetnames:
        del wb["Gráficos de integración"]
    ws = wb.create_sheet("Gráficos de integración")
    fila = rbq.titulo_hoja(
        ws, "GRÁFICOS DE LA INTEGRACIÓN — BLOQUE %s"
        % integrado["codigo_bloque"], ancho=8)

    fila_cab, fila = rbq.escribir_tabla(
        ws, fila, "Hechos declarados por fuente",
        ["Fuente", "Hechos"],
        [[ETIQUETA_FUENTE[f], integrado["conteo_fuente"].get(f, 0)]
         for f in FUENTES])
    libre = rbq.grafico_torta_excel(ws, "Procedencia de los hechos declarados",
                                    fila_cab, len(FUENTES), 1, 2,
                                    "D%d" % fila_cab)
    fila = max(fila, libre) + 2

    fila_cab, fila = rbq.escribir_tabla(
        ws, fila, "Estado del cruce entre fuentes",
        ["Estado", "Hechos"],
        [[e, integrado["conteo_estado"].get(e, 0)] for e in ESTADOS])
    libre = rbq.grafico_barras_excel(ws, "Estado del cruce entre fuentes",
                                     fila_cab, len(ESTADOS), 1, 2,
                                     "D%d" % fila_cab, eje_y="Hechos")
    fila = max(fila, libre) + 2

    estratos = [
        ("Dosel", integrado["por_clave"]["cobertura_dosel"]["valor"]),
        ("Arbustiva", integrado["por_clave"]["cobertura_arbustiva"]["valor"]),
        ("Herbácea", integrado["por_clave"]["cobertura_herbacea"]["valor"]),
        ("Hojarasca", integrado["por_clave"]["cobertura_hojarasca"]["valor"]),
        ("Suelo desnudo", integrado["por_clave"]["suelo_desnudo_pct"]["valor"]),
    ]
    estratos = [(e, _num(v)) for e, v in estratos if _num(v) is not None]
    if estratos:
        fila_cab, fila = rbq.escribir_tabla(
            ws, fila, "Cobertura por estrato — campo (%)",
            ["Estrato", "% del área"], [[e, v] for e, v in estratos])
        libre = rbq.grafico_barras_excel(ws, "Cobertura por estrato (%)",
                                         fila_cab, len(estratos), 1, 2,
                                         "D%d" % fila_cab, eje_y="%")
        fila = max(fila, libre) + 2

    causas = integrado["causas_activas"][:10]
    if causas:
        fila_cab, fila = rbq.escribir_tabla(
            ws, fila, "Causas activas de degradación (F-DT-04)",
            ["Causa", "Intensidad (1 ligera - 4 muy fuerte)"],
            [[c["causa"], c["peso"]] for c in causas])
        libre = rbq.grafico_barras_excel(ws, "Causas activas por intensidad",
                                         fila_cab, len(causas), 1, 2,
                                         "D%d" % fila_cab, eje_y="Intensidad")
        fila = max(fila, libre) + 2

    calificaciones = [(c, integrado["consistencia_resumen"].get(c, 0))
                      for c in rbq.CALIFICACIONES]
    fila_cab, fila = rbq.escribir_tabla(
        ws, fila, "Control de consistencia por calificación",
        ["Calificación", "Verificaciones"],
        [[c, n] for c, n in calificaciones])
    rbq.grafico_torta_excel(ws, "Control de consistencia", fila_cab,
                               len(calificaciones), 1, 2, "D%d" % fila_cab)

    for columna, ancho in zip("AB", (44, 18)):
        ws.column_dimensions[columna].width = ancho
    return ws


def actualizar_libro(contenido_v6, integrado):
    """Ficha de resumen V7: el libro V6 con la ficha de campo integrada.

    Conserva las cinco hojas del formato V6 —solo reescribe los valores que
    la verificacion de campo actualiza y regenera el control de
    consistencia— y agrega tres hojas: la integracion, el registro de campo
    y los graficos. Devuelve los bytes del libro.
    """
    wb = load_workbook(io.BytesIO(contenido_v6))
    hoja_resumen = _hoja_por_clave(wb, ("resumen",))
    escritos = []
    if hoja_resumen is not None:
        escritos = _actualizar_pares_resumen(hoja_resumen, integrado)
        _sellar_version(hoja_resumen, integrado)
    _reescribir_consistencia(wb, integrado)
    _hoja_integracion(wb, integrado)
    _hoja_registro_campo(wb, integrado)
    _hoja_graficos_integracion(wb, integrado)
    salida = io.BytesIO()
    wb.save(salida)
    integrado["campos_reescritos"] = sorted(set(escritos))
    return salida.getvalue()


def _sellar_version(ws, integrado):
    """Deja constancia de la version integrada en el subtitulo de la hoja."""
    for fila in range(1, 12):
        celda = ws.cell(fila, 1)
        texto = _txt(celda.value)
        if texto.lower().startswith("diagn"):
            marca = ("Integrado con la ficha de campo F-DT vigente (%s) — %s"
                     % (integrado["nombre_archivo_campo"] or "plantilla DT",
                        VERSION_INTEGRACION))
            if marca not in texto:
                celda.value = "%s · %s" % (texto, marca)
            return


# ══════════════════════════════════════════════════════════════════════════
# Agregacion por provincia, distrito, microcuenca o bloque
# ══════════════════════════════════════════════════════════════════════════

AGRUPACIONES = {
    "provincia": "Provincia",
    "distrito": "Distrito",
    "microcuenca": "Microcuenca",
    "codigo_bloque": "Bloque",
}


def _promedio(valores):
    valores = [v for v in valores if v is not None]
    return round(sum(valores) / len(valores), 4) if valores else None


def _moda(valores):
    """Valor mas frecuente de una lista de textos; '' si no hay ninguno."""
    conteo = {}
    for valor in valores:
        texto = _txt(valor)
        if texto:
            conteo[texto] = conteo.get(texto, 0) + 1
    if not conteo:
        return ""
    return max(conteo.items(), key=lambda kv: (kv[1], kv[0]))[0]


def agrupar(lista_integrados, agrupacion="provincia"):
    """Resume los bloques integrados por provincia, distrito o microcuenca.

    Devuelve una lista de diccionarios ordenada por el nombre del grupo, con
    los agregados que alimentan las tablas y los graficos de los reportes.
    """
    if agrupacion not in AGRUPACIONES:
        raise ValueError("Agrupación no reconocida: %s" % agrupacion)

    grupos = {}
    for integrado in lista_integrados:
        nombre = _txt(integrado.get(agrupacion)) or "Sin declarar"
        grupos.setdefault(nombre, []).append(integrado)

    filas = []
    for nombre in sorted(grupos):
        bloques = grupos[nombre]
        consistencia = [b.get("consistencia_resumen") or {} for b in bloques]
        filas.append({
            "grupo": nombre,
            "n_bloques": len(bloques),
            "area_ha": round(sum(b.get("area_ha_num") or 0 for b in bloques), 2),
            "brecha_ha": round(sum(b.get("msavi_bajo_umbral_ha") or 0
                                   for b in bloques), 2),
            "msavi_promedio": _promedio([b.get("msavi_2024_num") for b in bloques]),
            "bajo_umbral_pct": _promedio([b.get("bajo_umbral_pct_num")
                                          for b in bloques]),
            "pendiente_pct": _promedio([b.get("pendiente_pct_num") for b in bloques]),
            "cobertura_total_pct": _promedio([b.get("cobertura_total_pct_num")
                                              for b in bloques]),
            "suelo_desnudo_pct": _promedio([b.get("suelo_desnudo_pct_num")
                                            for b in bloques]),
            "n_carcavas": sum(int(b.get("n_carcavas_num") or 0) for b in bloques),
            "n_taxones": sum(int(b.get("n_taxones_num") or 0) for b in bloques),
            "cobertura_integracion_pct": _promedio([b.get("cobertura_pct")
                                                    for b in bloques]),
            "n_actualizados": sum(b["conteo_estado"].get(ACTUALIZADO, 0)
                                  for b in bloques),
            "n_pendientes": sum(b["conteo_estado"].get(PENDIENTE, 0)
                                for b in bloques),
            "n_sustantivas": sum(c.get("SUSTANTIVA", 0) for c in consistencia),
            "n_verificaciones": sum(c.get("total", 0) for c in consistencia),
            "hechos_campo": sum(b["conteo_fuente"].get(CAMPO, 0) for b in bloques),
            "hechos_gabinete": sum(b["conteo_fuente"].get(GABINETE, 0)
                                   for b in bloques),
            "hechos_oficial": sum(b["conteo_fuente"].get(OFICIAL, 0)
                                  for b in bloques),
            "estado_conservacion": _moda([b.get("estado_conservacion")
                                          for b in bloques]),
            "nivel_erosion": _moda([b.get("nivel_erosion") for b in bloques]),
            "urgencia_intervencion": _moda([b.get("urgencia_intervencion")
                                            for b in bloques]),
            "tipo_ecosistema": _moda([b.get("tipo_ecosistema") for b in bloques]),
            "codigos": [b.get("codigo_bloque", "") for b in bloques],
        })
    return filas


def totales(lista_integrados):
    """Totales del universo integrado, para la cabecera de los reportes."""
    consistencia = [b.get("consistencia_resumen") or {} for b in lista_integrados]
    return {
        "n_bloques": len(lista_integrados),
        "area_ha": round(sum(b.get("area_ha_num") or 0
                             for b in lista_integrados), 2),
        "brecha_ha": round(sum(b.get("msavi_bajo_umbral_ha") or 0
                               for b in lista_integrados), 2),
        "msavi_promedio": _promedio([b.get("msavi_2024_num")
                                     for b in lista_integrados]),
        "n_provincias": len({_txt(b.get("provincia")) for b in lista_integrados
                             if _txt(b.get("provincia"))}),
        "n_distritos": len({_txt(b.get("distrito")) for b in lista_integrados
                            if _txt(b.get("distrito"))}),
        "n_microcuencas": len({_txt(b.get("microcuenca")) for b in lista_integrados
                               if _txt(b.get("microcuenca"))}),
        "hechos_campo": sum(b["conteo_fuente"].get(CAMPO, 0)
                            for b in lista_integrados),
        "hechos_gabinete": sum(b["conteo_fuente"].get(GABINETE, 0)
                               for b in lista_integrados),
        "hechos_oficial": sum(b["conteo_fuente"].get(OFICIAL, 0)
                              for b in lista_integrados),
        "n_actualizados": sum(b["conteo_estado"].get(ACTUALIZADO, 0)
                              for b in lista_integrados),
        "n_pendientes": sum(b["conteo_estado"].get(PENDIENTE, 0)
                            for b in lista_integrados),
        "n_sustantivas": sum(c.get("SUSTANTIVA", 0) for c in consistencia),
        "n_verificaciones": sum(c.get("total", 0) for c in consistencia),
        "cobertura_pct": _promedio([b.get("cobertura_pct")
                                    for b in lista_integrados]),
        "bloques_con_campo": sum(1 for b in lista_integrados if b.get("tiene_campo")),
    }


# ══════════════════════════════════════════════════════════════════════════
# Reporte Excel consolidado
# ══════════════════════════════════════════════════════════════════════════

_MIME_XLSX = ("application/vnd.openxmlformats-officedocument"
              ".spreadsheetml.sheet")


def _ajustar_anchos(ws, anchos):
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def generar_excel_consolidado(lista_integrados, agrupacion="provincia"):
    """Libro consolidado de la integracion, con graficos nativos de Excel.

    Trae la sintesis del universo, el agregado por la agrupacion pedida
    (provincia, distrito, microcuenca o bloque), el detalle bloque a bloque,
    los hechos que la verificacion de campo actualizo y el control de
    consistencia completo.
    """
    lista_integrados = sorted(
        lista_integrados, key=lambda b: _orden_codigo(b.get("codigo_bloque", "")))
    resumen = totales(lista_integrados)
    grupos = agrupar(lista_integrados, agrupacion)
    etiqueta = AGRUPACIONES[agrupacion]

    wb = Workbook()
    wb.remove(wb.active)

    # ── Hoja 1: sintesis ──
    ws = wb.create_sheet("Síntesis")
    fila = rbq.titulo_hoja(
        ws, "INTEGRACIÓN DEL DIAGNÓSTICO TERRITORIAL — SÍNTESIS", ancho=6)
    _, fila = rbq.escribir_tabla(
        ws, fila, "UNIVERSO INTEGRADO", ["Indicador", "Valor"],
        [["Bloques integrados", resumen["n_bloques"]],
         ["Bloques con ficha de campo", resumen["bloques_con_campo"]],
         ["Provincias", resumen["n_provincias"]],
         ["Distritos", resumen["n_distritos"]],
         ["Microcuencas", resumen["n_microcuencas"]],
         ["Superficie de catálogo (ha)", resumen["area_ha"]],
         ["Superficie bajo umbral MSAVI %s — brecha (ha)" % rbq.UMBRAL_MSAVI,
          resumen["brecha_ha"]],
         ["MSAVI 2024 promedio", resumen["msavi_promedio"]],
         ["Cobertura promedio de la integración (%)", resumen["cobertura_pct"]],
         ["Hechos actualizados por campo", resumen["n_actualizados"]],
         ["Hechos sin declarar en ninguna fuente", resumen["n_pendientes"]],
         ["Verificaciones de consistencia", resumen["n_verificaciones"]],
         ["Discrepancias sustantivas", resumen["n_sustantivas"]]])
    fila += 1

    fila_cab, fila = rbq.escribir_tabla(
        ws, fila, "PROCEDENCIA DE LOS HECHOS DECLARADOS",
        ["Fuente", "Hechos", "Alcance"],
        [[ETIQUETA_FUENTE[CAMPO], resumen["hechos_campo"],
          DESCRIPCION_FUENTE[CAMPO]],
         [ETIQUETA_FUENTE[GABINETE], resumen["hechos_gabinete"],
          DESCRIPCION_FUENTE[GABINETE]],
         [ETIQUETA_FUENTE[OFICIAL], resumen["hechos_oficial"],
          DESCRIPCION_FUENTE[OFICIAL]]])
    rbq.grafico_torta_excel(ws, "Procedencia de los hechos declarados",
                               fila_cab, 3, 1, 2, "E%d" % fila_cab)
    _ajustar_anchos(ws, [46, 18, 70])

    # ── Hoja 2: agregado por la agrupacion pedida ──
    ws = wb.create_sheet("Por %s" % etiqueta.lower())
    fila = rbq.titulo_hoja(ws, "INTEGRACIÓN POR %s" % etiqueta.upper(), ancho=9)
    cabeceras = [etiqueta, "Bloques", "Superficie (ha)", "Brecha MSAVI (ha)",
                 "MSAVI prom.", "% bajo umbral", "Pendiente prom. (%)",
                 "Cobertura campo (%)", "Suelo desnudo (%)", "Cárcavas",
                 "Taxones", "Hechos de campo", "Hechos actualizados",
                 "Sustantivas", "Estado de conservación dominante",
                 "Urgencia dominante"]
    filas_grupo = [[g["grupo"], g["n_bloques"], g["area_ha"], g["brecha_ha"],
                    g["msavi_promedio"], g["bajo_umbral_pct"],
                    g["pendiente_pct"], g["cobertura_total_pct"],
                    g["suelo_desnudo_pct"], g["n_carcavas"], g["n_taxones"],
                    g["hechos_campo"], g["n_actualizados"], g["n_sustantivas"],
                    g["estado_conservacion"], g["urgencia_intervencion"]]
                   for g in grupos]
    fila_cab, fila = rbq.escribir_tabla(
        ws, fila, "AGREGADO POR %s" % etiqueta.upper(), cabeceras, filas_grupo)
    # Los graficos van en dos filas de dos: la segunda arranca donde
    # termina la primera, no a una distancia fija que podria solaparse.
    n = len(grupos)
    superior = fila + 2
    libre = rbq.grafico_barras_excel(
        ws, "Superficie por %s (ha)" % etiqueta.lower(), fila_cab, n, 1, 3,
        "A%d" % superior, eje_y="ha")
    rbq.grafico_barras_excel(
        ws, "Brecha MSAVI por %s (ha)" % etiqueta.lower(), fila_cab, n, 1, 4,
        "J%d" % superior, eje_y="ha")
    inferior = libre + 1
    rbq.grafico_barras_excel(
        ws, "Bloques por %s" % etiqueta.lower(), fila_cab, n, 1, 2,
        "A%d" % inferior, eje_y="Bloques")
    rbq.grafico_barras_excel(
        ws, "Discrepancias sustantivas por %s" % etiqueta.lower(), fila_cab,
        n, 1, 14, "J%d" % inferior, eje_y="Verificaciones")
    _ajustar_anchos(ws, [30, 10, 16, 18, 13, 14, 18, 18, 17, 11, 11, 16, 18,
                         13, 34, 20])
    ws.freeze_panes = ws.cell(fila_cab + 1, 2)

    # ── Hoja 3: detalle por bloque ──
    ws = wb.create_sheet("Bloques")
    fila = rbq.titulo_hoja(ws, "DETALLE POR BLOQUE PRELIMINAR DE INTERVENCIÓN",
                            ancho=9)
    fila_cab, fila = rbq.escribir_tabla(
        ws, fila, "BLOQUES INTEGRADOS", [
            "Bloque", "Microcuenca", "Provincia", "Distrito",
            "Centro poblado", "Superficie (ha)", "UTM ESTE", "UTM NORTE",
            "Pendiente (%)", "MSAVI 2024", "% bajo umbral",
            "Tipo de ecosistema", "Estado de conservación",
            "Nivel de erosión", "Urgencia", "Modalidad de acceso",
            "Hechos de campo", "Actualizados", "Sustantivas",
            "Cobertura integración (%)", "Ficha de campo"],
        [[b.get("codigo_bloque", ""), b.get("microcuenca", ""),
          b.get("provincia", ""), b.get("distrito", ""),
          b.get("centro_poblado", ""), b.get("area_ha_num"),
          b.get("utm_este_num"), b.get("utm_norte_num"),
          b.get("pendiente_pct_num"), b.get("msavi_2024_num"),
          b.get("bajo_umbral_pct_num"), b.get("tipo_ecosistema", ""),
          b.get("estado_conservacion", ""), b.get("nivel_erosion", ""),
          b.get("urgencia_intervencion", ""), b.get("modalidad_acceso", ""),
          b["conteo_fuente"].get(CAMPO, 0),
          b["conteo_estado"].get(ACTUALIZADO, 0),
          (b.get("consistencia_resumen") or {}).get("SUSTANTIVA", 0),
          b.get("cobertura_pct"), b.get("nombre_archivo_campo", "")]
         for b in lista_integrados])
    _ajustar_anchos(ws, [12, 18, 16, 20, 24, 15, 13, 13, 14, 12, 14, 34, 30,
                         24, 14, 30, 16, 14, 13, 20, 40])
    ws.freeze_panes = ws.cell(fila_cab + 1, 2)

    # ── Hoja 4: hechos actualizados por la verificacion de campo ──
    ws = wb.create_sheet("Hechos actualizados")
    fila = rbq.titulo_hoja(
        ws, "HECHOS ACTUALIZADOS POR LA VERIFICACIÓN DE CAMPO", ancho=7)
    actualizados = [
        [b.get("codigo_bloque", ""), b.get("provincia", ""),
         b.get("distrito", ""), r["seccion"], r["etiqueta"],
         r["valor_alterno"], r["valor"]]
        for b in lista_integrados for r in b["campos"]
        if r["estado"] == ACTUALIZADO]
    fila_cab, fila = rbq.escribir_tabla(
        ws, fila, "DIFERENCIAS ADOPTADAS DESDE LA FICHA DE CAMPO VIGENTE",
        ["Bloque", "Provincia", "Distrito", "Sección", "Hecho",
         "Valor en la ficha V6", "Valor de campo vigente"],
        actualizados or [["—", "", "", "",
                          "La verificación de campo confirma los libros V6",
                          "", ""]])
    _ajustar_anchos(ws, [12, 16, 20, 34, 38, 52, 52])
    ws.freeze_panes = ws.cell(fila_cab + 1, 2)

    # ── Hoja 5: control de consistencia consolidado ──
    ws = wb.create_sheet("Control de consistencia")
    fila = rbq.titulo_hoja(ws, "CONTROL DE CONSISTENCIA CONSOLIDADO", ancho=7)
    fila_cab, fila = rbq.escribir_tabla(
        ws, fila, "VERIFICACIONES POR BLOQUE",
        ["Bloque", "Provincia", "Distrito", "Cod.", "Campo afectado",
         "Discrepancia observada", "Calificación", "Tratamiento adoptado"],
        [[b.get("codigo_bloque", ""), b.get("provincia", ""),
          b.get("distrito", ""), r["codigo"], r["campo"], r["discrepancia"],
          r["calificacion"], r["tratamiento"]]
         for b in lista_integrados for r in b.get("consistencia", [])])
    fila += 1
    conteo = {c: 0 for c in rbq.CALIFICACIONES}
    for b in lista_integrados:
        for r in b.get("consistencia", []):
            if r["calificacion"] in conteo:
                conteo[r["calificacion"]] += 1
    fila_cab_cal, fila = rbq.escribir_tabla(
        ws, fila, "REPARTO POR CALIFICACIÓN", ["Calificación", "Verificaciones"],
        [[c, conteo[c]] for c in rbq.CALIFICACIONES])
    rbq.grafico_torta_excel(ws, "Control de consistencia consolidado",
                               fila_cab_cal, len(rbq.CALIFICACIONES), 1, 2,
                               "D%d" % fila_cab_cal)
    _ajustar_anchos(ws, [12, 16, 20, 8, 30, 64, 16, 64])
    ws.freeze_panes = ws.cell(fila_cab + 1, 2)

    # ── Hoja 6: nota metodologica ──
    ws = wb.create_sheet("Metodología")
    fila = rbq.titulo_hoja(ws, "NOTA METODOLÓGICA DE LA INTEGRACIÓN", ancho=2)
    _, fila = rbq.escribir_tabla(
        ws, fila, "CRITERIO DE INTEGRACIÓN", ["Aspecto", "Criterio adoptado"],
        [["Unidad de integración",
          "El bloque preliminar de intervención. Cada bloque cruza su "
          "plantilla DT de campo con su ficha de resumen de gabinete."],
         ["Regla contra la duplicidad",
          "Cada hecho territorial se declara una sola vez. El cruce recorre "
          "un registro de hechos, no los documentos de origen, de modo que "
          "el mismo dato no puede escribirse dos veces con dos procedencias."],
         ["Fuente que manda sobre el hecho",
          "Campo para lo observado en terreno; gabinete para lo derivado del "
          "MDE y de los compuestos Sentinel-2; fuente oficial para el "
          "catálogo maestro de bloques, la división política INEI y la "
          "codificación de microcuencas de la ANA."],
         ["Tratamiento de las diferencias",
          "Si las fuentes coinciden se declara un solo valor. Si difieren, "
          "prevalece la que manda sobre el hecho y la otra queda declarada "
          "como trazabilidad, con su verificación en el control de "
          "consistencia."],
         ["Valores ausentes",
          "No se estiman. Lo que ninguna fuente declara se consigna como "
          "pendiente y se lista en el control de consistencia."],
         ["Sistema de referencia",
          "UTM WGS 84 Zona 17S (EPSG:32717). Se valida el rango de cada "
          "coordenada antes de usarla."],
         ["Marco normativo",
          "Guía General para la IFE de PI (DGPMI-MEF, 2022), Anexo 2 "
          "GdR-CCC; R.M. N.° 00213-2024-MINAM para el indicador de brecha."],
         ["Versión", VERSION_INTEGRACION],
         ["Fecha de generación",
          datetime.now().strftime("%Y-%m-%d %H:%M:%S")]])
    _ajustar_anchos(ws, [34, 110])

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _orden_codigo(codigo):
    """Orden natural de los codigos: 2 < 10 < M1B1 < M10B4."""
    texto = _txt(codigo)
    if texto.isdigit():
        return (0, int(texto), "")
    m = _RE_MICRO_BLOQUE.match(texto)
    if m:
        return (1, int(m.group(1)) * 1000 + int(m.group(2)), m.group(3) or "")
    return (2, 0, texto)


# ══════════════════════════════════════════════════════════════════════════
# Reportes PDF
# ══════════════════════════════════════════════════════════════════════════

# Colores por calificacion, para que la severidad se lea igual en todos los
# graficos del proyecto.
_COLOR_CALIFICACION = {
    "SUSTANTIVA": (176, 58, 46),
    "NO SUSTANTIVA": (196, 160, 60),
    "CORREGIDO": (58, 124, 79),
    "CONFORME": (27, 77, 46),
}

_COLOR_FUENTE = {
    ETIQUETA_FUENTE[CAMPO].upper(): (27, 77, 46),
    ETIQUETA_FUENTE[GABINETE].upper(): (58, 124, 79),
    ETIQUETA_FUENTE[OFICIAL].upper(): (27, 79, 114),
}


def _pares_seccion(integrado, seccion):
    """Pares etiqueta/valor de una seccion, con la procedencia entre corchetes."""
    pares = []
    for registro in integrado["campos"]:
        if registro["seccion"] != seccion or registro["estado"] == PENDIENTE:
            continue
        valor = registro["valor"]
        if registro["unidad"] and _num(valor) is not None:
            valor = "%s %s" % (valor, registro["unidad"])
        marca = {CAMPO: "campo", GABINETE: "gabinete",
                 OFICIAL: "oficial"}[registro["fuente_valor"]]
        pares.append(("%s [%s]" % (registro["etiqueta"], marca), valor))
    return pares


def generar_pdf_bloque(integrado, incluir_graficos=True):
    """Ficha PDF del bloque integrado: un hecho, un valor, su procedencia."""
    codigo = integrado.get("codigo_bloque", "")
    pdf = rbq.PDFResumen(
        subtitulo="FICHA INTEGRADA DE DIAGNÓSTICO TERRITORIAL - BLOQUE %s"
                  % codigo)
    pdf.alias_nb_pages()
    pdf.add_page()

    resumen_cons = integrado.get("consistencia_resumen") or {}
    pdf.seccion("SÍNTESIS DE LA INTEGRACIÓN")
    pdf.campos([
        ("Bloque", codigo),
        ("Microcuenca", integrado.get("microcuenca", "")),
        ("Provincia", integrado.get("provincia", "")),
        ("Distrito", integrado.get("distrito", "")),
        ("Superficie de catálogo (ha)", _fmt_num(integrado.get("area_ha_num"))),
        ("MSAVI 2024 - media", _fmt_num(integrado.get("msavi_2024_num"))),
        ("Brecha bajo umbral %s (ha)" % rbq.UMBRAL_MSAVI,
         _fmt_num(integrado.get("msavi_bajo_umbral_ha"))),
        ("Hechos declarados", "%d de %d (%.1f %%)"
         % (integrado["n_declarados"], integrado["n_campos"],
            integrado["cobertura_pct"])),
        ("Hechos actualizados por campo",
         integrado["conteo_estado"].get(ACTUALIZADO, 0)),
        ("Verificaciones de consistencia", resumen_cons.get("total", 0)),
        ("Discrepancias sustantivas", resumen_cons.get("SUSTANTIVA", 0)),
        ("Ficha de campo integrada",
         integrado.get("nombre_archivo_campo", "") or "No disponible"),
    ], columnas=2)

    if incluir_graficos:
        pdf.grafico_torta(
            "Procedencia de los hechos declarados",
            [ETIQUETA_FUENTE[f] for f in FUENTES],
            [integrado["conteo_fuente"].get(f, 0) for f in FUENTES],
            colores=_COLOR_FUENTE)
        pdf.grafico_barras(
            "Control de consistencia por calificación",
            list(rbq.CALIFICACIONES),
            [resumen_cons.get(c, 0) for c in rbq.CALIFICACIONES],
            colores=_COLOR_CALIFICACION)

    for seccion in SECCIONES:
        pares = _pares_seccion(integrado, seccion)
        if not pares:
            continue
        pdf.seccion(seccion.upper())
        pdf.campos(pares, columnas=2)

    if incluir_graficos:
        _graficos_campo(pdf, integrado)

    _tablas_campo(pdf, integrado)

    consistencia = integrado.get("consistencia") or []
    if consistencia:
        pdf.seccion("CONTROL DE CONSISTENCIA (%d VERIFICACIONES)"
                    % len(consistencia))
        pdf.tabla(["Cod.", "Campo afectado", "Discrepancia observada",
                   "Calificación", "Tratamiento adoptado"],
                  [[r["codigo"], r["campo"], r["discrepancia"],
                    r["calificacion"], r["tratamiento"]] for r in consistencia],
                  anchos_rel=[0.5, 1.6, 3.4, 1.0, 3.2])

    observaciones = [(k, v) for k, v in
                     (integrado.get("observaciones_campo") or {}).items() if v]
    if observaciones:
        pdf.seccion("OBSERVACIONES DEL EVALUADOR")
        pdf.tabla(["Ficha", "Observación registrada en campo"],
                  [[k, v] for k, v in observaciones], anchos_rel=[0.8, 6])

    pdf.nota(NOTA_CONSISTENCIA)
    return rbq.pdf_bytes(pdf)


def _graficos_campo(pdf, integrado):
    """Graficos que solo la ficha de campo permite construir."""
    por_clave = integrado["por_clave"]
    estratos = [
        ("Dosel", por_clave["cobertura_dosel"]["valor"]),
        ("Arbustiva", por_clave["cobertura_arbustiva"]["valor"]),
        ("Herbácea", por_clave["cobertura_herbacea"]["valor"]),
        ("Hojarasca", por_clave["cobertura_hojarasca"]["valor"]),
        ("Suelo desnudo", por_clave["suelo_desnudo_pct"]["valor"]),
    ]
    estratos = [(e, _num(v)) for e, v in estratos if _num(v) is not None]
    if estratos:
        pdf.grafico_barras("Cobertura por estrato - campo (%)",
                           [e for e, _ in estratos], [v for _, v in estratos],
                           unidad=" %")

    causas = integrado.get("causas_activas") or []
    if causas:
        pdf.grafico_barras(
            "Causas activas de degradación (intensidad 1-4)",
            [c["causa"] for c in causas[:10]],
            [c["peso"] for c in causas[:10]], horizontal=True, alto=52)

    msavi = [c for c in (integrado.get("msavi_tabla") or [])
             if c.get("superficie_ha") is not None]
    if msavi:
        pdf.grafico_torta("Distribución areal del MSAVI 2024 por clase DN (ha)",
                          [c.get("clase", "") for c in msavi],
                          [c.get("superficie_ha") for c in msavi])


def _tablas_campo(pdf, integrado):
    """Tablas del formulario de campo que el resumen no reproduce."""
    if integrado.get("carcavas"):
        pdf.seccion("F-DT-02 - INVENTARIO DE CÁRCAVAS")
        pdf.tabla(["Cod.", "Tipo", "Long. (m)", "Prof. (m)", "Ancho (m)",
                   "Estado", "Causa"],
                  [[c.get("codigo", ""), c.get("tipo", ""),
                    c.get("longitud_m", ""), c.get("prof_m", ""),
                    c.get("ancho_m", ""), c.get("estado", ""),
                    c.get("causa", "")] for c in integrado["carcavas"]],
                  anchos_rel=[0.6, 1.6, 0.8, 0.8, 0.8, 1.2, 1.6])

    if integrado.get("floristica"):
        pdf.seccion("F-DT-03 - ELENCO FLORÍSTICO (%d TAXONES)"
                    % len(integrado["floristica"]))
        pdf.tabla(["N.", "Nombre común", "Nombre científico", "Familia",
                   "Estrato", "Origen", "Abundancia"],
                  [[f.get("n", ""), f.get("nombre_comun", ""),
                    f.get("nombre_cientifico", ""), f.get("familia", ""),
                    f.get("estrato", ""), f.get("origen", ""),
                    f.get("abundancia", "")] for f in integrado["floristica"]],
                  anchos_rel=[0.4, 1.5, 1.7, 1.2, 1.3, 0.9, 1.1])

    if integrado.get("especies_clave"):
        pdf.seccion("F-DT-03 - ESPECIES CLAVE GEORREFERENCIADAS")
        pdf.tabla(["N.", "Especie", "Categoría", "UICN", "UTM ESTE",
                   "UTM NORTE", "Observación"],
                  [[e.get("n", ""), e.get("nombre", ""), e.get("categoria", ""),
                    e.get("estado_uicn", ""), e.get("utm_e", ""),
                    e.get("utm_n", ""), e.get("observacion", "")]
                   for e in integrado["especies_clave"]],
                  anchos_rel=[0.4, 1.9, 1.5, 0.9, 0.9, 0.9, 2.0])

    if integrado.get("causas_activas"):
        pdf.seccion("F-DT-04 - CAUSAS ACTIVAS DE DEGRADACIÓN")
        pdf.tabla(["Causa", "Intensidad", "Extensión", "Antigüedad",
                   "Evidencia"],
                  [[c["causa"], c["intensidad"], c["extension"],
                    c["antiguedad"], c["evidencia"]]
                   for c in integrado["causas_activas"]],
                  anchos_rel=[2.0, 1.0, 1.2, 1.2, 2.2])

    if integrado.get("indicadores"):
        pdf.seccion("F-DT-04 - INDICADORES CUANTITATIVOS")
        pdf.tabla(["N.", "Indicador", "Unidad", "Valor", "Umbral", "Nivel"],
                  [[i.get("n", ""), i.get("indicador", ""), i.get("unidad", ""),
                    i.get("valor", ""), i.get("umbral", ""), i.get("nivel", "")]
                   for i in integrado["indicadores"]],
                  anchos_rel=[0.4, 2.2, 0.8, 0.8, 2.4, 0.9])

    if integrado.get("fuentes_agua"):
        pdf.seccion("F-DT-05 - FUENTES DE AGUA INVENTARIADAS")
        pdf.tabla(["N.", "Tipo de fuente", "UTM ESTE", "UTM NORTE", "Régimen",
                   "Calidad", "Dist. (m)", "Uso / observación"],
                  [[f.get("n", ""), f.get("tipo", ""), f.get("utm_e", ""),
                    f.get("utm_n", ""), f.get("regimen", ""),
                    f.get("calidad", ""), f.get("distancia_m", ""),
                    f.get("uso_obs", "")] for f in integrado["fuentes_agua"]],
                  anchos_rel=[0.4, 1.7, 0.9, 0.9, 1.1, 1.3, 0.7, 2.0])


def generar_pdf_consolidado(lista_integrados, agrupacion="provincia",
                            incluir_graficos=True):
    """Reporte PDF del universo integrado, agregado por el nivel pedido."""
    lista_integrados = sorted(
        lista_integrados, key=lambda b: _orden_codigo(b.get("codigo_bloque", "")))
    resumen = totales(lista_integrados)
    grupos = agrupar(lista_integrados, agrupacion)
    etiqueta = AGRUPACIONES[agrupacion]

    pdf = rbq.PDFResumen(
        subtitulo="INTEGRACIÓN DEL DIAGNÓSTICO TERRITORIAL - REPORTE POR %s"
                  % etiqueta.upper(), orientacion="L")
    pdf.alias_nb_pages()
    pdf.add_page()

    pdf.seccion("UNIVERSO INTEGRADO")
    pdf.campos([
        ("Bloques integrados", resumen["n_bloques"]),
        ("Bloques con ficha de campo", resumen["bloques_con_campo"]),
        ("Provincias", resumen["n_provincias"]),
        ("Distritos", resumen["n_distritos"]),
        ("Microcuencas", resumen["n_microcuencas"]),
        ("Superficie de catálogo (ha)", _fmt_num(resumen["area_ha"])),
        ("Brecha bajo umbral %s (ha)" % rbq.UMBRAL_MSAVI,
         _fmt_num(resumen["brecha_ha"])),
        ("MSAVI 2024 promedio", _fmt_num(resumen["msavi_promedio"])),
        ("Hechos declarados desde campo", resumen["hechos_campo"]),
        ("Hechos declarados desde gabinete", resumen["hechos_gabinete"]),
        ("Hechos declarados desde fuente oficial", resumen["hechos_oficial"]),
        ("Hechos actualizados por campo", resumen["n_actualizados"]),
        ("Verificaciones de consistencia", resumen["n_verificaciones"]),
        ("Discrepancias sustantivas", resumen["n_sustantivas"]),
    ], columnas=3)

    if incluir_graficos:
        pdf.grafico_torta(
            "Procedencia de los hechos declarados",
            [ETIQUETA_FUENTE[CAMPO], ETIQUETA_FUENTE[GABINETE],
             ETIQUETA_FUENTE[OFICIAL]],
            [resumen["hechos_campo"], resumen["hechos_gabinete"],
             resumen["hechos_oficial"]], colores=_COLOR_FUENTE)
        pdf.grafico_barras("Superficie por %s (ha)" % etiqueta.lower(),
                           [g["grupo"] for g in grupos],
                           [g["area_ha"] for g in grupos],
                           unidad=" ha", horizontal=len(grupos) > 6, alto=54)
        pdf.grafico_barras("Brecha MSAVI bajo umbral por %s (ha)"
                           % etiqueta.lower(),
                           [g["grupo"] for g in grupos],
                           [g["brecha_ha"] for g in grupos],
                           unidad=" ha", horizontal=len(grupos) > 6, alto=54)
        pdf.grafico_barras("Discrepancias sustantivas por %s" % etiqueta.lower(),
                           [g["grupo"] for g in grupos],
                           [g["n_sustantivas"] for g in grupos],
                           horizontal=len(grupos) > 6, alto=54)

    pdf.seccion("AGREGADO POR %s" % etiqueta.upper())
    pdf.tabla([etiqueta, "Bloques", "Superficie (ha)", "Brecha (ha)",
               "MSAVI prom.", "% bajo umbral", "Cobertura campo (%)",
               "Cárcavas", "Taxones", "Actualizados", "Sustantivas",
               "Estado de conservación dominante"],
              [[g["grupo"], g["n_bloques"], _fmt_num(g["area_ha"]),
                _fmt_num(g["brecha_ha"]), _fmt_num(g["msavi_promedio"]),
                _fmt_num(g["bajo_umbral_pct"]),
                _fmt_num(g["cobertura_total_pct"]), g["n_carcavas"],
                g["n_taxones"], g["n_actualizados"], g["n_sustantivas"],
                g["estado_conservacion"]] for g in grupos],
              anchos_rel=[1.8, 0.7, 1.0, 0.9, 0.8, 0.9, 1.0, 0.7, 0.7, 0.9,
                          0.9, 2.4],
              alineaciones=["L", "C", "R", "R", "R", "R", "R", "C", "C", "C",
                            "C", "L"])

    pdf.add_page()
    pdf.seccion("DETALLE POR BLOQUE")
    pdf.tabla(["Bloque", "Microcuenca", "Provincia", "Distrito",
               "Superficie (ha)", "MSAVI", "% bajo umbral",
               "Estado de conservación", "Erosión", "Urgencia",
               "Actualizados", "Sustantivas"],
              [[b.get("codigo_bloque", ""), b.get("microcuenca", ""),
                b.get("provincia", ""), b.get("distrito", ""),
                _fmt_num(b.get("area_ha_num")),
                _fmt_num(b.get("msavi_2024_num")),
                _fmt_num(b.get("bajo_umbral_pct_num")),
                b.get("estado_conservacion", ""), b.get("nivel_erosion", ""),
                b.get("urgencia_intervencion", ""),
                b["conteo_estado"].get(ACTUALIZADO, 0),
                (b.get("consistencia_resumen") or {}).get("SUSTANTIVA", 0)]
               for b in lista_integrados],
              anchos_rel=[0.8, 1.3, 1.2, 1.4, 1.0, 0.8, 0.9, 2.0, 1.4, 0.9,
                          1.15, 1.15],
              alineaciones=["L", "L", "L", "L", "R", "R", "R", "L", "L", "L",
                            "C", "C"])

    actualizados = [(b.get("codigo_bloque", ""), r)
                    for b in lista_integrados for r in b["campos"]
                    if r["estado"] == ACTUALIZADO]
    if actualizados:
        pdf.add_page()
        pdf.seccion("HECHOS ACTUALIZADOS POR LA VERIFICACIÓN DE CAMPO (%d)"
                    % len(actualizados))
        pdf.tabla(["Bloque", "Hecho", "Valor en la ficha V6",
                   "Valor de campo vigente"],
                  [[codigo, r["etiqueta"], _recorte(r["valor_alterno"], 190),
                    _recorte(r["valor"], 190)] for codigo, r in actualizados],
                  anchos_rel=[0.7, 2.0, 3.0, 3.0])

    sustantivas = [(b.get("codigo_bloque", ""), r)
                   for b in lista_integrados
                   for r in b.get("consistencia", [])
                   if r["calificacion"] == "SUSTANTIVA"]
    if sustantivas:
        pdf.add_page()
        pdf.seccion("DISCREPANCIAS SUSTANTIVAS PENDIENTES (%d)"
                    % len(sustantivas))
        pdf.tabla(["Bloque", "Campo afectado", "Discrepancia observada",
                   "Tratamiento adoptado"],
                  [[codigo, r["campo"], r["discrepancia"], r["tratamiento"]]
                   for codigo, r in sustantivas],
                  anchos_rel=[0.7, 1.8, 3.4, 3.0])

    pdf.nota(NOTA_CONSISTENCIA)
    return rbq.pdf_bytes(pdf)


# ══════════════════════════════════════════════════════════════════════════
# Integracion en lote
# ══════════════════════════════════════════════════════════════════════════

def parsear_lote_campo(archivos):
    """Parsea varias plantillas de campo. Devuelve (registros, fallidos)."""
    registros, fallidos = [], []
    for nombre, contenido in archivos:
        try:
            registros.append(parsear_ficha_campo(contenido, nombre))
        except Exception as exc:                      # noqa: BLE001
            fallidos.append((nombre, "%s: %s" % (type(exc).__name__, exc)))
    return registros, fallidos


def integrar_lote(resumenes, campos_por_codigo):
    """Integra una coleccion de fichas de resumen con sus fichas de campo.

    `resumenes` es una lista de diccionarios ya parseados y
    `campos_por_codigo` un diccionario {codigo de bloque: ficha de campo}.
    Los bloques que solo existan en campo tambien se integran, para que un
    levantamiento sin ficha de resumen no se pierda.
    """
    integrados, vistos = [], set()
    for resumen in resumenes:
        codigo = normalizar_codigo(resumen.get("codigo_bloque"))
        if not codigo or codigo in vistos:
            continue
        vistos.add(codigo)
        integrados.append(integrar_bloque(
            rbq.completar_sintesis_msavi(resumen),
            campos_por_codigo.get(codigo), codigo))
    for codigo, campo in campos_por_codigo.items():
        codigo = normalizar_codigo(codigo)
        if codigo and codigo not in vistos:
            vistos.add(codigo)
            integrados.append(integrar_bloque({}, campo, codigo))
    integrados.sort(key=lambda b: _orden_codigo(b.get("codigo_bloque", "")))
    return integrados


def actualizar_libros(libros, campos_por_codigo):
    """Regenera las fichas de resumen con la ficha de campo integrada.

    `libros` es [(nombre, contenido_v6)]. Devuelve
    (actualizados, integrados, fallidos), donde `actualizados` es
    [(nombre, contenido_v7)].
    """
    actualizados, integrados, fallidos = [], [], []
    for nombre, contenido in libros:
        try:
            resumen = rbq.completar_sintesis_msavi(
                rbq.parsear_resumen_bloque(contenido, nombre))
            codigo = normalizar_codigo(resumen.get("codigo_bloque")
                                       or rbq.codigo_desde_nombre(nombre))
            integrado = integrar_bloque(resumen, campos_por_codigo.get(codigo),
                                        codigo)
            actualizados.append((nombre, actualizar_libro(contenido, integrado)))
            integrados.append(integrado)
        except Exception as exc:                      # noqa: BLE001
            fallidos.append((nombre, "%s: %s" % (type(exc).__name__, exc)))
    integrados.sort(key=lambda b: _orden_codigo(b.get("codigo_bloque", "")))
    return actualizados, integrados, fallidos


def tabla_integrados(lista_integrados):
    """Filas planas del catalogo integrado, para mostrarlo en el aplicativo."""
    return [{
        "Bloque": b.get("codigo_bloque", ""),
        "Microcuenca": b.get("microcuenca", ""),
        "Provincia": b.get("provincia", ""),
        "Distrito": b.get("distrito", ""),
        "Centro poblado": b.get("centro_poblado", ""),
        "Área (ha)": b.get("area_ha_num"),
        "MSAVI 2024": b.get("msavi_2024_num"),
        "% bajo umbral": b.get("bajo_umbral_pct_num"),
        "Ecosistema": b.get("tipo_ecosistema", ""),
        "Conservación": b.get("estado_conservacion", ""),
        "Erosión": b.get("nivel_erosion", ""),
        "Urgencia": b.get("urgencia_intervencion", ""),
        "Hechos de campo": b["conteo_fuente"].get(CAMPO, 0),
        "Hechos de gabinete": b["conteo_fuente"].get(GABINETE, 0),
        "Hechos oficiales": b["conteo_fuente"].get(OFICIAL, 0),
        "Actualizados": b["conteo_estado"].get(ACTUALIZADO, 0),
        "Pendientes": b["conteo_estado"].get(PENDIENTE, 0),
        "Cobertura (%)": b.get("cobertura_pct"),
        "Sustantivas": (b.get("consistencia_resumen") or {}).get("SUSTANTIVA", 0),
        "Verificaciones": (b.get("consistencia_resumen") or {}).get("total", 0),
        "Ficha de campo": b.get("nombre_archivo_campo", ""),
    } for b in lista_integrados]
