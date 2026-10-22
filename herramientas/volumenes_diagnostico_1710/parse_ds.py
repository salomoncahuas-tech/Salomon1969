"""Lee el libro «Gráficos DS Consolidado» (VAL_PCT) y produce la estructura window.UPD (res, ds, T)."""
import openpyxl, re, json, sys

SHEETS = {'Cobertura': 'cob', 'F-DS-01': 'f01', 'F-DS-02': 'f02', 'F-DS-03': 'f03'}
TSHEETS = {'T Demografía': 'demo', 'T Actividades': 'act', 'T Tenencia': 'ten', 'T Actores': 'actores',
           'T Entrevistas': 'entrev', 'T Control': 'cal', 'T Fichas': 'fichas'}
SEC = re.compile(r'^[A-Z]\. ')

def cellv(c):
    v = c.value
    if isinstance(v, str) and v.startswith('='):
        return '__F__'
    if hasattr(v, 'strftime'):
        return v.strftime('%Y-%m-%d')
    return v

def rows_of(ws):
    out = []
    for r in ws.iter_rows():
        out.append([cellv(c) for c in r])
    return out

def trim(r):
    r = list(r)
    while r and (r[-1] is None or r[-1] == '__F__'):
        r.pop()
    return r

def mode_from_pct_desc(d):
    d = d or ''
    if 'cada columna' in d: return 'col'
    if 'cada fila' in d: return 'row'
    if 'total de la tabla' in d or 'del total' in d: return 'all'
    if 'centro(s) poblado(s) con dato' in d or 'base' in d.lower(): return 'base'
    return None

def parse_ds_sheet(ws):
    R = rows_of(ws)
    i = 0; tabs = []; pct_desc = {}
    # localizar secciones
    starts = [k for k, r in enumerate(R) if isinstance(r[0], str) and SEC.match(r[0]) and all(v is None for v in r[1:])]
    starts.append(len(R))
    for a, b in zip(starts, starts[1:]):
        blk = R[a:b]
        title = blk[0][0].strip()
        if title.endswith('(%)') and not re.search(r'\(%\)$', title.replace(' (%)', '', 1)) is None:
            pass
        desc = blk[1][0] if len(blk) > 1 and blk[1][0] and all(v is None for v in blk[1][1:]) else None
        # cabecera: primera fila tras la descripción con >=2 celdas no nulas
        k = 2 if desc else 1
        while k < len(blk) and sum(v is not None for v in blk[k]) < 2:
            k += 1
        if k >= len(blk):
            continue
        h = trim([v for v in blk[k]])
        ncol = len(h)
        data = []; notes = []; tot = False; base = None; calc = set()
        for r in blk[k + 1:]:
            c0 = r[0]
            if c0 is None and all(v is None for v in r):
                continue
            if isinstance(c0, str) and c0.startswith('Fuente'):
                notes.append(c0); continue
            if isinstance(c0, str) and all(v is None for v in r[1:]) and not data and False:
                continue
            if isinstance(c0, str) and c0.strip().upper() == 'TOTAL':
                tot = True; continue
            if isinstance(c0, str) and c0.startswith('Base:') and len(r) > 1:
                if isinstance(r[1], (int, float)): base = r[1]
                continue
            if isinstance(c0, str) and all(v is None for v in r[1:]) and len(c0) > 60:
                notes.append(c0); continue
            if not str(c0).startswith('PROMEDIO'):
                for j, v in enumerate(r[:ncol]):
                    if v == '__F__': calc.add(j)
            row = [None if v == '__F__' else v for v in r[:ncol]]
            data.append(trim(row) if not str(c0).startswith('PROMEDIO') else trim(row))
        # columnas calculadas también por encabezado
        tabs.append(dict(t=title, d=desc, h=h, r=data, n=notes, tot=tot, base=base, calc=sorted(calc)))
    # tablas gemelas en %: título «X (%)» cuando existe «X»
    titles = {t['t'] for t in tabs}
    pcts = {t['t'][:-4]: t for t in tabs if t['t'].endswith(' (%)') and t['t'][:-4] in titles}
    main = [t for t in tabs if not (t['t'].endswith(' (%)') and t['t'][:-4] in titles)]
    for t in main:
        p = pcts.get(t['t'])
        t['m'] = (mode_from_pct_desc(p['d']) or 'col') if p else None
    return main, pcts

def parse_T(ws):
    R = rows_of(ws)
    k = 5
    while k < len(R) and sum(v is not None for v in R[k]) < 2:
        k += 1
    h = trim(R[k]); n = len(h)
    rows = []
    for r in R[k + 1:]:
        if all(v is None for v in r): continue
        rows.append([None if v == '__F__' else v for v in r[:n]])
    return dict(h=h, r=rows)

def parse_res(ws):
    R = [[cellv(c) for c in r] for r in ws.iter_rows()]
    res = dict(id={}, cab=[], cont=[], av=[])
    sec = None
    for r in R:
        c0 = r[0]
        if c0 is None: continue
        if isinstance(c0, str) and re.match(r'^\d\. ', c0): sec = c0[0]; continue
        if sec == '1':
            if c0 == 'Campo': continue
            res['id'][c0] = r[1]
        elif sec == '2':
            if c0 == 'Indicador': continue
            res['cab'].append([c0, r[1], r[2]])
        elif sec == '3':
            if c0 == 'Sección': continue
            if r[1] is None:
                res['av'].append(c0)
            else:
                res['cont'].append([c0, r[1], r[2]])
    return res

def parse_book(path):
    wb = openpyxl.load_workbook(path)
    out = dict(res=None, ds={}, T={})
    for ws in wb.worksheets:
        nm = ws.title
        if nm.startswith('Resumen'):
            out['res'] = parse_res(ws); continue
        for p, k in SHEETS.items():
            if nm.startswith(p):
                main, pcts = parse_ds_sheet(ws); out['ds'][k] = main
        for p, k in TSHEETS.items():
            if nm.startswith(p):
                out['T'][k] = parse_T(ws)
    return out

if __name__ == '__main__':
    d = parse_book(sys.argv[1])
    json.dump(d, open(sys.argv[2], 'w'), ensure_ascii=False, default=str)
