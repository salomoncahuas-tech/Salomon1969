#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Prueba de `completar_msavi_plantillas.py` sobre un libro que replica la
estructura de las plantillas del Proyecto IN Piura (rotulos tomados literalmente
del libro `Plantilla_Excel_Bloque_38_IN_Piura.xlsx`).

    python test_completar_msavi.py
"""

import os
import shutil
import sys
import tempfile

import openpyxl
from openpyxl.styles import Alignment, Font

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import completar_msavi_plantillas as mod  # noqa: E402

CODIGO = "38"
AREA = 44.2


def construye_libro(ruta):
    libro = openpyxl.Workbook()

    # ---------------- Resumen ----------------
    hoja = libro.active
    hoja.title = "Resumen"
    hoja["A1"] = "AUTORIDAD NACIONAL DE INFRAESTRUCTURA - ANIN"
    hoja.merge_cells("A1:D1")                       # cabecera combinada
    hoja["A2"] = "FICHA RESUMEN — BLOQUE PRELIMINAR DE INTERVENCIÓN 38"
    hoja.merge_cells("A2:D2")
    hoja["A3"] = "1. IDENTIFICACIÓN Y LOCALIZACIÓN"
    hoja["A4"] = "Código del bloque";  hoja["B4"] = CODIGO
    hoja["A5"] = "Superficie de catálogo (V5/V6), ha"; hoja["B5"] = AREA
    hoja["C5"] = "Fuente de superficie"; hoja["D5"] = "Catálogo maestro Bloques V5/V6"
    hoja["A6"] = "3. ÍNDICES DE VEGETACIÓN (SENTINEL-2)"
    hoja["A7"] = "MSAVI 2024 — media del bloque"; hoja["B7"] = 0.6725
    hoja["C7"] = "MSAVI 2024 — clase de la media"; hoja["D7"] = "> 0.6139 (Vigor alto)"
    hoja["A8"] = "Condición frente al umbral 0.4976"; hoja["B8"] = "Sobre umbral"
    hoja["C8"] = "NDVI mediana 2025 — clase modal"; hoja["D8"] = "Vegetación alta (99.30 %)"
    hoja["A9"] = "Superficie clasificada NDVI 2025 (ha)"; hoja["B9"] = 44.212
    hoja["C9"] = "Desviación frente al catálogo (%)"; hoja["D9"] = 0.03
    for fila in range(4, 10):
        hoja.cell(row=fila, column=1).font = Font(name="Arial", bold=True, size=9)

    # ------------- Cobertura MSAVI-NDVI -------------
    cob = libro.create_sheet("Cobertura MSAVI-NDVI")
    cob["A1"] = "DISTRIBUCIÓN AREAL DE CLASES ESPECTRALES — BLOQUE 38"
    cob.merge_cells("A1:E1")
    cob["A2"] = "A. MSAVI 2024 — CLASIFICACIÓN POR UMBRALES DEL PROYECTO"
    cob.merge_cells("A2:E2")
    cob["A3"] = "Clase MSAVI"; cob["B3"] = "Superficie (ha)"
    cob["C3"] = "% del área clasificada"; cob["D3"] = "Interpretación"
    cob["E3"] = "Condición frente al umbral 0.4976"
    filas = [("> 0.6139", "Vigor alto", "Sobre umbral"),
             ("0.4976 - 0.6139", "Vigor moderado", "Sobre umbral"),
             ("0.3813 - 0.4976", "Vigor bajo", "BAJO umbral 0.4976"),
             ("0.2650 - 0.3813", "Vigor muy bajo", "BAJO umbral 0.4976"),
             ("<= 0.2650", "Suelo desnudo / no vegetal", "BAJO umbral 0.4976")]
    for i, (rot, interp, cond) in enumerate(filas, start=4):
        cob.cell(row=i, column=1, value=rot)
        cob.cell(row=i, column=2, value="Por determinar")
        cob.cell(row=i, column=3, value="Por determinar")
        cob.cell(row=i, column=4, value=interp)
        cob.cell(row=i, column=5, value=cond)
    cob["A9"] = "MSAVI 2024 — MEDIA DEL BLOQUE"; cob["B9"] = 0.672487
    cob["A10"] = ("NOTA METODOLÓGICA. La media del MSAVI 2024 del bloque (0.672487) "
                  "procede del catálogo maestro de bloques V5/V6 y es dato oficial. "
                  "La DISTRIBUCIÓN AREAL por clase de MSAVI no está disponible…")
    cob.merge_cells("A10:E10")                       # nota combinada
    cob["A10"].alignment = Alignment(wrap_text=True)

    # ------------- Control de consistencia -------------
    con = libro.create_sheet("Control de consistencia")
    con["A1"] = "CONTROL DE CONSISTENCIA DE LA INFORMACIÓN — BLOQUE 38"
    con.merge_cells("A1:E1")
    con["A2"] = "Cód."; con["B2"] = "Campo afectado"; con["C2"] = "Discrepancia observada"
    con["D2"] = "Calificación"; con["E2"] = "Tratamiento adoptado"
    disc = [("D-01", "Código de microcuenca", "…", "SUSTANTIVA", "…"),
            ("D-02", "UTM del punto de muestreo", "…", "CONFORME", "Sin acción."),
            ("D-03", "Superficie del polígono", "…", "CONFORME", "Sin acción."),
            ("D-04", "Inventario de cárcavas", "…", "NO SUSTANTIVA", "…"),
            ("D-05", "Representatividad", "…", "SUSTANTIVA", "…"),
            ("D-06", "Centro poblado asociado", "…", "CONFORME", "…"),
            ("D-07", "Peligro integrado (MCA-AHP)", "…", "SUSTANTIVA", "…")]
    for i, fila in enumerate(disc, start=3):
        for j, valor in enumerate(fila, start=1):
            con.cell(row=i, column=j, value=valor)
    con["A10"] = "RESUMEN"; con["B10"] = "7 verificaciones"
    con["C10"] = "CONFORME: 3 · NO SUSTANTIVA: 1 · SUSTANTIVA: 3"

    libro.create_sheet("Estaciones fotográficas")["A1"] = "…"
    libro.create_sheet("Microcuenca")["A1"] = "…"
    libro.save(ruta)


def main():
    tmp = tempfile.mkdtemp(prefix="msavi_test_")
    try:
        ruta = os.path.join(tmp, "Plantilla_Excel_Bloque_%s_IN_Piura.xlsx" % CODIGO)
        construye_libro(ruta)

        rc = mod.main(["--dir", tmp, "--solo", CODIGO])
        assert rc == 0, "el script devolvio codigo de salida %s" % rc

        libro = openpyxl.load_workbook(ruta)
        with open(mod.CSV_POR_DEFECTO, encoding="utf-8") as fh:
            import csv as _csv
            datos = {f["BLOQUE"]: f for f in _csv.DictReader(fh, delimiter=";")}[CODIGO]

        # --- hoja de cobertura ---
        cob = libro["Cobertura MSAVI-NDVI"]
        suma_ha = suma_pct = 0.0
        for i, (rot, col_csv, _x) in enumerate(mod.CLASES, start=4):
            assert cob.cell(row=i, column=1).value == rot, "se movio la fila %s" % rot
            ha = cob.cell(row=i, column=2).value
            pct = cob.cell(row=i, column=3).value
            assert isinstance(ha, (int, float)) and isinstance(pct, (int, float)), \
                "la clase %s no quedo numerica: %r / %r" % (rot, ha, pct)
            assert abs(pct - float(datos[col_csv])) < 0.011, \
                "%% distinto en %s: %s vs %s" % (rot, pct, datos[col_csv])
            suma_ha += ha
            suma_pct += pct
        assert abs(suma_ha - AREA) < 0.01, \
            "las hectareas no suman la superficie de catalogo: %.4f vs %.2f" % (suma_ha, AREA)
        assert abs(suma_pct - 100.0) < 0.05, "los porcentajes no suman 100: %.3f" % suma_pct
        assert "conteo de celdas" in cob["A10"].value.lower(), "no se reemplazo la nota"

        # --- hoja Resumen ---
        res = libro["Resumen"]
        pos = mod.busca_fila(res, "Condicion frente al umbral")
        assert pos, "desaparecio 'Condición frente al umbral'"
        texto = res.cell(row=pos[0], column=pos[1] + 1).value
        assert "distribucion areal" in mod.normaliza(texto), texto
        assert mod.busca_fila(res, "MSAVI 2024 - superficie BAJO el umbral"), \
            "no se agrego la fila de superficie bajo umbral"
        assert mod.busca_fila(res, "MSAVI 2024 - clase areal dominante"), \
            "no se agrego la fila de clase areal dominante"
        # las filas siguientes no se pisaron
        assert mod.busca_fila(res, "Superficie clasificada NDVI 2025"), \
            "se perdio la fila de NDVI 2025"
        assert mod.busca_fila(res, "Superficie de catalogo"), "se perdio el area de catalogo"

        # --- hoja Control de consistencia ---
        con = libro["Control de consistencia"]
        pos = mod.busca_fila(con, "D-08", col_max=2)
        assert pos, "no se agrego la verificacion D-08"
        assert con.cell(row=pos[0], column=4).value == "CORREGIDO"
        pos_res = mod.busca_fila(con, "RESUMEN", col_max=2)
        assert con.cell(row=pos_res[0], column=2).value == "8 verificaciones", \
            con.cell(row=pos_res[0], column=2).value
        assert "CORREGIDO: 1" in con.cell(row=pos_res[0], column=3).value
        assert os.path.exists(ruta + ".bak"), "no se dejo respaldo .bak"

        print("\nOK  todas las comprobaciones pasaron "
              "(%.3f ha repartidas, %.2f %% bajo umbral)"
              % (suma_ha, float(datos["PCT_BAJO_UMBRAL_0_4976"])))
        return 0
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    sys.exit(main())
