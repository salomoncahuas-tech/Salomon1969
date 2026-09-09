#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Comprueba, libro por libro, si la distribucion areal del MSAVI 2024 quedo
escrita correctamente en las plantillas Excel del Proyecto IN Piura.

AUTORIDAD NACIONAL DE INFRAESTRUCTURA - ANIN
DIRECCION DE INTERVENCIONES MULTISECTORIALES Y DE EMERGENCIA - DIME
SUBDIRECCION DE ESTUDIOS DE INVERSION - SESDI
Proyecto IN Piura | CUI 2669244 | UTM WGS 84 Zona 17S (EPSG:32717)

Sirve para verificar una carpeta despues de subirla o descargarla de Drive, y
distinguir sin ambiguedad un libro COMPLETADO de uno que sigue en su version
original con «Por determinar».

Comprueba en cada libro:
  1. el paquete .xlsx abre y sus XML estan bien formados;
  2. estan las cinco hojas con sus nombres exactos;
  3. las cinco clases MSAVI tienen superficie y porcentaje NUMERICOS, sin
     «Por determinar», y coinciden con el CSV de referencia;
  4. las cinco clases suman la superficie de catalogo (V5/V6) del propio libro
     y los porcentajes suman 100;
  5. existe la fila TOTAL CLASIFICADO de la seccion A con su formula;
  6. la nota metodologica es la nueva (declara el origen de la medicion);
  7. la hoja «Resumen» trae la distribucion areal y no perdio sus filas;
  8. la hoja «Control de consistencia» trae la verificacion calificada
     CORREGIDO y su fila RESUMEN cuadra con las filas D-xx;
  9. no hay rangos combinados solapados (sintoma de una insercion mal hecha).

Uso:
    python verificar_completado.py --dir <carpeta_con_los_xlsx>
    python verificar_completado.py --dir <carpeta> --detalle
    python verificar_completado.py --dir <carpeta> --csv-informe informe.csv

Codigo de salida 0 si todos los libros estan conformes; 1 en caso contrario.
"""

import argparse
import csv
import glob
import os
import re
import sys
import xml.etree.ElementTree as ET
import zipfile

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("Falta openpyxl.  Instalar con:  pip install openpyxl")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import completar_msavi_plantillas as mod  # noqa: E402

HOJAS = ["Resumen", "Cobertura MSAVI-NDVI", "Estaciones fotográficas",
         "Microcuenca", "Control de consistencia"]

PATRON = re.compile(r"Plantilla_Excel_Bloque_(.+)_IN_Piura\.xlsx$")


def revisa(ruta, datos_csv):
    """Devuelve (estado, [incidencias]).  Estado: OK / INCOMPLETO / INCIDENCIA."""
    fallos = []
    codigo = PATRON.search(os.path.basename(ruta))
    codigo = codigo.group(1) if codigo else None
    if codigo is None:
        return "INCIDENCIA", ["el nombre no sigue Plantilla_Excel_Bloque_<codigo>_IN_Piura.xlsx"]
    datos = datos_csv.get(codigo)

    try:
        archivo = zipfile.ZipFile(ruta)
        danado = archivo.testzip()
        if danado:
            return "INCIDENCIA", ["paquete .xlsx dañado en %s" % danado]
        for interno in archivo.namelist():
            if interno.endswith((".xml", ".rels")):
                ET.fromstring(archivo.read(interno))
    except Exception as exc:                             # noqa: BLE001
        return "INCIDENCIA", ["no se puede abrir el libro: %s" % exc]

    libro = openpyxl.load_workbook(ruta)
    if libro.sheetnames != HOJAS:
        fallos.append("hojas alteradas: %s" % libro.sheetnames)
    if "Cobertura MSAVI-NDVI" not in libro.sheetnames:
        return "INCIDENCIA", fallos or ["falta la hoja 'Cobertura MSAVI-NDVI'"]

    hoja = libro["Cobertura MSAVI-NDVI"]
    area = mod.superficie_catalogo(libro)

    # --- 3 y 4: las cinco clases ---
    incompleto = False
    suma_ha = suma_pct = 0.0
    filas = []
    for rotulo, col_csv, _x in mod.CLASES:
        pos = mod.busca_fila(hoja, rotulo, col_max=1)
        if not pos:
            fallos.append("falta la fila de clase %s" % rotulo)
            continue
        filas.append(pos[0])
        ha = hoja.cell(row=pos[0], column=2).value
        pct = hoja.cell(row=pos[0], column=3).value
        if not isinstance(ha, (int, float)) or not isinstance(pct, (int, float)):
            incompleto = True
            continue
        suma_ha += ha
        suma_pct += pct
        if datos and abs(pct - float(datos[col_csv])) > 0.02:
            fallos.append("%s: %.2f %% no coincide con el CSV (%s)"
                          % (rotulo, pct, datos[col_csv]))
    if incompleto:
        return "INCOMPLETO", ["la seccion A sigue sin metrado (¿es el libro original?)"]
    if filas and filas != list(range(min(filas), min(filas) + len(mod.CLASES))):
        fallos.append("las filas de clase no son contiguas")
    if area and abs(suma_ha - area) > 0.01:
        fallos.append("las clases suman %.4f ha frente a %.3f ha de catálogo"
                      % (suma_ha, area))
    if abs(suma_pct - 100.0) > 0.05:
        fallos.append("los porcentajes suman %.3f" % suma_pct)

    # --- 5: fila TOTAL de la seccion A ---
    if filas:
        siguiente = mod.normaliza(hoja.cell(row=max(filas) + 1, column=1).value)
        if not siguiente.startswith(mod.normaliza("TOTAL CLASIFICADO")):
            fallos.append("falta la fila TOTAL CLASIFICADO de la sección A")
        else:
            esperada = "=SUM(B%d:B%d)" % (min(filas), max(filas))
            if hoja.cell(row=max(filas) + 1, column=2).value != esperada:
                fallos.append("la fórmula del TOTAL no es %s" % esperada)

    # --- 6: nota metodologica ---
    pos = mod.busca_fila(hoja, "NOTA METODOLOGICA")
    if not pos:
        fallos.append("falta la NOTA METODOLÓGICA de la hoja de cobertura")
    elif "conteo de celdas" not in str(hoja.cell(row=pos[0], column=pos[1]).value).lower():
        fallos.append("la NOTA METODOLÓGICA sigue siendo la anterior")
    if mod.busca_fila(hoja, "Por determinar", col_max=3):
        fallos.append("queda «Por determinar» en la hoja de cobertura")

    # --- 7: hoja Resumen ---
    if "Resumen" in libro.sheetnames:
        res = libro["Resumen"]
        pos = mod.busca_fila(res, "Condición frente al umbral")
        if not pos:
            fallos.append("falta «Condición frente al umbral» en Resumen")
        elif "distribucion areal" not in mod.normaliza(res.cell(row=pos[0],
                                                                column=pos[1] + 1).value):
            fallos.append("«Condición frente al umbral» no declara la distribución areal")
        for rotulo in ["MSAVI 2024 — superficie BAJO el umbral",
                       "MSAVI 2024 — clase areal dominante",
                       "Superficie clasificada NDVI 2025",
                       "Superficie de catálogo"]:
            if not mod.busca_fila(res, rotulo):
                fallos.append("falta en Resumen: %s" % rotulo)

    # --- 8: Control de consistencia ---
    if "Control de consistencia" in libro.sheetnames:
        con = libro["Control de consistencia"]
        pos = mod.busca_fila(con, "Distribución areal de clases MSAVI 2024", col_max=2)
        if not pos:
            fallos.append("falta la verificación del MSAVI en Control de consistencia")
        else:
            if con.cell(row=pos[0], column=4).value != "CORREGIDO":
                fallos.append("la verificación del MSAVI no está calificada CORREGIDO")
            if not re.fullmatch(r"D-\d+", str(con.cell(row=pos[0], column=1).value or "").strip()):
                fallos.append("la verificación del MSAVI no tiene correlativo D-xx")
        pos_res = mod.busca_fila(con, "RESUMEN", col_max=1)
        if not pos_res:
            fallos.append("falta la fila RESUMEN en Control de consistencia")
        else:
            cuenta = {"CONFORME": 0, "NO SUSTANTIVA": 0, "SUSTANTIVA": 0, "CORREGIDO": 0}
            total = 0
            for f in range(1, pos_res[0]):
                if not re.fullmatch(r"D-\d+", str(con.cell(row=f, column=1).value or "").strip()):
                    continue
                total += 1
                cal = str(con.cell(row=f, column=4).value or "").strip()
                if cal in cuenta:
                    cuenta[cal] += 1
            m = re.match(r"(\d+) verificaciones",
                         str(con.cell(row=pos_res[0], column=2).value or ""))
            if not m or int(m.group(1)) != total:
                fallos.append("el recuento dice %r y hay %d filas D-xx"
                              % (con.cell(row=pos_res[0], column=2).value, total))
            esperado = ("CONFORME: %d · NO SUSTANTIVA: %d · SUSTANTIVA: %d "
                        "· CORREGIDO: %d" % (cuenta["CONFORME"], cuenta["NO SUSTANTIVA"],
                                                  cuenta["SUSTANTIVA"], cuenta["CORREGIDO"]))
            if str(con.cell(row=pos_res[0], column=3).value) != esperado:
                fallos.append("el desglose del RESUMEN no cuadra con las filas D-xx")

    # --- 9: rangos combinados ---
    for h in libro.worksheets:
        rangos = list(h.merged_cells.ranges)
        for i in range(len(rangos)):
            for j in range(i + 1, len(rangos)):
                a, b = rangos[i], rangos[j]
                if (a.min_row <= b.max_row and b.min_row <= a.max_row
                        and a.min_col <= b.max_col and b.min_col <= a.max_col):
                    fallos.append("%s: rangos combinados solapados %s/%s" % (h.title, a, b))
    libro.close()
    return ("OK" if not fallos else "INCIDENCIA"), fallos


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dir", required=True, help="carpeta con los libros .xlsx")
    ap.add_argument("--csv", default=mod.CSV_POR_DEFECTO,
                    help="CSV de distribución MSAVI de referencia")
    ap.add_argument("--detalle", action="store_true",
                    help="lista también los libros conformes")
    ap.add_argument("--csv-informe", help="escribe el resultado en un CSV")
    args = ap.parse_args(argv)

    with open(args.csv, encoding="utf-8") as fh:
        datos_csv = {f["BLOQUE"]: f for f in csv.DictReader(fh, delimiter=";")}

    rutas = sorted(glob.glob(os.path.join(args.dir, "Plantilla_Excel_Bloque_*.xlsx")))
    if not rutas:
        sys.exit("No se encontró ningún Plantilla_Excel_Bloque_*.xlsx en %s" % args.dir)

    resultados = []
    for ruta in rutas:
        estado, fallos = revisa(ruta, datos_csv)
        codigo = PATRON.search(os.path.basename(ruta))
        resultados.append((codigo.group(1) if codigo else "?", estado, fallos))
        if estado != "OK" or args.detalle:
            marca = {"OK": " OK ", "INCOMPLETO": "SIN ", "INCIDENCIA": " !! "}[estado]
            print("[%s] %-8s %s" % (marca, resultados[-1][0],
                                    "; ".join(fallos) if fallos else "completado y cuadrado"))

    ok = sum(1 for _c, e, _f in resultados if e == "OK")
    sin = sum(1 for _c, e, _f in resultados if e == "INCOMPLETO")
    inc = sum(1 for _c, e, _f in resultados if e == "INCIDENCIA")
    print("\nLibros revisados: %d" % len(resultados))
    print("  completados y cuadrados : %d" % ok)
    print("  sin metrado (originales): %d" % sin)
    print("  con incidencia          : %d" % inc)
    if sin:
        print("\nLos libros «sin metrado» son la versión anterior. Descargue los "
              "completados de in_piura_plantillas/salida/ en el repositorio.")

    faltan = sorted(set(datos_csv) - {c for c, _e, _f in resultados})
    if faltan:
        print("\nBloques del CSV que no aparecen en la carpeta (%d): %s"
              % (len(faltan), ", ".join(faltan)))

    if args.csv_informe:
        with open(args.csv_informe, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(["BLOQUE", "ESTADO", "INCIDENCIAS"])
            for codigo, estado, fallos in resultados:
                w.writerow([codigo, estado, " | ".join(fallos)])
        print("\nInforme escrito en %s" % args.csv_informe)

    return 0 if (ok == len(resultados) and not faltan) else 1


if __name__ == "__main__":
    sys.exit(main())
