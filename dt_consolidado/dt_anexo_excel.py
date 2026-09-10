# -*- coding: utf-8 -*-
"""
Proyecto IN Piura — CUI 2669244 · ANIN / DIME / SESDI
Generador del anexo Excel del Informe Tecnico Consolidado de Diagnostico
Territorial, una version por provincia.

Matrices incluidas:
  1. Matriz Base            7. Sintesis F-DT
  2. Lectura NDVI-MSAVI     8. AdR-CCC preliminar
  3. Integracion ISL         9. Consistencia
  4. Pendientes             10. Microcuencas
  5. Carcavas               11. Estaciones georreferenciadas
  6. Ecosistema y UP
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import dt_data as D

VERDE = "1B4D2E"        # verde institucional ANIN
AZUL = "1B4F72"         # azul institucional ANIN
GRIS = "F2F5F3"         # fila alterna
AMBAR = "FFF3CD"        # celda de atencion
ROJO = "F8D7DA"         # discrepancia sustantiva

FINO = Side(style="thin", color="BFBFBF")
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)

TITULO = "PROYECTO IN PIURA | CUI 2669244 | Recuperación del servicio de regulación de riesgos naturales y de ecosistemas degradados — Cuenca Alta del Río Piura"


def _encabezado(ws, provincia, titulo, subtitulo, ncols):
    """Escribe el encabezado institucional ANIN y devuelve la fila siguiente."""
    filas = [
        ("AUTORIDAD NACIONAL DE INFRAESTRUCTURA — ANIN", 11, True, VERDE),
        ("DIRECCIÓN DE INTERVENCIONES MULTISECTORIALES Y DE EMERGENCIA — DIME", 9, False, VERDE),
        ("SUBDIRECCIÓN DE ESTUDIOS DE INVERSIÓN — SESDI", 9, False, VERDE),
        (TITULO, 8, False, "555555"),
        (f"{titulo} — PROVINCIA DE {provincia.upper()}", 12, True, AZUL),
        (subtitulo, 8, False, "555555"),
    ]
    for i, (txt, size, bold, color) in enumerate(filas, start=1):
        ws.cell(row=i, column=1, value=txt)
        c = ws.cell(row=i, column=1)
        c.font = Font(name="Arial", size=size, bold=bold, color=color)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=max(ncols, 2))
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[6].height = 26
    return 8


def _tabla(ws, fila, cabeceras, filas, anchos=None, notas=None, resaltar=None):
    """Escribe una tabla con estilo ANIN. Devuelve la fila siguiente libre."""
    for j, h in enumerate(cabeceras, start=1):
        c = ws.cell(row=fila, column=j, value=h)
        c.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
    ws.row_dimensions[fila].height = 32

    for i, row in enumerate(filas):
        r = fila + 1 + i
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(name="Arial", size=9)
            c.border = BORDE
            c.alignment = Alignment(
                horizontal="right" if isinstance(v, (int, float)) else "left",
                vertical="center", wrap_text=not isinstance(v, (int, float)))
            if isinstance(v, float):
                c.number_format = "#,##0.00"
            if i % 2 == 1:
                c.fill = PatternFill("solid", fgColor=GRIS)
            if resaltar:
                col = resaltar(row, j - 1)
                if col:
                    c.fill = PatternFill("solid", fgColor=col)

    sig = fila + 1 + len(filas)
    if notas:
        sig += 1
        for nota in notas:
            c = ws.cell(row=sig, column=1, value=nota)
            c.font = Font(name="Arial", size=8, italic=True, color="444444")
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            ws.merge_cells(start_row=sig, start_column=1,
                           end_row=sig, end_column=max(len(cabeceras), 2))
            ws.row_dimensions[sig].height = 14 * (1 + len(nota) // 150)
            sig += 1

    if anchos:
        for j, w in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)
    return sig + 1


def _total(vals):
    return round(sum(v for v in vals if isinstance(v, (int, float))), 2)


# ---------------------------------------------------------------- matrices

def hoja_base(wb, prov, bs):
    ws = wb.create_sheet("1. Matriz Base")
    f = _encabezado(ws, prov, "MATRIZ BASE GEOESPACIAL CONSOLIDADA",
                    "Fuentes: catálogo maestro Bloques V5/V6 · estadística zonal sobre el MDE · "
                    "MSAVI 2024 y NDVI mediana 2025 (Sentinel-2) · fichas F-DT-01 a F-DT-05. "
                    "Sistema de referencia UTM WGS 84 Zona 17S (EPSG:32717).", 18)
    cab = ["Bloque", "Distrito", "Microcuenca", "Zona planif.", "Centro poblado asociado",
           "Superficie (ha)", "Este (m)", "Norte (m)", "Altitud mín. (msnm)",
           "Altitud máx. (msnm)", "Amplitud (m)", "Piso altitudinal dominante",
           "Pendiente prom. (%)", "Clase de pendiente", "Forma del terreno",
           "Posición fisiográfica", "MSAVI 2024", "Tipo de ecosistema (UP)"]
    filas = [[b["codigo"], b["distrito"], b["microcuenca"], b["zona_planificacion"],
              b["centro_poblado"], b["area_ha"], b["este"], b["norte"], b["alt_min"],
              b["alt_max"], b["amplitud"], b["piso"], b["pend_pct"], b["clase_pendiente"],
              b["forma_terreno"], b["posicion"], b["msavi"], b["ecosistema"]] for b in bs]
    cd = D.con_dato(bs, "alt_min")
    filas.append(["TOTAL / RANGO", f"{len({b['distrito'] for b in bs})} distritos",
                  f"{len({b['microcuenca'] for b in bs})} microcuencas", "", "",
                  _total([b["area_ha"] for b in bs]), "", "",
                  min(b["alt_min"] for b in cd) if cd else "s/d",
                  max(b["alt_max"] for b in cd) if cd else "s/d", "", "",
                  round(D.ponderada(bs, "pend_pct"), 2), "", "", "",
                  round(D.ponderada(bs, "msavi"), 4), ""])
    _tabla(ws, f, cab, filas,
           anchos=[10, 16, 14, 10, 34, 12, 11, 12, 11, 11, 10, 22, 12, 20, 16, 16, 10, 34],
           notas=["Las coordenadas del centroide se expresan en UTM WGS 84 Zona 17S. Todas validan en el rango del proyecto (Este 450 000–750 000 m; Norte 9 300 000–9 600 000 m).",
                  "La altitud y la pendiente proceden de la estadística zonal sobre el modelo digital de elevación, no de medición de campo. La pendiente se expresa en porcentaje.",
                  "PRINCIPIO DE NO INVENCIÓN: ningún valor ausente ha sido estimado o imputado. Los campos no sustentados en observación de campo, estadística zonal o catálogo oficial se consignan con la marca de origen «Por determinar» o «Por verificar»."])
    return ws


def hoja_vegetacion(wb, prov, bs):
    ws = wb.create_sheet("2. NDVI-MSAVI")
    f = _encabezado(ws, prov, "MATRIZ DE LECTURA INTEGRADA DE ÍNDICES DE VEGETACIÓN",
                    "MSAVI 2024 y NDVI mediana 2025 (Sentinel-2, mosaico Piura). "
                    "Umbral de brecha del proyecto: MSAVI = 0.4976. Marco del indicador de brecha: R.M. N.° 00213-2024-MINAM.", 17)
    cab = ["Bloque", "Distrito", "Superficie catálogo (ha)", "MSAVI 2024 (media)",
           "Clase de la media", "Condición vs. umbral 0.4976", "Clase DN dominante",
           "Sup. clasificada MSAVI (ha)", "Desv. MSAVI vs. catálogo (%)",
           "Sup. BAJO umbral (ha)", "% bajo umbral (brecha espectral)",
           "NDVI 2025 — clase modal", "NDVI Veg. alta (ha)", "NDVI Veg. mediana (ha)",
           "NDVI Veg. ligera (ha)", "NDVI Tierra desnuda (ha)", "Sup. clasificada NDVI (ha)"]
    filas = []
    for b in bs:
        n = b["ndvi_clases"]
        filas.append([
            b["codigo"], b["distrito"], b["area_ha"], b["msavi"], b["msavi_clase"],
            b["msavi_condicion"], b["dn_dominante"], b["sup_msavi"], b["desv_msavi"],
            b["ha_bajo_umbral"], b["pct_bajo_umbral"], b["ndvi_modal"],
            n.get("Vegetación alta", (None,))[0], n.get("Vegetación mediana", (None,))[0],
            n.get("Vegetación ligera", (None,))[0], n.get("Tierra desnuda", (None,))[0],
            b["sup_ndvi"]])
    ha = sum(b["area_ha"] for b in bs)
    brecha = _total([b["ha_bajo_umbral"] for b in bs])
    filas.append(["TOTAL / MEDIA", "", round(ha, 2),
                  round(sum(b["msavi"] for b in bs) / len(bs), 4), "", "", "",
                  _total([b["sup_msavi"] for b in bs]), "", brecha,
                  round(100 * brecha / ha, 2), "",
                  _total([b["ndvi_clases"].get("Vegetación alta", (0,))[0] or 0 for b in bs]),
                  _total([b["ndvi_clases"].get("Vegetación mediana", (0,))[0] or 0 for b in bs]),
                  _total([b["ndvi_clases"].get("Vegetación ligera", (0,))[0] or 0 for b in bs]),
                  _total([b["ndvi_clases"].get("Tierra desnuda", (0,))[0] or 0 for b in bs]),
                  _total([b["sup_ndvi"] for b in bs])])

    def marca(row, j):
        if j == 10 and isinstance(row[10], (int, float)) and row[10] >= 50:
            return AMBAR
        return None

    _tabla(ws, f, cab, filas,
           anchos=[10, 16, 13, 11, 24, 16, 30, 14, 14, 13, 14, 24, 13, 13, 13, 13, 14],
           notas=["LECTURA CRÍTICA. Los índices MSAVI y NDVI miden vigor y densidad de biomasa, no composición ni integridad ecosistémica. Un valor alto NO equivale a ausencia de degradación: en bloques con mosaico agrícola, pastizal cultivado o plantaciones, la respuesta espectral puede ser alta sobre una unidad productora sustituida en su composición.",
                  "La diferencia entre ambos productos NO debe leerse como mejora entre 2024 y 2025: son índices distintos (el NDVI satura ante biomasa densa, el MSAVI conserva sensibilidad al suelo de fondo) aplicados sobre compuestos temporales distintos.",
                  "ADVERTENCIA DE ESTACIONALIDAD. En los bloques de Bosque Estacionalmente Seco la caducifolia produce divergencia entre los compuestos de estiaje y las tomas de campo de temporada húmeda. Toda meta física dimensionada sobre índices de estiaje debe tratarse como cota superior.",
                  "Sombreado ámbar: bloques con 50 % o más de su superficie bajo el umbral MSAVI 0.4976 (brecha espectral mayoritaria).",
                  "La equivalencia entre el valor de clase DN del ráster MSAVI clasificado y los umbrales del proyecto fue establecida por contraste con la media MSAVI del catálogo maestro de los 117 bloques (r = 0.9965; error absoluto medio = 0.0091) y debe ser confirmada por el especialista SIG del proyecto."],
           resaltar=marca)
    return ws


def hoja_isl(wb, prov, bs):
    ws = wb.create_sheet("3. Integración ISL Geología")
    meta = D.PROV_META[prov]
    f = _encabezado(ws, prov, "INTEGRACIÓN CON EL ÍNDICE DE SUSCEPTIBILIDAD LITOLÓGICA (ISL)",
                    f"Fuente: {meta['volumen_geologia']}. ISL-MM (movimientos en masa) e "
                    "ISL-EH (erosión hídrica), escala 1 a 5 sobre la Carta Geológica Nacional "
                    "1:50 000 — INGEMMET. Clases: < 1.50 Muy baja · 1.50–2.49 Baja · "
                    "2.50–3.49 Media · 3.50–4.49 Alta · ≥ 4.50 Muy alta.", 14)
    cab = ["Bloque", "Distrito", "Microcuenca", "Superficie (ha)", "N.° unidades geológicas",
           "Unidad geológica dominante", "ISL-MM", "Clase MM", "ISL-EH", "Clase EH",
           "Pendiente prom. (%)", "MSAVI 2024", "% bajo umbral MSAVI",
           "Lectura integrada susceptibilidad–cobertura"]
    filas = []
    for b in bs:
        filas.append([b["codigo"], b["distrito"], b["microcuenca"], b["area_ha"],
                      b["geo_n_unidades"], b["geo_unidad"], b["isl_mm"], b["isl_clase_mm"],
                      b["isl_eh"], b["isl_clase_eh"], b["pend_pct"], b["msavi"],
                      b["pct_bajo_umbral"], _lectura_isl(b)])
    ha = sum(b["area_ha"] for b in bs)
    filas.append(["TOTAL / MEDIA PONDERADA", "", "", round(ha, 2), "", "",
                  round(D.ponderada(bs, "isl_mm"), 2), "",
                  round(D.ponderada(bs, "isl_eh"), 2), "",
                  round(D.ponderada(bs, "pend_pct"), 2),
                  round(D.ponderada(bs, "msavi"), 4), "", ""])

    def marca(row, j):
        if j in (7, 9) and isinstance(row[j], str) and row[j] == "Muy alta":
            return ROJO
        if j in (7, 9) and isinstance(row[j], str) and row[j] == "Alta":
            return AMBAR
        return None

    f = _tabla(ws, f, cab, filas,
               anchos=[10, 16, 14, 12, 12, 34, 9, 12, 9, 12, 12, 10, 12, 46],
               notas=["ALCANCE DEL ISL. El ISL es una jerarquización preliminar basada EXCLUSIVAMENTE en la competencia mecánica, la anisotropía estructural, el perfil de meteorización y el estado de consolidación del sustrato. NO constituye una evaluación de peligro ni de riesgo: no integra pendiente, precipitación detonante, cobertura vegetal ni condiciones hidrogeológicas.",
                      "Su integración con la pendiente, la precipitación detonante y la cobertura vegetal —para obtener el peligro propiamente dicho— corresponde al Entregable 8 (Análisis de riesgos para el proyecto). La columna «Lectura integrada» de esta matriz es una aproximación de gabinete que NO sustituye dicho modelamiento.",
                      "La verificación de campo y los ensayos geotécnicos se ejecutarán en la fase de estudio definitivo o expediente técnico.",
                      "Sombreado rojo: susceptibilidad Muy alta. Sombreado ámbar: susceptibilidad Alta."],
               resaltar=marca)

    # Agregado por clase de susceptibilidad
    for etiqueta, campo_clase, campo_val in (("ISL-MM (movimientos en masa)", "isl_clase_mm", "isl_mm"),
                                             ("ISL-EH (erosión hídrica)", "isl_clase_eh", "isl_eh")):
        ws.cell(row=f, column=1, value=f"AGREGADO POR CLASE DE SUSCEPTIBILIDAD — {etiqueta}")
        ws.cell(row=f, column=1).font = Font(name="Arial", size=10, bold=True, color=AZUL)
        f += 1
        agg = {}
        for b in bs:
            k = b[campo_clase]
            a = agg.setdefault(k, {"n": 0, "ha": 0.0})
            a["n"] += 1
            a["ha"] += b["area_ha"]
        orden = ["Muy alta", "Alta", "Media", "Baja", "Muy baja"]
        fl = [[k, agg[k]["n"], round(agg[k]["ha"], 2), round(100 * agg[k]["ha"] / ha, 2)]
              for k in orden if k in agg]
        fl.append(["TOTAL", len(bs), round(ha, 2), 100.0])
        f = _tabla(ws, f, ["Clase de susceptibilidad", "N.° de bloques", "Superficie (ha)",
                           "% de la superficie provincial"], fl, anchos=None)
    return ws


def _lectura_isl(b):
    """Aproximacion de gabinete que cruza susceptibilidad litologica y cobertura."""
    alta = b["isl_clase_mm"] in ("Alta", "Muy alta") or b["isl_clase_eh"] in ("Alta", "Muy alta")
    brecha = (b["pct_bajo_umbral"] or 0) >= 50
    pend = (b["pend_pct"] or 0) >= 25
    if alta and brecha and pend:
        return ("Convergencia de sustrato susceptible, cobertura bajo umbral y pendiente "
                "≥ 25 %: prioridad máxima para control de erosión y estabilización.")
    if alta and brecha:
        return ("Sustrato susceptible con cobertura mayoritariamente bajo umbral: "
                "la cobertura no compensa la susceptibilidad del sustrato.")
    if alta and pend:
        return ("Sustrato susceptible sobre pendiente ≥ 25 % con cobertura sobre umbral: "
                "la cobertura actual cumple función protectora; prioridad de conservación.")
    if alta:
        return ("Sustrato susceptible atenuado por cobertura sobre umbral y pendiente "
                "moderada: conservación del dosel funcional.")
    if brecha:
        return ("Sustrato de susceptibilidad media o baja con cobertura bajo umbral: "
                "la brecha es principalmente de cobertura, no de sustrato.")
    return "Sustrato de susceptibilidad media o baja con cobertura sobre umbral."


def hoja_pendientes(wb, prov, bs):
    ws = wb.create_sheet("4. Pendientes")
    f = _encabezado(ws, prov, "PENDIENTE PROMEDIO POR BLOQUE",
                    "Estadística zonal sobre el modelo digital de elevación. La pendiente se "
                    "reporta en porcentaje y en grados. Criterio de idoneidad del proyecto: "
                    "pendiente máxima 75 %.", 12)
    cab = ["Bloque", "Distrito", "Superficie (ha)",
           "LECTURA A — valor del catálogo como % (rótulo de las plantillas V6)",
           "LECTURA A — equivalente en grados",
           "LECTURA B — valor del catálogo como grados (Informe DT de Frías)",
           "LECTURA B — equivalente en %", "Clase de pendiente (lectura A)",
           "Rango declarado en campo", "¿Coincide gabinete–campo?",
           "Altitud mín. (msnm)", "Altitud máx. (msnm)", "Amplitud (m)",
           "Forma predominante del terreno"]
    filas = []
    for b in bs:
        coincide = ("s/d" if b["pend_pct"] is None else
                    "Sí" if D.norm(b["clase_pendiente"]) == D.norm(b["pendiente_campo"]) else "No")
        filas.append([b["codigo"], b["distrito"], b["area_ha"], b["pend_pct"], b["pend_grados"],
                      b["pend_valor_catalogo"], b["pend_si_grados_pct"],
                      b["clase_pendiente"], b["pendiente_campo"], coincide,
                      b["alt_min"], b["alt_max"], b["amplitud"], b["forma_terreno"]])
    ha = sum(b["area_ha"] for b in bs)
    cd = D.con_dato(bs, "alt_min")
    filas.append(["TOTAL / MEDIA PONDERADA", "", round(ha, 2),
                  round(D.ponderada(bs, "pend_pct"), 2),
                  round(D.ponderada(bs, "pend_grados"), 2),
                  round(D.ponderada(bs, "pend_valor_catalogo"), 2),
                  round(D.ponderada(bs, "pend_si_grados_pct"), 2), "", "", "",
                  min(b["alt_min"] for b in cd) if cd else "s/d",
                  max(b["alt_max"] for b in cd) if cd else "s/d",
                  round(D.ponderada(bs, "amplitud"), 1), ""])

    def marca(row, j):
        if j == 9 and row[9] == "No":
            return AMBAR
        if j in (3, 6) and isinstance(row[j], (int, float)) and row[j] > 75:
            return ROJO
        if row[7] == "Sin cobertura del MDE" and j in (3, 4, 5, 6, 10, 11, 12):
            return ROJO
        return None

    f = _tabla(ws, f, cab, filas,
               anchos=[10, 16, 12, 20, 16, 20, 16, 22, 22, 13, 11, 11, 10, 16],
               notas=["DISCREPANCIA D-P01 — UNIDAD DE LA PENDIENTE DEL CATÁLOGO. Esta hoja publica las DOS lecturas del mismo valor porque la unidad del campo no está resuelta. LECTURA A: las 117 plantillas de resumen V6 rotulan el valor del catálogo como PORCENTAJE y derivan los grados como arcotangente. LECTURA B: el Informe Técnico Consolidado de Diagnóstico Territorial del distrito de Frías concluye que ese mismo valor está en GRADOS, tras contrastarlo con la capa «Pendientes vector» (error cuadrático medio 1,14° en grados frente a 12,70° como porcentaje; r = +0,956).",
                      "Se ha verificado que el valor numérico es IDÉNTICO en ambas fuentes: el conflicto es de unidad, no de dato. No se resuelve en este anexo. Requiere pronunciamiento formal de SESDI antes de fijar metas físicas de movimiento de tierras.",
                      "Bajo AMBAS lecturas ningún bloque supera el criterio de idoneidad de pendiente máxima del proyecto (75 %): la elegibilidad no está en cuestión; sí lo está el dimensionamiento de partidas.",
                      "DISCREPANCIA D-M01 — Sombreado rojo en las columnas de altitud y pendiente: bloques sin cobertura del modelo digital de elevación. El reporte de estadística zonal devuelve 0 msnm y 0 % de pendiente, que son marcadores de ausencia y no mediciones. Los valores se anulan y se excluyen de todo promedio, mínimo y máximo.",
                      "Sombreado ámbar: discrepancia entre la clase de pendiente derivada de la estadística zonal y el rango declarado por la brigada en la ficha F-DT-01. No invalida ninguno de los dos registros: el valor zonal es de bloque y el de campo es de traza.",
                      "Las medias del pie de cuadro se ponderan por la superficie de cada bloque y se calculan solo sobre los bloques con dato."],
               resaltar=marca)

    # Distribucion por clase de pendiente
    ws.cell(row=f, column=1, value="DISTRIBUCIÓN DE BLOQUES POR CLASE DE PENDIENTE (ESTADÍSTICA ZONAL)")
    ws.cell(row=f, column=1).font = Font(name="Arial", size=10, bold=True, color=AZUL)
    f += 1
    agg = {}
    for b in bs:
        a = agg.setdefault(b["clase_pendiente"], {"n": 0, "ha": 0.0})
        a["n"] += 1
        a["ha"] += b["area_ha"]
    fl = [[k, v["n"], round(v["ha"], 2), round(100 * v["ha"] / ha, 2)]
          for k, v in sorted(agg.items(), key=lambda kv: -kv[1]["ha"])]
    fl.append(["TOTAL", len(bs), round(ha, 2), 100.0])
    _tabla(ws, f, ["Clase de pendiente (lectura A)", "N.° de bloques", "Superficie (ha)",
                   "% de la superficie provincial"], fl,
           notas=["La clase corresponde a la LECTURA A (valor del catálogo como porcentaje). "
                  "Bajo la lectura B la clasificación de los bloques cambia; véase la "
                  "discrepancia D-P01 en las notas del cuadro superior."])
    return ws


def hoja_carcavas(wb, prov, bs):
    ws = wb.create_sheet("5. Cárcavas")
    f = _encabezado(ws, prov, "INVENTARIO DE CÁRCAVAS CODIFICADAS",
                    "Fuente: capa vectorial «Cárcavas codificadas» y su caracterización "
                    "morfométrica. Coordenadas en UTM WGS 84 Zona 17S (EPSG:32717).", 14)
    cab = ["Código", "Bloque", "Distrito", "Microcuenca", "Este (m)", "Norte (m)",
           "Longitud (m)", "Alt. mín. (msnm)", "Alt. máx. (msnm)", "Alt. prom. (msnm)",
           "Pendiente prom. (%)", "Índice de irregularidad", "NDVI", "Clase morfológica"]
    filas = []
    for b in bs:
        for c in b["carcavas"]:
            filas.append([c["codigo"], b["codigo"], b["distrito"], b["microcuenca"],
                          c["este"], c["norte"], c["longitud"], c["alt_min"], c["alt_max"],
                          c["alt_prom"], c["pend_prom"], c["iirreg"], c["ndvi"],
                          c["clase_morf"]])
    if filas:
        filas.sort(key=lambda r: -r[6])
        filas.append(["TOTAL", f"{len({r[1] for r in filas})} bloques", "", "", "", "",
                      _total([r[6] for r in filas]), "", "", "", "", "", "", ""])
        _tabla(ws, f, cab, filas,
               anchos=[14, 10, 16, 14, 11, 12, 11, 11, 11, 11, 12, 12, 8, 20],
               notas=["El inventario recoge únicamente las cárcavas efectivamente digitalizadas y codificadas. LA AUSENCIA DE REGISTROS EN UN BLOQUE NO EQUIVALE A AUSENCIA VERIFICADA DE CÁRCAVAS: el levantamiento instrumental está pendiente para el resto de los bloques y es requisito previo a la fijación de metas físicas de infraestructura gris.",
                      "El NDVI consignado corresponde al valor sobre el eje de la cárcava y permite discriminar los rasgos activos (suelo expuesto) de los estabilizados por vegetación.",
                      "La atribución distrital y provincial de la capa de cárcavas se contrasta con la del catálogo maestro de bloques en la hoja «9. Consistencia»."])
    else:
        ws.cell(row=f, column=1, value=(
            "No se registran cárcavas digitalizadas y codificadas en los bloques de esta "
            "provincia dentro de la capa vectorial disponible. La ausencia de registros NO "
            "equivale a ausencia verificada de cárcavas: el levantamiento instrumental "
            "permanece pendiente."))
        ws.cell(row=f, column=1).font = Font(name="Arial", size=9, italic=True)
    return ws


def hoja_ecosistema(wb, prov, bs):
    ws = wb.create_sheet("6. Ecosistema y UP")
    f = _encabezado(ws, prov, "ECOSISTEMA, ESTADO DE CONSERVACIÓN Y UNIDAD PRODUCTORA",
                    "La Unidad Productora del proyecto es el ecosistema en su totalidad. "
                    "Marco del indicador de brecha: R.M. N.° 00213-2024-MINAM.", 13)
    cab = ["Bloque", "Distrito", "Superficie (ha)", "Tipo de ecosistema (UP)",
           "Estado de conservación", "Uso actual dominante", "Tipo de cobertura dominante",
           "Cobertura vegetal campo (%)", "Suelo desnudo campo (%)", "Regeneración natural",
           "Nivel general de erosión", "N.° de taxones registrados", "Estado sanitario"]
    filas = [[b["codigo"], b["distrito"], b["area_ha"], b["ecosistema"],
              b["estado_conservacion"], b["uso_actual"], b["cobertura_tipo"],
              b["cobertura_pct"], b["suelo_desnudo"], b["regeneracion"], b["erosion"],
              b["taxones"], b["estado_sanitario"]] for b in bs]
    f = _tabla(ws, f, cab, filas,
               anchos=[10, 16, 12, 40, 26, 20, 20, 12, 12, 16, 22, 11, 16],
               notas=["Las celdas «Por determinar» corresponden a campos que la ficha F-DT no consignó. No han sido imputados ni estimados.",
                      "El elenco florístico registrado es preliminar y en la mayoría de los bloques no alcanza el mínimo de 15 taxones que exige el formato V5; no constituye inventario botánico formal."])
    ha = sum(b["area_ha"] for b in bs)
    for etiqueta, campo in (("TIPO DE ECOSISTEMA (UNIDAD PRODUCTORA)", "ecosistema"),
                            ("ESTADO DE CONSERVACIÓN", "estado_conservacion"),
                            ("NIVEL GENERAL DE EROSIÓN", "erosion")):
        ws.cell(row=f, column=1, value=f"AGREGADO POR {etiqueta}")
        ws.cell(row=f, column=1).font = Font(name="Arial", size=10, bold=True, color=AZUL)
        f += 1
        fl = [[k, v["n"], round(v["ha"], 2), round(100 * v["ha"] / ha, 2)]
              for k, v in D.conteo(bs, campo)]
        fl.append(["TOTAL", len(bs), round(ha, 2), 100.0])
        f = _tabla(ws, f, [etiqueta.capitalize(), "N.° de bloques", "Superficie (ha)",
                           "% de la superficie provincial"], fl, anchos=[46, 14, 14, 16])
    return ws


def hoja_sintesis(wb, prov, bs):
    ws = wb.create_sheet("7. Síntesis F-DT")
    f = _encabezado(ws, prov, "SÍNTESIS DE LAS FICHAS F-DT-01 A F-DT-05 POR BLOQUE",
                    "Matriz comparativa transversal de las cinco fichas del Diagnóstico "
                    "Territorial (Plantilla DT Campo Check Validada V5).", 16)
    cab = ["Bloque", "Superficie (ha)", "F-DT-01 Forma", "F-DT-01 Posición",
           "F-DT-01 Exposición", "F-DT-01 Afloramientos", "F-DT-01 Escarpes activos",
           "F-DT-01 Remociones activas", "F-DT-02 Nivel de erosión",
           "F-DT-02 Urgencia control erosión", "F-DT-03 Estado de conservación",
           "F-DT-03 N.° taxones", "F-DT-04 Causa subyacente principal",
           "F-DT-04 Velocidad de degradación", "F-DT-04 Reversibilidad",
           "F-DT-05 Zona de recarga hídrica"]
    filas = [[b["codigo"], b["area_ha"], b["forma_terreno"], b["posicion"], b["exposicion"],
              b["afloramientos"], b["escarpes"], b["remociones"], b["erosion"],
              b["urgencia_erosion"], b["estado_conservacion"], b["taxones"], b["causa"],
              b["velocidad"], b["reversibilidad"], b["recarga"]] for b in bs]
    _tabla(ws, f, cab, filas,
           anchos=[10, 12, 14, 16, 12, 12, 12, 13, 22, 13, 26, 10, 60, 14, 20, 13],
           notas=["Los campos proceden de las fichas F-DT-01 a F-DT-05 tal como fueron diligenciadas. No se ha homologado la redacción de las causas de degradación entre brigadas."])
    return ws


def hoja_adr(wb, prov, bs):
    ws = wb.create_sheet("8. AdR-CCC preliminar")
    f = _encabezado(ws, prov, "ANÁLISIS PRELIMINAR DEL RIESGO EN CONTEXTO DE CAMBIO CLIMÁTICO",
                    "Guía General para la IFE de PI (DGPMI-MEF, 2022), Anexo 2 — GdR-CCC. "
                    "Ley 29664 (SINAGERD) · Ley 30754 · D.S. 017-2018-MINAM.", 13)
    cab = ["Bloque", "Distrito", "Superficie (ha)", "ISL-MM (sustrato)", "ISL-EH (sustrato)",
           "Pendiente prom. (%)", "% bajo umbral MSAVI (fragilidad de cobertura)",
           "Escarpes activos", "Remociones activas", "Velocidad de degradación",
           "Reversibilidad técnica", "Urgencia de intervención",
           "Peligro integrado (MCA-AHP)"]
    filas = [[b["codigo"], b["distrito"], b["area_ha"], b["isl_clase_mm"], b["isl_clase_eh"],
              b["pend_pct"], b["pct_bajo_umbral"], b["escarpes"], b["remociones"],
              b["velocidad"], b["reversibilidad"], b["urgencia"], b["peligro_integrado"]]
             for b in bs]
    f = _tabla(ws, f, cab, filas,
               anchos=[10, 16, 12, 14, 14, 12, 16, 12, 13, 15, 20, 13, 56],
               notas=["ADVERTENCIA SUSTANTIVA. El nivel de PELIGRO INTEGRADO (MCA-AHP) no está disponible en los insumos de este entregable. Debe tomarse del modelamiento de mesolocalización (PMM + EPH + PGI ponderados por AHP) antes de fijar la prioridad definitiva de intervención de cada bloque. Las columnas de esta matriz son descriptores de sus componentes, no el peligro integrado.",
                      "En este proyecto la Unidad Productora es el ecosistema en su totalidad. El análisis de riesgo se aplica tanto al ecosistema existente como a las MRR-CCC propuestas, que también deben evaluarse por exposición.",
                      "Peligro ≠ exposición ≠ fragilidad ≠ resiliencia. El ISL describe una condición del sustrato (factor condicionante); el porcentaje bajo umbral MSAVI describe la fragilidad de la cobertura; el factor desencadenante principal (precipitación acumulada multianual) no se evalúa en este entregable."])
    ha = sum(b["area_ha"] for b in bs)
    ws.cell(row=f, column=1, value="AGREGADO POR URGENCIA DE INTERVENCIÓN")
    ws.cell(row=f, column=1).font = Font(name="Arial", size=10, bold=True, color=AZUL)
    f += 1
    fl = [[k, v["n"], round(v["ha"], 2), round(100 * v["ha"] / ha, 2)]
          for k, v in D.conteo(bs, "urgencia")]
    fl.append(["TOTAL", len(bs), round(ha, 2), 100.0])
    _tabla(ws, f, ["Urgencia de intervención", "N.° de bloques", "Superficie (ha)",
                   "% de la superficie provincial"], fl, anchos=[30, 14, 14, 16])
    return ws


def hoja_consistencia(wb, prov, bs):
    ws = wb.create_sheet("9. Consistencia")
    f = _encabezado(ws, prov, "CONTROL DE CONSISTENCIA Y REGISTRO DE DISCREPANCIAS",
                    "Discrepancias detectadas entre el registro de campo, los productos de "
                    "gabinete y el catálogo maestro. Calificación: SUSTANTIVA / NO SUSTANTIVA "
                    "/ CORREGIDO / CONFORME.", 6)
    cab = ["Bloque", "Cód.", "Campo afectado", "Discrepancia observada", "Calificación",
           "Tratamiento adoptado"]
    filas = []
    for b in bs:
        for d in b["consistencia"]:
            filas.append([b["codigo"], d["codigo"], d["campo"], d["discrepancia"],
                          d["calificacion"], d["tratamiento"]])

    def marca(row, j):
        if j == 4 and row[4] == "SUSTANTIVA":
            return ROJO
        if j == 4 and row[4] == "CORREGIDO":
            return AMBAR
        return None

    f = _tabla(ws, f, cab, filas, anchos=[10, 8, 26, 60, 14, 60],
               notas=["Una calificación CONFORME acredita que la verificación se ejecutó y no arrojó discrepancia; NO equivale a validación de campo del dato.",
                      "Sombreado rojo: discrepancia SUSTANTIVA, requiere acción antes del cierre del entregable. Sombreado ámbar: discrepancia ya CORREGIDA en esta versión."],
               resaltar=marca)
    cnt = {}
    for r in filas:
        cnt[r[4]] = cnt.get(r[4], 0) + 1
    ws.cell(row=f, column=1, value="RESUMEN DE CALIFICACIONES")
    ws.cell(row=f, column=1).font = Font(name="Arial", size=10, bold=True, color=AZUL)
    f += 1
    fl = [[k, v] for k, v in sorted(cnt.items(), key=lambda kv: -kv[1])]
    fl.append(["TOTAL DE VERIFICACIONES", len(filas)])
    _tabla(ws, f, ["Calificación", "N.° de verificaciones"], fl, anchos=[30, 22])
    return ws


def hoja_microcuencas(wb, prov, bs):
    ws = wb.create_sheet("10. Microcuencas")
    f = _encabezado(ws, prov, "AGREGACIÓN POR MICROCUENCA",
                    "Los bloques de una misma microcuenca comparten régimen hídrico y "
                    "condicionantes de ladera; la microcuenca es la unidad de planificación "
                    "de las MRR-CCC.", 10)
    agg = {}
    for b in bs:
        a = agg.setdefault(b["microcuenca"], {"n": 0, "ha": 0.0, "bl": [], "dist": set()})
        a["n"] += 1
        a["ha"] += b["area_ha"]
        a["bl"].append(b)
        a["dist"].add(b["distrito"])
    filas = []
    for mc, a in sorted(agg.items(), key=lambda kv: -kv[1]["ha"]):
        bl = a["bl"]
        cdm = D.con_dato(bl, "alt_min")
        pm = D.ponderada(bl, "pend_pct")
        filas.append([mc, " · ".join(sorted(a["dist"])), a["n"], round(a["ha"], 2),
                      min(x["alt_min"] for x in cdm) if cdm else "s/d",
                      max(x["alt_max"] for x in cdm) if cdm else "s/d",
                      round(pm, 2) if pm is not None else "s/d",
                      round(D.ponderada(bl, "msavi"), 4),
                      round(sum(x["ha_bajo_umbral"] or 0 for x in bl), 2),
                      round(100 * sum(x["ha_bajo_umbral"] or 0 for x in bl) / a["ha"], 2)])
    ha = sum(b["area_ha"] for b in bs)
    cd = D.con_dato(bs, "alt_min")
    filas.append(["TOTAL / MEDIA PONDERADA", "", len(bs), round(ha, 2),
                  min(b["alt_min"] for b in cd) if cd else "s/d",
                  max(b["alt_max"] for b in cd) if cd else "s/d",
                  round(D.ponderada(bs, "pend_pct"), 2),
                  round(sum(b["msavi"] * b["area_ha"] for b in bs) / ha, 4),
                  round(sum(b["ha_bajo_umbral"] or 0 for b in bs), 2),
                  round(100 * sum(b["ha_bajo_umbral"] or 0 for b in bs) / ha, 2)])
    _tabla(ws, f, ["Microcuenca", "Distrito(s)", "N.° de bloques", "Superficie (ha)",
                   "Altitud mín. (msnm)", "Altitud máx. (msnm)",
                   "Pendiente ponderada (%)", "MSAVI ponderado",
                   "Sup. bajo umbral MSAVI (ha)", "% bajo umbral"],
           filas, anchos=[16, 30, 12, 13, 12, 12, 14, 13, 15, 12],
           notas=["Las medias se ponderan por la superficie de cada bloque. Únicamente se agregan los bloques del ámbito provincial: una microcuenca compartida entre provincias aparece en el anexo de cada una con su fracción correspondiente."])
    return ws


def hoja_estaciones(wb, prov, bs):
    ws = wb.create_sheet("11. Puntos georreferenciados")
    f = _encabezado(ws, prov, "PUNTOS GEORREFERENCIADOS DE VERIFICACIÓN DE CAMPO",
                    "Inventario de los puntos consignados en las fichas F-DT-01 a F-DT-05. "
                    "Sistema de referencia UTM WGS 84 Zona 17S (EPSG:32717).", 9)
    cab = ["Bloque", "Distrito", "Código", "Naturaleza del punto", "Este (m)", "Norte (m)",
           "Dist. al centroide (m)", "Estado / régimen", "Contenido registrado"]
    filas = []
    for b in bs:
        for p in b["puntos"]:
            filas.append([b["codigo"], b["distrito"], p["codigo"], p["naturaleza"],
                          p["este"], p["norte"], p["dist_centroide"], p["estado"],
                          p["contenido"]])
    _tabla(ws, f, cab, filas, anchos=[10, 16, 10, 40, 12, 12, 13, 24, 40],
           notas=["CONTROL GEOMÉTRICO. La distancia al centroide es referencial: no acredita inclusión ni exclusión respecto del polígono, que exige prueba de inclusión sobre la geometría del bloque. Las estaciones situadas a menos de ~30 m del límite deben considerarse SOBRE EL LÍMITE, dentro del margen de error de la georreferenciación (± 10-15 m)."])
    return ws


def hoja_portada(wb, prov, bs):
    ws = wb.active
    ws.title = "0. Portada y contenido"
    r = D.resumen(bs)
    f = _encabezado(ws, prov, "ANEXO DE MATRICES DEL DIAGNÓSTICO TERRITORIAL",
                    "Anexo del Informe Técnico Consolidado de Diagnóstico Territorial. "
                    "Fase de preinversión — Estudio de Perfil · Sistema Invierte.pe.", 4)
    fl = [["Provincia", D.PROV_META[prov]["nombre"]],
          ["N.° de bloques preliminares de intervención", r["n"]],
          ["Superficie total de los bloques (ha)", round(r["ha"], 2)],
          ["N.° de distritos", r["distritos"]],
          ["N.° de microcuencas", r["microcuencas"]],
          ["Rango altitudinal del ámbito (msnm)", f"{r['alt_min']:.0f} – {r['alt_max']:.0f}"],
          ["Pendiente promedio de los bloques — lectura A, % (media simple)",
           round(r["pend_media"], 2)],
          ["Pendiente promedio ponderada por superficie — lectura A (%)",
           round(r["pend_ponderada"], 2)],
          ["Pendiente equivalente bajo la lectura B, en grados (%) — véase D-P01",
           round(D.ponderada(bs, "pend_si_grados_pct"), 2)],
          ["Bloques sin cobertura del MDE (discrepancia D-M01)",
           f"{len(r['sin_mde'])} ({', '.join(r['sin_mde']) if r['sin_mde'] else 'ninguno'})"],
          ["MSAVI 2024 medio (media simple)", round(r["msavi_medio"], 4)],
          ["MSAVI 2024 medio (ponderado por superficie)", round(r["msavi_ponderado"], 4)],
          ["Superficie bajo umbral MSAVI 0.4976 (ha)", round(r["ha_brecha"], 2)],
          ["Brecha espectral sobre la superficie provincial (%)", round(r["pct_brecha"], 2)],
          ["Cárcavas codificadas inventariadas", r["n_carcavas"]],
          ["Longitud total de cárcavas inventariadas (m)", round(r["long_carcavas"], 2)],
          ["Bloques con cárcavas digitalizadas", r["bloques_con_carcavas"]],
          ["Sistema de coordenadas", "UTM WGS 84 Zona 17S (EPSG:32717)"],
          ["CUI del proyecto", "2669244"],
          ["Entidad formuladora", "ANIN — DIME — SESDI"],
          ["Marco del indicador de brecha", "R.M. N.° 00213-2024-MINAM"]]
    f = _tabla(ws, f, ["Parámetro", "Valor"], fl, anchos=[56, 34])
    ws.cell(row=f, column=1, value="CONTENIDO DEL ANEXO")
    ws.cell(row=f, column=1).font = Font(name="Arial", size=10, bold=True, color=AZUL)
    f += 1
    hojas = [
        ["1. Matriz Base", "Matriz base geoespacial consolidada de los bloques."],
        ["2. NDVI-MSAVI", "Matriz de lectura integrada de los índices de vegetación MSAVI 2024 y NDVI mediana 2025."],
        ["3. Integración ISL Geología", "Integración con los resultados del Índice de Susceptibilidad Litológica del E3 Estudio de Geología."],
        ["4. Pendientes", "Pendiente promedio por bloque y distribución por clase."],
        ["5. Cárcavas", "Inventario de cárcavas codificadas con caracterización morfométrica."],
        ["6. Ecosistema y UP", "Ecosistema, estado de conservación y Unidad Productora."],
        ["7. Síntesis F-DT", "Síntesis transversal de las fichas F-DT-01 a F-DT-05."],
        ["8. AdR-CCC preliminar", "Análisis preliminar del riesgo en contexto de cambio climático."],
        ["9. Consistencia", "Control de consistencia y registro de discrepancias."],
        ["10. Microcuencas", "Agregación por microcuenca."],
        ["11. Puntos georreferenciados", "Inventario de puntos georreferenciados de verificación."]]
    f = _tabla(ws, f, ["Hoja", "Contenido"], hojas, anchos=[30, 90],
               notas=["DECLARACIÓN DE INTEGRIDAD DE DATOS. Ningún valor ausente ha sido estimado o inferido sin declararlo. Los campos que no pudieron sustentarse en observación de campo, estadística zonal o catálogo oficial se conservan con la marca de origen «Por determinar» o «Por verificar». Las discrepancias detectadas entre campo, gabinete y catálogo se listan en la hoja «9. Consistencia»."])
    return ws


def generar(prov, bloques, destino):
    bs = D.por_provincia(bloques, prov)
    wb = Workbook()
    hoja_portada(wb, prov, bs)
    hoja_base(wb, prov, bs)
    hoja_vegetacion(wb, prov, bs)
    hoja_isl(wb, prov, bs)
    hoja_pendientes(wb, prov, bs)
    hoja_carcavas(wb, prov, bs)
    hoja_ecosistema(wb, prov, bs)
    hoja_sintesis(wb, prov, bs)
    hoja_adr(wb, prov, bs)
    hoja_consistencia(wb, prov, bs)
    hoja_microcuencas(wb, prov, bs)
    hoja_estaciones(wb, prov, bs)
    ruta = os.path.join(destino, f"Anexo_Matrices_DT_{D.PROV_META[prov]['archivo']}_IN_Piura.xlsx")
    wb.save(ruta)
    return ruta


if __name__ == "__main__":
    bl = D.cargar_bloques()
    out = os.path.join(BASE_OUT := os.path.dirname(os.path.abspath(__file__)), "entregables")
    os.makedirs(out, exist_ok=True)
    for p in D.PROVINCIAS:
        print("OK:", generar(p, bl, out))
