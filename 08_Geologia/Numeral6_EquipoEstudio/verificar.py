# -*- coding: utf-8 -*-
"""Comprueba que el numeral 6 generado lleve los datos corregidos del equipo de estudio."""
import sys

from docx import Document

ESPERADO = ["Christian Camacho Aponte", "Ing. Agrícola", "Asistente en SIG"]
PROSCRITO = ["Camacho Flores", "Bach. Christian", "Bach. en Ing. Agrícola",
             "Asistente SIG (Apoyo"]


def texto_completo(ruta):
    doc = Document(ruta)
    partes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for fila in t.rows:
            partes.extend(c.text for c in fila.cells)
    return "\n".join(partes)


def main():
    ruta = sys.argv[1] if len(sys.argv) > 1 else "E3_Geologia_6_Equipo_de_Estudio.docx"
    txt = texto_completo(ruta)
    fallos = 0
    for e in ESPERADO:
        ok = e in txt
        fallos += not ok
        print(f"{'OK  ' if ok else 'FALTA'}  presente: {e!r}")
    for p in PROSCRITO:
        ok = p not in txt
        fallos += not ok
        print(f"{'OK  ' if ok else 'ERROR'}  ausente : {p!r}")
    print("\nRESULTADO:", "correcto" if not fallos else f"{fallos} problema(s)")
    return 1 if fallos else 0


if __name__ == "__main__":
    sys.exit(main())
