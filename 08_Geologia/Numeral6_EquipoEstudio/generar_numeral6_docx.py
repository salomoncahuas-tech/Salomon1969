# -*- coding: utf-8 -*-
"""Genera el documento autonomo del numeral 6. EQUIPO DE ESTUDIO.

Entregable 3 - Estudio de Geologia. Proyecto IN Piura (CUI 2669244) - ANIN / DIME / SESDI.
Produce el mismo entregable que se inserta en cada informe provincial, pero como
archivo independiente en orientacion horizontal (A4 apaisado).

Uso:
    python3 generar_numeral6_docx.py [organigrama.png] [salida.docx]
"""
import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from estilo import AZUL_ANIN, VERDE_ANIN, leyenda, parrafo
from insertar_numeral6 import construir_numeral6

ENCABEZADO = [
    ("AUTORIDAD NACIONAL DE INFRAESTRUCTURA – ANIN", 17, True, VERDE_ANIN),
    ("DIRECCIÓN DE INTERVENCIONES MULTISECTORIALES Y DE EMERGENCIA – DIME", 14, True, VERDE_ANIN),
    ("SUBDIRECCIÓN DE ESTUDIOS DE INVERSIÓN – SESDI", 14, True, VERDE_ANIN),
]

SUBTITULO = ("Proyecto IN Piura (CUI 2669244) · Entregable 3 – Estudio de Geología · "
             "Modalidad de ejecución: Administración Directa")

PIE = "Entregable 3 – Estudio de Geología · Acápite 6. Equipo de Estudio"


def configurar_pagina(doc):
    """A4 apaisado con margenes estrechos, para que el cuadro 6.1 entre completo."""
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for lado in ("left_margin", "right_margin"):
        setattr(sec, lado, Cm(1.5))
    sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
    return sec


def fuente_base(doc):
    """Arial en todo el documento, conforme al estandar institucional ANIN."""
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8.5)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), "Arial")


def armar_encabezado(sec):
    hdr = sec.header
    hdr.paragraphs[0].text = ""
    cuerpo = hdr.paragraphs[0]._p.getparent()
    cuerpo.remove(hdr.paragraphs[0]._p)
    for texto, sz, bold, color in ENCABEZADO:
        cuerpo.append(parrafo(texto, sz=sz, bold=bold, color=color, jc="center", after=0))
    cuerpo.append(parrafo(SUBTITULO, sz=13, italic=True, color=VERDE_ANIN,
                          jc="center", after=60))


def armar_pie(sec):
    ftr = sec.footer
    ftr.paragraphs[0].text = ""
    cuerpo = ftr.paragraphs[0]._p.getparent()
    cuerpo.remove(ftr.paragraphs[0]._p)
    cuerpo.append(parrafo(PIE, sz=13, italic=True, color=AZUL_ANIN, jc="center", after=0))


def main():
    png = sys.argv[1] if len(sys.argv) > 1 else "Organigrama_E3_Geologia.png"
    salida = sys.argv[2] if len(sys.argv) > 2 else "E3_Geologia_6_Equipo_de_Estudio.docx"
    if not os.path.exists(png):
        raise SystemExit(f"ERROR: no existe {png}")

    doc = Document()
    sec = configurar_pagina(doc)
    fuente_base(doc)
    armar_encabezado(sec)
    armar_pie(sec)

    ancho_util = int((sec.page_width - sec.left_margin - sec.right_margin) / 635)  # EMU->dxa
    cuerpo = doc.element.body
    for el in construir_numeral6(doc, png, ancho_util):
        cuerpo.append(el)

    # el <w:sectPr> debe quedar al final del cuerpo
    sectPr = sec._sectPr
    sectPr.getparent().remove(sectPr)
    cuerpo.append(sectPr)

    doc.save(salida)
    print(f"OK -> {salida} ({os.path.getsize(salida):,} bytes)")


if __name__ == "__main__":
    main()
