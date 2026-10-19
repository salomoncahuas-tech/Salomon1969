"""Extrae los datos del Diagnóstico Integral IN Piura V1.

Lee los tres volúmenes (fuentes/vol*_v5.html) y el motor de costos
(fuentes/Costos_Parametros_IN_Piura_v1.3.xlsx) y escribe en datos/:

  data.json          datos que se incrustan en la portada integral
  costos_data.json   entradas del motor de costos (para probar_motor.js)
  ref.json           resultados del Excel con jornal 80, 100 y 120 (control)

También extrae src/base.css (hojas de estilo comunes de los volúmenes).

Los costos unitarios son lineales en el jornal: se toman del Excel con el
jornal vigente y de una copia recalculada con LibreOffice a S/ 120 por día.
Requiere openpyxl y LibreOffice (soffice) en el PATH.

Uso:  python extraer_datos.py
"""
import json
import os
import re
import shutil
import subprocess
import tempfile

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
FUENTES = os.path.join(BASE, 'fuentes')
DATOS = os.path.join(BASE, 'datos')
EXCEL = os.path.join(FUENTES, 'Costos_Parametros_IN_Piura_v1.3.xlsx')
VOLS = {'mor': 'vol1_morropon_v5.html', 'hua': 'vol2_huancabamba_v5.html', 'aya': 'vol3_ayabaca_v5.html'}


def leer_int(ruta):
    """Objeto window.INT de un volumen (datos por bloque e indicadores)."""
    s = open(ruta, encoding='utf8').read()
    i = s.index('window.INT=') + len('window.INT=')
    obj, _ = json.JSONDecoder().raw_decode(s[i:])
    return obj


def extraer_css(ruta, destino):
    s = open(ruta, encoding='utf8').read()
    parte = s[s.index('</head><body>') + len('</head><body>'):s.index('<header class="hero">')]
    css = '\n'.join(re.findall(r'<style>(.*?)</style>', parte, re.S))
    open(destino, 'w', encoding='utf8').write(css)


def recalcular(jornal, carpeta):
    """Copia del Excel con otro jornal, recalculada con LibreOffice."""
    wb = openpyxl.load_workbook(EXCEL)
    wb['PARAMETROS']['C8'] = jornal
    entrada = os.path.join(carpeta, f'costos_j{jornal}.xlsx')
    wb.save(entrada)
    salida = os.path.join(carpeta, 'recalc')
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML',
                    '--outdir', salida, entrada], check=True, capture_output=True, timeout=300)
    return openpyxl.load_workbook(os.path.join(salida, f'costos_j{jornal}.xlsx'), data_only=True)


def main():
    os.makedirs(DATOS, exist_ok=True)
    I = {k: leer_int(os.path.join(FUENTES, f)) for k, f in VOLS.items()}
    extraer_css(os.path.join(FUENTES, VOLS['hua']), os.path.join(BASE, 'src', 'base.css'))

    f = openpyxl.load_workbook(EXCEL)
    v = openpyxl.load_workbook(EXCEL, data_only=True)
    tmp = tempfile.mkdtemp()
    try:
        r120 = recalcular(120, tmp)
        r80 = recalcular(80, tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    # BLOQUES: entradas SIG y tipología por bloque
    B = v['BLOQUES']
    cols = 'A B C D E F G H I J K L M N O P Q R S T U V W X Y Z AA AB AC AD AE AF AG AH AI'.split()
    keys = ('id prov dist mc ha per ms bes and agr d1 d2 d3 d4 d5 p1 p2 p3 p4 p5 '
            'r1a r1b r2a r2b ca cb dvia dac via cml pil pbu grp deg tip').split()
    blocks = []
    for r in range(7, 124):
        o = {k: B[f'{c}{r}'].value for k, c in zip(keys, cols)}
        o['id'] = str(o['id'])
        blocks.append(o)

    # CU_COMP: costos unitarios compuestos con jornal vigente y con S/ 120
    cu = {}
    for row, name in [(60, 'PLANTON'), (61, 'PLANTA'), (62, 'HA'), (63, 'CERCO'), (64, 'ZANJA'), (65, 'TERRAZA'), (66, 'CARCAVA')]:
        cu[name] = {'h': [v['CU_COMP'][f'H{row}'].value, r120['CU_COMP'][f'H{row}'].value],
                    'nc': [v['CU_COMP'][f'I{row}'].value, r120['CU_COMP'][f'I{row}'].value],
                    'c': [v['CU_COMP'][f'J{row}'].value, r120['CU_COMP'][f'J{row}'].value],
                    'desc': v['CU_COMP'][f'C{row}'].value}

    P = v['PAQUETES']
    paq = [{'t': P[f'A{r}'].value, 'grp': P[f'C{r}'].value, 'deg': P[f'D{r}'].value, 'r1': P[f'E{r}'].value,
            'r2': P[f'F{r}'].value, 'ez': P[f'G{r}'].value, 'et': P[f'H{r}'].value, 'cer': P[f'I{r}'].value,
            'src': P[f'L{r}'].value, 'key': P[f'B{r}'].value} for r in range(7, 13)]

    O, Ov = f['OYM'], v['OYM']
    ind = {'=C_R1_T+C_R2_T': 'ha', '=C_PL_T': 'pl', '=C_KZ_T': 'kz', '=C_KT_T': 'kt', '=C_KC_T': 'kc'}
    oym = [{'comp': O[f'A{r}'].value, 'act': O[f'B{r}'].value, 'c': O[f'C{r}'].value, 'd': O[f'D{r}'].value,
            'u': O[f'E{r}'].value, 'anios': O[f'G{r}'].value, 'ind': ind[O[f'H{r}'].value],
            'f100': Ov[f'F{r}'].value, 'f120': r120['OYM'][f'F{r}'].value} for r in range(7, 21)]

    # Participación de MO no calificada en O&M (OYM!F23): sumas de PARTIDAS H e I
    claves = {f'REFORESTACION|04.02.0{i}' for i in range(1, 5)}
    H, Isum = [0, 0], [0, 0]
    for r in range(7, 133):
        if v['PARTIDAS'][f'A{r}'].value in claves:
            H[0] += v['PARTIDAS'][f'H{r}'].value
            Isum[0] += v['PARTIDAS'][f'I{r}'].value
            H[1] += r120['PARTIDAS'][f'H{r}'].value
            Isum[1] += r120['PARTIDAS'][f'I{r}'].value

    PR = v['PARAMETROS']
    par = {}
    for r in range(8, 49):
        n = PR[f'A{r}'].value
        if n and str(n).startswith('P_'):
            par[n] = PR[f'C{r}'].value
    par['P_HH_PEON_2020'] = PR['H10'].value
    flu = [v['FLUJO'].cell(8, c).value for c in range(3, 23)]

    costos = {'blocks': blocks, 'cu': cu, 'paq': paq, 'oym': oym, 'omH': H, 'omI': Isum, 'par': par, 'flu': flu}
    json.dump(costos, open(os.path.join(DATOS, 'costos_data.json'), 'w'), ensure_ascii=False)

    # Resultados de referencia del Excel (RESUMEN y CALCULO) para el control del motor
    ref = {}
    for j, wb in [(100, v), (120, r120), (80, r80)]:
        R = wb['RESUMEN']
        ref[j] = {'res': {f'{c}{r}': R[f'{c}{r}'].value for r in range(8, 30) for c in 'BC'},
                  'blk': {str(wb['CALCULO'][f'A{r}'].value): [wb['CALCULO'][f'{c}{r}'].value
                          for c in ['J', 'K', 'L', 'M', 'Y', 'AA', 'AF', 'AG', 'AH']] for r in range(7, 124)},
                  'tip': {R[f'A{r}'].value: R[f'E{r}'].value for r in range(39, 45)}}
    json.dump(ref, open(os.path.join(DATOS, 'ref.json'), 'w'))

    # data.json: volúmenes + motor de costos para la portada
    vols = {}
    for k in ['mor', 'hua', 'aya']:
        x = I[k]
        vols[k] = {kk: x[kk] for kk in ['prov', 'vol', 'kpi', 'regs', 'regTxt', 'dsTxt', 'gaps', 'agg', 'nd', 'nmc', 'nreg']}
        vols[k]['B'] = [{c: b.get(c) for c in ['b', 'd', 'reg', 'ha', 'e', 'n', 'alt', 'pend', 'ms', 'ipp', 'p45',
                                               'mm', 'eh', 'ncv', 'lcv', 'P', 'AI', 'pob', 'rk']} for b in x['B']]
    C = dict(costos)
    C['carcML'] = {b['b']: b['lcv'] for k in I for b in I[k]['B']}
    dist = {b['b']: b['d'] for k in I for b in I[k]['B']}
    sin_vol = [b['id'] for b in blocks if b['id'] not in dist]
    if sin_vol:
        raise SystemExit(f'Bloques del Excel sin volumen: {sin_vol}')
    corto = 'id prov dist distN mc ha per ms bes and agr r1a r1b r2a r2b ca cb dvia dac via cml pil pbu grp deg tip'.split()
    C['blocks'] = [{k: (dist[b['id']] if k == 'distN' else b[k]) for k in corto} for b in blocks]
    L = v['LEEME']
    C['leeme'] = [[L[f'A{r}'].value, L[f'B{r}'].value] for r in range(7, 19)]
    C['lect'] = [v['RESUMEN'][f'A{r}'].value for r in range(47, 53)]
    parT = []
    for r in range(7, 49):
        a, b, c, d, e = [PR[f'{x}{r}'].value for x in 'ABCDE']
        if a and not str(a).startswith('P_'):
            parT.append({'sec': a})
        elif a:
            parT.append({'n': a, 'p': b, 'v': c, 'u': d, 'f': e})
    C['parT'] = parT
    json.dump({'vols': vols, 'cost': C}, open(os.path.join(DATOS, 'data.json'), 'w'),
              ensure_ascii=False, separators=(',', ':'))

    print(f'{len(blocks)} bloques · {round(sum(b["ha"] for b in blocks), 3)} ha en el motor · '
          f'{sum(len(I[k]["B"]) for k in I)} bloques en los volúmenes')


if __name__ == '__main__':
    main()
