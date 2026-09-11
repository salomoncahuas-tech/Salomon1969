"""
Analitica del Diagnostico Territorial por bloque - Proyecto IN Piura.

Arma, sobre las fichas de resumen por bloque y su integracion con la ficha
DT de campo, el mismo informe analitico que ya existe para el Diagnostico
Social: cifras de cabecera, secciones con sus series graficables y tablas de
respaldo. De dibujarlo se encarga `analitica_series`, de modo que el grafico
interactivo del aplicativo, el grafico nativo del libro Excel y el anexo PDF
salen de la misma declaracion y no pueden contradecirse.

Este modulo NO lee ni escribe archivos ni base de datos: recibe el
diccionario que devuelve `resumenes_bloques.parsear_resumen_bloque` -y, si
esta disponible, el registro integrado de `dt_campo.integrar_bloque`- y
devuelve estructuras de datos puras. No recalcula ni estima ningun valor:
solo ordena, colorea y agrupa lo que ya viene declarado.

ANIN - DIME - SESDI | CUI 2669244 | UTM WGS 84 Zona 17S (EPSG:32717).
"""

import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

import analitica_series as se
import dt_campo as dtc
import resumenes_bloques as rbq
from analitica_series import (
    RAMPA_CRITICA, RAMPA_FAVORABLE, RAMPA_NEUTRA, SIN_DATO,
    _conteo, _escribir_tabla, _hoja_seccion, _hoja_tablas, _num,
    _pdf_barras_apiladas, _pdf_barras_simples, _serie, _titulo_hoja, _txt,
)


# ══════════════════════════════════════════════════════════════════════════
# ESCALAS ORDENADAS DEL DIAGNOSTICO TERRITORIAL
# ══════════════════════════════════════════════════════════════════════════
# Las listas de la plantilla V5 son escalas ordenadas, no identidades: el
# color debe leerse como severidad y no como etiqueta. Se fija aqui el color
# de cada clase sobre el dominio COMPLETO de la lista, de modo que una misma
# clase conserve su tono entre un bloque y otro y los graficos se puedan
# comparar entre si.

def _escala(clases, rampa, adversa=True):
    """Color por clase sobre una rampa, de la menos a la mas grave."""
    pasos = se._rampa(clases, rampa)
    return dict(zip(clases, pasos if adversa else list(reversed(pasos))))


ORDEN_EROSION = ["Nula", "Ligera (laminar)", "Moderada (surcos)",
                 "Fuerte (cárcavas incipientes)", "Severa (cárcavas activas)",
                 "Extrema"]
COLOR_EROSION = _escala(ORDEN_EROSION, RAMPA_CRITICA)

ORDEN_CONSERVACION = [
    "Conservado (sin intervención evidente)", "Levemente alterado",
    "Medianamente alterado", "Alterado (intervención marcada)",
    "Muy alterado / Degradado", "En restauración / Recuperación",
]
# «En restauracion» no es el extremo degradado: es una condicion favorable y
# se saca de la rampa para no leerse como la peor clase.
COLOR_CONSERVACION = _escala(ORDEN_CONSERVACION[:-1], RAMPA_CRITICA)
COLOR_CONSERVACION[ORDEN_CONSERVACION[-1]] = RAMPA_FAVORABLE[2]

ORDEN_URGENCIA = ["Baja", "Media", "Alta", "Crítica"]
COLOR_URGENCIA = _escala(ORDEN_URGENCIA, RAMPA_CRITICA)

ORDEN_VELOCIDAD = ["Estable", "Lenta", "Moderada", "Rápida", "Muy rápida"]
COLOR_VELOCIDAD = _escala(ORDEN_VELOCIDAD, RAMPA_CRITICA)

ORDEN_REVERSIBILIDAD = ["Totalmente reversible", "En recuperación",
                        "Parcialmente reversible", "Difícilmente reversible"]
COLOR_REVERSIBILIDAD = _escala(ORDEN_REVERSIBILIDAD, RAMPA_CRITICA)

ORDEN_NDVI = ["Vegetación alta", "Vegetación mediana", "Vegetación ligera",
              "Suelo desnudo / agua"]
COLOR_NDVI = _escala(ORDEN_NDVI, RAMPA_FAVORABLE, adversa=False)

# El control de consistencia comunica severidad, no identidad: la
# discrepancia sustantiva condiciona el cierre del entregable.
COLOR_CALIFICACION = {
    "SUSTANTIVA": RAMPA_CRITICA[4], "NO SUSTANTIVA": "#C4A03C",
    "CORREGIDO": RAMPA_FAVORABLE[2], "CONFORME": RAMPA_FAVORABLE[4],
}

# La procedencia si es identidad: tres fuentes que no se ordenan entre si.
COLOR_FUENTE = {
    dtc.ETIQUETA_FUENTE[dtc.CAMPO]: "#2E7D4F",
    dtc.ETIQUETA_FUENTE[dtc.GABINETE]: "#1F6FB2",
    dtc.ETIQUETA_FUENTE[dtc.OFICIAL]: "#C4A03C",
}

COLOR_ESTADO_CRUCE = {
    dtc.CONFORME: RAMPA_FAVORABLE[4], dtc.COMPLEMENTADO: RAMPA_FAVORABLE[2],
    dtc.ACTUALIZADO: "#1F6FB2", dtc.DISCREPANTE: "#C4A03C",
    dtc.PENDIENTE: SIN_DATO,
}

COLOR_UMBRAL = {
    "Bajo umbral %s (brecha)" % rbq.UMBRAL_MSAVI: RAMPA_CRITICA[4],
    "Sobre umbral %s" % rbq.UMBRAL_MSAVI: RAMPA_FAVORABLE[4],
}

FUENTE_CATALOGO = ("Catálogo maestro Bloques V5/V6 y estadística zonal sobre "
                   "el MDE y los compuestos Sentinel-2.")
FUENTE_CAMPO = "Fichas F-DT-01 a F-DT-05 de la plantilla DT de campo."


def _color_msavi(clase, condicion):
    """Color de una clase DN: crítico bajo el umbral, favorable sobre él.

    El umbral de la R.M. N.° 00213-2024-MINAM es lo que separa la brecha de
    lo que no lo es; el gráfico debe mostrar ese corte, no cinco tonos de un
    mismo verde donde ninguna clase se distingue de la siguiente.
    """
    bajo = "bajo" in se._clave(condicion)
    rampa = RAMPA_CRITICA if bajo else RAMPA_FAVORABLE
    m = re.search(r"DN\s*(\d)", clase or "")
    dn = int(m.group(1)) if m else 3
    # DN 1 es lo mas degradado y DN 5 lo mas vigoroso.
    paso = {1: 4, 2: 3, 3: 2}.get(dn, 0) if bajo else {5: 4, 4: 3}.get(dn, 2)
    return rampa[paso]


# ══════════════════════════════════════════════════════════════════════════
# SECCIONES DE UN BLOQUE
# ══════════════════════════════════════════════════════════════════════════
# Cada constructor recibe los datos ya parseados y devuelve una seccion
# {titulo, descripcion, series, tablas} o None si el bloque no trae con que
# construirla. Nunca se inventa una serie a partir de un valor ausente.

def _seccion(titulo, descripcion, series, tablas=()):
    series = [s for s in series if s and s.get("filas")]
    if not series and not tablas:
        return None
    return {"titulo": titulo, "descripcion": descripcion, "series": series,
            "tablas": [t for t in tablas if t[1]]}


def _seccion_vegetacion(datos):
    """Distribucion areal de los indices espectrales y brecha del indicador."""
    series = []

    ndvi = [r for r in (datos.get("ndvi_tabla") or [])
            if r.get("superficie_ha") is not None]
    if ndvi:
        series.append(_serie(
            "dt_ndvi", "NDVI mediana 2025 — distribución areal", "barras_h",
            [{"clase": r["clase"], "valor": r["superficie_ha"]} for r in ndvi],
            "clase", "valor", unidad="ha", decimales=2,
            eje_x="Superficie (ha)", eje_y="Clase NDVI",
            colores={r["clase"]: COLOR_NDVI.get(r["clase"], SIN_DATO)
                     for r in ndvi},
            descripcion="Reparto de la superficie del bloque entre las clases "
                        "de NDVI del compuesto Sentinel-2 de 2025.",
            nota="Estadística zonal del NDVI mediana 2025 sobre el polígono "
                 "del bloque."))

    msavi = [r for r in (datos.get("msavi_tabla") or [])
             if r.get("superficie_ha") is not None]
    if msavi:
        series.append(_serie(
            "dt_msavi", "MSAVI 2024 — superficie por clase DN", "barras_h",
            [{"clase": r["clase"], "valor": r["superficie_ha"]} for r in msavi],
            "clase", "valor", unidad="ha", decimales=3,
            eje_x="Superficie (ha)", eje_y="Clase DN",
            orden_cat=[r["clase"] for r in msavi],
            colores={r["clase"]: _color_msavi(r.get("clase"), r.get("condicion"))
                     for r in msavi},
            descripcion="Las clases bajo el umbral %s se muestran en tonos de "
                        "severidad y las que lo superan en tonos favorables."
                        % rbq.UMBRAL_MSAVI,
            nota="Estadística zonal del ráster MSAVI 2024 clasificado."))

    bajo = datos.get("msavi_bajo_umbral_ha")
    sobre = datos.get("msavi_sobre_umbral_ha")
    if bajo is not None and sobre is not None:
        etiqueta_bajo = "Bajo umbral %s (brecha)" % rbq.UMBRAL_MSAVI
        etiqueta_sobre = "Sobre umbral %s" % rbq.UMBRAL_MSAVI
        series.append(_serie(
            "dt_brecha", "Brecha espectral frente al umbral del indicador",
            "barras_h",
            [{"clase": etiqueta_bajo, "valor": bajo},
             {"clase": etiqueta_sobre, "valor": sobre}],
            "clase", "valor", unidad="ha", decimales=2,
            eje_x="Superficie (ha)", eje_y="Condición",
            colores=COLOR_UMBRAL,
            descripcion="Superficie que sustenta el indicador de brecha de la "
                        "R.M. N.° 00213-2024-MINAM.",
            nota="Derivado de la distribución areal del MSAVI 2024."))

    tablas = []
    if msavi:
        tablas.append(("MSAVI 2024 por clase DN", [{
            "Clase DN": r.get("clase", ""),
            "Superficie (ha)": r.get("superficie_ha"),
            "% del área": r.get("pct"),
            "Interpretación": r.get("interpretacion", ""),
            "Condición": r.get("condicion", ""),
        } for r in (datos.get("msavi_tabla") or [])]))
    if ndvi:
        tablas.append(("NDVI 2025 por clase", [{
            "Clase NDVI": r.get("clase", ""),
            "Superficie (ha)": r.get("superficie_ha"),
            "% del área": r.get("pct"),
            "Interpretación": r.get("interpretacion", ""),
        } for r in (datos.get("ndvi_tabla") or [])]))

    return _seccion(
        "Índices de vegetación",
        "Distribución areal de los compuestos Sentinel-2 sobre el polígono "
        "del bloque y superficie que sustenta el indicador de brecha.",
        series, tablas)


def _seccion_microcuenca(datos):
    """El bloque frente a los demas bloques de su microcuenca."""
    micro = [r for r in (datos.get("microcuenca_tabla") or [])
             if r.get("area_ha") is not None]
    if len(micro) < 2:
        return None
    codigo = _txt(datos.get("codigo_bloque"))
    nombre = _txt(datos.get("microcuenca"))

    def etiqueta(registro):
        return ("► " + _txt(registro.get("bloque")) if registro.get("es_actual")
                else _txt(registro.get("bloque")))

    # El bloque de la ficha se destaca; los demas quedan en un tono neutro
    # para que la comparacion se lea de un vistazo.
    colores = {etiqueta(r): ("#2E7D4F" if r.get("es_actual") else RAMPA_NEUTRA[1])
               for r in micro}
    orden = [etiqueta(r) for r in micro]

    series = [_serie(
        "dt_micro_area", "Superficie por bloque en la microcuenca %s" % nombre,
        "barras_h",
        [{"clase": etiqueta(r), "valor": r["area_ha"]} for r in micro],
        "clase", "valor", unidad="ha", decimales=2, orden_cat=orden,
        colores=colores, eje_x="Superficie (ha)", eje_y="Bloque",
        descripcion="El bloque %s aparece destacado entre los demás bloques "
                    "de su microcuenca." % codigo,
        nota=FUENTE_CATALOGO)]

    con_msavi = [r for r in micro if r.get("msavi") is not None]
    if len(con_msavi) > 1:
        series.append(_serie(
            "dt_micro_msavi", "MSAVI 2024 por bloque en la microcuenca %s" % nombre,
            "barras_h",
            [{"clase": etiqueta(r), "valor": r["msavi"]} for r in con_msavi],
            "clase", "valor", decimales=4, orden_cat=orden, colores=colores,
            eje_x="MSAVI 2024 (media del bloque)", eje_y="Bloque",
            descripcion="Vigor vegetal medio de cada bloque; el umbral del "
                        "indicador es %s." % rbq.UMBRAL_MSAVI,
            nota=FUENTE_CATALOGO))

    tablas = [("Contexto intramicrocuenca", [{
        "Bloque": etiqueta(r),
        "Área (ha)": r.get("area_ha"),
        "% microcuenca": r.get("pct_microcuenca"),
        "Rango altitudinal (msnm)": r.get("rango_altitudinal", ""),
        "Pendiente prom. (%)": r.get("pendiente_pct"),
        "MSAVI 2024": r.get("msavi"),
        "NDVI — veg. alta (%)": r.get("ndvi_veg_alta"),
    } for r in (datos.get("microcuenca_tabla") or [])])]

    return _seccion(
        "Contexto intramicrocuenca",
        "Posición relativa del bloque dentro de su microcuenca, en superficie "
        "y en vigor vegetal.", series, tablas)


def _seccion_consistencia(datos):
    """Reparto de las verificaciones del control de consistencia."""
    registros = datos.get("consistencia") or []
    if not registros:
        return None
    resumen = datos.get("consistencia_resumen") or {}
    filas = [{"clase": c, "valor": resumen.get(c, 0)}
             for c in rbq.CALIFICACIONES if resumen.get(c, 0)]

    series = [_serie(
        "dt_consistencia", "Verificaciones por calificación", "barras_h", filas,
        "clase", "valor", unidad="verificaciones",
        eje_x="N.° de verificaciones", eje_y="Calificación",
        colores=COLOR_CALIFICACION, orden_cat=list(rbq.CALIFICACIONES),
        descripcion="Resultado del cruce entre la ficha de campo, la de "
                    "gabinete y el catálogo maestro.",
        nota="Hoja «Control de consistencia» de la ficha de resumen.")]

    tablas = [("Control de consistencia", [{
        "Cód.": r.get("codigo", ""), "Campo afectado": r.get("campo", ""),
        "Discrepancia observada": r.get("discrepancia", ""),
        "Calificación": r.get("calificacion", ""),
        "Tratamiento adoptado": r.get("tratamiento", ""),
    } for r in registros])]

    return _seccion(
        "Control de consistencia",
        "Cada verificación declara una discrepancia observada y el "
        "tratamiento adoptado; ninguna se resuelve estimando valores.",
        series, tablas)


def _seccion_campo(integrado):
    """Lo que solo aporta la verificacion en campo (fichas F-DT)."""
    if not integrado or not integrado.get("tiene_campo"):
        return None
    por_clave = integrado.get("por_clave") or {}
    series = []

    estratos = [
        ("Dosel arbóreo", "cobertura_dosel"),
        ("Arbustiva", "cobertura_arbustiva"),
        ("Herbácea", "cobertura_herbacea"),
        ("Hojarasca", "cobertura_hojarasca"),
        ("Suelo desnudo", "suelo_desnudo_pct"),
    ]
    filas = [{"clase": etiqueta,
              "valor": _num((por_clave.get(clave) or {}).get("valor"))}
             for etiqueta, clave in estratos]
    filas = [f for f in filas if f["valor"] is not None]
    if filas:
        # El suelo desnudo es lo adverso del reparto: se colorea aparte para
        # que no se lea como un estrato mas de la cobertura.
        colores = {f["clase"]: RAMPA_FAVORABLE[3] for f in filas}
        colores["Suelo desnudo"] = RAMPA_CRITICA[4]
        series.append(_serie(
            "dt_estratos", "Cobertura por estrato declarada en campo (%)",
            "barras_h", filas, "clase", "valor", unidad="%", decimales=1,
            eje_x="% del área de la parcela", eje_y="Estrato",
            colores=colores, orden_cat=[f["clase"] for f in filas],
            descripcion="Lectura de la parcela de muestreo de F-DT-03.",
            nota=FUENTE_CAMPO))

    causas = integrado.get("causas_activas") or []
    if causas:
        series.append(_serie(
            "dt_causas", "Causas activas de degradación por intensidad",
            "barras_h",
            [{"clase": c["causa"], "valor": c["peso"] + 1} for c in causas],
            "clase", "valor", escala="critica",
            eje_x="Intensidad declarada (1 ligera – 4 muy fuerte)",
            eje_y="Causa", unidad="",
            orden_cat=[c["causa"] for c in causas],
            descripcion="Matriz de causas de F-DT-04; solo se listan las "
                        "declaradas presentes.",
            nota=FUENTE_CAMPO))

    fuentes = integrado.get("fuentes_agua") or []
    if fuentes:
        filas_fuente = _conteo([_txt(f.get("tipo")) for f in fuentes])
        if filas_fuente:
            series.append(_serie(
                "dt_fuentes_agua", "Fuentes de agua inventariadas por tipo",
                "barras_h", filas_fuente, "clase", "valor", escala="neutra",
                unidad="fuentes", eje_x="N.° de fuentes",
                eje_y="Tipo de fuente",
                descripcion="Inventario hídrico de F-DT-05.",
                nota=FUENTE_CAMPO))

    flora = integrado.get("floristica") or []
    if flora:
        filas_estrato = _conteo([_txt(f.get("estrato")) for f in flora])
        if filas_estrato:
            series.append(_serie(
                "dt_flora", "Elenco florístico por estrato", "barras_h",
                filas_estrato, "clase", "valor", escala="favorable",
                unidad="taxones", eje_x="N.° de taxones", eje_y="Estrato",
                descripcion="Composición del elenco levantado en la parcela "
                            "de muestreo.",
                nota=FUENTE_CAMPO))

    tablas = []
    for clave, titulo, columnas in (
            ("floristica", "Elenco florístico",
             [("n", "N.°"), ("nombre_comun", "Nombre común"),
              ("nombre_cientifico", "Nombre científico"), ("familia", "Familia"),
              ("estrato", "Estrato"), ("origen", "Origen"),
              ("abundancia", "Abundancia")]),
            ("especies_clave", "Especies clave georreferenciadas",
             [("n", "N.°"), ("nombre", "Especie"), ("categoria", "Categoría"),
              ("estado_uicn", "Estado UICN"), ("utm_e", "UTM ESTE"),
              ("utm_n", "UTM NORTE"), ("observacion", "Observación")]),
            ("carcavas", "Inventario de cárcavas",
             [("codigo", "Código"), ("tipo", "Tipo"),
              ("longitud_m", "Longitud (m)"), ("prof_m", "Prof. (m)"),
              ("ancho_m", "Ancho (m)"), ("estado", "Estado"),
              ("causa", "Causa")]),
            ("fuentes_agua", "Fuentes de agua",
             [("n", "N.°"), ("tipo", "Tipo de fuente"), ("utm_e", "UTM ESTE"),
              ("utm_n", "UTM NORTE"), ("regimen", "Régimen"),
              ("calidad", "Calidad"), ("uso_obs", "Uso / observación")]),
            ("indicadores", "Indicadores cuantitativos de F-DT-04",
             [("n", "N.°"), ("indicador", "Indicador"), ("unidad", "Unidad"),
              ("valor", "Valor"), ("umbral", "Umbral"), ("nivel", "Nivel")])):
        filas_tabla = integrado.get(clave) or []
        if filas_tabla:
            tablas.append((titulo, [{e: f.get(c, "") for c, e in columnas}
                                    for f in filas_tabla]))

    return _seccion(
        "Verificación de campo",
        "Hechos que solo levanta la ficha DT en terreno: estructura de la "
        "cobertura, causas de degradación, inventario hídrico y florístico.",
        series, tablas)


def _seccion_procedencia(integrado):
    """De donde sale cada hecho declarado y como cruzaron las dos fuentes."""
    if not integrado:
        return None
    conteo_fuente = integrado.get("conteo_fuente") or {}
    conteo_estado = integrado.get("conteo_estado") or {}
    series = []

    filas = [{"clase": dtc.ETIQUETA_FUENTE[f], "valor": conteo_fuente.get(f, 0)}
             for f in dtc.FUENTES if conteo_fuente.get(f, 0)]
    if filas:
        series.append(_serie(
            "dt_procedencia", "Procedencia de los hechos declarados",
            "barras_h", filas, "clase", "valor", unidad="hechos",
            eje_x="N.° de hechos", eje_y="Fuente", colores=COLOR_FUENTE,
            descripcion="Cada hecho territorial se declara una sola vez, con "
                        "la fuente que manda sobre él.",
            nota="Cruce de la ficha DT de campo con la ficha de resumen de "
                 "gabinete y el catálogo maestro."))

    filas = [{"clase": e, "valor": conteo_estado.get(e, 0)}
             for e in dtc.ESTADOS if conteo_estado.get(e, 0)]
    if filas:
        series.append(_serie(
            "dt_cruce", "Estado del cruce entre fuentes", "barras_h", filas,
            "clase", "valor", unidad="hechos", eje_x="N.° de hechos",
            eje_y="Estado", colores=COLOR_ESTADO_CRUCE,
            orden_cat=list(dtc.ESTADOS),
            descripcion="«Pendiente» son los hechos que ninguna fuente "
                        "declara: no se estiman.",
            nota="Registro de hechos integrados del bloque."))

    actualizados = [r for r in (integrado.get("campos") or [])
                    if r.get("estado") == dtc.ACTUALIZADO]
    tablas = []
    if actualizados:
        tablas.append(("Hechos actualizados por campo", [{
            "Sección": r["seccion"], "Hecho": r["etiqueta"],
            "Valor en la ficha de gabinete": r["valor_alterno"],
            "Valor de campo vigente": r["valor"],
        } for r in actualizados]))
    declarados = [r for r in (integrado.get("campos") or [])
                  if r.get("estado") != dtc.PENDIENTE]
    if declarados:
        tablas.append(("Procedencia hecho por hecho", [{
            "Sección": r["seccion"], "Hecho": r["etiqueta"],
            "Valor": r["valor"], "Unidad": r["unidad"],
            "Fuente": dtc.ETIQUETA_FUENTE[r["fuente_valor"]],
            "Estado del cruce": r["estado"],
            "Valor de la otra fuente": r["valor_alterno"],
        } for r in declarados]))

    return _seccion(
        "Procedencia de los datos",
        "Trazabilidad del bloque: qué fuente sustenta cada hecho y qué "
        "actualizó la verificación de campo.", series, tablas)


# ══════════════════════════════════════════════════════════════════════════
# SECCIONES CONSOLIDADAS
# ══════════════════════════════════════════════════════════════════════════
# La mirada consolidada agrega por provincia, distrito, microcuenca o bloque.
# Cada bloque entra con lo que declara; los que no declaran un hecho no se
# cuentan en su reparto en vez de imputarles una clase.

AGRUPACIONES = {
    "provincia": "Provincia", "distrito": "Distrito",
    "microcuenca": "Microcuenca", "codigo_bloque": "Bloque",
}


def _grupo(datos, agrupacion):
    return _txt(datos.get(agrupacion)) or "Sin declarar"


def _apiladas_por_grupo(lista, agrupacion, clave, orden, colores, id_, titulo,
                        descripcion, nota):
    """Barras apiladas: cuantos bloques de cada grupo caen en cada clase."""
    filas = []
    conteo = {}
    for datos in lista:
        clase = _txt(datos.get(clave))
        if not clase:
            continue
        conteo.setdefault(_grupo(datos, agrupacion), {}).setdefault(clase, 0)
        conteo[_grupo(datos, agrupacion)][clase] += 1
    for grupo in sorted(conteo):
        for clase in list(orden) + sorted(c for c in conteo[grupo]
                                          if c not in orden):
            if conteo[grupo].get(clase):
                filas.append({"cat": grupo, "sub": clase,
                              "valor": conteo[grupo][clase]})
    if not filas:
        return None

    # Una ficha puede declarar un valor que no esta en la lista oficial
    # ("FICHA F-DT-02 SIN CONTENIDO", por ejemplo). Esa clase se grafica
    # igual, al final del orden y en el gris de «sin dato»: tomarle prestado
    # un tono de severidad seria atribuirle una gravedad que nadie declaro.
    presentes = {f["sub"] for f in filas}
    orden_sub = [c for c in orden if c in presentes]
    fuera = sorted(presentes - set(orden))
    orden_sub += fuera
    tonos = dict(colores)
    for clase in fuera:
        tonos.setdefault(clase, SIN_DATO)

    return _serie(id_, titulo, "apiladas", filas, "cat", "valor", sub="sub",
                  orden_sub=orden_sub, colores=tonos, unidad="bloques",
                  eje_x="N.° de bloques", eje_y=AGRUPACIONES[agrupacion],
                  descripcion=descripcion, nota=nota)


def _seccion_territorio(lista, agrupacion):
    """Superficie, brecha y vigor vegetal agregados por el nivel pedido."""
    etiqueta = AGRUPACIONES[agrupacion]
    acumulado = {}
    for datos in lista:
        grupo = acumulado.setdefault(_grupo(datos, agrupacion),
                                     {"area": 0.0, "bajo": 0.0, "sobre": 0.0,
                                      "n": 0, "msavi": []})
        grupo["n"] += 1
        grupo["area"] += datos.get("area_ha_num") or 0.0
        grupo["bajo"] += datos.get("msavi_bajo_umbral_ha") or 0.0
        grupo["sobre"] += datos.get("msavi_sobre_umbral_ha") or 0.0
        if datos.get("msavi_2024_num") is not None:
            grupo["msavi"].append(datos["msavi_2024_num"])
    if not acumulado:
        return None
    grupos = sorted(acumulado, key=lambda g: -acumulado[g]["area"])

    etiqueta_bajo = "Bajo umbral %s (brecha)" % rbq.UMBRAL_MSAVI
    etiqueta_sobre = "Sobre umbral %s" % rbq.UMBRAL_MSAVI
    filas_umbral = []
    for grupo in grupos:
        for clase, valor in ((etiqueta_bajo, acumulado[grupo]["bajo"]),
                             (etiqueta_sobre, acumulado[grupo]["sobre"])):
            if valor:
                filas_umbral.append({"cat": grupo, "sub": clase,
                                     "valor": round(valor, 2)})

    series = [
        _serie("dtc_area", "Superficie de catálogo por %s" % etiqueta.lower(),
               "barras_h",
               [{"clase": g, "valor": round(acumulado[g]["area"], 2)}
                for g in grupos],
               "clase", "valor", unidad="ha", decimales=2, orden_cat=grupos,
               escala="neutra", eje_x="Superficie (ha)", eje_y=etiqueta,
               descripcion="Superficie declarada en el catálogo maestro.",
               nota=FUENTE_CATALOGO),
        _serie("dtc_bloques", "Bloques por %s" % etiqueta.lower(), "barras_h",
               [{"clase": g, "valor": acumulado[g]["n"]} for g in grupos],
               "clase", "valor", unidad="bloques", orden_cat=grupos,
               escala="neutra", eje_x="N.° de bloques", eje_y=etiqueta,
               descripcion="Reparto del universo integrado.",
               nota=FUENTE_CATALOGO),
    ]
    if filas_umbral:
        series.append(_serie(
            "dtc_umbral",
            "Superficie frente al umbral del indicador por %s" % etiqueta.lower(),
            "apiladas", filas_umbral, "cat", "valor", sub="sub",
            orden_sub=[etiqueta_bajo, etiqueta_sobre], colores=COLOR_UMBRAL,
            orden_cat=grupos, unidad="ha", decimales=2,
            eje_x="Superficie (ha)", eje_y=etiqueta,
            descripcion="La fracción en tono de severidad es la brecha de la "
                        "R.M. N.° 00213-2024-MINAM.",
            nota="Distribución areal del MSAVI 2024 por clase DN."))

    con_msavi = [g for g in grupos if acumulado[g]["msavi"]]
    if con_msavi:
        series.append(_serie(
            "dtc_msavi", "MSAVI 2024 promedio por %s" % etiqueta.lower(),
            "barras_h",
            [{"clase": g,
              "valor": round(sum(acumulado[g]["msavi"]) / len(acumulado[g]["msavi"]), 4)}
             for g in con_msavi],
            "clase", "valor", decimales=4, orden_cat=con_msavi,
            escala="favorable", eje_x="MSAVI 2024 (media de los bloques)",
            eje_y=etiqueta,
            descripcion="Promedio simple de la media de cada bloque; el "
                        "umbral del indicador es %s." % rbq.UMBRAL_MSAVI,
            nota=FUENTE_CATALOGO))

    return _seccion(
        "Territorio por %s" % etiqueta.lower(),
        "Superficie, número de bloques y vigor vegetal agregados por %s."
        % etiqueta.lower(), series)


def _seccion_estado(lista, agrupacion):
    """Estado del ecosistema agregado: conservacion, erosion y urgencia."""
    etiqueta = AGRUPACIONES[agrupacion]
    series = [
        _apiladas_por_grupo(
            lista, agrupacion, "estado_conservacion", ORDEN_CONSERVACION,
            COLOR_CONSERVACION, "dtc_conservacion",
            "Estado de conservación por %s" % etiqueta.lower(),
            "A más oscuro, mayor alteración declarada.",
            "Hoja Resumen de cada ficha de bloque."),
        _apiladas_por_grupo(
            lista, agrupacion, "nivel_erosion", ORDEN_EROSION, COLOR_EROSION,
            "dtc_erosion", "Nivel de erosión por %s" % etiqueta.lower(),
            "Nivel general declarado en la verificación de campo (F-DT-02).",
            FUENTE_CAMPO),
        _apiladas_por_grupo(
            lista, agrupacion, "urgencia_intervencion", ORDEN_URGENCIA,
            COLOR_URGENCIA, "dtc_urgencia",
            "Urgencia de intervención por %s" % etiqueta.lower(),
            "Urgencia declarada en F-DT-04, insumo para la priorización.",
            FUENTE_CAMPO),
    ]
    ecosistemas = _conteo([_txt(d.get("tipo_ecosistema")) for d in lista])
    if ecosistemas:
        series.append(_serie(
            "dtc_ecosistema", "Bloques por tipo de ecosistema (UP)", "barras_h",
            ecosistemas, "clase", "valor", unidad="bloques", escala="neutra",
            eje_x="N.° de bloques", eje_y="Tipo de ecosistema",
            descripcion="Unidad productora declarada para cada bloque.",
            nota=FUENTE_CAMPO))
    return _seccion(
        "Estado del ecosistema",
        "Cómo se reparte el universo integrado entre las escalas ordenadas "
        "de conservación, erosión y urgencia.", series)


def _seccion_consistencia_consolidada(lista, agrupacion):
    """Verificaciones y procedencia agregadas por el nivel pedido."""
    etiqueta = AGRUPACIONES[agrupacion]
    filas_cal, filas_fuente = [], []
    acumulado, fuentes = {}, {}
    for datos in lista:
        grupo = _grupo(datos, agrupacion)
        resumen = datos.get("consistencia_resumen") or {}
        for calificacion in rbq.CALIFICACIONES:
            if resumen.get(calificacion):
                acumulado.setdefault(grupo, {}).setdefault(calificacion, 0)
                acumulado[grupo][calificacion] += resumen[calificacion]
        conteo = datos.get("conteo_fuente") or {}
        for fuente in dtc.FUENTES:
            if conteo.get(fuente):
                fuentes.setdefault(grupo, {}).setdefault(fuente, 0)
                fuentes[grupo][fuente] += conteo[fuente]

    for grupo in sorted(acumulado):
        for calificacion in rbq.CALIFICACIONES:
            if acumulado[grupo].get(calificacion):
                filas_cal.append({"cat": grupo, "sub": calificacion,
                                  "valor": acumulado[grupo][calificacion]})
    for grupo in sorted(fuentes):
        for fuente in dtc.FUENTES:
            if fuentes[grupo].get(fuente):
                filas_fuente.append({"cat": grupo,
                                     "sub": dtc.ETIQUETA_FUENTE[fuente],
                                     "valor": fuentes[grupo][fuente]})

    series = []
    if filas_cal:
        series.append(_serie(
            "dtc_consistencia",
            "Verificaciones de consistencia por %s" % etiqueta.lower(),
            "apiladas", filas_cal, "cat", "valor", sub="sub",
            orden_sub=list(rbq.CALIFICACIONES), colores=COLOR_CALIFICACION,
            unidad="verificaciones", eje_x="N.° de verificaciones",
            eje_y=etiqueta,
            descripcion="La fracción roja son las discrepancias sustantivas "
                        "pendientes de resolver.",
            nota="Hoja «Control de consistencia» de cada ficha de bloque."))
    if filas_fuente:
        series.append(_serie(
            "dtc_procedencia",
            "Procedencia de los hechos declarados por %s" % etiqueta.lower(),
            "apiladas", filas_fuente, "cat", "valor", sub="sub",
            orden_sub=[dtc.ETIQUETA_FUENTE[f] for f in dtc.FUENTES],
            colores=COLOR_FUENTE, unidad="hechos", eje_x="N.° de hechos",
            eje_y=etiqueta,
            descripcion="Peso relativo del levantamiento de campo frente al "
                        "gabinete y a las fuentes oficiales.",
            nota="Registro de hechos integrados de cada bloque."))
    return _seccion(
        "Consistencia y procedencia",
        "Estado documental del universo integrado: qué está verificado, qué "
        "falta y de dónde viene cada dato.", series)


# ══════════════════════════════════════════════════════════════════════════
# ARMADO DEL INFORME
# ══════════════════════════════════════════════════════════════════════════

def _cifra(valor, decimales=2, unidad=""):
    if valor is None:
        return "s/d"
    texto = f"{valor:,.{decimales}f}".replace(",", " ")
    return texto + (" " + unidad if unidad else "")


def series_de_bloque(datos, integrado=None):
    """Todas las series de un bloque, indexadas por su id.

    Permite al aplicativo dibujar una serie suelta -la de MSAVI junto a su
    tabla en la ficha, por ejemplo- con la misma declaracion que usan el
    libro Excel y el anexo PDF, sin armar el informe completo.
    """
    datos = rbq.completar_sintesis_msavi(dict(datos or {}))
    series = {}
    for seccion in (_seccion_vegetacion(datos), _seccion_microcuenca(datos),
                    _seccion_consistencia(datos), _seccion_campo(integrado),
                    _seccion_procedencia(integrado)):
        for serie in (seccion or {}).get("series", []):
            series[serie["id"]] = serie
    return series


def _metricas_bloque(datos, integrado):
    """Cifras de cabecera del bloque, en el orden en que se leen."""
    resumen = datos.get("consistencia_resumen") or {}
    msavi = datos.get("msavi_2024_num")
    bajo_pct = datos.get("bajo_umbral_pct_num")
    metricas = [
        {"etiqueta": "Superficie de catálogo",
         "valor": _cifra(datos.get("area_ha_num"), 2, "ha"),
         "detalle": "Catálogo maestro Bloques V5/V6"},
        {"etiqueta": "MSAVI 2024 — media del bloque",
         "valor": _cifra(msavi, 4),
         "detalle": _txt(datos.get("msavi_clase")) or
                    "umbral del indicador %s" % rbq.UMBRAL_MSAVI,
         "tono": ("critico" if msavi is not None and msavi < rbq.UMBRAL_MSAVI
                  else "favorable" if msavi is not None else "neutro")},
        {"etiqueta": "Brecha bajo el umbral %s" % rbq.UMBRAL_MSAVI,
         "valor": _cifra(datos.get("msavi_bajo_umbral_ha"), 2, "ha"),
         "detalle": (_cifra(bajo_pct, 2, "% del bloque")
                     if bajo_pct is not None else "sin distribución areal"),
         "tono": "critico" if bajo_pct else "neutro"},
        {"etiqueta": "Pendiente promedio",
         "valor": _cifra(datos.get("pendiente_pct_num"), 2, "%"),
         "detalle": _txt(datos.get("clase_pendiente"))},
        {"etiqueta": "Rango altitudinal",
         "valor": "%s – %s msnm" % (_cifra(datos.get("altitud_min_num"), 0),
                                    _cifra(datos.get("altitud_max_num"), 0)),
         "detalle": _txt(datos.get("piso_altitudinal"))},
        {"etiqueta": "Estado de conservación",
         "valor": _txt(datos.get("estado_conservacion")) or "s/d",
         "detalle": _txt(datos.get("tipo_ecosistema"))},
        {"etiqueta": "Verificaciones de consistencia",
         "valor": str(resumen.get("total", 0)),
         "detalle": "%d sustantiva(s) pendiente(s)" % resumen.get("SUSTANTIVA", 0),
         "tono": "critico" if resumen.get("SUSTANTIVA") else "favorable"},
        {"etiqueta": "Estado de verificación de campo",
         "valor": _txt(datos.get("estado_verificacion")) or "s/d",
         "detalle": _txt(datos.get("fecha_evaluacion"))},
    ]
    if integrado:
        conteo = integrado.get("conteo_fuente") or {}
        metricas.append({
            "etiqueta": "Hechos declarados",
            "valor": "%d / %d" % (integrado.get("n_declarados", 0),
                                  integrado.get("n_campos", 0)),
            "detalle": "%d desde campo · cobertura %.1f %%"
                       % (conteo.get(dtc.CAMPO, 0),
                          integrado.get("cobertura_pct", 0.0))})
    return metricas


def indicadores_bloque(datos, integrado=None, tema="claro"):
    """Informe analitico del Diagnostico Territorial de UN bloque.

    `datos`     salida de `resumenes_bloques.parsear_resumen_bloque`, ya
                pasada por `completar_sintesis_msavi`.
    `integrado` registro de `dt_campo.integrar_bloque`, opcional: cuando se
                entrega, el informe suma las secciones de verificacion de
                campo y de procedencia.
    """
    datos = rbq.completar_sintesis_msavi(dict(datos or {}))
    secciones = []
    for constructor in (_seccion_vegetacion, _seccion_microcuenca):
        seccion = constructor(datos)
        if seccion:
            secciones.append(seccion)
    for constructor in (_seccion_campo, _seccion_procedencia):
        seccion = constructor(integrado)
        if seccion:
            secciones.append(seccion)
    seccion = _seccion_consistencia(datos)
    if seccion:
        secciones.append(seccion)

    avisos = list(datos.get("advertencias") or [])
    validacion = _txt(datos.get("validacion_utm"))
    if validacion and validacion != "Conforme":
        avisos.append("Validación UTM 17S: %s." % validacion)
    if not (datos.get("msavi_tabla") or []):
        avisos.append("El libro no trae la distribución areal del MSAVI 2024: "
                      "no se grafica y no se estima.")
    if integrado is None:
        avisos.append("Sin ficha DT de campo integrada: el informe muestra "
                      "solo lo que declara la ficha de resumen.")

    return {
        "alcance": "bloque",
        "codigo": _txt(datos.get("codigo_bloque")),
        "datos": datos,
        "integrado": integrado,
        "identificacion": [
            ("Código del bloque", _txt(datos.get("codigo_bloque"))),
            ("Microcuenca (catálogo ANA)", _txt(datos.get("microcuenca"))),
            ("Departamento", _txt(datos.get("departamento"))),
            ("Provincia", _txt(datos.get("provincia"))),
            ("Distrito", _txt(datos.get("distrito"))),
            ("Centro poblado asociado", _txt(datos.get("centro_poblado"))),
            ("Comunidad campesina", _txt(datos.get("comunidad_campesina"))),
            ("Centroide UTM ESTE (m)", _txt(datos.get("utm_este"))),
            ("Centroide UTM NORTE (m)", _txt(datos.get("utm_norte"))),
            ("Sistema de coordenadas",
             _txt(datos.get("sistema_coordenadas"))
             or "UTM WGS 84 Zona 17S (EPSG:32717)"),
            ("Tipo de intervención", _txt(datos.get("tipo_intervencion"))),
            ("Responsable de la evaluación", _txt(datos.get("evaluador"))),
        ],
        "metricas": _metricas_bloque(datos, integrado),
        "secciones": secciones,
        "avisos": avisos,
        "generado": datetime.now(),
        "tema": tema,
    }


def _metricas_consolidado(lista, agrupacion):
    total_area = sum(d.get("area_ha_num") or 0 for d in lista)
    total_bajo = sum(d.get("msavi_bajo_umbral_ha") or 0 for d in lista)
    medias = [d["msavi_2024_num"] for d in lista
              if d.get("msavi_2024_num") is not None]
    sustantivas = sum((d.get("consistencia_resumen") or {}).get("SUSTANTIVA", 0)
                      for d in lista)
    verificaciones = sum((d.get("consistencia_resumen") or {}).get("total", 0)
                         for d in lista)
    bajo_umbral = sum(1 for m in medias if m < rbq.UMBRAL_MSAVI)
    return [
        {"etiqueta": "Bloques en el ámbito", "valor": str(len(lista)),
         "detalle": "agrupados por %s" % AGRUPACIONES[agrupacion].lower()},
        {"etiqueta": "Superficie de catálogo", "valor": _cifra(total_area, 2, "ha"),
         "detalle": "suma de los bloques del ámbito"},
        {"etiqueta": "Brecha bajo el umbral %s" % rbq.UMBRAL_MSAVI,
         "valor": _cifra(total_bajo, 2, "ha"),
         "detalle": (_cifra(100.0 * total_bajo / total_area, 2, "% del ámbito")
                     if total_area else "sin superficie declarada"),
         "tono": "critico" if total_bajo else "neutro"},
        {"etiqueta": "MSAVI 2024 promedio",
         "valor": _cifra(sum(medias) / len(medias), 4) if medias else "s/d",
         "detalle": "%d bloque(s) bajo el umbral" % bajo_umbral,
         "tono": "critico" if bajo_umbral else "favorable"},
        {"etiqueta": "Provincias",
         "valor": str(len({_txt(d.get("provincia")) for d in lista
                           if _txt(d.get("provincia"))})),
         "detalle": "%d distrito(s) · %d microcuenca(s)"
                    % (len({_txt(d.get("distrito")) for d in lista
                            if _txt(d.get("distrito"))}),
                       len({_txt(d.get("microcuenca")) for d in lista
                            if _txt(d.get("microcuenca"))}))},
        {"etiqueta": "Verificaciones de consistencia", "valor": str(verificaciones),
         "detalle": "%d sustantiva(s) pendiente(s)" % sustantivas,
         "tono": "critico" if sustantivas else "favorable"},
    ]


def indicadores_consolidado(lista, etiqueta="", agrupacion="provincia",
                            tema="claro"):
    """Informe analitico de un conjunto de bloques, agregado por un nivel."""
    if agrupacion not in AGRUPACIONES:
        raise ValueError("Agrupacion no reconocida: %s" % agrupacion)
    lista = [rbq.completar_sintesis_msavi(dict(d)) for d in (lista or [])]
    lista.sort(key=lambda d: dtc._orden_codigo(_txt(d.get("codigo_bloque"))))

    secciones = []
    for constructor in (_seccion_territorio, _seccion_estado,
                        _seccion_consistencia_consolidada):
        seccion = constructor(lista, agrupacion)
        if seccion:
            secciones.append(seccion)

    if lista:
        secciones.append({
            "titulo": "Detalle por bloque",
            "descripcion": "Una fila por bloque con lo que declara su ficha "
                           "de resumen.",
            "series": [],
            "tablas": [("Bloques del ámbito", [{
                "Bloque": _txt(d.get("codigo_bloque")),
                "Microcuenca": _txt(d.get("microcuenca")),
                "Provincia": _txt(d.get("provincia")),
                "Distrito": _txt(d.get("distrito")),
                "Centro poblado": _txt(d.get("centro_poblado")),
                "Superficie (ha)": d.get("area_ha_num"),
                "UTM ESTE": d.get("utm_este_num"),
                "UTM NORTE": d.get("utm_norte_num"),
                "Pendiente (%)": d.get("pendiente_pct_num"),
                "MSAVI 2024": d.get("msavi_2024_num"),
                "% bajo umbral": d.get("bajo_umbral_pct_num"),
                "Brecha (ha)": d.get("msavi_bajo_umbral_ha"),
                "Tipo de ecosistema": _txt(d.get("tipo_ecosistema")),
                "Estado de conservación": _txt(d.get("estado_conservacion")),
                "Nivel de erosión": _txt(d.get("nivel_erosion")),
                "Urgencia": _txt(d.get("urgencia_intervencion")),
                "Verificaciones": (d.get("consistencia_resumen") or {}).get("total", 0),
                "Sustantivas": (d.get("consistencia_resumen") or {}).get("SUSTANTIVA", 0),
            } for d in lista])],
        })

    avisos = []
    sin_msavi = [_txt(d.get("codigo_bloque")) for d in lista
                 if not (d.get("msavi_tabla") or [])]
    if sin_msavi:
        avisos.append("%d bloque(s) sin distribución areal del MSAVI 2024: %s."
                      % (len(sin_msavi), ", ".join(sin_msavi[:12])
                         + (" …" if len(sin_msavi) > 12 else "")))
    if not lista:
        avisos.append("No hay bloques en el ámbito seleccionado.")

    return {
        "alcance": "consolidado",
        "codigo": etiqueta or "%d bloques" % len(lista),
        "agrupacion": agrupacion,
        "bloques": [_txt(d.get("codigo_bloque")) for d in lista],
        "datos": {}, "integrado": None,
        "identificacion": [
            ("Ámbito", etiqueta or "todos los bloques cargados"),
            ("Nivel de agregación", AGRUPACIONES[agrupacion]),
            ("Bloques incluidos", str(len(lista))),
            ("Sistema de coordenadas", "UTM WGS 84 Zona 17S (EPSG:32717)"),
            ("Marco del indicador de brecha", "R.M. N.° 00213-2024-MINAM"),
        ],
        "metricas": _metricas_consolidado(lista, agrupacion),
        "secciones": secciones,
        "avisos": avisos,
        "generado": datetime.now(),
        "tema": tema,
    }


# ══════════════════════════════════════════════════════════════════════════
# LIBRO EXCEL CON GRAFICOS NATIVOS
# ══════════════════════════════════════════════════════════════════════════
# La portada declara de que bloque se trata y que trae el libro; cada
# seccion ocupa su propia hoja con la tabla de datos al lado de su grafico,
# y las tablas de respaldo quedan filtrables. Nada de esto reemplaza a la
# ficha de resumen: es el anexo analitico que la acompana.

def _hoja_portada(wb, informe):
    """Portada del libro: identificacion, cifras de cabecera y avisos."""
    ws = wb.active
    ws.title = "Resumen DT"
    titulo = ("DIAGNOSTICO TERRITORIAL - BLOQUE %s" % informe["codigo"]
              if informe.get("alcance") == "bloque"
              else "DIAGNOSTICO TERRITORIAL CONSOLIDADO - %s" % informe["codigo"])
    fila = _titulo_hoja(ws, titulo, ancho=6)
    for col, ancho in zip("ABCDEF", (44, 28, 40, 16, 16, 16)):
        ws.column_dimensions[col].width = ancho

    _, fila = _escribir_tabla(
        ws, fila, "1. Identificación", ["Campo", "Valor"],
        [[etiqueta, valor] for etiqueta, valor in informe["identificacion"]
         if valor not in ("", None)])
    fila += 1

    _, fila = _escribir_tabla(
        ws, fila, "2. Cifras de cabecera", ["Indicador", "Valor", "Detalle"],
        [[m["etiqueta"], m["valor"], m.get("detalle", "")]
         for m in informe.get("metricas", [])])
    fila += 1

    inventario = [[s["titulo"], len(s.get("series", [])),
                   len(s.get("tablas", [])), s.get("descripcion", "")]
                  for s in informe.get("secciones", [])]
    if inventario:
        _, fila = _escribir_tabla(
            ws, fila, "3. Contenido del libro",
            ["Sección", "Gráficos", "Tablas de respaldo", "Alcance"],
            inventario,
            subtitulo="Cada sección ocupa una hoja con su tabla de datos y su "
                      "gráfico nativo de Excel, editable por el usuario.")
        fila += 1

    for aviso in informe.get("avisos", []):
        celda = ws.cell(fila, 1, "Aviso: " + aviso)
        celda.font = Font(name="Arial", size=9, italic=True, color="9C5700")
        celda.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila,
                       end_column=6)
        fila += 1
    if informe.get("avisos"):
        fila += 1

    celda = ws.cell(fila, 1, NOTA_INTEGRIDAD)
    celda.font = Font(name="Arial", size=8, italic=True)
    celda.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila + 2,
                   end_column=6)
    ws.sheet_view.showGridLines = False


NOTA_INTEGRIDAD = (
    "Libro generado por el aplicativo IN Piura sobre las fichas de resumen "
    "por bloque y su integración con la ficha DT de campo. Solo se grafica lo "
    "declarado: ningún valor ausente se estima ni se completa por analogía, "
    "conforme a la declaración de integridad de datos del proyecto. Sistema "
    "de referencia UTM WGS 84 Zona 17S (EPSG:32717); umbral de brecha "
    "MSAVI %s conforme a la R.M. N.° 00213-2024-MINAM." % rbq.UMBRAL_MSAVI)


def generar_excel(informe):
    """Libro Excel del informe territorial, con graficos nativos. Bytes."""
    wb = Workbook()
    usados = {"resumen dt"}
    _hoja_portada(wb, informe)
    for seccion in informe.get("secciones", []):
        if seccion.get("series"):
            _hoja_seccion(wb, seccion, usados)
    for seccion in informe.get("secciones", []):
        _hoja_tablas(wb, seccion, usados)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def nombre_excel(informe):
    marca = informe["generado"].strftime("%Y%m%d_%H%M%S")
    if informe.get("alcance") == "bloque":
        return "Analitica_DT_Bloque_%s_IN_Piura_%s.xlsx" % (informe["codigo"],
                                                            marca)
    nivel = AGRUPACIONES[informe.get("agrupacion", "provincia")]
    return "Analitica_DT_Consolidado_por_%s_IN_Piura_%s.xlsx" % (nivel, marca)


# ══════════════════════════════════════════════════════════════════════════
# ANEXO GRAFICO EN PDF
# ══════════════════════════════════════════════════════════════════════════

def generar_pdf(informe):
    """Anexo grafico en PDF del informe territorial. Devuelve bytes."""
    subtitulo = ("ANEXO ANALITICO DEL DIAGNOSTICO TERRITORIAL - BLOQUE %s"
                 % informe["codigo"] if informe.get("alcance") == "bloque"
                 else "ANEXO ANALITICO DEL DIAGNOSTICO TERRITORIAL - %s"
                      % informe["codigo"])
    pdf = rbq.PDFResumen(subtitulo=subtitulo)
    pdf.alias_nb_pages()
    pdf.add_page()

    pdf.seccion("1. IDENTIFICACION")
    pdf.tabla(["Campo", "Valor"],
              [[e, v] for e, v in informe["identificacion"] if v],
              anchos_rel=[1, 2], alineaciones=["L", "L"])

    pdf.seccion("2. CIFRAS DE CABECERA")
    pdf.tabla(["Indicador", "Valor", "Detalle"],
              [[m["etiqueta"], m["valor"], m.get("detalle", "")]
               for m in informe.get("metricas", [])],
              anchos_rel=[3, 2, 4], alineaciones=["L", "R", "L"])

    for i, seccion in enumerate(informe.get("secciones", []), start=3):
        series = seccion.get("series") or []
        if not series:
            continue
        # El titulo de seccion no puede quedar solo al pie de la pagina.
        pdf._salto_si_falta(48)
        pdf.seccion("%d. %s" % (i, seccion["titulo"].upper()))
        for serie in series:
            if serie.get("forma") in ("apiladas", "agrupadas", "mapa_calor"):
                _pdf_barras_apiladas(pdf, serie)
            else:
                _pdf_barras_simples(pdf, serie)

    if informe.get("avisos"):
        pdf.seccion("AVISOS")
        for aviso in informe["avisos"]:
            pdf.nota(aviso)
    pdf.nota(NOTA_INTEGRIDAD)
    return rbq.pdf_bytes(pdf)


def nombre_pdf(informe):
    marca = informe["generado"].strftime("%Y%m%d_%H%M%S")
    if informe.get("alcance") == "bloque":
        return "Anexo_Analitico_DT_Bloque_%s_IN_Piura_%s.pdf" % (
            informe["codigo"], marca)
    nivel = AGRUPACIONES[informe.get("agrupacion", "provincia")]
    return "Anexo_Analitico_DT_Consolidado_por_%s_IN_Piura_%s.pdf" % (nivel,
                                                                      marca)
