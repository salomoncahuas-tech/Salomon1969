"""Excel ANIN de bloques V6 y área de influencia (AI) por distrito — Proyecto IN Piura.

Lee datos/gis/IN_Piura_area_influencia_bloques_v6.gpkg (EPSG:32717) y escribe
mapas/salidas/IN_Piura_Bloques_AI_por_Distrito_V6.xlsx con cinco hojas:
  Resumen_Distrito, Resumen_Provincia  -> totales con fórmulas (COUNTIFS / SUMIFS sobre las hojas de detalle)
  Bloques_AI                           -> un registro por bloque (AI sumado con SUMIF desde AI_Poligonos)
  AI_Poligonos                         -> los 130 polígonos del área de influencia aprobada
  Notas                                -> fuentes, exclusiones y controles

Uso:  python mapas/excel_bloques_ai_distrito.py
"""
import sys
from pathlib import Path

import geopandas as gpd
import openpyxl
from openpyxl.styles import Alignment, Font

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / ".claude" / "skills" / "anin-in-piura" / "scripts"))
from anin_catalogo import DISTRITOS, PROVINCIAS, AREA_TOTAL_HA, cargar_bloques  # noqa: E402
from anin_excel import VERDE, encabezado_institucional, escribir_tabla  # noqa: E402
from anin_utm import validar_utm  # noqa: E402

GPKG = RAIZ / "datos" / "gis" / "IN_Piura_area_influencia_bloques_v6.gpkg"
SALIDA = RAIZ / "mapas" / "salidas" / "IN_Piura_Bloques_AI_por_Distrito_V6.xlsx"
FUENTE = ("Bloques V6 (Bloques V6.kml) y área de influencia aprobada (AI_aprobado_2.shp), ANIN-DIME-SESDI 2026; "
          "sin los bloques retirados 74 y 75 ni sus AI")
HA = "#,##0.000"
PCT = "0.0%"


def oficial(nombre, tabla):
    """'SAN MIGUEL DE EL FAIQUE' -> 'San Miguel de El Faique' con el catálogo de la skill."""
    return tabla[nombre.strip().lower().replace(" ", "_")]


def cargar():
    b = gpd.read_file(GPKG, layer="bloques_v6")
    ai = gpd.read_file(GPKG, layer="area_influencia_v6")
    for g in (b, ai):
        g["PROV"] = g.NOMBPROV.map(lambda n: oficial(n, PROVINCIAS))
        g["DIST"] = g.NOMBDIST.map(lambda n: oficial(n, DISTRITOS))

    # controles contra el catálogo V6 vigente
    cat = {x["codigo"]: x for x in cargar_bloques()}
    assert set(b.BLOQUE) == set(cat), "La capa de bloques no coincide con el catálogo V6"
    assert set(ai.BLOQUE) == set(b.BLOQUE), "Hay bloques sin AI o AI sin bloque"
    assert abs(b.AREA_HA.sum() - AREA_TOTAL_HA) < 0.01
    difs = [(r.BLOQUE, r.DIST, cat[r.BLOQUE]["distrito"]) for r in b.itertuples()
            if r.DIST != cat[r.BLOQUE]["distrito"] or abs(r.AREA_HA - cat[r.BLOQUE]["area_ha"]) > 0.001]
    assert not difs, f"Distrito o área distinta del catálogo: {difs}"

    c = b.geometry.representative_point()
    b["ESTE"], b["NORTE"] = c.x.round(1), c.y.round(1)
    malos = [r.BLOQUE for r in b.itertuples() if not validar_utm(r.ESTE, r.NORTE)[0]]
    assert not malos, f"Coordenadas fuera de rango: {malos}"
    return b, ai


def orden(df, *extra):
    return df.sort_values(["PROV", "DIST", *extra], key=lambda s: s.astype(str)).reset_index(drop=True)


def hoja_ai(wb, ai):
    ws = wb.create_sheet("AI_Poligonos")
    cols = ["N.°", "Código AI", "Bloque", "Microcuenca", "Provincia", "Distrito",
            "Área AI (ha)", "Área geométrica (ha)", "Diferencia (ha)"]
    fila0 = encabezado_institucional(ws, "ÁREA DE INFLUENCIA APROBADA — POLÍGONOS (AI_aprobado_2)", len(cols), FUENTE)
    d = orden(ai, "BLOQUE", "AREA_HA")
    filas = []
    for i, r in d.iterrows():
        f = fila0 + 1 + i
        filas.append([i + 1, r.CODIGO, r.BLOQUE, r.MICROC, r.PROV, r.DIST, r.AREA_HA, r.AREA_GEOM_HA,
                      f"=H{f}-G{f}"])
    escribir_tabla(ws, fila0, cols, filas,
                   anchos={"N.°": 6, "Código AI": 14, "Bloque": 10, "Microcuenca": 14, "Provincia": 14,
                           "Distrito": 26, "Área AI (ha)": 14, "Área geométrica (ha)": 16, "Diferencia (ha)": 14},
                   totales={"Área AI (ha)": "SUM", "Área geométrica (ha)": "SUM", "Diferencia (ha)": "SUM"},
                   formato_num={"Área AI (ha)": HA, "Área geométrica (ha)": HA, "Diferencia (ha)": HA})
    return ws, fila0, fila0 + len(filas)


def hoja_bloques(wb, b, ai_rng):
    ws = wb.create_sheet("Bloques_AI")
    cols = ["N.°", "Bloque", "Microcuenca", "Provincia", "Distrito", "Área bloque (ha)", "Polígonos AI",
            "Área AI (ha)", "Bloque + AI (ha)", "AI / bloque", "ESTE (m)", "NORTE (m)"]
    fila0 = encabezado_institucional(ws, "BLOQUES V6 Y SU ÁREA DE INFLUENCIA", len(cols), FUENTE)
    a0, a1 = ai_rng
    rc = f"AI_Poligonos!$C${a0 + 1}:$C${a1}"
    rg = f"AI_Poligonos!$G${a0 + 1}:$G${a1}"
    d = orden(b, "BLOQUE")
    filas = []
    for i, r in d.iterrows():
        f = fila0 + 1 + i
        filas.append([i + 1, r.BLOQUE, r.MICROC, r.PROV, r.DIST, r.AREA_HA,
                      f"=COUNTIF({rc},B{f})", f"=SUMIF({rc},B{f},{rg})", f"=F{f}+H{f}",
                      f"=IF(F{f}=0,0,H{f}/F{f})", r.ESTE, r.NORTE])
    escribir_tabla(ws, fila0, cols, filas,
                   anchos={"N.°": 6, "Bloque": 10, "Microcuenca": 14, "Provincia": 14, "Distrito": 26,
                           "Área bloque (ha)": 15, "Polígonos AI": 11, "Área AI (ha)": 14,
                           "Bloque + AI (ha)": 15, "AI / bloque": 11, "ESTE (m)": 13, "NORTE (m)": 14},
                   totales={"Área bloque (ha)": "SUM", "Polígonos AI": "SUM", "Área AI (ha)": "SUM",
                            "Bloque + AI (ha)": "SUM"},
                   formato_num={"Área bloque (ha)": HA, "Área AI (ha)": HA, "Bloque + AI (ha)": HA,
                                "AI / bloque": "0.00", "ESTE (m)": "#,##0.0", "NORTE (m)": "#,##0.0"})
    tot = fila0 + len(filas) + 1
    ws.cell(tot, 10, f"=IF(F{tot}=0,0,H{tot}/F{tot})").number_format = "0.00"
    ws.cell(tot + 2, 1, "ESTE / NORTE: punto interior del polígono del bloque, UTM WGS84 Zona 17S (EPSG:32717). "
            "AI / bloque: hectáreas de AI por hectárea de bloque.").font = Font(name="Arial", size=8, italic=True)
    return ws, fila0, fila0 + len(filas)


def hoja_resumen(wb, nombre, titulo, grupos, claves, b_rng, ai_rng, por_distrito):
    """grupos: lista de tuplas (provincia[, distrito]) ya ordenadas. Fórmulas sobre las hojas de detalle."""
    ws = wb.create_sheet(nombre, 0)
    cols = (["N.°", "Provincia", "Distrito"] if por_distrito else ["N.°", "Provincia"]) + [
        "N.° bloques", "Área bloques (ha)", "Polígonos AI", "Área AI (ha)", "Bloques + AI (ha)",
        "AI / bloques", "% del área de bloques", "% del área AI"]
    L = {c: openpyxl.utils.get_column_letter(j) for j, c in enumerate(cols, start=1)}
    fila0 = encabezado_institucional(ws, titulo, len(cols), FUENTE)
    b0, b1 = b_rng
    a0, a1 = ai_rng
    bp, bd, ba = (f"Bloques_AI!${c}${b0 + 1}:${c}${b1}" for c in "DEF")
    ap, ad, aa = (f"AI_Poligonos!${c}${a0 + 1}:${c}${a1}" for c in "EFG")
    tot = fila0 + len(grupos) + 1
    filas = []
    for i, g in enumerate(grupos):
        f = fila0 + 1 + i
        cp = f"{L['Provincia']}{f}"
        cb = f"{bp},{cp}" + (f",{bd},{L['Distrito']}{f}" if por_distrito else "")
        ca = f"{ap},{cp}" + (f",{ad},{L['Distrito']}{f}" if por_distrito else "")
        fila = [i + 1, *g,
                f"=COUNTIFS({cb})", f"=SUMIFS({ba},{cb})", f"=COUNTIFS({ca})", f"=SUMIFS({aa},{ca})",
                f"={L['Área bloques (ha)']}{f}+{L['Área AI (ha)']}{f}",
                f"=IF({L['Área bloques (ha)']}{f}=0,0,{L['Área AI (ha)']}{f}/{L['Área bloques (ha)']}{f})",
                f"={L['Área bloques (ha)']}{f}/{L['Área bloques (ha)']}${tot}",
                f"={L['Área AI (ha)']}{f}/{L['Área AI (ha)']}${tot}"]
        filas.append(fila)
    sumas = ["N.° bloques", "Área bloques (ha)", "Polígonos AI", "Área AI (ha)", "Bloques + AI (ha)",
             "% del área de bloques", "% del área AI"]
    escribir_tabla(ws, fila0, cols, filas,
                   anchos={"N.°": 6, "Provincia": 15, "Distrito": 26, "N.° bloques": 11, "Área bloques (ha)": 16,
                           "Polígonos AI": 11, "Área AI (ha)": 15, "Bloques + AI (ha)": 16, "AI / bloques": 11,
                           "% del área de bloques": 13, "% del área AI": 12},
                   totales={c: "SUM" for c in sumas},
                   formato_num={"Área bloques (ha)": HA, "Área AI (ha)": HA, "Bloques + AI (ha)": HA,
                                "AI / bloques": "0.00", "% del área de bloques": PCT, "% del área AI": PCT})
    ws.cell(tot, 2, f"{len(grupos)} {'distritos' if por_distrito else 'provincias'}").font = \
        Font(name="Arial", size=10, bold=True)
    c = ws.cell(tot, cols.index("AI / bloques") + 1,
                f"=IF({L['Área bloques (ha)']}{tot}=0,0,{L['Área AI (ha)']}{tot}/{L['Área bloques (ha)']}{tot})")
    c.number_format = "0.00"
    # controles: el total debe cuadrar con las hojas de detalle
    r = tot + 2
    ws.cell(r, 2, "Control").font = Font(name="Arial", size=9, bold=True, color=VERDE)
    controles = [
        (f"Total bloques = 117", f"=IF({L['N.° bloques']}{tot}=117,\"OK\",\"REVISAR\")"),
        (f"Área bloques = Bloques_AI", f"=IF(ABS({L['Área bloques (ha)']}{tot}-SUM({ba}))<0.001,\"OK\",\"REVISAR\")"),
        (f"Área AI = AI_Poligonos", f"=IF(ABS({L['Área AI (ha)']}{tot}-SUM({aa}))<0.001,\"OK\",\"REVISAR\")"),
    ]
    for k, (txt, fml) in enumerate(controles, start=1):
        ws.cell(r + k, 2, txt).font = Font(name="Arial", size=9)
        ws.cell(r + k, 4 if por_distrito else 3, fml).font = Font(name="Arial", size=9, bold=True)
    ws.cell(r + len(controles) + 2, 2, "AI / bloques: hectáreas de área de influencia por hectárea de bloque. "
            "El AI se asigna al distrito de su bloque.").font = Font(name="Arial", size=8, italic=True)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def hoja_notas(wb, b, ai):
    ws = wb.create_sheet("Notas")
    fila = encabezado_institucional(ws, "NOTAS, FUENTES Y CONTROLES", 2)
    notas = [
        ("Bloques", f"117 bloques V6 vigentes, {b.AREA_HA.sum():,.3f} ha; coinciden en código, distrito y área "
                    "con el catálogo datos/unidades_liberacion_areas.csv."),
        ("Área de influencia", f"Capa AI_aprobado_2.shp (CODIGO = AI_<bloque>): {len(ai)} polígonos vigentes, "
                               f"{ai.AREA_HA.sum():,.3f} ha. Los 117 bloques tienen AI; algunos tienen más de un "
                               "polígono (p. ej. M28B2 con 4)."),
        ("Exclusiones", "Se excluyen los bloques retirados en V6, 74 y 75, y sus áreas de influencia AI_74 "
                        "(30.073 ha) y AI_75 (125.819 ha). Por eso el AI vigente no es 17,295.547 ha (total del "
                        "shapefile), sino 17,139.655 ha."),
        ("Áreas", "Se usa el campo AREA_HA de la capa. El área recalculada de las geometrías suma "
                  f"{ai.AREA_GEOM_HA.sum():,.3f} ha ({ai.AREA_GEOM_HA.sum() - ai.AREA_HA.sum():+,.3f} ha); la "
                  "diferencia por polígono está en AI_Poligonos. Conviene recalcular AREA_HA en el SIG."),
        ("Asignación territorial", "Cada polígono AI hereda la provincia y el distrito de su bloque. Cada AI "
                                   "colinda con su bloque sin superponerse."),
        ("Coordenadas", "UTM WGS84 Zona 17S (EPSG:32717). ESTE/NORTE de cada bloque = punto interior del "
                        "polígono, validado en los rangos del proyecto (ESTE 450,000–750,000; NORTE "
                        "9,300,000–9,600,000)."),
        ("Mapas", "Láminas A3 en mapas/salidas/IN_Piura_Mapa_Area_Influencia_Bloques_V6.pdf."),
        ("Generado con", "mapas/excel_bloques_ai_distrito.py (los totales son fórmulas sobre las hojas de detalle)."),
    ]
    for k, (t, v) in enumerate(notas):
        a = ws.cell(fila + k, 1, t)
        a.font = Font(name="Arial", size=10, bold=True, color=VERDE)
        a.alignment = Alignment(vertical="top")
        c = ws.cell(fila + k, 2, v)
        c.font = Font(name="Arial", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[fila + k].height = 44
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110


def main():
    b, ai = cargar()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _, a0, a1 = hoja_ai(wb, ai)
    _, b0, b1 = hoja_bloques(wb, b, (a0, a1))
    dist = orden(b)[["PROV", "DIST"]].drop_duplicates()
    hoja_resumen(wb, "Resumen_Provincia", "BLOQUES V6 Y ÁREA DE INFLUENCIA POR PROVINCIA",
                 [(p,) for p in dist.PROV.unique()], None, (b0, b1), (a0, a1), por_distrito=False)
    hoja_resumen(wb, "Resumen_Distrito", "BLOQUES V6 Y ÁREA DE INFLUENCIA POR DISTRITO",
                 list(dist.itertuples(index=False, name=None)), None, (b0, b1), (a0, a1), por_distrito=True)
    hoja_notas(wb, b, ai)
    wb.move_sheet("Bloques_AI", offset=-1)  # detalle por bloque antes que el de polígonos
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True  # Excel calcula las fórmulas al abrir (openpyxl no guarda resultados)
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    print(SALIDA)


if __name__ == "__main__":
    main()
