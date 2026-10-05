"""Mapas de ecosistemas y los 117 bloques de intervención V6 — Proyecto IN Piura.

Cruza los bloques V6 (datos/gis/IN_Piura_area_influencia_bloques_v6.gpkg, capa bloques_v6) con una capa de
ecosistemas (Mapa Nacional de Ecosistemas del Perú, MINAM 2018, R.M. N.° 440-2018-MINAM, o un recorte de ella)
y genera en mapas/salidas/:
  - IN_Piura_Mapa_Ecosistemas_Bloques_V6.pdf y PNG por lámina (A3 horizontal):
      lámina 1  mapa general de ecosistemas del ámbito (15 distritos) con los 117 bloques y cuadro de superficies
      láminas 2-4  detalle por provincia (Ayabaca, Huancabamba, Morropón) con el código de cada bloque
  - IN_Piura_Ecosistemas_por_Bloque_V6.xlsx (formato ANIN, totales y porcentajes con fórmulas)
  - IN_Piura_ecosistemas_bloques_v6.gpkg (capas ecosistemas_ambito y ecosistemas_bloques_v6, EPSG:32717)

La capa de ecosistemas no viene en el repositorio. Se busca en datos/gis/ un archivo cuyo nombre contenga
«ecosistem» (.gpkg, .shp, .geojson o un .zip con el shapefile), o se indica con --ecosistemas. El campo del nombre
y el del símbolo se detectan solos (ECOSISTEMA, SIMBOLO, SIMB_ECOS, ...) o se fijan con --campo / --campo-simbolo.

Uso:
  python mapas/mapa_ecosistemas_bloques_v6.py
  python mapas/mapa_ecosistemas_bloques_v6.py --ecosistemas datos/gis/ecosistemas_minam_2018.zip --campo ECOSISTEMA
Requiere geopandas, matplotlib, adjustText y openpyxl.
"""
import argparse
import sys
import textwrap
import zipfile
from pathlib import Path

import geopandas as gpd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
import pandas as pd
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.lines import Line2D
from matplotlib.patches import Patch, Rectangle
from openpyxl.styles import Alignment, Font
from shapely.geometry import box

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(RAIZ / ".claude" / "skills" / "anin-in-piura" / "scripts"))
from mapa_area_influencia_bloques_v6 import (  # noqa: E402
    BLOQUE_EDGE, DORADO, NOMBRE_PROV, TOTAL_V6_HA, VERDE, cargar, ejes_utm, encuadre, etiquetas_distritos,
    mapa_ubicacion, marco, norte_y_escala, titulo_cdl)
from anin_excel import encabezado_institucional, escribir_tabla  # noqa: E402

DIR_GIS = RAIZ / "datos" / "gis"
SALIDA = RAIZ / "mapas" / "salidas"
CAMPOS_NOMBRE = ["ECO_LAYER", "ECO_DETALL", "ECOSISTEMA", "ECOSISTEMAS", "NOM_ECOS", "NOMB_ECOS", "ECOSIS", "NOMBRE", "DESCRIPCIO", "LEYENDA"]
CAMPOS_SIMBOLO = ["SIMBOLO", "SIMB_ECOS", "SIMB", "SIMBOL", "COD_ECOS", "COD_ECO", "CODIGO"]
SIN_DATO = "Sin información de ecosistema"
HA = "#,##0.000"
PCT = "0.0%"

# color por palabra clave del nombre (la primera coincidencia manda); lo que no coincide toma la paleta tab20
COLORES_CLAVE = [
    ("ribere", "#8DB596"), ("llanura", "#E8D79A"), ("estacionalmente seco", "#C9A55A"), ("seco", "#C9A55A"),
    ("matorral", "#D8B48F"), ("relicto", "#2E7D4F"), ("altimontano", "#3E8E5E"), ("basimontano", "#7FB77E"),
    ("montan", "#58A06A"), ("yunga", "#58A06A"), ("bosque", "#4F9A5B"), ("páramo", "#9C8FC4"),
    ("paramo", "#9C8FC4"), ("jalca", "#B9AADB"), ("pajonal", "#C9C27A"), ("bofedal", "#5DA9A6"),
    ("agr", "#F3E58A"), ("urban", "#A6A6A6"), ("río", "#6FA8DC"), ("rio", "#6FA8DC"), ("lag", "#6FA8DC"),
    ("cuerpo", "#6FA8DC"), ("minera", "#C9B1A0"), ("desierto", "#EADCC8"), (SIN_DATO.lower(), "#FFFFFF"),
]


# ---------------------------------------------------------------- datos

def buscar_capa(ruta):
    if ruta:
        p = Path(ruta)
        if not p.exists():
            sys.exit(f"No existe la capa de ecosistemas: {p}")
        return p
    cands = sorted(p for p in DIR_GIS.rglob("*")
                   if "ecosistem" in p.name.lower() and p.suffix.lower() in (".gpkg", ".shp", ".geojson", ".zip"))
    if not cands:
        sys.exit("No se encontró la capa de ecosistemas.\n"
                 f"Copie el Mapa Nacional de Ecosistemas (MINAM 2018) o su recorte al ámbito en {DIR_GIS}\n"
                 "con un nombre que contenga «ecosistemas» (p. ej. ecosistemas_minam_2018.zip) o use --ecosistemas.")
    return cands[0]


def leer_ecosistemas(ruta, campo, campo_simb, ambito, bloques):
    fuente = ruta
    if ruta.suffix.lower() == ".zip":  # el .shp puede estar dentro de una subcarpeta del zip
        with zipfile.ZipFile(ruta) as z:
            shp = [n for n in z.namelist() if n.lower().endswith(".shp")]
        fuente = f"zip://{ruta}!{shp[0]}" if len(shp) == 1 else f"zip://{ruta}"
    eco = gpd.read_file(fuente)
    if eco.crs is None:
        sys.exit(f"{ruta.name} no tiene sistema de coordenadas (.prj). No se asume ninguno: agregue el .prj.")
    print(f"Capa de ecosistemas: {ruta.name} · {len(eco)} polígonos · CRS {eco.crs.to_string()}")
    cols = {c.upper(): c for c in eco.columns}
    campo = campo or next((cols[c] for c in CAMPOS_NOMBRE if c in cols), None)
    campo_simb = campo_simb or next((cols[c] for c in CAMPOS_SIMBOLO if c in cols and cols[c] != campo), None)
    if not campo and not campo_simb:
        sys.exit(f"No se reconoce el campo del ecosistema. Campos disponibles: {list(eco.columns)}. Use --campo.")
    eco["ECOSISTEMA"] = eco[campo or campo_simb].astype(str).str.strip()
    eco["SIMBOLO"] = eco[campo_simb].astype(str).str.strip() if campo_simb else ""
    print(f"  campo de nombre: {campo or campo_simb} · campo de símbolo: {campo_simb or '(ninguno)'}")
    eco = eco[["ECOSISTEMA", "SIMBOLO", "geometry"]].to_crs(32717)
    eco["geometry"] = eco.geometry.make_valid()
    eco = eco[eco.geometry.notna() & ~eco.geometry.is_empty]
    # recorte al ámbito de 15 distritos más los bloques (la capa INEI simplificada deja fuera bordes de algunos)
    amb = gpd.GeoDataFrame(geometry=[ambito.union_all().union(bloques.union_all())], crs=ambito.crs)
    eco = eco[eco.intersects(amb.geometry.iloc[0])]
    if eco.empty:
        sys.exit("La capa de ecosistemas no cubre el ámbito del proyecto (¿CRS o zona equivocados?).")
    eco_amb = gpd.clip(eco, amb).explode(index_parts=False)
    eco_amb = eco_amb[eco_amb.geom_type.isin(["Polygon", "MultiPolygon"])]
    eco_amb = eco_amb.dissolve(["ECOSISTEMA", "SIMBOLO"]).reset_index()
    eco_amb["AREA_HA"] = (eco_amb.area / 10_000).round(3)
    return eco_amb


def cruzar(b, eco):
    """Intersección bloque × ecosistema; lo que el mapa de ecosistemas no cubre queda como SIN_DATO."""
    bb = b[["BLOQUE", "MICROC", "NOMBDIST", "NOMBPROV", "AREA_HA", "AREA_GEOM_HA", "geometry"]].rename(
        columns={"AREA_HA": "AREA_BLOQUE_HA"})
    x = gpd.overlay(bb, eco[["ECOSISTEMA", "SIMBOLO", "geometry"]], how="intersection", keep_geom_type=True)
    falta = gpd.overlay(bb, eco[["geometry"]].dissolve(), how="difference", keep_geom_type=True)
    falta = falta[falta.area > 10]  # < 0.001 ha: ruido de bordes
    if len(falta):
        falta["ECOSISTEMA"], falta["SIMBOLO"] = SIN_DATO, ""
        x = pd.concat([x, falta], ignore_index=True)
    x = x.dissolve(["BLOQUE", "ECOSISTEMA", "SIMBOLO"], aggfunc="first").reset_index()
    x["HA_GEOM"] = (x.area / 10_000).round(4)
    x = x[x.HA_GEOM > 0]
    suma = x.groupby("BLOQUE").HA_GEOM.sum()
    dif = (suma - b.set_index("BLOQUE").AREA_GEOM_HA).abs()
    assert dif.max() < 0.05, f"La intersección no reconstruye el área de los bloques: {dif[dif >= 0.05].to_dict()}"
    assert set(x.BLOQUE) == set(b.BLOQUE), "Hay bloques sin registro en el cruce con ecosistemas"
    x["PCT"] = x.HA_GEOM / x.BLOQUE.map(suma)
    x["HA_OFICIAL"] = x.PCT * x.AREA_BLOQUE_HA  # prorrateo al área oficial del catálogo (suma 12,270.235 ha)
    return gpd.GeoDataFrame(x, geometry="geometry", crs=b.crs)


def paleta(nombres):
    tab = plt.get_cmap("tab20").colors
    out, k = {}, 0
    for n in nombres:
        c = next((col for clave, col in COLORES_CLAVE if clave in n.lower()), None)
        if c is None or c in out.values() and n != SIN_DATO:
            c = matplotlib.colors.to_hex(tab[k % 20]); k += 1
        out[n] = c
    return out


def etiqueta(r):
    return f"{r.ECOSISTEMA} ({r.SIMBOLO})" if r.SIMBOLO and r.SIMBOLO != r.ECOSISTEMA else r.ECOSISTEMA


# ---------------------------------------------------------------- láminas

def pie(fig, fuente_eco):
    fig.text(0.02, 0.034, "Sistema de coordenadas: UTM WGS84 Zona 17S (EPSG:32717)", fontsize=7.5,
             color="#333333", fontweight="bold")
    fig.text(0.02, 0.019, f"Fuente: Ecosistemas — {fuente_eco}. Bloques de intervención V6 (ANIN-DIME-SESDI, 2026; "
             "sin los retirados 74 y 75). Límites distritales: INEI (capa simplificada, referencial). "
             "Elaboración: ANIN-DIME-SESDI.", fontsize=6.5, color="#555555")


def dibujar(ax, b, eco, amb, ctx, colores, xlim, ylim):
    ax.set_facecolor("white")
    ctx.plot(ax=ax, facecolor="#F2F2F2", edgecolor="#BFBFBF", linewidth=0.3, zorder=1)
    amb.plot(ax=ax, facecolor="white", edgecolor="none", zorder=1.5)  # lo que la capa no cubre queda en blanco
    eco.plot(ax=ax, color=eco.ECOSISTEMA.map(colores), edgecolor="none", zorder=2, rasterized=True)
    amb.boundary.plot(ax=ax, color="#7F8C86", linewidth=0.45, zorder=3)
    ctx[ctx.NOMBDEP == "PIURA"].dissolve("NOMBPROV").boundary.plot(ax=ax, color="#4D5656", linewidth=1.1, zorder=3)
    amb.dissolve().boundary.plot(ax=ax, color=VERDE, linewidth=2.2, zorder=4)
    b.plot(ax=ax, facecolor="none", edgecolor="white", linewidth=2.0, zorder=6)
    b.plot(ax=ax, facecolor="none", edgecolor="#C0392B", linewidth=0.9, zorder=6.5)
    ax.set_xlim(*xlim)
    ax.set_ylim(*ylim)


def leyenda(panel, y, eco_orden, colores, tam=7.6):
    elementos = [Patch(facecolor=colores[e], edgecolor="#7F8C86", linewidth=0.4, label=lab)
                 for e, lab in eco_orden if e != SIN_DATO]  # el vacío va como «Ámbito sin cobertura»
    elementos += [
        Patch(facecolor="none", edgecolor="#C0392B", linewidth=1.4, label="Bloque de intervención V6"),
        Line2D([], [], color=VERDE, linewidth=2.2, label="Límite del ámbito (15 distritos)"),
        Line2D([], [], color="#4D5656", linewidth=1.1, label="Límite provincial"),
        Line2D([], [], color="#7F8C86", linewidth=0.5, label="Límite distrital"),
        Patch(facecolor="white", edgecolor="#7F8C86", linewidth=0.4,
              label="Ámbito sin cobertura de la capa de ecosistemas"),
        Patch(facecolor="#F2F2F2", edgecolor="#BFBFBF", label="Fuera del ámbito (contexto)"),
    ]
    panel.legend(handles=elementos, loc="upper left", bbox_to_anchor=(0.0, y), fontsize=tam, frameon=False,
                 title="LEYENDA — ECOSISTEMAS", title_fontproperties={"weight": "bold", "size": 9},
                 alignment="left", handlelength=2.0, handleheight=0.9, labelspacing=0.3)


def cuadro(panel, filas, y_top, titulo, cab):
    alto = 0.0185 * (len(filas) + 1)
    t = panel.table(cellText=filas, colLabels=cab, colWidths=[0.52, 0.16, 0.16, 0.16], cellLoc="left",
                    bbox=[0.0, y_top - alto, 1.0, alto])
    t.auto_set_font_size(False)
    t.set_fontsize(6.6)
    for (i, j), c in t.get_celld().items():
        c.set_edgecolor("#B0B7B3"); c.set_linewidth(0.4)
        if j >= 1:
            c.set_text_props(ha="right")
        if i == 0:
            c.set_facecolor(VERDE); c.set_text_props(color="white", fontweight="bold", ha="center")
        elif i == len(filas):
            c.set_facecolor("#D5E8D4"); c.set_text_props(fontweight="bold")
        elif i % 2 == 0:
            c.set_facecolor("#F2F7F2")
    panel.text(0.0, y_top + 0.01, titulo, fontsize=9, fontweight="bold", color=VERDE, transform=panel.transAxes)
    return y_top - alto


def corto(texto, n=44):
    return texto if len(texto) <= n else texto[:n - 1] + "…"


def tabla_ecosistemas(x, filtro_prov=None, ancho=44):
    xs = x if filtro_prov is None else x[x.NOMBPROV == filtro_prov]
    r = xs.groupby("ECOSISTEMA").agg(ha=("HA_OFICIAL", "sum"), n=("BLOQUE", "nunique")).sort_values("ha",
                                                                                                    ascending=False)
    tot = r.ha.sum()
    filas = [[corto(e, ancho), f"{int(v.n)}", f"{v.ha:,.2f}", f"{v.ha / tot:.1%}"] for e, v in r.iterrows()]
    filas.append(["TOTAL", f"{xs.BLOQUE.nunique()}", f"{tot:,.3f}", "100.0%"])
    return filas


def lamina_general(b, x, eco, amb, ctx, colores, eco_orden, fuente_eco):
    fig = plt.figure(figsize=(16.54, 11.69))
    marco(fig, "MAPA DE ECOSISTEMAS Y BLOQUES DE INTERVENCIÓN V6",
          f"Ecosistemas del ámbito de 15 distritos y su presencia en los 117 bloques de intervención "
          f"({TOTAL_V6_HA:,.3f} ha) — Ayabaca, Huancabamba y Morropón, Región Piura", 1)
    ax = fig.add_axes([0.04, 0.075, 0.62, 0.755])
    xlim, ylim = encuadre(amb, 0.03, aspecto=0.62 * 16.54 / (0.755 * 11.69))
    dibujar(ax, b, eco, amb, ctx, colores, xlim, ylim)
    etiquetas_distritos(ax, amb, xlim, ylim, tam=7.2)
    ejes_utm(ax, paso=20000)
    norte_y_escala(ax, 20)

    panel = fig.add_axes([0.69, 0.065, 0.29, 0.77]); panel.axis("off")
    mapa_ubicacion(fig.add_axes([0.70, 0.65, 0.27, 0.18]), ctx, amb)
    y = cuadro(panel, tabla_ecosistemas(x), 0.70, "ECOSISTEMAS EN LOS BLOQUES V6",
               ["Ecosistema", "Bloques", "Área (ha)", "%"])
    leyenda(panel, y - 0.01, eco_orden, colores, tam=7.2 if len(eco_orden) > 12 else 7.6)
    pie(fig, fuente_eco)
    return fig


def lamina_provincia(b, x, eco, amb, ctx, colores, eco_orden, fuente_eco, prov, lamina):
    bp, aip = b[b.NOMBPROV == prov], amb[amb.NOMBPROV == prov]
    xp = x[x.NOMBPROV == prov]
    fig = plt.figure(figsize=(16.54, 11.69))
    marco(fig, f"ECOSISTEMAS Y BLOQUES V6 — PROVINCIA DE {NOMBRE_PROV[prov].upper()}",
          f"{len(bp)} bloques de intervención ({bp.AREA_HA.sum():,.3f} ha) en {aip.NOMBDIST.nunique()} distrito(s) "
          f"· {xp.ECOSISTEMA.nunique()} ecosistema(s) presentes en los bloques", lamina)
    ax = fig.add_axes([0.04, 0.075, 0.68, 0.755])
    xmin, ymin, xmax, ymax = aip.total_bounds
    caja = gpd.GeoSeries([box(xmin, ymin - (ymax - ymin) * 0.10, xmax, ymax)], crs=aip.crs)
    xlim, ylim = encuadre(caja, 0.04, aspecto=0.68 * 16.54 / (0.755 * 11.69))
    dibujar(ax, b, eco, amb, ctx, colores, xlim, ylim)
    etiquetas_distritos(ax, amb, xlim, ylim, tam=8.0)
    from adjustText import adjust_text
    pts = bp.geometry.representative_point()
    textos = [ax.text(pt.x, pt.y, cod, fontsize=6.6, fontweight="bold", color="#5B1A12", ha="center",
                      va="center", zorder=8,
                      bbox=dict(boxstyle="round,pad=0.12", facecolor="white", edgecolor=BLOQUE_EDGE,
                                linewidth=0.4, alpha=0.9))
              for cod, pt in zip(bp.BLOQUE, pts)]
    adjust_text(textos, x=list(pts.x), y=list(pts.y), ax=ax, expand=(1.25, 1.4), force_text=(0.4, 0.6),
                force_static=(0.3, 0.4), ensure_inside_axes=True, max_move=None,
                arrowprops=dict(arrowstyle="-", color=BLOQUE_EDGE, linewidth=0.45, zorder=7))
    w = xlim[1] - xlim[0]
    ejes_utm(ax, paso=10000 if w < 120000 else 20000)
    norte_y_escala(ax, next(k for k in (20, 10, 5, 2) if k * 1000 <= w / 7))

    panel = fig.add_axes([0.745, 0.065, 0.235, 0.77]); panel.axis("off")
    ax_ub = fig.add_axes([0.75, 0.665, 0.225, 0.165])
    mapa_ubicacion(ax_ub, ctx, amb)
    ax_ub.set_title("Ubicación de la provincia", fontsize=8.5, fontweight="bold", color=VERDE, pad=3)
    ax_ub.add_patch(Rectangle((xmin, ymin), xmax - xmin, ymax - ymin, fill=False, edgecolor="#C0392B",
                              linewidth=1.3, zorder=5))
    y = cuadro(panel, tabla_ecosistemas(x, prov, 34), 0.715, f"ECOSISTEMAS EN LOS BLOQUES — {NOMBRE_PROV[prov].upper()}",
               ["Ecosistema", "Bloques", "Área (ha)", "%"])
    presentes = set(xp.ECOSISTEMA) | set(gpd.clip(eco, aip.dissolve()).ECOSISTEMA)
    leyenda(panel, y - 0.01, [(e, l) for e, l in eco_orden if e in presentes], colores)
    pie(fig, fuente_eco)
    return fig


# ---------------------------------------------------------------- Excel

def excel(b, x, eco, ruta, fuente_eco):
    wb = openpyxl.Workbook()
    fuente = f"Ecosistemas: {fuente_eco}. Bloques V6, ANIN-DIME-SESDI 2026 (sin 74 y 75)"
    prov = lambda n: NOMBRE_PROV[n]  # noqa: E731

    # Detalle bloque × ecosistema
    wd = wb.active; wd.title = "Bloque_Ecosistema"
    cols = ["N.°", "Bloque", "Microcuenca", "Provincia", "Distrito", "Ecosistema", "Símbolo",
            "Área geométrica (ha)", "% del bloque", "Área del bloque, catálogo (ha)", "Área prorrateada (ha)"]
    f0 = encabezado_institucional(wd, "ECOSISTEMAS POR BLOQUE DE INTERVENCIÓN V6 — DETALLE", len(cols), fuente)
    d = x.sort_values(["NOMBPROV", "NOMBDIST", "BLOQUE", "HA_GEOM"], ascending=[True, True, True, False])
    n = len(d)
    rng = lambda c: f"${c}${f0 + 1}:${c}${f0 + n}"  # noqa: E731
    filas = []
    for i, r in enumerate(d.itertuples(), start=1):
        f = f0 + i
        filas.append([i, r.BLOQUE, r.MICROC, prov(r.NOMBPROV), titulo_cdl(r.NOMBDIST), r.ECOSISTEMA, r.SIMBOLO,
                      round(r.HA_GEOM, 3), f"=H{f}/SUMIF({rng('B')},B{f},{rng('H')})", r.AREA_BLOQUE_HA,
                      f"=I{f}*J{f}"])
    escribir_tabla(wd, f0, cols, filas,
                   anchos={"N.°": 6, "Bloque": 10, "Microcuenca": 14, "Provincia": 13, "Distrito": 24,
                           "Ecosistema": 48, "Símbolo": 11, "Área geométrica (ha)": 14, "% del bloque": 11,
                           "Área del bloque, catálogo (ha)": 16, "Área prorrateada (ha)": 15},
                   totales={"Área geométrica (ha)": "SUM", "Área prorrateada (ha)": "SUM"},
                   formato_num={"Área geométrica (ha)": HA, "% del bloque": PCT,
                                "Área del bloque, catálogo (ha)": HA, "Área prorrateada (ha)": HA})
    det = {c: f"Bloque_Ecosistema!{rng(c)}" for c in "BDFK"}

    # Resumen por ecosistema
    wr = wb.create_sheet("Resumen_Ecosistema", 0)
    ecos = x.groupby("ECOSISTEMA").HA_OFICIAL.sum().sort_values(ascending=False).index
    simb = x.drop_duplicates("ECOSISTEMA").set_index("ECOSISTEMA").SIMBOLO
    eco_amb = eco.groupby("ECOSISTEMA").AREA_HA.sum()
    cols = ["N.°", "Ecosistema", "Símbolo", "N.° de bloques", "Área en bloques (ha)", "% de los bloques",
            "Área en el ámbito (ha)", "% del ecosistema intervenido"]
    f1 = encabezado_institucional(wr, "SUPERFICIE DE ECOSISTEMAS EN LOS 117 BLOQUES V6", len(cols), fuente)
    tot_row = f1 + len(ecos) + 1
    filas = []
    for i, e in enumerate(ecos, start=1):
        f = f1 + i
        filas.append([i, e, simb.get(e, ""),
                      f"=COUNTIF({det['F']},B{f})",  # un registro por bloque y ecosistema
                      f"=SUMIF({det['F']},B{f},{det['K']})",
                      f"=E{f}/E${tot_row}", round(float(eco_amb.get(e, 0)), 3) if e != SIN_DATO else None,
                      f'=IF(N(G{f})=0,"",E{f}/G{f})'])
    escribir_tabla(wr, f1, cols, filas,
                   anchos={"N.°": 6, "Ecosistema": 52, "Símbolo": 11, "N.° de bloques": 11,
                           "Área en bloques (ha)": 15, "% de los bloques": 12, "Área en el ámbito (ha)": 16,
                           "% del ecosistema intervenido": 15},
                   totales={"Área en bloques (ha)": "SUM", "% de los bloques": "SUM", "Área en el ámbito (ha)": "SUM"},
                   formato_num={"Área en bloques (ha)": HA, "% de los bloques": PCT, "Área en el ámbito (ha)": HA,
                                "% del ecosistema intervenido": "0.00%"})
    c = wr.cell(tot_row + 2, 2, f'=IF(ABS(E{tot_row}-{TOTAL_V6_HA})<0.01,"OK: suma {TOTAL_V6_HA:,.3f} ha",'
                                f'"REVISAR: no suma {TOTAL_V6_HA:,.3f} ha")')
    c.font = Font(name="Arial", size=10, bold=True, color=VERDE.lstrip("#"))

    # Ecosistema por provincia (ha prorrateadas)
    wp = wb.create_sheet("Ecosistema_Provincia", 1)
    provs = ["AYABACA", "HUANCABAMBA", "MORROPON"]
    cols = ["Ecosistema"] + [NOMBRE_PROV[p] + " (ha)" for p in provs] + ["Total (ha)"]
    f2 = encabezado_institucional(wp, "ECOSISTEMAS EN LOS BLOQUES V6 POR PROVINCIA (ha)", len(cols), fuente)
    filas = []
    for i, e in enumerate(ecos, start=1):
        f = f2 + i
        filas.append([e] + [f'=SUMIFS({det["K"]},{det["F"]},$A{f},{det["D"]},"{NOMBRE_PROV[p]}")' for p in provs] + [f"=SUM(B{f}:D{f})"])
    escribir_tabla(wp, f2, cols, filas, anchos={"Ecosistema": 52},
                   totales={c: "SUM" for c in cols[1:]}, formato_num={c: HA for c in cols[1:]})

    # Un registro por bloque: ecosistema dominante
    wbq = wb.create_sheet("Bloques", 2)
    cols = ["N.°", "Bloque", "Microcuenca", "Provincia", "Distrito", "Área catálogo (ha)", "N.° de ecosistemas",
            "Ecosistema dominante", "% del dominante", "Otros ecosistemas"]
    f3 = encabezado_institucional(wbq, "ECOSISTEMA DOMINANTE POR BLOQUE V6", len(cols), fuente)
    filas = []
    for i, (cod, g) in enumerate(d.groupby("BLOQUE", sort=False), start=1):
        g = g.sort_values("HA_GEOM", ascending=False)
        r0 = g.iloc[0]
        f = f3 + i
        filas.append([i, cod, r0.MICROC, prov(r0.NOMBPROV), titulo_cdl(r0.NOMBDIST), r0.AREA_BLOQUE_HA,
                      f"=COUNTIF({det['B']},B{f})", r0.ECOSISTEMA, round(float(r0.PCT), 4),
                      "; ".join(f"{e} ({p:.0%})" for e, p in zip(g.ECOSISTEMA[1:], g.PCT[1:]))])
    escribir_tabla(wbq, f3, cols, filas,
                   anchos={"N.°": 6, "Bloque": 10, "Microcuenca": 14, "Provincia": 13, "Distrito": 24,
                           "Área catálogo (ha)": 14, "N.° de ecosistemas": 12, "Ecosistema dominante": 46,
                           "% del dominante": 11, "Otros ecosistemas": 70},
                   totales={"Área catálogo (ha)": "SUM"}, formato_num={"Área catálogo (ha)": HA, "% del dominante": PCT})

    # Notas
    wn = wb.create_sheet("Notas")
    notas = [
        "NOTAS Y CONTROLES",
        f"Capa de ecosistemas: {fuente_eco}.",
        "Los ecosistemas se recortaron al ámbito de 15 distritos (INEI, capa simplificada) y se reproyectaron a "
        "UTM WGS84 Zona 17S (EPSG:32717). Áreas calculadas en UTM.",
        "Bloque_Ecosistema: intersección de cada bloque V6 con los ecosistemas. «% del bloque» = área geométrica "
        "del tramo / área geométrica del bloque (fórmula SUMIF).",
        f"«Área prorrateada» = % del bloque × área del bloque en el catálogo V6, para que el total coincida con "
        f"{TOTAL_V6_HA:,.3f} ha. Los resúmenes usan esta área.",
        f"«{SIN_DATO}»: parte del bloque que la capa de ecosistemas no cubre (si aparece, revise la capa).",
        "% del ecosistema intervenido = área del ecosistema en bloques / área del ecosistema en el ámbito "
        "(15 distritos).",
        "Se excluyen los bloques retirados 74 y 75 (ver README_LIBERACION_AREAS.md).",
        "Escala: el Mapa Nacional de Ecosistemas es de escala regional; el ecosistema de cada bloque debe "
        "confirmarse en campo (F-DT-03, «Tipo de ecosistema (UP)»).",
    ]
    for i, t in enumerate(notas, start=1):
        c = wn.cell(i, 1, t)
        c.font = Font(name="Arial", size=11 if i == 1 else 10, bold=(i == 1), color=VERDE.lstrip("#") if i == 1 else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wn.column_dimensions["A"].width = 120
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.save(ruta)


# ---------------------------------------------------------------- principal

def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--ecosistemas", help="capa de ecosistemas (.gpkg, .shp, .geojson o .zip)")
    ap.add_argument("--campo", help="campo con el nombre del ecosistema")
    ap.add_argument("--campo-simbolo", help="campo con el símbolo del ecosistema")
    ap.add_argument("--fuente", default="Mapa Nacional de Ecosistemas del Perú, MINAM 2018 "
                                         "(R.M. N.° 440-2018-MINAM)", help="texto de la fuente para el pie")
    ap.add_argument("--salida", default=str(SALIDA), help="carpeta de salida (por defecto mapas/salidas)")
    a = ap.parse_args()
    salida = Path(a.salida)

    b, _, amb, ctx = cargar()
    eco = leer_ecosistemas(buscar_capa(a.ecosistemas), a.campo, a.campo_simbolo, amb, b)
    x = cruzar(b, eco)
    assert abs(x.HA_OFICIAL.sum() - TOTAL_V6_HA) < 0.01

    orden = (x.groupby(["ECOSISTEMA", "SIMBOLO"]).HA_OFICIAL.sum().reset_index()
             .merge(eco.groupby("ECOSISTEMA").AREA_HA.sum().rename("AMB").reset_index(), how="outer")
             .fillna({"HA_OFICIAL": 0, "SIMBOLO": ""}).sort_values(["HA_OFICIAL", "AMB"], ascending=False))
    colores = paleta(list(orden.ECOSISTEMA))
    eco_orden = [(r.ECOSISTEMA, textwrap.fill(etiqueta(r), 52, subsequent_indent="  ")) for r in orden.itertuples()]
    if SIN_DATO in x.ECOSISTEMA.values:
        sd = x[x.ECOSISTEMA == SIN_DATO].sort_values("PCT", ascending=False)
        print(f"AVISO: {sd.HA_OFICIAL.sum():,.3f} ha de {len(sd)} bloques sin ecosistema en la capa:")
        for r in sd.itertuples():
            print(f"   {r.BLOQUE:>8}  {r.NOMBDIST:<24} {r.HA_OFICIAL:9.3f} ha  ({r.PCT:.1%} del bloque)")

    salida.mkdir(parents=True, exist_ok=True)
    figs = [("01_ecosistemas_general", lamina_general(b, x, eco, amb, ctx, colores, eco_orden, a.fuente))]
    for i, p in enumerate(["AYABACA", "HUANCABAMBA", "MORROPON"], start=2):
        figs.append((f"0{i}_ecosistemas_{p.lower()}",
                     lamina_provincia(b, x, eco, amb, ctx, colores, eco_orden, a.fuente, p, i)))
    with PdfPages(salida / "IN_Piura_Mapa_Ecosistemas_Bloques_V6.pdf") as pp:
        for nombre, fig in figs:
            fig.savefig(salida / f"IN_Piura_{nombre}.png", dpi=200)
            pp.savefig(fig, dpi=200)  # ecosistemas en raster: PDF liviano
            plt.close(fig)

    excel(b, x, eco, salida / "IN_Piura_Ecosistemas_por_Bloque_V6.xlsx", a.fuente)
    gpkg = salida / "IN_Piura_ecosistemas_bloques_v6.gpkg"
    eco.to_file(gpkg, layer="ecosistemas_ambito", driver="GPKG")
    x.drop(columns=["AREA_GEOM_HA"]).to_file(gpkg, layer="ecosistemas_bloques_v6", driver="GPKG")

    print(f"{len(figs)} láminas, Excel y GeoPackage → {salida}")
    print(x.groupby("ECOSISTEMA").HA_OFICIAL.sum().sort_values(ascending=False).round(3).to_string())


if __name__ == "__main__":
    main()
