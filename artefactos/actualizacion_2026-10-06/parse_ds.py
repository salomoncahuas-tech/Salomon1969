"""Extrae las tablas de un libro Graficos_DS_Consolidado (aplicativo IN Piura) a JSON."""
import openpyxl, json, sys, re, datetime
f, out = sys.argv[1], sys.argv[2]
wv = openpyxl.load_workbook(f, data_only=True)
wf = openpyxl.load_workbook(f, data_only=False)
def clean(v):
    if isinstance(v, (datetime.date, datetime.datetime)): return v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v.is_integer(): return int(v)
    if isinstance(v, str): return v.strip()
    return v
TIT = re.compile(r'^[A-Z]\. \S')
def is_title(s):
    return bool(TIT.match(s)) or (s.upper() == s and len(s) >= 10 and '/' not in s and any(ch.isalpha() for ch in s))
res = {'sheets': []}
for ws in wv.worksheets:
    wsf = wf[ws.title]
    rows = []
    for r in range(1, ws.max_row + 1):
        vals = []
        for c in range(1, ws.max_column + 1):
            fv = wsf.cell(r, c).value
            vals.append(None if (isinstance(fv, str) and fv.startswith('=')) else clean(ws.cell(r, c).value))
        while vals and vals[-1] in (None, ''): vals.pop()
        rows.append(vals)
    sh = {'name': ws.title, 'pre': [], 'tables': []}
    cur = None; state = None
    for vals in rows[4:]:
        if not vals: continue
        first = vals[0] if vals[0] is not None else ''
        single = len(vals) == 1
        s = str(first)
        if single and is_title(s) and s != 'TOTAL' and not s.startswith('PROMEDIO'):
            if cur: sh['tables'].append(cur)
            cur = {'title': s, 'desc': '', 'header': None, 'rows': [], 'notes': [], 'total': False}; state = 'pre'
            continue
        if cur is None:
            sh['pre'].append(vals); continue
        if state == 'pre':
            if single: cur['desc'] = (cur['desc'] + ' ' + s).strip()
            else: cur['header'] = vals; state = 'data'
            continue
        # data
        if single and s == 'TOTAL': cur['total'] = True; continue
        if single and (s.startswith('Fuente') or s.startswith('Nota') or s.startswith('Aviso') or s.startswith('Libro') or s.startswith('Cada')):
            cur['notes'].append(s); continue
        cur['rows'].append(vals)
    if cur: sh['tables'].append(cur)
    res['sheets'].append(sh)
json.dump(res, open(out, 'w'), ensure_ascii=False, indent=1)
for sh in res['sheets']:
    print('##', sh['name'], '| pre', len(sh['pre']))
    for t in sh['tables']:
        print('   -', t['title'][:80], '|', len(t['header'] or []), 'cols', len(t['rows']), 'rows', 'TOTAL' if t['total'] else '', '| notes', len(t['notes']))
