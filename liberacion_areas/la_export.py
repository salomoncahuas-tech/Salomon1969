# -*- coding: utf-8 -*-
"""Exportación de la matriz predial y el avance a Excel con formato ANIN."""
from __future__ import annotations

import io
from datetime import date, datetime, timedelta, timezone

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL, ALT = "1B4F72", "EAF2F8"
_thin = Side(style="thin", color="A6ACAF")
_B = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
ENC = ["AUTORIDAD NACIONAL DE INFRAESTRUCTURA – ANIN",
       "DIRECCIÓN DE INTERVENCIONES MULTISECTORIALES Y DE EMERGENCIA – DIME",
       "SUBDIRECCIÓN DE ESTUDIOS DE INVERSIÓN"]
SEMAFORO = {"VERDE": "C6EFCE", "AMBAR": "FFEB9C", "ROJO": "FFC7CE"}
_LIMA = timezone(timedelta(hours=-5), "PET")      # Perú no aplica horario de verano


def _celda(v):
    """Excel no admite zonas horarias: las marcas de tiempo de PostgreSQL (timestamptz) se pasan a hora de Lima."""
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return None
    if isinstance(v, pd.Timestamp):
        v = v.to_pydatetime()
    if isinstance(v, datetime) and v.tzinfo is not None:
        return v.astimezone(_LIMA).replace(tzinfo=None)
    if isinstance(v, (dict, list)):
        return str(v)
    return v


def _hoja(wb, nombre: str, titulo: str, datos: pd.DataFrame, sumar: list[str] | None = None):
    ws = wb.create_sheet(nombre[:31])
    for i, t in enumerate(ENC + [titulo, f"Proyecto IN Piura – CUI 2669244 · Generado: {date.today():%d/%m/%Y}"], 1):
        ws.cell(i, 1, t).font = Font(name="Arial", bold=i in (1, 4), size=12 if i == 4 else 10, color=AZUL)
    fila0 = 7
    cols = list(datos.columns)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(fila0, j, c)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        cell.fill = PatternFill("solid", fgColor=AZUL)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = _B
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(40, len(str(c)) + 4))
    for i, fila in enumerate(datos.itertuples(index=False), fila0 + 1):
        for j, v in enumerate(fila, 1):
            cell = ws.cell(i, j, _celda(v))
            cell.font = Font(name="Arial", size=9)
            cell.border = _B
            if i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=ALT)
            if isinstance(v, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(v, date):
                cell.number_format = "dd/mm/yyyy"
            if isinstance(v, str) and v in SEMAFORO:
                cell.fill = PatternFill("solid", fgColor=SEMAFORO[v])
    ult = fila0 + len(datos)
    if sumar and len(datos):
        t = ult + 1
        ws.cell(t, 1, "TOTAL").font = Font(name="Arial", bold=True, size=9)
        for c in sumar:
            if c in cols:
                j = cols.index(c) + 1
                L = get_column_letter(j)
                cell = ws.cell(t, j, f"=SUM({L}{fila0 + 1}:{L}{ult})")
                cell.font = Font(name="Arial", bold=True, size=9)
                cell.border = _B
    ws.freeze_panes = ws.cell(fila0 + 1, 2)
    ws.auto_filter.ref = f"A{fila0}:{get_column_letter(max(1, len(cols)))}{max(ult, fila0)}"
    return ws


def exportar(hojas: dict[str, tuple[str, pd.DataFrame, list[str] | None]]) -> bytes:
    """hojas = {nombre: (titulo, dataframe, columnas_a_sumar)} → bytes del .xlsx"""
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, (titulo, datos, sumar) in hojas.items():
        _hoja(wb, nombre, titulo, datos, sumar)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
