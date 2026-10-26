# -*- coding: utf-8 -*-
"""
Ficha Técnica Predial imprimible – Proyecto IN Piura (CUI 2669244) · Liberación de Áreas.

Adapta la «Ficha técnica predial» de campo al proyecto y a los formularios KoboToolbox F-LA-02 (ficha predial y
socio-territorial) y F-LA-03 (inspección in situ y croquis):

- Las listas desplegables usan las MISMAS etiquetas de `la_campos.OPCIONES` (las que muestra Kobo), de modo que lo
  anotado en papel / Excel se transcribe o importa sin traducciones.
- El código de la unidad se elige de una lista que depende del asistente predial (catálogo V6 con asignación,
  `datos/unidades_liberacion_areas.csv`); ESP ve todas las unidades, incluidas las «por asignar».
- Provincia → distrito del predio con lista dependiente (ámbito: 3 provincias y 15 distritos), y control contra el
  distrito que el catálogo V6 registra para la unidad.
- Coordenadas UTM WGS84 Zona 17S validadas (Este 450,000–750,000 m; Norte 9,300,000–9,600,000 m).
- Una columna fuera del área de impresión indica el campo Kobo (XLSForm) de cada dato.

Uso:
    python -m liberacion_areas.la_ficha_predial [destino.xlsx]
"""
from __future__ import annotations

import csv
import math
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

from . import la_campos as lc

RAIZ = Path(__file__).resolve().parents[1]
CATALOGO_CSV = RAIZ / "datos" / "unidades_liberacion_areas.csv"
DESTINO = RAIZ / "Ficha_Predial_IN_Piura_F-LA.xlsx"
VERSION = "FP-2026.10"

NOMBRE_PROYECTO = ("Recuperación del servicio de regulación de riesgos naturales y recuperación de ecosistemas "
                   "degradados en la Cuenca Alta del Río Piura – Proyecto IN Piura (CUI 2669244)")
ENC = ["AUTORIDAD NACIONAL DE INFRAESTRUCTURA - ANIN",
       "DIRECCIÓN DE INTERVENCIONES MULTISECTORIALES Y DE EMERGENCIA - DIME",
       "SUBDIRECCIÓN DE ESTUDIOS DE INVERSIÓN"]

PROVINCIAS = {"ayabaca": "Ayabaca", "huancabamba": "Huancabamba", "morropon": "Morropón"}
DISTRITOS = {
    "frias": "Frías", "canchaque": "Canchaque", "huarmaca": "Huarmaca", "huancabamba": "Huancabamba",
    "lalaquiz": "Lalaquiz", "san_miguel_de_el_faique": "San Miguel de El Faique", "buenos_aires": "Buenos Aires",
    "chalaco": "Chalaco", "chulucanas": "Chulucanas", "morropon": "Morropón", "salitral": "Salitral",
    "san_juan_de_bigote": "San Juan de Bigote", "santa_catalina_de_mossa": "Santa Catalina de Mossa",
    "santo_domingo": "Santo Domingo", "yamango": "Yamango",
}
TIPO_UNIDAD = {"bloque": "Bloque", "lote_sus": "Lote SUS", "vivero": "Vivero"}
ESTE_MIN, ESTE_MAX, NORTE_MIN, NORTE_MAX = 450000, 750000, 9300000, 9600000

# Colores ANIN
VERDE, VERDE_MED, VERDE_CLARO, ETIQ, LISTA, CALC, GRIS = (
    "1B4D2E", "C6D9CC", "E8F0EA", "F2F5F3", "E2F0D9", "EDEDED", "7F7F7F")
_fino = Side(style="thin", color="8C9A90")
BORDE = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)
ARIAL = "Arial"

# Columnas de la ficha: B..Y (24 columnas); Z separador; AA = campo Kobo (no se imprime)
C0, C1 = 2, 25
COL_KOBO = 27
MITAD_ETQ, MITAD_VAL = (2, 6, 7, 13), (14, 18, 19, 25)    # (etq_ini, etq_fin, val_ini, val_fin)
TERCIOS = ((7, 12), (13, 18), (19, 25))                    # tres listas en una fila (selección múltiple de Kobo)

# Listas propias de la ficha (no existen como select en Kobo; se dejan en texto o se derivan)
LISTAS_FICHA: dict[str, list[str]] = {
    "personeria": ["Natural", "Jurídica"],
    "sit_registral": ["Inscrito en SUNARP", "No inscrito", "En trámite de inscripción", "Por verificar"],
    "departamento": ["Piura"],
    "dep_domicilio": ["Piura", "Tumbes", "Lambayeque", "Cajamarca", "Lima", "Otro"],
    "prov_domicilio": ["Ayabaca", "Huancabamba", "Morropón", "Piura", "Sullana", "Paita", "Talara", "Sechura", "Otra"],
    "cargo_rep": ["Presidente comunal", "Representante legal", "Apoderado", "Funcionario público", "Otro"],
    "tipo_predio": ["Rústico", "Eriazo", "Urbano", "Por verificar"],
    "edificaciones": ["Ninguna", "Vivienda", "Corral / cobertizo", "Cultivos anuales", "Cultivos permanentes / frutales",
                      "Cercos", "Pozo / reservorio", "Otro"],
    "pendiente": ["Plana (0–4 %)", "Moderada (4–15 %)", "Fuerte (15–25 %)", "Empinada (25–50 %)",
                  "Muy empinada (50–75 %)", "Extremadamente empinada (> 75 %)"],
    "erosion": ["Sin erosión visible", "Laminar", "En surcos", "En cárcavas", "Movimientos en masa"],
    "cobertura": ["Bosque seco conservado", "Bosque seco intervenido / ralo", "Matorral", "Pastizal",
                  "Suelo desnudo / degradado"],
    "tipo_colindante": ["Predio privado", "Comunidad campesina", "Terreno del Estado", "Quebrada / río",
                        "Camino / trocha", "Divisoria / cerro", "Otro"],
    "und_medida": ["m", "m²", "m³", "und", "ha"],
    "material": ["Madera", "Adobe", "Tapial", "Piedra", "Concreto", "Acero / metal", "Alambre de púas", "Malla",
                 "Calamina", "Otro"],
    "conservacion": ["Bueno", "Regular", "Malo"],
    "construccion": ["Terminado", "En construcción", "Abandonado"],
    "intervencion": ["Revegetación con especies nativas", "Reforestación con especies nativas",
                     "Zanjas de infiltración", "Terrazas de formación lenta", "Diques para control de cárcavas",
                     "Clausura temporal y manejo de pastizales", "Barreras vivas / cercos vivos",
                     "Vivero temporal o de aclimatación", "Obras complementarias (estabilización / drenaje)",
                     "Por definir"],
    # Condiciones del prediante que exigen una personería determinada (etiquetas Kobo de tipo_titularidad)
    "cond_natural": [lc.etiqueta_opcion("tipo_titularidad", c) for c in ("posesionario", "ocupante", "sucesion")],
    "cond_juridica": [lc.etiqueta_opcion("tipo_titularidad", c) for c in ("comunal", "estatal")],
}
LISTAS_KOBO = ["si_no", "aceptacion", "tipo_titularidad", "sucesion_estado", "estado_civil", "docs_exhibidos",
               "uso_actual", "ocupacion", "estado_la", "interferencias", "conclusion_campo", "tipo_acta", "pt_tipo",
               "modalidad", "acceso_tipo"]


# ------------------------------------------------------------------ catálogo
def _orden_natural(codigo: str):
    tipo = 2 if codigo.upper().startswith("SUS") else (0 if codigo.isdigit() else 1)
    return (tipo, [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", codigo)])


def cargar_catalogo(ruta: Path = CATALOGO_CSV) -> list[dict]:
    with open(ruta, encoding="utf-8-sig", newline="") as f:
        filas = list(csv.DictReader(f))
    salida = [{
        "codigo": r["name"].strip(),
        "tipo": TIPO_UNIDAD.get(r["tipo_unidad"], r["tipo_unidad"]),
        "provincia": PROVINCIAS[r["provincia"]],
        "distrito": DISTRITOS[r["distrito"]],
        "area_ha": float(r["area_ha"]),
        "bloque_ref": r.get("bloque_ref") or "",
        "posicion_sus": r.get("posicion_sus") or "",
        "asistente": r["asistente"].strip(),
        "asistente_nombre": r["asistente_nombre"].strip(),
    } for r in filas]
    salida.sort(key=lambda u: _orden_natural(u["codigo"]))
    return salida


def catalogo_desde_unidades(unidades: dict[str, dict] | list[dict]) -> list[dict]:
    """Catálogo del aplicativo (la_unidades / unidades.csv: provincia y distrito en slug) → formato de la ficha."""
    filas = unidades.values() if isinstance(unidades, dict) else unidades
    salida = []
    for r in filas:
        cod = str(r.get("codigo") or r.get("name") or "").strip()
        if not cod or r.get("tipo_unidad") == "vivero":
            continue
        area = r.get("area_ha")
        salida.append({
            "codigo": cod,
            "tipo": TIPO_UNIDAD.get(r.get("tipo_unidad"), r.get("tipo_unidad")),
            "provincia": PROVINCIAS.get(r.get("provincia"), str(r.get("provincia") or "").title()),
            "distrito": DISTRITOS.get(r.get("distrito"), str(r.get("distrito") or "").replace("_", " ").title()),
            "area_ha": float(area) if area not in (None, "") else None,
            "bloque_ref": str(r.get("bloque_ref") or ""),
            "posicion_sus": str(r.get("posicion_sus") or ""),
            "asistente": str(r.get("asistente") or "").strip(),
            "asistente_nombre": str(r.get("asistente_nombre") or "").strip(),
        })
    salida.sort(key=lambda u: _orden_natural(u["codigo"]))
    return salida


def _slug_lado(lado: str) -> str:
    return lado.lower()


def _nombre_rango(prefijo: str, clave: str) -> str:
    return prefijo + re.sub(r"[^A-Za-z0-9]", "_", clave.replace("ó", "o"))


# ------------------------------------------------------------------ libro
class _Ficha:
    """Arma la hoja de la ficha fila por fila y registra la correspondencia con Kobo."""

    def __init__(self, ws):
        self.ws = ws
        self.r = 1
        self.mapa: list[tuple[str, str, str, str, str]] = []   # sección, dato, celda, campo Kobo, tipo de entrada
        self.seccion = ""
        self._dv: dict[tuple, DataValidation] = {}
        self.celdas: dict[str, str] = {}                      # clave interna → celda

    # ---- estilo
    def _rango(self, r1, c1, r2, c2, valor=None, *, fill=None, font=None, align=None, borde=True, fmt=None):
        ws = self.ws
        if (r1, c1) != (r2, c2):
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        if valor is not None:
            ws.cell(r1, c1, valor)
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                c = ws.cell(rr, cc)
                if borde:
                    c.border = BORDE
                if fill:
                    c.fill = PatternFill("solid", fgColor=fill)
        a = ws.cell(r1, c1)
        a.font = font or Font(name=ARIAL, size=9)
        a.alignment = align or Alignment(vertical="center", wrap_text=True)
        if fmt:
            a.number_format = fmt
        return a

    def _validacion(self, celda: str, *, tipo="list", f1=None, f2=None, operador=None, estilo="stop",
                    titulo="Valor no válido", error="Elija un valor de la lista desplegable.", ayuda=None):
        clave = (tipo, f1, f2, operador, estilo, error, ayuda)
        dv = self._dv.get(clave)
        if dv is None:
            dv = DataValidation(type=tipo, formula1=f1, formula2=f2, operator=operador, allow_blank=True,
                                showErrorMessage=True, errorStyle=estilo, errorTitle=titulo, error=error)
            if ayuda:
                dv.showInputMessage, dv.promptTitle, dv.prompt = True, "Ayuda", ayuda[:250]
            self._dv[clave] = dv
            self.ws.add_data_validation(dv)
        dv.add(celda)

    def _entrada(self, r, c1, c2, kind, *, lista=None, f1=None, fmt=None, h=None, estilo="stop", error=None,
                 ayuda=None, formula=None, clave=None):
        """Celda de entrada (o fórmula) con su validación; devuelve su referencia.
        `clave` = nombre del campo Kobo (XLSForm FP) que se lee / escribe en esa celda (hoja _MAPA)."""
        ref = f"{get_column_letter(c1)}{r}"
        if clave:
            self.celdas[clave] = ref
        fill = {"lista": LISTA, "calc": CALC}.get(kind)
        font = Font(name=ARIAL, size=9, italic=kind == "calc", color="1F3B2A" if kind == "calc" else "000000")
        self._rango(r, c1, r, c2, formula, fill=fill, font=font, fmt=fmt,
                    align=Alignment(vertical="center", horizontal=h or "left", wrap_text=True))
        if kind == "lista":
            self._validacion(ref, f1=f1 or lista, estilo=estilo,
                             error=error or "Elija un valor de la lista desplegable.", ayuda=ayuda,
                             titulo="Valor no válido" if estilo == "stop" else "Valor fuera de la lista")
        elif kind in ("entero", "decimal", "fecha", "dni", "ruc", "celular", "texto"):
            if fmt is None and kind in ("dni", "ruc", "celular", "texto"):
                self.ws[ref].number_format = "@"
            if kind == "entero":
                mn, mx = f1
                self._validacion(ref, tipo="whole", f1=str(mn), f2=str(mx), operador="between",
                                 titulo="Fuera de rango", error=error or f"Número entero entre {mn} y {mx}.",
                                 ayuda=ayuda)
            elif kind == "decimal":
                mn, mx = f1
                self._validacion(ref, tipo="decimal", f1=str(mn), f2=str(mx), operador="between",
                                 titulo="Fuera de rango", error=error or f"Valor entre {mn:,} y {mx:,}.", ayuda=ayuda)
            elif kind == "fecha":
                self.ws[ref].number_format = "dd/mm/yyyy"
                self._validacion(ref, tipo="date", f1="DATE(2024,1,1)", f2="DATE(2035,12,31)", operador="between",
                                 titulo="Fecha", error="Ingrese una fecha válida (dd/mm/aaaa).", ayuda=ayuda)
            elif kind == "dni":
                self._validacion(ref, tipo="custom", f1=f"AND(LEN({ref})=8,ISNUMBER(--{ref}))", titulo="DNI",
                                 error="El DNI tiene 8 dígitos numéricos.", ayuda=ayuda)
            elif kind == "ruc":
                self._validacion(ref, tipo="custom",
                                 f1=f"AND(LEN({ref})=11,ISNUMBER(--{ref}),OR(LEFT({ref},2)=\"10\",LEFT({ref},2)=\"20\"))",
                                 titulo="RUC", error="El RUC tiene 11 dígitos y empieza con 10 o 20.", ayuda=ayuda)
            elif kind == "celular":
                self._validacion(ref, tipo="custom", f1=f"AND(LEN({ref})=9,LEFT({ref},1)=\"9\",ISNUMBER(--{ref}))",
                                 titulo="Celular", error="El celular tiene 9 dígitos y empieza con 9.", ayuda=ayuda)
        return ref

    def _kobo(self, r, campo, dato, ref, tipo):
        if campo:
            c = self.ws.cell(r, COL_KOBO, campo)
            c.font = Font(name=ARIAL, size=7, italic=True, color=GRIS)
            c.alignment = Alignment(vertical="center")
        self.mapa.append((self.seccion, dato, ref, campo or "—", tipo))

    # ---- bloques de la ficha
    def alto(self, h):
        self.ws.row_dimensions[self.r].height = h

    def espacio(self, h=4):
        self.alto(h)
        self.r += 1

    def seccion_titulo(self, texto):
        self.seccion = texto
        self._rango(self.r, C0, self.r, C1, texto, fill=VERDE,
                    font=Font(name=ARIAL, size=10, bold=True, color="FFFFFF"))
        self.alto(18)
        self.r += 1

    def subtitulo(self, texto):
        self._rango(self.r, C0, self.r, C1, texto, fill=VERDE_MED, font=Font(name=ARIAL, size=9, bold=True, color=VERDE))
        self.r += 1

    def etiqueta(self, r, c1, c2, texto, centrado=False):
        self._rango(r, c1, r, c2, texto, fill=ETIQ, font=Font(name=ARIAL, size=8, bold=True),
                    align=Alignment(vertical="center", horizontal="center" if centrado else "left", wrap_text=True,
                                    indent=0 if centrado else 1))
        # altura explícita según el largo de la etiqueta: Excel no ajusta solo las celdas combinadas (la etiqueta se
        # cortaría) y LibreOffice sí lo hace, lo que desplaza las fotos insertadas
        lineas = math.ceil(len(str(texto)) / ((c2 - c1 + 1) * 5.0 - (0 if centrado else 2)))
        actual = self.ws.row_dimensions[r].height or 15
        self.ws.row_dimensions[r].height = max(actual, 11 * lineas + 4)

    def campo(self, texto, kind="libre", *, clave=None, kobo=None, mitad=None, h=None, alto=None, **kw):
        """Fila de un dato. `mitad` = 0 (izquierda) o 1 (derecha) para dos datos por fila; None = fila completa."""
        if mitad is None:
            e0, e1, v0, v1 = 2, 6, 7, 25
        else:
            e0, e1, v0, v1 = (MITAD_ETQ, MITAD_VAL)[mitad]
        self.etiqueta(self.r, e0, e1, texto)
        ref = self._entrada(self.r, v0, v1, kind, h=h, clave=clave, **kw)
        self._kobo(self.r, kobo, texto, ref, {"libre": "Texto", "lista": "Lista desplegable", "calc": "Automático",
                                              "entero": "Número entero", "decimal": "Número", "fecha": "Fecha",
                                              "dni": "DNI (8 dígitos)", "ruc": "RUC (11 dígitos)",
                                              "celular": "Celular (9 dígitos)", "texto": "Texto"}[kind])
        if alto:
            self.alto(alto)
        if mitad is None or mitad == 1:
            self.r += 1
        return ref

    def fila_doble(self, izq: dict, der: dict):
        a = self.campo(mitad=0, **izq)
        b = self.campo(mitad=1, **der)
        return a, b


def _hoja_listas(wb, catalogo, asistentes):
    ws = wb.create_sheet("LISTAS")
    ws["A1"] = "LISTAS DESPLEGABLES DE LA FICHA PREDIAL – no cambie el orden de las columnas ni los nombres de la fila 3"
    ws["A1"].font = Font(name=ARIAL, bold=True, color=VERDE)
    ws["A2"] = ("Fila 3: nombre del rango con nombre usado por la validación. Las listas de Kobo (prefijo L_) llevan "
                "las mismas etiquetas que los formularios F-LA de KoboToolbox (liberacion_areas/la_campos.py).")
    ws["A2"].font = Font(name=ARIAL, italic=True, size=8, color=GRIS)
    rangos: dict[str, str] = {}
    columnas: list[tuple[str, list[str]]] = []
    columnas.append(("L_asistente", [a for a, _ in asistentes]))
    for k in LISTAS_KOBO:
        columnas.append(("L_" + k, [e for _, e in lc.OPCIONES[k]]))
    for k, v in LISTAS_FICHA.items():
        columnas.append(("L_" + k, v))
    columnas.append(("L_provincia", list(PROVINCIAS.values())))
    for p in PROVINCIAS.values():
        columnas.append((_nombre_rango("D_", p), sorted({u["distrito"] for u in catalogo if u["provincia"] == p})))
    columnas.append(("L_distrito_15", sorted(DISTRITOS.values())))
    for a, _ in asistentes:
        unidades = [u["codigo"] for u in catalogo if a == "ESP" or u["asistente"] == a]
        columnas.append((_nombre_rango("U_", a), unidades))
    for j, (nombre, valores) in enumerate(columnas, 1):
        L = get_column_letter(j)
        h = ws.cell(3, j, nombre)
        h.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=8)
        h.fill = PatternFill("solid", fgColor=VERDE)
        for i, v in enumerate(valores, 4):
            c = ws.cell(i, j, v)
            c.number_format = "@"
            c.font = Font(name=ARIAL, size=8)
        ref = f"LISTAS!${L}$4:${L}${3 + max(1, len(valores))}"
        wb.defined_names[nombre] = DefinedName(nombre, attr_text=ref)
        rangos[nombre] = nombre
        ws.column_dimensions[L].width = max(10, min(34, max(len(str(v)) for v in valores + [nombre]) + 2))
    ws.freeze_panes = "A4"
    ws.sheet_properties.tabColor = GRIS
    return rangos


def _hoja_catalogo(wb, catalogo):
    ws = wb.create_sheet("CATALOGO_V6")
    for i, t in enumerate(ENC, 1):
        ws.cell(i, 1, t).font = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=9)
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=VERDE)
    ws.cell(4, 1, "CATÁLOGO V6 DE UNIDADES (bloques y lotes SUS) CON ASISTENTE PREDIAL ASIGNADO").font = Font(
        name=ARIAL, bold=True, size=11, color=VERDE)
    ws.cell(5, 1, "Fuente: datos/unidades_liberacion_areas.csv (mismo catálogo de los formularios Kobo F-LA). "
                  "Proyecto IN Piura – CUI 2669244 · UTM WGS84 Zona 17S").font = Font(name=ARIAL, italic=True, size=8,
                                                                                        color=GRIS)
    cab = ["Código", "Tipo de unidad", "Provincia", "Distrito", "Área (ha)", "Bloque de referencia", "Posición SUS",
           "Asistente", "Nombre del asistente"]
    f0 = 7
    for j, t in enumerate(cab, 1):
        c = ws.cell(f0, j, t)
        c.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor=VERDE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
    for i, u in enumerate(catalogo, f0 + 1):
        vals = [u["codigo"], u["tipo"], u["provincia"], u["distrito"], u["area_ha"], u["bloque_ref"],
                u["posicion_sus"], u["asistente"], u["asistente_nombre"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v)
            c.font = Font(name=ARIAL, size=9)
            c.border = BORDE
            if j in (1, 6):
                c.number_format = "@"
            if j == 5:
                c.number_format = "#,##0.000"
            if i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=VERDE_CLARO)
    ult = f0 + len(catalogo)
    t = ult + 2
    resumen = [("Bloques (n.°)", f'=COUNTIF(B{f0 + 1}:B{ult},"Bloque")', "0"),
               ("Área de bloques (ha)", f'=SUMIF(B{f0 + 1}:B{ult},"Bloque",E{f0 + 1}:E{ult})', "#,##0.000"),
               ("Lotes SUS (n.°)", f'=COUNTIF(B{f0 + 1}:B{ult},"Lote SUS")', "0"),
               ("Área de lotes SUS (ha)", f'=SUMIF(B{f0 + 1}:B{ult},"Lote SUS",E{f0 + 1}:E{ult})', "#,##0.000")]
    for k, (txt, fml, fmt) in enumerate(resumen):
        a = ws.cell(t + k, 4, txt)
        a.font = Font(name=ARIAL, bold=True, size=9)
        a.border = BORDE
        b = ws.cell(t + k, 5, fml)
        b.font = Font(name=ARIAL, bold=True, size=9)
        b.number_format = fmt
        b.border = BORDE
    ws.cell(t + 4, 4, "Control: 117 bloques que suman 12,270.235 ha (catálogo V6).").font = Font(
        name=ARIAL, italic=True, size=8, color=GRIS)
    for L, w in zip("ABCDEFGHI", (11, 12, 13, 22, 11, 12, 12, 10, 30)):
        ws.column_dimensions[L].width = w
    ws.freeze_panes = ws.cell(f0 + 1, 2)
    ws.auto_filter.ref = f"A{f0}:I{ult}"
    for nombre, L in (("CAT_COD", "A"), ("CAT_TIPO", "B"), ("CAT_PROV", "C"), ("CAT_DIST", "D"), ("CAT_AREA", "E"),
                      ("CAT_ASIS", "H"), ("CAT_ASIS_NOM", "I")):
        wb.defined_names[nombre] = DefinedName(nombre, attr_text=f"CATALOGO_V6!${L}${f0 + 1}:${L}${ult}")
    ws.print_title_rows = f"{f0}:{f0}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = "1B4F72"


def _hoja_ficha(wb, asistentes):
    ws = wb.active
    ws.title = "FICHA PREDIAL"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 1.5
    for c in range(C0, C1 + 1):
        ws.column_dimensions[get_column_letter(c)].width = 4.6
    ws.column_dimensions["Z"].width = 2
    ws.column_dimensions["AA"].width = 30
    f = _Ficha(ws)

    # ---------- encabezado institucional
    for t in ENC:
        f._rango(f.r, C0, f.r, C1, t, fill=VERDE, borde=False,
                 font=Font(name=ARIAL, size=10 if f.r == 1 else 9, bold=True, color="FFFFFF"),
                 align=Alignment(horizontal="center", vertical="center"))
        f.r += 1
    f.espacio(5)
    r = f.r
    f._rango(r, C0, r + 1, 18, "FICHA TÉCNICA PREDIAL", font=Font(name=ARIAL, size=15, bold=True, color=VERDE),
             align=Alignment(horizontal="center", vertical="center"))
    f._rango(r, 19, r, C1, "CÓDIGO DE PREDIO", fill=VERDE,
             font=Font(name=ARIAL, size=8, bold=True, color="FFFFFF"), align=Alignment(horizontal="center"))
    codigo_sup = f"S{r + 1}"
    f.r += 2
    f._rango(f.r, C0, f.r, C1, "Liberación de áreas · Equivalente a los formularios KoboToolbox F-LA-02 (ficha predial "
                               "y socio-territorial) y F-LA-03 (inspección in situ y croquis)",
             font=Font(name=ARIAL, size=8, italic=True, color=VERDE), align=Alignment(horizontal="center"), borde=False)
    f.r += 1
    f._rango(f.r, C0, f.r, C1, f"Proyecto IN Piura – CUI 2669244 · Coordenadas UTM WGS84 Zona 17S (EPSG:32717) · "
                               f"Versión {VERSION}",
             font=Font(name=ARIAL, size=7, italic=True, color=GRIS), align=Alignment(horizontal="center"), borde=False)
    f.r += 1
    ws.cell(f.r - 4, COL_KOBO, "CAMPO KOBO (XLSForm F-LA)").font = Font(name=ARIAL, size=8, bold=True, color=VERDE)
    f.espacio()

    # ---------- 0. datos generales
    f.seccion_titulo("  DATOS GENERALES")
    f.campo("PROYECTO", "calc", formula=NOMBRE_PROYECTO, alto=26)
    f.campo("SOLICITANTE", "calc", formula="AUTORIDAD NACIONAL DE INFRAESTRUCTURA – ANIN (DIME – SESDI)")
    asis, _ = f.fila_doble(
        dict(texto="ASISTENTE PREDIAL", kind="lista", lista="L_asistente", clave="asistente", kobo="asistente",
             ayuda="Código del asistente predial (AP-1 … AP-6) o ESP. Define la lista de unidades."),
        dict(texto="NOMBRE DEL ASISTENTE", kind="calc"))
    nombre_asis = f"S{f.r - 1}"
    unidad, tipo_u = f.fila_doble(
        dict(texto="CÓDIGO DE LA UNIDAD (BLOQUE / LOTE SUS)", kind="lista", clave="unidad", kobo="unidad",
             f1=f'INDIRECT("U_"&SUBSTITUTE({asis},"-","_"))',
             error="Elija una unidad asignada al asistente seleccionado (catálogo V6).",
             ayuda="Elija primero el ASISTENTE PREDIAL: la lista muestra solo sus bloques y lotes SUS asignados."),
        dict(texto="TIPO DE UNIDAD", kind="calc", kobo="tipo_unidad (calculado)"))
    ws[unidad].number_format = "@"
    m = f"MATCH({unidad}&\"\",CAT_COD,0)"
    ubic, area_u = f.fila_doble(
        dict(texto="UBICACIÓN SEGÚN CATÁLOGO V6", kind="calc"),
        dict(texto="ÁREA DE LA UNIDAD (ha)", kind="calc", fmt="#,##0.000", h="right"))
    n_pred, cod_pred = f.fila_doble(
        dict(texto="N° DE PREDIO EN LA UNIDAD", kind="entero", f1=(0, 99), kobo="n_predio", clave="n_predio", h="center",
             ayuda="Correlativo del predio dentro de la unidad (1–99). Use 0 solo en actas comunales A-03."),
        dict(texto="CÓDIGO DE PREDIO", kind="calc", kobo="cod_predio (calculado)"))
    f.fila_doble(dict(texto="FECHA DE INSPECCIÓN", kind="fecha", kobo="fecha_insp (F-LA-03)", clave="fecha_insp", h="center"),
                 dict(texto="FECHA DE REGISTRO", kind="fecha", kobo="hoy", clave="hoy", h="center"))
    _, control = f.fila_doble(
        dict(texto="CONSENTIMIENTO USO DE DATOS PERSONALES (LEY 29733)", kind="lista", lista="L_si_no",
             kobo="consentimiento", clave="consentimiento", h="center"),
        dict(texto="CONTROL DE ASIGNACIÓN", kind="calc"))
    ws[nombre_asis] = (f'=IF({asis}="","",IF({asis}="ESP","Especialista (ESP)",'
                       f'IFERROR(INDEX(CAT_ASIS_NOM,MATCH({asis},CAT_ASIS,0)),"")))')
    ws[tipo_u] = f'=IF({unidad}="","",IFERROR(INDEX(CAT_TIPO,{m}),"⚠ No figura en el catálogo V6"))'
    ws[ubic] = f'=IF({unidad}="","",IFERROR(INDEX(CAT_DIST,{m})&" – "&INDEX(CAT_PROV,{m})&" – Piura",""))'
    ws[area_u] = f'=IF({unidad}="","",IFERROR(INDEX(CAT_AREA,{m}),""))'
    ws[cod_pred] = f'=IF(OR({unidad}="",{n_pred}=""),"",{unidad}&"-P"&TEXT({n_pred},"00"))'
    ws[codigo_sup] = f"={cod_pred}"
    f._rango(int(codigo_sup[1:]), 19, int(codigo_sup[1:]), C1, None, fill=CALC,
             font=Font(name=ARIAL, size=12, bold=True, color=VERDE), align=Alignment(horizontal="center",
                                                                                     vertical="center"))
    ws[control] = (f'=IF(OR({asis}="",{unidad}=""),"",IF(ISNA({m}),"⚠ Unidad fuera del catálogo V6",'
                   f'IF({asis}="ESP","Especialista: acceso a todas las unidades",'
                   f'IF(INDEX(CAT_ASIS,{m})={asis},"✔ Unidad asignada al asistente",'
                   f'"⚠ Asignada a "&INDEX(CAT_ASIS,{m})))))')
    f.espacio()

    # ---------- 1. predio matriz (titularidad registral)
    f.seccion_titulo("  1. CONTROL E IDENTIFICACIÓN DEL PREDIO MATRIZ (TITULARIDAD REGISTRAL)")
    f.fila_doble(dict(texto="SITUACIÓN REGISTRAL", kind="lista", lista="L_sit_registral", clave="sit_registral"),
                 dict(texto="PARTIDA ELECTRÓNICA / N° TÍTULO", kind="texto", kobo="doc_partida", clave="doc_partida"))
    f.fila_doble(dict(texto="TIPO DE PERSONERÍA (TITULAR REGISTRAL)", kind="lista", lista="L_personeria", h="center",
                      clave="reg_personeria"),
                 dict(texto="OFICINA REGISTRAL", kind="texto", clave="oficina_registral"))
    f.campo("TITULAR REGISTRAL (NOMBRE / RAZÓN SOCIAL)", "texto", clave="reg_titular")
    r = f.r
    f.etiqueta(r, 2, 6, "DOCUMENTOS EXHIBIDOS (hasta 3)")
    for k, (c1, c2) in enumerate(TERCIOS, 1):
        f._entrada(r, c1, c2, "lista", lista="L_docs_exhibidos", clave=f"docs_exhibidos_{k}")
    f._kobo(r, "docs_exhibidos (selección múltiple)", "Documentos exhibidos", f"G{r}:Y{r}", "Lista desplegable")
    f.r += 1
    f.espacio()

    # ---------- 2. prediante
    f.seccion_titulo("  2. IDENTIFICACIÓN DEL PREDIANTE Y TENENCIA")
    cond, pers = f.fila_doble(
        dict(texto="CONDICIÓN JURÍDICA DEL PREDIANTE", kind="lista", lista="L_tipo_titularidad",
             kobo="tipo_titularidad", clave="tipo_titularidad"),
        dict(texto="TIPO DE PERSONERÍA", kind="lista", lista="L_personeria", h="center", clave="personeria"))
    _, alerta = f.fila_doble(
        dict(texto="ESTADO DE LA SUCESIÓN (si aplica)", kind="lista", lista="L_sucesion_estado",
             kobo="sucesion_estado", clave="sucesion_estado"),
        dict(texto="ALERTA / COHERENCIA", kind="calc"))
    ws[alerta] = (f'=IF({cond}="","",IF({cond}="{lc.etiqueta_opcion("tipo_titularidad", "ocupante")}",'
                  f'"⚠ Ocupante: ficha OBSERVADA – identificar titular o posesionario (Esp. Legal)",'
                  f'IF({pers}="","Indique la personería",'
                  f'IF(AND(ISNUMBER(MATCH({cond},L_cond_natural,0)),{pers}<>"Natural"),"⚠ Corresponde persona natural",'
                  f'IF(AND(ISNUMBER(MATCH({cond},L_cond_juridica,0)),{pers}<>"Jurídica"),"⚠ Corresponde persona jurídica",'
                  f'"✔ Coherente")))))')
    f.subtitulo("   2.1 PERSONA NATURAL (privado, posesionario, ocupante, sucesión)")
    f.fila_doble(dict(texto="NOMBRES", kind="texto", kobo="tit_nombres", clave="tit_nombres"),
                 dict(texto="APELLIDOS", kind="texto", kobo="tit_apellidos", clave="tit_apellidos"))
    f.fila_doble(dict(texto="DNI", kind="dni", kobo="tit_dni", clave="tit_dni", h="center", ayuda="8 dígitos, sin puntos ni espacios."),
                 dict(texto="CELULAR", kind="celular", kobo="tit_celular", clave="tit_celular", h="center", ayuda="9 dígitos, empieza con 9."))
    f.fila_doble(dict(texto="ESTADO CIVIL", kind="lista", lista="L_estado_civil", kobo="tit_estado_civil",
                      clave="tit_estado_civil"),
                 dict(texto="CORREO ELECTRÓNICO (opcional)", kind="texto", clave="tit_correo"))
    f.fila_doble(dict(texto="CÓNYUGE / CONVIVIENTE", kind="texto", kobo="cony_nombre", clave="cony_nombre"),
                 dict(texto="DNI DEL CÓNYUGE", kind="dni", kobo="cony_dni", clave="cony_dni", h="center"))
    f.subtitulo("   2.2 PERSONA JURÍDICA (comunidad campesina, entidad estatal, empresa o asociación)")
    f.campo("COMUNIDAD / ENTIDAD / RAZÓN SOCIAL", "texto", kobo="cc_nombre / est_entidad", clave="cc_nombre")
    f.fila_doble(dict(texto="RUC", kind="ruc", kobo="cc_ruc", clave="cc_ruc", h="center"),
                 dict(texto="PARTIDA REGISTRAL (SUNARP)", kind="texto", kobo="cc_partida", clave="cc_partida"))
    f.fila_doble(dict(texto="REPRESENTANTE / PRESIDENTE COMUNAL", kind="texto", kobo="cc_presidente / est_contacto",
                      clave="cc_presidente"),
                 dict(texto="CARGO", kind="lista", lista="L_cargo_rep", clave="rep_cargo"))
    f.fila_doble(dict(texto="DNI DEL REPRESENTANTE", kind="dni", kobo="cc_presidente_dni", clave="cc_presidente_dni", h="center"),
                 dict(texto="CELULAR DEL REPRESENTANTE", kind="celular", kobo="cc_presidente_cel", clave="cc_presidente_cel", h="center"))
    f.subtitulo("   2.3 DOMICILIO DEL PREDIANTE")
    r = f.r
    for (c1, c2), t in zip(((2, 11), (12, 16), (17, 21), (22, 25)), ("DIRECCIÓN", "DISTRITO", "PROVINCIA",
                                                                      "DEPARTAMENTO")):
        f.etiqueta(r, c1, c2, t, centrado=True)
    f.r += 1
    r = f.r
    aviso = "El valor no está en la lista. ¿Desea mantenerlo? (válido para domicilios fuera del ámbito)"
    f._entrada(r, 2, 11, "texto", clave="dom_direccion")
    f._entrada(r, 12, 16, "lista", lista="L_distrito_15", estilo="warning", error=aviso, clave="dom_distrito")
    f._entrada(r, 17, 21, "lista", lista="L_prov_domicilio", estilo="warning", error=aviso, clave="dom_provincia")
    f._entrada(r, 22, 25, "lista", lista="L_dep_domicilio", estilo="warning", error=aviso, clave="dom_departamento")
    f._kobo(r, None, "Domicilio (dirección, distrito, provincia, departamento)", f"B{r}:Y{r}", "Texto / lista abierta")
    f.alto(24)
    f.r += 1
    f.espacio()

    # ---------- 3. descripción general del predio
    corte_1 = f.r
    f.seccion_titulo("  3. DESCRIPCIÓN GENERAL DEL PREDIO INTERVENIDO")
    f.campo("NOMBRE DEL PREDIO", "texto", kobo="predio_nombre", clave="predio_nombre")
    f.fila_doble(dict(texto="CENTRO POBLADO / SECTOR", kind="texto", clave="centro_poblado"),
                 dict(texto="COMUNIDAD CAMPESINA (si aplica)", kind="texto", clave="comunidad"))
    f.campo("REFERENCIA DE ACCESO", "texto", alto=26, clave="ref_acceso")
    f.fila_doble(dict(texto="TIPO DE ACCESO", kind="lista", lista="L_acceso_tipo", clave="acceso_tipo"),
                 dict(texto="TIEMPO DESDE EL C.P. (min)", kind="entero", f1=(0, 600), h="center",
                      clave="tiempo_acceso_min"))
    f.campo("SISTEMA DE COORDENADAS", "calc", formula="UTM WGS84 – Zona 17 Sur (EPSG:32717) · punto interior del predio")
    f.fila_doble(dict(texto="ESTE UTM (m)", kind="decimal", f1=(ESTE_MIN, ESTE_MAX), kobo="punto_interior (Este)", clave="utm_este",
                      fmt="#,##0.0", h="right", error=f"Este UTM 17S entre {ESTE_MIN:,} y {ESTE_MAX:,} m.",
                      ayuda="Coordenada Este en UTM 17S (no lat/lon)."),
                 dict(texto="NORTE UTM (m)", kind="decimal", f1=(NORTE_MIN, NORTE_MAX), kobo="punto_interior (Norte)", clave="utm_norte",
                      fmt="#,##0.0", h="right", error=f"Norte UTM 17S entre {NORTE_MIN:,} y {NORTE_MAX:,} m.",
                      ayuda="Coordenada Norte en UTM 17S (no lat/lon)."))
    f.fila_doble(dict(texto="ALTITUD (m s.n.m.)", kind="decimal", f1=(0, 4500), kobo="punto_interior (alt)", clave="utm_altitud",
                      fmt="#,##0", h="right"),
                 dict(texto="PRECISIÓN GPS (m)", kind="decimal", f1=(0, 100), kobo="punto_interior (prec)", clave="utm_precision",
                      fmt="0.0", h="right"))
    dep, prov = f.fila_doble(
        dict(texto="DEPARTAMENTO", kind="lista", lista="L_departamento", h="center", clave="predio_departamento"),
        dict(texto="PROVINCIA", kind="lista", lista="L_provincia", h="center", clave="predio_provincia",
             ayuda="Ayabaca, Huancabamba o Morropón (ámbito del proyecto)."))
    ws[dep] = "Piura"
    dist, verif = f.fila_doble(
        dict(texto="DISTRITO", kind="lista", h="center", clave="predio_distrito", f1=f'INDIRECT("D_"&SUBSTITUTE({prov},"ó","o"))',
             error="Elija primero la PROVINCIA; el distrito debe pertenecer a ella (15 distritos del ámbito).",
             ayuda="Lista dependiente de la provincia elegida."),
        dict(texto="VERIFICACIÓN CON CATÁLOGO V6", kind="calc"))
    ws[verif] = (f'=IF(OR({dist}="",{unidad}=""),"",IFERROR(IF(INDEX(CAT_DIST,{m})={dist},'
                 f'"✔ Coincide con el distrito de la unidad","⚠ La unidad figura en "&INDEX(CAT_DIST,{m})&" – verificar"),""))')
    f.espacio()

    # ---------- 4. predio matriz
    f.seccion_titulo("  4. DESCRIPCIÓN DEL PREDIO MATRIZ")
    f.fila_doble(dict(texto="USO ACTUAL DEL SUELO", kind="lista", lista="L_uso_actual", kobo="uso_actual / uso_observado",
                      clave="uso_actual"),
                 dict(texto="ESPECIFIQUE SI ES «OTRO»", kind="texto", clave="uso_otro"))
    f.fila_doble(dict(texto="OCUPACIÓN OBSERVADA", kind="lista", lista="L_ocupacion", kobo="ocupacion",
                      clave="ocupacion"),
                 dict(texto="DETALLE DE LA OCUPACIÓN", kind="texto", clave="ocupacion_det"))
    f.fila_doble(dict(texto="EDIFICACIONES / CULTIVOS", kind="lista", lista="L_edificaciones",
                      clave="edificaciones"),
                 dict(texto="DETALLE / ESPECIFIQUE «OTRO»", kind="texto", clave="edif_det"))
    f.fila_doble(dict(texto="TOPOGRAFÍA / PENDIENTE", kind="lista", lista="L_pendiente", clave="pendiente_clase",
                      ayuda="Criterio de idoneidad del proyecto: pendiente máxima 75 %."),
                 dict(texto="EROSIÓN OBSERVADA", kind="lista", lista="L_erosion", clave="erosion"))
    f.fila_doble(dict(texto="ESTADO DE LA COBERTURA VEGETAL", kind="lista", lista="L_cobertura",
                      clave="cobertura"),
                 dict(texto="EVIDENCIA DE MOVIMIENTOS EN MASA", kind="lista", lista="L_si_no", h="center",
                      clave="mov_masa"))
    f.campo("OTROS / DESCRIPCIÓN DEL PREDIO", "texto", kobo="descripcion_predio (F-LA-03)", alto=40, clave="descripcion_predio",
            ayuda="Cobertura, pendiente, accesos, linderos, signos de degradación o de movimientos en masa.")
    f.subtitulo("   4.1 DATOS TÉCNICOS")
    _, area_decl = f.fila_doble(
        dict(texto="TIPO DE PREDIO", kind="lista", lista="L_tipo_predio", clave="tipo_predio"),
        dict(texto="ÁREA MATRIZ DEL PREDIO (ha)", kind="decimal", f1=(0, 100000), kobo="predio_area_decl_ha", clave="predio_area_decl_ha",
             fmt="#,##0.0000", h="right"))
    _, m2 = f.fila_doble(dict(texto="PERÍMETRO (m)", kind="decimal", f1=(0, 1000000), fmt="#,##0.00", h="right",
                              clave="perimetro_m"),
                         dict(texto="ÁREA MATRIZ DEL PREDIO (m²)", kind="calc", fmt="#,##0.00", h="right"))
    ws[m2] = f'=IF({area_decl}="","",{area_decl}*10000)'
    f.subtitulo("   4.2 COLINDANCIAS")
    r = f.r
    for (c1, c2), t in zip(((2, 5), (6, 18), (19, 25)), ("LÍMITE", "COLINDANTE (nombre / referencia)",
                                                         "TIPO DE COLINDANTE")):
        f.etiqueta(r, c1, c2, t, centrado=True)
    f.r += 1
    for lado in ("NORTE", "ESTE", "SUR", "OESTE"):
        r = f.r
        f.etiqueta(r, 2, 5, lado)
        f._entrada(r, 6, 18, "texto", clave=f"col_{_slug_lado(lado)}")
        f._entrada(r, 19, 25, "lista", lista="L_tipo_colindante", clave=f"col_{_slug_lado(lado)}_tipo")
        f._kobo(r, None, f"Colindancia {lado.lower()}", f"F{r}:Y{r}", "Texto / lista")
        f.alto(20)
        f.r += 1
    f.campo("CONFLICTOS DE LINDEROS / OBSERVACIONES", "texto", kobo="conflicto_linderos_det", alto=30,
            clave="conflicto_linderos_det")
    f.espacio()
    corte_2 = f.r

    # ---------- 5. área a intervenir
    f.seccion_titulo("  5. DESCRIPCIÓN DEL ÁREA A INTERVENIR (DENTRO DE LA UNIDAD)")
    f.subtitulo("   5.1 DATOS TÉCNICOS")
    area_int, m2i = f.fila_doble(
        dict(texto="ÁREA DEL PREDIO DENTRO DE LA UNIDAD (ha)", kind="decimal", f1=(0, 100000),
             kobo="predio_area_unidad_ha", clave="predio_area_unidad_ha", fmt="#,##0.0000", h="right"),
        dict(texto="ÁREA DENTRO DE LA UNIDAD (m²)", kind="calc", fmt="#,##0.00", h="right"))
    ws[m2i] = f'=IF({area_int}="","",{area_int}*10000)'
    pct, chk = f.fila_doble(dict(texto="% DEL PREDIO MATRIZ", kind="calc", fmt="0.0%", h="right"),
                            dict(texto="CONTROL DE ÁREAS", kind="calc"))
    ws[pct] = f'=IF(OR({area_int}="",{area_decl}="",{area_decl}=0),"",{area_int}/{area_decl})'
    ws[chk] = (f'=IF(OR({area_int}="",{area_decl}=""),"",IF({area_int}<={area_decl},'
               f'"✔ No supera el área matriz","⚠ Supera el área matriz – verificar"))')
    f.subtitulo("   5.2 CARACTERÍSTICAS")
    r = f.r
    f.etiqueta(r, 2, 6, "INTERFERENCIAS (hasta 3)")
    for k, (c1, c2) in enumerate(TERCIOS, 1):
        f._entrada(r, c1, c2, "lista", lista="L_interferencias", clave=f"interferencias_{k}")
    f._kobo(r, "interferencias (selección múltiple)", "Interferencias", f"G{r}:Y{r}", "Lista desplegable")
    f.r += 1
    f.campo("DETALLE DE INTERFERENCIAS", "texto", kobo="interf_detalle", alto=30, clave="interf_detalle")
    r = f.r
    f.etiqueta(r, 2, 6, "INTERVENCIÓN PREVISTA (referencial)")
    f._entrada(r, 7, 15, "lista", lista="L_intervencion", clave="intervencion_prevista_1")
    f._entrada(r, 16, 25, "lista", lista="L_intervencion", clave="intervencion_prevista_2")
    f._kobo(r, None, "Intervención prevista", f"G{r}:Y{r}", "Lista desplegable")
    f.alto(24)
    f.r += 1
    f.subtitulo("   5.3 DESCRIPCIÓN DE AFECTACIONES (obras y bienes existentes en el área)")
    r = f.r
    cols = ((2, 3, "CÓD."), (4, 8, "DESCRIPCIÓN / USO"), (9, 10, "METRADO"), (11, 12, "UND."),
            (13, 15, "MATERIAL"), (16, 18, "EST. CONSERV."), (19, 21, "EST. CONSTR."), (22, 25, "UBICACIÓN"))
    for c1, c2, t in cols:
        f.etiqueta(r, c1, c2, t, centrado=True)
    f.alto(22)
    f.r += 1
    for k in range(1, 6):
        r = f.r
        f._rango(r, 2, r, 3, f"OC-{k:02d}", fill=ETIQ, font=Font(name=ARIAL, size=8, bold=True),
                 align=Alignment(horizontal="center", vertical="center"))
        f._entrada(r, 4, 8, "texto", clave=f"af_descripcion_{k}")
        f._entrada(r, 9, 10, "decimal", f1=(0, 1000000), fmt="#,##0.00", h="right", clave=f"af_metrado_{k}")
        f._entrada(r, 11, 12, "lista", lista="L_und_medida", h="center", clave=f"af_und_{k}")
        f._entrada(r, 13, 15, "lista", lista="L_material", clave=f"af_material_{k}")
        f._entrada(r, 16, 18, "lista", lista="L_conservacion", clave=f"af_conservacion_{k}")
        f._entrada(r, 19, 21, "lista", lista="L_construccion", clave=f"af_construccion_{k}")
        f._entrada(r, 22, 25, "texto", clave=f"af_ubicacion_{k}")
        f.alto(20)
        f.r += 1
    f._kobo(f.r - 5, None, "Afectaciones OC-01 … OC-05", f"B{f.r - 5}:Y{f.r - 1}", "Tabla")
    f.espacio()

    # ---------- 6. resultado
    f.seccion_titulo("  6. DISPONIBILIDAD Y CONDICIÓN DE LIBERACIÓN")
    f.fila_doble(dict(texto="ACEPTACIÓN DEL TITULAR / PREDIANTE", kind="lista", lista="L_aceptacion",
                      kobo="aceptacion", clave="aceptacion"),
                 dict(texto="MODALIDAD DE DISPONIBILIDAD", kind="lista", lista="L_modalidad", clave="modalidad"))
    f.fila_doble(dict(texto="CONCLUSIÓN DE CAMPO", kind="lista", lista="L_conclusion_campo",
                      kobo="conclusion_campo (F-LA-03)", clave="conclusion_campo"),
                 dict(texto="ESTADO LA PROPUESTO", kind="lista", lista="L_estado_la", kobo="estado_la_propuesto",
                      clave="estado_la_propuesto", h="center"))
    f.fila_doble(dict(texto="ACTA PREVISTA", kind="lista", lista="L_tipo_acta", clave="tipo_acta_prevista"),
                 dict(texto="CÓDIGO DE TITULAR PREVIO (T0000)", kind="texto", kobo="cod_titular_prev", clave="cod_titular_prev",
                      h="center"))
    f.campo("OBSERVACIONES / ALERTAS SOCIALES", "texto", alto=36, clave="alertas")
    f.espacio()

    # ---------- 7. croquis y puntos
    f.seccion_titulo("  7. CROQUIS DEL PREDIO Y PUNTOS DE CONTROL (UTM WGS84 Zona 17S)")
    r = f.r
    f.celdas["croquis"] = f"B{r}"
    f._rango(r, C0, r + 13, C1, "Croquis: indicar el norte, linderos, colindantes, accesos, vértices numerados e "
                                "interferencias.", font=Font(name=ARIAL, size=7, italic=True, color=GRIS),
             align=Alignment(vertical="top", horizontal="left", wrap_text=True))
    f.r += 14
    r = f.r
    cols = ((2, 3, "N°"), (4, 8, "TIPO DE PUNTO"), (9, 13, "ESTE (m)"), (14, 18, "NORTE (m)"),
            (19, 21, "ALTITUD"), (22, 25, "OBSERVACIÓN"))
    for c1, c2, t in cols:
        f.etiqueta(r, c1, c2, t, centrado=True)
    f.r += 1
    for k in range(1, 9):
        r = f.r
        f._rango(r, 2, r, 3, k, fill=ETIQ, font=Font(name=ARIAL, size=8, bold=True),
                 align=Alignment(horizontal="center", vertical="center"))
        f._entrada(r, 4, 8, "lista", lista="L_pt_tipo", clave=f"pt_tipo_{k}")
        f._entrada(r, 9, 13, "decimal", f1=(ESTE_MIN, ESTE_MAX), fmt="#,##0.0", h="right",
                   error=f"Este UTM 17S entre {ESTE_MIN:,} y {ESTE_MAX:,} m.", clave=f"pt_este_{k}")
        f._entrada(r, 14, 18, "decimal", f1=(NORTE_MIN, NORTE_MAX), fmt="#,##0.0", h="right",
                   error=f"Norte UTM 17S entre {NORTE_MIN:,} y {NORTE_MAX:,} m.", clave=f"pt_norte_{k}")
        f._entrada(r, 19, 21, "decimal", f1=(0, 4500), fmt="#,##0", h="right", clave=f"pt_alt_{k}")
        f._entrada(r, 22, 25, "texto", clave=f"pt_obs_{k}")
        f.r += 1
    f._kobo(f.r - 8, "r_puntos: pt_gps, pt_tipo", "Puntos de control / vértices", f"B{f.r - 8}:Y{f.r - 1}", "Tabla")
    f.espacio()

    # ---------- 8. panel fotográfico
    corte_3 = f.r
    f.seccion_titulo("  8. PANEL FOTOGRÁFICO")
    fotos = [("foto_1", "Foto 1 – Vista general del predio"), ("foto_2", "Foto 2 – Cobertura vegetal / uso actual"),
             ("foto_3", "Foto 3 – Linderos, hitos o accesos"), ("foto_4", "Foto 4 – Interferencias o detalle relevante")]
    for par in (fotos[:2], fotos[2:]):
        r = f.r
        for (campo, leyenda), (c1, c2) in zip(par, ((2, 13), (14, 25))):
            f.celdas[campo] = f"{get_column_letter(c1)}{r}"
            f._rango(r, c1, r + 11, c2, "Insertar fotografía", font=Font(name=ARIAL, size=7, italic=True, color=GRIS),
                     align=Alignment(horizontal="center", vertical="center"))
            f._rango(r + 12, c1, r + 12, c2, leyenda, fill=ETIQ, font=Font(name=ARIAL, size=8, bold=True),
                     align=Alignment(horizontal="center", vertical="center"))
            f._rango(r + 13, c1, r + 13, c1 + 3, "Archivo / N°:", fill=ETIQ, font=Font(name=ARIAL, size=7, bold=True))
            f._entrada(r + 13, c1 + 4, c2, "texto", clave=f"{campo}_archivo")
            f._kobo(r + 12 if c1 == 2 else r + 13, campo, leyenda, f"{get_column_letter(c1)}{r}", "Imagen")
        f.r += 14
        f.espacio(6)

    # ---------- 9. firmas
    f.seccion_titulo("  9. CONFORMIDAD Y FIRMAS")
    r = f.r
    cajas = ((2, 8, "PREDIANTE / TITULAR (conformidad de la información)"),
             (10, 17, "RESPONSABLE DE LA INFORMACIÓN"), (19, 25, "RESPONSABLE DE LA REVISIÓN"))
    for c1, c2, t in cajas:
        f._rango(r, c1, r, c2, t, fill=ETIQ, font=Font(name=ARIAL, size=7, bold=True),
                 align=Alignment(horizontal="center", vertical="center", wrap_text=True))
        f._rango(r + 1, c1, r + 5, c2, None)
    ws.row_dimensions[r].height = 24
    n = r + 6
    nombres = (f'=IF({f.celdas["tit_nombres"]}<>"",{f.celdas["tit_nombres"]}&" "&{f.celdas["tit_apellidos"]},'
               f'IF({f.celdas["cc_nombre"]}<>"",{f.celdas["cc_nombre"]},"Nombre:"))',
               f'=IF({nombre_asis}<>"",{nombre_asis},"Nombre:")', "Nombre:")
    roles = ("Huella digital (si no firma)", f'=IF({asis}="","Asistente predial","Asistente predial "&{asis})', "Especialista Legal / Técnico – SESDI")
    for (c1, c2, _), nm, rol in zip(cajas, nombres, roles):
        f._rango(n, c1, n, c2, nm, font=Font(name=ARIAL, size=8, bold=True), align=Alignment(horizontal="center"))
        f._rango(n + 1, c1, n + 1, c2, rol, font=Font(name=ARIAL, size=7), align=Alignment(horizontal="center"))
    f.r = n + 2
    f.espacio(6)
    f.fila_doble(dict(texto="FECHA DE ELABORACIÓN", kind="fecha", h="center", clave="fecha_elaboracion"),
                 dict(texto="LUGAR", kind="texto", clave="lugar"))
    f._rango(f.r, C0, f.r, C1, "Los datos personales se registran con el consentimiento del titular y se tratan "
                               "conforme a la Ley N.° 29733, Ley de Protección de Datos Personales. Uso exclusivo del "
                               "proceso de liberación de áreas del Proyecto IN Piura.",
             font=Font(name=ARIAL, size=6, italic=True, color=GRIS), borde=False)
    f.alto(18)
    ultima = f.r
    for fila in range(1, ultima + 1):          # todas las filas con altura fija (posiciones iguales en Excel y LibreOffice)
        if ws.row_dimensions[fila].height is None:
            ws.row_dimensions[fila].height = 15

    # ---------- formato condicional: alertas en ámbar, conformes en verde
    calc_refs = [control, alerta, verif, chk, tipo_u]
    for ref in calc_refs:
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'LEFT({ref},1)="⚠"'],
                                                       font=Font(name=ARIAL, bold=True, color="9C5700"),
                                                       fill=PatternFill("solid", fgColor="FFEB9C")))
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'LEFT({ref},1)="✔"'],
                                                       font=Font(name=ARIAL, bold=True, color="1B4D2E"),
                                                       fill=PatternFill("solid", fgColor="C6EFCE")))

    # ---------- impresión
    ws.print_area = f"A1:Y{ultima}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top, ws.page_margins.bottom = 0.5, 0.6
    for corte in (corte_1, corte_2, corte_3):
        ws.row_breaks.append(Break(id=corte - 1))
    ws.oddFooter.left.text = "Proyecto IN Piura – CUI 2669244 · Ficha técnica predial"
    ws.oddFooter.left.size = 7
    ws.oddFooter.right.text = "Página &P de &N"
    ws.oddFooter.right.size = 7
    ws.freeze_panes = None
    ws.sheet_view.zoomScale = 110
    ws.sheet_properties.tabColor = VERDE
    return f


def _hoja_instrucciones(wb, ficha: _Ficha, asistentes):
    ws = wb.create_sheet("INSTRUCCIONES", 1)
    ws.sheet_view.showGridLines = False
    for i, t in enumerate(ENC, 1):
        c = ws.cell(i, 1, t)
        c.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=VERDE)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
    ws.cell(4, 1, "FICHA TÉCNICA PREDIAL – INSTRUCCIONES Y CORRESPONDENCIA CON KOBOTOOLBOX").font = Font(
        name=ARIAL, bold=True, size=12, color=VERDE)
    pasos = [
        "1. Elija el ASISTENTE PREDIAL: la lista de CÓDIGO DE LA UNIDAD mostrará solo los bloques y lotes SUS que "
        "tiene asignados en el catálogo V6 (hoja CATALOGO_V6). ESP (especialista) ve todas las unidades, incluidas "
        "las «por asignar».",
        "2. Elija la PROVINCIA del predio antes del DISTRITO: el distrito se filtra por provincia (15 distritos del "
        "ámbito). La ficha avisa si el distrito difiere del que registra el catálogo para la unidad.",
        "3. Las coordenadas van en UTM WGS84 Zona 17S: Este entre 450,000 y 750,000 m; Norte entre 9,300,000 y "
        "9,600,000 m. Si tiene lat/lon, conviértalas antes (no se aceptan grados).",
        "4. DNI: 8 dígitos; RUC: 11 dígitos (10… o 20…); celular: 9 dígitos que empiezan con 9. Se guardan como texto "
        "para no perder ceros iniciales.",
        "5. Las listas con etiqueta de KoboToolbox (condición jurídica, estado civil, uso actual, ocupación, aceptación, "
        "estado LA, etc.) usan exactamente las opciones de los formularios F-LA-02 y F-LA-03, de modo que la ficha se "
        "transcribe al formulario Kobo o al aplicativo sin reinterpretar valores.",
        "6. Las celdas grises son automáticas (no las sobrescriba). Los avisos ⚠ señalan datos a verificar antes de "
        "cerrar la ficha; ✔ indica conformidad.",
        "7. El domicilio del prediante acepta valores fuera de la lista (Excel pedirá confirmación), porque puede "
        "residir fuera del ámbito del proyecto.",
        "8. Para imprimir: la hoja FICHA PREDIAL está configurada en A4 vertical (4 páginas). La columna «Campo Kobo» "
        "queda fuera del área de impresión.",
        "9. Sincronización con el aplicativo: una vez llenada, importe esta ficha en Liberación de Áreas › «2 · Plantillas "
        "Excel / Kobo» › Ficha técnica predial completa. Entra como F-LA-02 + F-LA-03, igual que un envío del formulario "
        "Kobo «FP – Ficha técnica predial completa». No cambie el nombre de las hojas ni borre la hoja oculta _MAPA.",
    ]
    r = 6
    for p in pasos:
        c = ws.cell(r, 1, p)
        c.font = Font(name=ARIAL, size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    ws.cell(r, 1, "LEYENDA DE CELDAS").font = Font(name=ARIAL, bold=True, color=VERDE)
    r += 1
    for fill, txt in ((ETIQ, "Etiqueta del dato"), ("FFFFFF", "Entrada libre (texto, número, fecha)"),
                      (LISTA, "Lista desplegable (elija un valor)"), (CALC, "Automático (fórmula, no editar)")):
        a = ws.cell(r, 1, "")
        a.fill = PatternFill("solid", fgColor=fill)
        a.border = BORDE
        ws.cell(r, 2, txt).font = Font(name=ARIAL, size=9)
        r += 1
    r += 1
    ws.cell(r, 1, "ASISTENTES PREDIALES (catálogo V6)").font = Font(name=ARIAL, bold=True, color=VERDE)
    r += 1
    for j, t in enumerate(("Código", "Nombre", "Unidades asignadas"), 1):
        c = ws.cell(r, j, t)
        c.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor=VERDE)
        c.border = BORDE
    r += 1
    for a, nom in asistentes:
        ws.cell(r, 1, a)
        ws.cell(r, 2, nom)
        ws.cell(r, 3, "=COUNTA(CAT_COD)" if a == "ESP" else f'=COUNTIF(CAT_ASIS,"{a}")')
        for j in (1, 2, 3):
            ws.cell(r, j).font = Font(name=ARIAL, size=9)
            ws.cell(r, j).border = BORDE
        r += 1
    ws.cell(r, 2, "Por asignar (sin asistente; disponibles para ESP)").font = Font(name=ARIAL, size=9, italic=True)
    ws.cell(r, 3, '=COUNTIF(CAT_ASIS,"PENDIENTE")').font = Font(name=ARIAL, size=9, italic=True)
    r += 2
    ws.cell(r, 1, "CORRESPONDENCIA FICHA ↔ KOBOTOOLBOX (F-LA-02 / F-LA-03)").font = Font(name=ARIAL, bold=True,
                                                                                         color=VERDE)
    r += 1
    for j, t in enumerate(("Sección", "Dato de la ficha", "Celda", "Campo Kobo (XLSForm)", "Tipo de entrada"), 1):
        c = ws.cell(r, j, t)
        c.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor=VERDE)
        c.border = BORDE
    cab = r
    r += 1
    for k, fila in enumerate(ficha.mapa):
        for j, v in enumerate(fila, 1):
            c = ws.cell(r, j, v.strip())
            c.font = Font(name=ARIAL, size=8)
            c.border = BORDE
            c.alignment = Alignment(wrap_text=True, vertical="center")
            if k % 2:
                c.fill = PatternFill("solid", fgColor=VERDE_CLARO)
        r += 1
    ws.auto_filter.ref = f"A{cab}:E{r - 1}"
    for L, w in zip("ABCDE", (34, 44, 18, 34, 20)):
        ws.column_dimensions[L].width = w
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = "1B4F72"


HOJA_MAPA = "_MAPA"


def _hoja_mapa(wb, ficha: _Ficha):
    """Hoja oculta campo Kobo (XLSForm FP) → celda: permite leer la ficha llenada aunque cambie su diseño."""
    ws = wb.create_sheet(HOJA_MAPA)
    ws.append(["campo", "celda", "version"])
    for campo, celda in ficha.celdas.items():
        ws.append([campo, celda, VERSION])
    ws.sheet_state = "hidden"


def construir_libro(catalogo: list[dict]) -> tuple[Workbook, _Ficha]:
    """Libro completo de la ficha (FICHA PREDIAL, INSTRUCCIONES, CATALOGO_V6, LISTAS, _MAPA) para un catálogo."""
    nombres = {u["asistente"]: u["asistente_nombre"] for u in catalogo}
    asistentes = [(c, nombres.get(c, e)) for c, e in lc.OPCIONES["asistente"]]   # AP-1 … AP-6, ESP (como en Kobo)
    wb = Workbook()
    ficha = _hoja_ficha(wb, asistentes)
    _hoja_instrucciones(wb, ficha, asistentes)
    _hoja_catalogo(wb, catalogo)
    _hoja_listas(wb, catalogo, asistentes)
    _hoja_mapa(wb, ficha)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.active = 0
    wb.properties.title = "Ficha Técnica Predial – Proyecto IN Piura"
    wb.properties.creator = "ANIN – DIME – SESDI"
    return wb, ficha


def mapa_celdas() -> dict[str, str]:
    """Campo → celda de la versión vigente de la ficha (no depende del catálogo)."""
    global _MAPA_CACHE
    if _MAPA_CACHE is None:
        _MAPA_CACHE = dict(_hoja_ficha(Workbook(), [(c, e) for c, e in lc.OPCIONES["asistente"]]).celdas)
    return dict(_MAPA_CACHE)


_MAPA_CACHE: dict[str, str] | None = None


def ficha_bytes(catalogo: list[dict]) -> bytes:
    import io
    wb, _ = construir_libro(catalogo)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generar_ficha_predial(destino: Path = DESTINO, ruta_catalogo: Path = CATALOGO_CSV) -> Path:
    wb, _ = construir_libro(cargar_catalogo(ruta_catalogo))
    destino = Path(destino)
    wb.save(destino)
    return destino


if __name__ == "__main__":
    salida = generar_ficha_predial(Path(sys.argv[1]) if len(sys.argv) > 1 else DESTINO)
    print(f"Ficha predial generada: {salida}")
