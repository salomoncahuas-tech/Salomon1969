# -*- coding: utf-8 -*-
"""
XLSForm KoboToolbox de la Ficha Técnica Predial completa – Proyecto IN Piura (CUI 2669244).

Versión digital de `Ficha_Predial_IN_Piura_F-LA.xlsx` (la_ficha_predial.py) para recoger la ficha en el celular:
- Reutiliza los nombres de campo y las opciones de F-LA-02 / F-LA-03 (`la_campos`) donde el dato ya existe, y las
  listas propias de la ficha impresa (`la_ficha_predial.LISTAS_FICHA`) para los datos nuevos.
- Unidad filtrada por asistente con el mismo `unidades.csv` de los formularios F-LA (select_one_from_file).
- Provincia → distrito en cascada (choice_filter) para el predio y el domicilio.
- Restricciones: DNI, RUC, celular, código de titular, UTM 17S si se digita, área dentro de la unidad ≤ área matriz,
  personería coherente con la condición del prediante.
- El título NO lleva «F-LA»: el aplicativo no lo confunde con F-LA-02 al sincronizar KoboToolbox.

Uso:
    python -m liberacion_areas.la_xlsform_ficha [carpeta_destino]
"""
from __future__ import annotations

import re
import shutil
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import la_campos as lc
from . import la_ficha_predial as fp

FORM_ID = "fp_ficha_predial"
TITULO = "FP – Ficha técnica predial completa · IN Piura"
IDIOMA = "Español (es)"
CSV_UNIDADES = "unidades.csv"
FOTO_MAX_PX = 1280
DESTINO = fp.RAIZ / "kobo_ficha_predial"
ARCHIVO = "FP_Ficha_Predial_Completa_KOBO.xlsx"

_NATURAL = ("posesionario", "ocupante", "sucesion")
_JURIDICA = ("comunal", "estatal")
_PROV_AMBITO = tuple(fp.PROVINCIAS)                 # ayabaca, huancabamba, morropon


def _slug(texto: str) -> str:
    s = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")[:40]


def _o(var: str, valores) -> str:
    return "(" + " or ".join(f"${{{var}}}='{v}'" for v in valores) + ")"


# ------------------------------------------------------------------ listas (choices)
def _choices() -> list[dict]:
    filas: list[dict] = []
    for ln in ("asistente", "si_no", "posicion_sus", "tipo_titularidad", "sucesion_estado", "estado_civil", "docs_exhibidos",
               "uso_actual", "ocupacion", "acceso_tipo", "interferencias", "aceptacion", "modalidad",
               "conclusion_campo", "estado_la", "tipo_acta", "pt_tipo"):
        filas += [{"list_name": ln, "name": c, "label": e} for c, e in lc.OPCIONES[ln]]
    propias = ("personeria", "sit_registral", "cargo_rep", "tipo_predio", "edificaciones", "pendiente", "erosion",
               "cobertura", "tipo_colindante", "und_medida", "material", "conservacion", "construccion", "intervencion")
    for ln in propias:
        codigos = [_slug(e) for e in fp.LISTAS_FICHA[ln]]
        assert len(set(codigos)) == len(codigos), ln
        filas += [{"list_name": ln, "name": c, "label": e} for c, e in zip(codigos, fp.LISTAS_FICHA[ln])]
    filas += [{"list_name": "departamento", "name": "piura", "label": "Piura"}]
    filas += [{"list_name": "dep_domicilio", "name": _slug(e), "label": e} for e in fp.LISTAS_FICHA["dep_domicilio"]]
    filas += [{"list_name": "prov_domicilio", "name": _slug(e), "label": e} for e in fp.LISTAS_FICHA["prov_domicilio"]]
    filas += [{"list_name": "provincia", "name": c, "label": e} for c, e in fp.PROVINCIAS.items()]
    # distrito → provincia, con los mismos códigos (slugs) del catálogo V6 / unidades.csv
    cat = fp.cargar_catalogo()
    inv_p = {v: k for k, v in fp.PROVINCIAS.items()}
    inv_d = {v: k for k, v in fp.DISTRITOS.items()}
    pares = sorted({(inv_p[u["provincia"]], inv_d[u["distrito"]], u["distrito"]) for u in cat}, key=lambda t: t[2])
    filas += [{"list_name": "distrito", "name": d, "label": et, "provincia": p} for p, d, et in pares]
    return filas


# ------------------------------------------------------------------ preguntas (survey)
def _survey() -> list[dict]:
    S: list[dict] = []

    def q(tipo, nombre, etiqueta="", **kw):
        S.append({"type": tipo, "name": nombre, "label": etiqueta, **kw})

    def grupo(nombre, etiqueta, relevante="", lista=False):
        q("begin_group", nombre, etiqueta, relevant=relevante, appearance="field-list" if lista else "")

    def fin(nombre):
        S.append({"type": "end_group", "name": nombre})

    dni = dict(constraint="regex(., '^[0-9]{8}$')", constraint_message="El DNI tiene 8 dígitos")
    cel = dict(constraint="regex(., '^9[0-9]{8}$')", constraint_message="El celular tiene 9 dígitos y empieza con 9")

    q("start", "inicio")
    q("end", "fin")
    q("note", "nota_inicio", "**FICHA TÉCNICA PREDIAL** – Liberación de áreas · Proyecto IN Piura (CUI 2669244)",
      hint="ANIN – DIME – SESDI. Coordenadas en UTM WGS84 Zona 17S.")

    # ---- datos generales
    grupo("g_generales", "Datos generales")
    q("date", "hoy", "Fecha de registro", required="yes", default="today()")
    q("select_one asistente", "asistente", "Asistente predial", required="yes",
      hint="La lista de unidades muestra solo las asignadas al asistente (ESP ve todas).")
    q(f"select_one_from_file {CSV_UNIDADES}", "unidad", "Código de la unidad (bloque / lote SUS)", required="yes",
      choice_filter="asistente=${asistente} or ${asistente}='ESP'", appearance="autocomplete")
    q("calculate", "cod_unidad", calculation="${unidad}")
    for k in ("label", "tipo_unidad", "provincia", "distrito", "bloque_ref", "area_ha", "area_bloque_ha",
              "asistente", "asistente_nombre"):
        nombre = {"label": "unidad_desc", "asistente": "asistente_unidad"}.get(k, k)
        q("calculate", nombre, calculation=f"instance('unidades')/root/item[name=${{unidad}}]/{k}")
    q("note", "nota_unidad", "Unidad: **${unidad_desc}** · asignada a ${asistente_unidad} (${asistente_nombre})",
      relevant="${unidad} != ''")
    q("note", "aviso_asignacion", "⚠ La unidad ${unidad} está asignada a ${asistente_unidad}, no a ${asistente}.",
      relevant="${unidad} != '' and ${asistente} != 'ESP' and ${asistente_unidad} != ${asistente}")
    q("integer", "n_predio", "N.° de predio dentro de la unidad", required="yes", constraint=". >= 0 and . <= 99",
      constraint_message="Entre 0 y 99", hint="Genera el código {UNIDAD}-P{nn}.")
    q("calculate", "cod_predio", calculation="if(${n_predio} > 0, concat(${unidad}, '-P', "
                                            "if(${n_predio} < 10, concat('0', ${n_predio}), ${n_predio})), '')")
    q("note", "nota_cod_predio", "Código de predio: **${cod_predio}**", relevant="${cod_predio} != ''")
    q("date", "fecha_insp", "Fecha de inspección", required="yes", default="today()")
    q("select_one si_no", "consentimiento", "Consentimiento para el uso de datos personales (Ley 29733)",
      required="yes")
    q("note", "aviso_consentimiento", "⚠ Sin consentimiento no registre datos personales del prediante; complete "
      "solo la descripción del predio.", relevant="${consentimiento}='no'")
    fin("g_generales")

    # ---- 1. predio matriz (registral)
    grupo("g_registral", "1. Control e identificación del predio matriz (titularidad registral)")
    q("select_one sit_registral", "sit_registral", "Situación registral")
    q("text", "doc_partida", "Partida electrónica / N.° de título")
    q("text", "oficina_registral", "Oficina registral")
    q("select_one personeria", "reg_personeria", "Tipo de personería del titular registral")
    q("text", "reg_titular", "Titular registral (nombre / razón social)", relevant="${consentimiento}='si'")
    q("select_multiple docs_exhibidos", "docs_exhibidos", "Documentos exhibidos")
    fin("g_registral")

    # ---- 2. prediante
    grupo("g_prediante", "2. Identificación del prediante y tenencia", relevante="${consentimiento}='si'")
    q("select_one tipo_titularidad", "tipo_titularidad", "Condición jurídica del prediante", required="yes")
    q("note", "aviso_ocupante", "⚠ Ocupante sin titularidad ni posesión acreditada: la ficha queda OBSERVADA. "
      "Identifique al titular o posesionario (Especialista Legal).", relevant="${tipo_titularidad}='ocupante'")
    q("select_one sucesion_estado", "sucesion_estado", "Estado de la sucesión",
      relevant="${tipo_titularidad}='sucesion'")
    q("select_one personeria", "personeria", "Tipo de personería", required="yes",
      constraint=f"not({_o('tipo_titularidad', _NATURAL)} and .='juridica') and "
                 f"not({_o('tipo_titularidad', _JURIDICA)} and .='natural')",
      constraint_message="No corresponde a la condición del prediante (posesionario, ocupante y sucesión son "
                         "persona natural; comunidad y Estado, persona jurídica)")
    grupo("g_natural", "2.1 Persona natural", relevante="${personeria}='natural'")
    q("text", "tit_nombres", "Nombres", required="yes")
    q("text", "tit_apellidos", "Apellidos", required="yes")
    q("text", "tit_dni", "DNI", appearance="numbers", **dni)
    q("text", "tit_celular", "Celular", appearance="numbers", **cel)
    q("select_one estado_civil", "tit_estado_civil", "Estado civil")
    q("text", "tit_correo", "Correo electrónico (opcional)",
      constraint="regex(., '^[^@ ]+@[^@ ]+\\.[^@ ]+$')", constraint_message="Correo no válido")
    q("text", "cony_nombre", "Nombre del cónyuge / conviviente",
      relevant="${tit_estado_civil}='casado' or ${tit_estado_civil}='conviviente'")
    q("text", "cony_dni", "DNI del cónyuge", appearance="numbers",
      relevant="${tit_estado_civil}='casado' or ${tit_estado_civil}='conviviente'", **dni)
    fin("g_natural")
    grupo("g_juridica", "2.2 Persona jurídica", relevante="${personeria}='juridica'")
    q("text", "cc_nombre", "Comunidad / entidad / razón social", required="yes")
    q("text", "cc_ruc", "RUC", appearance="numbers", constraint="regex(., '^(10|20)[0-9]{9}$')",
      constraint_message="El RUC tiene 11 dígitos y empieza con 10 o 20")
    q("text", "cc_partida", "Partida registral (SUNARP)")
    q("text", "cc_presidente", "Representante / presidente comunal")
    q("select_one cargo_rep", "rep_cargo", "Cargo")
    q("text", "cc_presidente_dni", "DNI del representante", appearance="numbers", **dni)
    q("text", "cc_presidente_cel", "Celular del representante", appearance="numbers", **cel)
    fin("g_juridica")
    grupo("g_domicilio", "2.3 Domicilio del prediante")
    q("text", "dom_direccion", "Dirección")
    q("select_one dep_domicilio", "dom_departamento", "Departamento", default="piura")
    q("select_one prov_domicilio", "dom_provincia", "Provincia", relevant="${dom_departamento}='piura'")
    q("select_one distrito", "dom_distrito", "Distrito", choice_filter="provincia=${dom_provincia}",
      relevant=_o("dom_provincia", _PROV_AMBITO))
    q("text", "dom_lugar_otro", "Provincia / distrito (fuera del ámbito)",
      relevant=f"${{dom_departamento}} != '' and not({_o('dom_provincia', _PROV_AMBITO)})")
    fin("g_domicilio")
    fin("g_prediante")

    # ---- 3. descripción general
    grupo("g_predio", "3. Descripción general del predio intervenido")
    q("text", "predio_nombre", "Nombre del predio")
    q("text", "centro_poblado", "Centro poblado / sector")
    q("text", "comunidad", "Comunidad campesina (si aplica)")
    q("text", "ref_acceso", "Referencia de acceso", appearance="multiline")
    q("select_one acceso_tipo", "acceso_tipo", "Tipo de acceso")
    q("integer", "tiempo_acceso_min", "Tiempo desde el centro poblado (min)", constraint=". >= 0 and . <= 600",
      constraint_message="Entre 0 y 600 min")
    q("select_one metodo_coord", "metodo_coord", "Coordenadas del punto interior", required="yes", default="gps")
    q("geopoint", "punto_interior", "Punto interior del predio (GPS)", relevant="${metodo_coord}='gps'",
      required="yes", hint="Espere a que la precisión sea ≤ 10 m. El aplicativo lo convierte a UTM 17S.")
    q("decimal", "utm_este", "Este UTM 17S (m)", relevant="${metodo_coord}='utm'", required="yes",
      constraint=f". >= {fp.ESTE_MIN} and . <= {fp.ESTE_MAX}",
      constraint_message=f"Este UTM 17S entre {fp.ESTE_MIN:,} y {fp.ESTE_MAX:,} m")
    q("decimal", "utm_norte", "Norte UTM 17S (m)", relevant="${metodo_coord}='utm'", required="yes",
      constraint=f". >= {fp.NORTE_MIN} and . <= {fp.NORTE_MAX}",
      constraint_message=f"Norte UTM 17S entre {fp.NORTE_MIN:,} y {fp.NORTE_MAX:,} m")
    q("decimal", "utm_altitud", "Altitud (m s.n.m.)", relevant="${metodo_coord}='utm'",
      constraint=". >= 0 and . <= 4500", constraint_message="Entre 0 y 4,500 m s.n.m.")
    q("select_one departamento", "predio_departamento", "Departamento", required="yes", default="piura")
    q("select_one provincia", "predio_provincia", "Provincia", required="yes")
    q("select_one distrito", "predio_distrito", "Distrito", required="yes",
      choice_filter="provincia=${predio_provincia}")
    q("note", "aviso_distrito", "⚠ Según el catálogo V6 la unidad ${unidad} está en el distrito ${distrito}: verifique.",
      relevant="${predio_distrito} != '' and ${distrito} != '' and ${predio_distrito} != ${distrito}")
    # lote SUS: mismo polígono y posición que F-LA-03 (el aplicativo valida ≥ 1 ha, Σ ≤ 10 % del bloque, ≤ 50 m)
    q("select_one posicion_sus", "posicion_sus_campo", "Posición del lote SUS declarada en campo",
      relevant="${tipo_unidad}='lote_sus'")
    q("geoshape", "poligono_sus", "Lote SUS – polígono", relevant="${tipo_unidad}='lote_sus'", required="yes",
      hint="Recorra el perímetro del lote SUS; mínimo 1 ha, ≤ 10 % del bloque, dentro o a ≤ 50 m.")
    fin("g_predio")

    # ---- 4. predio matriz
    grupo("g_matriz", "4. Descripción del predio matriz")
    q("select_one uso_actual", "uso_actual", "Uso actual del suelo")
    q("text", "uso_otro", "Especifique el otro uso", relevant="${uso_actual}='otro'", required="yes")
    q("select_one ocupacion", "ocupacion", "Ocupación observada")
    q("text", "ocupacion_det", "Detalle de la ocupación")
    q("select_multiple edificaciones", "edificaciones", "Edificaciones / cultivos")
    q("text", "edif_det", "Detalle / especifique «Otro»", relevant="selected(${edificaciones}, 'otro')")
    q("select_one pendiente", "pendiente_clase", "Topografía / pendiente")
    q("note", "aviso_pendiente", "⚠ Pendiente mayor a 75 %: fuera del criterio de idoneidad del proyecto.",
      relevant=f"${{pendiente_clase}}='{_slug(fp.LISTAS_FICHA['pendiente'][-1])}'")
    q("select_multiple erosion", "erosion", "Erosión observada")
    q("select_one cobertura", "cobertura", "Estado de la cobertura vegetal")
    q("select_one si_no", "mov_masa", "Evidencia de movimientos en masa")
    q("text", "descripcion_predio", "Otros / descripción del predio", appearance="multiline",
      hint="Cobertura, pendiente, accesos, linderos, signos de degradación o de movimientos en masa.")
    q("select_one tipo_predio", "tipo_predio", "Tipo de predio")
    q("decimal", "predio_area_decl_ha", "Área matriz del predio (ha)", constraint=". >= 0",
      constraint_message="Valor no negativo")
    q("calculate", "predio_area_m2", calculation="if(${predio_area_decl_ha} != '', ${predio_area_decl_ha} * 10000, '')")
    q("decimal", "perimetro_m", "Perímetro (m)", constraint=". >= 0", constraint_message="Valor no negativo")
    fin("g_matriz")
    grupo("g_colindancias", "4.2 Colindancias", lista=True)
    for lado in ("norte", "este", "sur", "oeste"):
        q("text", f"col_{lado}", f"Colindante {lado.upper()}")
        q("select_one tipo_colindante", f"col_{lado}_tipo", f"Tipo de colindante {lado.upper()}")
    fin("g_colindancias")
    q("text", "conflicto_linderos_det", "Conflictos de linderos / observaciones", appearance="multiline")

    # ---- 5. área a intervenir
    grupo("g_area", "5. Descripción del área a intervenir (dentro de la unidad)")
    q("decimal", "predio_area_unidad_ha", "Área del predio dentro de la unidad (ha)",
      constraint="string-length(${predio_area_decl_ha}) = 0 or (. >= 0 and . <= ${predio_area_decl_ha})",
      constraint_message="No puede superar el área matriz del predio")
    q("calculate", "pct_predio", calculation="if(${predio_area_decl_ha} > 0 and ${predio_area_unidad_ha} != '', "
                                            "round(100 * ${predio_area_unidad_ha} div ${predio_area_decl_ha}, 1), '')")
    q("note", "nota_pct", "Área dentro de la unidad: ${pct_predio} % del predio matriz.", relevant="${pct_predio} != ''")
    q("select_multiple interferencias", "interferencias", "Interferencias",
      constraint="not(selected(., 'ninguna') and count-selected(.) > 1)",
      constraint_message="«Ninguna» no se combina con otras opciones")
    q("text", "interf_detalle", "Detalle de interferencias", appearance="multiline",
      relevant="${interferencias} != '' and not(selected(${interferencias}, 'ninguna'))")
    q("select_multiple intervencion", "intervencion_prevista", "Intervención prevista (referencial)")
    q("select_one si_no", "tiene_afectaciones", "¿Existen obras o bienes afectados en el área?")
    fin("g_area")
    q("begin_repeat", "r_afectaciones", "Afectaciones (OC-01, OC-02…)", relevant="${tiene_afectaciones}='si'")
    q("calculate", "af_codigo", calculation="concat('OC-', if(position(..) < 10, '0', ''), position(..))")
    q("note", "af_nota", "**${af_codigo}**")
    q("text", "af_descripcion", "Descripción / uso", required="yes")
    q("decimal", "af_metrado", "Metrado", constraint=". >= 0", constraint_message="Valor no negativo")
    q("select_one und_medida", "af_und", "Unidad de medida")
    q("select_one material", "af_material", "Material predominante")
    q("select_one conservacion", "af_conservacion", "Estado de conservación")
    q("select_one construccion", "af_construccion", "Estado de construcción")
    q("text", "af_ubicacion", "Ubicación en el predio")
    S.append({"type": "end_repeat", "name": "r_afectaciones"})

    # ---- 6. liberación
    grupo("g_liberacion", "6. Disponibilidad y condición de liberación")
    q("select_one aceptacion", "aceptacion", "Aceptación del titular / prediante", required="yes")
    q("select_one modalidad", "modalidad", "Modalidad de disponibilidad")
    q("select_one conclusion_campo", "conclusion_campo", "Conclusión de campo", required="yes")
    q("select_one estado_la", "estado_la_propuesto", "Estado LA propuesto")
    q("select_one tipo_acta", "tipo_acta_prevista", "Acta prevista")
    q("text", "cod_titular_prev", "Código de titular ya registrado (T0000), si existe",
      constraint="regex(., '^T[0-9]{4}$')", constraint_message="Formato T0000 (T y 4 dígitos)")
    q("text", "alertas", "Observaciones / alertas sociales", appearance="multiline")
    fin("g_liberacion")

    # ---- 7. croquis y puntos
    grupo("g_croquis", "7. Croquis del predio y puntos de control")
    q("image", "croquis", "Croquis del predio (dibuje: norte, linderos, accesos, vértices)", appearance="draw",
      parameters=f"max-pixels={FOTO_MAX_PX}")
    fin("g_croquis")
    # repeat con número fijo de filas: abierto, Kobo pedía «agregar otro» sin fin y no dejaba seguir la ficha
    q("integer", "n_puntos", "N.° de puntos de control / vértices del predio", required="yes", default="0",
      constraint=". >= 0 and . <= 50", constraint_message="Entre 0 y 50",
      hint="Indique cuántos registrará (0 a 50). Para corregir, vuelva aquí y cambie el número.")
    q("begin_repeat", "r_puntos", "Puntos de control / vértices del predio", relevant="${n_puntos} > 0",
      repeat_count="${n_puntos}")
    q("geopoint", "pt_gps", "Punto", required="yes", hint="Precisión ≤ 10 m.")
    q("select_one pt_tipo", "pt_tipo", "Tipo de punto")
    S.append({"type": "end_repeat", "name": "r_puntos"})

    # ---- 8. fotos
    grupo("g_fotos", "8. Panel fotográfico")
    for campo, leyenda in (("foto_1", "Foto 1 – Vista general del predio"),
                           ("foto_2", "Foto 2 – Cobertura vegetal / uso actual"),
                           ("foto_3", "Foto 3 – Linderos, hitos o accesos"),
                           ("foto_4", "Foto 4 – Interferencias o detalle relevante")):
        q("image", campo, leyenda, parameters=f"max-pixels={FOTO_MAX_PX}",
          hint="Tome la foto con la cámara del celular (o elija una imagen).")
    fin("g_fotos")

    # ---- 9. conformidad
    grupo("g_firmas", "9. Conformidad")
    q("image", "firma_prediante", "Firma del prediante / titular (conformidad)", appearance="signature",
      parameters="max-pixels=800",
      relevant="${consentimiento}='si'")
    q("text", "lugar", "Lugar")
    fin("g_firmas")
    q("note", "nota_fin", "Fin de la ficha predial. Pulse «Enviar»: el formulario queda en blanco y vuelve al inicio "
                         "para registrar el siguiente predio.",
      hint="Use el enlace «En línea-Fuera de línea (envíos múltiples)» de KoboToolbox o KoboCollect.")
    return S


_COLS_SURVEY = ["type", "name", "label", "hint", "required", "relevant", "constraint", "constraint_message",
                "calculation", "choice_filter", "appearance", "default", "parameters", "repeat_count"]
_TRAD = ("label", "hint", "constraint_message")


def hojas() -> dict[str, tuple[list[str], list[dict]]]:
    """survey, choices y settings como (columnas, filas). Las columnas traducibles llevan el idioma (label::Español
    (es)); sin sufijo junto a default_language, Kobo crea una traducción sin nombre y el editor no abre el formulario."""
    choices = _choices() + [{"list_name": "metodo_coord", "name": "gps", "label": "GPS del celular"},
                            {"list_name": "metodo_coord", "name": "utm", "label": "Digitar UTM 17S (Este / Norte)"}]
    settings = [{"form_title": TITULO, "form_id": FORM_ID, "version": datetime.now().strftime("%Y%m%d%H%M"),
                 "default_language": IDIOMA,
                 "instance_name": "concat(${cod_predio}, ' – ', ${tit_apellidos}, ${cc_nombre})"}]
    return {"survey": (_COLS_SURVEY, _survey()),
            "choices": (["list_name", "name", "label", "provincia"], choices),
            "settings": (list(settings[0]), settings)}


def generar_xlsform(destino) -> Path | None:
    """Escribe el XLSForm en `destino` (ruta o archivo abierto / BytesIO)."""
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, (cols, filas) in hojas().items():
        ws = wb.create_sheet(nombre)
        cab = [f"{c}::{IDIOMA}" if c in _TRAD else c for c in cols]
        ws.append(cab)
        for f in filas:
            ws.append([f.get(c, "") or None for c in cols])
        for j, _ in enumerate(cab, 1):
            ws.cell(1, j).font = Font(name="Arial", bold=True, color="FFFFFF")
            ws.cell(1, j).fill = PatternFill("solid", fgColor=fp.VERDE)
            ws.column_dimensions[get_column_letter(j)].width = 18
        ws.freeze_panes = "A2"
    for L, w in (("A", 34), ("B", 24), ("C", 60), ("D", 40), ("F", 30), ("G", 40), ("H", 36), ("I", 40), ("J", 32)):
        wb["survey"].column_dimensions[L].width = w
    wb["choices"].column_dimensions["C"].width = 44
    if isinstance(destino, (str, Path)):
        destino = Path(destino)
    wb.save(destino)
    return destino if isinstance(destino, Path) else None


LEEME = f"""FICHA TÉCNICA PREDIAL COMPLETA – XLSForm para KoboToolbox
Proyecto IN Piura (CUI 2669244) · ANIN – DIME – SESDI

Archivos
  {ARCHIVO}   formulario (XLSForm)
  {CSV_UNIDADES}                       catálogo V6 de unidades con asistente (el mismo de los formularios F-LA)

Cómo cargarlo en kobotoolbox.org
  1. Nuevo › Cargar un XLSForm › elija {ARCHIVO} › Crear proyecto.
  2. En el proyecto: Configuración › Multimedia › cargue {CSV_UNIDADES} (sin él la lista de unidades sale vacía).
  3. Desplegar. Para llenar: Formulario › Recolectar datos › «En línea-Fuera de línea (envíos múltiples)» o KoboCollect.
  4. Si cambian las asignaciones de asistentes, reemplace solo {CSV_UNIDADES} en Multimedia y vuelva a desplegar.

Notas
  - Las coordenadas del GPS se guardan en lat/lon (formato de Kobo); el aplicativo las convierte a UTM WGS84 17S.
    Si se digita UTM, Este debe estar entre {fp.ESTE_MIN:,} y {fp.ESTE_MAX:,} m y Norte entre {fp.NORTE_MIN:,} y {fp.NORTE_MAX:,} m.
  - Sincronización con el aplicativo (Liberación de Áreas): el aplicativo reconoce este formulario por su título
    («Ficha técnica predial») y cada envío entra como F-LA-02 (titular) + F-LA-03 (inspección), con las mismas
    validaciones; los datos propios de la ficha (colindancias, afectaciones, domicilio…) se conservan en el registro.
    El aviso «envíos de KoboToolbox aún no importados» de la página incluye este formulario.
  - Lotes SUS: el formulario pide el polígono del lote (geoshape) como el F-LA-03.
  - Puntos de control / vértices: primero se indica cuántos se registrarán (0 a 50) y el formulario pide solo
    esos puntos y continúa con las fotos. Para agregar o quitar puntos, vuelva a esa pregunta y cambie el número.
  - Se regenera con: python -m liberacion_areas.la_xlsform_ficha
"""


def paquete_zip(unidades_csv: bytes) -> bytes:
    """ZIP con el XLSForm, el unidades.csv indicado (catálogo vigente del aplicativo) y el LEEME."""
    import io
    import zipfile
    buf, xls = io.BytesIO(), io.BytesIO()
    generar_xlsform(xls)
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(ARCHIVO, xls.getvalue())
        z.writestr(CSV_UNIDADES, unidades_csv)
        z.writestr("LEEME.txt", LEEME.encode("utf-8"))
    return buf.getvalue()


def generar_paquete(carpeta: Path = DESTINO) -> Path:
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    generar_xlsform(carpeta / ARCHIVO)
    shutil.copyfile(fp.CATALOGO_CSV, carpeta / CSV_UNIDADES)
    (carpeta / "LEEME.txt").write_text(LEEME, encoding="utf-8")
    return carpeta


if __name__ == "__main__":
    salida = generar_paquete(Path(sys.argv[1]) if len(sys.argv) > 1 else DESTINO)
    print(f"XLSForm de la ficha predial generado en: {salida}")
