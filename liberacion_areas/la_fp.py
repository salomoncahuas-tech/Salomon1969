# -*- coding: utf-8 -*-
"""
Ficha Técnica Predial completa (FP) ↔ aplicativo de Liberación de Áreas – Proyecto IN Piura (CUI 2669244).

La ficha existe en tres formas sincronizadas por los mismos nombres de campo (XLSForm FP, la_xlsform_ficha):
  1. Formulario KoboToolbox «FP – Ficha técnica predial completa»  → sincronización / importación (la_kobo).
  2. Ficha Excel imprimible llenada (Ficha_Predial_IN_Piura_F-LA.xlsx) → importación en «Plantillas Excel / Kobo».
  3. Ficha Excel generada desde el aplicativo con los datos de un predio (y sus fotos)  → «Historial / Edición».

Cada ficha FP entra al aplicativo como DOS envíos: F-LA-02 (titular) y F-LA-03 (inspección), con identificador
{uuid}-F02 / {uuid}-F03, y pasa por core.validar_envio y la_db.importar sin lógica paralela. El payload de ambos
conserva todos los datos de la ficha (colindancias, afectaciones, domicilio, personería…).
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime

import pandas as pd

from . import la_campos as lc
from . import la_core as core
from . import la_ficha_predial as fpx
from . import la_xlsform_ficha as xf

FORM_ID = xf.FORM_ID                      # fp_ficha_predial
CODIGO = "FP"
NOMBRE = "Ficha técnica predial completa (F-LA-02 + F-LA-03)"
# Formulario desplegado en kf.kobotoolbox.org (además se reconoce por el título)
UIDS_FP = {"aCyBhgqeCTZ3UKzmBW8tx2"}
_PATRON_FP = re.compile(r"FICHA\s+T[EÉ]CNICA\s+PREDIAL|^\s*FP\s*[-–—:]", re.I)
PARTES = {"f_la_02_titular": "F02", "f_la_03_inspeccion": "F03"}
CAMPOS_FIRMA = {"firma_prediante"}
_RESERVADOS = ("_dni", "_celular", "_cel", "_ruc", "_correo")


# ------------------------------------------------------------------ reconocimiento del formulario Kobo
def uids_fp() -> set[str]:
    """UID conocidos del formulario FP (+ KOBO_FP_UID de los secrets, si existe)."""
    uids = set(UIDS_FP)
    try:
        from odk_kobo import secreto_kobo
        extra = secreto_kobo("KOBO_FP_UID")
        if extra:
            uids |= {u.strip() for u in str(extra).replace(";", ",").split(",") if u.strip()}
    except Exception:  # noqa: BLE001 – fuera de Streamlit o sin secrets
        pass
    return uids


def es_formulario_fp(nombre: str | None, uid: str | None = None) -> bool:
    return bool(uid and uid in uids_fp()) or bool(_PATRON_FP.search(str(nombre or "")))


def uuid_parte(uid: str, form_id: str) -> str:
    return f"{uid}-{PARTES[form_id]}"


def uuids_partes(uid: str | None) -> list[str]:
    return [uuid_parte(uid, f) for f in PARTES] if uid else []


# ------------------------------------------------------------------ ficha → envíos F-LA-02 / F-LA-03
def _campo_adjunto(a: dict) -> str:
    return str(a.get("question_xpath") or a.get("filename") or "").split("/")[-1]


def partes(reg: dict) -> dict[str, dict]:
    """Envío FP aplanado → {f_la_02_titular: envío, f_la_03_inspeccion: envío} con los nombres de campo F-LA."""
    base = dict(reg)
    uid = core.kobo_uuid(reg)
    unidad = base.get("unidad") or base.get("cod_unidad")
    if unidad is not None:
        unidad = str(unidad).strip()
        base["unidad"] = base["cod_unidad"] = unidad
    n = core._num(base.get("n_predio"))
    if not base.get("cod_predio"):
        base["cod_predio"] = lc.cod_predio(unidad, n) if n else None
    # punto interior: GPS (geopoint) o UTM 17S digitado
    if base.get("metodo_coord") == "utm" or (not base.get("punto_interior") and base.get("utm_este")):
        gp = lc.utm_a_geopoint(base.get("utm_este"), base.get("utm_norte"), base.get("utm_altitud"),
                               base.get("utm_precision"))
        if gp:
            base["punto_interior"] = gp
    # titular estatal: F-LA-02 lo guarda en est_entidad / est_contacto
    if base.get("tipo_titularidad") == "estatal":
        base.setdefault("est_entidad", base.get("cc_nombre"))
        base.setdefault("est_contacto", base.get("cc_presidente"))
    if base.get("uso_actual") and not base.get("uso_observado"):
        base["uso_observado"] = base["uso_actual"]
    base["fp_uuid"] = uid
    base["origen_ficha"] = CODIGO
    adjuntos = base.get("_attachments") or []
    propios = base.pop("_uuid_partes", None) or {}      # ficha exportada por el aplicativo: sus envíos de origen
    salida = {}
    for fid, suf in PARTES.items():
        d = dict(base)
        d.pop("meta/instanceID", None)
        d.pop("_uuid", None)
        d["meta_instanceID"] = propios.get(fid) or (f"{uid}-{suf}" if uid else None)
        d["_attachments"] = [a for a in adjuntos if (_campo_adjunto(a) in CAMPOS_FIRMA) == (fid == "f_la_02_titular")]
        salida[fid] = d
    return salida


def validar(reg: dict, catalogo: dict, geoms: dict, existentes: set[str], sus_areas: dict | None = None,
            omitir: set[str] | None = None) -> list[core.Resultado]:
    """Valida una ficha FP como F-LA-02 + F-LA-03 (core.validar_envio). Las partes cuyo identificador está en
    `omitir` (p. ej. eliminadas en el aplicativo) no se devuelven."""
    out = []
    for fid, d in partes(reg).items():
        if omitir and core.kobo_uuid(d) in omitir:
            continue
        cat = catalogo.get(d.get("cod_unidad")) or {}
        for k in ("tipo_unidad", "bloque_ref", "area_bloque_ha"):
            if d.get(k) in (None, "") and cat.get(k) not in (None, ""):
                d[k] = cat[k]
        r = core.validar_envio(fid, d, catalogo, geoms, existentes, sus_areas)
        r.form_id = fid
        if fid == "f_la_03_inspeccion" and r.estado_import != "DUPLICADO":
            _validar_ubicacion(r, d, cat)
        out.append(r)
    return out


def _validar_ubicacion(r: core.Resultado, d: dict, cat: dict):
    """Distrito del predio: debe pertenecer a la provincia elegida (en Excel y en la digitación la lista no se filtra
    por provincia); si difiere del distrito de la unidad en el catálogo V6 solo se avisa (un bloque puede abarcar
    dos distritos)."""
    prov, dist = d.get("predio_provincia"), d.get("predio_distrito")
    if prov and dist and _PROVINCIA_DE.get(dist) not in (None, prov):
        r.observar(f"Distrito del predio «{etiqueta('predio_distrito', dist)}» no pertenece a la provincia "
                   f"«{etiqueta('predio_provincia', prov)}»")
    if dist and cat.get("distrito") and dist != cat.get("distrito"):
        r.motivos.append(f"Aviso: la unidad {d.get('cod_unidad')} figura en el distrito "
                         f"{etiqueta('predio_distrito', cat.get('distrito'))} (catálogo V6)")


def importar(conn, resultados: list[core.Resultado], fuente: str, usuario: str) -> int:
    """Registra los resultados agrupados por formulario (F-LA-02 / F-LA-03). Devuelve los insertados."""
    from . import la_db as db
    n = 0
    for fid in PARTES:
        grupo = [r for r in resultados if r.form_id == fid]
        if grupo:
            n += db.importar(conn, grupo, fuente, fid, usuario)["insertados"]
    return n


# ------------------------------------------------------------------ listas (código ↔ etiqueta)
def _listas() -> tuple[dict[str, str], dict[str, dict[str, str]]]:
    """(campo → lista, lista → {código: etiqueta}) del XLSForm FP."""
    campo_lista = {}
    for f in xf.hojas()["survey"][1]:
        t = f["type"].split()
        if len(t) == 2 and t[0] in ("select_one", "select_multiple"):
            campo_lista[f["name"]] = t[1]
    opciones: dict[str, dict[str, str]] = {}
    for c in xf.hojas()["choices"][1]:
        opciones.setdefault(c["list_name"], {})[c["name"]] = c["label"]
    return campo_lista, opciones


_CAMPO_LISTA, _OPCIONES = _listas()
_PROVINCIA_DE = {c["name"]: c.get("provincia") for c in xf.hojas()["choices"][1] if c["list_name"] == "distrito"}
# Celdas de la ficha que repiten un campo (selección múltiple o tablas) → campo base del XLSForm
_BASE = {**{f"docs_exhibidos_{k}": "docs_exhibidos" for k in (1, 2, 3)},
         **{f"interferencias_{k}": "interferencias" for k in (1, 2, 3)},
         **{f"intervencion_prevista_{k}": "intervencion_prevista" for k in (1, 2)},
         **{f"af_{c}_{k}": f"af_{c}" for c in ("und", "material", "conservacion", "construccion") for k in range(1, 6)},
         **{f"pt_tipo_{k}": "pt_tipo" for k in range(1, 9)}}


def _lista_de(campo: str) -> str | None:
    return _CAMPO_LISTA.get(_BASE.get(campo, campo))


def etiqueta(campo: str, codigo) -> str | None:
    if codigo in (None, ""):
        return None
    lista = _lista_de(campo)
    return _OPCIONES.get(lista, {}).get(str(codigo), str(codigo)) if lista else codigo


def codigo(campo: str, valor) -> tuple[str | None, bool]:
    """Etiqueta o código → (código, reconocido)."""
    if valor in (None, ""):
        return None, True
    lista = _lista_de(campo)
    v = str(valor).strip()
    if not lista:
        return v, True
    for cod, et in _OPCIONES.get(lista, {}).items():
        if v == cod or v.lower() == cod.lower() or v.lower() == et.lower():
            return cod, True
    return v, False


# ------------------------------------------------------------------ lectura de la ficha Excel llenada
_FECHAS = ("hoy", "fecha_insp", "fecha_elaboracion")
_ENTEROS = ("n_predio", "tiempo_acceso_min")
_DECIMALES = ("predio_area_decl_ha", "predio_area_unidad_ha", "perimetro_m", "utm_este", "utm_norte", "utm_altitud",
              "utm_precision")
_TEXTO_NUM = ("tit_dni", "cony_dni", "cc_ruc", "cc_presidente_dni", "tit_celular", "cc_presidente_cel", "unidad",
              "doc_partida", "cc_partida")
_OBLIGATORIOS = {"asistente": "Asistente predial", "unidad": "Código de la unidad", "n_predio": "N° de predio",
                 "fecha_insp": "Fecha de inspección", "consentimiento": "Consentimiento (Ley 29733)",
                 "aceptacion": "Aceptación del titular", "conclusion_campo": "Conclusión de campo"}


def _fecha(v) -> str | None:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return lc._fecha_txt(v)


def _texto(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


_PREFIJO_UUID = "__uuid__"


def _mapa_libro(wb) -> tuple[dict[str, str], list[str], dict[str, str]]:
    """(campo → celda, avisos, {form_id: kobo_uuid} si la ficha la generó el aplicativo desde un predio)."""
    avisos = []
    if fpx.HOJA_MAPA in wb.sheetnames:
        filas = [r for r in wb[fpx.HOJA_MAPA].iter_rows(min_row=2, values_only=True) if r and r[0] and r[1]]
        uuids = {str(r[0])[len(_PREFIJO_UUID):]: str(r[1]) for r in filas if str(r[0]).startswith(_PREFIJO_UUID)}
        return {r[0]: r[1] for r in filas if not str(r[0]).startswith(_PREFIJO_UUID)}, avisos, uuids
    ws = wb["FICHA PREDIAL"]
    version = " ".join(str(c.value) for fila in ws.iter_rows(min_row=1, max_row=10) for c in fila
                       if isinstance(c.value, str) and "Versión" in c.value)
    if fpx.VERSION not in version:
        avisos.append(f"La ficha no indica la versión {fpx.VERSION}: verifique que sea la plantilla vigente.")
    return fpx.mapa_celdas(), avisos, {}


def leer_ficha_excel(archivo, nombre: str | None = None) -> dict:
    """Ficha Excel llenada → {"registro": envío FP aplanado, "avisos": [...], "archivo": nombre}."""
    from openpyxl import load_workbook
    wb = load_workbook(archivo, data_only=True)
    if "FICHA PREDIAL" not in wb.sheetnames:
        raise ValueError("El archivo no tiene la hoja «FICHA PREDIAL» (use la ficha técnica predial del proyecto).")
    ws = wb["FICHA PREDIAL"]
    mapa, avisos, uuids = _mapa_libro(wb)
    crudo = {k: ws[c].value for k, c in mapa.items() if k not in ("croquis", "foto_1", "foto_2", "foto_3", "foto_4")}
    reg: dict = {}
    multi: dict[str, list[str]] = {}
    for k, v in crudo.items():
        if v in (None, ""):
            continue
        if re.match(r"^(af|pt)_.+_\d+$", k):
            continue                                            # tablas: se arman abajo
        if k in _FECHAS:
            reg[k] = _fecha(v)
        elif k in _ENTEROS:
            n = core._num(v)
            reg[k] = int(n) if n is not None else None
        elif k in _DECIMALES:
            reg[k] = core._num(v)
        elif k in _TEXTO_NUM:
            reg[k] = _texto(v)
        else:
            cod, ok = codigo(k, v)
            if not ok:
                if k == "dom_distrito" or k == "dom_provincia":
                    reg["dom_lugar_otro"] = " / ".join(x for x in (reg.get("dom_lugar_otro"), str(v)) if x)
                    continue
                avisos.append(f"«{v}» no es una opción de la lista ({k}); se registra tal cual.")
            if k in _BASE:
                multi.setdefault(_BASE[k], []).append(cod)
            else:
                reg[k] = cod
    for k, vals in multi.items():
        reg[k] = " ".join(dict.fromkeys(v for v in vals if v))
    # afectaciones y puntos de control
    af = []
    for i in range(1, 6):
        fila = {c: crudo.get(f"af_{c}_{i}") for c in ("descripcion", "metrado", "und", "material", "conservacion",
                                                        "construccion", "ubicacion")}
        if any(v not in (None, "") for v in fila.values()):
            af.append({"af_codigo": f"OC-{i:02d}", "af_descripcion": _texto(fila["descripcion"]),
                       "af_metrado": core._num(fila["metrado"]), "af_ubicacion": _texto(fila["ubicacion"]),
                       **{f"af_{c}": codigo(f"af_{c}_{i}", fila[c])[0]
                          for c in ("und", "material", "conservacion", "construccion")}})
    if af:
        reg["tiene_afectaciones"], reg["r_afectaciones"] = "si", af
    pts = []
    for i in range(1, 9):
        e, n_, alt = crudo.get(f"pt_este_{i}"), crudo.get(f"pt_norte_{i}"), crudo.get(f"pt_alt_{i}")
        if core._num(e) is None and core._num(n_) is None:
            continue
        if not lc.utm_en_rango(e, n_):
            avisos.append(f"Punto de control {i}: Este {e} / Norte {n_} fuera del rango UTM 17S del ámbito.")
        pts.append({"pt_gps": lc.utm_a_geopoint(e, n_, alt, 0), "pt_tipo": codigo(f"pt_tipo_{i}", crudo.get(f"pt_tipo_{i}"))[0],
                    "pt_obs": _texto(crudo.get(f"pt_obs_{i}"))})
    if pts:
        reg["r_puntos"] = pts
    if reg.get("utm_este") is not None or reg.get("utm_norte") is not None:
        reg["metodo_coord"] = "utm"
        if not lc.utm_en_rango(reg.get("utm_este"), reg.get("utm_norte")):
            avisos.append(f"Punto interior: Este {reg.get('utm_este')} / Norte {reg.get('utm_norte')} fuera del rango UTM "
                          f"17S del ámbito (E {core.E_MIN:,}–{core.E_MAX:,}; N {core.N_MIN:,}–{core.N_MAX:,}).")
    reg.setdefault("hoy", reg.get("fecha_elaboracion") or reg.get("fecha_insp"))
    for k, et in _OBLIGATORIOS.items():
        if reg.get(k) in (None, ""):
            avisos.append(f"Falta {et}")
    if reg.get("consentimiento") == "si" and not reg.get("tipo_titularidad"):
        avisos.append("Falta Condición jurídica del prediante")
    if reg.get("unidad") and reg.get("n_predio"):
        reg["cod_predio"] = lc.cod_predio(reg["unidad"], reg["n_predio"])
    ident = {k: v for k, v in reg.items() if k not in ("hoy",)}
    reg["meta_instanceID"] = lc.uuid_contenido(FORM_ID, ident, prefijo="xls-fp")
    if uuids:          # ficha generada por el aplicativo: re-importarla no duplica el predio (DUPLICADO)
        reg["_uuid_partes"] = {fid: u for fid, u in uuids.items() if fid in PARTES}
        avisos.append("Ficha generada por el aplicativo: sus registros ya existen; para corregirlos use "
                      "«Historial / Edición».")
    return {"registro": reg, "avisos": avisos, "archivo": nombre or getattr(archivo, "name", "")}


# ------------------------------------------------------------------ ficha Excel llenada desde el aplicativo
def _unir(payloads: list[dict]) -> dict:
    d: dict = {}
    for p in payloads:
        for k, v in (p or {}).items():
            if v not in (None, "", []):
                d[k] = v
    if not d.get("utm_este") and d.get("punto_interior"):
        u = lc.geopoint_a_utm(d["punto_interior"])
        if u:
            d.update({"utm_este": u.get("este"), "utm_norte": u.get("norte"), "utm_altitud": u.get("alt"),
                      "utm_precision": u.get("prec")})
    if d.get("tipo_titularidad") == "estatal":
        d.setdefault("cc_nombre", d.get("est_entidad_otra") or d.get("est_entidad"))
        d.setdefault("cc_presidente", d.get("est_contacto"))
    if not d.get("personeria") and d.get("tipo_titularidad"):
        t = d["tipo_titularidad"]
        d["personeria"] = "juridica" if t in ("comunal", "estatal") else ("natural" if t != "sin_titular" else None)
    d["unidad"] = d.get("unidad") or d.get("cod_unidad")
    return d


def _imagen(contenido: bytes, ancho_px: int, alto_px: int):
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image
    im = Image.open(io.BytesIO(contenido))
    im = im.convert("RGB")
    im.thumbnail((ancho_px, alto_px))
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=85)
    buf.seek(0)
    img = XLImage(buf)
    img.width, img.height = im.size
    return img


_COL_PX, _FILA_PX = 37, 20          # columna de la ficha (4.6 caracteres) y fila estándar (15 pt), en píxeles


def _ancla_marco(celda: str, w: int, h: int, n_cols: int, n_filas: int):
    """Ancla de dos celdas para una imagen dentro de su marco (`n_cols` × `n_filas` desde `celda`): la imagen queda
    sujeta a las celdas, no a una posición en píxeles, y conserva aproximadamente su proporción."""
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
    col, fila = coordinate_from_string(celda)
    c0, r0 = column_index_from_string(col) - 1, fila - 1
    filas = n_filas - 2                                     # una fila libre arriba y abajo
    cols = max(2, min(n_cols - 1, round(w / h * filas * _FILA_PX / _COL_PX)))
    c_ini = c0 + (n_cols - cols) // 2
    ancla = TwoCellAnchor(editAs="oneCell")
    ancla._from = AnchorMarker(col=c_ini, colOff=0, row=r0 + 1, rowOff=0)
    ancla.to = AnchorMarker(col=c_ini + cols, colOff=0, row=r0 + 1 + filas, rowOff=0)
    return ancla


def ficha_excel_llena(catalogo: list[dict], payloads: list[dict], imagenes: dict[str, bytes] | None = None,
                      uuids: dict[str, str] | None = None) -> bytes:
    """Ficha Excel ANIN con los datos de un predio (payloads de F-LA-02 / F-LA-03 / FP) y sus fotos.
    `uuids` ({form_id: kobo_uuid} de los envíos de origen) se guarda en la hoja _MAPA para que re-importar la ficha
    no duplique el predio."""
    wb, ficha = fpx.construir_libro(catalogo)
    for fid, u in (uuids or {}).items():
        wb[fpx.HOJA_MAPA].append([f"{_PREFIJO_UUID}{fid}", u, fpx.VERSION])
    ws = wb["FICHA PREDIAL"]
    d = _unir(payloads)
    mapa = ficha.celdas

    def poner(clave, valor):
        if clave in mapa and valor not in (None, ""):
            ws[mapa[clave]] = valor

    for clave in mapa:
        if clave in ("croquis", "foto_1", "foto_2", "foto_3", "foto_4") or clave in _BASE or \
                re.match(r"^(af|pt)_.+_\d+$", clave):
            continue
        v = d.get(clave)
        if v in (None, ""):
            continue
        if clave in _FECHAS:
            f = _fecha(v)
            try:
                v = datetime.strptime(f, "%Y-%m-%d") if f else None
            except ValueError:
                v = f
        elif clave in _ENTEROS:
            n = core._num(v)
            v = int(n) if n is not None else v
        elif clave in _DECIMALES:
            v = core._num(v)
        elif clave in _TEXTO_NUM:
            v = _texto(v)
        else:
            v = etiqueta(clave, v)
        poner(clave, v)
    if not d.get("dom_distrito") and d.get("dom_lugar_otro"):
        poner("dom_distrito", d["dom_lugar_otro"])
    for base, n in (("docs_exhibidos", 3), ("interferencias", 3), ("intervencion_prevista", 2)):
        cods = lc.codigos_multiples(_CAMPO_LISTA.get(base, ""), d.get(base)) if d.get(base) else []
        for i, c in enumerate(cods[:n], 1):
            poner(f"{base}_{i}", etiqueta(f"{base}_{i}", c))
    for i, a in enumerate((d.get("r_afectaciones") or [])[:5], 1):
        poner(f"af_descripcion_{i}", a.get("af_descripcion"))
        poner(f"af_metrado_{i}", core._num(a.get("af_metrado")))
        poner(f"af_ubicacion_{i}", a.get("af_ubicacion"))
        for c in ("und", "material", "conservacion", "construccion"):
            poner(f"af_{c}_{i}", etiqueta(f"af_{c}_{i}", a.get(f"af_{c}")))
    for i, p in enumerate((d.get("r_puntos") or [])[:8], 1):
        u = lc.geopoint_a_utm(p.get("pt_gps"))
        poner(f"pt_tipo_{i}", etiqueta(f"pt_tipo_{i}", p.get("pt_tipo")))
        poner(f"pt_este_{i}", u.get("este"))
        poner(f"pt_norte_{i}", u.get("norte"))
        poner(f"pt_alt_{i}", u.get("alt"))
        poner(f"pt_obs_{i}", p.get("pt_obs"))
    for campo, contenido in (imagenes or {}).items():
        if campo not in mapa or not contenido:
            continue
        # marco de la foto: 12 columnas × 12 filas (croquis: 24 × 14); la imagen va centrada dentro del marco
        n_cols, n_filas = (24, 14) if campo == "croquis" else (12, 12)
        try:
            img = _imagen(contenido, 1000, 400)
        except Exception:  # noqa: BLE001 – adjunto que no es imagen
            continue
        ws.add_image(img, _ancla_marco(mapa[campo], img.width, img.height, n_cols, n_filas))
        ws[mapa[campo]] = None
        poner(f"{campo}_archivo", d.get(campo))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ficha_predio_bytes(conn, cod_predio: str) -> bytes:
    """Ficha Excel llena de un predio con sus F-LA-02 / F-LA-03 registrados (de cualquier vía, el más reciente
    prevalece) y sus fotos guardadas en la base."""
    import json
    from . import la_db as db
    envs = db.df(conn, """SELECT kobo_uuid, form_id, payload FROM la_envios_raw WHERE cod_predio = %s
                          AND form_id IN ('f_la_02_titular', 'f_la_03_inspeccion') ORDER BY creado""", (cod_predio,))
    payloads, imagenes, uuids = [], {}, {}
    for e in (envs.to_dict("records") if not envs.empty else []):
        uuids[e["form_id"]] = e["kobo_uuid"]
        p = e["payload"]
        payloads.append(json.loads(p) if isinstance(p, str) else p)
        for a in db.adjuntos_envio(conn, e["kobo_uuid"]).to_dict("records"):
            if a["contenido"] is not None and str(a["mimetype"] or "").startswith("image") and \
                    a["campo"] in ("foto_1", "foto_2", "foto_3", "foto_4", "croquis"):
                imagenes[a["campo"]] = bytes(a["contenido"])
    return ficha_excel_llena(fpx.catalogo_desde_unidades(db.catalogo(conn)), payloads, imagenes, uuids)


# ------------------------------------------------------------------ datos propios de la ficha (detalle)
def _etiquetas_campos() -> dict[str, str]:
    return {f["name"]: f.get("label") or f["name"] for f in xf.hojas()["survey"][1]
            if f["type"] not in ("begin_group", "end_group", "begin_repeat", "end_repeat", "note", "calculate",
                                 "start", "end")}


def datos_adicionales(payload: dict) -> pd.DataFrame:
    """Campos de la ficha FP que no están en F-LA-02 / F-LA-03 (colindancias, domicilio, personería…), legibles y
    sin datos reservados (DNI, RUC, celular, correo – Ley 29733)."""
    propios = {c.nombre for fid in PARTES for c in lc.FORMULARIOS[fid].campos} | {"unidad", "hoy", "asistente"}
    filas = []
    for k, et in _etiquetas_campos().items():
        if k in propios or k in ("poligono_sus", "croquis", "foto_1", "foto_2", "foto_3", "foto_4", "firma_prediante"):
            continue
        v = payload.get(k)
        if v in (None, "", []):
            continue
        if k.startswith("dom_"):
            et = f"Domicilio – {et}"
        elif k.startswith("predio_") and k in ("predio_departamento", "predio_provincia", "predio_distrito"):
            et = f"Predio – {et}"
        if k.endswith(_RESERVADOS):
            v = "[dato reservado – Ley 29733]"
        elif _lista_de(k):
            v = ", ".join(etiqueta(k, c) or "" for c in str(v).split())
        filas.append({"campo": et, "valor": str(v)})
    return pd.DataFrame(filas)


def tabla_afectaciones(payload: dict) -> pd.DataFrame:
    filas = []
    for i, a in enumerate(payload.get("r_afectaciones") or [], 1):
        filas.append({"código": a.get("af_codigo") or f"OC-{i:02d}", "descripción": a.get("af_descripcion"),
                      "metrado": a.get("af_metrado"), "und.": etiqueta("af_und", a.get("af_und")),
                      "material": etiqueta("af_material", a.get("af_material")),
                      "conservación": etiqueta("af_conservacion", a.get("af_conservacion")),
                      "construcción": etiqueta("af_construccion", a.get("af_construccion")),
                      "ubicación": a.get("af_ubicacion")})
    return pd.DataFrame(filas).fillna("")


# ------------------------------------------------------------------ digitación en el aplicativo («1 · Registro en campo»)
def _registrar_opciones() -> dict[str, str]:
    """Listas propias de la ficha en lc.OPCIONES con prefijo fp_ (mismas opciones y códigos que el XLSForm FP)."""
    nombres = {}
    for lista in ("personeria", "sit_registral", "cargo_rep", "tipo_predio", "edificaciones", "pendiente", "erosion",
                  "cobertura", "tipo_colindante", "und_medida", "material", "conservacion", "construccion",
                  "intervencion", "departamento", "dep_domicilio", "prov_domicilio", "provincia", "distrito"):
        clave = f"fp_{lista}"
        lc.OPCIONES[clave] = list(_OPCIONES[lista].items())
        nombres[lista] = clave
    return nombres


_L = _registrar_opciones()
_SI = {"consentimiento": ["si"]}


def _formulario_app() -> lc.Formulario:
    C, R = lc.Campo, lc.Repeat
    nat = {**_SI, "personeria": ["natural"]}
    jur = {**_SI, "personeria": ["juridica"]}
    campos = [
        C("hoy", "Fecha de registro", "date", requerido=True),
        C("asistente", "Asistente predial", "select_one", "asistente", requerido=True,
          ayuda="La lista de unidades muestra solo las asignadas al asistente (ESP ve todas)."),
        C("unidad", "Código de la unidad (bloque / lote SUS)", "unidad", requerido=True),
        C("n_predio", "N.° de predio dentro de la unidad", "int", requerido=True, minimo=1, maximo=99,
          ayuda="Genera el código {UNIDAD}-P{nn}."),
        C("fecha_insp", "Fecha de inspección", "date", requerido=True),
        C("consentimiento", "Consentimiento para el uso de datos personales (Ley 29733)", "select_one", "si_no",
          requerido=True),
        # 1. predio matriz (titularidad registral)
        C("sit_registral", "Situación registral", "select_one", _L["sit_registral"]),
        C("doc_partida", "Partida electrónica / N.° de título", "text"),
        C("oficina_registral", "Oficina registral", "text"),
        C("reg_personeria", "Personería del titular registral", "select_one", _L["personeria"]),
        C("reg_titular", "Titular registral (nombre / razón social)", "text", relevante=_SI),
        C("docs_exhibidos", "Documentos exhibidos", "select_multiple", "docs_exhibidos"),
        # 2. prediante
        C("tipo_titularidad", "Condición jurídica del prediante", "select_one", "tipo_titularidad", requerido=True,
          relevante=_SI),
        C("sucesion_estado", "Estado de la sucesión", "select_one", "sucesion_estado",
          relevante={**_SI, "tipo_titularidad": ["sucesion"]}),
        C("personeria", "Tipo de personería", "select_one", _L["personeria"], requerido=True, relevante=_SI,
          ayuda="Posesionario, ocupante y sucesión: persona natural. Comunidad y Estado: persona jurídica."),
        C("tit_nombres", "Nombres", "text", requerido=True, relevante=nat),
        C("tit_apellidos", "Apellidos", "text", requerido=True, relevante=nat),
        C("tit_dni", "DNI (8 dígitos)", "text", relevante=nat),
        C("tit_celular", "Celular (9 dígitos)", "text", relevante=nat),
        C("tit_estado_civil", "Estado civil", "select_one", "estado_civil", relevante=nat),
        C("tit_correo", "Correo electrónico (opcional)", "text", relevante=nat),
        C("cony_nombre", "Nombre del cónyuge / conviviente", "text",
          relevante={**nat, "tit_estado_civil": ["casado", "conviviente"]}),
        C("cony_dni", "DNI del cónyuge", "text", relevante={**nat, "tit_estado_civil": ["casado", "conviviente"]}),
        C("cc_nombre", "Comunidad / entidad / razón social", "text", requerido=True, relevante=jur),
        C("cc_ruc", "RUC (11 dígitos)", "text", relevante=jur),
        C("cc_partida", "Partida registral (SUNARP)", "text", relevante=jur),
        C("cc_presidente", "Representante / presidente comunal", "text", relevante=jur),
        C("rep_cargo", "Cargo del representante", "select_one", _L["cargo_rep"], relevante=jur),
        C("cc_presidente_dni", "DNI del representante", "text", relevante=jur),
        C("cc_presidente_cel", "Celular del representante", "text", relevante=jur),
        C("dom_direccion", "Domicilio – dirección", "text", relevante=_SI),
        C("dom_departamento", "Domicilio – departamento", "select_one", _L["dep_domicilio"], relevante=_SI),
        C("dom_provincia", "Domicilio – provincia", "select_one", _L["prov_domicilio"],
          relevante={**_SI, "dom_departamento": ["piura"]}),
        C("dom_distrito", "Domicilio – distrito", "select_one", _L["distrito"],
          relevante={**_SI, "dom_provincia": list(PROVINCIAS_AMBITO)}),
        C("dom_lugar_otro", "Domicilio – provincia / distrito fuera del ámbito", "text", relevante=_SI),
        # 3. descripción general
        C("predio_nombre", "Nombre del predio", "text"),
        C("centro_poblado", "Centro poblado / sector", "text"),
        C("comunidad", "Comunidad campesina (si aplica)", "text"),
        C("acceso_tipo", "Tipo de acceso", "select_one", "acceso_tipo"),
        C("tiempo_acceso_min", "Tiempo desde el centro poblado (min)", "int", minimo=0, maximo=600),
        C("ref_acceso", "Referencia de acceso", "textarea"),
        C("punto_interior", "Punto interior del predio", "utm", requerido=True),
        C("predio_provincia", "Provincia del predio", "select_one", _L["provincia"], requerido=True),
        C("predio_distrito", "Distrito del predio", "select_one", _L["distrito"], requerido=True,
          ayuda="Debe pertenecer a la provincia elegida (15 distritos del ámbito)."),
        C("posicion_sus_campo", "Posición del lote SUS declarada en campo", "select_one", "posicion_sus",
          relevante={"tipo_unidad": ["lote_sus"]}),
        # 4. predio matriz
        C("uso_actual", "Uso actual del suelo", "select_one", "uso_actual"),
        C("uso_otro", "Especifique el otro uso", "text", relevante={"uso_actual": ["otro"]}),
        C("ocupacion", "Ocupación observada", "select_one", "ocupacion"),
        C("ocupacion_det", "Detalle de la ocupación", "text"),
        C("edificaciones", "Edificaciones / cultivos", "select_multiple", _L["edificaciones"]),
        C("edif_det", "Detalle / especifique «Otro»", "text"),
        C("pendiente_clase", "Topografía / pendiente", "select_one", _L["pendiente"],
          ayuda="Criterio de idoneidad del proyecto: pendiente máxima 75 %."),
        C("erosion", "Erosión observada", "select_multiple", _L["erosion"]),
        C("cobertura", "Estado de la cobertura vegetal", "select_one", _L["cobertura"]),
        C("mov_masa", "Evidencia de movimientos en masa", "select_one", "si_no"),
        C("tipo_predio", "Tipo de predio", "select_one", _L["tipo_predio"]),
        C("predio_area_decl_ha", "Área matriz del predio (ha)", "decimal", minimo=0),
        C("perimetro_m", "Perímetro (m)", "decimal", minimo=0),
        C("descripcion_predio", "Otros / descripción del predio", "textarea",
          ayuda="Cobertura, pendiente, accesos, linderos, signos de degradación o de movimientos en masa."),
        *[c for lado in ("norte", "este", "sur", "oeste") for c in (
            C(f"col_{lado}", f"Colindante {lado.upper()}", "text"),
            C(f"col_{lado}_tipo", f"Tipo de colindante {lado.upper()}", "select_one", _L["tipo_colindante"]))],
        C("conflicto_linderos_det", "Conflictos de linderos / observaciones", "textarea"),
        # 5. área a intervenir
        C("predio_area_unidad_ha", "Área del predio dentro de la unidad (ha)", "decimal", minimo=0),
        C("interferencias", "Interferencias", "select_multiple", "interferencias"),
        C("intervencion_prevista", "Intervención prevista (referencial)", "select_multiple", _L["intervencion"]),
        C("interf_detalle", "Detalle de interferencias", "textarea"),
        # 6. liberación
        C("aceptacion", "Aceptación del titular / prediante", "select_one", "aceptacion", requerido=True),
        C("modalidad", "Modalidad de disponibilidad", "select_one", "modalidad"),
        C("conclusion_campo", "Conclusión de campo", "select_one", "conclusion_campo", requerido=True),
        C("estado_la_propuesto", "Estado LA propuesto", "select_one", "estado_la"),
        C("tipo_acta_prevista", "Acta prevista", "select_one", "tipo_acta"),
        C("cod_titular_prev", "Código de titular ya registrado (T0000), si existe", "text"),
        C("alertas", "Observaciones / alertas sociales", "textarea"),
        # 8. fotos
        C("foto_1", "Foto 1 – Vista general del predio", "foto"),
        C("foto_2", "Foto 2 – Cobertura vegetal / uso actual", "foto"),
        C("foto_3", "Foto 3 – Linderos, hitos o accesos", "foto"),
        C("foto_4", "Foto 4 – Interferencias o detalle relevante", "foto"),
        C("croquis", "Croquis del predio (foto del dibujo)", "foto"),
    ]
    repeats = [
        R("r_afectaciones", "Afectaciones (obras y bienes existentes en el área)",
          [C("af_descripcion", "Descripción / uso", "text", requerido=True), C("af_metrado", "Metrado", "decimal"),
           C("af_und", "Und.", "select_one", _L["und_medida"]), C("af_material", "Material", "select_one", _L["material"]),
           C("af_conservacion", "Est. conservación", "select_one", _L["conservacion"]),
           C("af_construccion", "Est. construcción", "select_one", _L["construccion"]),
           C("af_ubicacion", "Ubicación", "text")], "FP_afectaciones"),
        R("r_puntos", "Puntos de control / vértices del predio",
          [C("pt_gps", "Punto", "utm", requerido=True), C("pt_tipo", "Tipo de punto", "select_one", "pt_tipo")],
          "FP_puntos"),
        R("poligono_sus", "Vértices del lote SUS (en orden, UTM 17S)",
          [C("orden", "Orden", "int", requerido=True), C("pt_gps", "Vértice", "utm", requerido=True)],
          "FP_vertices_SUS", relevante={"tipo_unidad": ["lote_sus"]}),
    ]
    return lc.Formulario(FORM_ID, CODIGO, "Ficha técnica predial completa (registra F-LA-02 + F-LA-03)", campos,
                         repeats, con_predio=True)


PROVINCIAS_AMBITO = tuple(fpx.PROVINCIAS)               # ayabaca, huancabamba, morropon
FORMULARIO = None


def _registrar_formulario():
    global FORMULARIO
    FORMULARIO = _formulario_app()
    lc.FORMULARIOS_EXTRA[FORM_ID] = FORMULARIO


_registrar_formulario()
