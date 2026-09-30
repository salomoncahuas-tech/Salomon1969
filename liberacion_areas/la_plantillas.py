# -*- coding: utf-8 -*-
"""
Plantillas de Liberación de Áreas – Proyecto IN Piura (CUI 2669244).

1. Plantilla Excel ANIN (una hoja por formulario F-LA-01 … F-LA-06, hojas de repeats, catálogo de unidades y
   listas desplegables) para llenar en gabinete o campo y luego IMPORTAR en el aplicativo.
2. Lectura de esa plantilla → envíos iguales a los de Kobo (la_campos.construir_registro).
3. XLSForm para KoboToolbox con los MISMOS nombres de campo (paquete ZIP con unidades.csv).
Todo se deriva de la_campos.FORMULARIOS: un cambio de campo se refleja en las tres vías.
"""
from __future__ import annotations

import io
import math
import zipfile
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import la_campos as lc
from . import la_core as core

VERDE, AZUL, ALT, GRIS = "1B4D2E", "1B4F72", "EAF2F8", "808B96"
_thin = Side(style="thin", color="A6ACAF")
_B = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
ENC = ["AUTORIDAD NACIONAL DE INFRAESTRUCTURA - ANIN",
       "DIRECCIÓN DE INTERVENCIONES MULTISECTORIALES Y DE EMERGENCIA - DIME",
       "SUBDIRECCIÓN DE ESTUDIOS DE INVERSIÓN"]
VERSION_PLANTILLA = "LA-2026.10"
FILA_ETIQUETA, FILA_CLAVE, FILA_DATOS = 6, 7, 8
FILAS_PREPARADAS = 300
CLAVE_ID = "id_registro"
MARCA_SI = "SI"


# ------------------------------------------------------------------ columnas de cada hoja
def _columnas_campos(campos: list[lc.Campo]) -> list[dict]:
    """Expande los campos en columnas de Excel: UTM → Este/Norte/Altitud/Precisión; selección múltiple →
    una columna SI/— por opción."""
    cols = []
    for c in campos:
        cond = ""
        if c.relevante:
            cond = "Solo si " + " y ".join(f"{k} = {' / '.join(v)}" for k, v in c.relevante.items())
        nota = " · ".join(x for x in (c.ayuda, cond) if x)
        if c.tipo == "utm":
            for suf, et, rng in (("este", "Este (m)", (core.E_MIN, core.E_MAX)), ("norte", "Norte (m)", (core.N_MIN, core.N_MAX)),
                                 ("alt", "Altitud (m)", None), ("prec", "Precisión GPS (m)", (0, 1000))):
                cols.append({"clave": f"{c.nombre}_{suf}", "etiqueta": f"{c.etiqueta} – {et}", "tipo": "decimal",
                             "rango": rng, "req": c.requerido and suf in ("este", "norte"),
                             "nota": ("UTM WGS84 Zona 17S. " + nota) if suf in ("este", "norte") else nota})
        elif c.tipo == "select_multiple":
            for cod, et in lc.OPCIONES[c.lista]:
                cols.append({"clave": f"{c.nombre}__{cod}", "etiqueta": f"{c.etiqueta}: {et}", "tipo": "marca",
                             "req": False, "nota": f"Escriba {MARCA_SI} si aplica. {nota}".strip()})
        else:
            rng = (c.minimo, c.maximo) if (c.minimo is not None or c.maximo is not None) else None
            cols.append({"clave": c.nombre, "etiqueta": c.etiqueta, "tipo": c.tipo, "lista": c.lista, "rango": rng,
                         "req": c.requerido, "nota": nota})
    return cols


def _hojas_plantilla() -> list[dict]:
    hojas = []
    for fid, f in lc.FORMULARIOS.items():
        hojas.append({"hoja": f.codigo, "form_id": fid, "titulo": f"{f.codigo} – {f.nombre.upper()}",
                      "cols": [{"clave": CLAVE_ID, "etiqueta": "N.° de registro", "tipo": "int", "req": True,
                                "nota": "Correlativo único en esta hoja (1, 2, 3…). Enlaza con las hojas de detalle."}]
                      + _columnas_campos(f.campos), "repeat": None})
        for rp in f.repeats:
            hojas.append({"hoja": rp.hoja, "form_id": fid, "titulo": f"{f.codigo} – {rp.etiqueta.upper()}",
                          "cols": [{"clave": CLAVE_ID, "etiqueta": f"N.° de registro en la hoja {f.codigo}", "tipo": "int",
                                    "req": True, "nota": f"El mismo N.° de registro de la hoja {f.codigo}."}]
                          + _columnas_campos(rp.campos), "repeat": rp.nombre})
    return hojas


# ------------------------------------------------------------------ generación de la plantilla Excel
def _encabezado(ws, titulo: str, ncols: int, subtitulo: str = ""):
    for i, t in enumerate(ENC, 1):
        c = ws.cell(i, 1, t)
        c.font = Font(name="Arial", bold=(i == 1), size=11 if i == 1 else 9, color=VERDE)
    c = ws.cell(4, 1, titulo)
    c.font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=VERDE)
    for j in range(2, max(ncols, 8) + 1):
        ws.cell(4, j).fill = PatternFill("solid", fgColor=VERDE)
    ws.cell(5, 1, subtitulo or f"Proyecto IN Piura – CUI 2669244 · Plantilla {VERSION_PLANTILLA} · "
                               f"Coordenadas UTM WGS84 Zona 17S · Generada {date.today():%d/%m/%Y}"
            ).font = Font(name="Arial", italic=True, size=8, color=GRIS)


def _listas(wb, catalogo: dict[str, dict], asistente: str | None):
    """Hoja oculta _Listas: una columna por lista desplegable (etiquetas) + códigos de unidades."""
    ws = wb.create_sheet("_Listas")
    rangos = {}
    listas = dict(lc.OPCIONES)
    unidades = sorted(c for c, r in catalogo.items() if not asistente or asistente == "ESP" or r.get("asistente") == asistente)
    for j, (nombre, ops) in enumerate(list(listas.items()) + [("unidad", [(u, u) for u in unidades]),
                                                              ("marca", [(MARCA_SI, MARCA_SI)])], 1):
        ws.cell(1, j, nombre).font = Font(name="Arial", bold=True)
        for i, (_, et) in enumerate(ops, 2):
            ws.cell(i, j, et)
        L = get_column_letter(j)
        rangos[nombre] = f"'_Listas'!${L}$2:${L}${max(2, len(ops) + 1)}"
    ws.sheet_state = "hidden"
    return rangos


def _validacion(ws, col: dict, j: int, rangos: dict):
    L = get_column_letter(j)
    ref = f"{L}{FILA_DATOS}:{L}{FILA_DATOS + FILAS_PREPARADAS - 1}"
    dv = None
    t = col["tipo"]
    if t in ("select_one", "unidad", "marca"):
        lista = "unidad" if t == "unidad" else ("marca" if t == "marca" else col.get("lista"))
        dv = DataValidation(type="list", formula1=rangos[lista], allow_blank=True, showErrorMessage=True,
                            errorTitle="Valor no válido", error="Elija un valor de la lista desplegable.")
    elif t in ("int", "decimal") and col.get("rango"):
        mn, mx = col["rango"]
        dv = DataValidation(type="whole" if t == "int" else "decimal", operator="between",
                            formula1=str(mn if mn is not None else -1e12), formula2=str(mx if mx is not None else 1e12),
                            allow_blank=True, showErrorMessage=True, errorTitle="Fuera de rango",
                            error=f"Valor entre {mn} y {mx}.")
    elif t == "date":
        dv = DataValidation(type="date", operator="between", formula1="DATE(2024,1,1)", formula2="DATE(2035,12,31)",
                            allow_blank=True, showErrorMessage=True, errorTitle="Fecha", error="Ingrese una fecha válida (dd/mm/aaaa).")
    elif col["clave"].endswith("_dni") or col["clave"] == "tit_dni":
        dv = DataValidation(type="textLength", operator="equal", formula1="8", allow_blank=True, showErrorMessage=True,
                            errorTitle="DNI", error="El DNI tiene 8 dígitos.")
    if dv is not None:
        ws.add_data_validation(dv)
        dv.add(ref)


_TEXTO = ("unidad", "unidades", "text", "textarea")


def _hoja_formulario(wb, h: dict, rangos: dict):
    ws = wb.create_sheet(h["hoja"][:31])
    cols = h["cols"]
    _encabezado(ws, h["titulo"], len(cols),
                "Complete desde la fila 8. Celdas con (*) son obligatorias. No modifique las filas 6 y 7. "
                "Coordenadas en UTM WGS84 Zona 17S.")
    for j, col in enumerate(cols, 1):
        e = ws.cell(FILA_ETIQUETA, j, col["etiqueta"] + (" (*)" if col.get("req") else ""))
        e.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        e.fill = PatternFill("solid", fgColor=AZUL if not col.get("req") else VERDE)
        e.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        e.border = _B
        if col.get("nota"):
            e.comment = Comment(col["nota"], "ANIN")
        k = ws.cell(FILA_CLAVE, j, col["clave"])
        k.font = Font(name="Arial", size=7, color=GRIS, italic=True)
        k.border = _B
        ancho = 14 if col["tipo"] in ("int", "decimal", "date", "marca") else 22
        if col["tipo"] == "textarea":
            ancho = 40
        ws.column_dimensions[get_column_letter(j)].width = ancho
        _validacion(ws, col, j, rangos)
        fmt = "@" if col["tipo"] in _TEXTO or col["clave"].endswith(("_dni", "_ruc", "_celular", "_cel")) else (
            "dd/mm/yyyy" if col["tipo"] == "date" else None)
        for i in range(FILA_DATOS, FILA_DATOS + FILAS_PREPARADAS):
            c = ws.cell(i, j)
            c.border = _B
            c.font = Font(name="Arial", size=9)
            if fmt:
                c.number_format = fmt
            if i % 2 == 1:
                c.fill = PatternFill("solid", fgColor=ALT)
    ws.row_dimensions[FILA_ETIQUETA].height = 48
    ws.freeze_panes = ws.cell(FILA_DATOS, 2)
    return ws


def _hoja_catalogo(wb, catalogo: dict[str, dict], asistente: str | None):
    ws = wb.create_sheet("Catalogo_unidades")
    cols = ["codigo", "tipo_unidad", "provincia", "distrito", "bloque_ref", "area_ha", "area_bloque_ha", "posicion_sus",
            "asistente", "asistente_nombre"]
    _encabezado(ws, "CATÁLOGO DE UNIDADES V6 (bloques, lotes SUS)" + (f" – ASISTENTE {asistente}" if asistente else ""),
                len(cols), "Referencia de solo lectura: use el código en la columna «Unidad» de cada formulario.")
    for j, c in enumerate(cols, 1):
        e = ws.cell(FILA_ETIQUETA, j, c)
        e.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        e.fill = PatternFill("solid", fgColor=AZUL)
        e.border = _B
        ws.column_dimensions[get_column_letter(j)].width = 16
    filas = sorted(catalogo.values(), key=lambda r: (str(r.get("provincia")), str(r.get("distrito")), str(r.get("codigo") or r.get("name"))))
    i = FILA_DATOS - 1
    for r in filas:
        if asistente and asistente != "ESP" and r.get("asistente") != asistente:
            continue
        for j, c in enumerate(cols, 1):
            v = r.get(c) if c != "codigo" else (r.get("codigo") or r.get("name"))
            if isinstance(v, float) and math.isnan(v):
                v = None
            cell = ws.cell(i, j, float(v) if c.startswith("area") and core._num(v) is not None else v)
            cell.font = Font(name="Arial", size=9)
            cell.border = _B
        i += 1
    ws.freeze_panes = ws.cell(FILA_DATOS, 2)
    ws.auto_filter.ref = f"A{FILA_ETIQUETA}:{get_column_letter(len(cols))}{max(i - 1, FILA_ETIQUETA)}"


def _hoja_instrucciones(wb, asistente: str | None):
    ws = wb.active
    ws.title = "Instrucciones"
    _encabezado(ws, "PLANTILLA DE LIBERACIÓN DE ÁREAS – REGISTRO DE CAMPO", 8)
    lineas = [
        ("Tres vías de registro equivalentes", True),
        ("1. KoboToolbox: formularios F-LA desplegados en la cuenta del proyecto (el aplicativo los descarga por API).", False),
        ("2. Digitación directa en el aplicativo: página «Liberación de Áreas» › pestaña «Registro en campo».", False),
        ("3. Esta plantilla Excel: llenar y luego importar en «Liberación de Áreas» › «Plantillas Excel».", False),
        ("Las tres vías usan los mismos campos y las mismas validaciones (catálogo V6, asistente, UTM 17S, lotes SUS, actas).", False),
        ("", False),
        ("Cómo llenar", True),
        ("• Una fila por registro desde la fila 8. La columna «N.° de registro» es un correlativo único en cada hoja.", False),
        ("• Los detalles (puntos de control, vértices del lote SUS, firmantes) van en su hoja, con el mismo N.° de registro.", False),
        ("• Use las listas desplegables. Puede escribir el código o la etiqueta; el aplicativo los reconoce.", False),
        ("• Coordenadas en UTM WGS84 Zona 17S: Este entre 450,000 y 750,000 m; Norte entre 9,300,000 y 9,600,000 m.", False),
        ("• El código de predio se forma solo: {UNIDAD}-P{nn} (p. ej. 27-P01). En actas A-03 comunales use N.° de predio 0.", False),
        ("• Selección múltiple (documentos, interferencias, checklist): escriba SI en cada opción que aplique.", False),
        ("• No cambie los nombres de las hojas ni las filas 6 y 7 (etiquetas y claves).", False),
        ("", False),
        ("Al importar", True),
        ("• El aplicativo muestra cada fila como NUEVO, DUPLICADO u OBSERVADO antes de confirmar.", False),
        ("• Volver a importar el mismo archivo no duplica registros (las filas sin cambios se detectan como DUPLICADO).", False),
        ("• Una fila corregida en Excel se registra como un envío nuevo: para corregir un registro ya importado use "
         "«Historial / Edición» en el aplicativo.", False),
        ("• Datos personales (DNI, celular) son reservados – Ley 29733: no circule este archivo fuera del equipo.", False),
    ]
    if asistente:
        lineas.insert(0, (f"Plantilla preparada para el asistente {asistente}: el desplegable de unidades muestra solo sus unidades.", True))
    for i, (t, neg) in enumerate(lineas, 7):
        c = ws.cell(i, 1, t)
        c.font = Font(name="Arial", size=10, bold=neg, color=VERDE if neg else "000000")
    ws.column_dimensions["A"].width = 130
    fila = 7 + len(lineas) + 1
    ws.cell(fila, 1, "Hojas del libro").font = Font(name="Arial", bold=True, color=VERDE)
    for k, h in enumerate(_hojas_plantilla(), fila + 1):
        ws.cell(k, 1, f"• {h['hoja']}: {h['titulo']}").font = Font(name="Arial", size=9)


def generar_plantilla_excel(catalogo: dict[str, dict], asistente: str | None = None) -> bytes:
    """Libro .xlsx con formato ANIN listo para llenar. `catalogo` = {codigo: fila de la_unidades / unidades.csv}."""
    wb = Workbook()
    _hoja_instrucciones(wb, asistente)
    rangos = _listas(wb, catalogo, asistente)
    for h in _hojas_plantilla():
        _hoja_formulario(wb, h, rangos)
    _hoja_catalogo(wb, catalogo, asistente)
    wb.move_sheet("_Listas", offset=len(wb.sheetnames))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ lectura de la plantilla llenada
def _valor_celda(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _leer_hoja(ws) -> list[dict]:
    """Filas con datos (dict clave → valor) de una hoja de plantilla. Localiza la fila de claves por 'id_registro'."""
    fila_clave = None
    for i in range(1, 20):
        if str(ws.cell(i, 1).value or "").strip() == CLAVE_ID:
            fila_clave = i
            break
    if fila_clave is None:
        return []
    claves = [str(ws.cell(fila_clave, j).value or "").strip() for j in range(1, ws.max_column + 1)]
    filas = []
    for n, row in enumerate(ws.iter_rows(min_row=fila_clave + 1, values_only=True), fila_clave + 1):
        d = {k: _valor_celda(v) for k, v in zip(claves, row) if k}
        if any(v not in (None, "") for k, v in d.items() if k != CLAVE_ID):
            d["_fila"] = n
            filas.append(d)
    return filas


def _texto_dni(v):
    if isinstance(v, int):
        return str(v).zfill(8)
    return v


def _valores_desde_fila(campos: list[lc.Campo], fila: dict) -> dict:
    val = {}
    for c in campos:
        if c.tipo == "utm":
            for suf in lc.UTM_SUFIJOS:
                if f"{c.nombre}_{suf}" in fila:
                    val[f"{c.nombre}_{suf}"] = fila[f"{c.nombre}_{suf}"]
        elif c.tipo == "select_multiple":
            marcadas = [cod for cod, _ in lc.OPCIONES[c.lista]
                        if str(fila.get(f"{c.nombre}__{cod}") or "").strip().upper() in ("SI", "SÍ", "X", "1", "TRUE")]
            if any(f"{c.nombre}__" in k for k in fila):
                val[c.nombre] = marcadas
            elif c.nombre in fila:
                val[c.nombre] = fila[c.nombre]
        elif c.nombre in fila:
            v = fila[c.nombre]
            if c.nombre.endswith("_dni"):
                v = _texto_dni(v)
            if c.tipo in ("unidad", "unidades", "text") and isinstance(v, (int, float)):
                v = str(v)
            val[c.nombre] = v
    return val


def leer_plantilla_excel(archivo, catalogo: dict[str, dict]) -> dict:
    """Plantilla llenada → {"envios": {form_id: [dict(registro, avisos, fila, hoja)]},
                             "documentos": [dict(datos, avisos, fila)], "hojas_leidas": [...], "errores": [...]}.
    Cada registro es un envío aplanado como el de Kobo, con kobo_uuid determinista (xls-…)."""
    wb = load_workbook(archivo, data_only=True)
    out = {"envios": {}, "documentos": [], "hojas_leidas": [], "errores": []}
    for fid, f in lc.FORMULARIOS.items():
        if f.codigo not in wb.sheetnames:
            continue
        principales = _leer_hoja(wb[f.codigo])
        out["hojas_leidas"].append(f"{f.codigo} ({len(principales)})")
        detalle = {}
        for rp in f.repeats:
            filas = _leer_hoja(wb[rp.hoja]) if rp.hoja in wb.sheetnames else []
            for x in filas:
                detalle.setdefault((rp.nombre, x.get(CLAVE_ID)), []).append(x)
            ids_padre = {p.get(CLAVE_ID) for p in principales}
            huerfanas = sorted({str(x.get(CLAVE_ID)) for x in filas if x.get(CLAVE_ID) not in ids_padre})
            if huerfanas:
                out["errores"].append(f"Hoja {rp.hoja}: N.° de registro sin fila en {f.codigo}: {', '.join(huerfanas)}")
        vistos = set()
        for p in principales:
            rid = p.get(CLAVE_ID)
            if rid in (None, ""):
                out["errores"].append(f"Hoja {f.codigo}, fila {p['_fila']}: falta N.° de registro")
                continue
            if rid in vistos:
                out["errores"].append(f"Hoja {f.codigo}: N.° de registro {rid} repetido (fila {p['_fila']})")
                continue
            vistos.add(rid)
            valores = _valores_desde_fila(f.campos, p)
            reps = {}
            for rp in f.repeats:
                reps[rp.nombre] = [dict(_valores_desde_fila(rp.campos, x), **({"orden": x.get("orden")} if "orden" in x else {}))
                                   for x in detalle.get((rp.nombre, rid), [])]
            if fid == "f_la_05_documentos":
                reg, avisos = lc.construir_registro(fid, valores, catalogo)
                out["documentos"].append({"datos": _documento_desde_registro(reg), "avisos": avisos, "fila": p["_fila"]})
                continue
            uid = lc.uuid_contenido(fid, {"v": valores, "r": reps})
            reg, avisos = lc.construir_registro(fid, valores, catalogo, reps, kobo_uuid=uid)
            reg["_origen"] = "PLANTILLA"
            out["envios"].setdefault(fid, []).append({"registro": reg, "avisos": avisos, "fila": p["_fila"], "hoja": f.codigo})
    if not out["hojas_leidas"]:
        out["errores"].append("El archivo no contiene hojas F-LA-0x de la plantilla de Liberación de Áreas.")
    return out


def _documento_desde_registro(reg: dict) -> dict:
    cp = lc.cod_predio(reg.get("unidad"), reg.get("n_predio")) if core._num(reg.get("n_predio")) else None
    return {"cod_unidad": reg.get("unidad"), "cod_predio": cp, "tipo": reg.get("tipo"), "entidad": reg.get("entidad"),
            "fecha": reg.get("fecha"), "resultado": reg.get("resultado"), "n_partida": reg.get("n_partida"),
            "descripcion": reg.get("descripcion"), "archivo_url": reg.get("archivo_url"),
            "registrado_por": reg.get("registrado_por")}


# ------------------------------------------------------------------ XLSForm para KoboToolbox
_CSV_UNIDADES = "unidades.csv"
_CAMPOS_CATALOGO = ("tipo_unidad", "provincia", "distrito", "bloque_ref", "area_bloque_ha")


def _relevante_xls(rel: dict | None) -> str:
    if not rel:
        return ""
    return " and ".join("(" + " or ".join(f"${{{k}}}='{v}'" for v in vals) + ")" for k, vals in rel.items())


def _fila_xls(c: lc.Campo, relevante_extra: str = "") -> dict:
    tipo = {"text": "text", "textarea": "text", "int": "integer", "decimal": "decimal", "date": "date",
            "utm": "geopoint"}.get(c.tipo)
    fila = {"type": tipo, "name": c.nombre, "label": c.etiqueta, "hint": c.ayuda,
            "required": "yes" if c.requerido else "", "relevant": _relevante_xls(c.relevante) or relevante_extra}
    if c.tipo == "textarea":
        fila["appearance"] = "multiline"
    if c.tipo == "select_one":
        fila["type"] = f"select_one {c.lista}"
    elif c.tipo == "select_multiple":
        fila["type"] = f"select_multiple {c.lista}"
    elif c.tipo == "unidad":
        fila["type"] = f"select_one_from_file {_CSV_UNIDADES}"
        fila["choice_filter"] = "asistente=${asistente} or ${asistente}='ESP'"
        fila["appearance"] = "autocomplete"
    elif c.tipo == "unidades":
        fila["type"] = f"select_multiple_from_file {_CSV_UNIDADES}"
        fila["choice_filter"] = "asistente=${asistente} or ${asistente}='ESP'"
        fila["appearance"] = "autocomplete"
    elif c.tipo == "utm":
        fila["hint"] = (c.ayuda + " " if c.ayuda else "") + "Precisión ≤ 10 m."
    if c.nombre.endswith("_dni"):
        fila["constraint"] = "regex(., '^[0-9]{8}$')"
        fila["constraint_message"] = "El DNI tiene 8 dígitos"
    elif c.tipo in ("int", "decimal") and (c.minimo is not None or c.maximo is not None):
        cons = []
        if c.minimo is not None:
            cons.append(f". >= {c.minimo:g}")
        if c.maximo is not None:
            cons.append(f". <= {c.maximo:g}")
        fila["constraint"] = " and ".join(cons)
        fila["constraint_message"] = "Valor fuera de rango"
    if c.tipo == "date" and c.nombre == "hoy":
        fila["default"] = "today()"
    return fila


def xlsform_hojas(form_id: str) -> dict[str, pd.DataFrame]:
    """XLSForm (survey, choices, settings) de un formulario F-LA, con los nombres que espera el importador."""
    f = lc.FORMULARIOS[form_id]
    survey = [{"type": "start", "name": "inicio"}, {"type": "end", "name": "fin"}]
    listas_usadas = set()
    for c in f.campos:
        fila = _fila_xls(c)
        survey.append(fila)
        if c.lista:
            listas_usadas.add(c.lista)
        if c.tipo == "unidad":
            survey.append({"type": "calculate", "name": "cod_unidad", "calculation": "${unidad}"})
            for k in _CAMPOS_CATALOGO:
                survey.append({"type": "calculate", "name": k,
                               "calculation": f"instance('unidades')/root/item[name=${{unidad}}]/{k}"})
        if c.tipo == "unidades":
            survey.append({"type": "calculate", "name": "cod_unidad", "calculation": "${unidades}"})
            for k in ("provincia", "distrito"):
                survey.append({"type": "calculate", "name": k,
                               "calculation": f"instance('unidades')/root/item[name=selected-at(${{unidades}}, 0)]/{k}"})
        if c.nombre == "n_predio" and f.con_predio:
            survey.append({"type": "calculate", "name": "cod_predio",
                           "calculation": "if(${n_predio} > 0, concat(${unidad}, '-P', "
                                          "if(${n_predio} < 10, concat('0', ${n_predio}), ${n_predio})), '')"})
    if form_id == "f_la_06_vivero":
        survey.append({"type": "calculate", "name": "cod_unidad", "calculation": "${cod_vivero}"})
    for rp in f.repeats:
        rel = _relevante_xls(rp.relevante)
        if rp.nombre == "poligono_sus":
            survey.append({"type": "geoshape", "name": "poligono_sus", "label": rp.etiqueta.replace(" (en orden, UTM 17S)", ""),
                           "relevant": rel, "required": "yes" if rel else "",
                           "hint": "Recorra el perímetro del lote SUS; mínimo 1 ha, ≤ 10 % del bloque, dentro o a ≤ 50 m."})
            continue
        survey.append({"type": "begin_repeat", "name": rp.nombre, "label": rp.etiqueta, "relevant": rel})
        for c in rp.campos:
            survey.append(_fila_xls(c))
            if c.lista:
                listas_usadas.add(c.lista)
        survey.append({"type": "end_repeat", "name": rp.nombre})
    choices = [{"list_name": ln, "name": cod, "label": et} for ln in sorted(listas_usadas) for cod, et in lc.OPCIONES[ln]]
    cols_s = ["type", "name", "label", "hint", "required", "relevant", "constraint", "constraint_message", "calculation",
              "choice_filter", "appearance", "default"]
    settings = pd.DataFrame([{"form_title": f"{f.codigo} – {f.nombre}", "form_id": form_id,
                              "version": datetime.now().strftime("%Y%m%d%H%M"), "default_language": "Español (es)"}])
    return {"survey": pd.DataFrame(survey).reindex(columns=cols_s).fillna(""),
            "choices": pd.DataFrame(choices, columns=["list_name", "name", "label"]),
            "settings": settings}


def generar_xlsform(form_id: str) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for nombre, d in xlsform_hojas(form_id).items():
            d.to_excel(w, sheet_name=nombre, index=False)
    return buf.getvalue()


def unidades_csv(catalogo: dict[str, dict]) -> bytes:
    """unidades.csv para adjuntar como archivo multimedia en cada formulario Kobo (select_one_from_file)."""
    cols = ["name", "label", "provincia", "distrito", "tipo_unidad", "bloque_ref", "area_ha", "area_bloque_ha",
            "posicion_sus", "asistente", "asistente_nombre"]
    filas = []
    for cod, r in sorted(catalogo.items()):
        filas.append({c: (cod if c == "name" else r.get(c)) for c in cols})
    return pd.DataFrame(filas, columns=cols).to_csv(index=False).encode("utf-8")


def paquete_kobo(catalogo: dict[str, dict]) -> bytes:
    """ZIP con los XLSForm F-LA-01/02/03/04/06, unidades.csv y un LEEME."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for fid, f in lc.FORMULARIOS.items():
            if f.kobo:
                z.writestr(f"{f.codigo}_{fid}.xlsx", generar_xlsform(fid))
        z.writestr(_CSV_UNIDADES, unidades_csv(catalogo))
        z.writestr("LEEME.txt", (
            "Formularios KoboToolbox – Liberación de Áreas · Proyecto IN Piura (CUI 2669244)\n\n"
            "1. En KoboToolbox: Nuevo › Cargar XLSForm › elija cada archivo F-LA-0x_*.xlsx.\n"
            "2. En cada formulario: Configuración › Multimedia › cargue unidades.csv (el mismo para todos).\n"
            "3. Despliegue el formulario. El nombre debe contener 'F-LA' para que el aplicativo lo liste.\n"
            "4. En el aplicativo: Liberación de Áreas › Importar KoboToolbox › API › Listar formularios F-LA.\n\n"
            "Los nombres de campo son los mismos de la plantilla Excel y del registro en el aplicativo;\n"
            "las tres vías pasan por las mismas validaciones. F-LA-05 (búsqueda documental) se registra en el\n"
            "aplicativo o en la plantilla Excel.\n"
            "Si ya tiene formularios F-LA desplegados, compare los nombres de campo antes de reemplazarlos.\n").encode("utf-8"))
    return buf.getvalue()
