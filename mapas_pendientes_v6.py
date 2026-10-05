"""Mapas de pendientes de los 117 bloques V6 del Proyecto IN Piura (CUI 2669244).

Genera, a partir del DEM Copernicus GLO-30 (30 m) reproyectado a UTM 17S (EPSG:32717):
  - un mapa de pendientes por bloque (PNG, A4 horizontal, formato ANIN),
  - un mapa general con los 117 bloques,
  - un atlas PDF con el mapa general y los 117 mapas,
  - una matriz Excel con estadísticas de pendiente por bloque (formato ANIN, totales con fórmula),
  - el ráster de pendiente (%) y su clasificación en GeoTIFF.

Uso:
    python mapas_pendientes_v6.py --dem-dir CARPETA_DEM [--bloques poligonos.gpkg --campo-codigo codigo]
                                  [--salida salidas/mapas_pendientes_v6] [--descargar-dem]

Sin --bloques el repositorio solo tiene el centroide UTM de cada bloque (datos/atributos_bloques_v6.csv),
así que las estadísticas se calculan sobre un CÍRCULO DE ÁREA EQUIVALENTE centrado en el centroide.
Es una aproximación referencial: no es el límite del bloque. Cuando se cuente con la capa de polígonos
(GPKG o SHP con .prj), pásela con --bloques y el script usará el límite real.

Clases de pendiente: rangos del Reglamento de Clasificación de Tierras por su Capacidad de Uso Mayor
(A 0-2 % ... H > 75 %). El umbral de 75 % coincide con el criterio de idoneidad del Paso 4.
"""
import argparse
import csv
import io
import math
import sys
import textwrap
import zipfile
from datetime import date
from pathlib import Path

import numpy as np

RAIZ = Path(__file__).resolve().parent
sys.path.insert(0, str(RAIZ / ".claude" / "skills" / "anin-in-piura" / "scripts"))
from anin_catalogo import AREA_TOTAL_HA, cargar_bloques  # noqa: E402
from anin_utm import validar_utm  # noqa: E402

CRS_PROYECTO = "EPSG:32717"
RES = 30.0  # m
VERDE = "#1B4D2E"

# (código, mínimo %, máximo %, descripción, color)
CLASES = [
    ("A", 0, 2, "Plana o casi a nivel", "#1A9850"),
    ("B", 2, 4, "Ligeramente inclinada", "#66BD63"),
    ("C", 4, 8, "Moderadamente inclinada", "#A6D96A"),
    ("D", 8, 15, "Fuertemente inclinada", "#D9EF8B"),
    ("E", 15, 25, "Moderadamente empinada", "#FEE08B"),
    ("F", 25, 50, "Empinada", "#FDAE61"),
    ("G", 50, 75, "Muy empinada", "#F46D43"),
    ("H", 75, None, "Extremadamente empinada", "#A50026"),
]
LIMITES = [c[1] for c in CLASES[1:]]  # 2, 4, 8, 15, 25, 50, 75
UMBRAL_IDONEIDAD = 75.0

ENCABEZADOS = [
    "AUTORIDAD NACIONAL DE INFRAESTRUCTURA - ANIN",
    "DIRECCIÓN DE INTERVENCIONES MULTISECTORIALES Y DE EMERGENCIA - DIME",
    "SUBDIRECCIÓN DE ESTUDIOS DE INVERSIÓN - SESDI",
]
SUBTITULO = ("PROYECTO IN PIURA | CUI 2669244 | Cuenca Alta del Río Piura | "
             "UTM WGS 84 Zona 17S (EPSG:32717)")
FUENTE_DEM = ("DEM: Copernicus GLO-30 (ESA, 30 m, modelo de superficie). "
              "Bloques: catálogo V6 (datos/unidades_liberacion_areas.csv).")
URL_DEM = "https://copernicus-dem-30m.s3.amazonaws.com/{n}/{n}.tif"


# ---------------------------------------------------------------- datos de bloques
def cargar_centroides(ruta=None):
    """Centroides UTM por código. Lee datos/atributos_bloques_v6.csv o, si no está extraído,
    la copia dentro de Liberacion_Areas_vinculo_datos.zip."""
    ruta = Path(ruta) if ruta else RAIZ / "datos" / "atributos_bloques_v6.csv"
    if ruta.exists():
        texto = ruta.read_text(encoding="utf-8")
    else:
        with zipfile.ZipFile(RAIZ / "Liberacion_Areas_vinculo_datos.zip") as z:
            texto = z.read("datos/atributos_bloques_v6.csv").decode("utf-8")
    return {r["codigo"]: r for r in csv.DictReader(io.StringIO(texto))}


def preparar_bloques(ruta_centroides=None, ruta_poligonos=None, campo_codigo="codigo"):
    """Lista de bloques V6 con provincia, distrito, área del catálogo, centroide y geometría."""
    import geopandas as gpd
    from shapely.geometry import Point

    catalogo = cargar_bloques()
    cent = cargar_centroides(ruta_centroides)
    poligonos = {}
    if ruta_poligonos:
        g = gpd.read_file(ruta_poligonos)
        if g.crs is None:
            raise SystemExit(f"{ruta_poligonos} no tiene CRS (.prj); no se asume ninguno.")
        g = g.to_crs(CRS_PROYECTO)
        g[campo_codigo] = g[campo_codigo].astype(str).str.strip()
        poligonos = {k: v.geometry.union_all() for k, v in g.groupby(campo_codigo)}

    avisos, bloques = [], []
    for b in catalogo:
        c = cent.get(b["codigo"])
        if c is None:
            avisos.append(f"Bloque {b['codigo']}: sin centroide en atributos_bloques_v6.csv")
            continue
        este, norte = float(c["utm_este"]), float(c["utm_norte"])
        ok, msg = validar_utm(este, norte)
        if not ok:
            avisos.append(f"Bloque {b['codigo']}: {msg}")
            continue
        if b["codigo"] in poligonos:
            geom, metodo = poligonos[b["codigo"]], "Polígono del bloque"
            area_geom = geom.area / 1e4
            if abs(area_geom - b["area_ha"]) > 0.02 * b["area_ha"]:
                avisos.append(f"Bloque {b['codigo']}: área del polígono {area_geom:,.3f} ha difiere "
                              f"del catálogo {b['area_ha']:,.3f} ha")
        else:
            radio = math.sqrt(b["area_ha"] * 1e4 / math.pi)
            geom, metodo = Point(este, norte).buffer(radio, 128), "Círculo de área equivalente"
            if poligonos:
                avisos.append(f"Bloque {b['codigo']}: no figura en la capa de polígonos; se usa círculo")
        bloques.append({**b, "microcuenca": c.get("microcuenca", ""), "zona": c.get("zona", ""),
                        "este": este, "norte": norte, "geom": geom, "metodo": metodo})
    if poligonos:
        sobrantes = set(poligonos) - {b["codigo"] for b in catalogo}
        for k in sorted(sobrantes):
            avisos.append(f"Polígono '{k}' no pertenece al catálogo V6 (¿bloque retirado?); se ignora")
    return bloques, avisos


# ---------------------------------------------------------------- DEM y pendiente
def descargar_dem(carpeta, bloques, margen_m=8000):
    """Descarga las teselas Copernicus GLO-30 (1°x1°) que cubren el ámbito."""
    import urllib.request
    from pyproj import Transformer

    carpeta.mkdir(parents=True, exist_ok=True)
    t = Transformer.from_crs(CRS_PROYECTO, "EPSG:4326", always_xy=True)
    xs = [b["este"] for b in bloques]
    ys = [b["norte"] for b in bloques]
    lon0, lat0 = t.transform(min(xs) - margen_m, min(ys) - margen_m)
    lon1, lat1 = t.transform(max(xs) + margen_m, max(ys) + margen_m)
    for lat in range(math.floor(lat0), math.floor(lat1) + 1):
        for lon in range(math.floor(lon0), math.floor(lon1) + 1):
            ns = f"{'N' if lat >= 0 else 'S'}{abs(lat):02d}"
            ew = f"{'E' if lon >= 0 else 'W'}{abs(lon):03d}"
            n = f"Copernicus_DSM_COG_10_{ns}_00_{ew}_00_DEM"
            destino = carpeta / f"{n}.tif"
            if not destino.exists():
                print(f"Descargando {n} ...")
                urllib.request.urlretrieve(URL_DEM.format(n=n), destino)


def dem_utm(carpeta, bloques, margen_m=6000):
    """Mosaico del DEM reproyectado a EPSG:32717 a 30 m sobre el ámbito de los bloques."""
    import rasterio
    from rasterio.merge import merge
    from rasterio.transform import from_origin
    from rasterio.warp import Resampling, reproject

    teselas = sorted(Path(carpeta).glob("*.tif"))
    if not teselas:
        raise SystemExit(f"No hay GeoTIFF en {carpeta}. Use --descargar-dem.")
    fuentes = [rasterio.open(p) for p in teselas]
    mosaico, tr_src = merge(fuentes)
    crs_src = fuentes[0].crs
    for f in fuentes:
        f.close()

    xmin = math.floor((min(b["geom"].bounds[0] for b in bloques) - margen_m) / RES) * RES
    ymin = math.floor((min(b["geom"].bounds[1] for b in bloques) - margen_m) / RES) * RES
    xmax = math.ceil((max(b["geom"].bounds[2] for b in bloques) + margen_m) / RES) * RES
    ymax = math.ceil((max(b["geom"].bounds[3] for b in bloques) + margen_m) / RES) * RES
    ancho, alto = int((xmax - xmin) / RES), int((ymax - ymin) / RES)
    tr = from_origin(xmin, ymax, RES, RES)
    dem = np.full((alto, ancho), np.nan, dtype="float32")
    reproject(mosaico[0].astype("float32"), dem, src_transform=tr_src, src_crs=crs_src,
              dst_transform=tr, dst_crs=CRS_PROYECTO, resampling=Resampling.bilinear,
              src_nodata=-32767.0, dst_nodata=np.nan)
    return dem, tr


def pendiente_pct(dem, res=RES):
    """Pendiente en porcentaje por el método de Horn (ventana 3x3)."""
    z = np.pad(dem, 1, mode="edge")
    a, b, c = z[:-2, :-2], z[:-2, 1:-1], z[:-2, 2:]
    d, f = z[1:-1, :-2], z[1:-1, 2:]
    g, h, i = z[2:, :-2], z[2:, 1:-1], z[2:, 2:]
    dzdx = ((c + 2 * f + i) - (a + 2 * d + g)) / (8 * res)
    dzdy = ((g + 2 * h + i) - (a + 2 * b + c)) / (8 * res)
    return (np.hypot(dzdx, dzdy) * 100).astype("float32")


def sombreado(dem, res=RES, azimut=315, altitud=45):
    z = np.pad(dem, 1, mode="edge")
    dx = ((z[:-2, 2:] + 2 * z[1:-1, 2:] + z[2:, 2:]) - (z[:-2, :-2] + 2 * z[1:-1, :-2] + z[2:, :-2])) / (8 * res)
    dy = ((z[2:, :-2] + 2 * z[2:, 1:-1] + z[2:, 2:]) - (z[:-2, :-2] + 2 * z[:-2, 1:-1] + z[:-2, 2:])) / (8 * res)
    pend = np.arctan(np.hypot(dx, dy))
    asp = np.arctan2(-dx, dy)
    az, alt = np.radians(360 - azimut + 90), np.radians(altitud)
    hs = np.sin(alt) * np.cos(pend) + np.cos(alt) * np.sin(pend) * np.cos(az - asp)
    return np.clip(hs, 0, 1).astype("float32")


def clasificar(pend):
    cls = np.digitize(pend, LIMITES).astype("uint8") + 1  # 1..8
    cls[np.isnan(pend)] = 0
    return cls


def guardar_tif(ruta, arr, tr, nodata):
    import rasterio
    with rasterio.open(ruta, "w", driver="GTiff", height=arr.shape[0], width=arr.shape[1], count=1,
                       dtype=arr.dtype, crs=CRS_PROYECTO, transform=tr, nodata=nodata,
                       compress="deflate", tiled=True) as dst:
        dst.write(arr, 1)


# ---------------------------------------------------------------- estadísticas
def estadisticas(bloque, pend, tr):
    from rasterio.features import geometry_mask
    from rasterio.windows import from_bounds

    w = from_bounds(*bloque["geom"].bounds, transform=tr).round_offsets().round_lengths()
    w = w.intersection(type(w)(0, 0, pend.shape[1], pend.shape[0]))
    sub = pend[w.row_off:w.row_off + w.height, w.col_off:w.col_off + w.width]
    tr_w = tr * tr.translation(w.col_off, w.row_off)
    dentro = geometry_mask([bloque["geom"]], out_shape=sub.shape, transform=tr_w,
                           invert=True, all_touched=False)
    v = sub[dentro & ~np.isnan(sub)]
    if v.size == 0:
        return None
    cls = np.digitize(v, LIMITES)
    pct = [float((cls == k).mean() * 100) for k in range(len(CLASES))]
    return {"media": float(v.mean()), "mediana": float(np.median(v)), "p90": float(np.percentile(v, 90)),
            "max": float(v.max()), "pct_clases": pct, "pct_75": pct[-1],
            "celdas": int(v.size)}


# ---------------------------------------------------------------- dibujo
def _cmap():
    from matplotlib.colors import BoundaryNorm, ListedColormap
    cmap = ListedColormap([c[4] for c in CLASES])
    cmap.set_bad((1, 1, 1, 0))
    return cmap, BoundaryNorm(np.arange(0.5, len(CLASES) + 1.5), cmap.N)


def _rgb_clases(cls, hs):
    """Clases de pendiente con sombreado multiplicativo suave."""
    from matplotlib.colors import to_rgb
    paleta = np.array([(1, 1, 1)] + [to_rgb(c[4]) for c in CLASES])
    rgb = paleta[cls]
    rgb = rgb * (0.72 + 0.28 * hs[..., None])
    alfa = (cls > 0).astype(float)[..., None]
    return np.concatenate([np.clip(rgb, 0, 1), alfa], axis=2)


def _ejes_utm(ax, paso):
    from matplotlib.ticker import FuncFormatter, MultipleLocator
    ax.xaxis.set_major_locator(MultipleLocator(paso))
    ax.yaxis.set_major_locator(MultipleLocator(paso))
    fmt = FuncFormatter(lambda v, _: f"{v:,.0f}")
    ax.xaxis.set_major_formatter(fmt)
    ax.yaxis.set_major_formatter(fmt)
    ax.tick_params(labelsize=6.5, direction="in", top=True, right=True)
    plt_setp = ax.get_yticklabels()
    for t in plt_setp:
        t.set_rotation(90)
        t.set_va("center")
    ax.grid(color="#555555", lw=0.3, ls=":", alpha=0.6)
    ax.set_xlabel("ESTE (m)", fontsize=7)
    ax.set_ylabel("NORTE (m)", fontsize=7)


def _norte_escala(ax, ancho_m):
    x0, x1 = ax.get_xlim()
    y0, y1 = ax.get_ylim()
    # Norte
    ax.annotate("N", xy=(0.94, 0.95), xytext=(0.94, 0.85), xycoords="axes fraction",
                ha="center", va="center", fontsize=10, fontweight="bold",
                arrowprops=dict(facecolor="black", width=3, headwidth=9, headlength=8),
                bbox=dict(boxstyle="circle,pad=0.15", fc="white", ec="none", alpha=0.8))
    # Escala gráfica
    objetivo = ancho_m / 4
    base = 10 ** math.floor(math.log10(objetivo))
    largo = max(m * base for m in (1, 2, 5) if m * base <= objetivo)
    xs, ys = x0 + (x1 - x0) * 0.05, y0 + (y1 - y0) * 0.05
    ax.add_patch(__import__("matplotlib.patches", fromlist=["Rectangle"]).Rectangle(
        (xs - (x1 - x0) * 0.015, ys - (y1 - y0) * 0.02), largo + (x1 - x0) * 0.06, (y1 - y0) * 0.075,
        fc="white", ec="none", alpha=0.85, zorder=5))
    for k in range(4):
        ax.add_patch(__import__("matplotlib.patches", fromlist=["Rectangle"]).Rectangle(
            (xs + k * largo / 4, ys), largo / 4, (y1 - y0) * 0.012,
            fc="black" if k % 2 == 0 else "white", ec="black", lw=0.6, zorder=6))
    etiqueta = f"{largo / 1000:g} km" if largo >= 1000 else f"{largo:g} m"
    ax.text(xs, ys + (y1 - y0) * 0.02, "0", fontsize=6, ha="center", zorder=6)
    ax.text(xs + largo, ys + (y1 - y0) * 0.02, etiqueta, fontsize=6, ha="center", zorder=6)


def _pagina(titulo, subtitulo):
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    disponibles = {f.name for f in font_manager.fontManager.ttflist}
    plt.rcParams["font.family"] = next(f for f in ("Arial", "Liberation Sans", "DejaVu Sans")
                                       if f in disponibles or f == "DejaVu Sans")
    fig = plt.figure(figsize=(11.69, 8.27))  # A4 horizontal
    banda = fig.add_axes([0, 0.905, 1, 0.095])
    banda.set_facecolor(VERDE)
    banda.set_xticks([]), banda.set_yticks([])
    for s in banda.spines.values():
        s.set_visible(False)
    for k, t in enumerate(ENCABEZADOS):
        banda.text(0.012, 0.78 - k * 0.27, t, color="white", fontsize=8.5 if k == 0 else 7.5,
                   fontweight="bold" if k == 0 else "normal", va="center", transform=banda.transAxes)
    banda.text(0.988, 0.62, titulo, color="white", fontsize=13, fontweight="bold", ha="right",
               va="center", transform=banda.transAxes)
    banda.text(0.988, 0.25, subtitulo, color="#E8F0EA", fontsize=8, ha="right", va="center",
               transform=banda.transAxes)
    fig.text(0.012, 0.012, f"{SUBTITULO}   |   {FUENTE_DEM}   |   Elaboración: ANIN-DIME-SESDI, "
             f"{date.today():%d/%m/%Y}", fontsize=5.8, color="#333333")
    return fig


def _leyenda(ax_l, stats=None):
    """Leyenda de clases; con stats agrega la columna '% del área' alineada a la derecha."""
    from matplotlib.patches import Rectangle
    ax_l.set_axis_off()
    ax_l.text(0, 1, "PENDIENTE (%)", fontsize=8.5, fontweight="bold", color=VERDE, va="top")
    if stats:
        ax_l.text(1, 1, "% del área", fontsize=7, fontweight="bold", color=VERDE, va="top", ha="right")
    paso = 0.86 / len(CLASES)
    for k, (cod, lo, hi, desc, col) in enumerate(CLASES):
        y = 0.86 - k * paso
        ax_l.add_patch(Rectangle((0, y - paso * 0.4), 0.08, paso * 0.75, fc=col, ec="#444444", lw=0.4,
                                 transform=ax_l.transAxes))
        rango = f"> {lo}" if hi is None else f"{lo} – {hi}"
        ax_l.text(0.11, y, cod, fontsize=7, fontweight="bold", va="center", transform=ax_l.transAxes)
        ax_l.text(0.30, y, rango, fontsize=7, va="center", ha="right", transform=ax_l.transAxes)
        ax_l.text(0.34, y, desc, fontsize=7, va="center", transform=ax_l.transAxes)
        if stats:
            ax_l.text(1, y, f"{stats['pct_clases'][k]:.1f}", fontsize=7, va="center", ha="right",
                      color="#A50026" if cod == "H" and stats["pct_clases"][k] > 0 else "black",
                      transform=ax_l.transAxes)


def _barras(ax_b, stats):
    ax_b.barh(range(len(CLASES)), stats["pct_clases"], color=[c[4] for c in CLASES],
              edgecolor="#444444", lw=0.4)
    ax_b.set_yticks(range(len(CLASES)), [c[0] for c in CLASES], fontsize=6.5)
    ax_b.invert_yaxis()
    ax_b.tick_params(axis="x", labelsize=6.5)
    ax_b.set_xlabel("% del área de cálculo", fontsize=6.5)
    ax_b.set_title("Distribución por clase de pendiente", fontsize=7.5, color=VERDE,
                   fontweight="bold", loc="left")
    for s in ("top", "right"):
        ax_b.spines[s].set_visible(False)
    tope = max(stats["pct_clases"]) or 1
    ax_b.set_xlim(0, tope * 1.18)
    for k, v in enumerate(stats["pct_clases"]):
        if v >= 0.05:
            ax_b.text(v + tope * 0.015, k, f"{v:.1f}", va="center", fontsize=6)


def mapa_bloque(b, st, cls, hs, tr, ruta_png=None, pdf=None):
    import matplotlib.pyplot as plt

    gx0, gy0, gx1, gy1 = b["geom"].bounds
    cx, cy = (gx0 + gx1) / 2, (gy0 + gy1) / 2
    semi = max((gx1 - gx0), (gy1 - gy0)) / 2 * 1.6 + 300
    x0, x1, y0, y1 = cx - semi, cx + semi, cy - semi, cy + semi
    c0 = max(int((x0 - tr.c) / RES), 0)
    c1 = min(int(math.ceil((x1 - tr.c) / RES)), cls.shape[1])
    r0 = max(int((tr.f - y1) / RES), 0)
    r1 = min(int(math.ceil((tr.f - y0) / RES)), cls.shape[0])
    ext = (tr.c + c0 * RES, tr.c + c1 * RES, tr.f - r1 * RES, tr.f - r0 * RES)

    fig = _pagina(f"MAPA DE PENDIENTES – BLOQUE {b['codigo']}",
                  f"Prov. {b['provincia']} · Dist. {b['distrito']} · Microcuenca {b['microcuenca']} · "
                  f"{b['area_ha']:,.3f} ha")
    ax = fig.add_axes([0.06, 0.07, 0.56, 0.80])
    ax.imshow(_rgb_clases(cls[r0:r1, c0:c1], hs[r0:r1, c0:c1]), extent=ext, interpolation="nearest")
    es_circulo = b["metodo"].startswith("Círculo")
    geoms = [b["geom"]] if b["geom"].geom_type == "Polygon" else list(b["geom"].geoms)
    for gm in geoms:
        gx, gy = gm.exterior.xy
        ax.plot(gx, gy, color="black", lw=1.4, ls="--" if es_circulo else "-", zorder=4)
        ax.plot(gx, gy, color="white", lw=0.5, ls="--" if es_circulo else "-", zorder=4)
    ax.plot(b["este"], b["norte"], marker="+", color="black", ms=11, mew=1.8, zorder=5)
    ax.set_xlim(x0, x1)
    ax.set_ylim(y0, y1)
    ax.set_aspect("equal")
    paso = 500 if semi < 1500 else 1000 if semi < 3500 else 2000
    _ejes_utm(ax, paso)
    _norte_escala(ax, 2 * semi)

    # Panel derecho: leyenda, barras, ficha y nota
    ax_l = fig.add_axes([0.665, 0.69, 0.31, 0.18])
    _leyenda(ax_l, st)
    if st:
        _barras(fig.add_axes([0.69, 0.475, 0.28, 0.17]), st)
    ax_s = fig.add_axes([0.665, 0.07, 0.31, 0.36])
    ax_s.set_axis_off()
    filas = [
        ("Código de bloque", b["codigo"]),
        ("Provincia / Distrito", f"{b['provincia']} / {b['distrito']}"),
        ("Microcuenca", b["microcuenca"]),
        ("Área catálogo V6", f"{b['area_ha']:,.3f} ha"),
        ("Centroide (E, N)", f"{b['este']:,.0f} , {b['norte']:,.0f}"),
        ("Pendiente media", f"{st['media']:.1f} %" if st else "s/d"),
        ("Mediana / P90", f"{st['mediana']:.1f} % / {st['p90']:.1f} %" if st else "s/d"),
        ("Pendiente máxima", f"{st['max']:.1f} %" if st else "s/d"),
        ("Superficie > 75 %", f"{st['pct_75']:.1f} %  (≈ {st['pct_75'] / 100 * b['area_ha']:,.2f} ha)"
         if st else "s/d"),
        ("Ámbito de cálculo", b["metodo"]),
    ]
    tabla = ax_s.table(cellText=filas, colWidths=[0.42, 0.58], loc="upper left", cellLoc="left")
    tabla.auto_set_font_size(False)
    tabla.set_fontsize(7)
    tabla.scale(1, 1.25)
    for (r, c), cel in tabla.get_celld().items():
        cel.set_edgecolor("#999999")
        cel.set_linewidth(0.4)
        if c == 0:
            cel.set_facecolor("#E8F0EA")
            cel.get_text().set_fontweight("bold")
        if r == 8 and st and st["pct_75"] > 0:
            cel.get_text().set_color("#A50026")
    nota = ("Línea discontinua: círculo de área equivalente centrado en el centroide del bloque "
            "(referencial; NO es el límite del bloque). Estadísticas aproximadas hasta contar con "
            "el polígono oficial." if es_circulo else
            "Línea continua: límite del bloque (capa de polígonos).")
    nota += (" Umbral de idoneidad (Paso 4): pendiente máxima 75 %. Clases: rangos del Reglamento "
             "de Clasificación de Tierras por su Capacidad de Uso Mayor.")
    ax_s.text(0, 0.0, textwrap.fill(nota, 92), fontsize=6, va="bottom", color="#333333",
              transform=ax_s.transAxes)
    if ruta_png:
        fig.savefig(ruta_png, dpi=150)
    if pdf:
        pdf.savefig(fig)
    plt.close(fig)


def mapa_general(bloques, resultados, cls, hs, tr, ruta_png=None, pdf=None):
    import matplotlib.pyplot as plt

    f = 3  # submuestreo para el mapa general (90 m)
    ext = (tr.c, tr.c + cls.shape[1] * RES, tr.f - cls.shape[0] * RES, tr.f)
    fig = _pagina("MAPA DE PENDIENTES – 117 BLOQUES V6",
                  f"{len(bloques)} bloques · {sum(b['area_ha'] for b in bloques):,.3f} ha · "
                  "Ayabaca, Huancabamba y Morropón")
    ax = fig.add_axes([0.06, 0.07, 0.56, 0.80])
    ax.imshow(_rgb_clases(cls[::f, ::f], hs[::f, ::f]), extent=ext, interpolation="nearest")
    for b in bloques:
        st = resultados.get(b["codigo"])
        geoms = [b["geom"]] if b["geom"].geom_type == "Polygon" else list(b["geom"].geoms)
        for gm in geoms:
            gx, gy = gm.exterior.xy
            ax.fill(gx, gy, fc="none", ec="black", lw=0.6, zorder=4)
        rojo = st and st["pct_75"] >= 5
        ax.plot(b["este"], b["norte"], "o", ms=2.2, color="#A50026" if rojo else "black", zorder=5)
        ax.text(b["este"] + 450, b["norte"] + 250, b["codigo"], fontsize=3.6, zorder=6,
                bbox=dict(boxstyle="round,pad=0.08", fc="white", ec="none", alpha=0.7))
    xs = [b["este"] for b in bloques]
    ys = [b["norte"] for b in bloques]
    pad = 4000
    cx, cy = (min(xs) + max(xs)) / 2, (min(ys) + max(ys)) / 2
    semi = max(max(xs) - min(xs), max(ys) - min(ys)) / 2 + pad
    ax.set_xlim(max(cx - semi, ext[0]), min(cx + semi, ext[1]))
    ax.set_ylim(max(cy - semi, ext[2]), min(cy + semi, ext[3]))
    ax.set_aspect("equal")
    _ejes_utm(ax, 10000)
    _norte_escala(ax, 2 * semi)

    ax_l = fig.add_axes([0.665, 0.69, 0.31, 0.18])
    _leyenda(ax_l)
    ax_r = fig.add_axes([0.665, 0.07, 0.31, 0.55])
    ax_r.set_axis_off()
    vals = [r for r in resultados.values() if r]
    areas = {b["codigo"]: b["area_ha"] for b in bloques}
    a_tot = sum(areas[k] for k, r in resultados.items() if r)
    media_pond = sum(r["media"] * areas[k] for k, r in resultados.items() if r) / a_tot
    a75 = sum(r["pct_75"] / 100 * areas[k] for k, r in resultados.items() if r)
    n5 = sum(1 for r in vals if r["pct_75"] >= 5)
    filas = [
        ("Bloques cartografiados", f"{len(vals)} de {len(bloques)}"),
        ("Área total (catálogo V6)", f"{sum(areas.values()):,.3f} ha"),
        ("Pendiente media ponderada", f"{media_pond:.1f} %"),
        ("Superficie > 75 % (aprox.)", f"≈ {a75:,.1f} ha ({a75 / a_tot * 100:.1f} %)"),
        ("Bloques con ≥ 5 % del área > 75 %", f"{n5} (punto rojo)"),
    ]
    tabla = ax_r.table(cellText=filas, colWidths=[0.55, 0.45], loc="upper left", cellLoc="left")
    tabla.auto_set_font_size(False)
    tabla.set_fontsize(7)
    tabla.scale(1, 1.4)
    for (r, c), cel in tabla.get_celld().items():
        cel.set_edgecolor("#999999")
        cel.set_linewidth(0.4)
        if c == 0:
            cel.set_facecolor("#E8F0EA")
            cel.get_text().set_fontweight("bold")
    metodo = {b["metodo"] for b in bloques}
    nota = ("Contorno: " + " / ".join(sorted(metodo)) + ". Los círculos de área equivalente son "
            "referenciales y no representan el límite real de los bloques." if any(
                m.startswith("Círculo") for m in metodo) else "Contorno: polígonos de los bloques.")
    ax_r.text(0, 0.05, textwrap.fill(nota, 92), fontsize=6, va="bottom", color="#333333",
              transform=ax_r.transAxes)
    if ruta_png:
        fig.savefig(ruta_png, dpi=220)
    if pdf:
        pdf.savefig(fig)
    plt.close(fig)


# ---------------------------------------------------------------- Excel
def matriz_excel(ruta, bloques, resultados, avisos):
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from anin_excel import encabezado_institucional, escribir_tabla

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pendientes por bloque"
    cols = (["N°", "Código", "Provincia", "Distrito", "Microcuenca", "Área (ha)", "Este (m)",
             "Norte (m)", "Pend. media (%)", "Mediana (%)", "P90 (%)", "Máx. (%)"]
            + [f"{c[0]} {c[1]}–{c[2]} %" if c[2] else f"{c[0]} >{c[1]} %" for c in CLASES]
            + ["Área > 75 % (ha)", "Ámbito de cálculo"])
    fila = encabezado_institucional(ws, "MATRIZ DE PENDIENTES POR BLOQUE – CATÁLOGO V6", len(cols),
                                    fuente=FUENTE_DEM)
    filas = []
    ini = fila + 1
    i_area = cols.index("Área (ha)") + 1
    i_h = cols.index(next(c for c in cols if c.startswith("H "))) + 1
    for n, b in enumerate(bloques, start=1):
        st = resultados.get(b["codigo"])
        r = ini + n - 1
        base = [n, b["codigo"], b["provincia"], b["distrito"], b["microcuenca"], b["area_ha"],
                round(b["este"]), round(b["norte"])]
        if st:
            base += [round(st["media"], 2), round(st["mediana"], 2), round(st["p90"], 2), round(st["max"], 2)]
            base += [round(p, 2) for p in st["pct_clases"]]
        else:
            base += [None] * (4 + len(CLASES))
        base += [f"={get_column_letter(i_area)}{r}*{get_column_letter(i_h)}{r}/100", b["metodo"]]
        filas.append(base)
    fmt = {c: "0.00" for c in cols[8:-2]}
    fmt.update({"Área (ha)": "#,##0.000", "Este (m)": "#,##0", "Norte (m)": "#,##0",
                "Área > 75 % (ha)": "#,##0.00"})
    anchos = {"N°": 6, "Código": 10, "Provincia": 14, "Distrito": 24, "Microcuenca": 15,
              "Ámbito de cálculo": 28}
    sig = escribir_tabla(ws, fila, cols, filas, anchos=anchos, formato_num=fmt,
                         totales={"Área (ha)": "SUM", "Área > 75 % (ha)": "SUM"})
    ult = fila + len(filas)
    tot = ult + 1
    a = get_column_letter(i_area)
    m = get_column_letter(cols.index("Pend. media (%)") + 1)
    c = ws.cell(tot, cols.index("Pend. media (%)") + 1,
                f"=SUMPRODUCT({a}{fila + 1}:{a}{ult},{m}{fila + 1}:{m}{ult})/{a}{tot}")
    c.number_format = "0.00"
    c.font = Font(name="Arial", size=10, bold=True)
    ws.cell(sig, 1, "Pend. media de la fila TOTAL: promedio ponderado por área (SUMPRODUCT).").font = \
        Font(name="Arial", size=9, italic=True)

    # Resumen por distrito
    wr = wb.create_sheet("Resumen por distrito")
    cols_r = ["Provincia", "Distrito", "N° bloques", "Área (ha)", "Área > 75 % (ha)", "% > 75 %"]
    f2 = encabezado_institucional(wr, "RESUMEN DE PENDIENTES POR DISTRITO", len(cols_r), fuente=FUENTE_DEM)
    claves = sorted({(b["provincia"], b["distrito"]) for b in bloques})
    hoja = "'Pendientes por bloque'"
    rng = lambda col: f"{hoja}!${get_column_letter(cols.index(col) + 1)}${fila + 1}:" \
                      f"${get_column_letter(cols.index(col) + 1)}${ult}"
    filas_r = []
    for k, (p, d) in enumerate(claves):
        r = f2 + 1 + k
        filas_r.append([p, d, f'=COUNTIF({rng("Distrito")},B{r})',
                        f'=SUMIF({rng("Distrito")},B{r},{rng("Área (ha)")})',
                        f'=SUMIF({rng("Distrito")},B{r},{rng("Área > 75 % (ha)")})',
                        f"=IF(D{r}=0,0,E{r}/D{r}*100)"])
    escribir_tabla(wr, f2, cols_r, filas_r, anchos={"Distrito": 26, "Provincia": 16},
                   formato_num={"Área (ha)": "#,##0.000", "Área > 75 % (ha)": "#,##0.00", "% > 75 %": "0.00"},
                   totales={"N° bloques": "SUM", "Área (ha)": "SUM", "Área > 75 % (ha)": "SUM"})

    # Notas
    wn = wb.create_sheet("Notas")
    encabezado_institucional(wn, "NOTAS METODOLÓGICAS Y AVISOS", 2)
    notas = [
        "DEM: Copernicus GLO-30 (ESA), resolución 30 m, reproyectado a UTM 17S WGS84 (EPSG:32717) "
        "por interpolación bilineal. Es un modelo de SUPERFICIE: en áreas con dosel alto la pendiente "
        "puede diferir de la del terreno.",
        "Pendiente: método de Horn (ventana 3x3), expresada en porcentaje.",
        "Clases A–H: rangos de pendiente del Reglamento de Clasificación de Tierras por su Capacidad de "
        "Uso Mayor (verificar la versión vigente antes de citar). El límite de 75 % coincide con el "
        "criterio de idoneidad del Paso 4 de selección de bloques.",
        "Ámbito de cálculo: el repositorio no contiene polígonos de los bloques. Cuando la columna "
        "'Ámbito de cálculo' dice 'Círculo de área equivalente', las estadísticas se calcularon sobre "
        "un círculo centrado en el centroide del bloque con la misma área del catálogo V6. Son "
        "REFERENCIALES. [POR DEFINIR]: recalcular con la capa oficial de polígonos "
        "(python mapas_pendientes_v6.py --bloques bloques_v6.gpkg).",
        "Área > 75 % (ha) = Área del catálogo × % de la clase H / 100 (estimación).",
        f"Control: {len(bloques)} bloques vigentes del catálogo V6; área total esperada "
        f"{AREA_TOTAL_HA:,.3f} ha.",
    ] + [f"Aviso: {a}" for a in avisos]
    for k, t in enumerate(notas):
        cel = wn.cell(8 + k, 1, k + 1)
        cel.font = Font(name="Arial", size=10, bold=True)
        cel = wn.cell(8 + k, 2, t)
        cel.font = Font(name="Arial", size=10)
        cel.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
        wn.row_dimensions[8 + k].height = 15 * (1 + len(t) // 95)
    wn.column_dimensions["A"].width = 6
    wn.column_dimensions["B"].width = 110
    for hoja_ws in (ws, wr, wn):
        hoja_ws.page_setup.orientation = "landscape"
        hoja_ws.page_setup.fitToWidth = 1
    wb.save(ruta)


# ---------------------------------------------------------------- principal
def main(argv=None):
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.backends.backend_pdf import PdfPages

    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dem-dir", required=True, help="Carpeta con las teselas DEM (GeoTIFF)")
    ap.add_argument("--descargar-dem", action="store_true", help="Descarga las teselas Copernicus que falten")
    ap.add_argument("--bloques", help="Capa de polígonos de bloques (GPKG/SHP con CRS)")
    ap.add_argument("--campo-codigo", default="codigo", help="Campo con el código del bloque")
    ap.add_argument("--centroides", help="CSV de centroides (por defecto datos/atributos_bloques_v6.csv)")
    ap.add_argument("--salida", default=str(RAIZ / "salidas" / "mapas_pendientes_v6"))
    ap.add_argument("--solo", nargs="*", help="Procesar solo estos códigos (prueba)")
    a = ap.parse_args(argv)

    salida = Path(a.salida)
    (salida / "mapas_bloques").mkdir(parents=True, exist_ok=True)
    bloques, avisos = preparar_bloques(a.centroides, a.bloques, a.campo_codigo)
    total = sum(b["area_ha"] for b in bloques)
    print(f"{len(bloques)} bloques · {total:,.3f} ha (esperado {AREA_TOTAL_HA:,.3f})")
    if a.solo:
        bloques = [b for b in bloques if b["codigo"] in set(a.solo)]

    if a.descargar_dem:
        descargar_dem(Path(a.dem_dir), bloques)
    dem, tr = dem_utm(a.dem_dir, bloques)
    pend = pendiente_pct(dem)
    hs = sombreado(dem)
    cls = clasificar(pend)
    guardar_tif(salida / "pendiente_pct_utm17s.tif", np.nan_to_num(pend, nan=-9999).astype("float32"), tr, -9999)
    guardar_tif(salida / "clases_pendiente_utm17s.tif", cls, tr, 0)

    resultados = {}
    for b in bloques:
        st = estadisticas(b, pend, tr)
        if st is None:
            avisos.append(f"Bloque {b['codigo']}: sin celdas DEM válidas")
        resultados[b["codigo"]] = st

    with PdfPages(salida / "Atlas_Pendientes_Bloques_V6_IN_Piura.pdf") as pdf:
        mapa_general(bloques, resultados, cls, hs, tr, salida / "Mapa_General_Pendientes_V6.png", pdf)
        for k, b in enumerate(bloques, 1):
            nombre = f"Mapa_Pendientes_Bloque_{b['codigo']}.png"
            mapa_bloque(b, resultados[b["codigo"]], cls, hs, tr, salida / "mapas_bloques" / nombre, pdf)
            if k % 20 == 0:
                print(f"  {k}/{len(bloques)} mapas")
        info = pdf.infodict()
        info["Title"] = "Atlas de pendientes – Bloques V6 – Proyecto IN Piura"
        info["Author"] = "ANIN - DIME - SESDI"

    matriz_excel(salida / "Matriz_Pendientes_Bloques_V6_IN_Piura.xlsx", bloques, resultados, avisos)
    for av in avisos:
        print("AVISO:", av)
    print(f"Listo: {salida}")


if __name__ == "__main__":
    main()
